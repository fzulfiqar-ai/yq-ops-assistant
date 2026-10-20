"""Crosscheck the live DB metrics against the source Focus reports.

A regression guard for accuracy: each check parses the authoritative Focus export and
asserts the DB view total matches within tolerance. Run after any ingest.

  python -m scripts.verify_numbers

Reads (20-Oct-2026): where DATABASE_URL is set (the owner's PC) one read-only database session
answers every read of a run, exactly as the REST/RPC path would (_Session); without it (the API on
Render, CI) the REST/RPC path reads, as before.
"""
from __future__ import annotations

import contextvars
import glob
import logging
import math
import os
import re
import sys
from datetime import date, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from dotenv import load_dotenv  # noqa: E402
load_dotenv()

import warnings  # noqa: E402
warnings.filterwarnings("ignore")

from app.database import get_client  # noqa: E402
from scripts.ingest import (  # noqa: E402
    parse_order_lines, parse_orders, parse_pricebook, parse_receivables, parse_receivables_totals, parse_stock,
    parse_stock_balance, read_grid,
)
from scripts.load_supabase import PARTIAL_OK_MARKER, marker_present, pricebook_key, warehouse_set_guard  # noqa: E402

DATA_DIR = ROOT / "business_data"
CLEAN_DIR = ROOT / "data" / "clean"
log = logging.getLogger(__name__)

# Per-day and per-salesman sales are compared to the fils (R2, 24-Sep-2026). The file side is the
# SUM OF THE LINES, never Focus's Grand Total row (parse_* skip total rows): the 240926 preview
# documented that the report's Grand Total is rounded independently of the lines, so it can sit a
# few fils off the lines it totals. The lines are what loads, so the lines are what is checked.
ABS_TOL_BHD = 0.005

# Report types that have NO cross-check rule here. They still load (data/clean/<table>.csv), so a
# refresh made only of these must not be reported as a FAIL -- it is "loaded, nothing to compare".
# (Price books left this list on 24-Sep-2026: see _price_book_checks.)
_UNCHECKED = {
    "product_profitability": "Product_Profitability_Report",
    "ledger_entries": "Ledger",
}

# filename key (lower-cased substring, as classify() sees it) -> price_book
_PRICE_BOOKS = (("masellingpricebook", "MA_base"), ("moderntradesellerbook", "modern_trade"))


def _latest_source_dir() -> Path:
    """The most recent dated Focus drop, e.g. business_data/Focus ERP 2026-09-14.

    The default used to be hard-pinned to "Focus ERP Updated Reports" (a June export). Running
    this with no argument therefore crosschecked TODAY's database against MONTHS-old reports and
    reported FAIL on receivables and stock -- a scary red result caused purely by a stale default.
    Prefer the newest 'Focus ERP <date>' folder, which is what an upload actually loaded.
    """
    dated = sorted(
        (p for p in DATA_DIR.glob("Focus ERP 20*") if p.is_dir()),
        key=lambda p: p.name,
    )
    if dated:
        return dated[-1]
    return DATA_DIR / "Focus ERP Updated Reports"


NEW = _latest_source_dir()


def _nn(v) -> float:
    try:
        f = float(v) if v not in (None, "") else 0.0
        return 0.0 if math.isnan(f) else f
    except Exception:
        return 0.0


def _find(key: str, src_dir: Path) -> str | None:
    for p in sorted(glob.glob(str(src_dir / "*.xls*"))):
        if key in os.path.basename(p).lower():
            return p
    return None


def _find_newest_book(key: str, src_dir: Path) -> str | None:
    """The price book to check when a folder holds several exports of one book: the NEWEST by
    export_date (docProps.created / mtime), never the one that sorts first alphabetically (the
    older download counter). ingest de-duplicates by mtime, so this is the file that was loaded."""
    paths = [p for p in glob.glob(str(src_dir / "*.xls*")) if key in os.path.basename(p).lower()]
    if not paths:
        return None
    return max(sorted(paths), key=lambda p: (export_date(p), os.path.getmtime(p)))


# ── One read-only database session per run (load speed, release 2.4, 20-Oct-2026) ─────────────────
# Where DATABASE_URL is set (the owner's PC) one read-only psycopg transaction answers the run's reads
# instead of about 100 PostgREST requests made one after another. Each read is answered as the REST/RPC
# path answers it, so run_checks returns the same rows to the last digit:
#   * an RPC read runs the same SQL inside the RPC's own wrapper (_RPC_WRAPPER) on relations the probe
#     shows yq_readonly reading in full, and the same JSON is parsed the same way;
#   * where yq_readonly may not read (orders and stock_movements on production) the RPC fails and the REST
#     pages follow: the session refuses alike, then reads the same rows in the same unique order as the
#     same JSON and adds them with the same code -- or, when every value is a whole number and their
#     absolute sum is below 2**53, takes the database's exact sum: floating point adds such numbers
#     exactly in any order, so it is the very float the page loop reaches;
#   * a read the probe cannot vouch for, and every read after a session error, goes the REST/RPC way
#     (each read is identical on its own, so the mix is too).
# The API on Render never opens one (_open_session): its verify reads over the REST/RPC path, as before.
_SESSION: contextvars.ContextVar = contextvars.ContextVar("verify_numbers_session", default=None)
_RELATIONS = ("v_sales", "orders", "stock_movements", "stock_balance", "v_receivables", "ar_ageing_totals",
              "v_price_list_by_book", "selling_prices", "v_sales_by_channel")
# run_readonly_query and run_readonly_query_params (scripts/security_migration.sql) both run the SQL so
_RPC_WRAPPER = "SELECT COALESCE(json_agg(t), '[]'::json) FROM (%s) t"
_RPC_SOURCE = "'" + _RPC_WRAPPER.replace("'", "''") + "'"           # ... as their source text spells it
_SQL_RELATION = re.compile(r"\b(?:from|join)\s+([a-z_][a-z0-9_]*)", re.I)
_EXACT_BELOW = 2 ** 53
# asked of the session and of the RPC: the same database (server start time) and the same settings
_FINGERPRINT_SQL = ("SELECT current_database() || ' ' || extract(epoch FROM pg_postmaster_start_time())::text AS s, "
                    "current_setting('TimeZone') AS tz, current_setting('DateStyle') AS ds, "
                    "current_setting('extra_float_digits') AS efd")
# Per relation: the RPC (owned by yq_readonly) reads every row / is refused; PostgREST (service_role)
# reads every row; this session reads every row. A view reads with its owner's rights; a table under row
# security only through a permissive policy USING (true) and no restrictive one.
_PROBE_SQL = """
WITH want(rel) AS (SELECT unnest(%s::text[])),
     rpc AS (SELECT count(*) = 2 AS ok FROM pg_proc p
              WHERE p.pronamespace = 'public'::regnamespace AND p.prosecdef AND p.proowner = 'yq_readonly'::regrole
                AND p.oid::regprocedure::text IN ('run_readonly_query(text)', 'run_readonly_query_params(text,jsonb)')
                AND strpos(p.prosrc, %s) > 0 AND has_function_privilege('service_role', p.oid, 'EXECUTE'))
SELECT w.rel, c.oid IS NOT NULL AS present, rpc.ok AS rpc_ok,
       COALESCE(has_schema_privilege('yq_readonly', 'public', 'USAGE')
                AND has_table_privilege('yq_readonly', c.oid, 'SELECT')
                AND CASE WHEN c.relkind = 'v'
                         THEN NOT EXISTS (SELECT 1 FROM unnest(c.reloptions) o WHERE split_part(o, '=', 1) = 'security_invoker')
                         WHEN c.relkind IN ('r', 'p')
                         THEN NOT c.relrowsecurity
                              OR (EXISTS (SELECT 1 FROM pg_policy p WHERE p.polrelid = c.oid AND p.polpermissive
                                            AND p.polcmd IN ('r', '*') AND p.polroles && ARRAY[0::oid, 'yq_readonly'::regrole::oid]
                                            AND pg_get_expr(p.polqual, p.polrelid) = 'true')
                                  AND NOT EXISTS (SELECT 1 FROM pg_policy p WHERE p.polrelid = c.oid AND NOT p.polpermissive))
                         ELSE false END, false) AS rpc_all,
       COALESCE(NOT (has_schema_privilege('yq_readonly', 'public', 'USAGE')
                     AND has_any_column_privilege('yq_readonly', c.oid, 'SELECT')), false) AS rpc_refused,
       COALESCE(has_schema_privilege('service_role', 'public', 'USAGE') AND has_table_privilege('service_role', c.oid, 'SELECT')
                AND (c.relkind = 'v' OR NOT c.relrowsecurity
                     OR (SELECT r.rolbypassrls FROM pg_roles r WHERE r.rolname = 'service_role')), false) AS rest_all,
       COALESCE(has_table_privilege(c.oid, 'SELECT')
                AND (c.relkind = 'v' OR NOT c.relrowsecurity
                     OR (SELECT r.rolbypassrls FROM pg_roles r WHERE r.rolname = current_user)
                     OR (c.relowner = (SELECT r.oid FROM pg_roles r WHERE r.rolname = current_user)
                         AND NOT c.relforcerowsecurity)), false) AS mine_all
FROM want w CROSS JOIN rpc
LEFT JOIN pg_class c ON c.relnamespace = 'public'::regnamespace AND c.relname = w.rel
"""


class _Fallback(Exception):
    """The session cannot answer this read as the REST/RPC path would: that path reads it."""


def _why(exc: BaseException) -> str:
    """A failure in a few words, never the database address (psycopg's connection errors carry it)."""
    if type(exc).__module__.split(".")[0] == "psycopg":
        state = getattr(exc, "sqlstate", None)
        return type(exc).__name__ + (f" (SQLSTATE {state})" if state else "")
    return f"{type(exc).__name__}: {str(exc)[:200]}"


def _shape(probe: list[dict]) -> tuple[set, set, set]:
    """From the probe: the relations the RPC reads in full, those it refuses, and those the REST pages read in
    full. A full read through the session also needs the session itself to see every row."""
    mine = {r["rel"] for r in probe if r.get("present") and r.get("mine_all")}
    rpc = {r["rel"] for r in probe if r.get("present") and r.get("rpc_ok")}
    return ({r["rel"] for r in probe if r.get("rpc_all")} & mine & rpc,
            {r["rel"] for r in probe if r.get("rpc_refused")} & rpc,
            {r["rel"] for r in probe if r.get("rest_all")} & mine)


def _pg_sql(sql: str, params: list | None) -> tuple[str, list[str]]:
    """The RPC's $1..$8 as psycopg placeholders, each bound as text the way the RPC binds them (a jsonb array of
    text); a literal % is doubled for psycopg."""
    vals = [str(p) for p in (params or [])]
    out: list[str] = []

    def one(m):
        out.append(vals[int(m.group(1)) - 1])
        return "%s::text"
    return re.sub(r"\$(\d+)", one, sql.replace("%", "%%")), out


def _rest_where(span: tuple | None) -> tuple[str, list]:
    """PostgREST's gte / lte on the date column (the value compared as that column's type)."""
    if not span:
        return "", []
    datecol, dmin, dmax = span
    return f" WHERE {datecol} >= %s AND {datecol} <= %s", [dmin, dmax]


def _add_up(rows: list[dict], col: str) -> float:
    """The REST loops' sum: the rows' values in their order, a NULL or NaN as 0."""
    return sum(_nn(r.get(col)) for r in rows)


def _add_groups(out: dict, rows: list[dict], keycol: str, valcol: str) -> None:
    """The REST loops' grouping into {key: [sum, rows]}, adding the rows in their order."""
    for r in rows:
        k = str(r.get(keycol) or "(none)")[:64]
        acc = out.setdefault(k, [0.0, 0])
        acc[0] += _nn(r.get(valcol))
        acc[1] += 1


class _Session:
    """One read-only psycopg transaction on DATABASE_URL that answers run_checks' reads exactly as the RPC and
    the REST pages would (see the block above); a method raises _Fallback where it cannot."""

    def __init__(self, conn, probe: list[dict]):
        self.conn = conn
        self.rpc_reads, self.rpc_refuses, self.rest_reads = _shape(probe)
        self.ready = False                       # set by open(): later failures are logged
        self._rows: dict = {}                    # REST rows already read, by (view, span, ordered)

    @classmethod
    def open(cls, dsn: str) -> "_Session":
        """Connect read-only, probe how the RPC and PostgREST read each relation, and check that the RPC reads
        this very database with the same settings. Raises when any of it fails; the connection is closed."""
        import psycopg
        conn = psycopg.connect(dsn, autocommit=False, connect_timeout=10, application_name="yq-verify-numbers",
                               keepalives=1, keepalives_idle=30, keepalives_interval=10, keepalives_count=3)
        s = cls(conn, [])
        try:
            conn.read_only = True                                   # before the first statement: BEGIN READ ONLY
            s._fetch("set statement_timeout = '8s'")                # each RPC and REST request has 8 s (PostgREST)
            s._fetch("set idle_in_transaction_session_timeout = '300s'")
            state = s._fetch("show transaction_read_only")[0]["transaction_read_only"]
            if str(state).lower() != "on":
                raise RuntimeError(f"the session is not read-only (transaction_read_only = {state})")
            probe = s._fetch(_PROBE_SQL, [list(_RELATIONS), _RPC_SOURCE])
            if not probe or not all(r.get("rpc_ok") for r in probe):
                raise RuntimeError("the read RPCs are not the ones verify_numbers knows")
            mine = s._fetch(_FINGERPRINT_SQL)
            theirs = _rpc(_FINGERPRINT_SQL)                         # the RPC itself: no session is installed yet
            if [dict(r) for r in mine] != [dict(r) for r in (theirs or [])]:
                raise RuntimeError("DATABASE_URL does not reach the database the RPC reads, or its settings differ")
            s.rpc_reads, s.rpc_refuses, s.rest_reads = _shape(probe)
            s.ready = True
            return s
        except BaseException:
            s.close()
            raise

    def _fetch(self, sql: str, params: list | None = None) -> list[dict]:
        """One statement of the transaction. Any failure ends the session: this read and every later one go the
        REST/RPC way."""
        if self.conn is None:
            raise _Fallback("the session is closed")
        from psycopg.rows import dict_row
        try:
            with self.conn.cursor(row_factory=dict_row) as cur:
                cur.execute(sql, params if params is not None else [])
                return [dict(r) for r in cur.fetchall()] if cur.description else []
        except Exception as exc:  # noqa: BLE001
            why = _why(exc)
            if self.ready:
                log.warning("verify_numbers: the database session failed (%s); the REST/RPC path reads the rest", why)
            self.close()
            raise _Fallback(why) from exc

    def close(self) -> None:
        """Roll the read-only transaction back and close the connection (again: nothing happens)."""
        conn, self.conn = self.conn, None
        if conn is None:
            return
        for step in (conn.rollback, conn.close):
            try:
                step()
            except Exception:  # noqa: BLE001
                pass

    def rpc(self, sql: str, params: list | None = None) -> list[dict]:
        """exec_sql / exec_sql_params: the same SQL in the RPC's own wrapper where yq_readonly reads every row, the
        RPC's refusal where it may not read, _Fallback for a relation the probe did not vouch for."""
        rels = {m.group(1).lower() for m in _SQL_RELATION.finditer(sql)}
        if rels - self.rpc_reads - self.rpc_refuses:
            raise _Fallback("a relation the probe did not vouch for")
        refused = sorted(rels & self.rpc_refuses)
        if refused:
            raise RuntimeError(f"permission denied for table {refused[0]}")      # the RPC's own answer (42501)
        q, p = _pg_sql(sql, params)
        head, tail = _RPC_WRAPPER.split("%s")
        got = self._fetch(head + q + tail, p)
        return next(iter(got[0].values())) if got else []

    def rest_sum(self, view: str, col: str, span: tuple | None = None):
        """_rest_sum: the database's exact sum when every value is a whole number and their absolute sum is below
        2**53, else the rows in the REST order added by the same code. No rows: 0, as sum() of nothing."""
        if view not in self.rest_reads:
            raise _Fallback("not read in full over REST")
        where, params = _rest_where(span)
        agg = self._fetch(f"SELECT COUNT(*) AS n, SUM({col}) AS s, "
                          f"COALESCE(bool_and({col} = trunc({col}) AND {col} <> 'NaN'), true) "
                          f"AND COALESCE(SUM(abs({col})), 0) < {_EXACT_BELOW} AS exact FROM {view}{where}", params)[0]
        if agg["exact"]:
            return _nn(agg["s"]) if agg["n"] else 0
        return _add_up(self.rest_rows(view, (col,) + ((span[0],) if span else ()), span), col)

    def rest_groups(self, view: str, keycol: str, valcol: str, span: tuple) -> dict[str, tuple[float, int]]:
        """_db_groups' REST branch: the same rows, grouped by the same code."""
        out: dict[str, list] = {}
        _add_groups(out, self.rest_rows(view, (keycol, valcol, span[0]), span), keycol, valcol)
        return {k: (v[0], v[1]) for k, v in out.items()}

    def rest_rows(self, view: str, cols: tuple, span: tuple | None = None, ordered: bool = True) -> list[dict]:
        """The rows the REST pages return, with their unique order column, as the same JSON. One read serves every
        later call on the same view and span (the invoice total and the invoice days read the same rows)."""
        if view not in self.rest_reads:
            raise _Fallback("not read in full over REST")
        key = (view, span, ordered)
        hit = self._rows.get(key)
        if hit is not None and set(cols) <= hit[0]:
            return hit[1]
        order = _stable_order(view) if ordered else None
        names = list(dict.fromkeys(((order,) if order else ()) + tuple(cols)))
        where, params = _rest_where(span)
        agg = f"json_agg(t ORDER BY t.{order})" if order else "json_agg(t)"
        got = self._fetch(f"SELECT COALESCE({agg}, '[]'::json) FROM (SELECT {', '.join(names)} FROM {view}{where}) t",
                          params)
        rows = next(iter(got[0].values())) if got else []
        self._rows[key] = (set(names), rows)
        return rows


def _open_session() -> _Session | None:
    """This run's database session, or None (the REST/RPC path reads). Only where DATABASE_URL is set, and never
    on the API host (Render): the API's verify keeps the REST/RPC path."""
    dsn = (os.environ.get("DATABASE_URL") or "").strip()
    if not dsn:
        return None
    from app.config import is_production_env
    if is_production_env(os.environ):
        return None
    try:
        return _Session.open(dsn)
    except Exception as exc:  # noqa: BLE001 -- no session: the REST/RPC path reads, as before
        log.warning("verify_numbers: no database session (%s); reading over the REST/RPC path", _why(exc))
        return None


_OLD = object()      # the session did not answer: the REST/RPC path reads


def _session_call(name: str, *args):
    """A REST read through this run's session, or _OLD when there is none or it cannot answer exactly."""
    s = _SESSION.get()
    if s is None:
        return _OLD
    try:
        return getattr(s, name)(*args)
    except _Fallback:
        return _OLD


def _rpc(sql: str, params: list | None = None) -> list[dict]:
    """The read-only RPC (exec_sql / exec_sql_params), answered by this run's session where it can answer exactly."""
    s = _SESSION.get()
    if s is not None:
        try:
            return s.rpc(sql, params)
        except _Fallback:
            pass
    from app.db_read import exec_sql, exec_sql_params
    return exec_sql_params(sql, params) if params else exec_sql(sql)


def _sql_sum(sql: str, params: list | None = None) -> float | None:
    """Single-query SUM via the read-only RPC — deterministic. The old approach paged
    PostgREST with .range() and NO ORDER BY; Postgres gives no stable order without one,
    so pages could overlap/skip and the same sum came back different between runs
    (observed: v_sales 53,446 vs 52,909 seconds apart → spurious verify FAILs)."""
    try:
        rows = _rpc(sql, params)
        return _nn((rows or [{}])[0].get("s"))
    except Exception:  # noqa: BLE001 — RPC missing/ungranted → ordered REST fallback
        return None


def _stable_order(view: str) -> str:
    """The UNIQUE column REST pages are ordered by. Paging on a value or date column is not stable
    (ties straddle page boundaries, so rows are skipped or repeated): on 24-Sep-2026 the fallback
    read production's `orders` 57.10 short on one day and `stock_movements` 250 units short --
    inside the old 0.5% tolerance, outside the fils. yq_readonly has no SELECT on orders,
    order_lines or stock_movements (stock_balance is granted, read on production 20-Oct-2026), so
    on production this fallback IS the path for them."""
    return "line_id" if view == "v_sales" else "id"


def _rest_sum(view: str, col: str, order_col: str | None = None, filters=None) -> float:
    """Paginated REST fallback with an explicit, UNIQUE ORDER BY so pages are stable."""
    c = get_client()
    rows, off = [], 0
    while True:
        q = c.table(view).select(col).order(order_col or _stable_order(view))
        for f in (filters or []):
            q = f(q)
        b = q.range(off, off + 999).execute().data or []
        rows += b
        if len(b) < 1000:
            break
        off += 1000
    return _add_up(rows, col)


def _db_sum(view: str, col: str) -> float:
    s = _sql_sum(f"SELECT COALESCE(SUM({col}),0) AS s FROM {view}")
    if s is not None:
        return s
    s = _session_call("rest_sum", view, col)
    return s if s is not _OLD else _rest_sum(view, col)


def _db_sum_between(view: str, col: str, datecol: str, dmin: str, dmax: str) -> float:
    """Sum `col` over rows whose `datecol` is in [dmin, dmax]. Scopes the check to the uploaded
    file's own date range, so an incremental (partial) day-book validates against the same span
    instead of the whole cumulative DB (which made a small top-up look like a huge 'drift')."""
    s = _sql_sum(
        f"SELECT COALESCE(SUM({col}),0) AS s FROM {view} "
        f"WHERE {datecol} >= $1::date AND {datecol} <= $2::date", [dmin, dmax])
    if s is not None:
        return s
    s = _session_call("rest_sum", view, col, (datecol, dmin, dmax))
    if s is not _OLD:
        return s
    return _rest_sum(view, col, filters=[lambda q: q.gte(datecol, dmin), lambda q: q.lte(datecol, dmax)])


def group_sums(rows: list[dict], key, val: str) -> dict[str, tuple[float, int]]:
    """Pure: {group key: (sum of `val`, row count)} over parsed report rows. `key` is a callable or
    a column name; blank keys group under '(none)'."""
    out: dict[str, list] = {}
    for r in rows:
        k = key(r) if callable(key) else r.get(key)
        k = str(k).strip() if k not in (None, "") else "(none)"
        acc = out.setdefault(k, [0.0, 0])
        acc[0] += _nn(r.get(val))
        acc[1] += 1
    return {k: (round(v[0], 6), v[1]) for k, v in out.items()}


def compare_groups(file_g: dict, db_g: dict, abs_tol: float = ABS_TOL_BHD) -> dict:
    """Pure: file vs DB per group. Returns {groups, max_diff, mismatched: [(key, file, db)],
    count_mismatched: [(key, file_n, db_n)], file_rows, db_rows}. A group on one side only counts
    as a mismatch against 0."""
    keys = sorted(set(file_g) | set(db_g))
    mism, cnt, max_diff = [], [], 0.0
    for k in keys:
        fs, fn = file_g.get(k, (0.0, 0))
        ds, dn = db_g.get(k, (0.0, 0))
        d = abs(fs - ds)
        max_diff = max(max_diff, d)
        if d > abs_tol + 1e-9:
            mism.append((k, round(fs, 3), round(ds, 3)))
        if fn != dn:
            cnt.append((k, fn, dn))
    return {"groups": len(keys), "max_diff": round(max_diff, 3), "mismatched": mism, "count_mismatched": cnt,
            "file_rows": sum(n for _, n in file_g.values()), "db_rows": sum(n for _, n in db_g.values())}


def _db_groups(sql: str, params: list, rest: tuple[str, str, str, str] | None = None) -> dict[str, tuple[float, int]]:
    """{k: (s, n)} from a GROUP BY query returning k, s, n via the read-only RPC; `rest` =
    (view, keycol, valcol, datecol) drives the REST fallback, paged in a stable (unique-column)
    order and grouped in Python -- the path production takes for the base tables."""
    try:
        rows = _rpc(sql, params) or []
        return {str(r.get("k") if r.get("k") is not None else "(none)")[:64]: (_nn(r.get("s")), int(_nn(r.get("n"))))
                for r in rows}
    except Exception:  # noqa: BLE001
        if not rest:
            raise
    view, keycol, valcol, datecol, dmin, dmax = rest[0], rest[1], rest[2], rest[3], params[0], params[1]
    got = _session_call("rest_groups", view, keycol, valcol, (datecol, dmin, dmax))
    if got is not _OLD:
        return got
    c = get_client()
    out: dict[str, list] = {}
    off = 0
    while True:
        b = (c.table(view).select(f"{keycol},{valcol}").gte(datecol, dmin).lte(datecol, dmax)
             .order(_stable_order(view)).range(off, off + 999).execute().data or [])
        _add_groups(out, b, keycol, valcol)
        if len(b) < 1000:
            break
        off += 1000
    return {k: (v[0], v[1]) for k, v in out.items()}


def _note_list(pairs: list, label: str, limit: int = 6) -> str:
    if not pairs:
        return ""
    shown = "; ".join(f"{k}: file {a} vs DB {b}" for k, a, b in pairs[:limit])
    return f"{label}: {shown}" + (f" (+{len(pairs) - limit} more)" if len(pairs) > limit else "")


def _sales_detail_checks(ol: list[dict], od: list[dict] | None, checks: list) -> None:
    """Per-day and per-salesman gross vs the DB at ABS_TOL_BHD, plus row counts, for the day book
    (v_sales.gross_bhd, the same column the file carries) and per-day invoice gross + invoice
    counts for the register (orders). Each family adds one summary check row, never hundreds."""
    dts = sorted(str(r["line_date"])[:10] for r in ol if r.get("line_date"))
    if dts:
        dmin, dmax = dts[0], dts[-1]
        f_day = group_sums(ol, lambda r: str(r.get("line_date") or "")[:10], "gross_bhd")
        f_sm = group_sums(ol, "warehouse_name", "gross_bhd")
        d_day = _db_groups("SELECT sale_date::text AS k, COALESCE(SUM(gross_bhd),0) AS s, COUNT(*) AS n FROM v_sales "
                           "WHERE sale_date >= $1::date AND sale_date <= $2::date GROUP BY 1", [dmin, dmax],
                           rest=("v_sales", "sale_date", "gross_bhd", "sale_date"))
        d_sm = _db_groups("SELECT COALESCE(salesman_raw, '(none)') AS k, COALESCE(SUM(gross_bhd),0) AS s, COUNT(*) AS n "
                          "FROM v_sales WHERE sale_date >= $1::date AND sale_date <= $2::date GROUP BY 1", [dmin, dmax],
                          rest=("v_sales", "salesman_raw", "gross_bhd", "sale_date"))
        cd = compare_groups(f_day, d_day)
        checks.append((f"Sales per day: max |file - DB| BHD ({cd['groups']} days)", 0.0, cd["max_diff"], 0.0,
                       ABS_TOL_BHD, _note_list(cd["mismatched"], "days off")))
        checks.append((f"Sales lines {dmin}..{dmax} (file vs DB rows)", float(cd["file_rows"]), float(cd["db_rows"]),
                       0.0, 0.0, _note_list(cd["count_mismatched"], "days with a different line count")))
        cs = compare_groups(f_sm, d_sm)
        checks.append((f"Sales per salesman: max |file - DB| BHD ({cs['groups']} names)", 0.0, cs["max_diff"], 0.0,
                       ABS_TOL_BHD, _note_list(cs["mismatched"], "salesmen off")))
    if od:
        dts = sorted(str(r["order_date"])[:10] for r in od if r.get("order_date"))
        if dts:
            dmin, dmax = dts[0], dts[-1]
            f_day = group_sums(od, lambda r: str(r.get("order_date") or "")[:10], "gross_bhd")
            d_day = _db_groups("SELECT order_date::text AS k, COALESCE(SUM(gross_bhd),0) AS s, COUNT(*) AS n FROM orders "
                               "WHERE order_date >= $1::date AND order_date <= $2::date GROUP BY 1", [dmin, dmax],
                               rest=("orders", "order_date", "gross_bhd", "order_date"))
            cd = compare_groups(f_day, d_day)
            checks.append((f"Invoices per day: max |file - DB| BHD ({cd['groups']} days)", 0.0, cd["max_diff"], 0.0,
                           ABS_TOL_BHD, _note_list(cd["mismatched"], "days off")))
            checks.append((f"Invoices {dmin}..{dmax} (file vs DB rows)", float(cd["file_rows"]), float(cd["db_rows"]),
                           0.0, 0.0, _note_list(cd["count_mismatched"], "days with a different invoice count")))


def _stock_snapshot_checks(sb: list[dict], src: Path, checks: list) -> None:
    """The warehouse set of the newest snapshot vs the one before it (D5): a narrower export FAILS
    unless the folder carries PARTIAL_OK. Also the snapshot's row count, file vs DB."""
    new_set = {str(r.get("warehouse_name") or "(unassigned)") for r in sb}
    as_of = max((str(r.get("as_of_date") or "")[:10] for r in sb), default="")
    if not as_of:
        return
    prev_date, prev_set = None, set()
    try:
        r = _rpc("SELECT MAX(as_of_date)::text AS d FROM stock_balance WHERE as_of_date < $1::date", [as_of])
        prev_date = (r or [{}])[0].get("d")
        if prev_date:
            rows = _rpc("SELECT DISTINCT COALESCE(warehouse_name, '(unassigned)') AS w FROM stock_balance "
                        "WHERE as_of_date = $1::date", [prev_date])
            prev_set = {str(x.get("w")) for x in (rows or [])}
    except Exception:  # noqa: BLE001 -- RPC unavailable: REST
        c = get_client()
        r = (c.table("stock_balance").select("as_of_date").lt("as_of_date", as_of)
             .order("as_of_date", desc=True).limit(1).execute().data or [])
        prev_date = (r or [{}])[0].get("as_of_date")
        if prev_date:
            rows = (c.table("stock_balance").select("warehouse_name").eq("as_of_date", str(prev_date)[:10])
                    .order("id").limit(5000).execute().data or [])
            prev_set = {str(x.get("warehouse_name") or "(unassigned)") for x in rows}
    partial_ok = marker_present(src, PARTIAL_OK_MARKER)
    reason = warehouse_set_guard(prev_set, new_set, partial_ok)
    missing = 0.0 if reason is None else float(len(prev_set - new_set))
    checks.append((f"Stock snapshot {as_of}: previous warehouses missing" + (" (PARTIAL_OK)" if partial_ok else ""),
                   0.0, missing, 0.0, 0.0,
                   (reason or f"{len(new_set)} warehouse(s)" + (f"; previous {prev_date} had {len(prev_set)}" if prev_date else ""))))
    # rows and value of the snapshot, scoped to the warehouses the file covers: the loader replaces
    # per (as_of_date, warehouse), so other warehouses of the same day are kept on purpose.
    import json
    keys = {(str(r.get("item_name")), str(r.get("warehouse_name") or "(unassigned)")) for r in sb}
    whs = json.dumps(sorted(new_set))
    scope = ("FROM stock_balance WHERE as_of_date = $1::date "
             "AND COALESCE(warehouse_name, '(unassigned)') IN (SELECT jsonb_array_elements_text($2::jsonb))")
    n_db = _sql_sum(f"SELECT COUNT(*) AS s {scope}", [as_of, whs])
    v_db = _sql_sum(f"SELECT COALESCE(SUM(total_value_bhd), 0) AS s {scope}", [as_of, whs]) if n_db is not None else None
    if n_db is None:   # the RPC could not read stock_balance: page it over REST, in id order
        c = get_client()
        rows_db, off = [], 0
        while True:
            b = (c.table("stock_balance").select("warehouse_name,total_value_bhd").eq("as_of_date", as_of)
                 .in_("warehouse_name", sorted(new_set)).order("id").range(off, off + 999).execute().data or [])
            rows_db += b
            if len(b) < 1000:
                break
            off += 1000
        n_db, v_db = float(len(rows_db)), sum(_nn(r.get("total_value_bhd")) for r in rows_db)
    checks.append((f"Stock snapshot {as_of} rows (file vs DB, {len(new_set)} warehouse(s))", float(len(keys)), n_db, 0.0, 0.0, ""))
    checks.append((f"Stock snapshot {as_of} value BHD (file warehouses)", sum(_nn(r.get("total_value_bhd")) for r in sb),
                   v_db, 0.0, ABS_TOL_BHD, ""))


def _receivables_total_checks(f: str, ar: list[dict], checks: list) -> None:
    """Focus's Grand Total (stored in ar_ageing_totals) vs the file, and the rows-vs-Focus gap as a
    note: the gap is a known Focus-export trait (credits shown positive), not a load error."""
    tot = parse_receivables_totals(read_grid(f), os.path.basename(f))
    if not tot:
        return
    rows_sum = sum(_nn(r["balance_bhd"]) for r in ar)
    gap = round(rows_sum - _nn(tot["focus_total_bhd"]), 3)
    note = (f"rows sum {rows_sum:,.3f} vs Focus total {_nn(tot['focus_total_bhd']):,.3f}: gap {gap:+,.3f}"
            + (" -- credits may be shown as owed; ask accounts for a signed / Dr-Cr export" if abs(gap) > 0.005 else ""))
    stored = _sql_sum("SELECT COALESCE(MAX(focus_total_bhd),0) AS s FROM ar_ageing_totals WHERE as_of_date = $1::date",
                      [str(tot.get("as_of_date") or "")[:10]])
    if stored is None:
        checks.append((f"Receivables Focus total {tot.get('as_of_date')} (not stored: migration pending)",
                       _nn(tot["focus_total_bhd"]), _nn(tot["focus_total_bhd"]), 0.0, ABS_TOL_BHD, note))
    else:
        checks.append((f"Receivables Focus total {tot.get('as_of_date')} stored", _nn(tot["focus_total_bhd"]), stored,
                       0.0, ABS_TOL_BHD, note))


def export_date(path) -> str:
    """The day a Focus price book was exported, ISO. The workbook's docProps `created` stamp (the
    exporter writes it; a price book has no title block to read a date from), else the file's
    mtime -- the EARLIER of the two, because `created` carries the server's clock (measured
    2 h 30 ahead of the download time on the 14-Sep-2026 book) and could roll past midnight."""
    p = Path(path)
    cands: list[str] = []
    try:
        from openpyxl import load_workbook
        wb = load_workbook(p, read_only=True)
        try:
            c = wb.properties.created
        finally:
            wb.close()
        if c:
            cands.append(c.date().isoformat())
    except Exception:  # noqa: BLE001 — not an xlsx, or no properties part
        pass
    try:
        cands.append(datetime.fromtimestamp(p.stat().st_mtime).date().isoformat())
    except OSError:
        pass
    return min(cands) if cands else date.today().isoformat()


def current_prices(rows: list[dict], today: str | None = None) -> dict[str, float]:
    """Pure: the price v_price_list_by_book shows per SKU, derived from ONE parsed price-book file
    by the view's own rule -- Authorized, rate > 0, started, not ended; the base (blank-warehouse)
    layer first, then the latest start_date, then the row later in the file (= the higher id)."""
    today = today or date.today().isoformat()
    best: dict[str, tuple[tuple, float]] = {}
    for i, r in enumerate(rows):
        sku = r.get("sku_code")
        start, end = r.get("start_date"), r.get("end_date")
        rate = _nn(r.get("rate_bhd"))
        if not sku or (r.get("status") or "") != "Authorized" or rate <= 0 or not start:
            continue
        if str(start) > today or (end and str(end) < today):
            continue
        rank = (r.get("warehouse_name") in (None, ""), str(start), i)
        if sku not in best or rank > best[sku][0]:
            best[sku] = (rank, rate)
    return {s: v[1] for s, v in best.items()}


def _db_book_prices(book: str) -> dict[str, float]:
    """{sku: live price} for one book from v_price_list_by_book (RPC, ordered REST fallback)."""
    try:
        rows = _rpc("SELECT sku_code, price_bhd FROM v_price_list_by_book WHERE price_book = $1", [book])
    except Exception:  # noqa: BLE001
        rows = (get_client().table("v_price_list_by_book").select("sku_code,price_bhd")
                .eq("price_book", book).order("sku_code").limit(5000).execute().data or [])
    return {str(r["sku_code"]): _nn(r.get("price_bhd")) for r in (rows or []) if r.get("sku_code")}


def _db_today() -> str:
    """The database's CURRENT_DATE (UTC), which is what every price view compares start_date
    against. The local clock is 3 h ahead of it in Bahrain, so a price starting "today" would
    otherwise look live here and not there between midnight and 03:00."""
    try:
        d = (_rpc("SELECT CURRENT_DATE::text AS d") or [{}])[0].get("d")
        if d:
            return str(d)[:10]
    except Exception:  # noqa: BLE001
        pass
    return date.today().isoformat()


def _db_live_focus_rows(book: str, key: str) -> list[dict]:
    """Live rows of `book` that came from a FOCUS export of it (source_file starts with the book's
    file stem; workbook imports 'YQ_MRN%' are a separate source of truth and never count). Filters
    voided rows when the column exists (selling_prices_void_migration.sql), everything otherwise."""
    cols = "sku_code, price_book, customer_code, warehouse_name, start_date::text AS start_date, source_file"
    base = f"SELECT {cols} FROM selling_prices WHERE price_book = $1 AND source_file ILIKE $2"
    like = f"{key}%"
    try:
        try:
            return _rpc(base + " AND voided_at IS NULL", [book, like]) or []
        except Exception:  # noqa: BLE001 — no voided_at column yet
            return _rpc(base, [book, like]) or []
    except Exception:  # noqa: BLE001
        q = (get_client().table("selling_prices")
             .select("sku_code,price_book,customer_code,warehouse_name,start_date,source_file")
             .eq("price_book", book).ilike("source_file", like).order("id").limit(5000))
        return q.execute().data or []


def _price_book_checks(src: Path, checks: list) -> None:
    """Three source-backed checks per price book in the upload (none when no book was uploaded):
      * the current price per SKU derived from the FILE equals v_price_list_by_book, SKU by SKU
        (as of the database's own CURRENT_DATE);
      * no live Focus row of that book is dated after the file's export day unless the file itself
        carries it (a scheduled price Focus knows about). 14-Sep-2026: 165 day-first rows put
        prices on 5-Oct and 1-Nov that Focus never set -- this is the check that would have caught
        them;
      * no live Focus row of that book at all whose key the file does not carry: the file is the
        whole book, so such rows are older exports the loader's snapshot rule should have voided.
        This is the line that FAILS the refresh when the loader's mass-void guard skipped the void
        (the stale rows stay live), and until selling_prices_void_migration.sql has run.
    Workbook rows (YQ_MRN%) never count; the newest export in the folder is the one checked."""
    today = _db_today()
    for key, book in _PRICE_BOOKS:
        f = _find_newest_book(key, src)
        if not f:
            continue
        rows = parse_pricebook(read_grid(f), os.path.basename(f), book)
        if not rows:
            continue
        exported = export_date(f)
        want = current_prices(rows, today=today)
        live = _db_book_prices(book)
        agree = sum(1 for s, p in want.items() if s in live and abs(live[s] - p) < 0.0005)
        checks.append((f"{book} price = file ({len(want)} SKUs)", float(len(want)), float(agree), 0.0))
        keys = {pricebook_key(r) for r in rows}
        stale = [r for r in _db_live_focus_rows(book, key) if pricebook_key(r) not in keys]
        ghosts = [r for r in stale if str(r.get("start_date") or "")[:10] > exported]
        checks.append((f"{book} rows after {exported} not in file", 0.0, float(len(ghosts)), 0.0))
        checks.append((f"{book} live rows not in file (older exports)", 0.0, float(len(stale)), 0.0))


def evaluate_checks(checks: list) -> tuple[bool, list[dict]]:
    """Pure: (all_ok, rows). A check is (name, report, db, tol%) or (name, report, db, tol%,
    abs_tol_bhd, note): the absolute tolerance (the fils) decides when given, the percentage
    otherwise. Rows carry `note` when the check gave one."""
    rows, ok = [], True
    for chk in checks:
        name, report, db, tol = chk[0], float(chk[1]), float(chk[2]), chk[3]
        abs_tol = chk[4] if len(chk) > 4 else None
        note = chk[5] if len(chk) > 5 else None
        diff_pct = abs(report - db) / report * 100 if report else (0 if db == 0 else 100)
        passed = (abs(report - db) <= abs_tol + 1e-9) if abs_tol is not None else (diff_pct <= tol)
        ok = ok and passed
        row = {"metric": name, "report": report, "db": db, "diff_pct": diff_pct, "passed": passed}
        if note:
            row["note"] = note
        rows.append(row)
    return ok, rows


def run_checks(src_dir: Path | None = None, *, session: bool = True) -> tuple[bool, list[dict]]:
    """Crosscheck DB view totals against the source Focus reports in src_dir.

    Returns (all_ok, rows) where each row = {metric, report, db, diff_pct, passed}. A missing
    report file is skipped (not failed) so a partial refresh still validates what it loaded.
    Importable so the refresh engine can gate/annotate on the result.

    Where DATABASE_URL is set, one read-only database session answers every read, with the same
    rows as the REST/RPC path (_Session); session=False reads over the REST/RPC path only."""
    s = _open_session() if session else None
    token = _SESSION.set(s)
    try:
        return _run_checks(src_dir)
    finally:
        _SESSION.reset(token)
        if s is not None:
            s.close()


def _run_checks(src_dir: Path | None) -> tuple[bool, list[dict]]:
    """run_checks' body: every read goes through _rpc / _session_call or the REST client."""
    src = src_dir or NEW
    checks: list[tuple[str, float, float, float]] = []  # name, report, db, tol%

    ol: list[dict] | None = None
    od: list[dict] | None = None
    f = _find("sales_day_book", src)
    if f:  # Sales gross — scope the DB sum to the file's own date range (partial-upload safe)
        ol = parse_order_lines(read_grid(f), "x")
        report = sum(_nn(r["gross_bhd"]) for r in ol)
        dts = sorted(str(r["line_date"])[:10] for r in ol if r.get("line_date"))
        db = (_db_sum_between("v_sales", "revenue_bhd", "sale_date", dts[0], dts[-1])
              if dts else _db_sum("v_sales", "revenue_bhd"))
        checks.append(("Sales gross BHD", report, db, 0.5))
    f = _find("summary_sales_register", src)
    if f:  # Invoice headers -- gross per invoice, scoped to the file's own date range
        od = parse_orders(read_grid(f), "x")
        report = sum(_nn(r["gross_bhd"]) for r in od)
        dts = sorted(str(r["order_date"])[:10] for r in od if r.get("order_date"))
        db = (_db_sum_between("orders", "gross_bhd", "order_date", dts[0], dts[-1])
              if dts else _db_sum("orders", "gross_bhd"))
        checks.append(("Invoice headers gross BHD", report, db, 0.5))
    if ol is not None or od is not None:   # per day / per salesman to the fils + row counts
        _sales_detail_checks(ol or [], od, checks)
    f = _find("stock_ledger", src)
    if f:  # Stock movements -- units received + issued inside the file's date range. This is the
        # check that was missing on 21-Sep-2026: a Stock_ledger-only upload had nothing to compare
        # against and was reported as "NO SOURCE REPORTS FOUND" even though it loaded cleanly.
        sm = parse_stock(read_grid(f), "x")
        report = sum(_nn(r["received_qty"]) + _nn(r["issued_qty"]) for r in sm)
        dts = sorted(str(r["move_date"])[:10] for r in sm if r.get("move_date"))
        if dts:
            db = (_db_sum_between("stock_movements", "received_qty", "move_date", dts[0], dts[-1])
                  + _db_sum_between("stock_movements", "issued_qty", "move_date", dts[0], dts[-1]))
        else:
            db = _db_sum("stock_movements", "received_qty") + _db_sum("stock_movements", "issued_qty")
        checks.append(("Stock movements units", report, db, 0.5))
    f = _find("customer_summary_ageing_by_due_date", src)
    if f:  # Receivables total vs AR ageing report (v_receivables already = latest snapshot)
        ar = parse_receivables(read_grid(f), "x")
        checks.append(("Receivables BHD", sum(_nn(r["balance_bhd"]) for r in ar),
                       _db_sum("v_receivables", "outstanding_bhd"), 0.5))
        _receivables_total_checks(f, ar, checks)
    f = _find("stock_balance_by_warehouse", src)
    if f:  # Stock snapshot: warehouse set vs the previous snapshot, rows and selling value, all
        # scoped to the file's own as_of_date AND warehouses (_stock_snapshot_checks). The older
        # whole-day 'Stock value BHD' check summed every warehouse of the latest as_of_date, so a
        # same-day re-export that no longer lists a warehouse (Focus omits zero-stock warehouses,
        # the loader keeps that warehouse's earlier rows) gave a false FAIL beside the scoped one.
        sb = parse_stock_balance(read_grid(f), "x")
        _stock_snapshot_checks(sb, src, checks)
    _price_book_checks(src, checks)   # only when a price book is part of the upload

    # How many checks above are backed by an actual source FILE. Everything after this point
    # is DB-vs-DB and cannot detect a bad load, so this is the number that decides whether the
    # run proved anything at all.
    source_backed = len(checks)

    # Channel split vs total sales gross. NOTE: this is DB-vs-DB and very nearly a tautology —
    # v_sales_by_channel is `select channel, sum(revenue_bhd) from v_sales group by channel`
    # with no WHERE and no NULL branch in the channel CASE, so summing its groups is the same
    # number by construction. It is kept only as a guard against the view being redefined with
    # a filter or a channel that can go NULL. It must never be counted as evidence of a good
    # load, which is why it is added AFTER source_backed is taken.
    ch = _session_call("rest_rows", "v_sales_by_channel", ("revenue_bhd",), None, False)
    if ch is _OLD:
        c = get_client()
        ch = c.table("v_sales_by_channel").select("revenue_bhd").execute().data or []
    checks.append(("Channel = sales gross (DB-only)", _db_sum("v_sales", "revenue_bhd"),
                   sum(_nn(r["revenue_bhd"]) for r in ch), 0.5))

    ok, rows = evaluate_checks(checks)

    # No source file matched a cross-check rule. Reporting a bare PASS here would be the worst
    # outcome (a green light meaning "I found nothing to compare"), but a bare FAIL is wrong too
    # when the upload was made only of report types that have no rule (price books, profitability):
    # those loaded fine. So: PASS with an explicit "loaded, nothing to compare" line when the
    # ingest step produced rows for such a report; FAIL only when nothing at all was loaded.
    if source_backed == 0:
        loaded = _unchecked_loaded()
        if loaded:
            rows.append({"metric": "Loaded, no totals to cross-check: " + ", ".join(loaded),
                         "report": 0.0, "db": 0.0, "diff_pct": 0.0, "passed": True,
                         "note": "These report types have no verification rule; row counts are shown above."})
        else:
            ok = False
            rows.append({"metric": "NO SOURCE REPORTS FOUND", "report": 0.0, "db": 0.0,
                         "diff_pct": 100.0, "passed": False,
                         "note": "None of the uploaded files produced rows. Check the file names and contents."})
    return ok, rows


def _unchecked_loaded() -> list[str]:
    """Human labels of unchecked report types whose ingest CSV holds at least one row."""
    out: list[str] = []
    for table, label in _UNCHECKED.items():
        p = CLEAN_DIR / f"{table}.csv"
        try:
            if p.exists() and sum(1 for _ in p.open(encoding="utf-8")) > 1:
                out.append(label)
        except Exception:  # noqa: BLE001
            continue
    return out


def main() -> int:
    # Optional folder argument, same as scripts.refresh:  python -m scripts.verify_numbers "Planning 030726"
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else None
    if src is not None and not src.is_absolute():
        src = ROOT / src
    ok, rows = run_checks(src)
    print("=" * 72)
    print(f"{'METRIC':48} {'REPORT':>14} {'DB':>14}  RESULT")
    print("-" * 72)
    for r in rows:
        print(f"{r['metric'][:48]:48} {r['report']:>14,.3f} {r['db']:>14,.3f}  "
              f"{'PASS' if r['passed'] else 'FAIL'} ({r['diff_pct']:.2f}%)")
        if r.get("note"):
            print(f"    {r['note']}")
    print("=" * 72)
    print("ALL CHECKS PASS" if ok else "SOME CHECKS FAILED - investigate before trusting the dashboard.")
    return 0 if ok else 2


if __name__ == "__main__":
    raise SystemExit(main())
