"""Print an AI Head review to PDF and bundle the pack's tables as the appendix workbook (plan §19 / §22).

    python -m scripts.ai_head.render exports/ai_head/2026-09-26              # weekly.html -> weekly.pdf + weekly.xlsx
    python -m scripts.ai_head.render exports/ai_head/month-2026-09 --name monthly

Local files only: the PDF is weekly.html printed by headless Chromium (the same Playwright call as
scripts/weekly_report.py), the workbook is one sheet per pack CSV (openpyxl). Nothing is sent anywhere;
the folder is gitignored.
"""
from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

# the appendix order: what a reader looks for first
APPENDIX = ("shop_orders.csv", "shop_lines.csv", "rep_governance.csv", "merchants.csv", "regulars_due.csv",
            "items.csv", "search_demand.csv", "funnel_daily.csv", "market_signals.csv", "statements.csv",
            "focus_links.csv", "followups_holdout.csv", "insights_history.csv", "focus_lines_week.csv")


def sheet_name(fname: str) -> str:
    return fname.rsplit(".", 1)[0][:31]


def to_xlsx(folder: Path, out: Path) -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="6D4091")   # marketplace plum
    n = 0
    names = [f for f in APPENDIX if (folder / f).exists()] + sorted(
        p.name for p in folder.glob("*.csv") if p.name not in APPENDIX)
    for fname in names:
        with (folder / fname).open(encoding="utf-8", newline="") as f:
            rows = list(csv.reader(f))
        ws = wb.create_sheet(sheet_name(fname))
        for r in rows:
            ws.append(r)
        if rows:
            for c in ws[1]:
                c.font, c.fill = head_font, head_fill
            ws.freeze_panes = "A2"
            for col in ws.columns:
                width = max(len(str(c.value or "")) for c in list(col)[:200])
                ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 48)
        n += 1
    if n:
        wb.save(out)
    return n


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("folder", help="the pack folder (exports/ai_head/<date>)")
    ap.add_argument("--name", default="weekly", help="weekly (default) or monthly: <name>.html -> <name>.pdf / .xlsx")
    ap.add_argument("--no-pdf", action="store_true", help="only the workbook")
    a = ap.parse_args(argv)
    folder = Path(a.folder)
    if not folder.is_dir():
        print(f"render: {folder} is not a folder", file=sys.stderr)
        return 2
    page = folder / f"{a.name}.html"
    if not a.no_pdf:
        if not page.exists():
            print(f"render: {page.name} is missing -- the review writes it first", file=sys.stderr)
            return 2
        from scripts.weekly_report import to_pdf
        to_pdf(page.read_text(encoding="utf-8"), folder / f"{a.name}.pdf")
        print(f"wrote {a.name}.pdf")
    sheets = to_xlsx(folder, folder / f"{a.name}.xlsx")
    print(f"wrote {a.name}.xlsx ({sheets} sheets)" if sheets else "no CSV in the folder: no workbook")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
