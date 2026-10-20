"""His edited Stock & Order Plan, read back: the quantities and comments he typed, as his file holds them.

He works on the day's workbook (scripts.ops.workbook.stock_book) in Excel, on the PC or the phone: a quantity in a
yellow Your qty cell, a comment, an item added by its code under the order list. When he says an order is placed
(/yq-stock ordered <supplier>) or asks for the team copy (/yq-stock team), the file is read back here. Nothing in
this module changes his file or touches the database.

HOW THE FILE IS READ
  * Its bytes are read once, and the sha256 is theirs (a preview's token knows the same file again). A file that
    another program holds (PermissionError) is copied first, to a temporary name under exports/ of the data root,
    and the copy is read and removed. Excel leaves a file it has open readable: what is read is then its last
    save, and open_in_excel says so (Excel's "~$<name>" owner file is beside it).
  * The layout tag in Notes!B5 names the layout: stock-plan v1 and v2 are read. Another tag, or none, is refused
    in one sentence (PlanFileError), and so is a team copy (stock_team gives it its own tag): a team copy is
    never read as his plan.
  * The Order sheet's item table is found by its header texts (Supplier, Code, Item, Suggested, Your qty, Final,
    Comment), never by fixed letters. Its rows run from the header to the last row of its filter (to the end of
    the sheet when the filter is gone); empty rows and the notes under the table (text in its first column
    only) are passed over.
  * A typed cell is read as typed (openpyxl data_only=False): a formula typed in Your qty stays a formula. A
    formula cell is read as Excel last worked it out (data_only=True). A file saved by a program that does not
    calculate carries no value for its formulas: then the engine's own evaluator (formula.py) works the cell out
    from the file's formulas and cells, as Excel would.
  * The plan's own rows (their Item a plain value) are read whatever is typed in them. The 15 rows to add an
    item are the rows under them whose Item cell looks the code up on All items (found by that, so a sort that
    moved them moves nothing); one counts when a code and a quantity are typed in it (a code alone is a
    look-up, not an order line).

THE RULE FOR A ROW
  final = Your qty when it is a whole number, nought or more; else Suggested (a blank Your qty takes it), when
  that is a whole number, nought or more; else None.
  A problem is one plain sentence naming the row (the Excel row number), in sheet order. A problem never raises:
  the caller decides (/yq-stock ordered logs nothing while one stands). They are:
    Your qty that holds a formula, is not a number (text, a date, TRUE), is not whole or is below nought;
    a quantity typed with no code; an added code the All items sheet does not hold ("code not found");
    a row to order (final above nought) with no code, or naming no supplier; a blank Your qty whose Suggested is
    no whole number (no value, an error value); the same code on two rows, each with a quantity (typed on both,
    or a final above nought on both: an order logs one line per code).

CONTRACT (scripts.ops.stock and scripts.ops.workbook.stock_team rely on it)
  PlanFileError(ValueError)   the file cannot be read as his plan: one sentence ending in a full stop, the file's
                              name in it
  LAYOUTS                     ("stock-plan v1", "stock-plan v2"), the tags read, in TAG_SHEET!TAG_CELL (Notes!B5)
  TEAM_MARK                   " - Team": a team copy's name holds it, and newest_plan passes such a file over
  newest_plan(folder) -> Path | None
      The newest "YQ Stock and Order Plan <YYYY-MM-DD>*.xlsx" in folder: the latest data date in the name, then
      the highest copy number (" (2)", " (3)" ...), then the latest save. Excel's "~$" lock files, the engine's
      "~build-" temporary files and team copies are passed over; None when there is none, or no folder.
  read_plan(path) -> dict
      {"path": str, "sha256": the hex digest of the bytes read, "layout": the tag, "data_date": "YYYY-MM-DD" from
       the name (None when it holds none), "rows": [row], "problems": [str], "open_in_excel": bool,
       "saved_at": the file's modified time, ISO, Bahrain (None when it cannot be read)}
      row = {"row": the Excel row, "supplier": str ("" when blank), "code": str ("" when blank), "item": str,
             "suggested": as the file holds it (a number; None when it has no value, an error value's text),
             "your_qty": the typed value (None when blank; a formula as its text, "=..."),
             "final": int | None (the rule above), "comment": str | None, "added": True for an added row}
      Raises PlanFileError, nothing else, for a file that is missing, a folder, held by another program even for
      a copy, no workbook, a team copy, of an unknown layout, or with no Order item table.
  final_lines(plan, supplier) -> [(code, qty)]
      The rows of that supplier (case and spaces ignored) with a code and a final above nought, in sheet order,
      qty an int: what an order placed would log. It does not look at the problems: the caller does.
  load(path) -> PlanBook; plan_of(book) -> dict
      read_plan in two steps: stock_team reads the plan and its other sheets from the same read of the bytes.
  PlanBook.value(sheet, row, col)   a cell (1-based) as his file shows it (typed, saved or worked out); NoValue
                                    when it is a formula that cannot be worked out. value_or(...) gives a
                                    default instead.
  PlanBook.order                    (header row, {header folded: column}, last row) of the Order item table
  header_row(book, sheet, heads) -> (header row, {header folded: column}) | None
      The first row of that sheet holding every one of heads (folded: spaces collapsed, case ignored).
"""
from __future__ import annotations

import hashlib
import io
import os
import re
import uuid
import warnings
from datetime import date, datetime
from pathlib import Path

from scripts.ops import paths
from scripts.ops.workbook import book as bk
from scripts.ops.workbook import formula as fx

LAYOUTS = ("stock-plan v1", "stock-plan v2")
TEAM_TAG_START = "stock-plan team"          # stock_team.TEAM_TAG starts so
TAG_SHEET, TAG_CELL = "Notes", "B5"
STEM = "YQ Stock and Order Plan"
TEAM_MARK = " - Team"
ORDER_SHEET = "Order"
H_SUPPLIER, H_CODE, H_ITEM, H_SUGGESTED, H_YOURS, H_FINAL, H_COMMENT = (
    "Supplier", "Code", "Item", "Suggested", "Your qty", "Final", "Comment")
ORDER_HEADS = (H_SUPPLIER, H_CODE, H_ITEM, H_SUGGESTED, H_YOURS, H_FINAL, H_COMMENT)
NOT_FOUND = "code not found"                # the Item an added row shows for a code All items does not hold
HEAD_SCAN = 200                             # the rows a header is looked for in

_PLAN_NAME = re.compile(r"^YQ Stock and Order Plan (\d{4})-(\d{2})-(\d{2})(.*)\.xlsx$", re.I)
_COPY = re.compile(r" \((\d+)\)$")


class PlanFileError(ValueError):
    """The file cannot be read as his plan: one plain sentence, printed as it is."""


class NoValue(Exception):
    """A formula cell with no saved value that the evaluator cannot work out either."""


def fold(text) -> str:
    """A header, a code or a supplier as it is compared: spaces collapsed, case ignored ('' for nothing)."""
    return "" if text is None else " ".join(str(text).split()).casefold()


# ── the newest plan ───────────────────────────────────────────────────────────

def _plan_date(name: str) -> date | None:
    m = _PLAN_NAME.match(name)
    if not m:
        return None
    try:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except ValueError:
        return None


def _copy_number(name: str) -> int:
    m = _COPY.search(name[: -len(".xlsx")] if name.lower().endswith(".xlsx") else name)
    return int(m.group(1)) if m else 1


def _mtime_ns(p: Path) -> int:
    try:
        return os.stat(bk._long(p)).st_mtime_ns
    except OSError:
        return 0


def newest_plan(folder) -> Path | None:
    """His newest plan in folder (see the module docstring), or None."""
    folder = Path(folder)
    if not folder.is_dir():
        return None
    best, best_key = None, None
    for p in folder.iterdir():
        name = p.name
        if name.startswith("~$") or name.startswith("~build-") or TEAM_MARK.casefold() in name.casefold():
            continue
        d = _plan_date(name)
        if d is None or not p.is_file():
            continue
        key = (d, _copy_number(name), _mtime_ns(p), name)
        if best_key is None or key > best_key:
            best, best_key = p, key
    return best


# ── one read of the file ──────────────────────────────────────────────────────

def _read_bytes(path: Path) -> bytes:
    with open(bk._long(path), "rb") as f:
        return f.read()


def _copy_aside(path: Path, tmp: Path) -> None:
    paths.copy_file(path, tmp)


def _bytes_of(path: Path) -> bytes:
    name = path.name
    if path.is_dir():
        raise PlanFileError(f"{name} is a folder, not a Stock & Order Plan workbook.")
    try:
        return _read_bytes(path)
    except FileNotFoundError:
        raise PlanFileError(f"{name} is not there: it was moved or deleted.") from None
    except PermissionError:
        pass
    except OSError as exc:
        raise PlanFileError(f"{name} could not be read ({type(exc).__name__}).") from None
    # another program holds it: a copy (Windows copies a file some programs hold), read and removed at once
    tmp = paths.exports_dir() / f"~read-{uuid.uuid4().hex[:12]}.xlsx"
    try:
        _copy_aside(path, tmp)
        return _read_bytes(tmp)
    except OSError:
        raise PlanFileError(f"{name} is held by another program: save and close it in Excel, then try "
                            "again.") from None
    finally:
        try:
            paths.remove_file(tmp)
        except OSError:
            pass


def _saved_at(path: Path) -> str | None:
    from scripts.ops import status
    try:
        stamp = os.stat(bk._long(path)).st_mtime
    except OSError:
        return None
    return datetime.fromtimestamp(stamp, status.BAHRAIN).isoformat(timespec="seconds")


def _open_in_excel(path: Path) -> bool:
    """Excel's owner file beside it: "~$" and the name (Word's way, the name less its first two letters, too)."""
    name = path.name
    return any(bk.exists(path.with_name("~$" + n)) for n in (name, name[2:]) if n)


def _grid(ws, typed: bool) -> tuple[dict, set]:
    """{(row, col): value} (1-based) of the cells that hold something; with typed, the formula cells too."""
    values: dict = {}
    formulas: set = set()
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if v is None:
                continue
            key = (cell.row, cell.column)
            if typed and cell.data_type == "f":
                text = v if isinstance(v, str) else getattr(v, "text", None)    # an array formula keeps its text
                formulas.add(key)
                v = text if isinstance(text, str) and text else "="
                if not v.startswith("="):
                    v = "=" + v
            values[key] = v
    return values, formulas


def _as_saved(v):
    """A value the evaluator worked out, as a saved value would read: whole numbers as int, an error as its code,
    an empty text as nothing."""
    if isinstance(v, fx.Err):
        return v.code
    if v is None or v == "":
        return None
    if isinstance(v, bool) or isinstance(v, str):
        return v
    if isinstance(v, (int, float)):
        return bk.cached(v)
    return v


class PlanBook:
    """One read of a workbook: its sha, its cells as typed and as last saved, its filters, its layout."""

    def __init__(self, path: Path, data: bytes):
        import openpyxl

        self.path = Path(path)
        self.sha256 = hashlib.sha256(data).hexdigest()
        name = self.path.name
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")              # openpyxl's notes on parts it does not keep
                typed = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
                saved = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        except Exception as exc:  # noqa: BLE001  (a file that is no workbook fails in many ways: one sentence)
            raise PlanFileError(f"{name} cannot be read as an Excel workbook ({type(exc).__name__}).") from None
        self.sheetnames = list(typed.sheetnames)
        self.typed: dict[str, dict] = {}
        self.formulas: dict[str, set] = {}
        self.saved: dict[str, dict] = {}
        self.filters: dict[str, str | None] = {}
        for ws in typed.worksheets:
            self.typed[ws.title], self.formulas[ws.title] = _grid(ws, typed=True)
            self.filters[ws.title] = ws.auto_filter.ref if ws.auto_filter is not None else None
        for ws in saved.worksheets:
            self.saved[ws.title] = _grid(ws, typed=False)[0]
        self._engine = None
        self.layout = self._layout()
        self.order = self._order_table()

    # -- cells (1-based rows and columns)
    def has(self, sheet: str, r: int, c: int) -> bool:
        return (r, c) in self.typed.get(sheet, {})

    def is_formula(self, sheet: str, r: int, c: int) -> bool:
        return (r, c) in self.formulas.get(sheet, ())

    def typed_value(self, sheet: str, r: int, c: int):
        """What is in the cell as typed: a formula as its text ("=..."), None when blank."""
        return self.typed.get(sheet, {}).get((r, c))

    def value(self, sheet: str, r: int, c: int):
        """The cell as his file shows it: a typed value as typed; a formula as last saved, else worked out."""
        if not self.is_formula(sheet, r, c):
            return self.typed_value(sheet, r, c)
        v = self.saved.get(sheet, {}).get((r, c))
        if v is not None:
            return v
        return self._worked_out(sheet, r, c)

    def value_or(self, sheet: str, r: int, c: int, default=None):
        try:
            return self.value(sheet, r, c)
        except NoValue:
            return default

    def _worked_out(self, sheet: str, r: int, c: int):
        try:
            return _as_saved(self._engine_book().value(sheet, r - 1, c - 1))
        except (fx.FormulaError, KeyError, IndexError, ValueError, TypeError, ArithmeticError, RecursionError):
            raise NoValue(f"{sheet}!{fx.col_letters(c - 1)}{r}") from None

    def _engine_book(self) -> bk.Book:
        """The file's cells in the report engine's Book, so formula.py can work a formula out (made once)."""
        if self._engine is None:
            book = bk.Book()
            for name in self.sheetnames:
                ws = book.sheet(name)
                formulas = self.formulas[name]
                for (r, c), v in self.typed[name].items():
                    if (r, c) in formulas:
                        ws.formula(r - 1, c - 1, v[1:].replace("_xlfn.", "").replace("_xlws.", ""))
                    elif isinstance(v, (bool, int, float, date)):
                        ws.put(r - 1, c - 1, v)
                    else:
                        ws.text(r - 1, c - 1, str(v))
            self._engine = book
        return self._engine

    # -- the layout and the Order table
    def _layout(self) -> str:
        name = self.path.name
        tr, tc = fx.parse_cell(TAG_CELL)
        tag = self.typed.get(TAG_SHEET, {}).get((tr + 1, tc + 1)) if TAG_SHEET in self.typed else None
        tag = " ".join(str(tag).split()) if tag is not None else ""
        if tag.casefold().startswith(TEAM_TAG_START):
            raise PlanFileError(f"{name} is a team copy (no costs, his quantities as values): read your own plan "
                                "file instead.")
        if tag not in LAYOUTS:
            said = f"is {tag!r}" if tag else "is missing"
            raise PlanFileError(f"{name} is not a Stock & Order Plan this engine reads: its layout tag ({TAG_SHEET}!"
                                f"{TAG_CELL}) {said}, not {' or '.join(LAYOUTS)}.")
        return tag

    def _order_table(self) -> tuple[int, dict, int]:
        found = header_row(self, ORDER_SHEET, ORDER_HEADS)
        if found is None:
            raise PlanFileError(f"{self.path.name} has no order list this engine reads: the Order sheet's item table "
                                f"(its header {', '.join(ORDER_HEADS)}) is not found.")
        head, cols = found
        last = max((r for r, _c in self.typed[ORDER_SHEET]), default=head)
        ref = self.filters.get(ORDER_SHEET)
        if ref:
            try:
                rng = fx.parse_range(ref, ORDER_SHEET)
            except fx.FormulaError:
                rng = None
            if rng is not None and rng.r1 + 1 == head and rng.r2 + 1 > head:
                last = rng.r2 + 1                      # the filter covers the item table: its last row ends it
        return head, cols, last


def header_row(book: PlanBook, sheet: str, heads) -> tuple[int, dict] | None:
    """(header row, {header folded: column}) of the first row of sheet holding every one of heads, or None."""
    cells = book.typed.get(sheet)
    if not cells:
        return None
    by_row: dict[int, dict] = {}
    for (r, c), v in sorted(cells.items()):
        if r <= HEAD_SCAN and isinstance(v, str) and (r, c) not in book.formulas[sheet]:
            by_row.setdefault(r, {}).setdefault(fold(v), c)
    want = [fold(h) for h in heads]
    for r in sorted(by_row):
        if all(h in by_row[r] for h in want):
            return r, by_row[r]
    return None


def load(path) -> PlanBook:
    """One read of his plan: the bytes, the two views of the cells, the layout tag and the Order table checked."""
    path = Path(path)
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        raise PlanFileError("openpyxl is not installed on this PC (pip install openpyxl): the plan cannot be read "
                            "back.") from None
    return PlanBook(path, _bytes_of(path))


# ── the rows ──────────────────────────────────────────────────────────────────

def _text(v) -> str | None:
    """A cell as text: None when blank; a whole number without '.0'; a date as YYYY-MM-DD."""
    if v is None:
        return None
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = " ".join(str(v).split())
    return s or None


def _note(v) -> str | None:
    """A comment as he typed it: the ends trimmed, the line breaks inside kept; None when blank."""
    if v is None or isinstance(v, (bool, int, float, date)):
        return _text(v)
    s = str(v).strip()
    return s or None


def _whole(v) -> int | None:
    """A whole number of units, nought or more, else None."""
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        return None
    if isinstance(v, float) and (not v == v or v in (float("inf"), float("-inf"))):
        return None
    if v != int(v) or v < 0:
        return None
    return int(v)


def _shown(v) -> str:
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (datetime, date)):
        return f"a date, {_text(v)}"
    return repr(v) if isinstance(v, str) else str(v)


def _qty_problem(v) -> str | None:
    """Why a typed Your qty is not a whole number of units, nought or more (None when it is one)."""
    if isinstance(v, str) and v.startswith("="):
        return "holds a formula: type a number"
    if isinstance(v, bool) or isinstance(v, (datetime, date)) or not isinstance(v, (int, float)):
        return f"is not a number ({_shown(v)}): type a whole number of units, or clear the cell"
    if isinstance(v, float) and not (v == v and abs(v) != float("inf")):
        return f"is not a number ({v}): type a whole number of units"
    if v != int(v):
        return f"is not a whole number ({_shown(v)}): type a whole number of units"
    if v < 0:
        return f"is below nought ({_shown(v)}): type nought or more"
    return None


def _rows_words(rows: list[int]) -> str:
    rs = [str(r) for r in rows]
    return rs[0] if len(rs) == 1 else ", ".join(rs[:-1]) + " and " + rs[-1]


def plan_of(book: PlanBook) -> dict:
    """read_plan's answer for a PlanBook (see the CONTRACT)."""
    try:
        rows, problems = _order_rows(book)
    except (KeyError, TypeError, AttributeError, IndexError, ArithmeticError, ValueError) as exc:
        if isinstance(exc, PlanFileError):
            raise
        raise PlanFileError(f"{book.path.name} could not be read back ({type(exc).__name__}): its Order sheet holds "
                            "something this engine does not know.") from None
    d = _plan_date(book.path.name)
    return {"path": str(book.path), "sha256": book.sha256, "layout": book.layout,
            "data_date": d.isoformat() if d else None, "rows": rows, "problems": problems,
            "open_in_excel": _open_in_excel(book.path), "saved_at": _saved_at(book.path)}


def read_plan(path) -> dict:
    """His plan, read back (see the CONTRACT). Raises PlanFileError only."""
    return plan_of(load(path))


def _order_rows(book: PlanBook) -> tuple[list[dict], list[str]]:
    sheet = ORDER_SHEET
    head, cols, last = book.order
    col = {h: cols[fold(h)] for h in ORDER_HEADS}
    table_cols = sorted(set(cols.values()))
    first = table_cols[0]
    rows: list[dict] = []
    said: list[tuple[int, int, str]] = []                 # (row, order, words)

    def problem(r: int, words: str) -> None:
        said.append((r, len(said), words))

    no_code = "has a quantity but no code: type the item's code in its Code cell, or clear the quantity"
    for r in range(head + 1, last + 1):
        filled = [c for c in table_cols if book.has(sheet, r, c)]
        if not filled or (filled == [first] and not book.is_formula(sheet, r, first)):
            continue                                       # an empty row, or a note under the table
        added = book.is_formula(sheet, r, col[H_ITEM]) or book.is_formula(sheet, r, col[H_SUPPLIER])
        code = _text(book.value_or(sheet, r, col[H_CODE])) or ""
        typed = book.typed_value(sheet, r, col[H_YOURS])
        if added and (not code or typed is None):
            if not code and typed is not None:
                problem(r, f"Row {r} {no_code}")
            continue                                       # an empty row to add an item, or a code alone (a look-up)
        try:
            item, looked_up = _text(book.value(sheet, r, col[H_ITEM])) or "", True
        except NoValue:
            item, looked_up = "", False
        flagged = False                                    # an added row whose code is not looked up
        if added:
            if not looked_up:
                problem(r, f"Row {r}: the item of the code {code} could not be looked up: open the file in Excel, save "
                           "it and try again")
                flagged = True
            elif not item or item.casefold() == NOT_FOUND:
                problem(r, f"Row {r}: the code {code} is not on the All items sheet: check the code")
                item, flagged = NOT_FOUND, True
        elif not code and not item:
            if typed is not None:
                problem(r, f"Row {r} {no_code}")
            continue
        supplier = _text(book.value_or(sheet, r, col[H_SUPPLIER])) or ""
        try:
            suggested = book.value(sheet, r, col[H_SUGGESTED])
        except NoValue:
            suggested = None
        why = _qty_problem(typed) if typed is not None else None
        if why:
            problem(r, f"Your qty in row {r} {why}")
        if typed is not None and why is None:
            final = int(typed)
        else:
            final = _whole(suggested)
            if typed is None and final is None:
                if suggested is None:
                    problem(r, f"Suggested in row {r} has no value: open the file in Excel, save it and try again")
                else:
                    problem(r, f"Suggested in row {r} is {_shown(suggested)}, not a whole number: type your "
                               "quantity in Your qty")
        if final and not code:
            problem(r, f"Row {r} ({item}) has no code, so it cannot be logged: type its code in the Code cell, or 0 "
                       "in Your qty")
        elif final and not supplier and not flagged:
            problem(r, f"Row {r} ({code}) names no supplier: type it in the Supplier cell")
        rows.append({"row": r, "supplier": supplier, "code": code, "item": item, "suggested": suggested,
                     "your_qty": typed, "final": final, "comment": _note(book.value_or(sheet, r, col[H_COMMENT])),
                     "added": added})

    # the same code on two rows, each with a quantity: after that row's own problems
    groups: dict[str, list[dict]] = {}
    for row in rows:
        if row["code"]:
            groups.setdefault(fold(row["code"]), []).append(row)
    for group in groups.values():
        typed_n = sum(1 for x in group if x["your_qty"] is not None)
        ordered_n = sum(1 for x in group if x["final"])
        if len(group) > 1 and (typed_n > 1 or ordered_n > 1):
            rs = [x["row"] for x in group]
            problem(rs[0], f"Code {group[0]['code']} is on rows {_rows_words(rs)}, each with a quantity: keep it on "
                           "one row")
    return rows, [words for _r, _n, words in sorted(said)]


def final_lines(plan: dict, supplier) -> list[tuple[str, int]]:
    """[(code, qty)]: that supplier's rows with a code and a final above nought, in sheet order."""
    want = fold(supplier)
    if not want:
        return []
    out = []
    for row in plan.get("rows") or []:
        q = row.get("final")
        code = " ".join(str(row.get("code") or "").split())
        if fold(row.get("supplier")) != want or not code or isinstance(q, bool) or not isinstance(q, int) or q <= 0:
            continue
        out.append((code, q))
    return out
