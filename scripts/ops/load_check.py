"""python -m scripts.ops load check [ddmmyy] [again] [--pricebook load|skip] [--with-profitability]
                                    [--replace-ok] [--partial-ok] [--dev]
python -m scripts.ops load status [--run RUN] [--wait]

`check` is what /yq-load runs first. It reads one Focus drop folder (Uplaods/<ddmmyy>/) and production,
copies the reports that will be loaded into a NEW staging folder and prints one table

    File | Decision | Reason      LOAD, LOAD (main-store block only), HOLD BACK, LEAVE OUT,
                                  NEEDS YOUR DECISION, NOT LOADED, STOP

then the checks behind it, the figures to expect after the load, what else a load does and, when
nothing stops it, a token and the GO line of the writer (scripts/yq_load_go.py, another program).
The check changes nothing in the database: every read goes through scripts/ops/db.py. Its own files
go to business_data/staging/<run>/ and business_data/ops/loads/<run>/ (both git-ignored). The
owner's folder is only read.

TWO BLOCKS. The chat shows a short block, which the assistant retypes before the owner's box opens:
lists cut at CHAT_LIST_MAX with a count, and no list of what a load also does, of the staged files
or of the token (the GO line carries it). The whole check (its progress lines, the lists to LIST_MAX
and all of the above) is kept as loads/<run>/check.txt, which the short block names under its head
line.

NOTHING NEW. A drop that is already loaded is checked in full like any other. When the check finds
that every report the load would apply has nothing to insert, change, remove or bring back, and that
nothing else a load does would change anything (the stock ledger voucher by voucher, the sales by day
and salesman, the catalog sync, a last load that failed), it ends with the line "Nothing new in
<ddmmyy>: everything this load would apply is already loaded ... No load needed." instead of a token:
no GO line, no go.cmd (on 15-Oct-2026 one drop was loaded twice, and the second load sent a second
e-mail). The reports that are not part of the load (held back, left out, or waiting for his decision,
whose default is to leave them out) are named just above that line, with the words that load them:
he says them, the check runs again with that option, and then it has something to load. A row the
loader FOLDS from two lines of one key (FOLDED) is the one place the preview cannot see equality: the
loader adds the lines up in floating point, the preview in exact decimals, so such a row reads as
rewritten on every check of the same file. It is no change when it equals production's row in every
column, the figures that are added up differing only beyond the sixth decimal place (folded_noise);
everything else is compared exactly, by the preview. Any difference anywhere and the check ends as it
always did.

The preview's seconds are kept per report (scripts/ops/querytime.py): one progress line as each
report's comparison is done, and plan.json timings.preview_parts.

The steps, one printed line each as it finishes:
  1 preflight      the project folder (not a code worktree), the programs equal the live site, the
                   two database addresses name one project, no load is running
  2 the folder     one file per report, none open in Excel, none still being written, each titled
                   report carries its "As on date"; everything else is listed, never loaded
  3 staging        copies proven identical (sha256); the Ledger export is never copied
  4 stock balance  more than the main store in the file: a copy cut to the main-store block, proven
                   by re-reading it with the real parser and against Focus's own Sub Total
  5 profitability  held back when the report's cost of sales is short of the stock ledger's
  6 price book     compared key by key with the live rows of that book
  7 the preview    app.ingest_batch.run_preview, read-only: invoices, lines, days, the loader's guards
                   (every staged report but the stock ledger: step 8 compares that one)
  8 stock ledger   voucher by voucher against production (the loader replaces the file's whole span),
                   and its rows in one read: their count and one digest (Q_LEDGER_SAME)
  9 catalog        what the catalog sync that follows every load would switch on or off

Older reports are never loaded. A drop whose reports are all older than the last load stops at once
(the loader replaces the dates a file covers, so an older export would take out what was entered
after it); one older report in a current folder stops on its own line of the table.

What the check read must still be there when it ends: a report that changed in the folder while the
check ran, or a database that another load moved meanwhile, is a stop ("run the check again"), so
the token never binds a table that was worked out on something else.

One check counts for a folder: the newest. When a new check of a folder gets as far as staging,
every earlier check of THAT folder that never started loading is replaced by it: its plan.json gains
"superseded_by" (the writer refuses such a plan), its go.cmd is removed and its staging folder moves
to business_data/staging/superseded/<run>/. A run that started, finished or failed is never touched.
The run name is <ddmmyy>_<HHMM>, the minute the check was made: a second check inside one minute
waits for the next minute, so two checks never share a name.

--dev is for a proof run from a code worktree: it skips the worktree and live-site checks, issues no
token, no GO line and no go.cmd, records nothing on the action register (a check that stops
records nothing either: it may be looking at a broken file) and replaces no earlier check: a dev
run adds its own two folders and changes nothing else.

The word `again` after the folder is taken and ignored (the load program's too-fast answer says
"Type /yq-load <ddmmyy> again.", so he may type it); any other word there is refused.

An earlier load that was cut off (a lock whose worker is gone) stops the check at once: nothing new
is checked, and the last line is that load's own GO line, which is what resumes it (a dev run names
the run and prints no GO line). The dead lock of a REHEARSAL is not that: a rehearsal never reaches
the database, so the check goes on and says in one line that the load clears the lock. A load that
was CLOSED as failed after it had begun to write does not stop the check either: its first line
says so, and that this check starts the load that replaces it.

`status` follows the writer: the journal lines since the last call (the writer's own sentences: its
fields text, why and next, in full) and, as the last line, one of RUNNING at <step>, FINISHED,
FAILED at <step>, INTERRUPTED at <step>, NOT STARTED. With no --run it follows the run that holds
the lock, else the latest one. --wait asks again every 5 s and returns as soon as the load has
ended, or after 100 s.

Not done here, on purpose: the options are kept in plan.json and never written to the run's journal
(a journal line before the writer starts would read as a running load).
"""
from __future__ import annotations

import hashlib
import os
import re
import sys
import tempfile
import threading
import time
from collections import Counter
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

from scripts.ops import db  # noqa: F401  (first: it guards every channel before an app module loads)
from scripts.ops import alerts, config, fmt, gitinfo, loadplan, paths, status

STEPS = 9
SIZE_PAUSE_S = 0.4            # between the two size readings of every report file
DEADLINE_S = 540              # the tool call has 600 s: stop cleanly before it is cut
SESSION_TIMEOUT_S = 300
LIST_MAX = 20                 # vouchers or invoices listed before "and N more" (check.txt)
CHAT_LIST_MAX = 5             # ... in the block the chat shows
RECORD_MAX = 200              # items of one rule put on the register by one check
SCOPE_LOOKUPS_MAX = 40        # scoped price rows whose sales are looked up
REPS_MAX = 6                  # salesmen listed with what the load changes for each
OTHER_ROWS_MAX = 10
WAIT_POLL_S = 5
WAIT_MAX_S = 100
WAIT_START_GRACE_S = 20       # after the go, the worker needs a moment before its first journal line
RUN_NAME_WAITS = 2            # minutes a check waits for a run name that is free
# A staging folder with no plan.json beside it is a check that stopped before its table, or one that
# is still running. No check runs longer than DEADLINE_S, so one older than this is a left-over.
LEFTOVER_AFTER_S = DEADLINE_S + 60
SUPERSEDED = "superseded"     # business_data/staging/superseded/<run>/: the staging of a replaced check
EXCEL = (".xlsx", ".xls")
SYSTEM_FILES = ("desktop.ini", "thumbs.db", ".ds_store")
AUTO_SYNC = "price-book auto-sync"          # what app/catalog.py writes in catalog_items.updated_by
BOOK_STEMS = {"MA_base": "MASellingPriceBook", "modern_trade": "ModernTradeSellerBook"}
OFFSET_FILE = "status.offset"
DELIVERY_OFFSET_FILE = "delivery.offset"     # the delivery rows `load status` has shown
TIMING_SHOWN_FILE = "timing.shown"           # the load-time line is shown once
DELIVERY_WAIT_S = 600        # a finished load reads RUNNING at delivery for at most this long (the writer allows 480 s)
SEND_RETRY = "python -m scripts.yq_send --run {run}"     # the writer's SEND_LINE: the e-mail again, nothing reloaded
DELIVERY_ENDS = ("ended", "sent", "test_sent", "not_sent", "failed")   # rows that close a queued delivery
# Register rules of this check (scripts/ops/alerts.py keeps the register).
#   EVENT_RULES     something changed between two loads. Listed from the first time it is found
#                   (never a baseline) and never cleared by a later check: the item stays open
#                   until he marks it done or explained.
#   STATE_RULES     something that IS so, and clears by itself once it is gone.
#   BASELINE_RULES  the state rules whose FIRST run is a baseline ("already in the data": counted
#                   once, not listed). A report held back twice is an action from its first time.
EVENT_RULES = ("inv_changed", "voucher_changed")
STATE_RULES = ("price_scope", "profit_held")
BASELINE_RULES = ("price_scope",)
# The loader's marker files (scripts/load_supabase.py; scripts/ops/loadplan.py names them too) and
# the owner's option that alone may put each one into staging: marker -> (plan.json key, his word).
MARKER_OPTIONS = {"REPLACE_OK": ("replace_ok", "--replace-ok"), "PARTIAL_OK": ("partial_ok", "--partial-ok")}
REHEARSAL_LEFT = "A cut-off rehearsal left its lock; the load clears it."
NAMELESS_LOCK = "A left-over lock that names no load was found (no load was started under it); the load clears it."
LOCK_AGAIN_S = 2              # before a lock that cannot be read as one is taken for a left-over
# The three snapshot reports and the column of Q_ANCHORS that holds the newest one production has.
SNAPSHOT_HELD = {"stock_balance": "stock_as_of", "receivables": "ar_as_of", "product_profitability": "profit_as_of"}
OLDER = "older_than_loaded"                 # the flag of a drop or a report older than what is loaded
# The rows Focus itself closes a block of the stock balance with (lower case). The parser reads any
# first cell that holds such a word as a total row, an item name included.
CLOSING_ROWS = ("sub total", "total", "grand total")
# The parts of the database fingerprint (scripts/ops/loadplan.py) in the owner's words.
DB_WORDS = {"orders": "invoices", "order_lines": "sales lines", "stock_movements": "stock movements",
            "stock_balance": "stock", "ar_ageing": "receivables", "selling_prices": "prices",
            "product_profitability": "profitability", "ingest_runs": "the log of loads",
            "ingest_batches_open": "an upload waiting on the Data page"}

KIND_ORDER = ("order_lines", "orders", "stock_movements", "stock_balance", "receivables",
              "product_profitability", "selling_prices:MA_base", "selling_prices:modern_trade", "ledger_entries")
# the table each report loads into (app.ingest_batch: the preview's target); both price books are one table
KIND_TARGET = {"order_lines": "order_lines", "orders": "orders", "stock_movements": "stock_movements",
               "stock_balance": "stock_balance", "receivables": "ar_ageing",
               "product_profitability": "product_profitability", "selling_prices:MA_base": "selling_prices",
               "selling_prices:modern_trade": "selling_prices"}
# The snapshot tables whose rows the loader FOLDS when Focus lists one key on two lines (scripts/load_supabase.py:
# two accounts under one name in the receivables, one item twice in one warehouse of the stock balance), and the
# figures it then adds up (the stock balance averages its rate). The loader adds in floating point, the preview in
# exact decimals: production's folded row can differ from the file's in the last digits of these figures only.
FOLDED = {"ar_ageing": ("balance_bhd", "bucket_0_30", "bucket_31_60", "bucket_61_90", "bucket_91_120",
                        "bucket_121_150", "bucket_151_180", "bucket_181_210", "bucket_over_210", "total_bhd"),
          "stock_balance": ("net_qty", "total_value_bhd", "selling_rate_bhd")}
FOLD_DEFAULTS = {"stock_balance": {"warehouse_name": "(unassigned)"}}     # as the loader fills a blank key part
# One millionth: more than float rounding leaves on such a sum or mean, far less than a fils, a unit or any rate
# Focus prints (three decimals at most)
FOLD_NOISE = Decimal("0.000001")
FOLD_NOUN = {"ar_ageing": "one account", "stock_balance": "one item in one warehouse"}
# The reports that can sit outside a load and the owner's words that put them in it (the skill's Step 1 gives
# each with its option: --with-profitability, --pricebook load)
OPTION_SAY = {"product_profitability": "load profitability anyway", "selling_prices:MA_base": "load the price book",
              "selling_prices:modern_trade": "load the price book"}
DAILY = KIND_ORDER[:6]
ALWAYS_STAGED = ("order_lines", "orders", "stock_movements", "receivables")
LABEL = {
    "order_lines": "Sales day book", "orders": "Sales register", "stock_movements": "Stock ledger",
    "stock_balance": "Stock balance", "receivables": "Receivables ageing",
    "product_profitability": "Profitability report", "selling_prices:MA_base": "Price book",
    "selling_prices:modern_trade": "Price book (modern trade)", "ledger_entries": "Ledger export",
}
WORDS = {"load": "LOAD", "trim": "LOAD (main-store block only)", "hold_back": "HOLD BACK", "leave_out": "LEAVE OUT",
         "needs_decision": "NEEDS YOUR DECISION", "not_loaded": "NOT LOADED", "stop": "STOP"}
# What a blocking code of the preview means here (see read_preview).
HARD_CODES = ("db_diff_failed", "empty_report", "snapshot_without_date", "sales_join_below_80", "check_failed")
# The first words of the last line of a check that finds nothing new to load: /yq-load acts on them
# (.claude/skills/yq-load/SKILL.md, Step 1), and nothing after that line is a GO line.
NOTHING_NEW_HEAD = "Nothing new in"
# The chat shows a short block (lists cut at CHAT_LIST_MAX, no SIDE_EFFECTS, staged files or token). The whole check,
# its progress lines and lists to LIST_MAX included, is kept in the run's folder; this line, under the head line,
# names it.
WHOLE_CHECK = "The whole check, with the longer lists: business_data/ops/loads/{run}/check.txt"
PREVIEW_IN = "preview_in"     # loads/<run>/preview_in/: what the preview compares, removed again after it
SIDE_EFFECTS = (
    "it sorts the items into their categories again and syncs the catalog with the price book (see Catalog above);",
    "it adds a name alias for every new Focus item name;",
    "it clears the saved answers (cache), so the portal reads the new figures;",
    "within 15 minutes production runs its demand forecast: it may put purchase-order and price drafts in "
    "Approvals and send Telegram warnings.",
)

Q_ANCHORS = (
    "select current_date::text as today, "
    "(select max(line_date)::text from order_lines) as sales_to, "
    "(select max(order_date)::text from orders) as orders_to, "
    "(select max(as_of_date)::text from stock_balance) as stock_as_of, "
    "(select max(as_of_date)::text from ar_ageing) as ar_as_of, "
    "(select max(report_date)::text from product_profitability) as profit_as_of, "
    "(select max(move_date)::text from stock_movements) as moves_to"
)
Q_PROFIT_BEFORE = (
    "select item_name, net_amount_bhd, cogs_bhd, report_date::text as report_date from product_profitability "
    "where report_date = (select max(report_date) from product_profitability where report_date < %s)"
)
Q_LIVE_PRICES = (
    "select sku_code, item_name, price_book, customer_code, customer_name, warehouse_name, "
    "start_date::text as start_date, end_date::text as end_date, rate_bhd, status, source_file "
    "from selling_prices where price_book = %s and voided_at is null"
)
Q_SNAPSHOT = (            # through the read RPC ($1 = the main store)
    "select item_name, net_qty from stock_balance where as_of_date = (select max(as_of_date) from stock_balance) "
    "and warehouse_name = $1 and net_qty > 0"
)
Q_SCOPE_SALES = (         # through the read RPC: $1 item code, $2 start date, $3 the scoped price
    "select count(distinct invoice_no) as invoices, coalesce(sum(quantity), 0) as units from v_sales "
    "where sku_code = $1 and sale_date >= $2::date and quantity > 0 "
    "and abs(gross_bhd - quantity * $3::numeric) < 0.001"
)
Q_WAREHOUSES_BEFORE = (
    "select distinct coalesce(warehouse_name, '(unassigned)') as w from stock_balance "
    "where as_of_date = (select max(as_of_date) from stock_balance where as_of_date < %s)"
)
Q_ORDERS_SPAN = "select invoice_no, salesman from orders where order_date between %s and %s"
Q_LINES_SPAN = "select distinct invoice_no, warehouse_name from order_lines where line_date between %s and %s"
Q_LEDGER = (
    "select voucher, item_name, move_date::text as d, coalesce(voucher_type, '') as vt, "
    "coalesce(warehouse_name, '') as w, coalesce(to_warehouse_name, '') as tw, coalesce(received_qty, 0) as rq, "
    "coalesce(issued_qty, 0) as iq, count(*) as n from stock_movements where move_date between %s and %s "
    "group by 1, 2, 3, 4, 5, 6, 7, 8"
)
# step 8's one read of the rows themselves, the preview's rule for the ledger (its dates, within its items): are
# production's rows the file's rows? Their count and one md5 of their keys (voucher | item | row_hash; the hash covers
# every column of a row) against the file's (ledger_digest). So the preview never has to send the ledger's rows.
Q_LEDGER_SAME = (
    "select count(*) = %s and coalesce(md5(string_agg(k, E'\\n' order by k collate \"C\")), '') = %s as same from ("
    "select coalesce(voucher, '') || '|' || item_name || '|' || row_hash as k from stock_movements "
    "where move_date between %s and %s and item_name = any(%s::text[])) as s"
)
Q_CATALOG = "select item_code, is_active, coalesce(hidden, false) as hidden, updated_by from catalog_items"
Q_BOOK = "select sku_code from v_price_list_by_book where price_book = 'MA_base'"
Q_CANDIDATES = (
    "select distinct sku_code from selling_prices where price_book = 'MA_base' and coalesce(rate_bhd, 0) > 0 "
    "and sku_code is not null and trim(sku_code) <> ''"
)
# the days production holds a stock or a receivables snapshot for, after a day, each with its kind: a load made
# by hand leaves no journal, only its data (held_in_a_row)
Q_SNAPSHOT_DAYS = ("select distinct 'stock_balance' as k, as_of_date::text as d from stock_balance where as_of_date > "
                   "%s::date union select distinct 'ar_ageing' as k, as_of_date::text as d from ar_ageing "
                   "where as_of_date > %s::date")


class Stop(Exception):
    """The check cannot go on. The message is the sentence the owner reads; `flag` is one more word
    for the result's flags (for a caller that reads the JSON)."""

    def __init__(self, sentence: str, flag: str | None = None):
        super().__init__(sentence)
        self.flag = flag


class Interrupted(Stop):
    """An earlier load was cut off: nothing new is checked, and its own go line (kept in its plan) is
    what resumes it. `run` is that load's run."""

    def __init__(self, sentence: str, go_line: str | None, run: str = ""):
        super().__init__(sentence, "interrupted")
        self.go_line, self.run = go_line, run


class Deps:
    """Every contact of the check with the outside: git, the live site, the two database channels,
    the clock. Tests replace them; nothing else in this module reaches beyond the folders."""

    quiet = False

    def say(self, line: str) -> None:
        print(line, file=sys.stderr if self.quiet else sys.stdout, flush=True)

    def now(self) -> datetime:
        return datetime.now(status.BAHRAIN)

    def clock(self) -> float:
        return time.monotonic()

    def pause(self, seconds: float) -> None:
        time.sleep(seconds)

    def linked_worktree(self) -> bool:
        return gitinfo.is_linked_worktree(paths.CODE_ROOT)

    def live_commit(self) -> str:
        return str(db.health(config.get("api.health_url")).get("commit") or "")

    def site_state(self, live: str) -> dict:
        return status.site_state(paths.CODE_ROOT, live=live)

    def untracked(self) -> list[str]:
        return gitinfo.untracked(paths.CODE_ROOT, "scripts")

    def addresses(self) -> tuple[str, str]:
        import app.config  # noqa: F401  (loads the project folder's .env; the side channels stay blank)
        return os.environ.get("SUPABASE_URL", ""), os.environ.get("DATABASE_URL", "")

    def lock_state(self) -> dict:
        return status.lock_state()

    def failed_load(self) -> dict | None:
        self.unread_journals: list[str] = []         # runs whose journal could not be read (failed_load_state)
        return status.failed_load(self.unread_journals)

    def read_params(self, sql: str, params: list) -> list[dict]:
        return db.rpc_params(sql, params)

    def session(self):
        return db.session(with_backend=True, timeout_s=SESSION_TIMEOUT_S)

    def preview(self, folder: Path, backend, out_dir: Path, today: str | None, row_hook) -> dict:
        from app.ingest_batch import run_preview
        return run_preview(folder, backend, actor="ops_load_check", persist=False, out_dir=out_dir,
                           today=today, row_hook=row_hook)


# ── small helpers ─────────────────────────────────────────────────────────────

def dec(x) -> Decimal:
    """Exact Decimal of a parsed cell (blank, NaN and text read as 0)."""
    if x is None:
        return Decimal(0)
    if isinstance(x, float) and x != x:
        return Decimal(0)
    return fmt.dec(x)


def iso(x) -> str:
    return str(x or "")[:10]


def plain(x) -> str:
    """A quantity as plain digits: '1250' for a whole number, '12.5' otherwise."""
    d = dec(x)
    return str(int(d)) if d == d.to_integral_value() else format(d.normalize(), "f")


def span_text(lo, hi) -> str:
    """'Sun 21 Sep 2025 to Tue 06 Oct 2026', or without the years when both days are in one year."""
    lo, hi = iso(lo), iso(hi)
    if not lo or not hi:
        return "no dates"
    if lo[:4] != hi[:4]:
        return f"{fmt.day(lo)} {lo[:4]} to {fmt.day(hi)} {hi[:4]}"
    return f"{fmt.day(lo)} to {fmt.day(hi)}"


def short_name(name: str, width: int = 46) -> str:
    """A long Focus file name with its middle cut out, so the table stays readable."""
    if len(name) <= width:
        return name
    keep = width - 2
    head = keep * 2 // 3
    return name[:head] + ".." + name[-(keep - head):]


def listed(names, limit: int = 6) -> str:
    names = [str(n) for n in names]
    more = len(names) - limit
    return ", ".join(names[:limit]) + (f" and {more} more" if more > 0 else "")


def plural(n: int, word: str, many: str | None = None) -> str:
    return f"{n:,} {word if n == 1 else (many or word + 's')}"


class Cut:
    """A line of the Checks whose list is cut where it is shown: at CHAT_LIST_MAX in the chat's block, at
    LIST_MAX in check.txt. lines(limit) gives its line or lines at that length."""

    def __init__(self, lines):
        self.lines = lines


def shown_checks(checks: list, limit: int) -> list[str]:
    out: list[str] = []
    for line in checks:
        out += line.lines(limit) if isinstance(line, Cut) else [line]
    return out


def plain_error(exc: BaseException) -> str:
    """An error in words that are safe to print: a database or network error is named by its kind
    only (its message can spell the database address), anything else keeps its first words."""
    kind = type(exc).__name__
    if type(exc).__module__.split(".")[0] in ("psycopg", "psycopg_pool", "httpx", "httpcore", "postgrest", "socket"):
        return kind
    text = " ".join(str(exc).split())[:200]
    return f"{kind}: {text}" if text else kind


def allowed_output(path: Path) -> Path:
    """Refuse any place outside business_data/ and exports/ of the data root (the one rule of
    scripts/ops/paths.py for everything this package writes, copies, moves or removes)."""
    path = Path(path)
    if not paths.allowed_output(path):
        raise PermissionError(f"scripts/ops writes only under business_data/ and exports/: {path}")
    return path


def drop_folder(folder: Path) -> None:
    """Remove a folder of this package's own copies (loads/<run>/preview_in/): its files, then the folder."""
    if not folder.is_dir():
        return
    for path in folder.iterdir():
        paths.remove_file(path)
    try:
        allowed_output(folder).rmdir()
    except OSError:
        pass                                     # an empty folder left in the run's folder harms nothing


def copy_checked(src: Path, dst: Path) -> str:
    """Copy one report into staging and prove the copy: the file's sha256 before the copy, after it
    and of the copy are one value."""
    allowed_output(dst)
    try:
        before = loadplan.sha256_file(src)
        paths.copy_file(src, dst)
        same = before == loadplan.sha256_file(src) == loadplan.sha256_file(dst)
    except OSError as exc:
        raise Stop(f"{src.name} could not be copied into staging ({type(exc).__name__}: {exc.strerror or 'no detail'}). "
                   "Is the file open, or its path too long for Windows? Nothing was changed.") from None
    if not same:
        raise Stop(f"{src.name}: the copy is not identical to the file (was it saved again during the copy?). "
                   "Run the check again.")
    return before


def marker_of(name: str) -> str | None:
    """The loader's marker a file NAME reads as, by the loader's own rule (the text before the
    first dot, any case), else None. A report staged under such a name would switch one of the
    loader's guards off without the owner's word."""
    head = str(name).split(".")[0].upper()
    return head if head in MARKER_OPTIONS else None


def same_folder(a, b) -> bool:
    try:
        return Path(str(a)).resolve() == Path(str(b)).resolve()
    except OSError:
        return str(a) == str(b)


def lock_now(deps: Deps | None = None) -> dict:
    """The load lock as scripts.ops.status reads it (through `deps` inside a check). The worker
    replaces the file every 10 s; a file that still cannot be read after the reader's own tries
    answers {'state': 'unreadable'} instead of ending the command."""
    try:
        return deps.lock_state() if deps is not None else status.lock_state()
    except OSError as exc:
        return {"state": "unreadable", "error": type(exc).__name__}


# ── step 1: preflight ─────────────────────────────────────────────────────────

def project_ref(address: str) -> str | None:
    """The project reference inside an API address or a database address. Never printed."""
    a = (address or "").strip()
    m = re.match(r"^https?://([a-z0-9]{12,40})\.supabase\.[a-z]+", a, re.I)
    if m:
        return m.group(1).lower()
    m = re.match(r"^postgres(?:ql)?://(.+)$", a, re.I)
    if not m:
        return None
    userinfo, _, hostpart = m.group(1).rpartition("@")
    user = userinfo.split(":", 1)[0]
    host = re.split(r"[:/?]", hostpart, maxsplit=1)[0].lower()
    if "." in user:                                   # the pooler login is <role>.<project>
        return user.rsplit(".", 1)[1].lower() or None
    m = re.match(r"^db\.([a-z0-9]{12,40})\.supabase\.", host)
    return m.group(1) if m else None


def same_project(api_address: str, db_address: str) -> bool:
    a, b = project_ref(api_address), project_ref(db_address)
    return bool(a) and a == b


def preflight(dev: bool, deps: Deps) -> dict:
    """Raises Stop with the one sentence to act on; returns {live, line, notes}. `live` is the commit
    exactly as /health reports it (it is one part of the token; empty in a dev run)."""
    out: dict = {"live": "", "notes": []}
    parts: list[str] = []
    if dev:
        parts.append("dev run (worktree and live-site checks skipped)")
    else:
        if deps.linked_worktree():
            raise Stop("This is a code worktree (.claude/worktrees): a load never runs from here. Open the "
                       "project folder itself and type /yq-load there.")
        try:
            live = deps.live_commit()
        except Exception as exc:  # noqa: BLE001
            raise Stop("The live site did not answer, so the programs cannot be compared with it "
                       f"({type(exc).__name__}). Try again in a minute.") from None
        site = deps.site_state(live)
        state = site.get("state", "unknown")
        if state != "same":
            text = status.SITE_TEXT.get(state, status.SITE_TEXT["unknown"])
            if state == "changed_here":
                text = ("these program files were changed on this PC: " + listed(site.get("paths") or [], 8)
                        + ": tell Claude")
            elif state == "release_not_live" and site.get("paths"):
                text += " (" + listed(site["paths"], 5) + ")"
            elif state == "unknown" and site.get("error"):
                text += " (" + str(site["error"])[:120] + ")"
            raise Stop(text[0].upper() + text[1:] + ".")
        out["live"] = live
        parts.append(f"same programs as the live site ({live[:8]})")
        extra = [p for p in deps.untracked() if not p.startswith("scripts/ops/")]
        if extra:
            out["notes"].append("Files under scripts/ that git does not track (they do not stop a load): "
                                + listed(extra, 6) + ".")
    api_address, db_address = deps.addresses()
    if not same_project(api_address, db_address):
        raise Stop("The two database addresses in .env (SUPABASE_URL and DATABASE_URL) do not name the same "
                   "project, or one of them is missing. Nothing was read. Tell Claude.")
    parts.append("one database project")
    lock = lock_now(deps)
    if lock.get("state") == "interrupted" and lock.get("run") == "?":
        deps.pause(LOCK_AGAIN_S)                 # a lock that is being written at this very moment reads like this
        lock = lock_now(deps)
    state, run = lock.get("state"), str(lock.get("run") or "")
    left = lock.get("rehearsal_left")
    if state == "unreadable":
        raise Stop(f"The load lock could not be read ({lock.get('error')}), so it cannot be told whether a load is "
                   "running. Nothing was checked: try again in a moment.")
    away = lock.get("elsewhere")
    if away and state in ("running", "interrupted"):
        # the folder syncs through OneDrive: a lock taken on another PC is settled on that PC, never here
        raise Stop(f"The load lock was taken on another PC ({away}) for run {run}, at {lock.get('step')}. Nothing is "
                   f"checked, resumed or cleared from this PC: finish that load on {away}. If that PC can no longer "
                   "run it, ask Claude to clear the lock by hand.")
    if state == "interrupted" and run.startswith(status.REHEARSAL_PREFIX):
        state, left = "none", run                # a rehearsal has no go line to hand over and nothing to resume
    if state == "interrupted" and run == "?":
        # the file is there and is not a lock: its launcher ended before it had written it, so no worker
        # was ever started under it. Nothing to resume; the writer removes it at the next go
        state = "none"
        out["notes"].append(NAMELESS_LOCK)
    if state == "interrupted":
        token = None
        try:
            token = loadplan.read_plan(run).get("token") if loadplan.is_run(run) else None
        except (OSError, ValueError, AttributeError):
            pass
        raise Interrupted(f"LOAD INTERRUPTED at {lock.get('step')} (run {run}): an earlier load was cut off. Nothing "
                          "new is checked until it is resumed"
                          + (": its own go line does that (the box will say RESUME)." if token else
                             f", and the go line of that run could not be read (business_data/ops/loads/{run}/go.cmd): "
                             "tell Claude."),
                          loadplan.go_line(run, token) if token else None, run)
    if state == "running":
        what = "A rehearsal" if run.startswith(status.REHEARSAL_PREFIX) else "A load"
        raise Stop(f"{what} is running now (run {run}, at {lock.get('step')}). Wait for it: "
                   "python -m scripts.ops load status --wait")
    parts.append("no load running")
    if left:
        out["notes"].append(REHEARSAL_LEFT)
    out["line"] = ", ".join(parts)
    return out


def failed_load_state(deps: Deps) -> tuple[dict | None, str | None]:
    """({run, step} of the last load when it was closed as FAILED after it had begun to write and nothing has
    been loaded since (scripts.ops.status.failed_load), else None; why that could not be told, else None).
    Asked only while no lock is held: a load that is running or was cut off is said by the lock. It never
    stops a check. A journal that cannot be read at this moment (OneDrive holds it, or it is damaged) is
    "could not tell", never "no failed load": the check then never ends "nothing new" (what_is_new)."""
    try:
        if lock_now(deps).get("state") != "none":
            return None, None
        failed = deps.failed_load()
    except (OSError, ValueError, AttributeError, TypeError) as exc:
        return None, f"the load journals could not be read ({type(exc).__name__})"
    unread = [str(r) for r in (getattr(deps, "unread_journals", None) or [])]
    failed = failed if isinstance(failed, dict) and failed.get("run") else None
    return failed, (f"the journal of run {listed(unread, 3)} could not be read" if unread else None)


def last_failed_load(deps: Deps) -> dict | None:
    """The failed load alone (failed_load_state): a file that cannot be read at this moment answers None."""
    return failed_load_state(deps)[0]


def failed_line(failed: dict, starts: bool) -> str:
    """The line a check opens with after such a load. `starts`: this check ends with a GO line."""
    head = f"The last load (run {failed.get('run')}) failed at {failed.get('step') or '?'} after it had begun to write: "
    return head + ("this check starts a new load that replaces it." if starts else
                   "the database may hold part of it until a new load replaces it.")


def go_cmd_line(run: str) -> str | None:
    """Where the run's go.cmd lies, for the owner to double-click when Claude may not run the GO
    line itself; None when the file is not there."""
    try:
        path = loadplan.run_dir(run) / "go.cmd"
        return f"If Claude may not start it: double-click {path}" if path.is_file() else None
    except (OSError, ValueError):
        return None


# ── step 2: the folder ────────────────────────────────────────────────────────

def where_it_goes(name: str) -> str:
    low = name.lower()
    if low.endswith(".xml"):
        return "not loaded by this command; MRN XML: it goes on the Orders page"
    if low.endswith(".pdf"):
        return "not loaded by this command; PO or supplier invoice: it goes on the Orders page"
    return "not loaded by this command; it stays in the folder"


def scan_folder(folder: Path) -> dict:
    """What lies directly in a folder: {reports: {kind: [paths]}, others: [{name, reason}],
    locks: [names], folders: [{name, path, files, kinds, reason}]}."""
    from scripts import ingest as ing
    reports: dict[str, list[Path]] = {}
    others: list[dict] = []
    locks: list[str] = []
    entries = sorted(folder.iterdir(), key=lambda q: q.name.lower())
    for p in entries:
        if not p.is_file() or p.name.lower() in SYSTEM_FILES:
            continue
        if p.name.startswith("~$"):
            locks.append(p.name)
            continue
        kind = ing.classify(p.name) if p.suffix.lower() in EXCEL else None
        if kind is None:
            others.append({"name": p.name, "reason": where_it_goes(p.name)})
        elif kind.startswith("skip:"):
            others.append({"name": p.name,
                           "reason": "not loaded by this command (" + kind[5:] + "); it stays in the folder"})
        elif kind == "categories":
            others.append({"name": p.name,
                           "reason": "not loaded by this command (item-group report); it stays in the folder"})
        elif kind not in KIND_ORDER:                     # a report the loader may know and this check does not
            others.append({"name": p.name, "reason": f"not loaded by this command (a report this check does not know: "
                                                     f"{kind}); tell Claude"})
        else:
            reports.setdefault(kind, []).append(p)
    folders: list[dict] = []
    for q in entries:
        if not q.is_dir():
            continue
        inner = [f for f in q.rglob("*") if f.is_file()]
        kinds = sorted({k for k in (ing.classify(f.name) for f in q.iterdir()
                                    if f.is_file() and f.suffix.lower() in EXCEL and not f.name.startswith("~$"))
                        if k and not k.startswith("skip:") and k != "categories"})
        if kinds:
            reason = "not loaded by this command; another set of Focus reports"
        elif any(f.suffix.lower() in (".pdf", ".xml") for f in inner):
            reason = "not loaded by this command; PO, MRN XML or supplier invoice: these go on the Orders page"
        else:
            reason = "not loaded by this command; it stays in the folder"
        folders.append({"name": q.name, "path": q, "files": len(inner), "kinds": kinds, "reason": reason})
    return {"reports": reports, "others": others, "locks": locks, "folders": folders}


def title_rows(path: Path):
    import pandas as pd
    return pd.read_excel(path, header=None, dtype=object, engine="openpyxl", nrows=8)


def inventory(arg: str | None, deps: Deps) -> dict:
    """The drop folder as the check will use it. Raises Stop on anything that must be put right
    first. Returns {drop, day, files: {kind: path}, as_on: {kind: iso}, hashes: {kind: sha256},
    others, folders, note}."""
    from scripts import ingest as ing
    from scripts.verify_numbers import export_date

    drop = paths.find_drop(arg)
    if drop is None:
        raise Stop(f"No drop folder {arg!r} under Uplaods/." if arg else
                   "No dated drop folder (Uplaods/<ddmmyy>) was found.")
    found = scan_folder(drop)
    note = None
    if not found["reports"]:
        sets = [f for f in found["folders"] if f["kinds"]]
        if len(sets) > 1:
            raise Stop(f"No Focus report lies directly in {drop.name}, and more than one sub-folder holds reports ("
                       + listed([f["name"] for f in sets], 6) + "). Which set? Give the folder: "
                       "python -m scripts.ops load check \"<folder>\"")
        if not sets:
            raise Stop(f"No Focus report was found in {drop.name}.")
        note = f"The reports are in the sub-folder '{sets[0]['name']}' of {drop.name}: that set is checked."
        drop = sets[0]["path"]
        found = scan_folder(drop)
    if found["locks"]:
        raise Stop("A workbook of this folder is open in Excel (" + listed(found["locks"], 4) + "). Close it, "
                   "then run the check again.")
    twice = {k: v for k, v in found["reports"].items() if len(v) > 1}
    if twice:
        raise Stop("; ".join(f"{LABEL.get(k, k)}: {len(v)} files ({listed([p.name for p in v], 4)})"
                             for k, v in twice.items()) + ". Keep one file per report in the folder.")
    files = {k: v[0] for k, v in found["reports"].items()}
    for p in files.values():
        if marker_of(p.name):
            raise Stop(f"{p.name} is named like the loader's marker file {marker_of(p.name)}: staged under that name "
                       "it would switch one of the loader's guards off. Rename the file, then run the check again.")

    sizes = {p: (p.stat().st_size, p.stat().st_mtime_ns) for p in files.values()}
    deps.pause(SIZE_PAUSE_S)
    for p, was in sizes.items():
        if (p.stat().st_size, p.stat().st_mtime_ns) != was:
            raise Stop(f"{p.name} is still being written (its size changed while it was read). Wait until the "
                       "export has finished, then run the check again.")
        if was[0] == 0:
            raise Stop(f"{p.name} is empty. Export it again from Focus.")
    hashes: dict[str, str] = {}                      # what each report holds NOW: compared again at the end
    for kind, p in files.items():
        try:
            hashes[kind] = loadplan.sha256_file(p)
        except OSError as exc:
            raise Stop(f"{p.name} could not be read ({type(exc).__name__}). Is it open in another program? "
                       "Close it, then run the check again.") from None

    as_on: dict[str, str] = {}
    for kind, p in files.items():
        if kind == "ledger_entries":                 # never loaded, so an unreadable date never stops a load
            try:
                as_on[kind] = ing.report_date_from_title(title_rows(p)) or ""
            except Exception:  # noqa: BLE001
                as_on[kind] = ""
            continue
        if p.suffix.lower() != ".xlsx":
            raise Stop(f"{p.name} is not an .xlsx export. Export it again from Focus as Excel.")
        try:
            top = title_rows(p)
        except Exception as exc:  # noqa: BLE001
            raise Stop(f"{p.name} cannot be opened as a workbook ({type(exc).__name__}). Export it again.") from None
        if kind.startswith("selling_prices"):
            as_on[kind] = export_date(p)              # a price book has no title block
            continue
        if not ing.sniff_focus(p):
            raise Stop(f"{p.name} does not look like a Focus export (no company title in its first rows). "
                       "Remove it from the folder or export it again.")
        day = ing.report_date_from_title(top)
        if not day:
            raise Stop(f"{p.name} has no \"As on date\" in its title block, so it cannot be dated. Export it again "
                       "from Focus.")
        as_on[kind] = day

    day = paths.drop_date(drop.name) or paths.drop_date(drop.parent.name)
    if day is None:
        dated = (sorted(v for k, v in as_on.items() if v and not k.startswith("selling_prices"))
                 or sorted(v for v in as_on.values() if v))
        day = date.fromisoformat(dated[-1]) if dated else deps.now().date()
    return {"drop": drop, "day": day, "files": files, "as_on": as_on, "hashes": hashes, "others": found["others"],
            "folders": found["folders"], "note": note}


# ── step 3: staging ───────────────────────────────────────────────────────────

def claim_run(day: date, deps: Deps) -> tuple[str, Path]:
    """(run, its new staging folder). The run is <ddmmyy>_<HHMM>: the folder's date and the minute
    (Bahrain time) the check was made. A name is used once and never guessed ahead: when this
    minute's name is taken (a second check of the folder inside one minute) the check waits for the
    next minute. The staging folder is made HERE, and making it is the claim: of two checks that
    start at the same instant only one can make it, and the other waits like any late one."""
    taken = ""
    for attempt in range(RUN_NAME_WAITS + 1):
        now = deps.now()
        run = day.strftime("%d%m%y") + "_" + now.strftime("%H%M")
        staging = allowed_output(loadplan.staging_path(run))
        if not loadplan.run_dir(run).exists():
            try:
                staging.mkdir(parents=True, exist_ok=False)
                return run, staging
            except FileExistsError:
                pass
        taken = run
        if attempt < RUN_NAME_WAITS:
            wait = 60 - now.second - now.microsecond / 1_000_000 + 0.05
            deps.say(f"Run name {run} is already taken by an earlier check: waiting {int(wait) + 1} s for the next "
                     "minute.")
            deps.pause(wait)
    raise Stop(f"Other checks of this folder took the run names up to {taken}, one every minute. Nothing was staged: "
               "wait a minute, then run the check again.")


def read_plan_of(run: str) -> dict | None:
    """The plan of a run, or None when it has none that can be read as that run's plan."""
    try:
        plan = loadplan.read_plan(run)
    except (OSError, ValueError):
        return None
    return plan if isinstance(plan, dict) and plan.get("run") == run else None


def earlier_runs(run: str) -> list[str]:
    """Every other run named for the same date (a staging folder, a run folder, or both)."""
    prefix = run[:7]                                   # 'ddmmyy_'
    names: set[str] = set()
    for base in (paths.staging_dir(), paths.loads_dir()):
        if base.is_dir():
            names |= {p.name for p in base.iterdir()
                      if p.is_dir() and p.name.startswith(prefix) and loadplan.is_run(p.name)}
    names.discard(run)
    return sorted(names)


def move_aside(run: str, now: datetime) -> Path | None:
    """Move the staging folder of a replaced check to business_data/staging/superseded/<run>/ (a
    folder already there keeps its place: the new one gets a time suffix). Returns the new place,
    None when the run has no staging folder. Raises OSError when Windows refuses the move (a file
    of the folder is open)."""
    src = loadplan.staging_path(run)
    if not src.is_dir():
        return None
    base = allowed_output(paths.staging_dir() / SUPERSEDED)
    base.mkdir(parents=True, exist_ok=True)
    dst, n = base / run, 0
    while dst.exists():
        n += 1
        dst = base / (run + "_" + now.strftime("%Y%m%d-%H%M%S") + (f"-{n}" if n > 1 else ""))
    os.replace(allowed_output(src), allowed_output(dst))
    return dst


def retire_earlier(new_run: str, drop: Path, deps: Deps) -> dict:
    """A new check of a folder replaces every earlier check of THAT folder that never started
    loading (loadplan.run_result == NOT_STARTED): the plan gains "superseded_by" (the writer refuses
    such a plan), go.cmd is removed and the staging folder moves to staging/superseded/<run>/ (the
    earlier token then fails by itself as well: the staged files are part of it).
    Never touched: a run that started, finished or failed; a run that holds the lock; a check of
    ANOTHER folder with the same date (its plan names its folder).
    A staging folder with no plan beside it belongs to a check that stopped before its table: no go
    line ever existed for it, so only its copies are put away, and only once it is older than
    LEFTOVER_AFTER_S (a newer one may be a check that is still running).
    Returns {retired: [runs replaced], no_go: [those of them that never had a go line: a dev run, a
    check that stopped], leftovers: [runs with no table, put away], stuck: [runs whose staging
    folder could not be moved]}. Raises Stop when a check that still holds a token could be neither
    marked nor moved: two go lines for one folder would then both work."""
    out: dict = {"retired": [], "no_go": [], "leftovers": [], "stuck": []}
    lock = lock_now(deps)
    if lock.get("state") == "unreadable":
        raise Stop(f"The load lock could not be read ({lock.get('error')}), so it cannot be told whether a load "
                   "has just started. Nothing was staged: try again in a moment.")
    held = str(lock.get("run") or "") if lock.get("state") != "none" else ""
    now = deps.now()
    for old in earlier_runs(new_run):
        try:
            started = loadplan.run_result(old) != "NOT_STARTED"
        except (OSError, ValueError):
            started = True                                 # a journal that cannot be read: never touched
        if old == held or started:
            continue
        plan = read_plan_of(old)
        if plan is None:
            staging = loadplan.staging_path(old)
            try:
                if staging.is_dir() and time.time() - staging.stat().st_mtime >= LEFTOVER_AFTER_S:
                    move_aside(old, now)
                    out["leftovers"].append(old)
            except OSError:
                pass                                       # in use: a later check puts it away
            continue
        if not same_folder(plan.get("drop") or "", drop):
            continue
        before = bool(plan.get("superseded_by"))           # replaced by an earlier check already
        marked = False
        if not before:
            try:
                loadplan.write_plan(dict(plan, superseded_by=new_run))
                marked = True
            except OSError:
                pass
        moved = stuck = False
        try:
            moved = move_aside(old, now) is not None
        except OSError:
            stuck = True                                   # its folder is still where a load would look for it
        if stuck and not (before or marked):
            if plan.get("token"):
                raise Stop(f"An earlier check of this folder (run {old}) could not be put away: its plan could not "
                           "be rewritten and its staging folder could not be moved. Close whatever holds a file "
                           f"of business_data/staging/{old}, then run the check again (two go lines for one "
                           "folder must never both work).")
            continue                                       # it never had a go line; a later check puts it away
        try:
            paths.remove_file(loadplan.run_dir(old) / "go.cmd")
        except OSError:
            pass                                           # its plan or its moved folder has ended the line already
        if marked or (moved and not before):
            out["retired"].append(old)                     # counted by the check that ended its go line
            if not plan.get("token"):
                out["no_go"].append(old)
        if stuck:
            out["stuck"].append(old)
    return out


def retired_line(found: dict) -> str | None:
    """What the check says about the earlier checks it replaced; None when there was none. A check
    that had a go line is named first (that line is dead now); one that never had one (a dev run, a
    check that stopped) only had copies to put away."""
    runs, quiet, left, stuck = (list(found.get(k) or []) for k in ("retired", "no_go", "leftovers", "stuck"))
    where = f"business_data/staging/{SUPERSEDED}"
    said: list[str] = []
    live = [r for r in runs if r not in quiet]
    moved = [r for r in live if r not in stuck]
    if live:
        line = (f"This check replaces {plural(len(live), 'earlier check')} of this folder ({listed(live, 6)}): "
                + ("its go line no longer works" if len(live) == 1 else "their go lines no longer work"))
        if moved:
            line += (", and the staged copies moved to " if len(moved) == len(live) else
                     f"; the staged copies of {listed(moved, 6)} moved to ") + where
        said.append(line + ".")
    away = sorted(r for r in quiet + left if r not in stuck)
    if away:
        said.append(f"The staged copies of {plural(len(away), 'earlier check')} that never had a go line "
                    f"({listed(away, 6)}) were put away " + ("there as well." if moved else f"in {where}."))
    if stuck:
        said.append(f"The staging folder of {listed(stuck, 6)} could not be moved (a file in it is open): it stays "
                    "where it is and can no longer be loaded.")
    return " ".join(said) or None


def book_stage_name(book: str, run: str) -> str:
    """The Focus stem first, then the run: the loader's snapshot rule only lets a book void the rows
    of a DIFFERENTLY named export of the same book."""
    return f"{BOOK_STEMS[book]} {run}.xlsx"


def write_marker(staging: Path, name: str, run: str, option: str) -> Path:
    return paths.write_text(staging / name, f"written by load check {run} on the owner's option {option}\n")


# ── step 4: stock balance ─────────────────────────────────────────────────────

def canon_wh(name, aliases: dict) -> str:
    n = str(name or "").strip()
    return str(aliases.get(n, n))


def stock_block(grid, main: str, aliases: dict) -> dict:
    """Where the main store's block sits in a stock-balance grid, by the parser's own rules: the
    header row, the block's first row (its name row), the total row that closes it, that row's units
    and value, and every warehouse the file names. Focus closes a block with "Sub Total" when the
    export holds several warehouses and with "Total" when it was filtered to one. Raises Stop when
    the block cannot be cut."""
    from scripts import ingest as ing
    hdr = ing.check_header(grid, "stock_balance")
    start = end = None
    sub_units = sub_value = None
    label = "Sub Total"
    warehouses: list[str] = []
    for i in range(hdr + 1, len(grid)):
        r = list(grid.iloc[i])
        name = ing.txt(r[0])
        if name is None:
            continue
        if ing.is_total(name):
            if start is not None and end is None:        # the first total row after the block's items
                end, label = i, " ".join(name.split())
                sub_units = ing.norm_num(r[1]) if len(r) > 1 else None
                sub_value = ing.norm_num(r[3]) if len(r) > 3 else None
            continue
        if all(ing.norm_num(r[j]) is None for j in (1, 2, 3) if len(r) > j):
            wh = canon_wh(ing._clean_wh(name), aliases)          # a warehouse name row
            warehouses.append(wh)
            if wh == main:
                if start is not None:
                    raise Stop(f"The stock balance names the main store ({main}) twice: the block cannot be told "
                               "apart. Export it again.")
                start = i
            elif start is not None and end is None:
                raise Stop(f"Inside the main store's block ({main}), before its total row, the row '{name[:60]}' has no "
                           "quantity, rate or value: the loader would read it as another warehouse. Tell Claude.")
    if start is None:
        raise Stop(f"The stock balance does not hold the main store ({main}): "
                   + (listed(warehouses, 4) or "no warehouse")
                   + ". Export the stock balance again with the main store in it.")
    if end is None:
        raise Stop(f"The main store's block ({main}) has no total row after its items: the file's layout changed. "
                   "Tell Claude.")
    return {"header": hdr, "start": start, "end": end, "sub_units": sub_units, "sub_value": sub_value,
            "total_label": label, "warehouses": warehouses}


def write_trimmed(src: Path, dst: Path, grid, block: dict) -> Path:
    """A copy of the workbook that keeps the title block, the header row and the rows of one block
    (its name row down to and including its Sub Total row). Rows are DELETED from a copy of the
    original, so every kept cell keeps its value and style exactly."""
    from openpyxl import load_workbook
    allowed_output(dst)
    hdr, start, end = block["header"], block["start"], block["end"]
    wb = load_workbook(src)
    try:
        ws = wb.worksheets[0]
        for i in (hdr, start, end):                    # the sheet must line up with what the parser read
            if str(ws.cell(row=i + 1, column=1).value or "").strip() != str(grid.iat[i, 0] or "").strip():
                raise Stop(f"{src.name}: the sheet does not line up with what the parser read (row {i + 1}). "
                           "Tell Claude.")
        for rng in list(ws.merged_cells.ranges):
            if rng.max_row > hdr + 1:
                ws.unmerge_cells(str(rng))
        last = ws.max_row
        if last > end + 1:
            ws.delete_rows(end + 2, last - (end + 1))
        if start > hdr + 1:
            ws.delete_rows(hdr + 2, start - hdr - 1)
        try:
            dst.parent.mkdir(parents=True, exist_ok=True)
            fd, tmp = tempfile.mkstemp(prefix="trim.", suffix=".tmp", dir=str(dst.parent))
            os.close(fd)
            try:
                wb.save(tmp)
                os.replace(tmp, dst)
            finally:
                paths.remove_file(tmp)                 # gone already once it was moved into place
        except OSError as exc:
            raise Stop(f"the trimmed copy of {src.name} could not be written ({type(exc).__name__}: "
                       f"{exc.strerror or 'no detail'}). Is its path too long for Windows?") from None
    finally:
        wb.close()
    return dst


def block_rows(rows: list[dict], main: str, aliases: dict) -> list[tuple]:
    return [(r.get("item_name"), dec(r.get("net_qty")), dec(r.get("selling_rate_bhd")), dec(r.get("total_value_bhd")))
            for r in rows if canon_wh(r.get("warehouse_name"), aliases) == main]


def prove_block(original_rows: list[dict], staged: Path, main: str, aliases: dict, block: dict) -> dict:
    """The staged stock balance re-read with the real parser: one warehouse, the same item rows as
    the block in the original, the same date, and units and value equal to Focus's own total row of
    that block. Returns {items, units, value, as_on, names, raw_warehouses, total_label, problems}."""
    from scripts import ingest as ing
    rows = ing.parse_stock_balance(ing.read_grid(staged), staged.name)
    problems: list[str] = []
    names = sorted({canon_wh(r.get("warehouse_name"), aliases) for r in rows})
    if names != [main]:
        problems.append("the copy holds " + (listed(names, 4) or "no warehouse") + ", not the main store alone")
    want, got = block_rows(original_rows, main, aliases), block_rows(rows, main, aliases)
    if want != got:
        problems.append(f"the copy's item rows differ from the block in the original ({len(got)} against {len(want)})")
    dates = sorted({iso(r.get("as_of_date")) for r in rows})
    orig_dates = sorted({iso(r.get("as_of_date")) for r in original_rows})
    if len(dates) != 1 or not dates[0] or dates != orig_dates:
        problems.append("the copy's as-on date is not the original's")
    units = sum((g[1] for g in got), Decimal(0))
    value = sum((g[3] for g in got), Decimal(0))
    label = block.get("total_label") or "Sub Total"
    if block.get("sub_units") is None and block.get("sub_value") is None:
        problems.append(f"Focus's {label} row of the block is empty")
    else:
        if dec(block.get("sub_units")) != units:
            problems.append(f"the rows add up to {fmt.units(units)} units, Focus's {label} says "
                            f"{fmt.units(dec(block.get('sub_units')))}")
        if fmt.q3(block.get("sub_value")) != fmt.q3(value):
            problems.append(f"the rows add up to BHD {fmt.bhd(value)}, Focus's {label} says "
                            f"{fmt.bhd(block.get('sub_value'))}")
    return {"items": len(got), "units": units, "value": value, "as_on": dates[0] if dates else None,
            "names": [g[0] for g in got], "problems": problems, "total_label": label,
            "raw_warehouses": sorted({str(r.get("warehouse_name") or "(unassigned)") for r in rows})}


# ── step 5: profitability ─────────────────────────────────────────────────────

def bridge_names(ledger_names, report_names, exceptions: dict | None = None) -> dict[str, str]:
    """ledger item name -> report item name, by the engine's one bridge (scripts/ops/items.py): the
    stock ledger spells an item as its name, a space and its code, so the report name is a prefix of
    it and the longest one wins. `exceptions` (config items.bridge_exceptions) answers first. A
    ledger name that stands for none of the report's items is left out."""
    from scripts.ops.items import Bridge
    known = {str(n) for n in report_names if n}
    bridge = Bridge(known, dict(exceptions or {}))
    out: dict[str, str] = {}
    for led in ledger_names:
        hit = bridge.get(led) if led else None
        if hit in known:
            out[str(led)] = hit
    return out


def is_sales_invoice(row: dict) -> bool:
    return str(row.get("voucher_type") or "").strip().lower() == "sales invoice"


def ledger_cost(report_rows: list[dict], ledger_rows: list[dict], as_on: str, exceptions: dict | None = None) -> dict:
    """R = the report's cost of sales; L = the ledger's issued value of the Sales Invoice rows of the
    items the report lists, up to the report's date. gap = (L - R) / L."""
    names = {r.get("item_name") for r in report_rows if r.get("item_name")}
    sold = [r for r in ledger_rows if is_sales_invoice(r) and iso(r.get("move_date")) <= as_on]
    bridge = bridge_names({r.get("item_name") for r in sold}, names, exceptions)
    ledger = sum((dec(r.get("issued_value_bhd")) for r in sold if r.get("item_name") in bridge), Decimal(0))
    report = sum((dec(r.get("cogs_bhd")) for r in report_rows), Decimal(0))
    gap = (ledger - report) / ledger if ledger > 0 else None
    return {"ledger_cost": ledger, "report_cost": report, "gap": gap, "items": len(names),
            "matched": len(set(bridge.values()))}


def blank_cost_share(report_rows: list[dict]) -> dict:
    """Net sales of the sold items whose cost is blank or zero, as a share of net sales."""
    net = sum((dec(r.get("net_amount_bhd")) for r in report_rows), Decimal(0))
    blank = [r for r in report_rows if dec(r.get("net_amount_bhd")) > 0 and dec(r.get("cogs_bhd")) == 0]
    blank_net = sum((dec(r.get("net_amount_bhd")) for r in blank), Decimal(0))
    return {"net": net, "blank_net": blank_net, "blank_items": len(blank),
            "share": (blank_net / net) if net > 0 else Decimal(0)}


def profit_verdict(report_rows: list[dict], ledger_rows: list[dict] | None, as_on: str, gap_max: Decimal,
                   blank_max: Decimal, exceptions: dict | None = None, no_ledger: str | None = None) -> dict:
    """{hold, why[...], gap, ledger_cost, report_cost, blank: {...}}. Without a stock ledger the cost
    check cannot be made, and that is a hold; `no_ledger` says why there is none when it is not
    simply absent from the drop."""
    out: dict = {"hold": False, "why": [], "gap": None, "blank": blank_cost_share(report_rows)}
    if ledger_rows is None:
        out["hold"] = True
        out["why"].append((no_ledger or "no stock ledger in this drop") + ", so its cost of sales cannot be checked")
    else:
        out.update(ledger_cost(report_rows, ledger_rows, as_on, exceptions))
        if out["gap"] is None:
            out["hold"] = True
            out["why"].append("the stock ledger shows no sales cost for the report's items")
        elif out["gap"] > gap_max:
            out["hold"] = True
            out["why"].append(f"its cost of sales is {fmt.pct(out['gap'], 1)} short of the stock ledger's "
                              f"(limit {fmt.pct(gap_max, 1)})")
    if out["blank"]["share"] > blank_max:
        out["hold"] = True
        out["why"].append(f"{fmt.pct(out['blank']['share'], 1)} of net sales is on items with no cost "
                          f"(limit {fmt.pct(blank_max, 1)})")
    return out


def cost_movement(new_rows: list[dict], old_rows: list[dict]) -> tuple[int, int]:
    """(items whose net sales moved against the older report, of those: items whose cost did not)."""
    old = {r.get("item_name"): r for r in old_rows}
    moved = still = 0
    half_fils = Decimal("0.0005")
    for r in new_rows:
        o = old.get(r.get("item_name")) or {}
        if abs(dec(r.get("net_amount_bhd")) - dec(o.get("net_amount_bhd"))) >= half_fils:
            moved += 1
            if abs(dec(r.get("cogs_bhd")) - dec(o.get("cogs_bhd"))) < half_fils:
                still += 1
    return moved, still


def previous_load_plan() -> dict | None:
    """The plan of the load before this check (state.json names its run), or None when that load was
    not made through /yq-load and so left no plan."""
    run = str((status.read_state().get("last_load") or {}).get("run") or "")
    if not loadplan.is_run(run):
        return None
    try:
        plan = loadplan.read_plan(run)
    except (OSError, ValueError):
        return None
    return plan if isinstance(plan, dict) else None


def profit_before(plan: dict | None) -> tuple[bool | None, str]:
    """(did the load before hold its profitability report back, that report's as-on date), read from
    that load's own plan. A load whose drop brought no report held nothing back. (None, '') when
    there is no plan to read."""
    if not plan:
        return None, ""
    rows = [v for v in plan.get("verdicts") or [] if isinstance(v, dict) and v.get("kind") == "product_profitability"]
    dates = dict(plan.get("as_on") or {})
    dates.update(plan.get("as_on_all") or {})
    return any(v.get("verdict") == "hold_back" for v in rows), iso(dates.get("product_profitability"))


ORDINALS = ("first", "second", "third", "fourth", "fifth", "sixth", "seventh", "eighth", "ninth", "tenth")


def ordinal(n) -> str:
    """1 -> 'first' ... 10 -> 'tenth', then '11th' and so on."""
    try:
        n = int(n)
    except (TypeError, ValueError):
        return "next"
    return ORDINALS[n - 1] if 1 <= n <= len(ORDINALS) else f"{n}th"


def _pct(x) -> float | None:
    """A share as a percentage with two decimals, for plan.json."""
    return None if x is None else float(round(Decimal(x) * 100, 2))


def _money(x) -> str | None:
    return None if x is None else str(fmt.q3(x))


def second_hold(hold: bool, as_on: str, anchors: dict, before: tuple[bool | None, str] = (None, "")) -> bool | None:
    """Is this the second load in a row that holds the profitability report back?
      None   the report is held and is not newer than what the load before saw: the same drop is
             being checked again after its load. Nothing new can be said, so the register is left
             exactly as it is (neither added to nor cleared).
      True   the report is held and the load before held its own report back too. `before`
             (profit_before) says so from that load's plan; with no plan (a load made by hand) the
             data dates answer: production's report is older than production's last snapshot.
      False  everything else, a report that is not held included."""
    last = max(iso(anchors.get("stock_as_of")), iso(anchors.get("ar_as_of")))
    held_before, seen = before
    if not hold:
        return False
    if (last and as_on <= last) or (seen and as_on <= seen):
        return None
    if held_before is not None:
        return bool(held_before)
    return bool(last and iso(anchors.get("profit_as_of")) < last)


def finished_held_loads() -> tuple[dict[str, bool], set[str]]:
    """(data day -> did the load of that day hold its profitability report back, the snapshot days these loads
    brought) for every load the writer FINISHED (a 'finished' row in its journal; a rehearsal never counts). A
    day loaded twice is one entry, and the later load answers. A load whose drop brought no report held nothing
    back (profit_before). The snapshot days are each load's day and the as-on dates of the stock balance and the
    receivables it loaded, which can be another day. A journal or a plan that cannot be read is passed over."""
    base = paths.loads_dir()
    found: dict[str, tuple[str, bool]] = {}
    brought: set[str] = set()
    for folder in (base.iterdir() if base.is_dir() else []):
        run = folder.name
        if not folder.is_dir() or not loadplan.is_run(run) or run.startswith(status.REHEARSAL_PREFIX):
            continue
        try:
            ends = [str(r.get("at") or "") for r in loadplan.journal_read(run)
                    if isinstance(r, dict) and r.get("event") == "finished"]
            plan = read_plan_of(run) if ends else None
            day = iso((plan or {}).get("data_date"))
            if not day:
                continue
            held, _day = profit_before(plan)
            dates = plan.get("as_on") if isinstance(plan.get("as_on"), dict) else {}
        except (OSError, ValueError, AttributeError, TypeError):
            continue                          # held by another program this moment, or broken: it says nothing
        brought |= {day} | {iso(dates.get(k)) for k in ("stock_balance", "receivables") if iso(dates.get(k))}
        if day not in found or max(ends) > found[day][0]:
            found[day] = (max(ends), bool(held))
    return {d: held for d, (_at, held) in found.items()}, brought


def held_in_a_row(day: str, kept: str, loads: dict[str, bool], snapshots, explained=()) -> int:
    """How many loads in a row hold the profitability report back, the drop as on `day` counted once:
    the loads after production's report date `kept`, newest first, for as long as each one held its
    report back. `loads`: day -> held, the writer's finished loads; `explained`: the snapshot days those
    loads brought (finished_held_loads); `snapshots`: the days production holds a snapshot for, per kind
    ({kind: [days]}). A snapshot day that no finished load explains is a load made by hand: production's
    report stayed where it was, so it reads as held. One such load leaves one snapshot of each kind, and
    their dates may differ, so they count once per load, not per kind: as many as the kind that has the
    most of them (its days give the places). Without a report date in production the snapshots say
    nothing (they would count every load ever made)."""
    days = {d: bool(h) for d, h in loads.items() if (not kept or d > kept) and d < day}
    if kept:
        known = set(explained) | set(loads)
        kinds = snapshots.values() if isinstance(snapshots, dict) else [snapshots]
        by_hand = max((sorted({d for d in ds if kept < d < day and d not in known}) for ds in kinds),
                      key=lambda ds: (len(ds), ds[-1] if ds else ""), default=[])
        for d in by_hand:
            days.setdefault(d, True)
    n = 1
    for d in sorted(days, reverse=True):
        if not days[d]:
            break
        n += 1
    return n


# ── step 6: price book ────────────────────────────────────────────────────────

def is_general(row: dict) -> bool:
    return not (str(row.get("warehouse_name") or "").strip() or str(row.get("customer_code") or "").strip()
                or str(row.get("customer_name") or "").strip())


def scope_of(row: dict) -> str:
    return " / ".join(s for s in (str(row.get("warehouse_name") or "").strip(),
                                  str(row.get("customer_name") or row.get("customer_code") or "").strip()) if s)


def general_price(rows: list[dict], sku: str, on: str) -> Decimal | None:
    """The general (no warehouse, no customer) price of an item in force on a day, by the rule of the
    live price list: authorised, above zero, started, not ended; the latest start wins, then the
    later row."""
    best: tuple | None = None
    for i, r in enumerate(rows):
        if str(r.get("sku_code") or "").strip() != sku or not is_general(r):
            continue
        rate, start, end = dec(r.get("rate_bhd")), iso(r.get("start_date")), iso(r.get("end_date"))
        if (r.get("status") or "") != "Authorized" or rate <= 0 or not start or start > on or (end and end < on):
            continue
        if best is None or (start, i) > best[0]:
            best = ((start, i), rate)
    return best[1] if best else None


def status_word(row: dict | None) -> str:
    """A price row's status as the book spells it ('' when blank)."""
    return str((row or {}).get("status") or "").strip()


def pricebook_diff(file_rows: list[dict], live_rows: list[dict], book: str, exported: str) -> dict:
    """The file against the live rows of its book, key by key (scripts.load_supabase.pricebook_key):
    price, start date (it is part of the key), end date and status; the file a row came from is
    ignored. The status counts because the live price list holds a row only while it is Authorized:
    a row that turned to another status at the same price is not "the same as the live book". A live
    row counts as dropped only when it came from a Focus export of this book (the loader never voids
    another source).
    Returns {file_keys, live_keys, items: [{sku, scope, start, end, rate, what, tags, ...}]}."""
    from scripts.load_supabase import focus_book_kind, pricebook_key
    in_file: dict[tuple, dict] = {}
    for r in file_rows:
        in_file[pricebook_key(r)] = r                    # the loader keeps the last row of a key
    live = {pricebook_key(r): r for r in live_rows}
    rows_now, rows_live = list(in_file.values()), list(live_rows)
    items: list[dict] = []

    def entry(key: tuple, row: dict, what: str, old: dict | None) -> dict:
        sku = key[0]
        start = key[4]
        general = general_price(rows_now, sku, max(exported, start) if start else exported)
        scoped = not is_general(row)
        rate = dec(row.get("rate_bhd"))
        tags: list[str] = []
        if what == "dropped":
            tags.append("dropped")
        elif scoped:
            tags.append("scoped")
        else:
            tags.append("general")
        if what != "dropped" and start and start > exported:
            tags.append("scheduled")
        if old is not None and what == "changed":
            if dec(old.get("rate_bhd")) != rate:
                tags.append("price")
            if iso(old.get("end_date")) != iso(row.get("end_date")):
                tags.append("end_date")
            if status_word(old) != status_word(row):
                tags.append("status")
        return {"sku": sku, "scope": scope_of(row), "scoped": scoped, "start": start, "end": iso(row.get("end_date")),
                "rate": rate, "what": what, "tags": tags, "general": general,
                "live_general": general_price(rows_live, sku, exported),
                "old_rate": dec(old.get("rate_bhd")) if old else None,
                "old_end": iso(old.get("end_date")) if old else None,
                "status": status_word(row), "old_status": status_word(old) if old else None,
                "below_general": bool(scoped and what != "dropped" and general is not None and rate < general),
                "warehouse": str(row.get("warehouse_name") or "").strip(),
                "customer": str(row.get("customer_name") or "").strip(), "invoices": None, "units": None}

    for key, row in in_file.items():
        old = live.get(key)
        if old is None:
            items.append(entry(key, row, "added", None))
        elif (dec(old.get("rate_bhd")) != dec(row.get("rate_bhd")) or iso(old.get("end_date")) != iso(row.get("end_date"))
              or status_word(old) != status_word(row)):
            items.append(entry(key, row, "changed", old))
    for key, row in live.items():
        if key not in in_file and focus_book_kind(row.get("source_file")) == book:
            items.append(entry(key, row, "dropped", row))
    items.sort(key=lambda e: (e["sku"], e["scope"], e["start"]))
    return {"file_keys": len(in_file), "live_keys": len(live), "items": items}


def scope_sales(item: dict, deps: Deps) -> None:
    """How many invoices and units were already sold from the row's scope at its price since its
    start date (v_sales through the read RPC). A day-book rate is rounded to two decimals, so the
    line is matched on its gross = quantity x price."""
    sql, params = Q_SCOPE_SALES, [item["sku"], item["start"], str(item["rate"])]
    if item["warehouse"]:
        params.append(item["warehouse"])
        sql += f" and salesman_raw = ${len(params)}"
    if item["customer"]:
        params.append(item["customer"])
        sql += f" and coalesce(customer_name, customer_account) = ${len(params)}"
    rows = deps.read_params(sql, params)
    if rows:
        item["invoices"] = int(rows[0].get("invoices") or 0)
        item["units"] = dec(rows[0].get("units"))


def price_line(e: dict) -> str:
    """One key of the decision list, in plain words."""
    where = f"scoped to {e['scope']}" if e["scoped"] else "general price"
    since = f"from {fmt.day(e['start'])}" if e["start"] else "no start date"
    if e["what"] == "dropped":
        return f"{e['sku']}: {where} {fmt.bhd(e['rate'])} {since} is no longer in the book (it would be switched off)"
    bits: list[str] = []
    if e["what"] == "added":
        bits.append("new row")
    if "price" in e["tags"]:
        bits.append(f"was {fmt.bhd(e['old_rate'])}")
    if "end_date" in e["tags"]:
        bits.append(f"end date {fmt.day(e['old_end']) if e['old_end'] else 'none'} -> "
                    f"{fmt.day(e['end']) if e['end'] else 'none'}")
    if "status" in e["tags"]:
        was, now = e["old_status"] or "blank", e["status"] or "blank"
        bits.append(f"status {was} -> {now}" + (": it would no longer be in force" if was == "Authorized" else
                                                 ": it would come into force" if now == "Authorized" else ""))
    if "scheduled" in e["tags"]:
        bits.append("starts after the export day (a scheduled price)")
    if not e["scoped"] and e.get("live_general") is not None and e["live_general"] != e["rate"]:
        bits.append(f"the general price live now is {fmt.bhd(e['live_general'])}")
    if e["scoped"]:
        if e["general"] is None:
            bits.append("no general price in force")
        else:
            bits.append(("BELOW" if e["below_general"] else "against") + f" the general price {fmt.bhd(e['general'])}")
        if e["invoices"] is not None:
            bits.append(f"already sold there at this price: {plural(e['invoices'], 'invoice')}, "
                        f"{fmt.units(e['units'])} units")
    return f"{e['sku']}: {where} {fmt.bhd(e['rate'])} {since} ({'; '.join(bits)})"


def price_scope_items(diff_items: list[dict]) -> list[dict]:
    """Register items: a scoped row below the general price. What was already sold at that price is
    said only when it was looked up: a lookup that failed or was not made is "not looked up", never
    "0 units", and the item then carries no amount."""
    found: list[dict] = []
    for e in diff_items:
        if not e["below_general"]:
            continue
        key = alerts.make_key("price_scope", e["sku"].replace("|", "/"), alerts.ref_hash(e["scope"]), e["start"])
        sold = e["units"]
        detail = (f"{e['sku']}: {fmt.bhd(e['rate'])} from {fmt.day(e['start'])} for {e['scope']}, the general price is "
                  f"{fmt.bhd(e['general'])}; "
                  + (f"{fmt.units(sold)} units already sold there at this price" if sold is not None else
                     "what was already sold there at this price was not looked up"))
        found.append({"key": key, "level": "CHECK",
                      "title": "Price for one van or customer is below the general price: to check",
                      "detail": safe_words(detail, f"{e['sku']}: a scoped price is below the general price"),
                      "who": None, "due": None,
                      "bhd": (e["general"] - e["rate"]) * sold if sold is not None else None, "evidence": []})
    return found


def safe_words(text: str, fallback: str) -> str:
    """A register line may carry an item or file name; if such a name happens to hold a word the
    wording guard refuses, the plain fallback is used instead."""
    try:
        return fmt.careful(text)
    except ValueError:
        return fallback


# ── step 7: the preview ───────────────────────────────────────────────────────

LABEL_BY_TARGET = {"ar_ageing": "Receivables", "stock_balance": "Stock balance",
                   "product_profitability": "Profitability", "selling_prices": "Price book",
                   "orders": "Sales register", "order_lines": "Sales day book", "stock_movements": "Stock ledger"}


def read_preview(summary: dict, exceptions: list[dict]) -> dict:
    """What the check takes from run_preview's summary, without rebuilding any of its comparisons:
      actions           {target: {inserted, updated, unchanged, removed}}
      invoices          {invoice: 'changed' | 'removed'} from the samples of orders and lines. Only an
                        invoice whose HEADER is gone from the sales register is 'removed'; one that
                        lost a line (an edit in Focus renumbers the lines, so the last line number
                        goes) is still in the export and reads 'changed'
      more_invoices     changed or removed ROWS beyond the samples (rows, not invoices)
      stops             blocking codes that end the check, as sentences
      infos             blocking codes that are information here
      ledger_bulk       the code that only says many stock ledger rows are replaced (step 8 decides)
      warehouse_shrinks, void_guard   the preview's view of two loader guards (decided with the
                                      loader's own guards, and with the owner's options)"""
    out: dict = {"actions": {}, "invoices": {}, "stops": [], "infos": [], "ledger_bulk": None, "more_invoices": 0,
                 "warehouse_shrinks": None, "void_guard": None}
    targets = summary.get("targets") or {}
    for target, t in targets.items():
        a = t.get("actions") or {}
        out["actions"][target] = {"inserted": int(a.get("inserted", 0)), "updated": int(a.get("updated", 0)),
                                  "unchanged": int(a.get("unchanged", 0)),
                                  # a price book voids its rows, every other table deletes them: either way removed
                                  "removed": int(a.get("deleted") or 0) + int(a.get("voided") or 0)}
    for target in ("orders", "order_lines"):
        t = targets.get(target) or {}
        for sample in ("updated_sample", "removed_sample"):
            what = "removed" if (sample == "removed_sample" and target == "orders") else "changed"
            for row in t.get(sample) or []:
                inv = str(row.get("invoice_no") or "").strip()
                if inv and out["invoices"].get(inv) != "removed":
                    out["invoices"][inv] = what
        a = out["actions"].get(target) or {}
        beyond = (a.get("updated", 0) + a.get("removed", 0)
                  - len(t.get("updated_sample") or []) - len(t.get("removed_sample") or []))
        out["more_invoices"] = max(out["more_invoices"], beyond)
    for e in exceptions:
        if e.get("severity") != "blocking":
            continue
        code, target, msg = e.get("code"), e.get("target"), str(e.get("message") or "")
        if code in HARD_CODES:
            out["stops"].append(f"the preview cannot vouch for this drop ({code}): {msg[:200]}")
        elif code == "rows_removed" and target == "stock_movements":
            out["ledger_bulk"] = int(e.get("count") or 0)
        elif code == "rows_removed" and target in ("orders", "order_lines"):
            pass                                     # the invoices are listed; the loader's guards decide
        elif code == "days_shrink":
            out["infos"].append(f"{int(e.get('count') or 0)} day(s) have fewer sales lines in the file than are loaded "
                                "(the invoices are listed above)")
        elif code == "rows_removed":
            out["infos"].append(f"{LABEL_BY_TARGET.get(target, target)}: {int(e.get('count') or 0)} row(s) of the "
                                "snapshot already loaded for that day are not in this file and go with the load")
        elif code == "warehouse_set_shrinks":
            out["warehouse_shrinks"] = [str(x) for x in (e.get("items") or [])]
        elif code == "snapshot_void_guard":
            out["void_guard"] = msg
        else:
            out["stops"].append(f"the preview raised {code}: {msg[:200]}")
    return out


def sales_guards(captured: dict, db_rows: dict, replace_ok: bool) -> dict[str, list[str]]:
    """The loader's own span and salesman guards (scripts.load_supabase, pure) per sales table, fed
    with the rows production holds inside the file's dates: {target: [why it would fire]}. When a
    guard fires the loader skips the replace and the load then fails its verification, so the check
    stops instead. `captured` = the rows the preview prepared; `db_rows` = {target: [{invoice_no,
    salesman | warehouse_name}]}."""
    from scripts.load_supabase import salesman_guard, span_guard
    fired: dict[str, list[str]] = {}
    for target, col in (("orders", "salesman"), ("order_lines", "warehouse_name")):
        rows, held = captured.get(target), db_rows.get(target)
        if not rows or held is None:
            continue
        in_file = {str(r.get("invoice_no") or "").strip() for r in rows} - {""}
        in_db = {str(r.get("invoice_no") or "").strip() for r in held} - {""}
        seen: set[str] = set()
        gone: list[dict] = []
        for r in held:
            inv = str(r.get("invoice_no") or "").strip()
            if inv and inv not in in_file and inv not in seen:
                seen.add(inv)
                gone.append(r)
        names = {str(r.get(col)).strip() for r in rows if str(r.get(col) or "").strip()}
        why_names, missing = salesman_guard(held, names, col, replace_ok)
        out: list[str] = []
        if span_guard(gone, len(in_db), replace_ok):
            out.append(f"{len(gone)} of the {len(in_db)} invoices loaded for these dates are not in the file (more "
                       "than the loader replaces by itself)")
        if why_names:
            out.append(f"{plural(len(missing), 'salesman', 'salesmen')} with invoices loaded for these dates "
                       f"{'is' if len(missing) == 1 else 'are'} absent from the file ({listed(sorted(missing), 4)})")
        if out:
            fired[target] = out
    return fired


# ── step 8: stock ledger, voucher by voucher ──────────────────────────────────

def ledger_groups(rows: list[dict], from_db: bool = False) -> dict[tuple, tuple]:
    """(voucher, item) -> its lines as a sorted tuple of (date, type, warehouse, to-warehouse,
    received, issued, how many). The running balance, the rates and the row hash are left out: Focus
    recalculates them on every export."""
    acc: dict[tuple, Counter] = {}
    for r in rows:
        if from_db:
            line = (iso(r.get("d")), r.get("vt") or "", r.get("w") or "", r.get("tw") or "",
                    dec(r.get("rq")), dec(r.get("iq")))
            n = int(r.get("n") or 1)
        else:
            line = (iso(r.get("move_date")), r.get("voucher_type") or "", r.get("warehouse_name") or "",
                    r.get("to_warehouse_name") or "", dec(r.get("received_qty")), dec(r.get("issued_qty")))
            n = 1
        acc.setdefault((r.get("voucher") or "", r.get("item_name") or ""), Counter())[line] += n
    return {k: tuple(sorted(c.items())) for k, c in acc.items()}


def ledger_compare(file_rows: list[dict], db_rows: list[dict], last_loaded: str | None) -> dict:
    """The file against production inside the file's own dates.
      missing   items, warehouses and voucher types production holds there and the file lacks (STOP,
                unless every loaded line behind them is a SIM item's: ledger_sim)
      changed, removed, backdated   (voucher, item) pairs; backdated = new, dated before production's
                                    last movement date
      new       pairs the file adds from that date on (the normal case)"""
    f, p = ledger_groups(file_rows), ledger_groups(db_rows, from_db=True)

    def facets(groups: dict) -> tuple[set, set, set]:
        items = {k[1] for k in groups if k[1]}
        whs = {line[0][2] for sig in groups.values() for line in sig if line[0][2]}
        types = {line[0][1] for sig in groups.values() for line in sig if line[0][1]}
        return items, whs, types

    fi, fw, ft = facets(f)
    pi, pw, pt = facets(p)
    changed = sorted(k for k in f if k in p and f[k] != p[k])
    removed = sorted(k for k in p if k not in f)
    fresh = [k for k in f if k not in p]
    backdated = sorted(k for k in fresh if last_loaded and min(line[0][0] for line in f[k]) < last_loaded)
    return {"missing_items": sorted(pi - fi), "missing_warehouses": sorted(pw - fw), "missing_types": sorted(pt - ft),
            "changed": changed, "removed": removed, "backdated": backdated, "new": len(fresh) - len(backdated),
            "file": f, "db": p, "items": len(fi), "warehouses": len(fw), "types": len(ft), "pairs": len(f)}


def ledger_digest(rows: list[dict]) -> tuple[int, str]:
    """The file's side of Q_LEDGER_SAME: its rows' count and the md5 of their keys (voucher | item | row_hash),
    sorted as collate "C" sorts them (by code point). No rows: (0, ""), what production's empty set gives."""
    keys = sorted(f"{r.get('voucher') or ''}|{r.get('item_name')}|{r.get('row_hash')}" for r in rows)
    if not keys:
        return 0, ""
    return len(keys), hashlib.md5("\n".join(keys).encode("utf-8")).hexdigest()


# What may leave the stock ledger without a stop (the owner, 18-Oct-2026: SIM never counts): the
# movements of SIM items, told by the item's own name in whole words: a starter pack, or Batelco and
# SIM together (the Sim group of Focus held exactly those eight items then). Not the repository's
# looser SIM rules: the substring rule of v_sales also takes salesmen's names such as Qasim (the
# warehouses are named after salesmen), and SIM as a word alone takes accessories that name the SIM
# they serve (a cable with a SIM ejector); a rule that decides what a load may delete must not.
# Since 18-Oct-2026 the SIM items are Service items in Focus, which keeps them out of its stock
# reports altogether: every Stock ledger lacks them from then on.
SIM_STARTER = re.compile(r"\bstarter\s*pack", re.I)
SIM_CARRIER, SIM_WORD = re.compile(r"\bbatelco\b", re.I), re.compile(r"\bsim\b", re.I)


def sim_item(name) -> bool:
    text = str(name or "")
    return bool(SIM_STARTER.search(text) or (SIM_CARRIER.search(text) and SIM_WORD.search(text)))


def ledger_sim(cmp: dict) -> dict:
    """ledger_compare's missing items, warehouses and voucher types, each one SIM when every loaded
    line behind it is a SIM item's (a warehouse that held anything else is not): {"sim": {items,
    warehouses, types}, "other": {...}, "rows": the loaded rows of those SIM lines, which a load of
    the file takes away}."""
    wanted = {"items": set(cmp["missing_items"]), "warehouses": set(cmp["missing_warehouses"]),
              "types": set(cmp["missing_types"])}
    plain: dict[str, set] = {k: set() for k in wanted}           # the ones with a line that is not SIM
    rows = 0
    for (_voucher, item), sig in cmp["db"].items():
        sim = sim_item(item)
        for line, n in sig:
            behind = [(k, key) for k, key in (("items", item), ("warehouses", line[2]), ("types", line[1]))
                      if key in wanted[k]]
            if not behind:
                continue
            if sim:
                rows += n
            else:
                for k, key in behind:
                    plain[k].add(key)
    return {"sim": {k: sorted(wanted[k] - plain[k]) for k in wanted}, "other": {k: sorted(plain[k]) for k in wanted},
            "rows": rows}


SIM_PARTS = (("items", "SIM item", "SIM items", 12),
             ("warehouses", "warehouse with only SIM items", "warehouses with only SIM items", 60),
             ("types", "voucher type with only SIM items", "voucher types with only SIM items", 12))


def sim_words(sim: dict, named: bool = False) -> str:
    """'1 SIM item and 2 warehouses with only SIM items' for ledger_sim's SIM part ('' when none);
    named: each part with its names, every warehouse shown (a warehouse is a rep's van)."""
    parts = [plural(len(sim[k]), one, many) + (f" ({listed(sim[k], limit)})" if named else "")
             for k, one, many, limit in SIM_PARTS if sim[k]]
    if named:
        return "; ".join(parts)
    return ", ".join(parts[:-1]) + " and " + parts[-1] if len(parts) > 1 else "".join(parts)


def drop_sim_pairs(cmp: dict) -> tuple[dict, int]:
    """ledger_compare's changed, removed and earlier-dated vouchers without the (voucher, item) pairs
    of SIM items. SIM never counts: they are never listed to check, whether an export left them out
    or brought them back. -> (cmp, how many went)"""
    out, dropped = dict(cmp), 0
    for kind in ("changed", "removed", "backdated"):
        out[kind] = [key for key in cmp[kind] if not sim_item(key[1])]
        dropped += len(cmp[kind]) - len(out[kind])
    return out, dropped


def lines_text(sig: tuple | None) -> str:
    if not sig:
        return "nothing"
    parts: list[str] = []
    for (day, _vt, w, tw, rq, iq), n in sig[:3]:
        what = f"received {fmt.units(rq)}" if rq else (f"issued {fmt.units(iq)}" if iq else "no quantity")
        parts.append(f"{what} on {fmt.day(day)}" + (f" at {w}" if w else "") + (f" to {tw}" if tw else "")
                     + (f" x{n}" if n > 1 else ""))
    return "; ".join(parts) + (f"; and {len(sig) - 3} more lines" if len(sig) > 3 else "")


def voucher_items(cmp: dict) -> list[dict]:
    """Register items for the vouchers that changed, went or arrived with an earlier date."""
    found: list[dict] = []
    for kind, title in (("changed", "Stock voucher changed in Focus after it was loaded: to check"),
                        ("removed", "Stock voucher no longer in Focus's ledger: to check"),
                        ("backdated", "Stock voucher entered with an earlier date: to check")):
        for voucher, item in cmp[kind]:
            was, now = cmp["db"].get((voucher, item)), cmp["file"].get((voucher, item))
            detail = f"{voucher or '(no number)'}, {item}: loaded as {lines_text(was)}; the file has {lines_text(now)}"
            key = alerts.make_key("voucher_changed", (voucher or "(no number)").replace("|", "/"), item.replace("|", "/"))
            found.append({"key": key, "level": "CHECK", "title": title,
                          "detail": safe_words(detail, f"{voucher or '(no number)'}: differs from what was loaded"),
                          "who": None, "due": None, "bhd": None, "evidence": [voucher] if voucher else []})
    return found


# ── step 9: catalog ───────────────────────────────────────────────────────────

def catalog_effect(catalog_rows: list[dict], book_after: set[str], candidates: set[str]) -> dict:
    """What app.catalog.sync_from_price_book would do after this load: active items outside the
    price list go off, inactive items inside it come back on, items of the book the catalog does not
    know are added. {off, on, new}: codes with their marks."""
    existing = {str(r.get("item_code")): r for r in catalog_rows if r.get("item_code")}

    def mark(code: str) -> str:
        r = existing.get(code) or {}
        marks = []
        if r.get("hidden"):
            marks.append("hidden")
        if str(r.get("updated_by") or "").strip() not in ("", AUTO_SYNC):
            marks.append("last set by hand")
        return code + (" (" + ", ".join(marks) + ")" if marks else "")

    off = sorted(c for c, r in existing.items() if r.get("is_active") and book_after and c not in book_after)
    on = sorted(c for c, r in existing.items() if not r.get("is_active") and c in book_after)
    new = sorted(c for c in candidates if c and c not in existing)
    return {"off": [mark(c) for c in off], "on": [mark(c) for c in on], "new": new}


def catalog_line(effect: dict) -> str:
    if not (effect["off"] or effect["on"] or effect["new"]):
        return "Catalog: the sync that follows the load would switch nothing on or off."
    bits: list[str] = []
    if effect["off"]:
        bits.append(f"switch OFF {plural(len(effect['off']), 'item')}: {listed(effect['off'], 12)}")
    if effect["on"]:
        bits.append(f"switch ON {plural(len(effect['on']), 'item')}: {listed(effect['on'], 12)}")
    if effect["new"]:
        bits.append(f"add {plural(len(effect['new']), 'new item')}: {listed(effect['new'], 12)}")
    return "Catalog: the sync that follows the load would " + "; ".join(bits) + "."


# ── nothing new: a folded row ─────────────────────────────────────────────────

def key_of(row: dict, keys, defaults: dict | None = None) -> tuple:
    """A row's natural key as the loader matches it: each part as trimmed text (a blank part as the loader
    fills it, `defaults`), a date part as its day."""
    out: list[str] = []
    for k in keys:
        v = row.get(k)
        if v is None and defaults:
            v = defaults.get(k)
        text = "" if v is None else (v.isoformat() if isinstance(v, (date, datetime)) else str(v)).strip()
        out.append(text[:10] if k.endswith("_date") else text)
    return tuple(out)


def folded_keys(target: str, raw_rows: list[dict]) -> set[tuple]:
    """The keys of `target` that the report lists on more than one line: the loader folds each of them into
    one row (FOLDED). A table the loader never folds has none."""
    if target not in FOLDED:
        return set()
    from app.ingest_batch import TARGETS
    keys, defaults = TARGETS[target]["keys"], FOLD_DEFAULTS.get(target)
    seen = Counter(key_of(r, keys, defaults) for r in raw_rows if str(r.get(keys[0]) or "").strip())
    return {k for k, n in seen.items() if n > 1}


def same_cell(file_value, db_value, folded_figure: bool) -> bool:
    """Is one cell of a file row the same as production's, as the preview compares them (`is distinct
    from`)? Blank equals only blank; a day its day; text its text; a number its number exactly, and in a
    figure the loader adds up for a folded row (`folded_figure`), to within one millionth (FOLD_NOISE).
    Anything that is not a finite number on either side of a number is never the same (and never an error)."""
    if file_value is None or db_value is None:
        return file_value is None and db_value is None
    if isinstance(db_value, (date, datetime)):
        return iso(file_value) == db_value.isoformat()[:10]
    if isinstance(db_value, (Decimal, int, float)) and not isinstance(db_value, bool):
        if isinstance(file_value, bool):
            return False
        try:
            a, b = Decimal(str(file_value)), Decimal(str(db_value))
            if not (a.is_finite() and b.is_finite()):
                return False
            return abs(a - b) < FOLD_NOISE if folded_figure else a == b
        except (InvalidOperation, ValueError, ArithmeticError):
            return False
    return str(file_value) == str(db_value)


def noise_line(target: str, n: int) -> str:
    """The check line for the folded rows of a table that are no change (worded for any figure: an amount, a
    quantity or a rate)."""
    one = n == 1
    return (f"{LABEL_BY_TARGET.get(target, target)}: {plural(n, 'row')} the preview counts as rewritten "
            + ("is one" if one else "are rows") + " the loader adds up from two or more lines of "
            + FOLD_NOUN.get(target, "one key") + (", and it differs" if one else ", and they differ")
            + " from what is loaded only beyond the sixth decimal place (the loader keeps such a sum with "
              "floating-point rounding): not a change.")


def last_load_words() -> str:
    """' (last load <run> at <day> <HH:MM>)' from state.json, or '' when no load is on record here."""
    try:
        last = status.read_state().get("last_load") or {}
    except OSError:
        return ""
    run = str(last.get("run") or "") if isinstance(last, dict) else ""
    if not loadplan.is_run(run):
        return ""
    at = status.when(last.get("finished_at"))
    return f" (last load {run}" + (f" at {at.astimezone(status.BAHRAIN).strftime('%a %d %b %H:%M')}" if at else "") + ")"


# ── older than what is loaded; the register ───────────────────────────────────

def older_reason(mine: str, held: str, what: str) -> str:
    """Why a report older than what production holds stops, in the table's words. `what` names the
    thing production holds: 'this report' (a snapshot of the same kind) or 'a load'."""
    return (f"it is as on {fmt.day(mine)}; production already holds {what} as on {fmt.day(held)}. An older report "
            "is never loaded (it would put older figures back): export it again from Focus")


def ends_early_reason(noun: str, ends: str, held: str) -> str:
    """Why a report whose rows end before production's stops (`noun`: lines, invoices, movements).
    The loader only replaces the dates a file covers, so what production holds after them would
    stay as it is: either the file is not the newest export, or Focus no longer has those rows and
    somebody has to look."""
    return (f"its {noun} end on {fmt.day(ends)}; production already holds {noun} to {fmt.day(held)}. A report that "
            "ends earlier is never loaded: export it again from Focus, up to today (if Focus holds nothing after "
            f"{fmt.day(ends)}, tell Claude)")


def register_line(rule: str, counts: dict, still_open: int) -> str | None:
    """What one rule did on the action register, in a line; None when nothing moved. `still_open`:
    open items of an event rule that this check did not find again (they stay until he marks them)."""
    new, base, cleared, same = (int(counts.get(k) or 0) for k in ("new", "baseline", "cleared", "same"))
    if not (new or base or cleared):
        return None
    if rule in EVENT_RULES:
        line = f"Action list ({rule}): {new} added, {same} as before"
        if still_open:
            line += f", {still_open} from earlier checks still open"
    else:
        line = f"Action list ({rule}): {new} added, {cleared} cleared, {same} as before"
        if base:
            line += (f"; {base} counted once as already there (the first run of this rule: they are not listed "
                     "day by day)")
    if counts.get("explained"):
        line += f"; {int(counts['explained'])} you explained stay off the list"
    return line + "."


def rep_change(r: dict) -> str:
    """One salesman of the preview's comparison (the file against what is loaded, inside the file's
    dates), as lines / gross: by how much his sales change, then the two sides."""
    file_rows, db_rows = int(r.get("file_rows") or 0), int(r.get("db_rows") or 0)
    file_gross, db_gross = dec(r.get("file_gross")), dec(r.get("db_gross"))
    gross = file_gross - db_gross
    return (f"{r.get('key') or '(no salesman)'}: {file_rows - db_rows:+,} / {'+' if gross > 0 else ''}{fmt.bhd(gross)} "
            f"(file {file_rows:,} / {fmt.bhd(file_gross)}, loaded {db_rows:,} / {fmt.bhd(db_gross)})")


# ── the check ─────────────────────────────────────────────────────────────────

class Check:
    """One run of the check, step by step. Every step adds table rows (`rows`), check lines
    (`checks`) and, when the drop cannot be loaded as it is, the reasons (`stops`)."""

    def __init__(self, options: dict, deps: Deps):
        self.options, self.deps = options, deps
        self.rows: dict[str, dict] = {}              # report kind -> its line of the table
        self.checks: list[str] = []
        self.stops: list[str] = []
        self.flags: list[str] = []
        self.files: list[str] = []
        self.records: dict[str, list[dict]] = {}     # register rule -> what this check found
        self.guards: list[str] = []
        self.expected: dict = {}
        self.run: str | None = None
        self.inv: dict = {}
        self.pre: dict = {}
        self.anchors: dict = {}
        self.staged: dict[str, Path] = {}            # report kind -> its copy in staging
        self.staging: Path | None = None
        self.session = None
        self.ledger_rows: list[dict] | None = None
        self.ledger_stop: str | None = None          # why the stock ledger left the check early
        self.last_loaded = ""                        # the newest as-on date of production's snapshots
        self.fingerprint: dict | None = None         # the database before the first comparison
        self.before: tuple[bool | None, str] = (None, "")    # the load before: profit_before()
        self.stock: dict | None = None
        self.profit: dict | None = None
        self.profit_evidence: dict | None = None     # what the profitability step found (plan.json: profitability)
        self.books: dict[str, dict] = {}
        self.summary: dict = {}
        self.targets: dict = {}
        self.pv: dict = {"actions": {}, "invoices": {}, "stops": [], "infos": [], "more_invoices": 0}
        self.captured: dict[str, list] = {}
        self.today: str | None = None
        self.step = 0
        self.failed: dict | None = None              # the last load, when it failed after it began to write
        self.failed_unsure: str | None = None        # why that could not be told (failed_load_state)
        self.retired: dict = {}                      # the earlier checks of this folder this one replaced
        self.timings: dict[str, float] = {}          # seconds per part of the check (plan.json: timings)
        self.preview_parts: dict | None = None       # the preview's seconds per report and helper (querytime)
        self.ledger_cmp: dict | None = None          # step 8's counts: changed, removed, backdated, new
        self.ledger_same: bool | None = None         # step 8: production's rows are the file's (Q_LEDGER_SAME)
        self.said: list[str] = []                    # the progress lines, for check.txt
        self.catalog: dict | None = None             # step 9: what the catalog sync would switch
        self.noise: dict[str, int] = {}              # target -> its folded rows that are no change (folded_noise)
        self.new_in_drop: list[str] = []             # what makes this drop more than "nothing new" (what_is_new)
        from scripts.ops.gridmemo import GridMemo
        # one parse per report file (YQ_CHECK_NO_MEMO=1 parses every read again: a diagnostic switch)
        self.grids = GridMemo(memo=os.environ.get("YQ_CHECK_NO_MEMO") != "1")

    # ── small tools ───────────────────────────────────────────────────────────
    def say(self, line: str) -> None:
        """A progress line: printed as it comes, and kept for check.txt."""
        self.said.append(line)
        self.deps.say(line)

    def done(self, line: str) -> None:
        self.step += 1
        self.say(f"[{self.step}/{STEPS}] {line}")

    def row(self, kind: str, verdict: str, reason: str, staged: str | None = None) -> None:
        path = self.inv["files"].get(kind)
        self.rows[kind] = {"file": path.name if path else None, "kind": kind, "label": LABEL.get(kind, kind),
                           "verdict": verdict, "decision": WORDS[verdict], "reason": reason, "staged": staged}

    def flag(self, name: str) -> None:
        if name not in self.flags:
            self.flags.append(name)

    def reading(self, read):
        """One read on the session. A statement that fails ends the check: after an error the
        transaction is void, so nothing read later could be trusted."""
        try:
            return read()
        except db.ReadOnlyViolation:
            raise
        except Exception as exc:  # noqa: BLE001  (the message may name the database: never printed)
            raise Stop(f"A read of the database failed ({type(exc).__name__}). Nothing was changed. Run the check "
                       "again.") from None

    def q(self, sql: str, params=None) -> list[dict]:
        return self.reading(lambda: self.session.q(sql, params))

    def unstage(self, kind: str) -> None:
        path = self.staged.pop(kind, None)
        if path is not None:
            paths.remove_file(path)

    def stops_already(self, kind: str) -> bool:
        return (self.rows.get(kind) or {}).get("verdict") == "stop"

    def absent(self, kind: str) -> str:
        """Why a report takes no part in a comparison: not in the drop, or stopped on its own line."""
        name = LABEL.get(kind, kind).lower()
        return f"the {name} stops (see its line)" if kind in self.inv["files"] else f"no {name} in this drop"

    # ── the steps, in order ───────────────────────────────────────────────────
    def timed(self, name: str, fn, *args):
        """Run one part of the check and add its seconds to self.timings (plan.json: timings)."""
        t = time.perf_counter()
        try:
            return fn(*args)
        finally:
            self.timings[name] = round(self.timings.get(name, 0.0) + time.perf_counter() - t, 2)

    def go(self, drop_arg: str | None) -> dict:
        self.main = str(config.get("warehouses.main"))
        self.aliases = dict(config.get("warehouses.aliases"))
        self.gap_max = config.number("load.profit_gap_max")
        self.blank_max = config.number("load.blank_cost_share_max")
        self.bridge = dict(config.get("items.bridge_exceptions"))
        # said first, whatever the check finds; "could not tell" is kept apart from "none failed"
        self.failed, self.failed_unsure = self.timed("failed_load", failed_load_state, self.deps)
        self.timed("preflight", self.step_preflight)
        self.timed("folder", self.step_folder, drop_arg)
        self.timed("staging", self.step_staging)
        try:
            self.session = self.timed("connect", self.deps.session)
        except Exception as exc:  # noqa: BLE001  (the message may name the database address: never printed)
            raise Stop(f"The database could not be reached ({type(exc).__name__}). Nothing was changed. Try again "
                       "in a minute.") from None
        try:
            # the report parser is loaded by now (the folder step reads the title blocks): from here on each
            # file's grid is parsed once (gridmemo)
            with self.grids.active():
                for name, step in (("anchors", self.read_anchors), ("older", self.older_reports),
                                   ("ledger_read", self.read_ledger), ("stock_balance", self.step_stock_balance),
                                   ("profitability", self.step_profitability), ("price_books", self.step_price_books),
                                   ("preview", self.step_preview), ("ledger_compare", self.step_ledger),
                                   ("guards", self.loader_guards), ("catalog", self.step_catalog)):
                    self.timed(name, step)
                return self.finish()
        finally:
            try:
                self.session.close()
            except Exception:  # noqa: BLE001
                pass

    def step_preflight(self) -> None:
        self.pre = preflight(self.options["dev"], self.deps)
        self.checks.append("Preflight: " + self.pre["line"] + ".")
        self.checks += self.pre["notes"]
        self.done("Preflight: " + self.pre["line"] + ".")

    def step_folder(self, drop_arg: str | None) -> None:
        inv = self.inv = inventory(drop_arg, self.deps)
        titled = sorted({v for k, v in inv["as_on"].items() if v and k in DAILY})
        self.titled_day = titled[-1] if titled else None        # the newest as-on date of the daily reports
        self.report_day = (titled[-1] if titled else
                           (sorted(v for v in inv["as_on"].values() if v) or [self.deps.now().date().isoformat()])[-1])
        if inv["note"]:
            self.checks.append(inv["note"])
        if len(titled) > 1:
            self.flag("mixed_dates")
            self.checks.append("ATTENTION: the reports carry different as-on dates ("
                               + listed([fmt.day(d) for d in titled], 6) + "): are they one export?")
        if "ledger_entries" in inv["files"]:
            day = inv["as_on"]["ledger_entries"]
            self.row("ledger_entries", "not_loaded", "the Ledger export is never loaded by this command ("
                     + (f"as on {fmt.day(day)}" if day else "its as-on date could not be read")
                     + "); the month-end reports read it from this folder")
        self.done(f"Folder {inv['drop'].name}: {plural(len(inv['files']), 'Focus report')} as on "
                  f"{fmt.day(self.report_day)}"
                  + (f", {plural(len(inv['others']), 'other file')}" if inv["others"] else "")
                  + (f", {plural(len(inv['folders']), 'sub-folder')}" if inv["folders"] else "") + ".")

    def step_staging(self) -> None:
        """A new folder (its name is the run), the earlier checks of this folder put away, and the
        reports that load whatever the checks say. The stock balance, the profitability report and
        the price book are added by their own steps; the Ledger export never.
        A marker file of the loader is written here and nowhere else, each only on the owner's own
        option (finish() proves it again on the folder as it then is)."""
        from scripts.load_supabase import PARTIAL_OK_MARKER, REPLACE_OK_MARKER
        opt, files = self.options, self.inv["files"]
        run, staging = claim_run(self.inv["day"], self.deps)
        self.run, self.staging = run, staging
        self.files.append(str(staging))
        if not opt["dev"]:                           # a dev run adds its own folders and changes nothing else
            self.retired = retire_earlier(run, self.inv["drop"], self.deps)
            said = retired_line(self.retired)
            if said:
                self.checks.append(said)
                self.flag("replaced_earlier")
        for kind in ALWAYS_STAGED:
            if kind in files:
                copy_checked(files[kind], staging / files[kind].name)
                self.staged[kind] = staging / files[kind].name
        markers = []
        if opt["replace_ok"]:
            write_marker(staging, REPLACE_OK_MARKER, run, "--replace-ok")
            markers.append(REPLACE_OK_MARKER)
        if opt["partial_ok"]:
            write_marker(staging, PARTIAL_OK_MARKER, run, "--partial-ok")
            markers.append(PARTIAL_OK_MARKER)
        replaced = len(self.retired.get("retired") or [])
        self.done(f"Staging {run}: {plural(len(self.staged), 'report')} copied, each identical to its original"
                  + (("; marker " + " and ".join(markers) + " written on your word") if markers else "")
                  + (f"; {plural(replaced, 'earlier check')} of this folder replaced" if replaced else "") + ".")

    def marker_story(self) -> list[str]:
        """The loader's marker files in the staged folder against the owner's options: the two must
        tell one story (the writer compares them again before it shows its box). A file the loader
        would read as a marker nobody asked for, or an option whose marker is missing, is a stop."""
        from scripts.load_supabase import marker_present
        out: list[str] = []
        for marker, (key, word) in MARKER_OPTIONS.items():
            there, asked = marker_present(self.staging, marker), bool(self.options.get(key))
            if there and not asked:
                out.append(f"the staged folder holds a file the loader reads as its marker {marker}, and {word} was "
                           "not given: tell Claude")
            elif asked and not there:
                out.append(f"{word} was given, but its marker file {marker} is not in the staged folder: run the "
                           "check again")
        return out

    def read_anchors(self) -> None:
        """The database as it is before the first comparison (its fingerprint is read again when the
        token is made: see moved_meanwhile), then the dates production holds. A drop whose daily
        reports are ALL older than production's newest snapshot ends the check here: the loader
        replaces the dates a file covers, so an older export would take out every invoice and
        voucher entered after it, and comparing the two would only list those as 'removed'.
        Nothing is compared and nothing is recorded."""
        self.fingerprint = self.reading(lambda: loadplan.fingerprint(self.session))
        self.anchors = (self.q(Q_ANCHORS) or [{}])[0]
        self.today = iso(self.anchors.get("today")) or self.deps.now().date().isoformat()
        last = self.last_loaded = max(iso(self.anchors.get("stock_as_of")), iso(self.anchors.get("ar_as_of")))
        if last and self.titled_day and self.titled_day < last:
            raise Stop(f"These reports are as on {fmt.day(self.titled_day)}; production already holds a load as on "
                       f"{fmt.day(last)}. An older drop is never loaded: it would put older figures back for the "
                       "dates it covers. Check the newest folder instead, or export the reports again from Focus.",
                       OLDER)
        self.before = profit_before(previous_load_plan())

    def older_reports(self) -> None:
        """One older report in a current folder: its line stops, its copy leaves staging and its own
        step is skipped. A snapshot (stock balance, receivables, profitability) is older when
        production holds a newer one of its own kind. The sales reports and the stock ledger have no
        as-on date in the database: they are older when their title is dated before production's
        newest snapshot, which is the as-on date of the last load."""
        as_on, files = self.inv["as_on"], self.inv["files"]
        for kind in DAILY:
            mine = as_on.get(kind)
            if kind not in files or not mine:
                continue
            if kind in SNAPSHOT_HELD:
                held, what = iso(self.anchors.get(SNAPSHOT_HELD[kind])), "this report"
            else:
                held, what = self.last_loaded, "a load"
            if held and mine < held:
                self.row(kind, "stop", older_reason(mine, held, what))
                self.unstage(kind)
                self.flag(OLDER)
                if kind == "stock_movements":
                    self.ledger_stop = "older than what is loaded"

    def read_ledger(self) -> None:
        """The staged stock ledger, prepared exactly as the loader loads it (steps 5 and 8 read it).
        A ledger whose movements end before production's is not used: the loader would leave the
        later movements where they are, and the file is then not the newest export."""
        if "stock_movements" not in self.staged:
            return
        from app.ingest_batch import prepare
        from scripts import ingest as ing
        path = self.staged["stock_movements"]
        try:
            rows = prepare("stock_movements", ing.parse_stock(ing.read_grid(path), path.name))[0]
        except ing.HeaderMismatch as exc:
            self.row("stock_movements", "stop", "its header row is not the one the loader reads: " + str(exc)[:160])
            self.unstage("stock_movements")
            self.ledger_stop = "its header row changed"
            return
        ends = max((iso(r.get("move_date")) for r in rows if r.get("move_date")), default="")
        held = iso(self.anchors.get("moves_to"))
        if ends and held and ends < held:
            self.row("stock_movements", "stop", ends_early_reason("movements", ends, held))
            self.unstage("stock_movements")
            self.ledger_stop = "it ends before what is loaded"
            self.flag(OLDER)
            return
        self.ledger_rows = rows

    # 4 ── stock balance
    def step_stock_balance(self) -> None:
        from scripts import ingest as ing
        path = self.inv["files"].get("stock_balance")
        if path is None:
            self.done("Stock balance: not in this drop.")
            return
        if self.stops_already("stock_balance"):
            self.done("Stock balance: STOP (older than what is loaded).")
            return
        main, aliases, staging = self.main, self.aliases, self.staging
        try:
            grid = ing.read_grid(path)
            rows = ing.parse_stock_balance(grid, path.name)
            block = stock_block(grid, main, aliases)
            names = sorted(set(block["warehouses"]))
            trim = len(names) > 1
            if trim:
                dst = staging / (re.sub(r"[^A-Za-z0-9]+", "_", main).strip("_") + "_only_" + path.name)
                if ing.classify(dst.name) != "stock_balance":
                    raise Stop("the name of the trimmed copy would not be read as a stock balance")
                write_trimmed(path, dst, grid, block)
            else:
                dst = staging / path.name
                copy_checked(path, dst)
        except ing.HeaderMismatch as exc:
            self.row("stock_balance", "stop", "its header row is not the one the loader reads: " + str(exc)[:160])
            self.done("Stock balance: STOP (its header row changed).")
            return
        except Stop as stop:
            self.row("stock_balance", "stop", str(stop))
            self.done("Stock balance: STOP.")
            return
        proof = prove_block(rows, dst, main, aliases, block)
        gone: list[str] = []
        try:
            have = set(proof["names"])
            gone = sorted(str(r.get("item_name")) for r in self.deps.read_params(Q_SNAPSHOT, [main])
                          if r.get("item_name") not in have)
        except Exception as exc:  # noqa: BLE001
            proof["problems"].append(f"production's last snapshot could not be read ({type(exc).__name__})")
        if proof["problems"]:
            paths.remove_file(dst)
            hint = ""
            if str(proof["total_label"]).lower() not in CLOSING_ROWS:
                # the block was closed at a row that only reads as a total row because of its name (the parser
                # takes any name that holds such a word for one, and so would the loader)
                hint = (f". The block was closed at the row '{str(proof['total_label'])[:60]}', which is read as a "
                        "total row because of its name: if it is an item, the loader would leave it out. Tell Claude")
            self.row("stock_balance", "stop", "the staged copy could not be proven: " + "; ".join(proof["problems"])
                     + hint)
            self.done("Stock balance: STOP (the copy could not be proven).")
            return
        self.staged["stock_balance"] = dst
        figures = f"{plural(proof['items'], 'item')} / {fmt.units(proof['units'])} units / BHD {fmt.bhd(proof['value'])}"
        if trim:
            self.row("stock_balance", "trim", f"{len(names)} warehouses in the file; only {main} loads: {figures}",
                     dst.name)
            self.checks.append(f"Stock balance: the copy cut to {main} was read again with the loader's parser: "
                               f"{figures} as on {fmt.day(proof['as_on'])}, the same rows as that block in the "
                               f"original and equal to Focus's own {proof['total_label']} row.")
        else:
            self.row("stock_balance", "load", f"{main} only: {figures}", dst.name)
            self.checks.append(f"Stock balance: {figures} as on {fmt.day(proof['as_on'])}, equal to Focus's own "
                               f"{proof['total_label']} row.")
        self.checks.append("Stock balance: every item with stock in the last loaded snapshot is in this file." if not gone
                           else f"Stock balance: {plural(len(gone), 'item')} with stock in the last loaded snapshot "
                                f"{'is' if len(gone) == 1 else 'are'} gone from this file (sold out or moved): "
                                f"{listed(gone, 8)}.")
        proof.update(trim=trim, gone=gone)
        self.stock = proof
        self.done(f"Stock balance: {WORDS['trim' if trim else 'load']}: {figures}"
                  + (f"; {plural(len(gone), 'item')} gone since the last snapshot" if gone else "") + ".")

    # 5 ── profitability
    def step_profitability(self) -> None:
        from scripts import ingest as ing
        path = self.inv["files"].get("product_profitability")
        if path is None:
            self.done("Profitability: not in this drop.")
            return
        if self.stops_already("product_profitability"):
            self.done("Profitability: STOP (older than what is loaded).")
            return
        opt, anchors, as_on = self.options, self.anchors, self.inv["as_on"]["product_profitability"]
        try:
            rows = ing.parse_profitability(ing.read_grid(path), path.name)
        except ing.HeaderMismatch as exc:
            self.row("product_profitability", "stop", "its header row is not the one the loader reads: " + str(exc)[:160])
            self.done("Profitability: STOP (its header row changed).")
            return
        v = profit_verdict(rows, self.ledger_rows, as_on, self.gap_max, self.blank_max, self.bridge,
                           no_ledger=self.absent("stock_movements"))
        line = "Profitability: "
        if v.get("gap") is not None:
            line += (f"the stock ledger's cost of the invoiced items is BHD {fmt.bhd(v['ledger_cost'])}, the report's "
                     f"cost of sales BHD {fmt.bhd(v['report_cost'])}: gap {fmt.pct(v['gap'], 1)} "
                     f"(limit {fmt.pct(self.gap_max, 1)})")
        else:
            line += "the gap to the stock ledger could not be worked out"
        before = self.q(Q_PROFIT_BEFORE, [as_on])
        if before:
            day = iso(before[0].get("report_date"))
            if self.ledger_rows is not None:
                prev = ledger_cost(before, self.ledger_rows, day, self.bridge)
                if prev["gap"] is not None:
                    v["previous_gap"], v["previous_day"] = prev["gap"], day
                    line += (f"; the report already loaded (as on {fmt.day(day)}) shows {fmt.pct(prev['gap'], 1)} "
                             "against this ledger")
            v["moved"], v["still"] = cost_movement(rows, before)
            line += (f"; {v['still']} of the {v['moved']} items whose sales moved since {fmt.day(day)} show no cost "
                     "movement" if v["moved"] else f"; no item's sales moved since {fmt.day(day)}")
        line += (f". Items sold with no cost: {fmt.pct(v['blank']['share'], 1)} of net sales "
                 f"(limit {fmt.pct(self.blank_max, 1)}).")
        self.checks.append(line)
        held = v["held"] = bool(v["hold"] and not opt["with_profitability"])
        if held:
            since = iso(anchors.get("profit_as_of"))
            self.row("product_profitability", "hold_back", "; ".join(v["why"])
                     + (f"; production keeps the report as on {fmt.day(since)}" if since else ""))
            self.flag("hold_back")
        else:
            dst = self.staging / path.name
            copy_checked(path, dst)
            self.staged["product_profitability"] = dst
            if v["hold"]:
                reason = ("loaded on your word (--with-profitability) although the check says hold back: "
                          + "; ".join(v["why"]))
            elif v["gap"] < -self.gap_max:
                # the rule holds a report back when its cost is SHORT of the ledger's. One whose cost is far
                # above it is not that case and loads, but it does not "agree" either: the line says what it is
                reason = (f"its cost of sales is {fmt.pct(-v['gap'], 1)} above the stock ledger's, not short of it: to "
                          f"check (a report is held back only when its cost is more than {fmt.pct(self.gap_max, 1)} "
                          "short)")
                self.flag("profit_above_ledger")
            else:
                reason = (f"its cost of sales agrees with the stock ledger (gap {fmt.pct(v['gap'], 1)}, limit "
                          f"{fmt.pct(self.gap_max, 1)})")
            self.row("product_profitability", "load", reason, dst.name)
        second = second_hold(held, as_on, anchors, self.before)
        loads = None
        if held:                            # this hold counted with the held loads before it (loads_held)
            loads = 1 if second is False else self.loads_held(anchors.get("profit_as_of"), self.report_day)
            if second:
                loads = max(2, loads)       # second_hold has it from the load before: never fewer than two
        if second is None and loads and loads > 1:
            # this drop is already loaded: counted as itself, and the register is left exactly as it is
            self.checks.append(f"Profitability: held back for the {ordinal(loads)} load in a row (this drop is already "
                               "loaded: the action list stays as it is).")
        if second is not None:              # None: this drop is already loaded, the register stays as it is
            found: list[dict] = []
            if second:
                short = (v.get("ledger_cost") or Decimal(0)) - (v.get("report_cost") or Decimal(0))
                kept = iso(anchors.get("profit_as_of"))
                since = (date.fromisoformat(kept) + timedelta(days=1)).isoformat() if kept else ""
                moved = (f"; {v['still']} of the {v['moved']} items sold since {fmt.day(kept)} show no cost movement"
                         if v.get("moved") else "")
                found.append({"key": alerts.make_key("profit_held", as_on), "level": "ACT",
                              "title": ("Ask accounts to post cost of sales in Focus for the invoices since "
                                        + (fmt.day(since) if since else "the last loaded report")
                                        + ", then export the profitability report again"),
                              "detail": "; ".join(v["why"]) + moved
                                        + f"; not loaded for {loads} loads; production's report is as on "
                                        + fmt.day(kept),
                              "who": "Accounts", "due": None, "bhd": short if short > 0 else None, "evidence": []})
                self.flag("profit_held_twice")
                self.checks.append(f"Profitability: held back for the {ordinal(loads)} load in a row: one item for the "
                                   "action list (cost of sales looks unposted in Focus: accounts post it, then the "
                                   "report is exported again).")
            self.records["profit_held"] = found
        self.profit = v
        self.profit_evidence = {
            "as_on": as_on, "held": held, "verdict": "hold_back" if held else "load", "why": list(v["why"]),
            "gap_pct": _pct(v.get("gap")), "limit_pct": _pct(self.gap_max),
            "ledger_cost_bhd": _money(v.get("ledger_cost")), "report_cost_bhd": _money(v.get("report_cost")),
            "previous_day": v.get("previous_day"), "previous_gap_pct": _pct(v.get("previous_gap")),
            "items_sold_since": v.get("moved"), "with_no_cost_movement": v.get("still"),
            "blank_cost_share_pct": _pct((v.get("blank") or {}).get("share")), "items": v.get("items"),
            "matched": v.get("matched"), "loads_held": loads, "production_as_of": iso(anchors.get("profit_as_of")),
            "line": line}
        gap = f" (gap {fmt.pct(v['gap'], 1)}, limit {fmt.pct(self.gap_max, 1)})" if v.get("gap") is not None else ""
        self.done("Profitability: " + WORDS["hold_back" if held else "load"] + gap + ".")

    def loads_held(self, kept, day: str) -> int:
        """How many loads in a row have held the profitability report back, the drop as on `day` counted
        once whether or not it is loaded already (held_in_a_row): the loads the writer finished, each day
        once, and for a load made by hand (it left no journal) the snapshot days production holds after
        its report date. Loads are counted, not snapshot dates: a held load whose drop brought no stock
        balance, or one that gave a date again, counts all the same."""
        kept = iso(kept)
        snapshots: dict[str, list[str]] = {}
        for r in (self.q(Q_SNAPSHOT_DAYS, [kept, kept]) if kept else []):
            if r.get("d"):
                snapshots.setdefault(str(r.get("k") or ""), []).append(iso(r.get("d")))
        loads, explained = finished_held_loads()
        return held_in_a_row(day, kept, loads, snapshots, explained)

    # 6 ── price book
    def step_price_books(self) -> None:
        kinds = [k for k in ("selling_prices:MA_base", "selling_prices:modern_trade") if k in self.inv["files"]]
        if not kinds:
            self.done("Price book: not in this drop (live prices stay as they are).")
            return
        self.done("; ".join(self.price_book(kind) for kind in kinds) + ".")

    def price_book(self, kind: str) -> str:
        from scripts import ingest as ing
        from scripts.load_supabase import focus_book_kind
        opt, run, path = self.options, self.run, self.inv["files"][kind]
        book, exported = kind.split(":", 1)[1], self.inv["as_on"][kind]
        try:
            rows = ing.parse_pricebook(ing.read_grid(path), path.name, book)
        except ing.HeaderMismatch as exc:
            self.row(kind, "stop", "its header row is not the one the loader reads: " + str(exc)[:160])
            return f"{LABEL[kind]}: STOP (its header row changed)"
        live = self.q(Q_LIVE_PRICES, [book])
        diff = pricebook_diff(rows, live, book, exported)
        items = diff["items"]
        # what was already sold at a scoped price: looked up for SCOPE_LOOKUPS_MAX rows at most (each is
        # one read), the rows below the general price first: those are the ones for the action list
        scoped = sorted((e for e in items if e["scoped"] and e["what"] != "dropped" and e["start"]),
                        key=lambda e: not e["below_general"])
        unanswered = 0
        for e in scoped[:SCOPE_LOOKUPS_MAX]:
            try:
                scope_sales(e, self.deps)
            except Exception:  # noqa: BLE001  (its line then says nothing about sales, and the count says why)
                unanswered += 1
        not_looked = len(scoped) - len(scoped[:SCOPE_LOOKUPS_MAX])
        out = self.books[book] = {"book": book, "kind": kind, "file_rows": rows, "live_rows": live, "items": items,
                                  "exported": exported, "staged": None}
        n = Counter(e["what"] for e in items)
        below = [e for e in items if e["below_general"]]
        what = ", ".join(f"{n[w]} {w}" for w in ("added", "changed", "dropped") if n[w])
        choice = opt["pricebook"]
        if choice == "load":
            name = book_stage_name(book, run)
            if focus_book_kind(name) != book or ing.classify(name) != kind:
                self.row(kind, "stop", "the name of the staged copy would not be read as this price book")
                return f"{LABEL[kind]}: STOP"
            copy_checked(path, self.staging / name)
            self.staged[kind] = self.staging / name
            out["staged"] = name
            self.row(kind, "load", "loaded on your word (--pricebook load): "
                     + (what if items else "it is the same as the live book"), name)
        elif not items:
            self.row(kind, "leave_out", f"same as the live price book ({diff['file_keys']:,} rows): nothing to load")
        elif choice == "skip":
            self.row(kind, "leave_out", f"left out on your word (--pricebook skip): {what}")
        else:
            self.row(kind, "needs_decision", f"{plural(len(items), 'row')} differ from the live book ({what})"
                     + (f"; {len(below)} for one van or customer, below the general price" if below else "")
                     + ": left out unless you say load")
            self.flag("needs_decision")
        if items:
            self.checks.append(f"{LABEL[kind]} (exported {fmt.day(exported)}) against the live book, key by key:")
            self.checks.append(Cut(lambda limit: ["    " + price_line(e) for e in items[:limit]]
                                   + ([f"    and {len(items) - limit} more rows"] if len(items) > limit else [])))
            if not_looked:
                self.checks.append(f"    What was already sold at a scoped price was looked up for {SCOPE_LOOKUPS_MAX} of "
                                   f"the {len(scoped):,} scoped rows (those below the general price first); for the "
                                   f"other {not_looked:,} it was not looked up.")
            if unanswered:
                self.checks.append("    What was already sold at the scoped price could not be read for "
                                   f"{plural(unanswered, 'row')} (the database did not answer): those lines say "
                                   "nothing about sales.")
            if self.rows[kind]["verdict"] == "needs_decision":
                self.checks.append("    The decision is for the whole file: say \"load the price book\" "
                                   "(--pricebook load) or \"leave it out\" (--pricebook skip). With no answer it is "
                                   "left out and the rest loads.")
        else:
            self.checks.append(f"{LABEL[kind]} (exported {fmt.day(exported)}): no key added, changed or dropped against "
                               "the live book.")
        self.records.setdefault("price_scope", []).extend(price_scope_items(items))
        if below:
            self.flag("price_scope")
        return f"{LABEL[kind]}: {self.rows[kind]['decision']}" + (f" ({what})" if items else "")

    # 7 ── the preview
    def step_preview(self) -> None:
        as_on, staged = self.inv["as_on"], self.staged

        def hook(target, row):
            # the sales rows for the loader's guards; the rows of a table the loader folds for folded_noise
            if target in ("orders", "order_lines") or target in FOLDED:
                self.captured.setdefault(target, []).append(row)
            return row

        preview_dir = allowed_output(loadplan.run_dir(self.run) / "preview")
        # the preview's own backend inside a timing proxy (scripts/ops/querytime.py): seconds per report, and a
        # line as each report is compared. YQ_CHECK_NO_TIMING=1 hands the backend over bare (a diagnostic switch)
        backend, timer = self.session.backend, None
        if os.environ.get("YQ_CHECK_NO_TIMING") != "1":
            from scripts.ops.querytime import TimedBackend
            backend = timer = TimedBackend(self.session.backend, targets=LABEL_BY_TARGET, on_report=self.compared)
        into = allowed_output(loadplan.run_dir(self.run) / PREVIEW_IN)
        try:
            self.fill_preview_folder(into)
            result = self.deps.preview(into, backend, preview_dir, self.today, hook)
        except Exception as exc:  # noqa: BLE001
            raise Stop(f"The preview against the database failed ({plain_error(exc)}). Nothing was changed.") from None
        finally:
            if timer is not None:
                self.preview_parts = timer.timing_summary()
            drop_folder(into)
        self.files.append(str(preview_dir))
        summary = self.summary = result.get("summary") or {}
        targets = self.targets = summary.get("targets") or {}
        pv = self.pv = read_preview(summary, result.get("exceptions") or [])
        acts = pv["actions"]

        def counts(target: str) -> str:
            a = acts.get(target) or {}
            return f"{a.get('inserted', 0):,} new, {a.get('updated', 0):,} changed, {a.get('removed', 0):,} removed"

        fired = self.sales_guards()
        for kind, noun, held_to in (("order_lines", "lines", "sales_to"), ("orders", "invoices", "orders_to")):
            if kind not in staged:
                continue
            t = targets.get(kind) or {}
            scope = t.get("scope") or {}
            ends, held = iso(scope.get("to")), iso(self.anchors.get(held_to))
            if fired.get(kind):
                self.row(kind, "stop", "; ".join(fired[kind]) + ". The loader would not replace these dates and the "
                         "load would fail its verification. If the file is right, say so: the check then runs with "
                         "--replace-ok")
                self.guards += fired[kind]
            elif ends and held and ends < held:          # the file stops before what production holds
                self.row(kind, "stop", ends_early_reason(noun, ends, held))
                self.flag(OLDER)
            else:
                self.row(kind, "load", f"{int(t.get('file_rows') or 0):,} {noun} to {fmt.day(scope.get('to'))}: "
                         f"{counts(kind)}", staged[kind].name)
        if "receivables" in staged:
            t = targets.get("ar_ageing") or {}
            a = acts.get("ar_ageing") or {}
            again = int(t.get("db_rows") or 0) > 0
            same_total = bool(t.get("db_sums")) and t.get("db_sums") == t.get("sums")
            self.row("receivables", "load", f"{int(t.get('file_rows') or 0):,} accounts, BHD "
                     f"{fmt.bhd((t.get('sums') or {}).get('balance_bhd'))} as on {fmt.day(as_on.get('receivables'))}"
                     + (f" (that day is already loaded{', with the same total' if same_total else ''}: "
                        f"{a.get('inserted', 0)} new, {a.get('updated', 0)} rewritten, {a.get('removed', 0)} removed)"
                        if again else ""), staged["receivables"].name)

        changed = pv["invoices"]
        if "orders" in staged or "order_lines" in staged:
            self.checks.append(
                "Invoices: " + (counts("orders") if "orders" in staged else self.absent("orders"))
                + ". Lines: " + (counts("order_lines") if "order_lines" in staged else self.absent("order_lines")) + ".")
            if changed:
                self.flag("invoices_changed")
                names = sorted(changed)
                beyond = max(pv["more_invoices"], 0)               # rows the preview's own samples leave out

                def invoices(limit: int) -> list[str]:
                    unlisted = max(len(names) - limit, 0)          # invoices the list is too short for
                    return ["    Invoices that differ from what was loaded (the load replaces them; each one is for "
                            "the action list, to check): " + ", ".join(f"{i} ({changed[i]})" for i in names[:limit])
                            + (f" and {plural(unlisted, 'more invoice')}" if unlisted else "")
                            + (f"; and {plural(beyond, 'more changed row')} not listed" if beyond else "")]
                self.checks.append(Cut(invoices))
            day = fmt.day(as_on.get("order_lines") or as_on.get("orders"))
            self.records["inv_changed"] = [
                {"key": alerts.make_key("inv_changed", inv.replace("|", "/")), "level": "CHECK",
                 "title": ("Invoice no longer in Focus's export: to check" if what == "removed" else
                           "Invoice changed in Focus after it was loaded: to check"),
                 "detail": f"{inv}: {what} between the last load and the reports as on {day}",
                 "who": None, "due": None, "bhd": None, "evidence": [inv]}
                for inv, what in sorted(changed.items())][:RECORD_MAX]
            sales = summary.get("sales") or {}
            days, reps = sales.get("per_day_changed") or [], sales.get("per_salesman_changed") or []
            new_days = [d for d in days if not int(d.get("db_rows") or 0)]
            old_days = [d for d in days if int(d.get("db_rows") or 0)]
            if days or reps:
                self.checks.append(
                    f"Sales by day: {plural(len(new_days), 'new day')}"
                    + (" (" + listed([fmt.day(d.get("key")) for d in new_days], 8) + ")" if new_days else "")
                    + f", {plural(len(old_days), 'loaded day')} with different lines"
                    + (" (" + listed([f"{fmt.day(d.get('key'))}: file {d.get('file_rows')} lines / "
                                      f"{fmt.bhd(d.get('file_gross'))}, loaded {d.get('db_rows')} / "
                                      f"{fmt.bhd(d.get('db_gross'))}" for d in old_days], 6) + ")" if old_days else "")
                    + f"; {plural(len(reps), 'salesman', 'salesmen')} with new or changed lines.")
                if reps:
                    # whose sales the load changes: the largest change in gross first
                    by_change = sorted(reps, key=lambda r: (-abs(dec(r.get("file_gross")) - dec(r.get("db_gross"))),
                                                            str(r.get("key") or "")))
                    self.checks.append("Sales by salesman, the file against what is loaded (lines / gross BHD): "
                                       + "; ".join(rep_change(r) for r in by_change[:REPS_MAX])
                                       + (f"; and {len(by_change) - REPS_MAX} more" if len(by_change) > REPS_MAX else "")
                                       + ".")
            elif "order_lines" in staged and "per_day_changed" in sales:
                self.checks.append("Sales by day and by salesman: the file equals what is loaded.")
        stock = summary.get("stock") or {}
        for w in (stock.get("per_warehouse") or [])[:6]:
            self.checks.append(
                f"Stock against the last loaded snapshot (as on {fmt.day(stock.get('db_latest_as_of'))}), "
                f"{w.get('warehouse')}: file {int(w.get('file_items') or 0):,} items / {fmt.units(w.get('file_qty'))} "
                f"units / BHD {fmt.bhd(w.get('file_value'))}; loaded {int(w.get('db_items') or 0):,} / "
                f"{fmt.units(w.get('db_qty'))} / {fmt.bhd(w.get('db_value'))}; difference "
                f"{fmt.units(w.get('delta_qty'))} units / BHD {fmt.bhd(w.get('delta_value'))}.")
        for text in pv["infos"]:
            self.checks.append("Preview: " + text + ".")
        for text in pv["stops"]:
            self.stops.append("Preview: " + text)
        self.done("Preview against the database: "
                  + ("invoices " + counts("orders") if "orders" in staged else self.absent("orders"))
                  + ("; lines " + counts("order_lines") if "order_lines" in staged else "") + ".")

    def fill_preview_folder(self, into: Path) -> None:
        """What the preview compares: a copy of every staged file but the stock ledger, in the run's folder (the
        staging folder itself stays as the token binds it). Step 8 compares the ledger, voucher by voucher and in
        one read row by row (Q_LEDGER_SAME), so the preview never sends its rows to the database: that was most of
        a check's time."""
        ledger = self.staged.get("stock_movements")
        into.mkdir(parents=True, exist_ok=True)      # an empty one too, as an empty staging folder was handed over
        for path in sorted(self.staging.iterdir()):
            if path.is_file() and path != ledger:
                paths.copy_file(path, into / path.name)

    def compared(self, target: str, seconds: float, ok: bool) -> None:
        """One progress line as the preview's comparison of a report is done (querytime.TimedBackend)."""
        self.say(f"      preview: {LABEL_BY_TARGET.get(target, target)} compared in {seconds:.1f} s"
                 + ("" if ok else " (the comparison failed)"))

    def sales_guards(self) -> dict[str, list[str]]:
        """The loader's own span and salesman guards, fed with the rows production holds inside the
        file's dates."""
        db_rows: dict[str, list[dict]] = {}
        for target, sql in (("orders", Q_ORDERS_SPAN), ("order_lines", Q_LINES_SPAN)):
            scope = (self.targets.get(target) or {}).get("scope") or {}
            if self.captured.get(target) and scope.get("from") and scope.get("to"):
                db_rows[target] = self.q(sql, [scope["from"], scope["to"]])
        return sales_guards(self.captured, db_rows, self.options["replace_ok"])

    # 8 ── stock ledger, voucher by voucher
    def step_ledger(self) -> None:
        if "stock_movements" not in self.inv["files"]:
            self.done("Stock ledger: not in this drop.")
            return
        rows = self.ledger_rows
        if rows is None or "stock_movements" not in self.staged:
            self.done(f"Stock ledger: STOP ({self.ledger_stop or 'its header row changed'}).")
            return
        dates = sorted(iso(r.get("move_date")) for r in rows if r.get("move_date"))
        if not dates:
            self.row("stock_movements", "stop", "the file holds no dated movement")
            self.unstage("stock_movements")
            self.done("Stock ledger: STOP (no dated movement).")
            return
        lo, hi = dates[0], dates[-1]
        cmp = ledger_compare(rows, self.q(Q_LEDGER, [lo, hi]), iso(self.anchors.get("moves_to")) or None)
        # what is missing only of SIM never stops it (the owner, 18-Oct-2026): SIM never counts
        split = ledger_sim(cmp)
        sim_gone = sim_words(split["sim"])
        sim_verb = "is" if sum(len(v) for v in split["sim"].values()) == 1 else "are"
        missing = [(split["other"]["items"], "item"), (split["other"]["warehouses"], "warehouse"),
                   (split["other"]["types"], "voucher type")]
        if any(m for m, _ in missing):
            why = "; ".join(f"{plural(len(m), word)} loaded for these dates {'is' if len(m) == 1 else 'are'} not in the "
                            f"file ({listed(m, 4)})" for m, word in missing if m)
            if sim_gone:
                why += f" (besides {sim_gone}, which alone would not stop it: SIM never counts)"
            self.row("stock_movements", "stop", why + ". The loader replaces the file's whole date span, so these "
                     "movements would be lost: export the full Stock ledger again")
            self.done("Stock ledger: STOP ("
                      + ", ".join(f"{plural(len(m), word)} missing" for m, word in missing if m) + ").")
            return
        cmp, sim_pairs = drop_sim_pairs(cmp)
        count, digest = ledger_digest(rows)
        same = self.q(Q_LEDGER_SAME, [count, digest, lo, hi, sorted({r["item_name"] for r in rows})])
        self.ledger_same = bool(same and same[0].get("same"))
        tally = (f"changed {len(cmp['changed'])}, removed {len(cmp['removed'])}, entered with an earlier date "
                 f"{len(cmp['backdated'])}")
        differ = len(cmp["changed"]) + len(cmp["removed"]) + len(cmp["backdated"])
        self.ledger_cmp = {"changed": len(cmp["changed"]), "removed": len(cmp["removed"]),
                           "backdated": len(cmp["backdated"]), "new": int(cmp["new"])}
        nothing = f"nothing missing outside SIM; {sim_gone} {sim_verb} not in the file (SIM never counts)" if sim_gone \
            else "nothing missing"
        self.row("stock_movements", "load", f"{len(rows):,} movements to {fmt.day(hi)}: {nothing}; "
                 f"{plural(differ, 'voucher')} to check", self.staged["stock_movements"].name)
        self.checks.append(f"Stock ledger: {'outside SIM, ' if sim_gone else ''}every item ({cmp['items']}), warehouse "
                           f"({cmp['warehouses']}) and voucher type ({cmp['types']}) production holds for "
                           f"{span_text(lo, hi)} is in the file. Voucher by voucher ({cmp['pairs']:,} voucher-and-item "
                           f"pairs): {tally}, new {cmp['new']:,}.")
        if sim_gone:
            self.flag("sim_ledger_out")
            self.checks.append(f"Stock ledger, SIM: not in the file: {sim_words(split['sim'], named=True)}. "
                               f"{plural(split['rows'], 'loaded row')} of theirs inside these dates would go with a load "
                               "of this file and come back with an export that has them (SIM never counts; Focus keeps "
                               "Service items out of its stock reports).")
        if sim_pairs:
            self.checks.append(f"Stock ledger, SIM: {plural(sim_pairs, 'voucher-and-item pair')} of SIM items differ "
                               "from what is loaded; they are not listed to check (SIM never counts).")
        shown = ([("changed", k) for k in cmp["changed"]] + [("removed", k) for k in cmp["removed"]]
                 + [("earlier date", k) for k in cmp["backdated"]])
        if shown:
            self.flag("vouchers_changed")

            def vouchers(limit: int) -> list[str]:
                return [f"    {what}: {voucher or '(no number)'}, {item}: loaded as "
                        f"{lines_text(cmp['db'].get((voucher, item)))}; the file has "
                        f"{lines_text(cmp['file'].get((voucher, item)))}" for what, (voucher, item) in shown[:limit]] + (
                    [f"    and {len(shown) - limit} more (each one is for the action list, to check)"]
                    if len(shown) > limit else [])
            self.checks.append(Cut(vouchers))
        if self.pv.get("ledger_bulk") is not None:
            self.checks.append(f"Stock ledger: the loader reloads the file's whole date span ({self.pv['ledger_bulk']:,} "
                               "loaded rows are replaced). Focus recalculates its running balance on every export, so "
                               "the rows look changed although the vouchers are as listed above: expected.")
        self.records["voucher_changed"] = voucher_items(cmp)[:RECORD_MAX]
        self.done(f"Stock ledger: nothing missing{f' outside SIM ({sim_gone} not in the file)' if sim_gone else ''}; "
                  f"vouchers {tally}.")

    def loader_guards(self) -> None:
        """The loader's guards for the stock snapshot and for a staged price book (its sales guards
        were read with the preview). A guard predicted to fire is a STOP on that file."""
        from scripts.load_supabase import focus_book_kind, snapshot_guard, stale_pricebook_rows, warehouse_set_guard
        opt, pv = self.options, self.pv
        if self.stock and "stock_balance" in self.staged:
            before = {str(r.get("w")) for r in self.q(Q_WAREHOUSES_BEFORE, [self.stock["as_on"]])}
            why = warehouse_set_guard(before, set(self.stock["raw_warehouses"]), opt["partial_ok"])
            seen = pv.get("warehouse_shrinks")
            if why or (seen is not None and not opt["partial_ok"]):
                text = ("it covers fewer warehouses than the snapshot before it" if why else
                        "it covers fewer warehouses than the snapshot already loaded for that day"
                        ) + (f" ({listed(seen, 5)} missing)" if seen else f" ({len(before)} then)")
                self.guards.append("stock balance: " + text)
                self.row("stock_balance", "stop", text + ": the loader warns and the load fails its verification. If the "
                         "narrower file is intended, say so: the check then runs with --partial-ok")
            elif seen is not None:
                self.checks.append("Stock balance: fewer warehouses than the last snapshot (" + listed(seen, 5)
                                   + " missing): allowed by --partial-ok.")
        for book, b in self.books.items():
            if not b.get("staged"):
                continue
            loaded = [dict(r, source_file=b["staged"]) for r in b["file_rows"]]
            names = {p.name for p in self.staging.iterdir() if p.is_file() and p.suffix.lower() in EXCEL}
            live_focus = [r for r in b["live_rows"] if focus_book_kind(r.get("source_file")) == book]
            why = snapshot_guard(stale_pricebook_rows(loaded, b["live_rows"], names), live_focus,
                                 {str(r.get("sku_code") or "").strip() for r in loaded} - {""})
            if why or pv.get("void_guard"):
                text = "it would switch off too much of the live book (" + (why or str(pv.get("void_guard"))[:160]) + ")"
                self.guards.append(f"price book {book}: " + text)
                self.row(b["kind"], "stop", text + ". Is the export complete? Export the whole book again, or leave it "
                         "out (--pricebook skip)")
        self.checks.append("Loader guards: none is predicted to fire." if not self.guards else
                           "Loader guards predicted to fire: " + "; ".join(self.guards) + ".")

    # 9 ── catalog
    def step_catalog(self) -> None:
        from scripts.load_supabase import focus_book_kind
        from scripts.verify_numbers import current_prices
        catalog_rows = self.q(Q_CATALOG)
        live_book = {str(r.get("sku_code")).strip() for r in self.q(Q_BOOK) if r.get("sku_code")}
        candidates = {str(r.get("sku_code")).strip() for r in self.q(Q_CANDIDATES) if r.get("sku_code")}
        b = self.books.get("MA_base")
        if b and b.get("staged") and self.rows.get(b["kind"], {}).get("verdict") == "load":
            others = [r for r in b["live_rows"] if focus_book_kind(r.get("source_file")) != "MA_base"]
            book_after = set(current_prices(list(b["file_rows"]) + others, today=self.today))
            candidates |= {str(r.get("sku_code")).strip() for r in b["file_rows"]
                           if r.get("sku_code") and dec(r.get("rate_bhd")) > 0}
            basis = "the price list this book would give"
        else:
            book_after, basis = live_book, "the live price list (no price book loads)"
        effect = self.catalog = catalog_effect(catalog_rows, book_after, candidates)
        line = catalog_line(effect)
        self.checks.append(line + f" Compared with {basis}.")
        if effect["off"] or effect["on"]:
            self.flag("catalog_switch")
        self.done(line)

    # 10 ── the output
    def expected_figures(self, fingerprint: dict | None) -> dict:
        """What the owner should see after the load: from the preview and, for what this drop does
        not bring, from what production holds now."""
        anchors, as_on, staged, targets, pv = self.anchors, self.inv["as_on"], self.staged, self.targets, self.pv
        lines: list[str] = []
        out: dict = {}

        def total_after(table: str) -> int | None:
            if not fingerprint or table not in fingerprint:
                return None
            a = pv["actions"].get(table) or {}
            return int(fingerprint[table][0]) + a.get("inserted", 0) - a.get("removed", 0)

        line_scope = (targets.get("order_lines") or {}).get("scope") or {}
        sales_to = max(iso(line_scope.get("to")) if "order_lines" in staged else "", iso(anchors.get("sales_to"))) or None
        out["sales_to"] = sales_to
        n_inv = out["invoices"] = total_after("orders") if "orders" in staged else None
        n_lines = out["invoice_lines"] = total_after("order_lines") if "order_lines" in staged else None
        if "order_lines" in staged or "orders" in staged:
            counts = [x for x in (f"{n_inv:,} invoices" if n_inv is not None else "",
                                  f"{n_lines:,} lines" if n_lines is not None else "") if x]
            lines.append(f"Sales to {fmt.day(sales_to)}" + (": " + ", ".join(counts) + " in all." if counts else "."))
        else:
            here = any(k in self.inv["files"] for k in ("order_lines", "orders"))
            lines.append(("Sales: not loaded; " if here else "Sales: no sales report in this drop; ")
                         + f"they stay as loaded to {fmt.day(sales_to)}.")
        stock = self.stock
        if stock and "stock_balance" in staged:
            out["stock"] = {"as_on": stock["as_on"], "items": stock["items"], "units": plain(stock["units"]),
                            "value_bhd": str(fmt.q3(stock["value"]))}
            lines.append(f"Stock (main store) as on {fmt.day(stock['as_on'])}: {stock['items']:,} items / "
                         f"{fmt.units(stock['units'])} units / BHD {fmt.bhd(stock['value'])} at selling rate.")
        else:
            lines.append(f"Stock: not loaded; it stays as on {fmt.day(anchors.get('stock_as_of'))}.")
        if "receivables" in staged:
            t = targets.get("ar_ageing") or {}
            rc = self.summary.get("receivables") or {}
            total = (t.get("sums") or {}).get("balance_bhd")
            out["receivables"] = {"as_on": as_on.get("receivables"), "accounts": int(t.get("file_rows") or 0),
                                  "total_bhd": total, "focus_total_bhd": rc.get("focus_grand_total")}
            lines.append(f"Receivables as on {fmt.day(as_on.get('receivables'))}: BHD {fmt.bhd(total)} on "
                         f"{int(t.get('file_rows') or 0):,} accounts"
                         + (f" (Focus's own total: BHD {fmt.bhd(rc.get('focus_grand_total'))})"
                            if rc.get("focus_grand_total") else "") + ".")
        else:
            lines.append(f"Receivables: not loaded; they stay as on {fmt.day(anchors.get('ar_as_of'))}.")
        if "stock_movements" in staged:              # the preview never sees the ledger: the file's own rows
            rows = self.ledger_rows or []
            to = max((iso(r.get("move_date")) for r in rows if r.get("move_date")), default=None)
            out["stock_movements"] = {"rows": len(rows), "to": to}
            lines.append(f"Stock movements to {fmt.day(to)}: {len(rows):,} rows.")
        else:
            lines.append(f"Stock movements: not loaded; they stay as loaded to {fmt.day(anchors.get('moves_to'))}.")
        if "product_profitability" in staged:
            out["profitability_as_on"] = as_on.get("product_profitability")
            lines.append(f"Profitability as on {fmt.day(as_on.get('product_profitability'))}.")
        else:
            out["profitability_as_on"] = iso(anchors.get("profit_as_of")) or None
            lines.append("Profitability: not loaded; it stays as on " + fmt.day(anchors.get("profit_as_of"))
                         + (" (held back)." if self.profit and self.profit.get("held") else "."))
        loaded = [b for b in self.books.values()
                  if b.get("staged") and self.rows.get(b["kind"], {}).get("verdict") == "load"]
        for b in loaded:
            n = Counter(e["what"] for e in b["items"])
            lines.append(f"Price book {b['book']}: {n['added']} rows added, {n['changed']} changed, {n['dropped']} "
                         "switched off.")
        if not loaded:
            lines.append("Prices: no price book loads; live prices stay as they are.")
        out["lines"] = lines
        return out

    def not_received(self) -> None:
        """A partial drop is accepted: the table names what stays on older data."""
        a = self.anchors
        stays = {"order_lines": f"sales lines stay as loaded to {fmt.day(a.get('sales_to'))}",
                 "orders": f"invoices stay as loaded to {fmt.day(a.get('sales_to'))}",
                 "stock_movements": f"stock movements stay as loaded to {fmt.day(a.get('moves_to'))}",
                 "stock_balance": f"stock stays as on {fmt.day(a.get('stock_as_of'))}",
                 "receivables": f"receivables stay as on {fmt.day(a.get('ar_as_of'))}",
                 "product_profitability": f"profitability stays as on {fmt.day(a.get('profit_as_of'))}"}
        for kind in DAILY:
            if kind not in self.inv["files"]:
                self.row(kind, "not_loaded", "not received in this drop: " + stays[kind])
                self.flag("partial_drop")

    def table_text(self) -> str:
        rows: list[list[str]] = []
        for kind in KIND_ORDER:
            r = self.rows.get(kind)
            if r:
                rows.append([short_name(r["file"]) if r["file"] else f"(no {r['label']})", r["decision"], r["reason"]])
        extra = [[short_name(o["name"]), WORDS["not_loaded"], o["reason"]] for o in self.inv.get("others", [])]
        extra += [[short_name(f["name"] + "/"), WORDS["not_loaded"],
                   f"sub-folder with {plural(f['files'], 'file')}: {f['reason']}"] for f in self.inv.get("folders", [])]
        if len(extra) > OTHER_ROWS_MAX:
            more = len(extra) - (OTHER_ROWS_MAX - 1)
            extra = extra[:OTHER_ROWS_MAX - 1] + [[f"(+{more} more files and folders)", WORDS["not_loaded"],
                                                   "not loaded by this command"]]
        return fmt.grid(["File", "Decision", "Reason"], rows + extra, "lll")

    def record_items(self, data_date: str) -> list[str]:
        """Put what the check found on the action register (scripts/ops/alerts.py): only a check
        that could go ahead does (a stopped check may be looking at a broken file), and never a dev
        run. An invoice or a voucher that changed is an EVENT: it is listed from the first time it
        is found, and no later check clears it (the load itself makes the two sides equal, so the
        next check finds nothing): it stays open until he marks it done or explained. A report held
        back twice is listed from its first time too and clears once the report loads. A scoped
        price below the general one is a state whose first run is the baseline."""
        lines: list[str] = []
        for rule in EVENT_RULES + STATE_RULES:
            if rule not in self.records:
                continue                                   # that check did not run on this drop
            found = list(self.records[rule])
            if self.options["dev"] or self.stops:
                if found:
                    lines.append(f"{plural(len(found), 'item')} for the action list ({rule}): not recorded "
                                 + ("in a dev run." if self.options["dev"] else "while the check stops."))
                continue
            event = rule in EVENT_RULES
            try:
                counts = alerts.record(rule, found, data_date, baseline=rule in BASELINE_RULES,
                                       clear_missing=not event)
                still_open = 0
                if event:
                    now = {f["key"] for f in found}
                    still_open = sum(1 for row in alerts.current().values() if row.get("rule") == rule
                                     and row.get("state") == "open" and row.get("key") not in now)
            except (OSError, ValueError) as exc:
                lines.append(f"Action list ({rule}): could not be written ({type(exc).__name__}); what was found is "
                             "listed above.")
                continue
            said = register_line(rule, counts, still_open)
            if said:
                lines.append(said)
        return lines

    def moved_meanwhile(self, parts: dict) -> list[str]:
        """What changed under the check while it ran, as stops. A report of the folder: its hash at
        the folder step against its hash now (a verdict is worked out from the file and the copy is
        made afterwards, so a file saved again in between would be staged unjudged). The database:
        the fingerprint read before the first comparison against the one the token takes (the
        table would bind a database it was not worked out on)."""
        out: list[str] = []
        now = parts.get("drop") or {}
        changed = sorted(p.name for kind, p in self.inv["files"].items()
                         if (now.get(p.name) or {}).get("sha256") != self.inv["hashes"].get(kind))
        if changed:
            out.append(f"{listed(changed, 4)} changed in the folder while the check ran, so the table is not about "
                       "the files as they are now. Run the check again")
        before, after = self.fingerprint or {}, parts.get("fingerprint") or {}
        tables = sorted(k for k in set(before) | set(after) if before.get(k) != after.get(k))
        if tables:
            out.append(f"the database changed while the check ran ({listed([DB_WORDS.get(t, t) for t in tables], 4)}): "
                       "another load or upload was at work. Run the check again")
        return out

    def what_is_new(self) -> list[str]:
        """What keeps this drop from being "nothing new", in short words ([] = nothing at all: every report
        the load would apply matches production row for row, and nothing else a load does would change
        anything: the stock ledger voucher by voucher, the sales by day and salesman, the catalog sync, a
        last load that failed after it had begun to write). A line that waits for his decision is not part
        of the load by default (its default leaves it out), so it is named, not counted (not_in_load). Asked
        when the steps are done and before the database's fingerprint is read again, so the one read it may
        make (folded_noise, and only when nothing else is new) is inside what moved_meanwhile watches."""
        if self.stops:
            return ["the check stops"]
        out: list[str] = []
        if self.failed:
            out.append("the last load failed after it had begun to write: a load replaces it")
        if self.failed_unsure:                       # it could not be told: never "nothing new" on that
            out.append("whether the last load failed could not be told (" + self.failed_unsure + ")")
        if not self.staged:
            out.append("nothing is staged")
        rewritten: dict[str, int] = {}
        for target in dict.fromkeys(KIND_TARGET[k] for k in KIND_ORDER if k in self.staged and k in KIND_TARGET):
            if target == "stock_movements":
                continue                             # never in the preview: step 8 compares it (below)
            a = (self.targets.get(target) or {}).get("actions")
            label = LABEL_BY_TARGET.get(target, target)
            if not (isinstance(a, dict) and {"inserted", "updated", "unchanged"} <= set(a)
                    and ("deleted" in a or "voided" in a)):
                out.append(f"{label}: not compared by the preview")      # empty or partial counts prove nothing
                continue
            moved = {word: n for word, n in (("new", int(a.get("inserted") or 0)),
                                             ("removed", int(a.get("deleted") or 0) + int(a.get("voided") or 0)),
                                             ("back in force", int(a.get("revived") or 0))) if n}
            changed = int(a.get("updated") or 0)
            if moved:
                out.append(f"{label}: " + ", ".join(f"{n:,} {word}" for word, n in moved.items())
                           + (f", {changed:,} changed" if changed else ""))
            elif changed:
                rewritten[target] = changed              # perhaps a folded row only: looked at last
        if "stock_movements" in self.staged:
            c = self.ledger_cmp
            if c is None:
                out.append("Stock ledger: not compared voucher by voucher")
            elif any(c.values()):
                out.append("Stock ledger: " + ", ".join(f"{n:,} {word}" for word, n in c.items() if n))
            elif self.ledger_same is not True:       # the vouchers are the same, a row's figures are not
                out.append("Stock ledger: rows differ from what is loaded (their figures, not their vouchers)")
        if self.pv.get("invoices"):
            out.append(f"{plural(len(self.pv['invoices']), 'invoice')} differ from what is loaded")
        sales = self.summary.get("sales") or {}
        if "order_lines" in self.staged and (sales.get("per_day_changed") or sales.get("per_salesman_changed")):
            out.append("sales by day or by salesman differ from what is loaded")
        effect = self.catalog
        if effect is None:
            out.append("Catalog: not compared")
        elif effect["off"] or effect["on"] or effect["new"]:
            out.append("Catalog: the sync after a load would change it")
        for target, n in rewritten.items():
            if not out and self.folded_noise(target) == n:
                self.noise[target] = n
            else:
                out.append(f"{LABEL_BY_TARGET.get(target, target)}: {n:,} changed")
        return out

    def not_in_load(self) -> list[str]:
        """The reports of the table that are not part of this load, each with the words that put it in
        (OPTION_SAY): held back, left out, or waiting for his decision (whose default leaves it out). A price
        book that is the same as the live one has nothing to load, and is named as such."""
        out: list[str] = []
        for kind in KIND_ORDER:
            r = self.rows.get(kind)
            if not r or r["verdict"] not in ("hold_back", "leave_out", "needs_decision"):
                continue
            say = OPTION_SAY.get(kind)
            book = self.books.get(kind.split(":", 1)[1]) if kind.startswith("selling_prices:") else None
            if r["verdict"] == "hold_back":
                line = f"{r['label']}: {r['decision']}." + (f" To load it anyway, say: {say}." if say else "")
            elif r["verdict"] == "needs_decision":
                line = (f"{r['label']}: {r['decision']} (left out unless you say load)."
                        + (f" To load it, say: {say}." if say else ""))
            elif book is not None and not book.get("items"):
                line = f"{r['label']}: {r['decision']} (the same as the live price book: nothing to load)."
            else:
                line = f"{r['label']}: {r['decision']} (left out on your word)." + (f" To load it, say: {say}." if say else "")
            out.append(line)
        return out

    def folded_noise(self, target: str) -> int | None:
        """The rows of `target` the preview counts as rewritten, when each of them is a row the loader folds
        from two or more lines of the report (FOLDED) and equals production's row in every column the
        preview compares: exactly, and in the figures the loader adds up to within a millionth of a dinar
        (same_cell). Then their number: they are no change, the load would write the same sums again.
        None when that cannot be shown for every one of them: they stay changes, as the preview says. Reads
        production's rows of the file's snapshot day(s), one read each."""
        t = self.targets.get(target) or {}
        updated = int((t.get("actions") or {}).get("updated") or 0)
        sample = t.get("updated_sample") or []
        cols = [c for c in (t.get("payload_cols") or []) if c != "source_file"]
        scope = t.get("scope") or {}
        kind = next((k for k, tg in KIND_TARGET.items() if tg == target and k in self.staged), None)
        if target not in FOLDED or kind is None or not updated or updated > len(sample) or not cols:
            return None
        from app.ingest_batch import TARGETS
        from scripts import ingest as ing
        cfg = TARGETS[target]
        if scope.get("kind") != "partition" or list(scope.get("cols") or []) != list(cfg["partition"]):
            return None
        keys, defaults = cfg["keys"], FOLD_DEFAULTS.get(target)
        wanted = {key_of(r, keys, defaults) for r in sample}
        path = self.staged[kind]
        parse = ing.parse_receivables if target == "ar_ageing" else ing.parse_stock_balance
        try:
            folded = folded_keys(target, parse(ing.read_grid(path), path.name))
        except Exception:  # noqa: BLE001  (a copy the parser cannot read again proves nothing)
            return None
        if len(wanted) != updated or not wanted <= folded:
            return None
        mine = {key_of(r, keys, defaults): r for r in self.captured.get(target) or []}
        held: dict[tuple, dict] = {}
        where = " and ".join(f"{c}::text = %s" for c in cfg["partition"])
        for part in scope.get("partitions") or []:
            if len(part) != len(cfg["partition"]):
                return None
            for r in self.q(f"select * from {target} where {where}", [str(x) for x in part]):
                held[key_of(r, keys, defaults)] = r
        figures = set(FOLDED[target])
        for key in wanted:
            f, p = mine.get(key), held.get(key)
            if f is None or p is None or not all(same_cell(f.get(c), p.get(c), c in figures) for c in cols):
                return None
        return updated

    def finish(self) -> dict:
        opt, inv, run, staging = self.options, self.inv, self.run, self.staging
        self.not_received()
        self.stops = [f"{r['label']}: {r['reason']}" for k in KIND_ORDER for r in [self.rows.get(k)]
                      if r and r["verdict"] == "stop"] + self.stops
        data_date = self.report_day
        if self.failed_unsure:
            reason = self.failed_unsure[0].upper() + self.failed_unsure[1:]
            self.checks.append(f"{reason} (open in another program, or damaged): whether the last load failed is not "
                               "known, so the folder is not taken as loaded already.")
        # what makes this drop more than "nothing new", worked out BEFORE the fingerprint is read again: the
        # one read it may make (folded_noise) is then inside the window moved_meanwhile watches
        self.new_in_drop = self.what_is_new()
        self.checks += [noise_line(target, n) for target, n in self.noise.items()]
        parts = None
        try:
            parts = loadplan.collect_parts(inv["drop"], run, paths.CODE_ROOT, self.pre.get("live") or "", self.session)
        except Exception as exc:  # noqa: BLE001
            self.stops.append(f"the database fingerprint could not be read ({type(exc).__name__}), so no token can "
                              "be made")
        if parts:
            moved = self.moved_meanwhile(parts)
            if moved:
                self.flag("moved_meanwhile")
                self.stops += moved
        self.expected = self.expected_figures((parts or {}).get("fingerprint"))
        if not self.staged:
            self.stops.append("nothing in this drop is ready to load")
        odd = self.marker_story()
        if odd:
            self.flag("marker_mismatch")
            self.stops += odd
        # nothing new: the whole check ran and found nothing for a load to do, so it issues no token
        nothing_new = not self.stops and not self.new_in_drop
        if nothing_new:
            self.flag("nothing_new")
        register_lines = self.record_items(data_date)
        token = loadplan.token_of(parts) if (parts and not self.stops and not opt["dev"] and not nothing_new) else None
        staged_names = sorted(p.name for p in staging.iterdir() if p.is_file())
        if self.failed:
            self.flag("last_load_failed")
        # as_on: the reports that LOAD, and no other. The writer names the published set after the
        # newest of these dates, so the Ledger export (never loaded; at month end it is dated apart
        # from the daily reports) and a report held back or left out must not be among them.
        # as_on_all keeps the date of every report of the folder.
        # options: what the owner said; the marker files in staging tell the same story (marker_story).
        # supersedes: the earlier checks of this folder that this one replaced (each of their plans
        # names this run under superseded_by).
        # nothing_new / new_in_drop: what_is_new; timings.preview_parts: the preview's seconds per report
        # and per helper (scripts/ops/querytime.py).
        timings = dict(self.timings, parse=self.grids.summary())
        if self.preview_parts is not None:
            timings["preview_parts"] = self.preview_parts
        timings["total"] = round(sum(self.timings.values()), 2)
        plan = {"run": run, "drop": str(inv["drop"]), "created_at": self.deps.now().isoformat(timespec="seconds"),
                "verdicts": [self.rows[k] for k in KIND_ORDER if k in self.rows], "expected": self.expected,
                "options": opt, "token": token, "token_parts": parts if token else None,
                "as_on": {k: v for k, v in inv["as_on"].items() if k in self.staged}, "as_on_all": inv["as_on"],
                "data_date": data_date, "staging": str(staging), "staged": staged_names, "stops": list(self.stops),
                "flags": list(self.flags), "guards": list(self.guards),
                "supersedes": list(self.retired.get("retired") or []),
                "not_loaded": [o["name"] for o in inv["others"]] + [f["name"] + "/" for f in inv["folders"]],
                "profitability": self.profit_evidence, "nothing_new": nothing_new,
                "new_in_drop": list(self.new_in_drop), "timings": timings}
        self.files.append(str(loadplan.write_plan(plan)))
        if token:
            self.files.append(str(loadplan.write_go_cmd(run, token, paths.CODE_ROOT)))

        if self.stops:
            self.flag("stop")
        elif opt["dev"]:
            self.flag("dev")
        elif not nothing_new:
            self.flag("token")
        aside = self.not_in_load() if nothing_new else []

        def block(limit: int, whole: bool) -> list[str]:
            """The printed block: the chat's (lists cut at CHAT_LIST_MAX, the line that names check.txt under the
            head line) or the whole one in check.txt (lists to LIST_MAX, what a load also does, staged, token)."""
            out = [failed_line(self.failed, starts=bool(token)), ""] if self.failed else []
            out.append(f"LOAD CHECK  {inv['drop'].name}  (run {run}, reports as on {fmt.day_long(data_date)})"
                       + ("  [DEV RUN]" if opt["dev"] else ""))
            if not whole:
                out.append(WHOLE_CHECK.format(run=run))
            out += ["", self.table_text(), "", "Checks"]
            out += [line if line.startswith("    ") else "- " + line
                    for line in shown_checks(self.checks, limit) + register_lines]
            if nothing_new:                # no load follows: what production holds, and nothing a load would do
                out += ["", "Production already holds (the same as these reports)"]
                out += ["- " + line for line in self.expected["lines"]]
                if aside:                  # ... and what stays outside it, with the words that put it in
                    out += ["", "Not part of this load:"] + ["- " + line for line in aside]
            else:
                out += ["", "After the load you should see"] + ["- " + line for line in self.expected["lines"]]
                if whole:
                    out += ["", "A load also does this, every time"] + ["- " + line for line in SIDE_EFFECTS]
            out.append("")
            if self.stops:
                out.append("STOP: this drop cannot be loaded as it is.")
                out += ["- " + (s if s.endswith(("?", "!")) else s.rstrip(".") + ".") for s in self.stops]
                out.append("No token was issued. Nothing was changed.")
            elif nothing_new:
                out.append(f"{NOTHING_NEW_HEAD} {inv['drop'].name}: everything this load would apply is already loaded"
                           f"{last_load_words()}. No load needed.")
            elif opt["dev"]:
                out.append("DEV RUN: nothing stops this drop, but a dev run issues no token and no GO line. To load, run "
                           "the check from the project folder.")
            else:
                if whole:
                    out.append(f"Staged in business_data/staging/{run}: " + listed(staged_names, 12) + ".")
                    out.append(f"Token {token}: it binds this table to the files, the programs and the database as they "
                               "are now.")
                where = go_cmd_line(run)
                if where:
                    out.append(where)
                out.append("GO: " + loadplan.go_line(run, token))
            return out

        whole_check = allowed_output(loadplan.run_dir(run) / "check.txt")
        self.files.append(str(paths.write_text(whole_check, "\n".join(self.said + block(LIST_MAX, True)) + "\n")))
        ok = not self.stops
        return {"ok": ok, "command": "load", "data_date": data_date, "text": "\n".join(block(CHAT_LIST_MAX, False)),
                "flags": self.flags, "files": self.files, "warnings": [], "exit": 0 if ok else 2, "run": run,
                "token": token}


def options_of(args) -> dict:
    return {"pricebook": getattr(args, "pricebook", None),
            "with_profitability": bool(getattr(args, "with_profitability", False)),
            "replace_ok": bool(getattr(args, "replace_ok", False)),
            "partial_ok": bool(getattr(args, "partial_ok", False)), "dev": bool(getattr(args, "dev", False))}


def stopped(check: Check, stop: Stop) -> dict:
    """A stop before the table exists (preflight, the folder, the database not reachable, a drop
    older than what is loaded). For an interrupted load the last line is that load's own GO line:
    running it again is the resume (the line before it says where its go.cmd lies). A dev run
    prints no GO line, not even that one."""
    dev = check.options["dev"]
    out = ["LOAD CHECK: STOP" + ("  [DEV RUN]" if dev else ""), str(stop)]
    flags = ["stop"] + ([stop.flag] if stop.flag else [])
    if isinstance(stop, Interrupted):
        if stop.go_line and dev:
            out.append("DEV RUN: no GO line is printed here. Resume that load from the project folder: type "
                       "/yq-load there.")
        elif stop.go_line:
            where = go_cmd_line(stop.run)
            if where:
                out.append(where)
            out.append("GO: " + stop.go_line)
    else:
        out.append("No token was issued. Nothing was changed.")
        said = retired_line(check.retired)           # put away before the check stopped: he must know
        if said:
            flags.append("replaced_earlier")
            out.append(said)
        left = []
        if check.run:
            left = [f"business_data/{part}/{check.run}" for part, folder in
                    (("staging", loadplan.staging_path(check.run)), ("ops/loads", loadplan.run_dir(check.run)))
                    if folder.exists()]
        if left:
            out.append("(Left by this check: " + " and ".join(left) + ".)")
        if check.failed:
            flags.append("last_load_failed")
            out = [failed_line(check.failed, starts=False), ""] + out
    return {"ok": False, "command": "load", "data_date": None, "text": "\n".join(out), "flags": flags,
            "files": check.files, "warnings": [], "exit": 2, "run": check.run, "token": None}


def run_check(args, deps: Deps | None = None) -> dict:
    """The whole check under one deadline. A database or a network that stops answering must not
    hold the command until the tool call is cut, so the work runs in a helper thread: if it is not
    back in time the command says so and ends. Nothing can be left half done in the database: the
    check only reads."""
    deps = deps or Deps()
    deps.quiet = bool(getattr(args, "json", False))
    drop, wrong = drop_asked(args)
    if wrong:
        return {"ok": False, "command": "load", "data_date": None, "text": "LOAD CHECK: STOP\n" + wrong,
                "flags": ["stop"], "files": [], "warnings": [], "exit": 2, "run": None, "token": None}
    check = Check(options_of(args), deps)
    box: dict = {}

    def work() -> None:
        try:
            box["result"] = check.go(drop)
        except Stop as stop:
            box["result"] = stopped(check, stop)
        except BaseException as exc:  # noqa: BLE001  (raised again in the caller's thread)
            box["error"] = exc

    limit = getattr(deps, "deadline_s", DEADLINE_S)
    if not limit:
        work()
    else:
        worker = threading.Thread(target=work, name="load-check", daemon=True)
        worker.start()
        worker.join(limit)
        if worker.is_alive():
            head = failed_line(check.failed, starts=False) + "\n\n" if check.failed else ""
            return {"ok": False, "command": "load", "exit": 2, "flags": ["stop", "timeout"], "data_date": None,
                    "files": check.files, "warnings": [], "run": check.run, "token": None,
                    "text": f"{head}LOAD CHECK: STOP\nThe check did not finish in {int(limit) // 60} minutes (it was "
                            f"at step {check.step + 1} of {STEPS}): the database or the network stopped answering. "
                            "Nothing was changed and no token was issued. Run the check again."}
    if "error" in box:
        raise box["error"]
    return box["result"]


# ── status ────────────────────────────────────────────────────────────────────

EVENT_WORDS = {"started": "started", "step_start": "started", "step_ok": "done", "step_failed": "FAILED",
               "note": "note", "finished": "FINISHED", "failed": "FAILED"}
# The delivery journal's rows (loads/<run>/delivery.jsonl): an e-mail that did not go is never a FAILED load.
DELIVERY_WORDS = {"queued": "started", "sent": "e-mail sent", "test_sent": "TEST e-mail sent", "not_sent": "not sent",
                  "failed": "e-mail NOT sent"}
JOURNAL_SAYS = ("text", "why", "next")      # the fields the writer fills for the owner (scripts/yq_load_go.py)


def journal_says(row: dict) -> str:
    """The writer's own sentences of one journal row, in full and in its order: text, then why,
    then next. Nothing is cut; a sentence that does not end one is closed with a full stop, so the
    next one reads apart from it."""
    said: list[str] = []
    for key in JOURNAL_SAYS:
        piece = " ".join(str(row.get(key) or "").split())
        if piece:
            said.append(piece if piece[-1] in ".!?:" else piece + ".")
    return " ".join(said)


def journal_line(row: dict) -> str:
    """One journal row in plain words: the Bahrain time, the step, what happened, and the writer's
    own sentences (journal_says). Its machine fields (process ids, paths, archives, snapshots,
    counts, seconds) stay in the file."""
    at = str(row.get("at") or "")
    try:
        clock = datetime.fromisoformat(at).astimezone(status.BAHRAIN).strftime("%H:%M:%S")
    except ValueError:
        clock = at[11:19] or "--:--:--"
    event = str(row.get("event") or "")
    word = EVENT_WORDS.get(event) or DELIVERY_WORDS.get(event.removeprefix("delivery_"), event)
    if event == "step_failed" and row.get("stops") is False:
        word = "needs a look (the load went on)"
    elif event == "finished" and row.get("result") == "ok_needs_look":
        word = "FINISHED, and something needs a look"
    elif event == "failed" and row.get("result") == "check_again":
        word = "STOPPED before anything was loaded: check again"
    head = f"{row['step']}: {word}" if row.get("step") else word
    said = journal_says(row)
    return f"  {clock}  {head}" + (f" - {said}" if said else "")


def _open_delivery_age(run: str, closers: tuple, now: datetime | None = None) -> float | None:
    """Seconds since the last 'queued' row of the run's delivery journal when no row of `closers` came
    after it; None when there is no such row. An unreadable time counts as old."""
    rows = loadplan.delivery_read(run)
    queued = [i for i, r in enumerate(rows) if r.get("event") == "queued"]
    if not queued or any(r.get("event") in closers for r in rows[queued[-1] + 1:]):
        return None
    try:
        return ((now or datetime.now(timezone.utc)) - datetime.fromisoformat(str(rows[queued[-1]].get("at")))).total_seconds()
    except (TypeError, ValueError):
        return float("inf")


def delivery_pending(run: str, now: datetime | None = None) -> bool:
    """Is the delivery of this finished load still being made: a 'queued' row with no 'ended' row after
    it (the writer's last word, which carries the `Delivery:` line), younger than DELIVERY_WAIT_S? An
    older one was cut off: see delivery_cut_off."""
    age = _open_delivery_age(run, ("ended",), now)
    return age is not None and age <= DELIVERY_WAIT_S


def delivery_cut_off(run: str, now: datetime | None = None) -> bool:
    """A delivery that was queued and never closed, older than DELIVERY_WAIT_S: the PC slept or the
    worker was ended before the e-mail went. A send by hand afterwards closes it too (DELIVERY_ENDS).
    `load status` says so, with its retry, which reloads nothing (or why there is none: scripts/ops/noretry.py)."""
    age = _open_delivery_age(run, DELIVERY_ENDS, now)
    return age is not None and age > DELIVERY_WAIT_S


def run_state(run: str, now: datetime | None = None) -> tuple[str, str | None]:
    """(RUNNING | FINISHED | FAILED | INTERRUPTED | NOT STARTED, step). The journal says whether the
    run started and how it ended; the lock says whether its worker is still alive. A load that has
    closed (FINISHED, or FAILED after it began to write) reads RUNNING at 'delivery' while its brief or its
    note and the e-mail are being made (deliver_after of the writer): the skill then follows it a little
    longer, and its `Delivery:` line is shown before the state it ended in."""
    result = loadplan.run_result(run)
    step = loadplan.last_step(run)
    if result in ("FINISHED", "FAILED") and delivery_pending(run, now):
        return "RUNNING", "delivery"
    if result in ("FINISHED", "FAILED"):
        return result, step
    lock = lock_now()
    if lock.get("rehearsal_left") == run:          # a rehearsal whose worker is gone: cut off (nothing to resume)
        return "INTERRUPTED", lock.get("step") or step
    if lock.get("state") in ("running", "interrupted") and lock.get("run") == run:
        if lock["state"] == "interrupted":
            return "INTERRUPTED", lock.get("step") or step
        return "RUNNING", lock.get("step") or step
    if result == "NOT_STARTED":
        return "NOT STARTED", None
    if lock.get("state") == "unreadable":
        return "RUNNING", step                     # a lock is there and busy at this moment: the next call reads it
    rows = loadplan.journal_read(run)              # a journal with no end and no live lock of its own
    now = now or datetime.now(timezone.utc)
    try:
        age = (now - datetime.fromisoformat(str(rows[-1].get("at")))).total_seconds()
    except (TypeError, ValueError):
        age = None
    if age is not None and age <= status.HEARTBEAT_STALE_S:
        return "RUNNING", step                     # the worker is between its first line and its lock
    return "INTERRUPTED", step


def state_text(state: str, step: str | None) -> str:
    return state + (f" at {step or 'start'}" if state in ("RUNNING", "FAILED", "INTERRUPTED") else "")


def _delivery_windows(rows: list[dict]) -> list[tuple[int, int | None]]:
    """(queued, ended) row indexes of each delivery the writer ran after the load; None = not ended yet."""
    out, start = [], None
    for i, r in enumerate(rows):
        if r.get("event") == "queued":
            if start is not None:
                out.append((start, None))
            start = i
        elif r.get("event") == "ended" and start is not None:
            out.append((start, i))
            start = None
    if start is not None:
        out.append((start, None))
    return out


def new_delivery_lines(run: str) -> list[str]:
    """The rows of the run's delivery journal no earlier call has shown: what became of its brief and its
    e-mail after the load closed. Of a delivery the writer runs (deliver_after: 'queued' ... 'ended') only
    its first and its last row are shown: the last one is the `Delivery:` line, and the rows in between
    (the brief, the sender's own row) wait unshown until it is written. A send by hand (scripts/yq_send.py)
    shows its own row."""
    rows = loadplan.delivery_read(run)
    mark = loadplan.run_dir(run) / DELIVERY_OFFSET_FILE
    try:
        shown = int(mark.read_text(encoding="utf-8").strip() or 0) if mark.is_file() else 0
    except ValueError:
        shown = 0
    shown = min(max(shown, 0), len(rows))
    windows = _delivery_windows(rows)
    working = delivery_pending(run) and windows and windows[-1][1] is None
    out, upto = [], shown
    for i in range(shown, len(rows)):
        r, event = rows[i], rows[i].get("event")
        inside = next(((s, e) for s, e in windows if s < i and (e is None or i <= e)), None)
        if inside and inside[1] is None:
            if working and inside == windows[-1]:
                break                               # the writer is still at work: its last row says what became of it
            inside = None                           # it was cut off: each row says for itself what became of the e-mail
        upto = i + 1
        if event == "ended":
            out.append("  " + " ".join(str(r.get("text") or "").split()))      # the `Delivery:` line, as written
        elif event == "queued" or (event in DELIVERY_WORDS and not inside):
            out.append(delivery_line(dict(r, step="delivery", event="delivery_" + str(event))))
    if upto != shown and loadplan.run_dir(run).is_dir():
        paths.write_text(mark, f"{upto}\n")
    return out


def delivery_line(row: dict) -> str:
    """One row of the delivery journal, like journal_line, but its `next` (the line that sends the e-mail
    again) is left as written, with no full stop after it: it is run as it stands."""
    nxt = " ".join(str(row.get("next") or "").split())
    return journal_line(dict(row, next=None)) + (" " + nxt if nxt else "")


def timing_line(run: str) -> str | None:
    """One line on how long a closed load took, from its closing row (the writer keeps the seconds of
    every step there)."""
    closing = [r for r in loadplan.journal_read(run) if r.get("event") in ("finished", "failed")]
    if not closing or not isinstance(closing[-1].get("timings"), dict):
        return None
    t, total = closing[-1]["timings"], closing[-1].get("total_seconds")
    parts = [f"{k} {round(float(v))} s" for k, v in t.items() if isinstance(v, (int, float))]
    head = f"  Load time: {round(float(total) / 60, 1)} min" if isinstance(total, (int, float)) else "  Load time"
    return head + " (" + ", ".join(parts) + ")"


def new_journal_lines(run: str) -> list[str]:
    """The journal rows no earlier call has shown (a small count is kept in the run folder)."""
    rows = loadplan.journal_read(run)
    mark = loadplan.run_dir(run) / OFFSET_FILE
    try:
        shown = int(mark.read_text(encoding="utf-8").strip() or 0) if mark.is_file() else 0
    except ValueError:
        shown = 0
    shown = min(max(shown, 0), len(rows))
    if len(rows) != shown and loadplan.run_dir(run).is_dir():
        paths.write_text(mark, f"{len(rows)}\n")
    return [journal_line(r) for r in rows[shown:]]


def locked_run() -> str | None:
    """The run that holds the load lock now (running or cut off), if any."""
    lock = lock_now()
    run = str(lock.get("run") or "")
    return run if lock.get("state") in ("running", "interrupted") and loadplan.is_run(run) else None


def run_status(args, deps: Deps | None = None) -> dict:
    """The run to follow: the one named; else the one that holds the lock (the load that is running
    or was cut off is the one he asks about, whatever was checked after it); else the latest."""
    deps = deps or Deps()
    run = getattr(args, "run", None) or locked_run() or loadplan.latest_run()
    if not run:
        return {"ok": True, "command": "load", "data_date": None, "exit": 0, "flags": ["not_started"],
                "text": "No load has been checked yet.\nNOT STARTED", "state": "NOT STARTED", "run": None}
    if not loadplan.is_run(run):
        return {"ok": False, "command": "load", "data_date": None, "exit": 2, "flags": ["stop"],
                "text": f"'{run}' is not a run name (ddmmyy_HHMM).", "state": None, "run": run}
    lines: list[str] = []
    began = deps.clock()
    timed = False
    while True:
        lines += new_journal_lines(run)
        if not timed and loadplan.run_result(run) in ("FINISHED", "FAILED"):
            said = timing_line(run)
            if said and not (loadplan.run_dir(run) / TIMING_SHOWN_FILE).is_file():
                lines.append(said)
                paths.write_text(loadplan.run_dir(run) / TIMING_SHOWN_FILE, "1\n")
            timed = True
        lines += new_delivery_lines(run)
        state, step = run_state(run)
        waited = deps.clock() - began
        final = (state in ("FINISHED", "FAILED", "INTERRUPTED")
                 or (state == "NOT STARTED" and waited >= WAIT_START_GRACE_S))
        if not getattr(args, "wait", False) or final or waited >= WAIT_MAX_S:
            break
        deps.pause(WAIT_POLL_S)
    if state in ("FINISHED", "FAILED") and delivery_cut_off(run):
        # no retry where the sender would refuse it (an older load's brief), in its words: scripts/ops/noretry.py
        from scripts.ops.noretry import no_retry_needed
        lines.append("  Delivery: cut off before the e-mail went (the PC slept, or the load's window was closed). "
                     "Nothing was reloaded. " + (no_retry_needed(run) or f"Its retry: {SEND_RETRY.format(run=run)}"))
    text = [f"Load {run}"]
    newer = (read_plan_of(run) or {}).get("superseded_by")
    if newer and state == "NOT STARTED":
        text.append(f"  (this check was replaced by a newer one, run {newer}: its go line no longer works)")
    text += (lines or ["  (no new journal line)"]) + [state_text(state, step)]
    good = state in ("RUNNING", "FINISHED", "NOT STARTED")
    closing = [r for r in loadplan.journal_read(run) if r.get("event") in ("finished", "failed")]
    result = closing[-1].get("result") if closing else None        # ok | ok_needs_look | failed | check_again
    flags = [state.lower().replace(" ", "_")]
    if result and str(result) not in flags:
        flags.append(str(result))
    return {"ok": good, "command": "load", "data_date": None, "text": "\n".join(text), "exit": 0 if good else 2,
            "flags": flags, "run": run, "state": state, "step": step, "result": result}


# ── the command ───────────────────────────────────────────────────────────────

AGAIN = "again"      # the writer's too-fast answer says "Type /yq-load <ddmmyy> again.": he may type the word too


def drop_asked(args) -> tuple[str | None, str | None]:
    """(the drop folder a check is for, why the words cannot be taken) from the words after `check`. Taken: no
    word (the latest folder); the folder (ddmmyy, or a path); the folder and then the one word 'again', in any
    case, which is ignored; 'again' alone, which is /yq-load alone (he once typed exactly that after "Type
    /yq-load to check again."). Refused, with what to type: 'again' first and anything after it ('again again',
    'again <ddmmyy>'), and any other word after the folder. Nothing is guessed."""
    drop, extra = getattr(args, "drop", None), getattr(args, AGAIN, None)
    first_is_again = drop is not None and str(drop).strip().lower() == AGAIN
    if extra is not None and first_is_again:
        return None, ("'again' comes after the folder: type /yq-load <ddmmyy> again, or /yq-load alone for the latest "
                      "folder. Nothing was checked.")
    if extra is not None and str(extra).strip().lower() != AGAIN:
        return None, (f"'{str(extra)[:40]}' is not a word the check takes after the folder (only '{AGAIN}' is). "
                      "Nothing was checked.")
    return (None if first_is_again else drop), None


def add_arguments(parser) -> None:
    parser.add_argument("action", choices=("check", "status"), help="check a drop folder, or follow a running load")
    parser.add_argument("drop", nargs="?", help="check: the drop folder as ddmmyy (default: the latest) or a path")
    parser.add_argument(AGAIN, nargs="?", help="check: the word 'again' after the folder (as the load program's "
                                               "answer words it): taken and ignored; no other word is")
    parser.add_argument("--pricebook", choices=("load", "skip"),
                        help="check: the owner's decision on the price book (load it / leave it out)")
    parser.add_argument("--with-profitability", action="store_true",
                        help="check: load the profitability report although the check would hold it back")
    parser.add_argument("--replace-ok", action="store_true",
                        help="check: the sales files are right although invoices would go (writes REPLACE_OK)")
    parser.add_argument("--partial-ok", action="store_true",
                        help="check: a narrower stock balance is intended (writes PARTIAL_OK)")
    parser.add_argument("--dev", action="store_true",
                        help="check: a proof run from a code worktree: no token, no GO line, nothing recorded, "
                             "no earlier check replaced")
    parser.add_argument("--run", help="status: the run to follow (default: the one that holds the lock, else the "
                                      "latest)")
    parser.add_argument("--wait", action="store_true", help="status: wait up to 100 s for the load to end")


def run(args) -> dict:
    if args.action == "status":
        return run_status(args)
    return run_check(args)
