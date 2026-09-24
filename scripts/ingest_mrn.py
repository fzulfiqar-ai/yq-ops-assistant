"""Ingest Material Receipt Note (MRN) exports → real landed costs.

Each Transactions_*.xml is a Focus MRN. For every received line we take the TRUE all-in landed cost
per unit = StockValue ÷ Quantity (StockValue already includes vendor freight + 3rd-party freight +
customs + misc), plus the pre-freight product cost = Gross ÷ Quantity. Costs are keyed on the FULL
ProdCode (e.g. 'X01 UC') so cable variants never collapse. When a SKU appears in several MRNs we
keep the most RECENT receipt.

Receipt dates (R2, 24-Sep-2026). A cost row is dated by the voucher's REAL receipt date, never the
1st of its month: the ledger's own move date for 'MRN:<doc>' (stock_movements) first, then the
XML header's Date field (Focus packs it as year<<16 | month<<8 | day: 132778260 = 2026-09-20),
and only when neither is known the 1st of the doc month (flagged in the summary). D4: rows dated
'2026-09-01' lost to a stale extract dated 2026-09-14, so v_product_economics costed 16 of the 20
LC1716 SKUs from the wrong source.
Duplicate lines: a receipt that lists the same SKU twice (split boxes, two lots) is folded to one
line per (doc, SKU) with the quantities summed and a QUANTITY-WEIGHTED unit cost (Decimal, 4 dp,
ROUND_HALF_UP) -- the first line's cost is no longer taken for the whole quantity.

CLI:    python -m scripts.ingest_mrn                      # bulk-load the XML folders (writes)
        python -m scripts.ingest_mrn --from-workbook "…/YQ_MRN_Landing_Cost_Analysis (2).xlsx" --sheet YQ-25-11-1
                                                          # the owner's workbook for a receipt whose XML
                                                          # never reached disk: DRY RUN by default,
                                                          # prints what it would write; add --commit
API:    parse_mrn_bytes() + load_mrn_costs()              # used by POST /material-receipts/upload
"""
from __future__ import annotations

import argparse
import glob
import os
import re
import sys
from datetime import date
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

# Every place an MRN export has been dropped. Receipts YQ-26-06-1/-2 sat under Shipments/ and
# New order*/ and were missed by the old single glob, so a re-run "forgot" them and let the
# February receipts overwrite their newer landed costs (21-Sep-2026). Same DocNo in two folders
# is loaded once.
MRN_GLOBS = (
    "business_data/Order received/*/Transactions_*.xml",
    "business_data/Shipments/**/Transactions_*.xml",
    "business_data/New order*/Transactions_*.xml",
)

_Q4 = Decimal("0.0001")


def _doc_month(doc_no: str | None) -> str | None:
    """YQ-26-02-2 → '2026-02-01' (the LAST-resort date when no receipt date is known)."""
    m = re.match(r"YQ-(\d{2})-(\d{2})-", doc_no or "")
    return f"20{int(m.group(1)):02d}-{int(m.group(2)):02d}-01" if m else None


def decode_focus_date(v) -> str | None:
    """Pure: a Focus date serial (year<<16 | month<<8 | day) → ISO date, or None when it is not one.
    Header/Date 132778260 → 2026-09-20 (checked against the ledger's MRN:YQ-26-09-2 move date)."""
    try:
        n = int(float(str(v).strip()))
    except (TypeError, ValueError):
        return None
    y, m, d = n >> 16, (n >> 8) & 0xFF, n & 0xFF
    if not (2000 <= y <= 2100):
        return None
    try:
        return date(y, m, d).isoformat()
    except ValueError:
        return None


def _parse_root(root: ET.Element) -> list[dict]:
    out: list[dict] = []
    for trans in root.iter("Transaction"):
        header = trans.find("Header")
        mrn_no = header.findtext("DocNo") if header is not None else None
        xml_date = decode_focus_date(header.findtext("Date")) if header is not None else None
        # the MRN can be raised under a DIFFERENT number than the PO (received in a later month) —
        # follow the linked PO number from HeaderExtra so received lines attach to the right order.
        po_no = None
        hx = trans.find("HeaderExtra")
        if hx is not None:
            for pair in hx.findall("IdNamePair"):
                if (pair.findtext("Name") or "") == "PONo":
                    po_no = (pair.findtext("Tag") or "").replace("PO:", "").strip() or None
        doc_no = po_no or mrn_no          # KEY received data by the PO number (the order key)
        body = trans.find("BodyData")
        if body is None:
            continue
        for tb in body.findall("TransBody"):
            sales = tb.find("Sales")
            if sales is None:
                continue
            code = (sales.findtext("ProdCode") or "").strip()
            qty = float(sales.findtext("Quantity") or 0)
            stockval = float(sales.findtext("StockValue") or 0)   # all-in (incl. vendor freight)
            gross = float(sales.findtext("Gross") or 0)           # supplier price only
            if not code or qty <= 0:
                continue
            out.append({"code": code, "landed": stockval / qty, "product": gross / qty,
                        "qty": qty, "doc_no": doc_no, "mrn_no": mrn_no,
                        "eff": xml_date or _doc_month(mrn_no), "eff_source": "xml" if xml_date else "month",
                        "xml_date": xml_date})
    return out


def parse_mrn(path: str) -> list[dict]:
    return _parse_root(ET.parse(path).getroot())


def parse_mrn_bytes(data: bytes) -> list[dict]:
    return _parse_root(ET.fromstring(data))


def _d(v) -> Decimal:
    return Decimal(str(v if v is not None else 0))


def fold_duplicate_lines(rows: list[dict]) -> list[dict]:
    """Pure: one row per (doc_no, code): qty summed, landed / product unit cost = the QUANTITY-WEIGHTED
    average across the lines (Decimal, 4 dp, ROUND_HALF_UP). Every other field comes from the first
    line. Order of first appearance is kept."""
    acc: dict[tuple, dict] = {}
    for r in rows:
        key = (r.get("doc_no"), str(r.get("code") or "").upper())
        q = _d(r.get("qty"))
        if key not in acc:
            acc[key] = {**r, "_qty": q, "_landed": q * _d(r.get("landed")), "_product": q * _d(r.get("product")),
                        "lines": 1}
        else:
            a = acc[key]
            a["_qty"] += q
            a["_landed"] += q * _d(r.get("landed"))
            a["_product"] += q * _d(r.get("product"))
            a["lines"] += 1
    out: list[dict] = []
    for a in acc.values():
        q = a.pop("_qty")
        landed, product = a.pop("_landed"), a.pop("_product")
        a["qty"] = float(q)
        if q > 0:
            a["landed"] = float((landed / q).quantize(_Q4, rounding=ROUND_HALF_UP))
            a["product"] = float((product / q).quantize(_Q4, rounding=ROUND_HALF_UP))
        out.append(a)
    return out


def dedupe_latest(rows: list[dict]) -> list[dict]:
    """Keep the most recent receipt per SKU (so current cost = latest landed cost)."""
    best: dict[str, dict] = {}
    for r in rows:
        k = r["code"].upper()
        if k not in best or (r["eff"] or "") >= (best[k]["eff"] or ""):
            best[k] = r
    return list(best.values())


def ledger_receipt_dates(client, mrn_nos) -> dict[str, str]:
    """{mrn_no: first move date} of the 'MRN:<doc>' vouchers in stock_movements -- the real receipt
    date. A voucher the ledger does not carry is simply absent (select-only; never raises)."""
    out: dict[str, str] = {}
    for doc in sorted({d for d in mrn_nos if d}):
        try:
            rows = (client.table("stock_movements").select("move_date").eq("voucher", f"MRN:{doc}")
                    .order("move_date").limit(1).execute().data or [])
        except Exception:  # noqa: BLE001
            rows = []
        if rows and rows[0].get("move_date"):
            out[doc] = str(rows[0]["move_date"])[:10]
    return out


def apply_receipt_dates(rows: list[dict], ledger: dict[str, str]) -> dict[str, tuple[str, str]]:
    """Set every row's `eff` to its receipt date: ledger move date > XML header date > doc month.
    Returns {mrn_no: (date, source)} for the summary."""
    resolved: dict[str, tuple[str, str]] = {}
    for r in rows:
        mrn = r.get("mrn_no") or r.get("doc_no")
        if mrn in ledger:
            r["eff"], r["eff_source"] = ledger[mrn], "ledger"
        elif r.get("xml_date"):
            r["eff"], r["eff_source"] = r["xml_date"], "xml"
        else:
            r["eff"], r["eff_source"] = _doc_month(mrn), "month"
        if mrn and r["eff"]:
            resolved[mrn] = (r["eff"], r["eff_source"])
    return resolved


def build_payloads(rows: list[dict], have: dict[str, str]) -> dict:
    """Pure: the three table payloads from dated, folded rows. `have` = {SKU upper: effective_date}
    already in mrn_landed_costs -- a SKU's current cost is never replaced by an OLDER receipt."""
    final = dedupe_latest(rows)
    kept, skipped = [], []
    for r in final:
        if (r["eff"] or "") < have.get(r["code"].upper(), ""):
            skipped.append(r["code"])
            continue
        kept.append(r)
    cost_payload = [{
        "sku_code": r["code"], "landed_cost_bhd": round(r["landed"], 4),
        "product_cost_bhd": round(r["product"], 4), "last_qty": r["qty"],
        "doc_no": r["doc_no"], "effective_date": r["eff"],
    } for r in kept]
    lines_payload = [{
        "doc_no": r["doc_no"], "sku_code": r["code"], "qty": r["qty"],
        "landed_unit_bhd": round(r["landed"], 4), "product_unit_bhd": round(r["product"], 4),
    } for r in rows if r.get("doc_no")]
    # Data rule 1+2: purchase_costs is the cost source of truth and is VERSIONED by effective_date
    # ("latest cost" = MAX(id) per SKU). Mirror the newest receipt into it as a dated snapshot --
    # only the newest, because an older receipt inserted later would get a higher id and win the
    # MAX(id) rule with a stale cost. Dated by the real receipt date. ON CONFLICT DO NOTHING keeps
    # history intact on re-runs (idempotent).
    newest = max((r["eff"] or "") for r in kept) if kept else ""
    pc_payload = [{
        "sku_code": r["code"], "landed_cost_bhd": round(r["landed"], 4), "currency": "BHD",
        "effective_date": newest, "source_file": f"MRN {r['doc_no']}",
    } for r in kept if r["eff"] == newest and newest]
    return {"mrn_landed_costs": cost_payload, "mrn_lines": lines_payload, "purchase_costs": pc_payload,
            "skipped_older": skipped, "newest": newest}


def load_mrn_costs(rows: list[dict], client=None, dry_run: bool = False) -> dict:
    """Upsert via the Supabase client (used by the API): the deduped current-cost lookup
    (`mrn_landed_costs`), the full per-receipt detail (`mrn_lines`) and the dated purchase_costs
    mirror. dry_run: resolve dates and build every payload, write nothing. Returns a summary."""
    if client is None:
        from app.database import get_client
        client = get_client()

    dates = apply_receipt_dates(rows, ledger_receipt_dates(client, {r.get("mrn_no") or r.get("doc_no") for r in rows}))
    folded = fold_duplicate_lines(rows)
    # current cost = latest receipt per SKU -- and never older than what the table already holds.
    # The API path uploads ONE file at a time, and a bulk run may not see every folder: without
    # this guard an older receipt silently replaced a newer landed cost for the same SKU.
    try:
        have = {r["sku_code"].upper(): (r.get("effective_date") or "")
                for r in (client.table("mrn_landed_costs").select("sku_code,effective_date").execute().data or [])}
    except Exception:  # noqa: BLE001 -- table may not exist yet on a fresh DB
        have = {}
    p = build_payloads(folded, have)
    if p["skipped_older"]:
        print(f"  (kept {len(p['skipped_older'])} newer landed costs already in mrn_landed_costs)")
    if not dry_run:
        if p["mrn_landed_costs"]:
            client.table("mrn_landed_costs").upsert(p["mrn_landed_costs"], on_conflict="sku_code").execute()
        if p["mrn_lines"]:
            client.table("mrn_lines").upsert(p["mrn_lines"], on_conflict="doc_no,sku_code").execute()
        if p["purchase_costs"]:
            client.table("purchase_costs").upsert(p["purchase_costs"], on_conflict="sku_code,effective_date",
                                                  ignore_duplicates=True).execute()
    docs = sorted({r["doc_no"] for r in folded if r["doc_no"]})
    month_dated = sorted({m for m, (_, src) in dates.items() if src == "month"})
    return {"skus": len(p["mrn_landed_costs"]), "lines": len(rows), "folded_lines": len(folded), "docs": docs,
            "receipt_dates": {k: v[0] for k, v in dates.items()},
            "date_sources": {k: v[1] for k, v in dates.items()},
            "month_dated": month_dated,           # receipts whose real date is unknown (1st of month used)
            "purchase_costs_month": p["newest"], "purchase_costs_rows": len(p["purchase_costs"]),
            "payloads": p if dry_run else None, "dry_run": dry_run}


# ── the owner's landed-cost workbook (a receipt whose XML never reached disk) ──

WB_COLS = {"code": "item code", "qty": "qty", "landed": "landing cost / unit", "product": "purchase rate"}


def parse_workbook_sheet(path, sheet: str) -> list[dict]:
    """Rows of one MRN sheet of 'YQ_MRN_Landing_Cost_Analysis' in the XML row shape: Item Code, Qty,
    'Landing Cost / Unit' (= landed), 'Purchase Rate' (= product). The header row is found by its
    labels; reading stops at the TOTAL line. doc_no = mrn_no = the sheet name."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f"sheet '{sheet}' not in workbook (sheets: {', '.join(wb.sheetnames)})")
        ws = wb[sheet]
        cols: dict[str, int] | None = None
        out: list[dict] = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in row]
            if cols is None:
                low = [c.lower() for c in cells]
                found = {}
                for k, label in WB_COLS.items():
                    for j, c in enumerate(low):
                        if c.startswith(label):
                            found[k] = j
                            break
                if len(found) == len(WB_COLS):
                    cols = found
                continue
            code = cells[cols["code"]] if cols["code"] < len(cells) else ""
            first = cells[0].lower() if cells else ""
            if not code or first == "total" or code.lower() == "total":
                if first == "total":
                    break
                continue
            try:
                qty = float(str(row[cols["qty"]]).replace(",", ""))
                landed = float(str(row[cols["landed"]]).replace(",", ""))
                product = float(str(row[cols["product"]]).replace(",", ""))
            except (TypeError, ValueError):
                continue
            if qty <= 0:
                continue
            out.append({"code": code, "landed": landed, "product": product, "qty": qty,
                        "doc_no": sheet, "mrn_no": sheet, "eff": _doc_month(sheet), "eff_source": "month",
                        "xml_date": None})
        if cols is None:
            raise ValueError(f"sheet '{sheet}': header with {list(WB_COLS.values())} not found")
        return out
    finally:
        wb.close()


def floor_gains(codes, costed_upper: set[str]) -> list[str]:
    """Pure: the SKUs (as written) that have NO cost anywhere yet -- the ones a receipt gives a
    margin floor to (app/shop.py price_cart caps discounts only on covered lines)."""
    seen, out = set(), []
    for c in codes:
        k = str(c or "").strip()
        if k and k.upper() not in costed_upper and k.upper() not in seen:
            seen.add(k.upper())
            out.append(k)
    return out


def _costed_codes(client) -> set[str]:
    """Upper-cased SKUs with a cost in mrn_landed_costs or purchase_costs (> 0). Select-only."""
    out: set[str] = set()
    try:
        for r in (client.table("mrn_landed_costs").select("sku_code,landed_cost_bhd").limit(5000).execute().data or []):
            if r.get("sku_code") and (r.get("landed_cost_bhd") or 0):
                out.add(str(r["sku_code"]).strip().upper())
    except Exception as e:  # noqa: BLE001
        print(f"  (mrn_landed_costs unreadable: {str(e)[:100]})")
    try:
        off = 0
        while True:
            page = (client.table("purchase_costs").select("sku_code,landed_cost_bhd").order("id")
                    .range(off, off + 999).execute().data or [])
            for r in page:
                if r.get("sku_code") and float(r.get("landed_cost_bhd") or 0) > 0:
                    out.add(str(r["sku_code"]).strip().upper())
            if len(page) < 1000:
                break
            off += 1000
    except Exception as e:  # noqa: BLE001
        print(f"  (purchase_costs unreadable: {str(e)[:100]})")
    return out


def _active_catalog_codes(client) -> set[str]:
    try:
        rows = client.table("catalog_items").select("item_code").eq("is_active", True).limit(5000).execute().data or []
        return {str(r["item_code"]).strip().upper() for r in rows if r.get("item_code")}
    except Exception:  # noqa: BLE001
        return set()


def workbook_mode(path: str, sheet: str, commit: bool) -> int:
    from dotenv import load_dotenv
    load_dotenv(ROOT / ".env")
    from app.database import get_client
    client = get_client()
    rows = parse_workbook_sheet(path, sheet)
    if not rows:
        print(f"No lines on sheet {sheet}.")
        return 2
    costed_before = _costed_codes(client)
    active = _active_catalog_codes(client)
    summary = load_mrn_costs(rows, client=client, dry_run=not commit)
    p = summary["payloads"] or {}
    eff, src = summary["receipt_dates"].get(sheet), summary["date_sources"].get(sheet)
    print(f"MRN {sheet} from workbook sheet: {len(rows)} lines, {summary['folded_lines']} after folding duplicates, "
          f"{sum(r['qty'] for r in rows):,.0f} units")
    print(f"  receipt date {eff} (source: {src}{'' if src != 'month' else ' -- NOT in the ledger and no XML date; 1st of month'})")
    if commit:
        print("  WRITTEN: mrn_landed_costs %d, mrn_lines %d, purchase_costs %d (dated %s)" % (
            summary["skus"], len(rows), summary["purchase_costs_rows"], summary["purchase_costs_month"]))
    else:
        print(f"  would write mrn_landed_costs: {len(p.get('mrn_landed_costs', []))} rows"
              + (f" (skipped {len(p['skipped_older'])} SKUs with a newer receipt on file: "
                 + ", ".join(p["skipped_older"][:12]) + (" ..." if len(p["skipped_older"]) > 12 else "") + ")"
                 if p.get("skipped_older") else ""))
        for r in p.get("mrn_landed_costs", []):
            print(f"      {r['sku_code']:28} landed {r['landed_cost_bhd']:>9.4f}  product {r['product_cost_bhd']:>9.4f}  "
                  f"qty {r['last_qty']:>7,.0f}  {r['effective_date']}")
        print(f"  would write mrn_lines: {len(p.get('mrn_lines', []))} rows (doc {sheet})")
        print(f"  would write purchase_costs: {len(p.get('purchase_costs', []))} rows dated {p.get('newest')} "
              f"(source 'MRN {sheet}', existing keys untouched)")
    gains = floor_gains([r["code"] for r in rows], costed_before)
    if gains:
        print(f"  SKUs that GAIN a margin floor (no cost anywhere today): {len(gains)}")
        for g in gains:
            print(f"      {g:28} {'active catalog code' if g.upper() in active else 'not an active catalog code'}")
    else:
        print("  every SKU on this sheet already has a cost on file; no new margin floor")
    if not commit:
        print("(dry run -- nothing written; add --commit to write)")
    return 0


def main(argv: list[str] | None = None) -> None:
    ap = argparse.ArgumentParser(description="MRN landed costs: XML bulk load, or one sheet of the owner's workbook")
    ap.add_argument("--from-workbook", help="path to YQ_MRN_Landing_Cost_Analysis (…).xlsx")
    ap.add_argument("--sheet", help="the MRN sheet to load, e.g. YQ-25-11-1")
    ap.add_argument("--commit", action="store_true", help="workbook mode: write (default is a dry run)")
    args = ap.parse_args(sys.argv[1:] if argv is None else argv)
    if args.from_workbook:
        if not args.sheet:
            ap.error("--sheet is required with --from-workbook")
        raise SystemExit(workbook_mode(args.from_workbook, args.sheet, args.commit))

    import psycopg
    from dotenv import load_dotenv
    load_dotenv()

    files = sorted({f for g in MRN_GLOBS for f in glob.glob(g, recursive=True)})
    if not files:
        print("No MRN XML files found under business_data/ (Order received, Shipments, New order*).")
        return
    rows: list[dict] = []
    seen_docs: dict[str, str] = {}
    for f in files:
        try:
            got = parse_mrn(f)
        except Exception as e:  # noqa: BLE001
            print(f"  skip {f}: {type(e).__name__}: {e}")
            continue
        doc = got[0]["doc_no"] if got else None
        if doc and doc in seen_docs:
            print(f"  dup  {os.path.basename(os.path.dirname(f))}: {doc} already loaded from {seen_docs[doc]}")
            continue
        if doc:
            seen_docs[doc] = os.path.basename(os.path.dirname(f))
        rows.extend(got)
        print(f"  {os.path.basename(os.path.dirname(f))}: {len(got)} lines  ({os.path.basename(f)}) {doc or ''}")
    # ensure both tables exist, then reuse the shared (Supabase) loader for the upserts
    with psycopg.connect(os.environ["DATABASE_URL"], autocommit=True) as conn:
        conn.execute(open("scripts/mrn_costs_migration.sql", encoding="utf-8").read())
        conn.execute(open("scripts/mrn_lines_migration.sql", encoding="utf-8").read())
    summary = load_mrn_costs(rows)
    print(f"\nLoaded {summary['skus']} unique SKU landed costs + per-receipt lines "
          f"from {len(files)} MRN files ({summary['lines']} total lines).")
    for doc, d in sorted(summary["receipt_dates"].items()):
        print(f"  {doc}: receipt date {d} ({summary['date_sources'].get(doc)})")
    if summary["month_dated"]:
        print(f"  !! {len(summary['month_dated'])} receipt(s) dated to the 1st of the month (no ledger row, no XML date): "
              + ", ".join(summary["month_dated"]))


if __name__ == "__main__":
    main()
