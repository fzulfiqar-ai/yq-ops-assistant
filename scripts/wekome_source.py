"""WEKOME opening-order source: the invoice parser, the curated names/specs and the photo cleaner.

COPIED (24-Sep-2026, release R1b) from the pure, DB-free pieces of the untracked
scripts/make_management_reports.py — wps_cell_images, model_code, parse_variant, read_invoice,
build_wekome_products, price_list_picture, WK_CATALOG, WK_CATEGORY_ORDER, WK_PHOTO_FRAME,
WK_BOX_OVERRIDE, clean_photo, CONN, VARIANT_ORDER, NS — so scripts/import_upcoming.py can build
the "Coming soon" cards without the report script's DB/PDF/Excel half. The original is not
modified; keep the two in step by hand (a diff of the function bodies should be empty).

The shipment workbooks live in Wekome/Wekome Shipments/ (gitignored, never in the repo).
YQ_WEKOME_DIR overrides that folder (a worktree has no Wekome/ of its own).

Only the newest invoice copy is read (…1st(3).xlsx; the 17-Aug …1st.xlsx is ignored) and
shipment 2 merges into the same models. The Cost_Profit_Review workbook is never opened here.
"""
from __future__ import annotations

import io
import os
import re
import zipfile
from collections import OrderedDict
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

from PIL import Image, ImageChops, ImageOps

ROOT = Path(__file__).resolve().parents[1]

WK_DIR = Path(os.getenv("YQ_WEKOME_DIR") or (ROOT / "Wekome" / "Wekome Shipments"))
WK_SHIP1 = WK_DIR / "Copy of CI & PL INV#AS2026072701 - 20260812 updating -partial shipment 1st(3).xlsx"
WK_SHIP2 = WK_DIR / "Pending order" / "Copy of CI & PL INV#AS2026072701 - 20260902 updating -partial shipment 2nd(4).xlsx"
# The WK agent price list sits beside the shipments folder (Wekome/); only one packaging picture
# is read from it (WK_BOX_OVERRIDE), straight from the zip.
WK_PRICE_LIST = WK_DIR.parent / "WK Agent Full Price List on Sep.12 of 2026.xlsx"


# ── images ─────────────────────────────────────────────────────────────────────

def clean_photo(raw: bytes, box: int, frame: tuple[float, float] | None = None) -> bytes:
    """Flatten onto white, trim the empty margin, centre on a square with a little air,
    so every product sits at the same visual size. frame=(top, bottom) crops fractional
    bands first (vendor mark / caption strips). Returns JPEG bytes."""
    im = ImageOps.exif_transpose(Image.open(io.BytesIO(raw)))
    if frame:
        im = im.crop((0, int(im.height * frame[0]), im.width, int(im.height * (1 - frame[1]))))
    if im.mode in ("RGBA", "LA", "P"):
        im = im.convert("RGBA")
        flat = Image.new("RGB", im.size, (255, 255, 255))
        flat.paste(im, mask=im.split()[3])
        im = flat
    im = im.convert("RGB")
    diff = ImageChops.difference(im, Image.new("RGB", im.size, (255, 255, 255))).convert("L")
    bbox = diff.point(lambda v: 255 if v > 18 else 0).getbbox()
    if bbox:
        im = im.crop(bbox)
    side = int(max(im.size) * 1.1)
    square = Image.new("RGB", (side, side), (255, 255, 255))
    square.paste(im, ((side - im.width) // 2, (side - im.height) // 2))
    square.thumbnail((box, box), Image.LANCZOS)
    buf = io.BytesIO()
    square.save(buf, "JPEG", quality=92, optimize=True)
    return buf.getvalue()


# ══ WEKOME — new vendor opening order ══════════════════════════════════════════

# Names/specs condensed from "WK Agent Full Price List on Sep.12 of 2026" (col A name, col N spec).
# WDC-68 is not on that list; its name comes from the invoice + the May 2026 sample request.
WK_CATALOG: dict[str, tuple[str, str, str]] = {
    "WG-08":    ("Wireless Audio", "Neckband Wireless Earphones", "Bluetooth 6.0 · 200 mAh · up to 25 h playback"),
    "WS-33":    ("Wireless Audio", "TWS Wireless Earbuds", "Bluetooth 6.0 · 13 mm driver · about 4 h per charge"),
    "WS-31":    ("Wireless Audio", "Gen 3 TWS Wireless Earbuds", "Bluetooth 5.3 · 13 mm driver · 4–5 h per charge"),
    "WS-55":    ("Wireless Audio", "WK Design TWS Wireless Earbuds (Gen 5)", "Bluetooth 5.3 · 230 mAh case · about 4 h per charge"),
    "WS-56":    ("Wireless Audio", "WK Design TWS Wireless Earbuds Youth (Gen 6)", "Bluetooth 5.3 · 250 mAh case · about 4 h per charge"),
    "WS-57":    ("Wireless Audio", "WK Design TWS Wireless Earbuds (Gen 7)", "Bluetooth 5.3 · 250 mAh case · about 4 h per charge"),
    "WS-58":    ("Wireless Audio", "WK Design TWS Wireless Earbuds (Gen 8)", "Bluetooth 5.3 · 240 mAh case · about 4 h per charge"),
    "YB13":     ("Wired Earphones", "Wired Earphones", "10 mm driver · 1.2 m cable"),
    "YB03":     ("Wired Earphones", "Kingkong Series Wired Earphones", "14 mm driver · 1.2 m cable"),
    "YB05":     ("Wired Earphones", "Kingkong Series Wired Earphones", "14 mm driver · 1.2 m cable"),
    "WD-05":    ("Speakers", "Sleep Speaker with Built-in White Noise", "Bluetooth 5.3 · 5 W · 1200 mAh · 6 h+ playback"),
    "WDC-C37":  ("Data Cables", "Fast Charging Cable with Stand 240W (USB-C to USB-C)", "Braided + zinc alloy"),
    "WDC-C38":  ("Data Cables", "Fast Charging Cable with Stand 27W (USB-C to Lightning)", "Braided + zinc alloy"),
    "WDC-C34a": ("Data Cables", "Fast Charging Data Cable 66W (USB-A to USB-C)", "TPE"),
    "WDC-C34i": ("Data Cables", "Fast Charging Data Cable 12W (USB-A to Lightning)", "TPE"),
    "WDC-C34m": ("Data Cables", "Fast Charging Data Cable 12W (USB-A to Micro-USB)", "TPE"),
    "WDC-C35":  ("Data Cables", "Fast Charging Data Cable PD 65W (USB-C to USB-C)", "TPE"),
    "WDC-C31":  ("Data Cables", "Fast Charging Data Cable 66W (USB-A to USB-C)", "Nylon braided + aluminium alloy"),
    "WDC-C32":  ("Data Cables", "Fast Charging Data Cable 65W (USB-C to USB-C)", "Nylon braided + aluminium alloy"),
    "WDC-C33":  ("Data Cables", "Fast Charging Data Cable 35W (USB-C to Lightning)", "Nylon braided + aluminium alloy"),
    "WDC-67":   ("Data Cables", "Kingkong PVC Data Cable 66W (USB-A to USB-C)", "PVC"),
    "WDC-68":   ("Data Cables", "Kingkong PVC Data Cable 65W (USB-C to USB-C)", "PVC"),
    "WDC-21":   ("Data Cables", "Original Super Fast Charging Data Cable", "ABS + braided nylon"),
    "WDC-97":   ("Data Cables", "Kingkong Braided Fast Charging Cable 6A (USB-A to USB-C)", "Braided + TPE + aluminium alloy"),
    "WDC-98":   ("Data Cables", "Kingkong Braided Fast Charging Cable 65W (USB-C to USB-C)", "Braided + TPE + aluminium alloy"),
    "WDC-50":   ("Data Cables", "Braided 240W 4-in-1 Super Fast Charging Cable", "USB-A / USB-C to USB-C / Lightning · zinc alloy"),
    "WDC-51":   ("Data Cables", "Silicone 240W 4-in-1 Super Fast Charging Cable", "USB-A / USB-C to USB-C / Lightning · liquid silicone"),
    "WP-U100":  ("Wall Chargers", "PD 20W Fast Charger (UK plug)", "Single USB-C · PD 20W max"),
    "WP-U58":   ("Wall Chargers", "Flash Charge Series Charger 12W (UK plug)", "Dual USB-A · 12W max"),
    "WP-C56":   ("Car Chargers", "Transparent Car Charger 15W", "Dual USB-A · aluminium alloy + PC"),
    "WP-C53":   ("Car Chargers", "Transparent Car Charger 48W", "USB-A 18W + USB-C PD 30W"),
    "WP-C52":   ("Car Chargers", "MP3 Wireless Fast Charging Car Charger", "USB-A QC3.0 18W + USB-C 20W · MP3 player"),
    "WTP-137":  ("Screen Protectors", "6D Curved Dustproof Screen Protector (HD)", "Tempered glass 0.4 mm · 10 pcs per box"),
    "WTP-138":  ("Screen Protectors", "6D Curved Dustproof Screen Protector (Privacy)", "Tempered glass 0.4 mm · 10 pcs per box"),
}
# WK product shots put the WEKOME / WK DESIGN mark in the top ~8% and a Chinese caption in
# the bottom ~9%; cropping those bands lets the product fill the frame.
WK_PHOTO_FRAME = (0.085, 0.095)
WK_CATEGORY_ORDER = ["Wireless Audio", "Wired Earphones", "Speakers", "Data Cables", "Wall Chargers",
                     "Car Chargers", "Screen Protectors"]

NS = {"xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "rel": "http://schemas.openxmlformats.org/package/2006/relationships"}
CONN = {"A": "USB-A", "C": "USB-C", "L": "Lightning", "M": "Micro-USB"}
VARIANT_ORDER = ["colour", "device", "plug", "connector", "power", "type", "length"]


def wps_cell_images(path: Path) -> dict[str, bytes]:
    """{'ID_…': image bytes} for WPS in-cell pictures (=_xlfn.DISPIMG("ID_…",1))."""
    with zipfile.ZipFile(path) as z:
        if "xl/cellimages.xml" not in z.namelist():
            return {}
        rels = ET.fromstring(z.read("xl/_rels/cellimages.xml.rels"))
        target = {r.get("Id"): r.get("Target") for r in rels.findall("rel:Relationship", NS)}
        out = {}
        for pic in ET.fromstring(z.read("xl/cellimages.xml")).iter(f"{{{NS['xdr']}}}pic"):
            name = pic.find(".//xdr:cNvPr", NS).get("name")
            embed = pic.find(".//a:blip", NS).get(f"{{{NS['r']}}}embed")
            out[name] = z.read("xl/" + target[embed].lstrip("/").removeprefix("xl/"))
        return out


# The invoice's WDC-C34a packaging cell points at the C34i (12W Lightning) box; take the right one
# from the price list instead: (sheet, first row, last row of the model's block, column D = PACKAGING).
WK_BOX_OVERRIDE = {"WDC-C34a": ("CHARGING CABLE", 4, 10, "D")}


def price_list_picture(sheet: str, row_from: int, row_to: int, col: str) -> bytes | None:
    """First floating picture anchored in `col`, rows row_from..row_to (1-based) of a price-list
    sheet, read straight from the zip (the 467 MB workbook is never loaded)."""
    if not WK_PRICE_LIST.exists():
        return None
    col_idx = ord(col.upper()) - ord("A")

    def rels(z, path):
        base = path.rsplit("/", 1)[0]
        rel = f"{base}/_rels/{path.rsplit('/', 1)[1]}.rels"
        root = ET.fromstring(z.read(rel))
        out = {}
        for r in root.findall("rel:Relationship", NS):
            t = r.get("Target")
            out[r.get("Id")] = t.lstrip("/") if t.startswith("/") else str(Path(base, t).as_posix())
        return out

    def norm(path):  # resolve "xl/worksheets/../drawings/x.xml"
        parts = []
        for part in path.split("/"):
            parts.pop() if part == ".." else parts.append(part)
        return "/".join(parts)

    with zipfile.ZipFile(WK_PRICE_LIST) as z:
        wb = ET.fromstring(z.read("xl/workbook.xml"))
        main = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        rid = next(s.get(f"{{{NS['r']}}}id") for s in wb.iter(f"{main}sheet") if s.get("name").strip() == sheet)
        sheet_path = norm(rels(z, "xl/workbook.xml")[rid])
        drawing = next((norm(t) for t in rels(z, sheet_path).values() if "/drawings/" in t), None)
        if not drawing:
            return None
        pics = rels(z, drawing)
        found = []
        for anchor in ET.fromstring(z.read(drawing)):
            frm = anchor.find("xdr:from", NS)
            blip = anchor.find(".//a:blip", NS)
            if frm is None or blip is None:
                continue
            c, r = int(frm.find("xdr:col", NS).text), int(frm.find("xdr:row", NS).text) + 1
            if c == col_idx and row_from <= r <= row_to:
                found.append((r, norm(pics[blip.get(f"{{{NS['r']}}}embed")])))
        return z.read(min(found)[1]) if found else None


def model_code(raw: str) -> str | None:
    first = (raw or "").strip().splitlines()[0].strip() if raw else ""
    m = re.match(r"([A-Z]+-[A-Z]?\d+[a-z]?|[A-Z]+\d+)", first)
    return m.group(1) if m else None


def parse_variant(raw: str, category: str) -> dict[str, str]:
    s = re.sub(r"\s+", " ", (raw or "").replace("（", "(").replace("）", ")")).strip()
    comps: dict[str, str] = {}
    m = re.match(r"(?i)(white|black|blue|beige|grey|gray|yellow|silver|tarnish)\b", s)
    if m:
        comps["colour"] = m.group(1).capitalize()
        s = s[m.end():].strip(" -")
    if category == "Screen Protectors":
        m = re.search(r"[(\-]\s*(\d\.\d)\)?\s*$", s)
        size = m.group(1) if m else ""
        dev = s[:m.start()] if m else s
        dev = re.sub(r"IP\s*(\d+)", r"iPhone \1", dev, flags=re.I)
        dev = re.sub(r"\b(pro|max|fold)\b", lambda x: x.group(1).capitalize(), dev, flags=re.I)
        dev = re.sub(r"\s*/\s*(\d)", r" / iPhone \1", dev)
        dev = re.sub(r"\s*/\s*", " / ", dev).strip()
        comps["device"] = f'{dev} · {size}"' if size else dev
        return comps
    if m := re.search(r"\(([ACLM])-([ACLM])\)", s):
        comps["connector"] = f"{CONN[m.group(1)]} to {CONN[m.group(2)]}"
    if m := re.search(r"\((PD)?(\d+)W\)", s):
        comps["power"] = f"{m.group(2)}W" + (" PD" if m.group(1) else "")
    if m := re.search(r"\((\d+)A\)", s):
        comps["power"] = f"{m.group(1)}A"
    if m := re.search(r"(\d+(?:\.\d+)?)\s*M\b", s):
        comps["length"] = f"{m.group(1)} m"
    if "4-in-1" in s:
        comps["type"] = "4-in-1"
    if re.search(r"3\.5\s*mm", s, re.I):
        comps["plug"] = "3.5 mm jack"
    elif re.search(r"type-?c", s, re.I):
        comps["plug"] = "USB-C plug" + (" (digital)" if "digital" in s.lower() else "")
    elif re.search(r"\bUK\b", s):
        comps["plug"] = "UK plug"
    return comps


def read_invoice(path: Path, shipment: int) -> dict:
    from openpyxl import load_workbook
    wbf = load_workbook(path)                      # formulas: DISPIMG ids + floating pics
    wbv = load_workbook(path, data_only=True)      # cached values
    wsf, wsv = wbf["INVOICE"], wbv["INVOICE"]
    cell_imgs = wps_cell_images(path)
    floating = {(im.anchor._from.row + 1, im.anchor._from.col + 1): im._data() for im in wsf._images}

    inv_no = inv_date = None
    header = None
    for r in range(1, wsv.max_row + 1):
        h = wsv.cell(r, 8).value
        if isinstance(h, str) and h.strip().startswith("No."):
            inv_no = str(wsv.cell(r, 9).value).strip()
        if isinstance(h, str) and h.strip().startswith("Date"):
            inv_date = datetime.strptime(str(wsv.cell(r, 9).value).strip(), "%Y.%m.%d").date()
        if wsv.cell(r, 1).value == "Item" and str(wsv.cell(r, 3).value or "").startswith("MODEL"):
            header = r
            break
    if header is None or not inv_no:
        raise SystemExit(f"invoice layout not recognised: {path.name}")

    lines, foc = [], []
    for r in range(header + 1, wsv.max_row + 1):
        a = wsv.cell(r, 1).value
        if isinstance(a, str) and a.strip().upper() == "TTL":
            break
        qty = wsv.cell(r, 7).value
        if not qty:
            continue
        code = model_code(wsv.cell(r, 3).value)
        price = float(wsv.cell(r, 8).value or 0)
        if code is None or price == 0:
            foc.append({"desc": str(wsv.cell(r, 2).value or "").strip(), "qty": int(qty)})
            continue
        m = re.search(r'DISPIMG\("(ID_[0-9A-F]+)"', str(wsf.cell(r, 4).value or ""))
        pic = cell_imgs.get(m.group(1)) if m else floating.get((r, 4))
        mb = re.search(r'DISPIMG\("(ID_[0-9A-F]+)"', str(wsf.cell(r, 5).value or ""))
        box = cell_imgs.get(mb.group(1)) if mb else floating.get((r, 5))
        lines.append({
            "shipment": shipment, "invoice": inv_no, "date": inv_date, "line": a, "row": r,
            "code": code, "inv_desc": str(wsv.cell(r, 2).value or "").strip(),
            "model_raw": re.sub(r"\s+", " ", str(wsv.cell(r, 3).value)).strip(),
            "variant_raw": re.sub(r"\s+", " ", str(wsv.cell(r, 6).value or "")).strip(),
            "qty": int(qty), "price": price, "amount": round(int(qty) * price, 2), "pic": pic, "box": box,
        })

    pl = wbv["PACKING LIST"]
    pack = {}
    for r in range(1, pl.max_row + 1):
        if str(pl.cell(r, 1).value or "").strip().upper() == "TTL":
            pack = {"cartons": int(float(pl.cell(r, 9).value)), "cbm": float(pl.cell(r, 13).value),
                    "gross_kg": float(pl.cell(r, 15).value), "pcs": int(float(pl.cell(r, 7).value))}
            break
    foc_detail = []
    for r in range(1, pl.max_row + 1):
        d, f, q = pl.cell(r, 4).value, pl.cell(r, 6).value, pl.cell(r, 7).value
        if (str(d or "").strip().lower() == "stands" or "brand bag" in str(f or "").lower()) and q:
            foc_detail.append((str(f).strip(), int(float(q))))
    return {"invoice": inv_no, "date": inv_date, "lines": lines, "foc": foc, "foc_detail": foc_detail,
            "pack": pack, "file": path.name}


def build_wekome_products(ships: list[dict]) -> list[dict]:
    products: "OrderedDict[str, dict]" = OrderedDict()
    for ship in ships:
        for ln in ship["lines"]:
            code = ln["code"]
            if code not in WK_CATALOG:
                raise SystemExit(f"unknown WEKOME model {code!r} (row {ln['row']} of {ship['file']}) — add it to WK_CATALOG")
            cat, name, spec = WK_CATALOG[code]
            p = products.setdefault(code, {"code": code, "category": cat, "name": name, "spec": spec,
                                           "pic": None, "box": None, "variants": OrderedDict()})
            p["pic"] = p["pic"] or ln["pic"]
            p["box"] = p["box"] or ln["box"]
            comps = parse_variant(ln["variant_raw"], cat)
            key = tuple(sorted(comps.items()))
            v = p["variants"].setdefault(key, {"comps": comps, "raw": ln["variant_raw"], "qty1": 0, "qty2": 0,
                                               "value1": 0.0, "value2": 0.0, "prices": set()})
            v[f"qty{ln['shipment']}"] += ln["qty"]
            v[f"value{ln['shipment']}"] += ln["amount"]
            v["prices"].add(ln["price"])
    out = []
    for p in products.values():
        if not p["pic"]:
            raise SystemExit(f"no product picture found for {p['code']}")
        vs = list(p["variants"].values())
        keys = [k for k in VARIANT_ORDER if any(k in v["comps"] for v in vs)]
        varying = [k for k in keys if len({v["comps"].get(k) for v in vs}) > 1]
        common = [vs[0]["comps"][k] for k in ("plug", "length") if k in keys and k not in varying
                  and all(k in v["comps"] for v in vs)]
        show = [k for k in keys if k in varying or (k == "colour" and p["category"] != "Screen Protectors")]
        for v in vs:
            v["label"] = " · ".join(v["comps"][k] for k in show if k in v["comps"])
        p["common"] = " · ".join(common)
        p["variants"] = vs
        p["qty1"] = sum(v["qty1"] for v in vs)
        p["qty2"] = sum(v["qty2"] for v in vs)
        p["qty"] = p["qty1"] + p["qty2"]
        p["value"] = round(sum(v["value1"] + v["value2"] for v in vs), 2)
        prices = sorted({pr for v in vs for pr in v["prices"]})
        p["price_txt"] = f"{prices[0]:.2f}" if len(prices) == 1 else f"{prices[0]:.2f}–{prices[-1]:.2f}"
        out.append(p)
    out.sort(key=lambda p: (WK_CATEGORY_ORDER.index(p["category"]), -p["value"]))
    for i, p in enumerate(out, 1):
        p["no"] = i
    return out
