"""Monthly supplier workbook (Market Intel v2, Phase 2; plan Part F9-F11, build contract §3).

A reviewer picks a supplier and a month; build_draft() gathers that supplier's field reports of the
month and keeps a DRAFT supplier_reports row: counts by issue type, new vs recurring, open corrective
actions, and a narrative of template sentences in which every sentence about a finding cites the
sightings it rests on ([MI-<observation id>]) and every action its id ([CA-<action id>]). A draft's
narrative and rows can be edited; an approver (management included) approves it, which supersedes an
earlier approved report of the same supplier, month and edition; only then can it be downloaded as an
.xlsx (openpyxl), and every download is logged with the file's sha256. Nothing is ever e-mailed.

Which sightings belong to a supplier (each to one supplier at most): the first link found among
    the case's supplier (the office's reviewed link) -> the case's brand's supplier
    -> the sighting's own supplier_id -> its brand's supplier,
each followed through merged_into to the live supplier. Findings are the sightings whose issue type
(the case's, else the sighting's own) is of kind quality (sheet "Product issues") or packaging
("Packaging"); a v1 complaint without an issue type counts as quality. The internal edition adds the
competitor and demand sightings. A case found not substantiated (or closed without ever being
verified, i.e. below evidence level 3) is left out everywhere.

Rules:
    recurring      the same yq_item_code + issue type in >= 2 cases, or from >= 3 shops, within the 90
                   days that end with the month (a product nobody identified: its own case, >= 3 shops)
    insufficient   no photo, evidence level 1 and a single report (that product + issue in those 90
                   days): named in its own narrative paragraph and listed last, never counted, unless a
                   reviewer includes it by hand
    month          a Bahrain calendar month (UTC+3)

Privacy (a supplier edition leaves the company): no shop or customer name, no rep name ("field team"),
no YQ price, cost or margin, no photo flagged with people, no receipt photo, no other supplier's rows.
Every free-text value goes through app.ai_insights.mask_text (phones, e-mails) and then a scrubber
that replaces the shop, customer, rep and other-supplier names the database knows and any money
amount. Photos come from the private bucket through the app's storage client; one that cannot be
fetched is written as the words "photo not available".

The tables arrive with scripts/r8m_market_quality_migration.sql. Until then available() is False and
every call the routes make (build_draft, list_reports, get_report, update_draft, approve, download)
raises SupplierReportError(MISSING_HINT, 503), so the API may deploy first.
"""
from __future__ import annotations

import functools
import hashlib
import io
import logging
import re
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta, timezone

from app import market_intel as mi
from app.ai_insights import mask_text
from app.database import get_client

log = logging.getLogger(__name__)

# ── vocabulary ────────────────────────────────────────────────────────────────

EDITIONS = ("supplier", "internal")
EDITION_LABELS = {"supplier": "Supplier copy", "internal": "Internal edition: for YQ only, not for the supplier"}
STATUSES = ("draft", "approved", "superseded")
FINDING_KINDS = ("quality", "packaging")
MARKET_KINDS = ("competitor", "demand")
KIND_WORDS = {"quality": "Product issue", "packaging": "Packaging"}
# a v1 sighting without an issue type: its kind decides (new_product / other are not findings)
V1_KINDS = {"complaint": "quality", "competitor_price": "competitor", "promotion": "competitor",
            "shop_asked": "demand"}
EVIDENCE_WORDS = {1: "Reported", 2: "Evidence submitted", 3: "Inspected & verified by YQ",
                  4: "Supplier confirmed", 5: "Resolution verified"}
VERIFIED_LEVEL = 3
CASE_STATUS_WORDS = {"submitted": "Received", "under_review": "Under review",
                     "more_info_required": "More information requested", "verified": "Verified by YQ",
                     "not_substantiated": "Not substantiated", "action_requested": "Action requested",
                     "resolved": "Resolved", "closed": "Closed"}
OPEN_ACTION_STATUSES = ("open", "responded")
ACTION_STATUS_WORDS = {"open": "Awaiting your response", "responded": "Response received, YQ check pending",
                       "resolved": "Resolved", "cancelled": "Cancelled"}
BASIS_WORDS = {"customer_said": "Customer said", "rep_saw": "Field team saw"}
# a supplier edition shows product, packaging and label photos only: never a receipt (shop names,
# prices) or an 'other' photo; a photo flagged with people is never in either edition
SUPPLIER_PHOTO_ROLES = ("product", "packaging", "label_batch")
ROLE_WORDS = {"product": "Product", "packaging": "Packaging", "label_batch": "Label / batch", "receipt": "Receipt",
              "other": "Other"}
SIGNAL_WORDS = {"search_zero": "Marketplace search found nothing", "restock_ask": "Restock ask on the marketplace"}

RECUR_DAYS, RECUR_CASES, RECUR_SHOPS = 90, 2, 3
BAHRAIN = timezone(timedelta(hours=3))
EVIDENCE_PX = 640                 # the Evidence sheet's JPEG, longest side (shown at half size)
THUMB_PX = 160                    # a finding row's thumbnail
MAX_PHOTOS = 120                  # embedded per workbook (about 60 KB each); the rest are listed, not embedded
MAX_CITES = 12                    # ids cited in one sentence; the sheet lists them all
MAX_NARRATIVE = 20000
FIELD_TEAM = "field team"
SHOP_MASK, SUPPLIER_MASK, AMOUNT_MASK = "[shop]", "[another supplier]", "[amount]"
PHOTO_NOT_AVAILABLE = "photo not available"
PAGE_MAX = 50                     # pages of 1,000 rows per read

# Who may do what (the code's role helpers): reviewers build and edit; approvers (reviewers +
# management) read, approve and download. A salesman never reaches a supplier report.
REVIEW_ROLES = mi.REVIEW_ROLES
APPROVE_ROLES = mi.APPROVE_ROLES
# Management is read-only centrally (app.auth); approving needs this template on
# auth.READ_ONLY_WRITE_ALLOWLIST, like the Market Intel item approval.
APPROVE_ROUTE = "/market-intel/supplier-reports/{report_id}/approve"

MIGRATION = "scripts/r8m_market_quality_migration.sql"
MISSING_HINT = "Supplier reports are not set up yet: apply scripts/r8m_market_quality_migration.sql."

MONTHS = ("January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
          "November", "December")


class SupplierReportError(ValueError):
    """A refusal the route turns into an HTTP answer (400 unless `status` says otherwise)."""

    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.status = status


# ── availability (the migration may land after the API) ──────────────────────

_probe: dict = {"until": 0.0, "ok": False}
_PROBE_HIT_S, _PROBE_MISS_S = 600.0, 60.0


def available(force: bool = False) -> bool:
    """True when the r8m tables exist. A hit is remembered for 10 minutes, a miss for one."""
    now = time.monotonic()
    if not force and _probe["until"] > now:
        return bool(_probe["ok"])
    try:
        get_client().table("supplier_reports").select("id").limit(1).execute()
        ok = True
    except Exception as e:  # noqa: BLE001 — a missing table (or a blip) means "not yet"
        log.debug("supplier reports probe failed: %s", e)
        ok = False
    _probe.update(until=now + (_PROBE_HIT_S if ok else _PROBE_MISS_S), ok=ok)
    return ok


def forget_probe() -> None:
    _probe.update(until=0.0, ok=False)


def _not_set_up(e: Exception) -> bool:
    """A missing table (42P01 / PGRST205) or column (42703 / PGRST204): the migration is not (or no
    longer) there. The cached probe is forgotten."""
    msg = f"{getattr(e, 'code', '')} {e}"
    if mi.is_missing_table(e) or "42703" in msg or "PGRST204" in msg \
            or ("column" in msg and "does not exist" in msg):
        forget_probe()
        return True
    return False


def _tables(fn):
    """Every public call: 503 MISSING_HINT before the migration, or when a read meets a missing
    table or column half-way (the reverse ran)."""
    @functools.wraps(fn)
    def run(*a, **k):
        if not available():
            raise SupplierReportError(MISSING_HINT, 503)
        try:
            return fn(*a, **k)
        except SupplierReportError:
            raise
        except Exception as e:  # noqa: BLE001
            if _not_set_up(e):
                raise SupplierReportError(MISSING_HINT, 503) from None
            raise
    return run


def _need(role: str | None, allowed: frozenset, message: str) -> None:
    if (role or "") not in allowed:
        raise SupplierReportError(message, 403)


# ── small helpers ─────────────────────────────────────────────────────────────

def _now() -> datetime:
    return datetime.now(timezone.utc)


def _iso(dt: datetime | None = None) -> str:
    return (dt or _now()).astimezone(timezone.utc).isoformat()


def _int(v) -> int | None:
    try:
        return int(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


def _ts(v) -> datetime | None:
    return mi._parse_ts(v)


def _bh_date(v) -> date | None:
    t = _ts(v) if not isinstance(v, datetime) else v
    return t.astimezone(BAHRAIN).date() if t else None


def _date_of(v) -> date | None:
    """A date column ('2026-10-15') or a timestamp, as a Bahrain date."""
    if isinstance(v, date) and not isinstance(v, datetime):
        return v
    s = str(v or "").strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
        try:
            return date.fromisoformat(s)
        except ValueError:
            return None
    return _bh_date(s)


def _day(d: date | None) -> str:
    """'18 Sep 2026'."""
    return f"{d.day} {MONTHS[d.month - 1][:3]} {d.year}" if d else ""


def _plural(n: int, word: str) -> str:
    return f"{n} {word}{'' if n == 1 else 's'}"


def parse_month(v, *, now: datetime | None = None, future_ok: bool = False) -> date:
    """'YYYY-MM' (or the stored 'YYYY-MM-01') -> the month's first day. A month that has not started
    (Bahrain time) is refused unless future_ok (a list filter)."""
    m = re.fullmatch(r"\s*(\d{4})-(\d{1,2})(?:-01)?\s*", str(v or ""))
    if not m:
        raise SupplierReportError("The month is YYYY-MM, like 2026-09.")
    y, mo = int(m.group(1)), int(m.group(2))
    if not 1 <= mo <= 12 or y < 2024:
        raise SupplierReportError("The month is YYYY-MM, like 2026-09.")
    first = date(y, mo, 1)
    today = (now or _now()).astimezone(BAHRAIN).date()
    if not future_ok and first > today.replace(day=1):
        raise SupplierReportError("That month has not started yet.")
    return first


def month_key(first: date) -> str:
    return f"{first.year:04d}-{first.month:02d}"


def month_label(first: date) -> str:
    return f"{MONTHS[first.month - 1]} {first.year}"


def month_bounds(first: date) -> tuple[datetime, datetime]:
    """[the first day 00:00, the next month's first day 00:00), Bahrain time."""
    nxt = date(first.year + (1 if first.month == 12 else 0), first.month % 12 + 1, 1)
    return (datetime(first.year, first.month, 1, tzinfo=BAHRAIN),
            datetime(nxt.year, nxt.month, 1, tzinfo=BAHRAIN))


def _who(email) -> str | None:
    """A login as the API shows it: the local part only."""
    s = str(email or "").strip()
    return s.split("@")[0] or None if s else None


# ── the scrubber: what a supplier must never read ────────────────────────────

_AMOUNT = re.compile(
    r"(?i)(?:\b(?:bhd|bd|fils|dinars?)\.?\s*\d+(?:[.,]\d+)?"
    r"|\d+(?:[.,]\d+)?\s*(?:bhd|bd|fils|dinars?)\b"
    r"|\d+(?:[.,]\d+)?\s*(?:دينار|د\.ب)|(?:دينار|د\.ب)\s*\d+(?:[.,]\d+)?)")
_CURRENCY = re.compile(r"(?i)\b(?:bhd|bd|dinars?)\b|دينار|د\.ب")
_PRICE_WORD = re.compile(r"(?i)\b(price|cost|margin|rate|paid)\b(\s*(?:of|is|was|at|for|:|=|-)?\s*)\d+(?:[.,]\d+)?")
_WORD = re.compile(r"\w+")
# names that are system words, never a person or a shop (the ops selftest's list, plus a few)
_GENERIC = frozenset({"cash customer", "customer", "cash", "credit", "total", "sub total", "grand total",
                      "walk in", "walk-in", "walk in customer", "shop", "the shop", "office", "warehouse",
                      "field team", "supplier", "salesman", "yq bahrain"})


class Scrubber:
    """mask_text (phones -> [number], e-mails -> [email]), money amounts -> [amount], then every known
    name: a rep's -> 'field team', a shop's or a customer's -> [shop], another supplier's ->
    [another supplier]. Names are found by their first word, so thousands of names cost one dictionary
    lookup per word of text; the longest name is replaced first ("Shop Alpha Two" before "Shop Alpha").
    The `keep` names (the report's own supplier) are never touched, even when another name is part of
    them."""

    def __init__(self, *, reps=(), shops=(), others=(), keep=()):
        keep_names = sorted({" ".join(str(k).split()) for k in keep if k and len(str(k).strip()) >= 3},
                            key=len, reverse=True)
        keep_l = {k.lower() for k in keep_names}
        # a kept name is parked behind a private-use token (no word character, no digit) while the
        # other names are replaced, then put back
        self.keep = [(re.compile(r"(?<!\w)" + r"\s+".join(re.escape(w) for w in k.split(" ")) + r"(?!\w)",
                                 re.IGNORECASE), f"{chr(0xE100 + i)}") for i, k in enumerate(keep_names)]
        self.index: dict[str, list[tuple[int, re.Pattern, str]]] = {}
        seen: set[str] = set()
        for names, repl, min_len in ((reps, FIELD_TEAM, 3), (shops, SHOP_MASK, 4), (others, SUPPLIER_MASK, 4)):
            for raw in names:
                n = " ".join(str(raw or "").split())
                low = n.lower()
                words = _WORD.findall(low)
                if len(n) < min_len or not words or low in _GENERIC or low in keep_l or low in seen \
                        or n.replace(" ", "").isdigit():
                    continue
                seen.add(low)
                rx = re.compile(r"(?<!\w)" + r"\s+".join(re.escape(w) for w in n.split(" ")) + r"(?!\w)",
                                re.IGNORECASE)
                self.index.setdefault(words[0], []).append((len(n), rx, repl))

    def __call__(self, value):
        if not isinstance(value, str) or not value:
            return value
        t = mask_text(value)
        t = _AMOUNT.sub(AMOUNT_MASK, t)
        t = _PRICE_WORD.sub(lambda m: f"{m.group(1)}{m.group(2)}{AMOUNT_MASK}", t)
        t = _CURRENCY.sub(AMOUNT_MASK, t)
        parked: list[tuple[str, str]] = []
        for rx, token in self.keep:
            t = rx.sub(lambda m, token=token: parked.append((token, m.group(0))) or token, t)
        hits = [h for w in {x.lower() for x in _WORD.findall(t)} for h in self.index.get(w, ())]
        for _n, rx, repl in sorted(hits, key=lambda h: -h[0]):
            t = rx.sub(repl, t)
        for token, original in parked:
            t = t.replace(token, original, 1)
        return t


# ── reads (paged: PostgREST answers 1,000 rows at most) ──────────────────────

def _q(table: str, cols: str = "*"):
    return get_client().table(table).select(cols)


def _all(make, max_pages: int = PAGE_MAX) -> list[dict]:
    """Every row of make() (a fresh select builder), ordered by id, 1,000 at a time."""
    from app.shop_analytics import PAGE, paged
    rows, capped = paged(lambda a: (make().order("id").range(a, a + PAGE - 1).execute().data or []), max_pages)
    if capped:
        log.warning("supplier reports: a read stopped at %d rows", len(rows))
    return rows


def _by(table: str, col: str, ids, cols: str = "*") -> list[dict]:
    keys = sorted({i for i in (_int(x) for x in ids) if i is not None})
    out: list[dict] = []
    for i in range(0, len(keys), 200):
        chunk = keys[i:i + 200]
        out += _all(lambda chunk=chunk: _q(table, cols).in_(col, chunk))
    return out


def _names(make) -> list[dict]:
    """A name list for the scrubber. A table that does not exist is an empty list; any other failure
    stops the report (a list read half-way could let a name through)."""
    try:
        return _all(make)
    except Exception as e:  # noqa: BLE001
        if mi.is_missing_table(e):
            return []
        raise


def _live(sid, by_id: dict[int, dict], hops: int = 6) -> int | None:
    """Follow merged_into to the live row's id."""
    cur = _int(sid)
    seen: set[int] = set()
    while cur is not None and hops > 0 and cur not in seen:
        row = by_id.get(cur)
        nxt = _int((row or {}).get("merged_into"))
        if nxt is None:
            break
        seen.add(cur)
        cur, hops = nxt, hops - 1
    return cur


def _brand_supplier(bid, brands: dict[int, dict]) -> int | None:
    b = brands.get(_int(bid)) if _int(bid) is not None else None
    if not b:
        return None
    live = brands.get(_live(b["id"], brands)) or b
    return _int(live.get("supplier_id")) or _int(b.get("supplier_id"))


def _owner(o: dict, case: dict | None, brands: dict, suppliers: dict) -> int | None:
    """The live supplier a sighting belongs to (see the module doc), or None."""
    c = case or {}
    for sid in (_int(c.get("supplier_id")), _brand_supplier(c.get("brand_id"), brands),
                _int(o.get("supplier_id")), _brand_supplier(o.get("brand_id"), brands)):
        if sid is not None:
            return _live(sid, suppliers)
    return None


def _cases_for(ids) -> dict[int, dict]:
    """The cases by id, plus every case they were merged into."""
    out: dict[int, dict] = {}
    want = {i for i in (_int(x) for x in ids) if i is not None}
    for _ in range(6):
        need = sorted(want - set(out))
        if not need:
            break
        for r in _by("market_cases", "id", need):
            out[int(r["id"])] = r
        want = {_int(r.get("merged_into")) for r in out.values()} - {None}
    return out


def _live_case(cid, cases: dict[int, dict]) -> dict | None:
    c = cases.get(_int(cid)) if _int(cid) is not None else None
    hops = 6
    while c and _int(c.get("merged_into")) is not None and hops > 0:
        nxt = cases.get(_int(c["merged_into"]))
        if not nxt:
            break
        c, hops = nxt, hops - 1
    return c


def _level(case: dict | None) -> int:
    try:
        return max(1, min(5, int((case or {}).get("evidence_level") or 1)))
    except (TypeError, ValueError):
        return 1


def _dropped(case: dict | None) -> bool:
    """Not substantiated, or closed without ever being verified (verified sets level >= 3 and the level
    only moves up, so a closed case below 3 came from not_substantiated)."""
    if not case:
        return False
    st = case.get("status")
    return st == "not_substantiated" or (st == "closed" and _level(case) < VERIFIED_LEVEL)


def _kind(o: dict, case: dict | None, types: dict[int, dict]) -> str:
    for tid in (_int((case or {}).get("issue_type_id")), _int(o.get("issue_type_id"))):
        t = types.get(tid) if tid is not None else None
        if t and t.get("kind"):
            return str(t["kind"])
    return V1_KINDS.get(str(o.get("kind") or ""), "other")


def _shop_key(o: dict) -> str | None:
    if o.get("shop_customer_id"):
        return f"c{o['shop_customer_id']}"
    n = " ".join(str(o.get("shop_name") or "").lower().split())
    return f"n{n}" if n else None


def _eligible(p: dict, edition: str) -> bool:
    if p.get("has_people_flag"):
        return False
    return edition == "internal" or (p.get("role") or "product") in SUPPLIER_PHOTO_ROLES


# ── gather: everything one report is built from (read-only) ──────────────────

def gather(supplier_id, first: date, edition: str, *, now: datetime | None = None,
           as_of: datetime | None = None) -> dict:
    """The supplier, its sightings of the month (with the 90 days before for the recurring rule),
    their cases, photos and corrective actions, and the scrubber's name lists."""
    now = now or _now()
    if edition not in EDITIONS:
        raise SupplierReportError("The edition is 'supplier' or 'internal'.")
    suppliers = {int(r["id"]): r for r in _all(lambda: _q("suppliers"))}
    sid = _int(supplier_id)
    sup = suppliers.get(sid) if sid is not None else None
    if not sup:
        raise SupplierReportError("That supplier does not exist.", 404)
    live = _live(sid, suppliers)
    if live != sid:
        into = (suppliers.get(live) or {}).get("name") or "another supplier"
        raise SupplierReportError(f"{sup.get('name') or 'This supplier'} was merged into {into}: "
                                  "build that supplier's report.", 409)
    brands = {int(r["id"]): r for r in _all(lambda: _q("brands"))}
    types = {int(r["id"]): r for r in _all(lambda: _q("market_issue_types"))}
    start, end = month_bounds(first)
    since = end - timedelta(days=RECUR_DAYS)
    raw = _all(lambda: _q("market_observations").gte("observed_at", _iso(since)).lt("observed_at", _iso(end)))
    window: list[tuple[dict, datetime]] = []
    for o in raw:
        t = _ts(o.get("observed_at")) or _ts(o.get("created_at"))
        if t and since <= t < end:
            window.append((o, t))
    cases = _cases_for(o.get("case_id") for o, _t in window)

    entries: list[dict] = []
    for o, t in window:
        case = _live_case(o.get("case_id"), cases)
        sku = str((case or {}).get("yq_item_code") or o.get("yq_item_code") or "").strip() or None
        itype = _int((case or {}).get("issue_type_id")) or _int(o.get("issue_type_id"))
        key = ("sku", sku, itype) if sku else (("case", int(case["id"])) if case else ("obs", int(o["id"])))
        entries.append({"o": o, "ts": t, "case": case, "kind": _kind(o, case, types), "sku": sku, "itype": itype,
                        "key": key, "owner": _owner(o, case, brands, suppliers), "dropped": _dropped(case),
                        "shop": _shop_key(o)})
    # the recurring and single-report rules count every supplier's sightings of the same product + issue
    stats: dict[tuple, dict] = {}
    for e in entries:
        if e["kind"] in FINDING_KINDS and not e["dropped"]:
            s = stats.setdefault(e["key"], {"cases": set(), "shops": set(), "reports": 0})
            s["cases"].add(f"c{e['case']['id']}" if e["case"] else f"o{e['o']['id']}")
            if e["shop"]:
                s["shops"].add(e["shop"])
            s["reports"] += 1

    kinds = FINDING_KINDS + (MARKET_KINDS if edition == "internal" else ())
    mine = sorted((e for e in entries if e["ts"] >= start and e["owner"] == sid and not e["dropped"]
                   and e["kind"] in kinds), key=lambda e: (e["ts"], int(e["o"]["id"])))
    photos: dict[int, list[dict]] = {}
    for p in _by("market_photos", "observation_id", [e["o"]["id"] for e in mine]):
        photos.setdefault(int(p["observation_id"]), []).append(p)
    for v in photos.values():
        v.sort(key=lambda p: (_int(p.get("position")) or 0, _int(p.get("id")) or 0))

    # corrective actions: the supplier's (and its merged ones'), plus an action without a supplier
    # that is linked to one of this supplier's cases; never another supplier's
    family = sorted(i for i in suppliers if _live(i, suppliers) == sid)
    acts = {int(r["id"]): r for r in _all(lambda: _q("supplier_actions").in_("supplier_id", family))}
    case_ids = sorted({int(e["case"]["id"]) for e in mine if e["case"]})
    extra = {_int(r.get("action_id")) for r in _by("supplier_action_cases", "case_id", case_ids)} - set(acts) - {None}
    for r in _by("supplier_actions", "id", extra):
        if _int(r.get("supplier_id")) is None:
            acts[int(r["id"])] = r
    action_cases: dict[int, set[int]] = {}
    for r in _by("supplier_action_cases", "action_id", list(acts)):
        if _int(r.get("action_id")) is not None and _int(r.get("case_id")) is not None:
            action_cases.setdefault(int(r["action_id"]), set()).add(int(r["case_id"]))

    def listed(a: dict) -> bool:
        made = _ts(a.get("created_at"))
        if made and made >= end:
            return False
        if a.get("status") in OPEN_ACTION_STATUSES:
            return True
        changed = _ts(a.get("updated_at")) or made
        return bool(changed and start <= changed < end)
    actions = {i: a for i, a in acts.items() if listed(a)}

    # the scrubber: every rep, shop and customer name the database knows, the shop names typed on the
    # sightings, and every other supplier's name and code
    own = (sup.get("name"), sup.get("code"))
    reps = [r.get(k) for r in _names(lambda: _q("salesmen")) for k in ("name", "focus_name")]
    shops = [r.get(k) for r in _names(lambda: _q("shop_customers", "id,name,shop")) for k in ("name", "shop")]
    shops += [r.get("name") for r in _names(lambda: _q("customers", "id,name"))]
    shops += [o.get("shop_name") for o, _t in window]
    others = [r.get(k) for i, r in suppliers.items() if _live(i, suppliers) != sid for k in ("name", "code")]
    scrub = Scrubber(reps=reps, shops=shops, others=others, keep=own)
    return {"supplier": sup, "supplier_id": sid, "first": first, "start": start, "end": end, "since": since,
            "edition": edition, "now": now, "as_of": as_of or now, "types": types, "entries": mine,
            "stats": stats, "photos": photos, "items": (mi._catalog_ctx() or {}).get("items") or {},
            "actions": actions, "action_cases": action_cases, "scrub": scrub}


# ── assemble: rows, groups, counts (pure: no reads) ──────────────────────────

def _finding_row(ctx: dict, e: dict) -> dict:
    o, case, scrub = e["o"], e["case"] or {}, ctx["scrub"]
    oid = int(o["id"])
    t = ctx["types"].get(e["itype"]) or {}
    every = ctx["photos"].get(oid, [])
    shown = [p for p in every if _eligible(p, ctx["edition"])]
    level = _level(case) if e["case"] else (2 if every else 1)
    stat = ctx["stats"].get(e["key"]) or {"cases": set(), "shops": set(), "reports": 1}
    sku = e["sku"]
    product = (ctx["items"].get(sku) or {}).get("display_name") if sku else None
    product = scrub(product or case.get("title") or o.get("title") or (sku or "")) or None
    label = t.get("label_en") or ("Unspecified problem" if e["kind"] == "quality" else "Unspecified packaging problem")
    return {"id": oid, "ref": f"MI-{oid}", "date": e["ts"].astimezone(BAHRAIN).date().isoformat(),
            "sku": sku, "product": product,
            "issue": {"id": e["itype"], "code": t.get("code"), "label": scrub(label)},
            "kind": e["kind"], "description": scrub(mi.clean_text(o.get("note"), 1000)) or "",
            "qty": _int(o.get("qty_affected")),
            "batch_ref": scrub(mi.clean_text(o.get("batch_ref") or case.get("batch_ref"), 80)),
            "packaging_version": scrub(mi.clean_text(o.get("packaging_version") or case.get("packaging_version"), 80)),
            "basis": BASIS_WORDS.get(o.get("report_basis")),
            "evidence_level": level, "evidence_label": EVIDENCE_WORDS[level],
            "finding": "Verified" if level >= VERIFIED_LEVEL else "Unverified report",
            "status": CASE_STATUS_WORDS.get(case.get("status"), "Received") if e["case"] else "Received",
            "case_id": _int(case.get("id")) if e["case"] else None,
            "recurring": len(stat["cases"]) >= RECUR_CASES or len(stat["shops"]) >= RECUR_SHOPS,
            "safety": bool(o.get("safety_flag") or case.get("safety")),
            "weak": not every and level == 1 and int(stat["reports"]) <= 1,
            "photos_withheld": len(every) - len(shown),
            "_key": e["key"], "_shop": e["shop"], "_photos": shown,
            "_window": (len(stat["cases"]), len(stat["shops"]))}


def _market_row(ctx: dict, e: dict) -> dict:
    o, scrub = e["o"], ctx["scrub"]
    oid = int(o["id"])
    every = ctx["photos"].get(oid, [])
    shown = [p for p in every if _eligible(p, ctx["edition"])]
    row = {"id": oid, "ref": f"MI-{oid}", "date": e["ts"].astimezone(BAHRAIN).date().isoformat(), "kind": e["kind"],
           "sku": e["sku"], "note": scrub(mi.clean_text(o.get("note"), 1000)) or "",
           "area": scrub(mi.clean_text(o.get("area"), 60)), "weak": False,
           "photos_withheld": len(every) - len(shown), "_key": ("obs", oid), "_shop": e["shop"], "_photos": shown}
    if e["kind"] == "competitor":
        row.update(product_seen=scrub(mi.clean_text(o.get("title"), 160)),
                   brand=scrub(mi.clean_text(o.get("brand_text") or o.get("brand"), 80)),
                   competitor=scrub(mi.clean_text(o.get("competitor"), 80)), category=o.get("category"),
                   price_seen=mi.money_out(o.get("price_bhd")))
    else:
        src = o.get("source")
        row.update(source=SIGNAL_WORDS.get(o.get("signal"), "Marketplace") if src == "system" else "Field team",
                   asked_for=scrub(mi.clean_text(o.get("title"), 160)),
                   qty=_int(o.get("demand_qty")) or _int(o.get("qty_affected")), demand_level=o.get("demand_level"))
    return row


def _days_open(a: dict, as_of: datetime) -> int | None:
    made = _bh_date(a.get("created_at"))
    if not made:
        return None
    upto = as_of.astimezone(BAHRAIN).date() if a.get("status") in OPEN_ACTION_STATUSES \
        else (_bh_date(a.get("updated_at")) or as_of.astimezone(BAHRAIN).date())
    return max(0, (upto - made).days)


def assemble(ctx: dict, included=None) -> dict:
    """Rows (state included / insufficient / excluded), the finding groups, the actions and the counts.
    included=None is a fresh draft: every candidate but the insufficient ones."""
    rows = [_finding_row(ctx, e) if e["kind"] in FINDING_KINDS else _market_row(ctx, e) for e in ctx["entries"]]
    ids = {r["id"] for r in rows}
    wanted = None if included is None else {i for i in (_int(x) for x in included) if i is not None}
    inc = {r["id"] for r in rows if not r["weak"]} if wanted is None else wanted & ids
    for r in rows:
        r["state"] = "included" if r["id"] in inc else ("insufficient" if r["weak"] else "excluded")
    counted = [r for r in rows if r["state"] == "included" and r["kind"] in FINDING_KINDS]
    weak = [r for r in rows if r["state"] == "insufficient"]
    market = [r for r in rows if r["kind"] in MARKET_KINDS and r["state"] == "included"]

    as_of = ctx["as_of"]
    actions = []
    for aid, a in sorted(ctx["actions"].items(), key=lambda kv: (kv[1].get("status") not in OPEN_ACTION_STATUSES, kv[0])):
        linked_cases = ctx["action_cases"].get(aid, set())
        actions.append({"id": aid, "ref": f"CA-{aid}", "request": ctx["scrub"](mi.clean_text(a.get("request"), 1000)) or "",
                        "owner": ctx["scrub"](mi.clean_text(a.get("owner"), 120)),
                        "due_date": _date_of(a.get("due_date")).isoformat() if _date_of(a.get("due_date")) else None,
                        "status": a.get("status") or "open",
                        "status_words": ACTION_STATUS_WORDS.get(a.get("status") or "open", str(a.get("status") or "")),
                        "supplier_response": ctx["scrub"](mi.clean_text(a.get("supplier_response"), 2000)),
                        "resolution_check": ctx["scrub"](mi.clean_text(a.get("resolution_check"), 1000)),
                        "created_at": a.get("created_at"), "days_open": _days_open(a, as_of),
                        "linked_observation_ids": sorted(r["id"] for r in counted if r["case_id"] in linked_cases),
                        "_cases": linked_cases})
    open_cases = {c for a in actions if a["status"] in OPEN_ACTION_STATUSES for c in a["_cases"]}

    groups: dict[tuple, dict] = {}
    for r in counted:
        groups.setdefault(r["_key"], {"key": r["_key"], "rows": []})["rows"].append(r)
    for g in groups.values():
        rs, top = g["rows"], g["rows"][0]
        g.update(kind=top["kind"], issue=top["issue"], sku=top["sku"], product=top["product"],
                 recurring=top["recurring"], window=top["_window"], reports=len(rs),
                 shops=len({r["_shop"] for r in rs if r["_shop"]}),
                 verified=[r["id"] for r in rs if r["evidence_level"] >= VERIFIED_LEVEL],
                 safety=any(r["safety"] for r in rs), photos=sum(len(r["_photos"]) for r in rs),
                 versions=sorted({r["packaging_version"] for r in rs if r["packaging_version"]}),
                 batches=sorted({r["batch_ref"] for r in rs if r["batch_ref"]}),
                 open_action=any(r["case_id"] in open_cases for r in rs if r["case_id"] is not None))
    order = sorted(groups.values(), key=lambda g: (not g["recurring"], not g["safety"], -g["reports"],
                                                   str(g["sku"] or "~"), str(g["issue"]["label"])))

    by_issue: dict[tuple, dict] = {}
    for r in counted:
        b = by_issue.setdefault((r["issue"]["id"], r["kind"]), {
            "issue_type_id": r["issue"]["id"], "code": r["issue"]["code"], "label": r["issue"]["label"],
            "kind": r["kind"], "reports": 0, "shops": set(), "verified": 0, "recurring": False})
        b["reports"] += 1
        if r["_shop"]:
            b["shops"].add(r["_shop"])
        b["verified"] += 1 if r["evidence_level"] >= VERIFIED_LEVEL else 0
        b["recurring"] = b["recurring"] or r["recurring"]
    issues = [{**b, "shops": len(b["shops"])} for b in sorted(by_issue.values(), key=lambda b: (-b["reports"], str(b["label"])))]

    new = [g for g in order if not g["recurring"]]
    rec = [g for g in order if g["recurring"]]
    verified = sum(1 for r in counted if r["evidence_level"] >= VERIFIED_LEVEL)
    counts = {
        "version": 1, "month": month_key(ctx["first"]), "edition": ctx["edition"],
        "period": {"from": ctx["start"].date().isoformat(), "to": (ctx["end"] - timedelta(days=1)).date().isoformat(),
                   "timezone": "Bahrain (UTC+3)"},
        "supplier": {"id": ctx["supplier_id"], "name": ctx["supplier"].get("name")},
        "reports": len(counted), "product_issues": sum(1 for r in counted if r["kind"] == "quality"),
        "packaging": sum(1 for r in counted if r["kind"] == "packaging"),
        "shops": len({r["_shop"] for r in counted if r["_shop"]}),
        "verified": verified, "unverified": len(counted) - verified,
        "safety": sum(1 for r in counted if r["safety"]),
        "photos": sum(len(r["_photos"]) for r in counted),
        "new": {"problems": len(new), "reports": sum(g["reports"] for g in new)},
        "recurring": {"problems": len(rec), "reports": sum(g["reports"] for g in rec)},
        "by_issue": issues,
        "insufficient": len(weak), "insufficient_observation_ids": sorted(r["id"] for r in weak),
        "excluded": sum(1 for r in rows if r["state"] == "excluded"),
        "open_actions": sum(1 for a in actions if a["status"] in OPEN_ACTION_STATUSES),
        "action_ids": [a["id"] for a in actions],
        "competitor": sum(1 for r in market if r["kind"] == "competitor"),
        "demand": sum(1 for r in market if r["kind"] == "demand"),
        "rules": {"recurring": f"the same product and issue in >= {RECUR_CASES} cases or from >= {RECUR_SHOPS} "
                               f"shops within {RECUR_DAYS} days",
                  "insufficient": "no photo, evidence level 1 and a single report: listed, not counted"},
    }
    return {"rows": rows, "counted": counted, "insufficient": weak, "market": market, "groups": order,
            "actions": actions, "counts": counts, "included": sorted(inc),
            "missing": sorted((wanted or set()) - ids)}


# ── the narrative: template sentences, every finding with its ids ────────────

def _cite(ids, cap: int = MAX_CITES, sheet: str | None = None) -> str:
    ids = sorted({int(i) for i in ids})
    out = " ".join(f"[MI-{i}]" for i in ids[:cap])
    if len(ids) > cap:
        out += f" and {len(ids) - cap} more" + (f" on the {sheet} sheet" if sheet else "")
    return out


def _quoted(values: list[str], cap: int = 3) -> str:
    shown = ", ".join(f'"{v}"' for v in values[:cap])
    return shown + (f" and {len(values) - cap} more" if len(values) > cap else "")


def _where(issue_label, sku, product) -> str:
    issue = f'"{issue_label}"'
    if sku:
        return f"{issue} on {sku} ({product})" if product and product != sku else f"{issue} on {sku}"
    return f"{issue} on an unidentified product" + (f' ("{product}")' if product else "")


def narrative(ctx: dict, a: dict) -> str:
    """One sentence per line. Lines about a finding start 'Recurring:', 'New:', 'Insufficient
    evidence:', 'We ask:', 'Competitor:' or 'Demand:' and cite [MI-<id>]; an action line starts
    'Open action [CA-<id>]'."""
    c = a["counts"]
    sup = " ".join(str(ctx["supplier"].get("name") or "Supplier").split())
    lines = [f"{sup}: {month_label(ctx['first'])}."
             + (" Internal edition: for YQ only, not for the supplier." if ctx["edition"] == "internal" else "")]
    if c["reports"]:
        s = (f"Field reports counted this month: {c['reports']} ({_plural(c['product_issues'], 'product issue')}, "
             f"{c['packaging']} packaging)")
        if c["shops"]:
            s += f" from {_plural(c['shops'], 'shop')}"
        s += f"; {c['verified']} verified by YQ"
        if c["safety"]:
            s += f"; {c['safety']} flagged as a possible safety issue"
        lines.append(s + ".")
    else:
        lines.append("No product or packaging reports were counted this month.")
    for g in a["groups"]:
        sheet = "Packaging" if g["kind"] == "packaging" else "Product issues"
        s = f"{'Recurring' if g['recurring'] else 'New'}: {_where(g['issue']['label'], g['sku'], g['product'])}"
        if g["versions"]:
            s += f", packaging version {_quoted(g['versions'])}"
        if g["batches"]:
            s += f", batch {_quoted(g['batches'])}"
        s += f": {_plural(g['reports'], 'report')}"
        if g["shops"]:
            s += f" from {_plural(g['shops'], 'shop')}"
        s += " " + _cite([r["id"] for r in g["rows"]], sheet=sheet)
        if g["recurring"]:
            cases, shops = g["window"]
            s += f"; {_plural(cases, 'case')}" + (f" from {_plural(shops, 'shop')}" if shops else "") \
                + f" in the last {RECUR_DAYS} days"
        if g["verified"]:
            s += f"; {len(g['verified'])} verified by YQ {_cite(g['verified'], cap=6)}"
        else:
            s += "; not yet verified"
        if g["safety"]:
            s += "; flagged as a possible safety issue"
        s += f"; {_plural(g['photos'], 'photo')} on the Evidence sheet" if g["photos"] else "; no photos shared"
        lines.append(s + ".")
    opened = [x for x in a["actions"] if x["status"] in OPEN_ACTION_STATUSES]
    for x in opened:
        s = f"Open action [CA-{x['id']}]: \"{x['request']}\""
        due = _date_of(x["due_date"])
        if due:
            s += f", due {_day(due)}"
        since = _day(_bh_date(x.get("created_at")))
        if x["status"] == "open":
            s += "; awaiting your response" + (f" since {since}" if since else "")
        else:
            s += "; your response is in and YQ is checking it"
        if x["linked_observation_ids"]:
            s += "; linked reports " + _cite(x["linked_observation_ids"], cap=6)
        lines.append(s + ".")
    if not opened:
        lines.append("Open actions: none.")
    weak_groups: dict[tuple, list[dict]] = {}
    for r in a["insufficient"]:
        weak_groups.setdefault(r["_key"], []).append(r)
    for rs in weak_groups.values():
        top = rs[0]
        many = "one report" if len(rs) == 1 else f"{len(rs)} reports"
        lines.append(f"Insufficient evidence: {many} of {_where(top['issue']['label'], top['sku'], top['product'])} "
                     f"{_cite([r['id'] for r in rs])}; no photo, a single report and not verified, so not "
                     "included in the counts.")
    if ctx["edition"] == "internal":
        comp = [r["id"] for r in a["market"] if r["kind"] == "competitor"]
        dem = [r["id"] for r in a["market"] if r["kind"] == "demand"]
        if comp:
            lines.append(f"Competitor: {_plural(len(comp), 'sighting')} of competing products this month "
                         f"{_cite(comp, sheet='Competitor')}.")
        if dem:
            lines.append(f"Demand: {_plural(len(dem), 'request')} for this range this month {_cite(dem, sheet='Demand')}.")
    asks = [g for g in a["groups"] if (g["recurring"] or g["verified"]) and not g["open_action"]]
    if asks:
        parts = [f"{_where(g['issue']['label'], g['sku'], g['product'])} {_cite([g['rows'][0]['id']])}" for g in asks[:5]]
        more = f"; and {_plural(len(asks) - 5, 'more problem')}" if len(asks) > 5 else ""
        lines.append("We ask: the root cause and a corrective plan for " + "; ".join(parts) + more + ".")
    return "\n".join(lines)


_CITED_MI = re.compile(r"\bMI-(\d+)\b")
_CITED_CA = re.compile(r"\bCA-(\d+)\b")


def citation_problems(text: str, a: dict) -> list[str]:
    """Ids a narrative cites that are not in this report (every MI id must be one of its sightings,
    every CA id one of its actions)."""
    rows = {r["id"] for r in a["rows"]}
    acts = {x["id"] for x in a["actions"]}
    bad = sorted({f"MI-{i}" for i in map(int, _CITED_MI.findall(text or "")) if i not in rows}
                 | {f"CA-{i}" for i in map(int, _CITED_CA.findall(text or "")) if i not in acts})
    return bad


def _clean_narrative(v) -> str:
    s = str(v or "").replace("\r\n", "\n").replace("\r", "\n")
    s = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", s)
    s = "\n".join(" ".join(line.split()) for line in s.split("\n")).strip()
    if len(s) > MAX_NARRATIVE:
        raise SupplierReportError(f"The narrative is limited to {MAX_NARRATIVE:,} characters.")
    return s


# ── the API's shapes ──────────────────────────────────────────────────────────

def _photo_out(p: dict, signed: dict) -> dict:
    path = p.get("thumb_path") or p.get("path")
    return {"ref": f"PH-{p['id']}", "id": p.get("id"), "role": p.get("role") or "product",
            "role_words": ROLE_WORDS.get(p.get("role") or "product", "Product"),
            "caption": None, "thumb_url": signed.get((p.get("bucket"), path))}


def _row_out(r: dict, signed: dict, scrub) -> dict:
    out = {k: v for k, v in r.items() if not k.startswith("_") and k != "weak"}
    out["photos"] = [{**_photo_out(p, signed), "caption": scrub(mi.clean_text(p.get("caption"), 200))}
                     for p in r["_photos"]]
    return out


def _report_out(rep: dict, suppliers: dict | None = None) -> dict:
    counts = rep.get("counts") or {}
    sup = (suppliers or {}).get(_int(rep.get("supplier_id"))) or counts.get("supplier") or {}
    first = _date_of(rep.get("month"))
    downloads = [{"at": d.get("at"), "by": _who(d.get("by")), "sha256": d.get("sha256")}
                 for d in (rep.get("downloads") or []) if isinstance(d, dict)]
    return {"id": rep.get("id"), "supplier_id": rep.get("supplier_id"),
            "supplier": {"id": rep.get("supplier_id"), "name": sup.get("name"), "code": sup.get("code")},
            "month": month_key(first) if first else None, "month_label": month_label(first) if first else None,
            "edition": rep.get("edition"), "status": rep.get("status"), "narrative": rep.get("narrative") or "",
            "narrative_source": counts.get("narrative_source", "template"),
            "counts": {k: v for k, v in counts.items() if k not in ("supplier",)},
            "included_observation_ids": sorted(int(i) for i in (rep.get("included_observation_ids") or [])),
            "insufficient_observation_ids": counts.get("insufficient_observation_ids") or [],
            "approved_by": _who(rep.get("approved_by")), "approved_at": rep.get("approved_at"),
            "file_sha256": rep.get("file_sha256"), "downloads": downloads,
            "created_at": rep.get("created_at"), "created_by": _who(rep.get("created_by")),
            "updated_at": rep.get("updated_at"), "updated_by": _who(rep.get("updated_by"))}


def _preview(ctx: dict, a: dict, rep: dict) -> dict:
    """The report with its rows (every candidate, flagged included / insufficient / excluded) and its
    actions, as a supplier-safe preview: the same masked words the workbook prints."""
    pairs = []
    for r in a["rows"]:
        for p in r["_photos"]:
            pairs.append((p.get("bucket"), p.get("thumb_path") or p.get("path")))
    signed = mi.sign_many(pairs) if pairs else {}
    scrub = ctx["scrub"]
    rows: dict[str, list[dict]] = {"product_issues": [], "packaging": []}
    if ctx["edition"] == "internal":
        rows.update(competitor=[], demand=[])
    sheet = {"quality": "product_issues", "packaging": "packaging", "competitor": "competitor", "demand": "demand"}
    rank = {"included": 0, "insufficient": 1, "excluded": 2}
    for r in sorted(a["rows"], key=lambda r: (rank[r["state"]], r["date"], r["id"])):
        rows[sheet[r["kind"]]].append(_row_out(r, signed, scrub))
    warnings = [f"MI-{i} no longer belongs to this supplier's {month_label(ctx['first'])} reports; "
                "edit the rows or build a new report." for i in a["missing"]]
    return {"available": True, "report": _report_out(rep, {ctx["supplier_id"]: ctx["supplier"]}),
            "rows": rows, "groups": [{"sku": g["sku"], "product": g["product"], "issue": g["issue"], "kind": g["kind"],
                                      "recurring": g["recurring"], "reports": g["reports"], "shops": g["shops"],
                                      "verified": len(g["verified"]), "safety": g["safety"],
                                      "observation_ids": [r["id"] for r in g["rows"]]} for g in a["groups"]],
            "actions": [{k: v for k, v in x.items() if not k.startswith("_")} for x in a["actions"]],
            "live_counts": a["counts"], "warnings": warnings}


# ── the routes' work ──────────────────────────────────────────────────────────

def _report(report_id) -> dict:
    rid = _int(report_id)
    got: list[dict] = []
    if rid is not None:
        got = get_client().table("supplier_reports").select("*").eq("id", rid).limit(1).execute().data or []
    if not got:
        raise SupplierReportError("That supplier report does not exist.", 404)
    return got[0]


def _ctx_for(rep: dict, *, now: datetime | None = None) -> dict:
    first = _date_of(rep.get("month"))
    if not first:
        raise SupplierReportError("That report has no month.", 409)
    as_of = _ts(rep.get("approved_at")) if rep.get("status") != "draft" else None
    return gather(rep.get("supplier_id"), first.replace(day=1), str(rep.get("edition") or "supplier"),
                  now=now, as_of=as_of)


@_tables
def build_draft(supplier_id, month, edition: str, actor: str, role: str, *, now: datetime | None = None) -> dict:
    """Build the month's draft for a supplier and edition. An existing draft of the same supplier,
    month and edition is rebuilt in place (hand edits are replaced: "rebuilt": true)."""
    _need(role, REVIEW_ROLES, "Only the office can build supplier reports.")
    now = now or _now()
    edition = str(edition or "supplier").strip().lower()
    if edition not in EDITIONS:
        raise SupplierReportError("The edition is 'supplier' or 'internal'.")
    first = parse_month(month, now=now)
    ctx = gather(supplier_id, first, edition, now=now)
    a = assemble(ctx)
    counts = {**a["counts"], "narrative_source": "template", "built_at": _iso(now)}
    text = narrative(ctx, a)
    c = get_client()
    drafts = (c.table("supplier_reports").select("*").eq("supplier_id", ctx["supplier_id"])
              .eq("month", first.isoformat()).eq("edition", edition).eq("status", "draft")
              .order("id", desc=True).limit(1).execute().data or [])
    body = {"narrative": text, "included_observation_ids": a["included"], "counts": counts,
            "updated_by": actor or None, "updated_at": _iso(now)}
    if drafts:
        got = c.table("supplier_reports").update(body).eq("id", drafts[0]["id"]).eq("status", "draft").execute().data
        if got is not None and len(got) == 0:
            raise SupplierReportError("Someone changed this report meanwhile. Reload and try again.", 409)
        rep = (got or [{**drafts[0], **body}])[0]
        rebuilt = True
    else:
        row = {"supplier_id": ctx["supplier_id"], "month": first.isoformat(), "edition": edition, "status": "draft",
               "created_by": actor or None, **body}
        rep = (c.table("supplier_reports").insert(row).execute().data or [row])[0]
        rebuilt = False
    return {**_preview(ctx, a, rep), "rebuilt": rebuilt}


@_tables
def list_reports(role: str, *, supplier_id=None, month=None, edition=None, status=None, offset: int = 0,
                 limit: int = 50) -> dict:
    _need(role, APPROVE_ROLES, "Supplier reports are for the office and management.")
    lim = max(1, min(int(limit or 50), 200))
    off = max(0, int(offset or 0))
    q = get_client().table("supplier_reports").select("*", count="exact")
    if supplier_id not in (None, ""):
        if _int(supplier_id) is None:
            raise SupplierReportError("Unknown supplier.")
        q = q.eq("supplier_id", _int(supplier_id))
    if month not in (None, ""):
        q = q.eq("month", parse_month(month, future_ok=True).isoformat())
    if edition not in (None, ""):
        if edition not in EDITIONS:
            raise SupplierReportError("The edition is 'supplier' or 'internal'.")
        q = q.eq("edition", edition)
    if status not in (None, ""):
        if status not in STATUSES:
            raise SupplierReportError("The status is draft, approved or superseded.")
        q = q.eq("status", status)
    res = q.order("month", desc=True).order("id", desc=True).range(off, off + lim - 1).execute()
    rows = res.data or []
    sup_ids = sorted({_int(r.get("supplier_id")) for r in rows} - {None})
    suppliers = {int(s["id"]): s for s in _by("suppliers", "id", sup_ids)} if sup_ids else {}
    out = []
    for r in rows:
        o = _report_out(r, suppliers)
        o.pop("narrative", None)
        o["downloads"] = len(o["downloads"])
        out.append(o)
    total = res.count if getattr(res, "count", None) is not None else off + len(rows)
    return {"available": True, "rows": out, "count": int(total), "offset": off, "limit": lim}


@_tables
def get_report(report_id, role: str, *, now: datetime | None = None) -> dict:
    _need(role, APPROVE_ROLES, "Supplier reports are for the office and management.")
    rep = _report(report_id)
    ctx = _ctx_for(rep, now=now)
    return _preview(ctx, assemble(ctx, rep.get("included_observation_ids") or []), rep)


@_tables
def update_draft(report_id, changes: dict, actor: str, role: str, *, now: datetime | None = None) -> dict:
    """Reviewers, a draft only: {narrative?, included_observation_ids?}. Every id must be one of this
    supplier's sightings of the month; every id the narrative cites must be in the report. New rows
    with an untouched (template) narrative rewrite it; a narrative edited by hand is kept."""
    _need(role, REVIEW_ROLES, "Only the office can edit supplier reports.")
    now = now or _now()
    rep = _report(report_id)
    if rep.get("status") != "draft":
        raise SupplierReportError(f"Only a draft can be edited (this report is {rep.get('status')}).", 409)
    ctx = _ctx_for(rep, now=now)
    counts_before = rep.get("counts") or {}
    included = rep.get("included_observation_ids") or []
    if changes.get("included_observation_ids") is not None:
        want = []
        for x in changes["included_observation_ids"]:
            if _int(x) is None:
                raise SupplierReportError("included_observation_ids holds sighting ids (numbers).")
            want.append(int(x))
        known = {int(e["o"]["id"]) for e in ctx["entries"]}
        foreign = sorted(set(want) - known)
        if foreign:
            shown = ", ".join(f"MI-{i}" for i in foreign[:5])
            raise SupplierReportError(f"{shown} {'is' if len(foreign) == 1 else 'are'} not among this supplier's "
                                      f"reports for {month_label(ctx['first'])}.")
        included = sorted(set(want))
    a = assemble(ctx, included)
    source = counts_before.get("narrative_source", "template")
    text = rep.get("narrative") or ""
    if changes.get("narrative") is not None:
        text, source = _clean_narrative(changes["narrative"]), "edited"
    elif changes.get("included_observation_ids") is not None and source == "template":
        text = narrative(ctx, a)
    bad = citation_problems(text, a)
    if bad:
        raise SupplierReportError("The narrative cites " + ", ".join(bad[:5]) + ", which "
                                  + ("is" if len(bad) == 1 else "are") + " not in this report.")
    body = {"narrative": text, "included_observation_ids": a["included"],
            "counts": {**a["counts"], "narrative_source": source, "built_at": counts_before.get("built_at")},
            "updated_by": actor or None, "updated_at": _iso(now)}
    got = (get_client().table("supplier_reports").update(body).eq("id", rep["id"]).eq("status", "draft")
           .execute().data)
    if got is not None and len(got) == 0:
        raise SupplierReportError("Someone changed this report meanwhile. Reload and try again.", 409)
    out = _preview(ctx, a, (got or [{**rep, **body}])[0])
    if source == "edited" and changes.get("included_observation_ids") is not None and changes.get("narrative") is None:
        out["warnings"].append("The narrative was edited by hand, so it was kept: check that it still matches "
                               "the rows.")
    return out


def _still_true(ctx: dict, rep: dict) -> dict:
    """Approve and download re-read the data: a row that has since moved to another supplier (or month)
    or a narrative id that is no longer in the report stops them."""
    a = assemble(ctx, rep.get("included_observation_ids") or [])
    if a["missing"]:
        shown = ", ".join(f"MI-{i}" for i in a["missing"][:5])
        raise SupplierReportError(f"{shown} no longer belong{'s' if len(a['missing']) == 1 else ''} to this "
                                  "supplier's reports for the month: build a new report.", 409)
    bad = citation_problems(rep.get("narrative") or "", a)
    if bad:
        raise SupplierReportError("The narrative cites " + ", ".join(bad[:5]) + ", no longer in this report: "
                                  "edit it or build a new report.", 409)
    return a


@_tables
def approve(report_id, actor: str, role: str, *, now: datetime | None = None) -> dict:
    """Approvers: a draft -> approved (its counts refreshed from the data as it is now); an earlier
    approved report of the same supplier, month and edition becomes superseded."""
    _need(role, APPROVE_ROLES, "Only the office or management can approve supplier reports.")
    now = now or _now()
    rep = _report(report_id)
    if rep.get("status") != "draft":
        raise SupplierReportError(f"Only a draft can be approved (this report is {rep.get('status')}).", 409)
    ctx = _ctx_for(rep, now=now)
    a = _still_true(ctx, rep)
    c = get_client()
    olds = [r for r in (c.table("supplier_reports").select("id,status").eq("supplier_id", rep["supplier_id"])
                        .eq("month", rep["month"]).eq("edition", rep["edition"]).eq("status", "approved")
                        .execute().data or []) if int(r["id"]) != int(rep["id"])]
    # supersede first (a one-approved-per-month index can never be in the way), then approve; a lost
    # race puts the earlier ones back
    superseded = []
    for r in olds:
        got = (c.table("supplier_reports").update({"status": "superseded", "updated_by": actor or None,
                                                   "updated_at": _iso(now)})
               .eq("id", r["id"]).eq("status", "approved").execute().data)
        if got is None or len(got) > 0:
            superseded.append(int(r["id"]))
    counts = {**a["counts"], "narrative_source": (rep.get("counts") or {}).get("narrative_source", "template"),
              "built_at": (rep.get("counts") or {}).get("built_at")}
    got = (c.table("supplier_reports").update({"status": "approved", "approved_by": actor or None,
                                               "approved_at": _iso(now), "counts": counts,
                                               "updated_by": actor or None, "updated_at": _iso(now)})
           .eq("id", rep["id"]).eq("status", "draft").execute().data)
    if got is not None and len(got) == 0:
        for i in superseded:
            c.table("supplier_reports").update({"status": "approved"}).eq("id", i).eq("status", "superseded").execute()
        raise SupplierReportError("Someone changed this report meanwhile. Reload and try again.", 409)
    return {"ok": True, "id": int(rep["id"]), "status": "approved", "approved_at": _iso(now),
            "approved_by": _who(actor), "superseded": superseded}


def filename(rep: dict, supplier: dict) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "-", str(supplier.get("code") or supplier.get("name") or "")).strip("-")[:40]
    first = _date_of(rep.get("month"))
    return (f"YQ-supplier-report-{slug or 'supplier-' + str(rep.get('supplier_id'))}-"
            f"{month_key(first) if first else 'month'}-{rep.get('edition') or 'supplier'}.xlsx")


@_tables
def download(report_id, actor: str, role: str, *, now: datetime | None = None) -> dict:
    """Approvers and reviewers, an approved report only: the workbook, built now from the approved rows
    and narrative. The download is recorded on the row ({at, by, sha256}; file_sha256) before the
    bytes are handed out; a download that cannot be recorded is refused."""
    _need(role, APPROVE_ROLES, "Supplier reports are for the office and management.")
    now = now or _now()
    rep = _report(report_id)
    if rep.get("status") != "approved":
        raise SupplierReportError(f"Only an approved report can be downloaded (this report is {rep.get('status')}).",
                                  409)
    ctx = _ctx_for(rep, now=now)
    a = _still_true(ctx, rep)
    data = render_workbook(ctx, a, rep)
    sha = hashlib.sha256(data).hexdigest()
    entry = {"at": _iso(now), "by": actor or None, "sha256": sha}
    c = get_client()
    for _attempt in range(3):
        cur = _report(rep["id"])
        if cur.get("status") != "approved":
            raise SupplierReportError("This report was superseded meanwhile. Download the newer one.", 409)
        q = c.table("supplier_reports").update({"downloads": list(cur.get("downloads") or []) + [entry],
                                                "file_sha256": sha, "updated_at": _iso(_now())}) \
            .eq("id", cur["id"]).eq("status", "approved")
        if cur.get("updated_at"):
            q = q.eq("updated_at", cur["updated_at"])
        got = q.execute().data
        if got is None or len(got) > 0:
            return {"filename": filename(rep, ctx["supplier"]), "data": data, "sha256": sha, "bytes": len(data),
                    "report_id": int(rep["id"]), "edition": rep.get("edition"), "month": month_key(ctx["first"])}
    raise SupplierReportError("The download could not be recorded. Try again.", 409)


# ── photos for the workbook ───────────────────────────────────────────────────

def _jpeg(data: bytes, px: int) -> bytes | None:
    """A photo as a JPEG of at most px on its longest side, without metadata; None if unreadable."""
    from PIL import Image, ImageOps
    try:
        img = Image.open(io.BytesIO(data))
        if img.size[0] * img.size[1] > mi.MAX_PIXELS:
            return None
        img.draft("RGB", (px, px))
        img.load()
        img = ImageOps.exif_transpose(img)
        if img.mode in ("RGBA", "LA", "P", "PA"):
            rgba = img.convert("RGBA")
            base = Image.new("RGB", rgba.size, (255, 255, 255))
            base.paste(rgba, mask=rgba.getchannel("A"))
            img = base
        elif img.mode != "RGB":
            img = img.convert("RGB")
        img.thumbnail((px, px), Image.Resampling.LANCZOS)
        img.info = {}
        out = io.BytesIO()
        img.save(out, format="JPEG", quality=72, optimize=True)
        return out.getvalue()
    except Exception:  # noqa: BLE001 — truncated, corrupt or not a picture
        return None


def fetch_photos(photos: list[dict]) -> dict[int, dict | None]:
    """{photo id: {"evidence": 640 px JPEG, "thumb": 160 px JPEG, "sha": hex | None} or None} — the
    master from the private bucket (its 320 px thumbnail when the master cannot be read). None = the
    photo could not be fetched: the sheet says "photo not available"."""
    def one(p: dict):
        bucket = p.get("bucket") if p.get("bucket") in mi.BUCKETS else None
        for path, master in ((p.get("path"), True), (p.get("thumb_path"), False)):
            if not bucket or not path:
                continue
            try:
                raw = get_client().storage.from_(bucket).download(path)
            except Exception as e:  # noqa: BLE001
                log.info("supplier report: photo %s not fetched: %s", p.get("id"), e)
                continue
            if not raw:
                continue
            ev = _jpeg(bytes(raw), EVIDENCE_PX)
            if ev is None:
                continue
            sha = p.get("sha256") or (hashlib.sha256(bytes(raw)).hexdigest() if master else None)
            return int(p["id"]), {"evidence": ev, "thumb": _jpeg(bytes(raw), THUMB_PX), "sha": sha}
        return int(p["id"]), None
    if not photos:
        return {}
    with ThreadPoolExecutor(max_workers=4, thread_name_prefix="sr-photo") as ex:
        return dict(ex.map(one, photos))


# ── the workbook ──────────────────────────────────────────────────────────────

FINDING_COLUMNS = [("ID", 10), ("Date", 11), ("SKU", 14), ("Product", 30), ("Issue", 24), ("Finding", 17),
                   ("Description", 46), ("Qty", 6), ("Batch / packaging", 20), ("Basis", 15),
                   ("Evidence level", 22), ("Status", 18), ("Photo", 16), ("Evidence refs", 18)]
PACKAGING_COLUMNS = FINDING_COLUMNS[:9] + [("Packaging version", 16)] + FINDING_COLUMNS[9:]
EVIDENCE_COLUMNS = [("Photo ref", 11), ("Report ID", 11), ("SKU", 14), ("Issue", 24), ("Role", 14), ("Caption", 34),
                    ("SHA-256 (short)", 16), ("Photo", 48)]
ACTION_COLUMNS = [("Action", 9), ("Request", 42), ("Owner", 16), ("Due", 12), ("Supplier response", 36),
                  ("YQ check", 30), ("Status", 24), ("Days open", 10), ("Linked reports", 22)]
COMPETITOR_COLUMNS = [("ID", 10), ("Date", 11), ("Product seen", 30), ("Brand", 14), ("Competitor", 18),
                      ("Category", 14), ("Price seen (BHD)", 14), ("YQ SKU", 14), ("Note", 40), ("Area", 14),
                      ("Photo", 16)]
DEMAND_COLUMNS = [("ID", 10), ("Date", 11), ("Source", 26), ("Asked for", 30), ("YQ SKU", 14), ("Qty asked", 9),
                  ("Demand level", 12), ("Note", 40), ("Area", 14)]
SHEETS = {"supplier": ["Summary", "Product issues", "Packaging", "Evidence", "Corrective actions"],
          "internal": ["Summary", "Product issues", "Packaging", "Evidence", "Corrective actions", "Competitor",
                       "Demand"]}
SHOWN_THUMB_PX = 96               # a finding row's picture as drawn in the sheet
SHOWN_EVIDENCE_PX = 320           # the Evidence sheet's picture as drawn (the file keeps 640 px)


def render_workbook(ctx: dict, a: dict, rep: dict) -> bytes:
    """The supplier (or internal) workbook. Every free-text cell goes through the scrubber; the bytes
    are stable for the same data (every date in the file is the approval time)."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    scrub, edition = ctx["scrub"], ctx["edition"]
    stamp = _ts(rep.get("approved_at")) or ctx["now"]
    head_fill = PatternFill("solid", fgColor="6D28D9")
    head_font = Font(color="FFFFFF", bold=True)
    bold = Font(bold=True)
    top = Alignment(vertical="top", wrap_text=True)

    finding_rows = [r for r in a["rows"] if r["kind"] in FINDING_KINDS and r["state"] in ("included", "insufficient")]
    finding_rows.sort(key=lambda r: (r["state"] != "included", r["date"], r["id"]))
    market_rows = sorted(a["market"], key=lambda r: (r["date"], r["id"]))
    shown_photos = [p for r in finding_rows + market_rows for p in r["_photos"]]
    pics = fetch_photos(shown_photos[:MAX_PHOTOS])

    wb = Workbook()
    wb.properties.creator = "YQ Bahrain"
    wb.properties.lastModifiedBy = "YQ Bahrain"
    wb.properties.title = f"Supplier report {month_key(ctx['first'])}"
    wb.properties.created = stamp.astimezone(timezone.utc).replace(tzinfo=None)

    def put(ws, row: int, col: int, value, *, text: bool = False):
        v = scrub(value) if text and isinstance(value, str) else value
        if isinstance(v, str):
            v = ILLEGAL_CHARACTERS_RE.sub("", v)      # a control character would make openpyxl refuse the file
        cell = ws.cell(row=row, column=col, value=v)
        if isinstance(v, str) and v.startswith("="):
            cell.data_type = "s"          # typed text is never a formula
        cell.alignment = top
        return cell

    def head(ws, row: int, cols: list[tuple[str, int]], *, widths: bool = True) -> None:
        for j, (h, w) in enumerate(cols, start=1):
            cell = ws.cell(row=row, column=j, value=h)
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if widths:
                ws.column_dimensions[get_column_letter(j)].width = w

    def picture(ws, row: int, col: int, jpeg: bytes, shown: int) -> float:
        im = XLImage(io.BytesIO(jpeg))
        scale = shown / max(im.width, im.height, 1)
        im.width, im.height = max(1, int(im.width * scale)), max(1, int(im.height * scale))
        ws.add_image(im, f"{get_column_letter(col)}{row}")
        return im.height * 0.75 + 6

    def refs(r: dict) -> str:
        return ", ".join(f"PH-{p['id']}" for p in r["_photos"]) or "none"

    def photo_cell(ws, row: int, col: int, r: dict) -> None:
        if not r["_photos"]:
            put(ws, row, col, "no photo shared" if r["photos_withheld"] else "no photo")
            return
        pid = int(r["_photos"][0]["id"])
        got = pics.get(pid)
        if got and got.get("thumb"):
            ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 15,
                                                picture(ws, row, col, got["thumb"], SHOWN_THUMB_PX))
        else:
            put(ws, row, col, PHOTO_NOT_AVAILABLE if pid in pics else "see the Evidence sheet")

    # 1. Summary
    ws = wb.active
    ws.title = "Summary"
    for col, w in zip("ABCDEF", (36, 30, 12, 12, 12, 18)):
        ws.column_dimensions[col].width = w
    c = a["counts"]
    sup_name = " ".join(str(ctx["supplier"].get("name") or "Supplier").split())
    put(ws, 1, 1, f"{sup_name}: product feedback, {month_label(ctx['first'])}", text=True).font = Font(bold=True, size=14)
    status = (f"Approved {_day(_bh_date(rep.get('approved_at')))}" if rep.get("status") == "approved"
              else "Draft, not approved")
    facts = [("Period", f"{_day(ctx['start'].date())} to {_day((ctx['end'] - timedelta(days=1)).date())} (Bahrain time)"),
             ("Edition", EDITION_LABELS[edition]), ("Status", status),
             ("Reports counted", c["reports"]), ("Product issues", c["product_issues"]), ("Packaging", c["packaging"]),
             ("Shops reporting", c["shops"]),
             ("New problems", f"{c['new']['problems']} ({_plural(c['new']['reports'], 'report')})"),
             ("Recurring problems", f"{c['recurring']['problems']} ({_plural(c['recurring']['reports'], 'report')})"),
             ("Verified by YQ", c["verified"]), ("Possible safety issues", c["safety"]),
             ("Open corrective actions", c["open_actions"]),
             ("Not counted: insufficient evidence", c["insufficient"])]
    if edition == "internal":
        facts += [("Competitor sightings", c["competitor"]), ("Demand requests", c["demand"])]
    r = 3
    for label, value in facts:
        put(ws, r, 1, label).font = bold
        put(ws, r, 2, value)
        r += 1
    r += 1
    put(ws, r, 1, "Counts by issue").font = Font(bold=True, size=12)
    r += 1
    head(ws, r, [("Issue", 0), ("Kind", 0), ("Reports", 0), ("Shops", 0), ("Verified", 0), ("New or recurring", 0)],
         widths=False)
    r += 1
    if not c["by_issue"]:
        put(ws, r, 1, "No reports counted this month.")
        r += 1
    for b in c["by_issue"]:
        put(ws, r, 1, b["label"], text=True)
        put(ws, r, 2, KIND_WORDS.get(b["kind"], b["kind"]))
        put(ws, r, 3, b["reports"])
        put(ws, r, 4, b["shops"])
        put(ws, r, 5, b["verified"])
        put(ws, r, 6, "Recurring" if b["recurring"] else "New")
        r += 1
    r += 1
    put(ws, r, 1, "Narrative").font = Font(bold=True, size=12)
    r += 1
    for line in [x for x in str(rep.get("narrative") or "").split("\n") if x.strip()] or ["(no narrative)"]:
        put(ws, r, 1, line, text=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        ws.row_dimensions[r].height = min(409, 15 * max(1, -(-len(line) // 105)))   # Excel's tallest row
        r += 1
    r += 1
    put(ws, r, 1, "Open corrective actions").font = Font(bold=True, size=12)
    r += 1
    head(ws, r, [("Action", 0), ("Request", 0), ("Due", 0), ("Days open", 0), ("Status", 0), ("Linked reports", 0)],
         widths=False)
    r += 1
    opened = [x for x in a["actions"] if x["status"] in OPEN_ACTION_STATUSES]
    if not opened:
        put(ws, r, 1, "None.")
        r += 1
    for x in opened:
        put(ws, r, 1, x["ref"])
        put(ws, r, 2, x["request"], text=True)
        put(ws, r, 3, x["due_date"] or "")
        put(ws, r, 4, x["days_open"])
        put(ws, r, 5, x["status_words"])
        put(ws, r, 6, ", ".join(f"MI-{i}" for i in x["linked_observation_ids"]) or "none")
        r += 1
    r += 1
    put(ws, r, 1, "How to read this report").font = Font(bold=True, size=12)
    r += 1
    notes = ["IDs: MI = a field report, PH = a photo, CA = a corrective action. Every sentence about a finding "
             "cites the reports it rests on.",
             "Evidence levels: " + " > ".join(EVIDENCE_WORDS[i] for i in sorted(EVIDENCE_WORDS)) + ". "
             "'Verified' = inspected and verified by YQ or better; 'Unverified report' = not verified yet.",
             f"Recurring = the same product and issue in {RECUR_CASES} or more cases, or from {RECUR_SHOPS} or "
             f"more shops, within {RECUR_DAYS} days. Insufficient evidence (no photo, a single report, not "
             "verified) is listed but not counted.",
             "Reports come from the YQ field team. Shop and staff names are not shared; contact details and "
             "amounts are masked. Original photos are available on request."]
    for line in notes:
        put(ws, r, 1, line)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        ws.row_dimensions[r].height = 30
        r += 1

    # 2-3. Product issues, Packaging
    for title, kind, cols in (("Product issues", "quality", FINDING_COLUMNS), ("Packaging", "packaging", PACKAGING_COLUMNS)):
        ws = wb.create_sheet(title)
        head(ws, 1, cols)
        ws.freeze_panes = "A2"
        names = [h for h, _w in cols]
        row = 2
        for x in [y for y in finding_rows if y["kind"] == kind]:
            batch = "; ".join(v for v in (f"batch {x['batch_ref']}" if x["batch_ref"] else "",
                                          f"pack {x['packaging_version']}" if x["packaging_version"] else "") if v)
            values = {"ID": x["ref"], "Date": x["date"], "SKU": x["sku"] or "not identified",
                      "Product": x["product"] or "", "Issue": x["issue"]["label"],
                      "Finding": "Insufficient evidence (not counted)" if x["state"] == "insufficient" else x["finding"],
                      "Description": x["description"], "Qty": x["qty"], "Batch / packaging": batch,
                      "Packaging version": x["packaging_version"] or "", "Basis": x["basis"] or "",
                      "Evidence level": x["evidence_label"], "Status": x["status"],
                      "Evidence refs": refs(x)}
            for j, name in enumerate(names, start=1):
                if name == "Photo":
                    photo_cell(ws, row, j, x)
                else:
                    put(ws, row, j, values.get(name), text=name in ("Product", "Issue", "Description",
                                                                  "Batch / packaging", "Packaging version"))
            row += 1
        if row == 2:
            put(ws, 2, 1, "No reports this month.")

    # 4. Evidence: one row per photo
    ws = wb.create_sheet("Evidence")
    head(ws, 1, EVIDENCE_COLUMNS)
    ws.freeze_panes = "A2"
    row = 2
    for x in finding_rows + market_rows:
        for p in x["_photos"]:
            put(ws, row, 1, f"PH-{p['id']}")
            put(ws, row, 2, x["ref"])
            put(ws, row, 3, x["sku"] or "not identified")
            put(ws, row, 4, (x.get("issue") or {}).get("label") or KIND_WORDS.get(x["kind"], x["kind"].title()), text=True)
            put(ws, row, 5, ROLE_WORDS.get(p.get("role") or "product", "Product"))
            put(ws, row, 6, mi.clean_text(p.get("caption"), 200) or "", text=True)
            if int(p["id"]) not in pics:
                put(ws, row, 7, (str(p.get("sha256") or "")[:12]) or "-")
                put(ws, row, 8, f"not embedded: the workbook holds at most {MAX_PHOTOS} photos")
            elif pics[int(p["id"])] is None:
                put(ws, row, 7, (str(p.get("sha256") or "")[:12]) or "-")
                put(ws, row, 8, PHOTO_NOT_AVAILABLE)
            else:
                got = pics[int(p["id"])]
                put(ws, row, 7, (got["sha"] or "")[:12] or "-")
                ws.row_dimensions[row].height = picture(ws, row, 8, got["evidence"], SHOWN_EVIDENCE_PX)
            row += 1
    if row == 2:
        put(ws, 2, 1, "No photos this month.")

    # 5. Corrective actions
    ws = wb.create_sheet("Corrective actions")
    head(ws, 1, ACTION_COLUMNS)
    ws.freeze_panes = "A2"
    row = 2
    for x in a["actions"]:
        for j, (v, text) in enumerate(((x["ref"], False), (x["request"], True), (x["owner"] or "", True),
                                       (x["due_date"] or "", False), (x["supplier_response"] or "", True),
                                       (x["resolution_check"] or "", True), (x["status_words"], False),
                                       (x["days_open"], False),
                                       (", ".join(f"MI-{i}" for i in x["linked_observation_ids"]) or "none", False)),
                                      start=1):
            put(ws, row, j, v, text=text)
        row += 1
    if row == 2:
        put(ws, 2, 1, "No corrective actions for this period.")

    # 6-7. Internal edition: Competitor, Demand
    if edition == "internal":
        ws = wb.create_sheet("Competitor")
        head(ws, 1, COMPETITOR_COLUMNS)
        ws.freeze_panes = "A2"
        row = 2
        for x in [y for y in market_rows if y["kind"] == "competitor"]:
            for j, (v, text) in enumerate(((x["ref"], False), (x["date"], False), (x["product_seen"] or "", True),
                                           (x["brand"] or "", True), (x["competitor"] or "", True),
                                           (x["category"] or "", False), (x["price_seen"] or "", False),
                                           (x["sku"] or "", False), (x["note"], True), (x["area"] or "", True)),
                                          start=1):
                put(ws, row, j, v, text=text)
            photo_cell(ws, row, len(COMPETITOR_COLUMNS), x)
            row += 1
        if row == 2:
            put(ws, 2, 1, "No competitor sightings this month.")
        ws = wb.create_sheet("Demand")
        head(ws, 1, DEMAND_COLUMNS)
        ws.freeze_panes = "A2"
        row = 2
        for x in [y for y in market_rows if y["kind"] == "demand"]:
            for j, (v, text) in enumerate(((x["ref"], False), (x["date"], False), (x["source"], False),
                                           (x["asked_for"] or "", True), (x["sku"] or "", False), (x["qty"], False),
                                           (x["demand_level"] or "", False), (x["note"], True),
                                           (x["area"] or "", True)), start=1):
                put(ws, row, j, v, text=text)
            row += 1
        if row == 2:
            put(ws, 2, 1, "No demand requests this month.")

    buf = io.BytesIO()
    wb.save(buf)
    return _stable_zip(buf.getvalue(), stamp)


def _stable_zip(data: bytes, when: datetime) -> bytes:
    """The saved workbook with every zip entry and both document dates set to `when`, so the same
    approved report and data rebuild byte for byte (the logged sha256 then names the content)."""
    t = when.astimezone(timezone.utc)
    stamp = (max(1980, t.year), t.month, t.day, t.hour, t.minute, t.second)
    iso = t.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            body = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                body = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:)",
                              lambda m: m.group(1) + iso + m.group(2), body)
            zi = zipfile.ZipInfo(info.filename, date_time=stamp)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = info.external_attr
            dst.writestr(zi, body)
    return out.getvalue()
