"""Ops engine: the four month-end reconciliations (python -m scripts.ops recon [YYYY-MM], scripts/ops/recon_blocks.py)
and the month workbook's Reconciliation sheet (scripts/ops/workbook/sheets_recon.py), on an invented world.

    python -m tests.test_ops_recon

Same lightweight runner as the other suites (no pytest). No database, no network, no .env: SUPABASE_URL,
SUPABASE_KEY and DATABASE_URL are pinned to ci.invalid BEFORE any import and scripts.ops.db is imported first. The
reads go through a stand-in Source (every method of recon_blocks.Source answered from the world below); the count
results are made by counts.compare itself over the invented stock ledger; one Ledger export is a real workbook read
by the real reader and parser; the kickback tiers are app.shop's own rule. All names and numbers are invented; every
folder is a temporary one.

THE WORLD (the month is September 2026; the data reaches Fri 2 Oct):
  * the stock ledger: two supplier receipts into the main store in September (T-26-09-1: four items, two of them
    variants of one model code; T-26-09-2: one item), one into another store, one in August; transfers to two vans
    and an outlet; sales; three Sales Return vouchers (one item not sold in the month); a van that moved only in
    August (it is on no list);
  * the counts: the main store, Van A, Van B and the outlet, each as counts.compare keeps it;
  * the day book, the register and v_sales: two invoices in September (three lines, BHD 38.500 VAT included);
  * the statement: two reps on one target (tiers 200 / 400 / 600);
  * the Ledger export: four customer accounts (one ends in credit, one is paid off), the month's roll-forward
    opening 275 + SI 275 - RCTV 148 - SRV 110 - CN 20 + other 5 = closing 277, and one return that takes Rep Two
    below tier two.

Covered:
  * every status of every block: main store PROVEN, DIFFERS, NOT COUNTED (no count, a book side that waits, no
    line counted); vans PROVEN, DIFFERS, BOOKS ONLY (no count, not signed, waits; none to count), the books-only
    figures from /yq-moves in its shapes, "not worked out this time" without it; sales PROVEN, DIFFERS (the
    invoicing, an account, a tier), COLLECTION NOT PROVEN (no export, an unreadable one, no ageing, names that do
    not meet); purchases PROVEN, DIFFERS, PAPERS MISSING, PROVEN with no receipt;
  * the load control invoice by invoice, and stated as the load's own verify when the files do not reach the month;
  * the roll-forward arithmetic per account and in total, a running balance that breaks it, the closing against
    the ageing (as on the month end; the nearest after it, the Ledger read to it or stopping before it; a credit the
    ageing prints unsigned; an account the ageing lacks, one the export lacks);
  * the tiers the month's returns would change (app.shop's rule), returns naming no rep; without the export the
    estimate from returned units, nothing judged on it;
  * the purchases: the order by its number, by its codes (Jaccard at least one half, within 200 days), the papers
    on file, a missing paper as an action, cost lines n of m, a variant not provable from the PO, the order's flags,
    the supplier invoice by the lead-time rule, receipts into another store;
  * a part that fails stands aside; a block that cannot be read reads "not worked out this time";
  * the command: recon hands the month over, the default month, a month the data has not reached, the files under
    exports/ops/<D>/ only, the wording guard, masked ID-shaped names, no long digit run, no name typed on a sheet;
    purchases(first, last) for the weekly workbook;
  * the sheet: totals only (no customer, item, voucher, invoice, rep or counter name; stores as the month workbook
    names salesmen), every formula evaluates, the roll-forward's check reads nought, it saves and reads back.
"""
from __future__ import annotations

import inspect
import io
import json
import os
import re
import sys
import tempfile
import time
import traceback
import types
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

os.environ["SUPABASE_URL"] = "https://ci.invalid"
os.environ["SUPABASE_KEY"] = "ci-key"
os.environ["DATABASE_URL"] = "postgresql://ci:ci@ci.invalid:5432/ci"

from scripts.ops import db  # noqa: E402,F401  (first, as every command does)
from scripts.ops import config, counts, fmt, ledger, paths, recon, stock  # noqa: E402
from scripts.ops import ledger_export as le  # noqa: E402
from scripts.ops import recon_blocks as rb  # noqa: E402
from scripts.ops.workbook import book as bk  # noqa: E402
from scripts.ops.workbook import sheets_recon  # noqa: E402

TESTS: list[tuple[str, object]] = []

NINE = re.compile(r"(?<![0-9])[0-9]{9}(?![0-9])")            # shaped like a personal ID number
LONG = re.compile(r"(?<![0-9])[0-9]{15,}(?![0-9])")          # shaped like a card or account number

MAIN, VAN_A, VAN_B, VAN_C, OUTLET = "Main Store", "Van A - Acc WH", "Van B - Acc WH", "Van C - Acc WH", "Causeway"
DAMAGE, SIM_STORE = "Damage Store", "Sim Store"
FIXTURE = {"warehouses.main": MAIN, "warehouses.aliases": {}, "warehouses.not_counted": [DAMAGE],
           "warehouses.damage": DAMAGE, "items.bridge_exceptions": {}, "items.not_accessories": []}
MONTH, FIRST, LAST = "2026-09", date(2026, 9, 1), date(2026, 9, 30)
D = date(2026, 10, 2)                       # the data date
NAMES = {"A": "ITEM-A Cable Alpha", "B": "ITEM-B Charger Beta", "C": "ITEM-C Holder Gamma",
         "DR": "ITEM-D Speaker Red", "DB": "ITEM-D Speaker Blue", "E": "ITEM-E Fan Echo"}
CODES = {"A": "CA1", "B": "CB2", "C": "HG3", "DR": "DS1 RD", "DB": "DS1 BL", "E": "SE5"}
SIM = "SIM-01 Starter Pack SP9"
COSTS = {"A": Decimal("0.4"), "B": Decimal("2.5"), "C": Decimal("2"), "E": Decimal("2.5")}   # the variants: none
COUNTER, CHECKER = "Counter Person", "Checker Person"     # typed on the sheets: never printed, never kept here
REP1, REP2 = "Rep One", "Rep Two"
CUST_A, CUST_B, CUST_C, CUST_D, CUST_Z = "Customer A", "Customer B", "Customer C", "Customer D", "Customer Z"
SALES_AC, VAT_AC, RETURNS_AC, CASH_AC = "Sales Accessories", "Output VAT", "Sales Returns", "Cash in Hand"
NO_EXPORT = "no drop folder is dated 2026-09-30 to 2026-10-10, where the Ledger of 2026-09 is looked for"


def L(key: str) -> str:
    """The ledger's name of an item: its stock name, a space, its code."""
    return f"{NAMES[key]} {CODES[key]}"


def test(name):
    def deco(fn):
        TESTS.append((name, fn))
        return fn
    return deco


def raises(kind, fn) -> Exception:
    try:
        fn()
    except kind as exc:
        return exc
    raise AssertionError(f"expected {kind.__name__}")


# ── the stock ledger ──────────────────────────────────────────────────────────

class Book:
    """Movements in the shape scripts.ops.ledger gives them."""

    def __init__(self):
        self.rows: list[dict] = []
        self.n = 0

    def _row(self, item, day, vtype, wh, voucher, qin=0, qout=0, vin=0, vout=0, to_wh=None):
        self.rows.append({"item": item, "date": day, "voucher": voucher, "vtype": vtype, "wh": wh, "to_wh": to_wh,
                          "qty_in": Decimal(str(qin)), "qty_out": Decimal(str(qout)), "value_in": Decimal(str(vin)),
                          "value_out": Decimal(str(vout)), "narr": None})

    def receipt(self, item, day, units, value, voucher, wh=MAIN):
        self._row(item, day, ledger.SUPPLIER_RECEIPT, wh, voucher, qin=units, vin=value)
        return self

    def move(self, item, day, units, frm, to):
        self.n += 1
        self._row(item, day, ledger.ISSUE, frm, f"SIO:T-26-{day.month:02d}-{self.n}", qout=units, to_wh=to)
        self._row(item, day, ledger.RECEIVE, to, f"STRV:T-26-{day.month:02d}-{self.n}", qin=units)
        return self

    def sale(self, item, day, units, wh):
        self.n += 1
        self._row(item, day, ledger.SALE, wh, f"SI:T-{self.n}", qout=units)
        return self

    def back(self, item, day, units, wh, voucher):
        """A Sales Return: units back into `wh`."""
        self._row(item, day, ledger.SALES_RETURN, wh, voucher, qin=units)
        return self


def world() -> Book:
    b = Book()
    b.receipt(L("A"), date(2026, 8, 10), 100, 40, "MRN:T-26-08-4").receipt(L("B"), date(2026, 8, 10), 50, 125,
                                                                         "MRN:T-26-08-4")
    b.move(L("B"), date(2026, 8, 12), 4, MAIN, VAN_C).move(L("B"), date(2026, 8, 20), 4, VAN_C, MAIN)   # Van C: August
    for k, units, value in (("A", 200, 80), ("C", 30, 60), ("DR", 10, 30), ("DB", 10, 30)):
        b.receipt(L(k), date(2026, 9, 5), units, value, "MRN:T-26-09-1")
    b.receipt(L("E"), date(2026, 9, 20), 40, 100, "MRN:T-26-09-2")
    b.receipt(SIM, date(2026, 9, 21), 500, 50, "MRN:T-26-09-7", wh=SIM_STORE)       # into another store
    b.move(L("A"), date(2026, 9, 6), 30, MAIN, VAN_A).move(L("B"), date(2026, 9, 6), 10, MAIN, VAN_B)
    b.move(L("C"), date(2026, 9, 6), 5, MAIN, OUTLET)
    b.sale(L("A"), date(2026, 9, 10), 10, VAN_A).sale(L("B"), date(2026, 9, 11), 2, VAN_B)
    b.sale(SIM, date(2026, 9, 12), 5, SIM_STORE)
    b.back(L("A"), date(2026, 9, 15), 3, VAN_A, "SRV:T-1")
    b.back(L("B"), date(2026, 9, 16), 1, VAN_B, "SRV:T-2").back(L("B"), date(2026, 9, 16), 1, VAN_B, "SRV:T-2")
    b.back(L("C"), date(2026, 9, 18), 2, OUTLET, "SRV:T-3")                         # not sold in the month
    b.back(SIM, date(2026, 9, 19), 1, SIM_STORE, "SRV:T-4")                          # no accessory
    b.sale(L("A"), date(2026, 10, 1), 1, VAN_A)
    b.move(L("A"), D, 5, MAIN, VAN_A)                                                # the ledger reaches D
    return b


ROWS = world().rows
# At the end of 28 Sep (the book side of a count on the 29th): the main store A 270, B 40, C 25, DR 10, DB 10,
# E 40; Van A A 23 (30 - 10 + 3); Van B B 10 (10 - 2 + 2); the outlet C 7 (5 + 2).


# ── the counts, as counts.compare keeps them ──────────────────────────────────

def rules():
    return counts.store_rules()


def result(store: str, kind: str, counted: dict, *, day: date = date(2026, 9, 29), by=COUNTER, checked=CHECKER,
           added: dict | None = None, prep_kind: str = "cycle", rows=None, costs: dict | None = None) -> dict:
    """One store's count result: counts.compare over the world's ledger. `counted` {item key: units or None}."""
    books = counts.Books(rows if rows is not None else ROWS, rules())
    lines = [{"code": CODES[k], "item": NAMES[k], "ledger": L(k), "counted": n, "note": None, "row": 7 + i,
              "added": False} for i, (k, n) in enumerate(counted.items())]
    extra = [{"code": CODES[k], "item": NAMES[k], "ledger": L(k), "counted": n, "note": None, "row": 30 + i,
              "added": True} for i, (k, n) in enumerate((added or {}).items())]
    sheet = {"sheet": "Sheet", "store": store, "kind": kind, "slug": counts.slug(store),
             "prep_id": "2026-09-27-cycle-ghijkl", "prep_kind": prep_kind, "prep_date": date(2026, 9, 27),
             "count_date": day, "counted_by": by, "checked_by": checked, "lines": lines, "added": extra,
             "sheet_notes": []}
    meta = {"file": "YQ Count Sheets 2026-09-27 cycle.xlsx", "sha256": "sha256:" + "ab" * 32, "saved_at": None,
            "read_at": f"{day.isoformat()}T18:00:00+03:00", "ledger_source": "file", "ledger_last_date": D,
            "ledger_note": "test"}
    unit = {L(k): v for k, v in (COSTS if costs is None else costs).items()}
    return counts.compare(sheet, books, unit, meta)


def main_ok() -> dict:
    return result(MAIN, "main", {"A": 270, "C": 25, "E": 40})


def main_differs() -> dict:
    return result(MAIN, "main", {"A": 268, "C": 26, "E": None})         # 2 short (0.800), 1 over (2.000)


def van_counts(**over) -> dict:
    out = {VAN_A: result(VAN_A, "van", {"A": 23}), VAN_B: result(VAN_B, "van", {"B": 9}),
           OUTLET: result(OUTLET, "outlet", {"C": 7})}
    out.update(over)
    return {k: v for k, v in out.items() if v is not None}


BOOKS = {VAN_A: {"units": 23, "bhd": "9.2", "held": 0},
         VAN_B: {"units": 10, "bhd": "25", "held": [{"item": L("B"), "units": 10, "bhd": "25"}]},
         OUTLET: {"stock": {"units": 7, "bhd": "14"}, "held": {"items": 1, "units": 7, "bhd": "14"}}}
NOT_THERE = object()


# ── sales and receivables ─────────────────────────────────────────────────────

def ln(invoice, day, gross):
    return {"invoice": invoice, "day": day, "gross": gross}


DAYBOOK = (D, [ln("SI : 90", date(2026, 8, 20), 110.0), ln("SI : 101", date(2026, 9, 3), 11.0),
               ln("SI : 101", date(2026, 9, 3), 22.0), ln("SI : 102", date(2026, 9, 25), 5.5),
               ln("SI : 110", date(2026, 10, 1), 7.7)])
REGISTER = (D, [ln("SI : 90", date(2026, 8, 20), 110.0), ln("SI : 101", date(2026, 9, 3), 33.0),
                ln("SI : 102", date(2026, 9, 25), 5.5), ln("SI : 110", date(2026, 10, 1), 7.7)])
PORTAL = {"SI : 101": (2, Decimal("33.000")), "SI : 102": (1, Decimal("5.500"))}
SPAN = (date(2026, 8, 20), date(2026, 10, 1))                      # v_sales: the first and the last sale
LOAD = {"run": "021026_0800", "finished_at": "2026-10-02T08:10:00+03:00", "result": "ok"}
TARGET = {"target_bhd": 200, "tier2_bhd": 400, "tier3_bhd": 600, "kickback_t1": 0.005, "kickback_t2": 0.0075,
          "kickback_t3": 0.01}


def statement(rep2_sales: str = "450.000") -> list[dict]:
    rows = []
    for rep, sales in ((REP1, "250.000"), (REP2, rep2_sales)):
        tier, kick = rb.Source().tier(TARGET, Decimal(sales), "2026-09-30")
        rows.append({"salesman": rep, "salesman_id": None, "period": MONTH, "basis": "net_ex_vat",
                     "data_through": "2026-09-30", "sales_bhd": sales, "returns_bhd": None, "tier_reached": tier,
                     "rate": 0.0, "kickback_bhd": str(kick), "target_snapshot": TARGET})
    return rows


PRICES = {NAMES["A"]: (Decimal(10), Decimal("15.000")), NAMES["B"]: (Decimal(2), Decimal("9.000"))}

POSTINGS = [
    (date(2026, 8, 20), "SI : 90", [(CUST_A, 110, 0, REP1), (SALES_AC, 0, 100), (VAT_AC, 0, 10)]),
    (date(2026, 8, 21), "SI : 85", [(CUST_B, 44, 0, REP2), (SALES_AC, 0, 40), (VAT_AC, 0, 4)]),
    (date(2026, 8, 22), "SI : 86", [(CUST_C, 88, 0, REP1), (SALES_AC, 0, 80), (VAT_AC, 0, 8)]),
    (date(2026, 8, 23), "SI : 87", [(CUST_D, 33, 0, REP1), (SALES_AC, 0, 30), (VAT_AC, 0, 3)]),
    (date(2026, 9, 5), "SI : 100", [(CUST_A, 220, 0, REP1), (SALES_AC, 0, 200), (VAT_AC, 0, 20)]),
    (date(2026, 9, 8), "SI : 101", [(CUST_B, 55, 0, REP2), (SALES_AC, 0, 50), (VAT_AC, 0, 5)]),
    (date(2026, 9, 10), "RCTV : 50", [(CASH_AC, 110, 0), (CUST_A, 0, 110, REP1)]),
    (date(2026, 9, 12), "SRV : 7", [(RETURNS_AC, 100, 0), (VAT_AC, 10, 0), (CUST_B, 0, 110, REP2)]),
    (date(2026, 9, 14), "CN : 3", [(SALES_AC, 20, 0), (CUST_C, 0, 20)]),
    (date(2026, 9, 20), "JV : 1", [(CUST_D, 5, 0), (CASH_AC, 0, 5)]),
    (date(2026, 9, 25), "RCTV : 51", [(CASH_AC, 38, 0), (CUST_D, 0, 38, REP1)]),
    (date(2026, 10, 1), "SI : 110", [(CUST_A, 77, 0, REP1), (SALES_AC, 0, 70), (VAT_AC, 0, 7)]),
]
# per account at the end of 30 Sep: A 110 + 220 - 110 = 220; B 44 + 55 - 110 = -11 (a credit); C 88 - 20 = 68;
# D 33 + 5 - 38 = 0. Totals: 275 + 275 - 148 - 110 - 20 + 5 = 277.
AR_EXACT = {CUST_A: Decimal("220"), CUST_B: Decimal("11"), CUST_C: Decimal("68")}
AR_NAMES = {le.fold(a) for a in (CUST_A, CUST_B, CUST_C, CUST_D, CUST_Z)}


def legs_by_account(postings, as_on: date) -> dict:
    out: dict = {}
    for n, (day, voucher, legs) in enumerate(postings):
        if day > as_on:
            continue
        for k, leg in enumerate(legs):
            account, debit, credit = leg[0], Decimal(str(leg[1])), Decimal(str(leg[2]))
            salesman = leg[3] if len(leg) > 3 else None
            counter = next((o[0] for j, o in enumerate(legs) if j != k), None)
            out.setdefault(account, []).append((day, n, k, voucher, counter, debit, credit, salesman))
    return out


def export_of(postings=POSTINGS, as_on: date = LAST, broken: str | None = None) -> le.Export:
    """A Ledger export in memory, in the reader's shape: one section per account, a running balance (`broken`
    names an account whose stated balance is one fils off from its second row on)."""
    rows = []
    for account, legs in legs_by_account(postings, as_on).items():
        bal = Decimal(0)
        for i, (day, _n, _k, voucher, counter, debit, credit, salesman) in enumerate(sorted(legs, key=lambda t: t[:3])):
            bal += debit - credit
            stated = bal + (Decimal("0.001") if broken == account and i >= 1 else 0)
            rows.append({"account": account, "date": day, "voucher": voucher, "prefix": le.prefix_of(voucher),
                         "kind": le.kind_of(voucher), "counter": counter, "debit": debit, "credit": credit,
                         "balance": stated, "currency": "BHD", "payment_mode": None, "salesman": salesman})
    days = [r["date"] for r in rows]
    return le.Export(rows=rows, source="Ledger_test.xlsx", as_on=as_on, first_date=min(days), last_date=max(days))


HEADER = ["Date", "Voucher", "Account", "Debit", "Credit", "Balance", "Currency", "Payment Mode", "Salesman",
          "Narration"]


def write_ledger(path: Path, postings=POSTINGS, as_on: date = LAST) -> Path:
    """The Ledger as Focus exports it (title block, header, a section per account with its Sub Total, a Grand
    Total); the narration holds an ID-shaped run that must never come through."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    pad = [None] * 7
    ws.append(pad + ["Example Trading W.L.L"])
    ws.append(pad + ["Ledger"])
    ws.append([])
    ws.append(pad + [f"[As on date {as_on.strftime('%d/%m/%Y')}]"])
    ws.append(HEADER)
    total_d = total_c = Decimal(0)
    for account, legs in legs_by_account(postings, as_on).items():
        ws.append([account])
        bal = sub_d = sub_c = Decimal(0)
        for day, _n, _k, voucher, counter, debit, credit, salesman in sorted(legs, key=lambda t: t[:3]):
            bal += debit - credit
            sub_d += debit
            sub_c += credit
            ws.append([datetime(day.year, day.month, day.day), voucher, counter, float(debit) or None,
                       float(credit) or None, float(bal), "BHD", None, salesman, "paid by " + "12345" + "6789"])
        ws.append(["Sub Total", None, None, float(sub_d), float(sub_c), float(bal)])
        total_d += sub_d
        total_c += sub_c
    ws.append(["Grand Total", None, None, float(total_d), float(total_c)])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()
    return path


# ── purchases ─────────────────────────────────────────────────────────────────

def po(no, day, *codes):
    return [{"po_no": no, "po_date": day, "code": c} for c in codes]


PO_ROWS = (po("T-26-09-1", "2026-08-01", "CA1", "HG3", "DS1") + po("T-26-07-3", "2026-07-15", "SE5", "ZZ9")
           + po("T-25-12-1", "2025-12-01", "SE5") + po("T-26-08-4", "2026-07-01", "CA1", "CB2"))
FILE_ROWS = [{"po_no": "T-26-09-1", "kind": k} for k in ("po", "mrn", "invoice")] + [
    {"po_no": "T-26-07-3", "kind": "po"}, {"po_no": "T-26-07-3", "kind": "photo"}]
MRN_ROWS = [{"doc_no": "T-26-09-1", "code": c} for c in ("CA1", "HG3", "DS1 RD", "DS1 BL")]
INVOICE_ROWS = [{"invoice_no": "INV-9", "invoice_date": "2026-08-01", "model": "ITEM-E", "qty": "40", "vendor": "V"}]


def detail(*lines, status="received"):
    return {"found": True, "status": status, "lines": [{"code": c, "flags": list(f)} for c, f in lines],
            "reconciliation": {"summary": "test"}}


DETAILS = {"T-26-09-1": detail(("CA1", ()), ("HG3", ()), ("DS1", ("thin_margin",))),
           "T-26-07-3": detail(("SE5", ("cost_estimated",)), ("ZZ9", ()), status="pending")}


# ── the stand-in ──────────────────────────────────────────────────────────────

class Fake(rb.Source):
    """recon_blocks.Source answered from the world; `fail` {method name: exception} makes a read fail."""

    def __init__(self, **kw):
        super().__init__("auto")
        self.book_rows = kw.pop("rows", ROWS)
        self.results = kw.pop("results", {})
        self.books = kw.pop("books", NOT_THERE)
        self.files_ = kw.pop("files", ("Focus ERP 2026-10-02", Path("Sales_day_book_t.xlsx"),
                                       Path("Summary_sales_register_t.xlsx")))
        self.daybook, self.register = kw.pop("daybook", DAYBOOK), kw.pop("register", REGISTER)
        self.span = kw.pop("span", SPAN)
        self.portal = kw.pop("portal", PORTAL)
        self.load = kw.pop("load", LOAD)
        self.statement = kw.pop("statement", None)
        self.export = kw.pop("export", None)
        self.ar = kw.pop("ar", {})
        self.names = kw.pop("names", AR_NAMES)
        self.prices = kw.pop("prices", PRICES)
        self.pos, self.files = kw.pop("pos", PO_ROWS), kw.pop("order_files", FILE_ROWS)
        self.mrn, self.invoices = kw.pop("mrn", MRN_ROWS), kw.pop("invoices", INVOICE_ROWS)
        self.details = kw.pop("details", DETAILS)
        self.fail = kw.pop("fail", {})
        assert not kw, kw
        self.calls: list = []

    def _call(self, name, *args):
        self.calls.append((name,) + args)
        if name in self.fail:
            raise self.fail[name]

    def stock_ledger(self):
        self._call("stock_ledger")
        rows = self.book_rows
        led = ledger.Ledger(rows=rows, source="Stock_ledger_test.xlsx", first_date=ledger.first_date(rows),
                            last_date=ledger.last_date(rows), sha=None, origin="file")
        return led, {"source": "file", "parity": None,
                     "note": "The stock ledger is read from the file Stock_ledger_test.xlsx: it agrees with what the "
                             "portal holds."}

    def accessories(self, rows):
        self._call("accessories")
        return {L(k): CODES[k] for k in NAMES}, []

    def count_results(self, month):
        self._call("count_results", month)
        return dict(self.results)

    def books_only(self, as_of):
        self._call("books_only", as_of)
        if self.books is NOT_THERE:
            raise rb.NotYet("/yq-moves cannot give them yet")
        return self.books

    def sales_files(self):
        self._call("sales_files")
        return self.files_

    def read_sales_file(self, path, kind):
        self._call("read_sales_file", kind)
        got = self.daybook if kind == "order_lines" else self.register
        if isinstance(got, Exception):
            raise got
        return got

    def sales_span(self):
        self._call("sales_span")
        return self.span

    def portal_invoices(self, first, last):
        self._call("portal_invoices", first, last)
        return dict(self.portal)

    def last_load(self):
        return dict(self.load)

    def statement_rows(self, month):
        self._call("statement_rows", month)
        return statement() if self.statement is None else self.statement

    def ledger_export(self, month):
        self._call("ledger_export", month)
        if self.export == "real":
            return super().ledger_export(month)
        if isinstance(self.export, Exception):
            raise self.export
        if self.export is None:
            return None, NO_EXPORT
        return self.export, f"the Ledger as on {self.export.as_on.isoformat()}, in 300926/Ledger_test.xlsx"

    def ar_dates(self, lo, hi):
        self._call("ar_dates", lo, hi)
        return [d for d in sorted(self.ar) if lo <= d <= hi]

    def ar_rows(self, as_of):
        self._call("ar_rows", as_of)
        return dict(self.ar[as_of])

    def ar_names(self):
        self._call("ar_names")
        return set(self.names)

    def month_prices(self, first, last):
        self._call("month_prices", first, last)
        return dict(self.prices)

    def po_lines(self):
        self._call("po_lines")
        return list(self.pos)

    def order_files(self):
        self._call("order_files")
        return list(self.files)

    def mrn_lines(self):
        self._call("mrn_lines")
        return list(self.mrn)

    def supplier_invoices(self):
        self._call("supplier_invoices")
        return list(self.invoices)

    def order_detail(self, po_no):
        self._call("order_detail", po_no)
        return self.details.get(po_no, {"found": False})


class data_root:
    """A temporary data root with the config fixture. One opened inside another (BOOKS_VIEW() in a test's own) puts
    the outer one back when it closes, so a config read after it never falls through to the project folder's."""

    def __enter__(self) -> Path:
        self._outer = paths._override               # noqa: SLF001  (the root a test opened around this one, or None)
        self._tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.root = Path(self._tmp.name).resolve()
        paths.set_data_root(self.root)
        paths.write_text(paths.config_file(), json.dumps(FIXTURE))
        config.use(None)
        return self.root

    def __exit__(self, *exc):
        paths.set_data_root(self._outer)
        config.use(None)
        for _ in range(20):
            try:
                self._tmp.cleanup()
                break
            except OSError:
                time.sleep(0.1)
        return False


class Args:
    def __init__(self, *spec, ledger_="auto"):
        self.spec, self.ledger, self.json, self.data_root, self.date = list(spec), ledger_, False, None, None


def blocks(**kw) -> dict:
    return rb.month_blocks(MONTH, Fake(**kw))


def all_lines(b: dict) -> str:
    return "\n".join(b["lines"])


# ── 1: the main store ─────────────────────────────────────────────────────────

@test("main store: PROVEN when every counted line agrees with its book; lines counted n of m")
def _():
    with data_root():
        b = rb.main_block({MAIN: main_ok()}, rules(), MONTH)
        assert b["status"] == rb.PROVEN, b
        v = b["count"]
        assert (v["lines"], v["counted"], v["added"], v["differ"]) == (3, 3, 0, 0)
        assert v["short_units"] == 0 and v["over_units"] == 0 and v["short_bhd"] == 0
        text = all_lines(b)
        assert "3 of 3 lines counted" in text and "Every counted line agrees with its book" in text, text
        assert "Tue 29 Sep" in text and "the end of Mon 28 Sep" in text and "the cycle list" in text, text


@test("main store: DIFFERS (to check) with the units and BHD short and over at landed cost")
def _():
    with data_root():
        b = rb.main_block({MAIN: main_differs()}, rules(), MONTH)
        assert b["status"] == rb.DIFFERS
        v = b["count"]
        assert (v["counted"], v["lines"], v["differ"]) == (2, 3, 2)
        assert (v["short_units"], v["over_units"]) == (Decimal(2), Decimal(1))
        assert (v["short_bhd"], v["over_bhd"]) == (Decimal("0.800"), Decimal("2.000"))
        text = all_lines(b)
        assert "2 of 3 lines counted" in text and "2 lines differ from the book, to check" in text, text
        assert "2 units short and 1 over" in text and "BHD 0.800 short and BHD 2.000 over at landed cost" in text
        assert "agrees" not in text


@test("main store: NOT COUNTED without a count, with a book side that waits, with no line counted; never 'agrees'")
def _():
    with data_root():
        none = rb.main_block({VAN_A: result(VAN_A, "van", {"A": 23})}, rules(), MONTH)
        assert none["status"] == rb.NOT_COUNTED and none["count"] is None
        assert "No count of the main store for September 2026 is on file" in all_lines(none)
        waits = rb.main_block({MAIN: result(MAIN, "main", {"A": 270}, day=date(2026, 10, 9))}, rules(), MONTH)
        assert waits["status"] == rb.NOT_COUNTED and waits["count"]["waits"]
        assert "waits for a load with data to 8 October 2026" in all_lines(waits)
        empty = rb.main_block({MAIN: result(MAIN, "main", {"A": None, "C": None})}, rules(), MONTH)
        assert empty["status"] == rb.NOT_COUNTED and "has no line counted" in all_lines(empty)
        for b in (none, waits, empty):
            assert "agree" not in all_lines(b).replace("never", ""), all_lines(b)


@test("main store: of two counts of the month the newest stands; an added row counts; the count-day movers named")
def _():
    with data_root():
        older = result(MAIN, "main", {"A": 260}, day=date(2026, 9, 27))
        newer = result(MAIN, "main", {"A": 270}, day=date(2026, 9, 29), added={"B": 40})
        b = rb.main_block({"x": older, "y": newer}, rules(), MONTH)
        assert b["status"] == rb.PROVEN and b["count"]["added"] == 1 and "(+1 added)" in all_lines(b)
        moved = result(MAIN, "main", {"A": 270}, day=date(2026, 9, 6))     # 30 went to Van A on the count day
        said = all_lines(rb.main_block({MAIN: moved}, rules(), MONTH))
        assert "moved in the store on the count day" in said, said


# ── 2: the vans and outlets ───────────────────────────────────────────────────

@test("vans: the stores on the books are those that held an accessory at the month end or moved one in it")
def _():
    with data_root():
        acc = {L(k) for k in NAMES}
        assert rb.stores_on_books(ROWS, rules(), acc, FIRST, LAST) == [OUTLET, VAN_A, VAN_B]
        assert VAN_C not in rb.stores_on_books(ROWS, rules(), acc, FIRST, LAST)       # August only, back to nought
        assert rb.stores_on_books(ROWS, rules(), set(), FIRST, LAST) == []


@test("vans: PROVEN when every store's signed count agrees; DIFFERS when one does not")
def _():
    with data_root():
        ok = rb.vans_block(van_counts(**{VAN_B: result(VAN_B, "van", {"B": 10})}), [VAN_A, VAN_B, OUTLET], None,
                           "not worked out this time (test)", rules(), MONTH, LAST)
        assert ok["status"] == rb.PROVEN and [r["status"] for r in ok["rows"]] == [rb.PROVEN] * 3, ok["rows"]
        assert [r["store"] for r in ok["rows"]] == [VAN_A, VAN_B, OUTLET]          # vans first, then the outlets
        assert [r["label"] for r in ok["rows"]] == ["Van A", "Van B", OUTLET]
        bad = rb.vans_block(van_counts(), [VAN_A, VAN_B, OUTLET], None, None, rules(), MONTH, LAST)
        assert bad["status"] == rb.DIFFERS
        b_row = next(r for r in bad["rows"] if r["store"] == VAN_B)
        assert b_row["status"] == rb.DIFFERS and b_row["count"]["short_units"] == 1
        assert b_row["count"]["short_bhd"] == Decimal("2.500")
        assert bad["totals"]["short_units"] == 1 and bad["totals"]["signed"] == 3
        text = all_lines(bad)
        assert "Signed counts" in text and "DIFFERS (to check)" in text and COUNTER not in text and CHECKER not in text


@test("vans: BOOKS ONLY without a signed count (none, Checked by empty, a book side that waits, no line counted)")
def _():
    with data_root():
        res = {VAN_A: result(VAN_A, "van", {"A": 23}, checked=None),
               VAN_B: result(VAN_B, "van", {"B": 10}, day=date(2026, 10, 9)),
               OUTLET: result(OUTLET, "outlet", {"C": None})}
        b = rb.vans_block(res, [VAN_A, VAN_B, OUTLET, "Van D - Acc WH"], BOOKS_VIEW(), None, rules(), MONTH, LAST)
        assert b["status"] == rb.BOOKS_ONLY and all(r["status"] == rb.BOOKS_ONLY for r in b["rows"])
        why = {r["store"]: r["why"] for r in b["rows"]}
        assert "not signed (Checked by is empty)" in why[VAN_A] and "read the file again" in why[VAN_A]
        assert "waits for a load with data to 8 October 2026" in why[VAN_B]
        assert "no line counted" in why[OUTLET] and why["Van D - Acc WH"] == "no signed count of the month"
        text = all_lines(b)
        assert "Books only (4 stores with no signed count of September 2026)" in text, text
        assert "Van A - Acc WH: its count of Tue 29 Sep is not signed" in text
        none = rb.vans_block({}, [], None, None, rules(), MONTH, LAST)
        assert none["status"] == rb.BOOKS_ONLY and "nothing to count" in all_lines(none)


def BOOKS_VIEW():
    with data_root():
        return rb.books_view(BOOKS, rules())


@test("vans: the books-only figures from /yq-moves in its shapes, else 'not worked out this time'")
def _():
    with data_root():
        view = rb.books_view(BOOKS, rules())
        a, bb, o = (view[rules().canon(s)] for s in (VAN_A, VAN_B, OUTLET))
        assert (a["units"], a["bhd"], a["held_items"]) == (Decimal(23), Decimal("9.200"), 0)
        assert (bb["held_items"], bb["held_units"], bb["held_bhd"]) == (1, Decimal(10), Decimal("25.000"))
        assert (o["units"], o["bhd"], o["held_items"], o["held_bhd"]) == (Decimal(7), Decimal("14.000"), 1,
                                                                          Decimal("14.000"))
        listed = rb.books_view([{"store": VAN_A, "units": 5, "cost_bhd": "2"}], rules())
        assert listed[rules().canon(VAN_A)]["bhd"] == Decimal("2.000")
        obj = rb.books_view({VAN_A: types.SimpleNamespace(units=3, bhd=Decimal("1.5"), held=2)}, rules())
        assert obj[rules().canon(VAN_A)]["held_items"] == 2
        raises(ValueError, lambda: rb.books_view({VAN_A: {"colour": "red"}}, rules()))
        # through the month: absent -> the line says so; present -> the table
        m = blocks()
        v = m["blocks"]["vans"]
        assert v["status"] == rb.BOOKS_ONLY and v["books_note"] == "not worked out this time (/yq-moves cannot give them yet)"
        assert "not worked out this time (/yq-moves cannot give them yet)" in all_lines(v)
        m2 = blocks(books=BOOKS)
        text = all_lines(m2["blocks"]["vans"])
        assert "Held over 30 days unsold" in text and "BHD 25.000" in text and "9.200" in text, text
        m3 = blocks(books={VAN_A: {"colour": "red"}})
        assert "not worked out this time (the books-only figures could not be read: ValueError)" in all_lines(
            m3["blocks"]["vans"])


@test("vans: /yq-moves' books_only is read lazily, as on the month's end, held_days=30")
def _():
    with data_root():
        name = rb.MOVES_MODULE
        missing = object()
        saved = sys.modules.get(name, missing)            # /yq-moves itself may be there (release A): kept aside
        sys.modules[name] = None                          # not built: importing it raises ModuleNotFoundError
        try:
            raises(rb.NotYet, lambda: rb.Source().books_only(LAST))
            seen = []
            fake = types.ModuleType(name)
            fake.books_only = lambda as_of, held_days=99: seen.append((as_of, held_days)) or {VAN_A: {"units": 1,
                                                                                                       "bhd": 1}}
            sys.modules[name] = fake
            assert rb.Source().books_only(LAST) == {VAN_A: {"units": 1, "bhd": 1}}
            assert seen == [(LAST, 30)]
            del fake.books_only
            assert "no books_only yet" in str(raises(rb.NotYet, lambda: rb.Source().books_only(LAST)))
        finally:
            if saved is missing:
                sys.modules.pop(name, None)
            else:
                sys.modules[name] = saved


# ── 3: sales and receivables ──────────────────────────────────────────────────

def invoicing(**kw):
    args = dict(first=FIRST, reach=LAST, folder="Focus ERP 2026-10-02", daybook=DAYBOOK, register=REGISTER,
                portal=PORTAL, first_sale=SPAN[0], last_load=LOAD)
    args.update(kw)
    return rb.invoicing_part(**args)


@test("invoicing: the day book = the register = v_sales, invoice by invoice: PROVEN")
def _():
    p = invoicing()
    assert p["status"] == rb.PROVEN and p["basis"] == "files" and not p["differences"]
    assert p["daybook"] == {"invoices": 2, "lines": 3, "gross": Decimal("38.500")}
    assert p["register"] == {"invoices": 2, "lines": None, "gross": Decimal("38.500")}
    assert p["portal"] == p["daybook"]
    said = "\n".join(p["lines"])
    assert "2 invoices, 3 lines, BHD 38.500 VAT included" in said and "Focus ERP 2026-10-02" in said, said


@test("invoicing: DIFFERS for an invoice in one source only, other lines, or a register that says another figure")
def _():
    portal = dict(PORTAL, **{"SI : 103": (1, Decimal("4.000")), "SI : 102": (2, Decimal("5.500"))})
    register = (D, [r if r["invoice"] != "SI : 101" else ln("SI : 101", date(2026, 9, 3), 33.1) for r in REGISTER[1]])
    p = invoicing(portal=portal, register=register)
    assert p["status"] == rb.DIFFERS
    said = {d["invoice"]: d["said"] for d in p["differences"]}
    assert said["SI : 103"] == "in the portal, not in the day book; not in the register"
    assert "the day book has 1 line for BHD 5.500, the portal 2 lines for BHD 5.500" in said["SI : 102"]
    assert said["SI : 101"] == "the register says BHD 33.100, its day-book lines BHD 33.000"
    text = "\n".join(p["lines"])
    assert "3 invoices to check" in text and "SI : 103: in the portal" in text, text
    only_book = invoicing(portal={"SI : 101": PORTAL["SI : 101"]})
    assert only_book["differences"] == [{"invoice": "SI : 102", "said": "in the day book, not in the portal"}]
    only_reg = invoicing(register=(D, REGISTER[1] + [ln("SI : 104", date(2026, 9, 9), 1.0)]))
    assert {"invoice": "SI : 104", "said": "in the register only"} in only_reg["differences"]


@test("invoicing: stated as the load's own verify when the files do not reach the whole month, or are not there")
def _():
    late = (D, [r for r in DAYBOOK[1] if r["day"] >= date(2026, 9, 10)])
    p = invoicing(daybook=late)
    assert p["basis"] == "load_verify" and p["status"] is None
    text = "\n".join(p["lines"])
    assert "its day book it starts on 25 September 2026" in text and "each load's own verify" in text, text
    assert "021026_0800" in text and "The portal holds 2 invoices, 3 lines, BHD 38.500 VAT included" in text, text
    # starting late is fine when the portal's first sale is no earlier
    assert invoicing(daybook=late, first_sale=date(2026, 9, 25), register=(D, [])) ["basis"] == "load_verify"
    assert invoicing(daybook=(D, late[1]), register=(D, [r for r in REGISTER[1] if r["day"] >= date(2026, 9, 25)]),
                     first_sale=date(2026, 9, 25), portal={"SI : 102": PORTAL["SI : 102"]})["status"] == rb.PROVEN
    old = invoicing(daybook=(date(2026, 9, 20), DAYBOOK[1]))
    assert "it is as on 20 September 2026" in old["why"]
    assert "has no day book and no register" in invoicing(daybook=None, register=None)["why"]
    assert invoicing(folder=None)["why"] == "no loaded set of Focus files is on this PC"
    nothing = invoicing(folder=None, last_load={})
    assert "no load is on record here" in "\n".join(nothing["lines"])


@test("statement: the reps, their sales, tiers and kickback as build_rows drafts them; returns not deducted")
def _():
    st = rb.statement_part(statement())
    assert st["count"] == 2 and st["sales"] == Decimal("700.000") and st["kickback"] == Decimal("4.625")
    assert [(r["rep"], r["tier"]) for r in st["reps"]] == [(REP1, 1), (REP2, 2)]
    assert "returns not deducted" in st["lines"][0] and "2 reps with a target" in st["lines"][0]


@test("collection: opening + SI - RCTV - SRV - CN +/- other = closing per customer account and in total")
def _():
    part = rb.collection_part(export_of(), "the Ledger as on 2026-09-30", FIRST, LAST, [LAST], lambda d: dict(AR_EXACT),
                              AR_NAMES)
    by = {r["account"]: r for r in part["rows"]}
    assert set(by) == {CUST_A, CUST_B, CUST_C, CUST_D}          # the cash, sales and VAT accounts are no customers
    a, bb, c, dd = by[CUST_A], by[CUST_B], by[CUST_C], by[CUST_D]
    assert (a["opening"], a["si"], a["rctv"], a["closing"]) == (110, 220, 110, 220)
    assert (bb["opening"], bb["si"], bb["srv"], bb["closing"]) == (44, 55, 110, -11)
    assert (c["opening"], c["cn"], c["closing"]) == (88, 20, 68)
    assert (dd["opening"], dd["other"], dd["rctv"], dd["closing"]) == (33, 5, 38, 0)
    assert all(r["rolled"] == 0 for r in part["rows"])
    t = part["totals"]
    assert (t["opening"], t["si"], t["rctv"], t["srv"], t["cn"], t["other"], t["closing"]) == (275, 275, 148, 110, 20,
                                                                                              5, 277)
    assert t["opening"] + t["si"] - t["rctv"] - t["srv"] - t["cn"] + t["other"] == t["closing"]
    assert part["status"] == rb.PROVEN and part["agree"] == 4 and part["credits"] == 1 and not part["to_check"]
    text = "\n".join(part["lines"])
    assert ("opening BHD 275.000 + sales invoices BHD 275.000 - receipts BHD 148.000 - returns BHD 110.000 - credit "
            "notes BHD 20.000 + other BHD 5.000 = closing BHD 277.000") in text, text
    assert "as on 30 September 2026: 4 of the 4 accounts agree (1 of them a credit balance" in text, text


@test("collection: the closing against the ageing: a difference, one the ageing lacks, one the export lacks")
def _():
    ar = dict(AR_EXACT, **{CUST_C: Decimal("60"), CUST_Z: Decimal("12")})
    ar_sign = dict(ar, **{CUST_B: Decimal("-11")})                     # a signed ageing agrees as well
    part = rb.collection_part(export_of(), "x", FIRST, LAST, [LAST], lambda d: dict(ar), AR_NAMES)
    assert part["status"] == rb.DIFFERS and part["agree"] == 3
    said = {x["account"]: (x["ledger"], x["ageing"], x["said"]) for x in part["to_check"]}
    assert said == {CUST_C: (Decimal(68), Decimal(60), "differs"), CUST_Z: (None, Decimal(12), "not in the Ledger export")}
    text = "\n".join(part["lines"])
    assert "2 to check:" in text and "Customer C: the Ledger 68.000, the ageing 60.000 (differs)" in text, text
    assert rb.collection_part(export_of(), "x", FIRST, LAST, [LAST], lambda d: dict(ar_sign), AR_NAMES)["agree"] == 3
    missing_a = {CUST_B: Decimal(11), CUST_C: Decimal(68)}
    p2 = rb.collection_part(export_of(), "x", FIRST, LAST, [LAST], lambda d: dict(missing_a), AR_NAMES)
    assert [(x["account"], x["said"]) for x in p2["to_check"]] == [(CUST_A, "not in the ageing")]


@test("collection: the ageing nearest after the month end; the Ledger read to it, or stopping before it")
def _():
    oct1 = date(2026, 10, 1)
    ar = dict(AR_EXACT, **{CUST_A: Decimal("297")})                     # 220 + SI 110 of 1 Oct
    early = rb.collection_part(export_of(as_on=LAST), "x", FIRST, LAST, [date(2026, 9, 28), oct1], lambda d: dict(ar),
                               AR_NAMES)
    assert early["ar_day"] == oct1 and early["compared_at"] == LAST and early["status"] == rb.DIFFERS
    text = "\n".join(early["lines"])
    assert "as on 1 October 2026 (none is as on 30 September 2026)" in text
    assert "the Ledger stops at 30 September 2026, so what was entered after it shows as a difference" in text
    late = rb.collection_part(export_of(as_on=date(2026, 10, 3)), "x", FIRST, LAST, [oct1], lambda d: dict(ar), AR_NAMES)
    assert late["compared_at"] == oct1 and late["status"] == rb.PROVEN, late["to_check"]
    assert late["totals"]["closing"] == 277                             # the roll-forward stays the month's
    assert "the Ledger is read to 1 October 2026" in "\n".join(late["lines"])
    far = rb.collection_part(export_of(), "x", FIRST, LAST, [date(2026, 10, 11)], lambda d: dict(ar), AR_NAMES)
    assert far["status"] == rb.COLLECTION_NOT_PROVEN and far["ar_day"] is None
    assert "no receivables ageing as on 30 September 2026, nor in the 10 days after it, is loaded" in far["why"]


@test("collection: names that do not meet, a running balance that breaks, sales to accounts in no ageing")
def _():
    other = {"CUST-A (C001)": Decimal(220), "CUST-B (C002)": Decimal(11), CUST_C: Decimal(68)}
    p = rb.collection_part(export_of(), "x", FIRST, LAST, [LAST], lambda d: dict(other), AR_NAMES)
    assert p["status"] == rb.COLLECTION_NOT_PROVEN and "do not meet (1 of the 3 accounts with a balance found)" in p["why"]
    broken = rb.collection_part(export_of(broken=CUST_A), "x", FIRST, LAST, [LAST], lambda d: dict(AR_EXACT), AR_NAMES)
    assert broken["breaks"] == [CUST_A] and broken["status"] == rb.DIFFERS
    assert "does not follow its own rows for 1 account, to check: Customer A" in "\n".join(broken["lines"])
    shop = POSTINGS + [(date(2026, 9, 9), "SI : 120", [("Counter Sales Point", 22, 0), (SALES_AC, 0, 20),
                                                       (VAT_AC, 0, 2)])]
    q = rb.collection_part(export_of(shop), "x", FIRST, LAST, [LAST], lambda d: dict(AR_EXACT), AR_NAMES)
    assert q["outside"] == 1 and "1 account carrying a sales invoice is in no receivables ageing: left out" in "\n".join(
        q["lines"])


@test("returns: with the export, the month's returns ex-VAT per rep and the kickback tiers they would change")
def _():
    st = rb.statement_part(statement())
    r = rb.returns_from_export(export_of(), FIRST, LAST, st, rb.Source().tier)
    assert r["total"] == Decimal(100) and r["by_rep"] == {REP2: Decimal(100)} and r["unmatched"] == 0
    assert r["changed"] == [{"rep": REP2, "returns": Decimal(100), "net": Decimal(350), "tier": 2, "tier_net": 1,
                             "kickback": Decimal("3.375"), "kickback_net": Decimal("1.750")}]
    text = "\n".join(r["lines"])
    assert "BHD 100.000 in the month" in text and "change the kickback tier of 1 rep, to check" in text
    assert "Rep Two: tier 2 -> 1 (returns BHD 100.000; kickback BHD 3.375 -> 1.750)" in text, text
    same = rb.returns_from_export(export_of(), FIRST, LAST, rb.statement_part(statement("300.000")), rb.Source().tier)
    assert same["changed"] == [] and "They change no rep's kickback tier." in same["lines"]
    loose = POSTINGS + [(date(2026, 9, 26), "SRV : 8", [(RETURNS_AC, 10, 0), (VAT_AC, 1, 0), (CUST_C, 0, 11)])]
    u = rb.returns_from_export(export_of(loose), FIRST, LAST, st, rb.Source().tier)
    assert u["unmatched"] == Decimal(10) and "BHD 10.000 of it naming no rep with a target" in u["lines"][0]
    assert rb._rep_of("Rep Two - Acc WH", [REP1, REP2]) == REP2 and rb._rep_of("Rep Three", [REP1]) is None


@test("returns: without the export, an estimate from returned units; nothing judged on it")
def _():
    acc = {L(k) for k in NAMES}
    with data_root():                                 # the item bridge reads config: the fixture's, never a real one
        r = rb.returns_estimate(ROWS, acc, FIRST, LAST, PRICES)
        unpriced = rb.returns_estimate(ROWS, acc, FIRST, LAST, None)
    assert (r["units"], r["vouchers"], r["value"], r["unpriced"]) == (Decimal(7), 3, Decimal("13.500"), Decimal(2))
    said = r["lines"][0]
    assert said.startswith("Returns: net of returns is an estimate from returned units: 7 units came back on 3 Sales "
                           "Return vouchers"), said
    assert "about BHD 13.500 ex-VAT" in said and "2 units of items not sold in the month not priced" in said
    assert "no kickback tier is judged on an estimate" in said
    assert unpriced["value"] is None and "their price: not worked out this time" in unpriced["lines"][0]


@test("sales block: PROVEN with the export beside an ageing that agrees, the invoicing and no tier changed")
def _():
    with data_root():
        m = blocks(export=export_of(), ar={LAST: dict(AR_EXACT)}, statement=statement("300.000"))
        s = m["blocks"]["sales"]
        assert s["status"] == rb.PROVEN, (s["status"], s["lines"])
        assert s["collection"]["status"] == rb.PROVEN and s["invoicing"]["status"] == rb.PROVEN
        assert s["returns"]["basis"] == "export" and not s["returns"]["changed"]


@test("sales block: DIFFERS for the invoicing, for an account, for a tier the returns would change")
def _():
    with data_root():
        good = dict(export=export_of(), ar={LAST: dict(AR_EXACT)}, statement=statement("300.000"))
        assert blocks(**good)["blocks"]["sales"]["status"] == rb.PROVEN
        inv = blocks(**dict(good, portal=dict(PORTAL, **{"SI : 103": (1, Decimal(1))})))["blocks"]["sales"]
        acct = blocks(**dict(good, ar={LAST: dict(AR_EXACT, **{CUST_C: Decimal(60)})}))["blocks"]["sales"]
        tier = blocks(**dict(good, statement=statement()))["blocks"]["sales"]
        for s, why in ((inv, "invoicing"), (acct, "collection"), (tier, "returns")):
            assert s["status"] == rb.DIFFERS, (why, s["status"])
        assert inv["invoicing"]["status"] == rb.DIFFERS and acct["collection"]["status"] == rb.DIFFERS
        assert tier["returns"]["changed"] and tier["collection"]["status"] == rb.PROVEN


@test("sales block: COLLECTION NOT PROVEN without the export, with one that cannot be read, with no ageing")
def _():
    with data_root():
        s = blocks()["blocks"]["sales"]
        assert s["status"] == rb.COLLECTION_NOT_PROVEN and s["collection"]["why"] == NO_EXPORT
        text = all_lines(s)
        assert "Collection: COLLECTION NOT PROVEN: " + NO_EXPORT in text
        assert "Put Focus's Ledger report (all accounts, as on 30 September 2026) in the next drop folder" in text
        assert "net of returns is an estimate from returned units" in text and s["returns"]["basis"] == "estimate"
        bad = blocks(export=le.LedgerExportError("the Ledger export X.xlsx has no \"As on date\" in its title block"))
        sb = bad["blocks"]["sales"]
        assert sb["status"] == rb.COLLECTION_NOT_PROVEN and "could not be read: the Ledger export X.xlsx" in sb[
            "collection"]["why"]
        odd = blocks(export=OSError("held by another program"))["blocks"]["sales"]
        assert "could not be read (OSError)" in odd["collection"]["why"]
        no_ar = blocks(export=export_of(), ar={}, statement=statement("300.000"))["blocks"]["sales"]
        assert no_ar["status"] == rb.COLLECTION_NOT_PROVEN and "no receivables ageing" in no_ar["collection"]["why"]
        # the tiers are still judged from the export's own returns: a change is still to check
        changed = blocks(export=export_of(), ar={})["blocks"]["sales"]
        assert changed["returns"]["changed"] and changed["status"] == rb.DIFFERS
        assert changed["collection"]["status"] == rb.COLLECTION_NOT_PROVEN


@test("sales block: a part that cannot be read stands aside; the rest stands")
def _():
    with data_root():
        m = blocks(fail={"portal_invoices": RuntimeError("down")})
        s = m["blocks"]["sales"]
        assert s["invoicing"] is None and "v_sales could not be read: RuntimeError" in s["notes"]["invoicing"]
        assert s["status"] == rb.COLLECTION_NOT_PROVEN and s["statement"] is not None
        assert any("v_sales could not be worked out" in w for w in m["warnings"])
        m2 = blocks(fail={"statement_rows": RuntimeError("down")}, export=export_of(), ar={LAST: dict(AR_EXACT)})
        assert m2["statuses"]["sales"] is None and "Statement: not worked out this time" in all_lines(m2["blocks"]["sales"])
        unread = blocks(daybook=OSError("locked"))["blocks"]["sales"]["invoicing"]
        assert unread["basis"] == "load_verify" and "could not be read (OSError)" in unread["why"]


@test("collection through a real Ledger workbook in a drop folder: find and read, the narration never comes through")
def _():
    with data_root() as root:
        write_ledger(root / "Uplaods" / "300926" / "Ledger_test.xlsx")
        m = blocks(export="real", ar={LAST: dict(AR_EXACT)}, statement=statement("300.000"))
        s = m["blocks"]["sales"]
        assert s["collection"]["file"] == "Ledger_test.xlsx" and s["collection"]["status"] == rb.PROVEN, s["lines"]
        assert s["collection"]["totals"]["closing"] == 277 and s["status"] == rb.PROVEN
        blob = json.dumps(m, default=str)
        assert "paid by" not in blob and not NINE.search(blob)


# ── 4: purchases ──────────────────────────────────────────────────────────────

def purchases(**kw) -> dict:
    with data_root():
        return blocks(**kw)["blocks"]["purchases"]


@test("purchases: the order by its number, the papers, cost lines n of m, variants, the order's notes")
def _():
    p = purchases()
    rows = {r["receipt"]: r for r in p["rows"]}
    assert list(rows) == ["MRN:T-26-09-1", "MRN:T-26-09-2"]                       # August's is outside the month
    one = rows["MRN:T-26-09-1"]
    assert (one["po"], one["po_basis"], one["xml"], one["invoice"], one["invoice_basis"]) == (
        "T-26-09-1", "number", True, True, "vault")
    assert (one["cost_lines"], one["lines"], one["status"]) == (4, 4, rb.PROVEN)
    assert one["variants"] == ["DS1"] and not one["actions"]
    notes = "; ".join(one["notes"])
    assert "1 model code with several variants (DS1): variant not provable from the PO" in notes
    assert "margin on arrival below 20%: DS1" in notes


@test("purchases: the order by its codes (Jaccard one half, within 200 days); a missing paper is an action")
def _():
    p = purchases()
    two = next(r for r in p["rows"] if r["receipt"] == "MRN:T-26-09-2")
    assert (two["po"], two["po_basis"], two["overlap"]) == ("T-26-07-3", "codes", "1/2")    # not T-25-12-1: too old
    assert two["xml"] is False and two["invoice"] and two["invoice_basis"] == "matched" and two["invoices"] == ["INV-9"]
    assert two["status"] == rb.PAPERS_MISSING and two["actions"] == [rb.PAPER_ACTIONS["xml"]]
    assert not two["differs"]
    assert "its order T-26-07-3 is matched by its model codes, not by its number: the two share 1 of their 2 codes" in \
        two["notes"]
    assert p["status"] == rb.PAPERS_MISSING and p["totals"]["papers_missing"] == 1 and p["totals"]["elsewhere"] == 1
    text = all_lines(p)
    assert "To do, MRN:T-26-09-2: upload its MRN XML on the Orders page." in text, text
    assert "1 supplier receipt went into another store" in text and "DIFFERS" not in text
    # under half shared, or too long before: no order
    leads = {"SE5"}
    pos = rb._po_table(rb_rows(("T-1", "2026-07-15", "SE5", "ZZ8", "ZZ9")))
    assert rb.match_po("T-9", date(2026, 9, 20), leads, pos) == (None, None, None)
    pos = rb._po_table(rb_rows(("T-2", "2026-03-01", "SE5")))           # 203 days before
    assert rb.match_po("T-9", date(2026, 9, 20), leads, pos)[0] is None
    pos = rb._po_table(rb_rows(("T-3", "2026-03-04", "SE5")))           # 200 days: still in
    assert rb.match_po("T-9", date(2026, 9, 20), leads, pos)[0] == "T-3"
    later = rb._po_table(rb_rows(("T-4", "2026-09-21", "SE5")))         # dated after the receipt
    assert rb.match_po("T-9", date(2026, 9, 20), leads, later)[0] is None


def rb_rows(*orders):
    return [r for no, day, *codes in orders for r in po(no, day, *codes)]


@test("purchases: every paper missing, each one an action; the vault's PO with no lines; cost lines from a workbook")
def _():
    p = purchases(pos=[], order_files=[{"po_no": "T-26-09-2", "kind": "po"}], invoices=[], details={},
                  mrn=[{"doc_no": "T-26-09-2", "code": "SE5"}])
    two = next(r for r in p["rows"] if r["receipt"] == "MRN:T-26-09-2")
    assert two["po"] is None and two["po_file_only"]
    assert two["actions"][0].startswith("its purchase order is in the vault but its lines are not loaded")
    assert two["actions"][1:] == [rb.PAPER_ACTIONS["xml"], rb.PAPER_ACTIONS["invoice"]]
    assert "its landed costs are loaded without the XML on file (from a workbook)" in two["notes"]
    one = next(r for r in p["rows"] if r["receipt"] == "MRN:T-26-09-1")
    assert one["actions"] == [rb.PAPER_ACTIONS[k] for k in ("po", "xml", "invoice")] and one["cost_lines"] == 0
    assert p["status"] == rb.PAPERS_MISSING and p["totals"]["po"] == 0
    partial = purchases(mrn=MRN_ROWS[:3])
    one = next(r for r in partial["rows"] if r["receipt"] == "MRN:T-26-09-1")
    assert one["cost_lines"] == 3 and one["status"] == rb.PAPERS_MISSING
    assert one["actions"] == ["landed costs are loaded for 3 of its 4 lines: load the rest from its MRN XML"]


@test("purchases: DIFFERS (to check) when the order's flags say so; PROVEN with every paper and nothing to check")
def _():
    files = FILE_ROWS + [{"po_no": "T-26-07-3", "kind": "mrn"}]
    mrn = MRN_ROWS + [{"doc_no": "T-26-07-3", "code": "SE5"}]
    received = dict(DETAILS, **{"T-26-07-3": detail(("SE5", ()), ("ZZ9", ("not_received",)))})
    ok = purchases(order_files=files, mrn=mrn, details=dict(DETAILS, **{"T-26-07-3": detail(("SE5", ()), ("ZZ9", ()))}))
    assert ok["status"] == rb.PROVEN and [r["status"] for r in ok["rows"]] == [rb.PROVEN, rb.PROVEN]
    gap = purchases(order_files=files, mrn=mrn, details=received)
    two = next(r for r in gap["rows"] if r["receipt"] == "MRN:T-26-09-2")
    assert two["status"] == rb.DIFFERS and two["flags"] == {"not_received": ["ZZ9"]} and gap["status"] == rb.DIFFERS
    # the XML filed under the receipt's own number, its order showing nothing received: attach it to the order
    elsewhere = purchases(order_files=FILE_ROWS + [{"po_no": "T-26-09-2", "kind": "mrn"}],
                          mrn=MRN_ROWS + [{"doc_no": "T-26-09-2", "code": "SE5"}])
    two = next(r for r in elsewhere["rows"] if r["receipt"] == "MRN:T-26-09-2")
    assert two["xml"] and two["cost_lines"] == 1 and two["status"] == rb.PAPERS_MISSING
    assert two["actions"] == ["its order T-26-07-3 shows nothing received on the Orders page: attach the MRN XML to "
                              "that order"], two["actions"]
    details = dict(DETAILS, **{"T-26-09-1": detail(("CA1", ("over",)), ("HG3", ("short", "cost_overrun")),
                                                   ("DS1", ()))})
    bad = purchases(order_files=files, mrn=mrn, details=dict(details, **{"T-26-07-3": detail(("SE5", ()))}))
    one = next(r for r in bad["rows"] if r["receipt"] == "MRN:T-26-09-1")
    assert one["status"] == rb.DIFFERS and one["flags"] == {"over": ["CA1"], "short": ["HG3"], "cost_overrun": ["HG3"]}
    assert bad["status"] == rb.DIFFERS and bad["totals"]["differs"] == 1
    text = all_lines(bad)
    assert ("To check, MRN:T-26-09-1 against its order T-26-09-1: received short of the order: HG3; received over the "
            "order: CA1; supplier price more than 5% above the order: HG3.") in text, text
    # a paper missing beside a difference: the row asks for the paper first, the block still reads DIFFERS
    both = purchases(details=details, mrn=MRN_ROWS[:3])
    one = next(r for r in both["rows"] if r["receipt"] == "MRN:T-26-09-1")
    assert one["status"] == rb.PAPERS_MISSING and one["differs"] and both["status"] == rb.DIFFERS


@test("purchases: PROVEN with no receipt in the month (nothing is due); a read that fails stands the block aside")
def _():
    with data_root():
        quiet = [r for r in ROWS if r["vtype"] != ledger.SUPPLIER_RECEIPT or r["date"].month != 9]
        p = blocks(rows=quiet)["blocks"]["purchases"]
        assert p["status"] == rb.PROVEN and p["rows"] == []
        assert "No supplier receipt into the main store from 1 September 2026 to 30 September 2026: no paper is due." in \
            all_lines(p)
        m = blocks(fail={"po_lines": RuntimeError("down")})
        assert m["statuses"]["purchases"] is None and "not worked out this time" in all_lines(m["blocks"]["purchases"])
        m2 = blocks(fail={"order_detail": RuntimeError("down"), "supplier_invoices": RuntimeError("down")})
        p2 = m2["blocks"]["purchases"]
        assert p2["status"] is not None and "flags are not worked out this time" in all_lines(p2)
        assert "the supplier invoices are not worked out this time" in all_lines(p2)


@test("purchases: without the item codes the cost lines are not counted (never read as missing); the order still found")
def _():
    with data_root():
        m = blocks(fail={"accessories": RuntimeError("down")})
        p = m["blocks"]["purchases"]
        one = next(r for r in p["rows"] if r["receipt"] == "MRN:T-26-09-1")
        assert one["cost_lines"] is None and one["po"] == "T-26-09-1" and one["status"] == rb.PROVEN
        assert not any("landed costs" in a for r in p["rows"] for a in r["actions"])
        assert p["totals"]["cost_lines"] is None
        text = all_lines(p)
        assert "the item codes are not worked out this time: the cost lines are not counted" in text, text
        _b, ws, _r, errors = sheet_of(m)
        assert not errors and "Cost lines: not worked out this time (the item codes could not be read)." in \
            sheets_recon.texts_of(ws)


@test("purchases(first, last): any window, rows with their statuses, as the weekly workbook asks for them")
def _():
    with data_root():
        sig = inspect.signature(rb.purchases)
        assert list(sig.parameters)[:2] == ["first", "last"] and sig.parameters["source"].default is None
        got = rb.purchases(date(2026, 8, 1), date(2026, 9, 30), source=Fake())
        assert [r["receipt"] for r in got["rows"]] == ["MRN:T-26-08-4", "MRN:T-26-09-1", "MRN:T-26-09-2"]
        assert all(r["status"] in rb.STATUSES for r in got["rows"])
        missing = sum(1 for r in got["rows"] if "PAPERS MISSING" in r["status"])
        assert missing == 2 and got["totals"]["papers_missing"] == 2
        assert rb.purchases("2026-09-01", "2026-09-30", source=Fake())["totals"]["receipts"] == 2
        raises(rb.ReconInput, lambda: rb.purchases(LAST, FIRST, source=Fake()))


# ── the month and the command ─────────────────────────────────────────────────

@test("the month: four blocks, each status one of the six; a month the data has not reached is one sentence")
def _():
    with data_root():
        m = blocks(results={MAIN: main_ok(), **van_counts()}, books=BOOKS)
        assert m["statuses"] == {"main": rb.PROVEN, "vans": rb.DIFFERS, "sales": rb.COLLECTION_NOT_PROVEN,
                                 "purchases": rb.PAPERS_MISSING}, m["statuses"]
        assert (m["data_date"], m["books_as_of"], m["sales_to"], m["whole"]) == (D, LAST, LAST, True)
        assert set(rb.STATUSES) == {"PROVEN", "DIFFERS", "NOT COUNTED", "BOOKS ONLY", "COLLECTION NOT PROVEN",
                                    "PAPERS MISSING"}
        e = raises(rb.ReconInput, lambda: rb.month_blocks("2026-11", Fake()))
        assert str(e) == "No data reaches November 2026 yet: the data runs to 2 October 2026."
        e = raises(rb.ReconInput, lambda: rb.month_blocks("2026-07", Fake()))
        assert "The stock ledger starts on 10 August 2026" in str(e)
        assert "is written YYYY-MM" in str(raises(rb.ReconInput, lambda: rb.month_blocks("2026-9", Fake())))
        part = rb.month_blocks("2026-10", Fake())
        assert not part["whole"] and part["books_as_of"] == D and part["sales_to"] == date(2026, 10, 1)


@test("the month: the counts are read for the month; a count-results read that fails stands both count blocks aside")
def _():
    with data_root():
        f = Fake(results={MAIN: main_ok()})
        rb.month_blocks(MONTH, f)
        assert ("count_results", MONTH) in f.calls and ("books_only", LAST) in f.calls
        m = blocks(fail={"count_results": OSError("locked")})
        assert m["statuses"]["main"] is None and m["statuses"]["vans"] is None
        assert m["statuses"]["sales"] is not None and m["statuses"]["purchases"] is not None


@test("the month: the real count results on disk (counts.month_results), the newest of each store")
def _():
    with data_root():
        for r in (result(MAIN, "main", {"A": 260}, day=date(2026, 9, 25)), main_ok(), result(VAN_A, "van", {"A": 23})):
            paths.write_text(counts.result_file(r["month"], r["prep_id"], r["slug"] + r["count_date"][-2:]),
                             json.dumps(r))
        got = rb.Source().count_results(MONTH)
        assert set(got) == {MAIN, VAN_A} and got[MAIN]["count_date"] == "2026-09-29"
        assert rb.main_block(got, rules(), MONTH)["status"] == rb.PROVEN


def run(**kw) -> dict:
    return rb.run(Args(MONTH), kw.pop("month", MONTH), source=Fake(**kw))


@test("the command: the screen, the files under exports/ops/<D>/ only, the json, the wording, no long digit run")
def _():
    with data_root() as root:
        res = run(results={MAIN: main_differs(), **van_counts()}, books=BOOKS, export=export_of(),
                  ar={LAST: dict(AR_EXACT, **{"Customer 123456789": Decimal(5)})})
        assert res["exit"] == 0 and res["ok"] and res["data_date"] == D.isoformat()
        text = res["text"]
        head = text.splitlines()
        assert head[0] == ("RECONCILIATIONS  September 2026 (data to Fri 02 Oct; the books as at the end of 30 September "
                           "2026)"), head[0]
        assert head[1].split() == ["MAIN", "STORE", "DIFFERS", "(to", "check)"]
        for n, title in enumerate(("1. MAIN STORE: DIFFERS (to check)", "2. VANS AND OUTLETS: DIFFERS (to check)",
                                   "3. SALES AND RECEIVABLES: DIFFERS (to check)", "4. PURCHASES: PAPERS MISSING")):
            assert title in head, title
        assert text.rstrip().endswith("it agrees with what the portal holds.")
        files = [Path(f) for f in res["files"]]
        assert [f.name for f in files] == ["recon-2026-09.txt", "recon-2026-09.json"]
        assert all(f.parent == root / "exports" / "ops" / D.isoformat() for f in files)
        assert files[0].read_text(encoding="utf-8") == text + "\n"
        facts = json.loads(files[1].read_text(encoding="utf-8"))
        assert facts["statuses"]["sales"] == rb.DIFFERS and facts["month"] == MONTH
        for blob in (text, files[1].read_text(encoding="utf-8")):
            assert not NINE.search(blob) and not LONG.search(blob), "an ID-shaped run"
            assert COUNTER not in blob and CHECKER not in blob
        assert "Customer ******789" in text and "Customer ******789" in facts["blocks"]["sales"]["collection"][
            "to_check"][0]["account"]
        fmt.careful(text)
        assert res["flags"]["statuses"] == facts["statuses"]
        made = sorted(p.relative_to(root).as_posix() for p in root.rglob("*") if p.is_file())
        assert made == ["business_data/ops/config.json", f"exports/ops/{D.isoformat()}/recon-2026-09.json",
                        f"exports/ops/{D.isoformat()}/recon-2026-09.txt"], made


@test("the command: a word from the data that the guard refuses is blanked; a month not reached is exit 1")
def _():
    with data_root():
        ar = dict(AR_EXACT, **{"Theft Guard Traders": Decimal(5)})
        res = run(export=export_of(), ar={LAST: ar}, names=AR_NAMES | {le.fold("Theft Guard Traders")})
        assert "Theft" not in res["text"] and "----- Guard Traders" in res["text"]
        fmt.careful(res["text"])
        late = run(month="2026-12")
        assert late["exit"] == 1 and not late["ok"] and late["text"].startswith("No data reaches December 2026 yet")


@test("the command: no month is the last whole month before the data date; a block not worked out is exit 2")
def _():
    with data_root():
        res = rb.run(Args(), None, source=Fake())
        assert res["flags"]["month"] == "2026-09" and res["exit"] == 0
        assert rb._default_month(date(2026, 1, 3)) == "2025-12"
        broken = run(fail={"po_lines": RuntimeError("down")})
        assert broken["exit"] == 2 and not broken["ok"] and "4. PURCHASES: NOT WORKED OUT" in broken["text"]
        assert rb.run(Args(), None, source=Fake(span=(None, None)))["exit"] == 1


@test("recon hands `recon YYYY-MM` to recon_blocks.run(args, month); python -m scripts.ops recon prints it")
def _():
    with data_root() as root:
        saved = rb.Source
        rb.Source = lambda prefer="auto": Fake()
        try:
            res = recon.run(Args(MONTH))
            assert res["exit"] == 0 and res["text"].startswith("RECONCILIATIONS  September 2026"), res["text"][:80]
            from contextlib import redirect_stdout
            from scripts.ops import __main__ as cli
            buf = io.StringIO()
            with redirect_stdout(buf):
                code = cli.main(["recon", MONTH, "--data-root", str(root)])
            assert code == 0 and "4. PURCHASES: PAPERS MISSING" in buf.getvalue()
        finally:
            rb.Source = saved


@test("the Source's own pieces: the accessories through /yq-stock's seam, the day book through the loader's parser")
def _():
    with data_root() as root:
        saved = {n: getattr(stock, n) for n in ("cost_reads", "accessory_table")}
        stock.cost_reads = lambda: {"price_book": [], "snapshot": [], "mrn": {}, "sheet": {}, "unread": ["a read (X)"]}
        stock.accessory_table = lambda rows, book, snap: ({L("A"): types.SimpleNamespace(code="CA1"),
                                                           L("B"): types.SimpleNamespace(code=None)}, None, [])
        try:
            src = rb.Source()
            assert src.accessories(ROWS) == ({L("A"): "CA1", L("B"): None}, ["a read (X)"])
            assert src.accessories([]) is src.accessories(ROWS)               # read once per Source
        finally:
            for n, fn in saved.items():
                setattr(stock, n, fn)
        folder = root / "business_data" / "Focus ERP 2026-10-02"
        folder.mkdir(parents=True)
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        for _ in range(3):
            ws.append([])
        ws.append([None, None, None, "[As on date 02/10/2026]"])
        ws.append([])
        ws.append(["Date", "Voucher", "Customer Account", "Item", "Quantity", "Rate", "Gross", "Discount", "Taxable",
                   "VAT Amount", "Total Amount", "Narration", "Warehouse Name"])
        ws.append([datetime(2026, 9, 3), "SI : 101", "Customer A", NAMES["A"], 2, 5.5, 11, 0, 10, 1, None,
                   "id " + "98765" + "4321", VAN_A])
        wb.save(folder / "Sales_day_book_test.xlsx")
        (folder / "Summary_sales_register_test.xlsx").write_bytes(b"")
        (folder / "~$Sales_day_book_test.xlsx").write_bytes(b"")
        name, book_file, register_file = src.sales_files()
        assert name == folder.name and book_file.name == "Sales_day_book_test.xlsx"
        assert register_file.name == "Summary_sales_register_test.xlsx"
        as_on, rows = src.read_sales_file(book_file, "order_lines")
        assert as_on == D and rows == [{"invoice": "SI : 101", "day": date(2026, 9, 3), "gross": 11}]


# ── the sheet ─────────────────────────────────────────────────────────────────

def sheet_of(m: dict):
    book = bk.Book()
    ws = book.sheet(sheets_recon.SHEET_NAME)
    refs = sheets_recon.recon_sheet(ws, m)
    errors = book.evaluate()
    return book, ws, refs, errors


FORBIDDEN = ([CUST_A, CUST_B, CUST_C, CUST_D, CUST_Z, REP1, REP2, COUNTER, CHECKER, " - Acc WH", "SI : 1", "MRN:",
              "T-26-09", "T-26-07", "INV-9", "SRV:"] + list(NAMES.values()) + list(CODES.values()) + ["ITEM-"])
VOUCHER = re.compile(r"\b[A-Za-z]{1,6}(?:-[A-Za-z]{1,6})*-\d{2}-\d{2}-\d+\b")    # the weekly workbook's own rule


@test("the sheet: totals only, no customer, item, voucher, rep or counter name; stores as the workbook names salesmen")
def _():
    with data_root() as root:
        m = blocks(results={MAIN: main_differs(), **van_counts()}, books=BOOKS, export=export_of(),
                   ar={LAST: dict(AR_EXACT, **{CUST_C: Decimal(60)})})
        book, ws, refs, errors = sheet_of(m)
        assert not errors, errors
        texts = sheets_recon.texts_of(ws)
        blob = "\n".join(texts)
        for word in FORBIDDEN:
            assert word not in blob, word
        assert "Van A" in texts and "Van B" in texts and OUTLET in texts
        for word in ("DIFFERS (to check)", "PAPERS MISSING", "Reconciliation", "Without its second side"):
            assert word in blob, word
        assert not VOUCHER.search(blob), VOUCHER.search(blob)
        nums = sheets_recon.numbers_of(ws)
        assert all(isinstance(x, (int, float)) and not isinstance(x, bool) for x in nums)
        assert 277 in nums and 0.8 in nums and 38.5 in nums
        assert not any(c.kind == "s" and re.fullmatch(r"-?[0-9.,]+", c.value or "") for c in ws.cells.values())
        from scripts.ops.workbook import formula as fx
        check = ws.cells[fx.parse_cell(refs["roll_check"])]
        assert check.kind == "f" and check.value == 0, check.value   # the roll-forward's check reads nought
        sums = {c.formula[4]: c.value for c in ws.cells.values() if c.kind == "f" and c.formula.startswith("SUM(")}
        # lines, lines counted, units short (Van B's one), over, BHD short (2.500), over; the books-only columns
        # are blank for a store with a signed count
        assert sums == {"D": 3, "E": 3, "F": 1, "G": 0, "H": 2.5, "I": 0, "J": 0, "K": 0, "L": 0}, sums
        path = book.save(root / "exports" / "reports" / "monthly" / "recon-sheet-test.xlsx")
        import openpyxl
        wb = openpyxl.load_workbook(path)
        cells = [c.value for row in wb[sheets_recon.SHEET_NAME].iter_rows() for c in row if c.value is not None]
        text = "\n".join(str(v) for v in cells)
        for word in FORBIDDEN:
            assert word not in text, word
        assert not NINE.search(text) and not LONG.search(text) and not VOUCHER.search(text)


@test("the sheet: every status and every part, the vans' total, a block not worked out, a month with nothing")
def _():
    with data_root():
        m = blocks(results={MAIN: main_ok(), **van_counts(**{VAN_B: result(VAN_B, "van", {"B": 10}), OUTLET: None})},
                   books=BOOKS)
        assert m["statuses"] == {"main": rb.PROVEN, "vans": rb.BOOKS_ONLY, "sales": rb.COLLECTION_NOT_PROVEN,
                                 "purchases": rb.PAPERS_MISSING}
        _book, ws, refs, errors = sheet_of(m)
        assert not errors and "statuses" in refs and "roll_check" not in refs
        texts = sheets_recon.texts_of(ws)
        status_cells = [ws.cells[(r, 2)].value for r in range(6, 10)]
        assert status_cells == ["PROVEN", "BOOKS ONLY", "COLLECTION NOT PROVEN", "PAPERS MISSING"], status_cells
        assert texts.count("PROVEN") == 3 and "no count" in texts           # the main store, Van A, Van B; the outlet
        assert any(t.startswith("Collection not proven: no drop folder") for t in texts)
        assert "Net of returns is an estimate from returned units: the Ledger export holds their values." in texts
        sums = {c.formula[4]: c.value for c in ws.cells.values() if c.kind == "f" and c.formula.startswith("SUM(")}
        assert sums == {"D": 2, "E": 2, "F": 0, "G": 0, "H": 0, "I": 0, "J": 7, "K": 14, "L": 1}, sums
        inner = bk.Book().sheet(sheets_recon.SHEET_NAME)                  # handed the four blocks alone
        sheets_recon.recon_sheet(inner, m["blocks"])
        assert [inner.cells[(r, 2)].value for r in range(6, 10)] == status_cells
        nothing = blocks(rows=[r for r in ROWS if r["vtype"] != ledger.SUPPLIER_RECEIPT or r["date"].month != 9],
                         fail={"count_results": OSError("x")})
        _b2, ws2, _r2, errors2 = sheet_of(nothing)
        t2 = sheets_recon.texts_of(ws2)
        assert not errors2 and t2.count("Not worked out this time.") == 2 and "NOT WORKED OUT" in t2
        assert "No supplier receipt into the main store in the month: no paper is due." in t2


def main() -> int:
    passed = failed = 0
    for name, fn in TESTS:
        try:
            fn()
            print(f"PASS  {name}")
            passed += 1
        except Exception:  # noqa: BLE001
            print(f"FAIL  {name}")
            traceback.print_exc()
            failed += 1
    print(f"\n{passed} passed, {failed} failed")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
