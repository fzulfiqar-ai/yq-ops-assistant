"""Market Intel v2, Phase 2: the monthly supplier workbook (app/supplier_reports.py and its routes).

    python -m tests.test_r8m_supplier_reports

Same lightweight runner as tests/test_r7b_market_intel.py, on ITS in-memory PostgREST + Storage
stand-in (imported, not copied: it pages 1,000 rows at a time and honours .range()'s offset), wrapped
here to refuse what the real schema refuses: a column a table does not have, in a select list, a
filter, an order or a written row (the column lists of scripts/r7b_market_intel_migration.sql and
scripts/r8m_market_quality_migration.sql: supplier_action_cases has no id), and a second approved
report of the same supplier, month and edition (the partial unique index). SUPABASE_URL, SUPABASE_KEY
and DATABASE_URL are pinned to ci.invalid BEFORE any import, so nothing here can reach a database, a
bucket or the network. Every supplier, brand, shop, rep, product, phone and amount is invented.

Covered:
  * the month: a Bahrain calendar month (a sighting at 01:30 on the 1st is that month's), YYYY-MM
    parsing, a month that has not started refused, approval only from the month's end plus 7 days
    (a draft before that warns); paged reads past 1,000 rows, and a read too long to finish refused;
  * which sightings belong to a supplier: every link (case, case brand, sighting, sighting brand; merges
    followed; a brand's CURRENT supplier) must agree. A row whose links disagree is left out of every
    draft with a warning and blocks the approval; never another supplier's row; not substantiated left
    out; a merged supplier's own report refused;
  * counts by issue type, new vs recurring (>= 2 cases or >= 3 shops in the 90 days that end with the
    month; a case 100 days old does not count), the insufficient-evidence rule, open actions (a
    cancelled one never in a supplier copy), a SKU the catalog does not confirm printed as unverified;
  * the narrative: every sentence about a finding cites real [MI-...] ids of this supplier, every
    action line its [CA-...]; an edited one may cite only listed rows, each finding sentence must cite,
    and it is approved only while the data is unchanged; a template one is rewritten at approval;
  * edits: rows included / excluded, a row or a photo withheld (honoured, frozen at approval);
  * approval by management or an admin only (a reviewer 403), superseding, the frozen manifest: the
    download prints only what was approved (a later sighting, photo, answer or verification never
    reaches it) and is refused when a frozen row changed supplier or was withdrawn, or a frozen photo
    was flagged with people, re-classed, deleted or changed in storage; a racing approval is a 409 and
    a failed one puts the superseded report back;
  * download: approved only; {at, by, sha256} logged on the row and in audit_log; the same manifest
    gives the same bytes; a download that cannot be recorded is never served;
  * the workbook opened with openpyxl: sheet names, header rows, embedded 640 px JPEGs and thumbnails,
    "photo not available", and a privacy scan of every cell of the supplier edition, fed with phones
    written with dashes, bare 3-decimal prices, price words, a shop or supplier name that contains the
    supplier's own name, other suppliers' brands and aliases;
  * roles on every route; management's preview of a hand-edited narrative is masked;
  * before the r8m migration every route answers 503 with the hint, and so does a table missing
    half-way; a column we do not have is a 500 that names it, never "not set up".
"""
from __future__ import annotations

import os

os.environ["SUPABASE_URL"] = "https://ci.invalid"
os.environ["SUPABASE_KEY"] = "ci-key"
os.environ["DATABASE_URL"] = "postgresql://ci:ci@ci.invalid:5432/ci"

import tests.test_r7b_market_intel as R  # noqa: E402  first: stdout, sys.path and the fake are set up there

import hashlib  # noqa: E402
import io  # noqa: E402
import json  # noqa: E402
import re  # noqa: E402
import traceback  # noqa: E402
import uuid  # noqa: E402
from datetime import datetime, timezone  # noqa: E402

ROOT = R.ROOT
TESTS: list[tuple[str, object]] = []


def test(name):
    def deco(fn):
        TESTS.append((name, fn))
        return fn
    return deco


NOW = datetime(2026, 10, 14, 9, 0, tzinfo=timezone.utc)
SEP = "2026-09"


# ── the r7b stand-in, refusing what the real schema refuses ────────────────────

_R7B_OBS = {"id", "item_id", "client_uuid", "source", "signal", "kind", "salesman_id", "created_by", "shop_customer_id",
            "shop_name", "area", "title", "brand", "competitor", "category", "note", "price_bhd", "demand_level",
            "demand_qty", "barcode_raw", "yq_item_code", "result", "near_duplicate_of", "legacy_ref", "meta",
            "observed_at", "created_at"}
_R8M_OBS = {"case_id", "issue_type_id", "report_basis", "qty_affected", "batch_ref", "packaging_version", "safety_flag",
            "client_observed_at", "brand_id", "supplier_id", "brand_text", "supplier_text"}
_AUDIT = {"created_at", "created_by", "updated_at", "updated_by"}
SCHEMA = {
    "suppliers": {"id", "name", "code", "active", "merged_into"} | _AUDIT,
    "brands": {"id", "name", "supplier_id", "aliases", "renamed_from", "active", "merged_into"} | _AUDIT,
    "market_issue_types": {"id", "kind", "code", "label_en", "label_ar", "safety_default", "fields", "active", "sort"} | _AUDIT,
    "market_cases": {"id", "yq_item_code", "brand_id", "supplier_id", "issue_type_id", "batch_ref", "packaging_version",
                     "status", "evidence_level", "safety", "severity", "assignee", "due_date", "verification_method",
                     "verification_reason", "title", "first_seen", "last_seen", "obs_count", "shop_count",
                     "merged_into", "status_reason", "changed_at", "changed_by"} | _AUDIT,
    "market_case_events": {"id", "case_id", "at", "actor", "event", "from_status", "to_status", "detail"},
    "supplier_actions": {"id", "supplier_id", "request", "owner", "due_date", "status", "supplier_response",
                         "resolution_check"} | _AUDIT,
    "supplier_action_cases": {"action_id", "case_id", "created_at", "created_by"},      # no id: (action_id, case_id)
    "supplier_reports": {"id", "supplier_id", "month", "edition", "status", "narrative", "included_observation_ids",
                         "counts", "approved_by", "approved_at", "file_sha256", "downloads"} | _AUDIT,
    "market_observations": _R7B_OBS | _R8M_OBS,
    "market_photos": {"id", "observation_id", "bucket", "path", "position", "width", "height", "bytes", "phash",
                      "has_people_flag", "created_at", "role", "sha256", "thumb_path", "caption", "blurred_path"},
    "salesmen": {"id", "name", "phone", "email", "whatsapp", "user_email", "focus_name", "referral_code", "is_active",
                 "sort_order", "notify_email", "notify_whatsapp", "created_at", "updated_at", "public_profile",
                 "public_whatsapp", "title", "photo_url", "territory", "focus_aliases"},
    "shop_customers": {"id", "phone", "name", "shop", "area", "email", "salesman_id", "assigned_by", "assigned_at",
                       "focus_salesman_name", "sticky_salesman_id", "first_ref", "first_order_at", "last_order_at",
                       "orders_count", "total_bhd", "focus_customer_id", "device_ids", "verified_at", "opted_out",
                       "notes", "created_at", "updated_at"},
    "customers": {"id", "name", "created_at", "area", "segment"},
}


def _err(code: str, message: str) -> RuntimeError:
    return RuntimeError(json.dumps({"code": code, "message": message}))


class _Q(R._Query):
    """r7b's query, refusing a column the table does not have (42703 in a select list, a filter or an
    order; PGRST204 in a written row) and keeping r8m's one-approved-report-per-month index (23505)."""

    def __init__(self, db, table):
        super().__init__(db, table)
        self.unknown: list[str] = []

    def _known(self, col):
        cols = SCHEMA.get(self.table)
        if cols is not None and col not in cols:
            self.unknown.append(col)

    def select(self, cols="*", count=None, head=False):
        if cols and cols != "*":
            for c in str(cols).split(","):
                self._known(c.strip())
        return super().select(cols, count=count, head=head)

    def _add(self, kind, col, val):
        self._known(col)
        return super()._add(kind, col, val)

    def order(self, col, desc=False, **kw):
        self._known(col)
        return super().order(col, desc=desc, **kw)

    def execute(self):
        if self.unknown:
            self.db.calls.append(("refused", self.table, None, [("column", c, None) for c in self.unknown]))
            raise _err("42703", f"column {self.table}.{self.unknown[0]} does not exist")
        cols = SCHEMA.get(self.table)
        if cols is not None and self.op in ("insert", "update", "upsert"):
            for p in (self.payload if isinstance(self.payload, list) else [self.payload]):
                extra = sorted(set(p) - cols)
                if extra:
                    raise _err("PGRST204", f"Could not find the '{extra[0]}' column of '{self.table}' in the schema cache")
        if self.table != "supplier_reports" or self.op not in ("insert", "update"):
            return super().execute()
        rows = self.db.tables.setdefault(self.table, [])
        before = [dict(r) for r in rows]
        res = super().execute()
        seen = set()
        for r in rows:
            if r.get("status") == "approved":
                k = (str(r.get("supplier_id")), str(r.get("month")), str(r.get("edition")))
                if k in seen:
                    rows[:] = before
                    raise _err("23505", 'duplicate key value violates unique constraint "supplier_reports_one_approved_idx"')
                seen.add(k)
        return res


class _DB(R._FakeDB):
    def table(self, name):
        return _Q(self, name)


class _env(R._env):
    """r7b's environment (the fake behind every client, tokens = the e-mail's local part, the limiter
    off, the catalog served from memory) plus this module's client and a fixed clock."""

    def __enter__(self):
        fake = super().__enter__()
        from app import supplier_reports as sr
        self.sr = sr
        self.p2 = R._Patched((sr, "get_client", lambda: fake), (sr, "_now", lambda: NOW))
        self.p2.__enter__()
        sr.forget_probe()
        return fake

    def __exit__(self, *exc):
        self.p2.__exit__(*exc)
        self.sr.forget_probe()
        return super().__exit__(*exc)


class _Hook:
    """Run fn(query) once, just before the first query for which when(query) is true executes."""

    def __init__(self, fake, when, fn):
        self.fake, self.when, self.fn, self.done = fake, when, fn, False

    def __enter__(self):
        orig = self.fake.table

        def table(name, _orig=orig):
            q = _orig(name)
            real = q.execute

            def ex():
                if not self.done and self.when(q):
                    self.done = True
                    self.fn(q)
                return real()
            q.execute = ex
            return q
        self.fake.table = table
        return self

    def __exit__(self, *exc):
        del self.fake.table               # the class's table() again
        return False


# ── an invented world: supplier X (the subject), supplier Y, a supplier merged into X ──

USERS = [
    {"email": "boss@example.com", "role": "admin", "features": [], "status": "active"},
    {"email": "clerk@example.com", "role": "member", "features": ["Dashboard", "Market Intel"], "status": "active"},
    {"email": "mgmt@example.com", "role": "management", "features": ["Dashboard", "Sales", "Market Intel"],
     "status": "active"},
    {"email": "rep1@example.com", "role": "salesman", "features": ["Catalog", "Shop Orders", "Market Intel"],
     "status": "active"},
    {"email": "rep3@example.com", "role": "salesman", "features": ["Catalog", "Shop Orders"], "status": "active"},
    {"email": "store@example.com", "role": "storekeeper", "features": ["Storekeeper", "Market Intel"],
     "status": "active"},
]
SALESMEN = [{"id": 1, "name": "Rep One", "focus_name": "Rep One - Acc", "focus_aliases": ["R. One Old"],
             "referral_code": "rep-one", "is_active": True, "user_email": "rep1@example.com"},
            {"id": 2, "name": "Rep Two", "focus_name": None, "referral_code": "rep-two", "is_active": True,
             "user_email": "rep2@example.com"}]
SHOPS = [{"id": i, "phone": f"9730000000{i}", "name": f"Owner {n}", "shop": f"Shop {n}", "area": "Manama"}
         for i, n in ((1, "Alpha"), (2, "Beta"), (3, "Gamma"), (4, "Delta"), (5, "Epsilon"))]
CUSTOMERS = [{"id": 1, "name": "Focus Trader Kappa"}]
SUPPLIERS = [{"id": 1, "name": "Supplier X", "code": "SUPX", "active": True, "merged_into": None},
             {"id": 2, "name": "Supplier Y", "code": "SUPY", "active": True, "merged_into": None},
             {"id": 3, "name": "Supplier Old", "code": "SUPO", "active": False, "merged_into": 1}]
BRANDS = [{"id": 10, "name": "Brand X", "supplier_id": 1, "aliases": [], "active": True, "merged_into": None},
          {"id": 11, "name": "Brand Y", "supplier_id": 2, "aliases": ["BrandY Classic"], "renamed_from": "Ylite",
           "active": True, "merged_into": None},
          {"id": 12, "name": "Brand X Old", "supplier_id": None, "aliases": [], "active": False, "merged_into": 10}]
TYPES = [{"id": i, "kind": k, "code": c, "label_en": lab, "label_ar": None, "safety_default": s, "fields": [],
          "active": True, "sort": i}
         for i, k, c, lab, s in ((100, "quality", "not_charging", "Not charging", False),
                                 (101, "quality", "cable_breaks", "Cable breaks", False),
                                 (102, "packaging", "print_fading", "Print fading", False),
                                 (103, "quality", "loose_connector", "Loose connector", False),
                                 (104, "quality", "swelling", "Swelling or heat", True),
                                 (105, "competitor", "competitor_price", "Competitor price", False),
                                 (106, "demand", "customer_asked", "Customer asked for", False))]
CATALOG = {"items": {c: {"item_code": c, "display_name": n, "category": "CABLE", "brand": "Brand X"}
                     for c, n in (("C10", "Cable C10 Type-C"), ("C11", "Cable C11 Lightning"), ("P22", "Power Bank P22"),
                                  ("E05", "Earphone E05"), ("E06", "Earphone E06"), ("B30", "Battery Case B30"),
                                  ("B31", "Battery Pack B31"), ("Y99", "Charger Y99"), ("Y98", "Cable Y98"))}}


def _case(cid, sku, itype, **kw):
    row = {"id": cid, "yq_item_code": sku, "issue_type_id": itype, "brand_id": None, "supplier_id": None,
           "status": "submitted", "evidence_level": 1, "safety": False, "title": f"{sku} problem", "obs_count": 1,
           "merged_into": None, "batch_ref": None, "packaging_version": None, "verification_method": None,
           "verification_reason": None}
    row.update(kw)
    return row


VERIFIED = {"status": "verified", "evidence_level": 3, "verification_method": "sample_inspected",
            "verification_reason": "returned samples inspected"}
CASES = [
    _case(1, "C10", 100, supplier_id=1, obs_count=5, **VERIFIED),
    _case(2, "P22", 102, brand_id=10, status="under_review", evidence_level=2, packaging_version="v2"),
    _case(3, "E05", 103, supplier_id=1),                                   # one report, no photo: insufficient
    _case(4, "B30", 101, status="under_review", evidence_level=2),         # its supplier: the sighting's (merged)
    _case(5, "B30", 101, supplier_id=1, **dict(VERIFIED, status="closed")),  # July: the second case (verified, closed)
    _case(6, "E05", 103, supplier_id=1),                                   # June: older than 90 days
    _case(7, "C10", 101, supplier_id=1, status="not_substantiated", evidence_level=2),
    _case(8, "C11", 100, supplier_id=1, evidence_level=2),
    _case(9, "Y99", 100, supplier_id=2, evidence_level=2),
    _case(10, "Y98", 101, supplier_id=2, evidence_level=2),
    _case(11, "B31", 104, supplier_id=1, status="under_review", evidence_level=2, safety=True),
    _case(12, "E06", 103, supplier_id=1, obs_count=2),
    _case(13, "C10", 100, merged_into=1),
]


def _o(oid, when, **kw):
    row = {"id": oid, "client_uuid": str(uuid.uuid4()), "source": "rep", "kind": "complaint", "observed_at": when,
           "created_at": when, "created_by": "rep1@example.com", "salesman_id": 1, "item_id": None, "case_id": None,
           "shop_customer_id": None, "shop_name": None, "area": "Manama", "title": None, "brand": None,
           "competitor": None, "category": None, "note": None, "price_bhd": None, "demand_level": None,
           "demand_qty": None, "barcode_raw": None, "yq_item_code": None, "issue_type_id": None, "report_basis": None,
           "qty_affected": None, "batch_ref": None, "packaging_version": None, "safety_flag": False, "brand_id": None,
           "supplier_id": None, "brand_text": None, "supplier_text": None, "signal": None}
    row.update(kw)
    return row


NOTE_101 = ("Shop Alpha owner says Rep One sold it; call +973 3900 1234; paid 2.500 BD. "
            "Supplier Y makes a better one.")
# the review's strings, fed through a sighting's note end to end
NOTE_102 = ("Focus Trader Kappa returned 3 pieces; call 3900-1234 or +973-3900-1234; we sell it at 2.500, "
            "priced at 3.250, costs 1.900, MRP 4.000; Brand Y lasts longer and BrandY Classic too")
OBS = [
    # supplier X, September (Bahrain): the findings
    _o(101, "2026-09-03T08:00:00+00:00", case_id=1, yq_item_code="C10", issue_type_id=100, shop_customer_id=1,
       note=NOTE_101, report_basis="customer_said", qty_affected=2, batch_ref="2608"),
    _o(102, "2026-09-05T08:00:00+00:00", case_id=1, yq_item_code="C10", issue_type_id=100, shop_customer_id=2,
       note=NOTE_102, created_by="rep2@example.com", salesman_id=2),
    _o(103, "2026-09-08T08:00:00+00:00", case_id=1, yq_item_code="C10", issue_type_id=100, shop_customer_id=3,
       report_basis="rep_saw", note="Seen by R. One Old at the counter"),
    _o(104, "2026-09-09T08:00:00+00:00", case_id=13, yq_item_code="C10", issue_type_id=100, shop_customer_id=1),
    _o(111, "2026-08-31T22:30:00+00:00", case_id=2, yq_item_code="P22", issue_type_id=102, shop_customer_id=1,
       packaging_version="v2", note="Owner Alpha called about the faded box"),
    _o(112, "2026-09-12T08:00:00+00:00", case_id=2, yq_item_code="P22", issue_type_id=102, shop_customer_id=2,
       packaging_version="v2"),
    _o(121, "2026-09-14T08:00:00+00:00", case_id=3, yq_item_code="E05", issue_type_id=103, shop_customer_id=4),
    _o(131, "2026-09-16T08:00:00+00:00", case_id=4, yq_item_code="B30", issue_type_id=101, shop_customer_id=5,
       supplier_id=3),
    _o(141, "2026-09-18T08:00:00+00:00", case_id=8, yq_item_code="C11", issue_type_id=100, shop_customer_id=4),
    _o(151, "2026-09-20T08:00:00+00:00", case_id=11, yq_item_code="B31", issue_type_id=104, shop_name="Shop Zeta",
       safety_flag=True, note="Seen at Shop Zeta, the pack was hot"),
    _o(161, "2026-09-22T08:00:00+00:00", case_id=12, yq_item_code="E06", issue_type_id=103, shop_customer_id=1),
    _o(162, "2026-09-24T08:00:00+00:00", case_id=12, yq_item_code="E06", issue_type_id=103, shop_customer_id=5),
    # supplier X, September: market sightings (the internal edition only)
    _o(171, "2026-09-25T08:00:00+00:00", kind="competitor_price", issue_type_id=105, brand_id=12, shop_customer_id=2,
       title="Rival cable seen", competitor="Rival Co", price_bhd="1.250", note="Shop Beta stocks it"),
    _o(181, "2026-09-26T08:00:00+00:00", kind="shop_asked", issue_type_id=106, supplier_id=1, shop_customer_id=3,
       title="Asked for C10 in white", demand_qty=20),
    # never in X's September report
    _o(191, "2026-09-10T08:00:00+00:00", case_id=7, yq_item_code="C10", issue_type_id=101, shop_customer_id=2),
    _o(201, "2026-09-11T08:00:00+00:00", case_id=9, yq_item_code="Y99", issue_type_id=100, shop_customer_id=1),
    _o(202, "2026-09-13T08:00:00+00:00", case_id=10, yq_item_code="Y98", issue_type_id=101, shop_customer_id=2),
    _o(203, "2026-09-15T08:00:00+00:00", kind="competitor_price", brand_id=11, title="Rival charger",
       price_bhd="2.750"),
    _o(301, "2026-07-20T09:00:00+00:00", case_id=5, yq_item_code="B30", issue_type_id=101, shop_customer_id=2),
    _o(311, "2026-06-20T09:00:00+00:00", case_id=6, yq_item_code="E05", issue_type_id=103, shop_customer_id=1),
    _o(321, "2026-09-30T21:30:00+00:00", case_id=1, yq_item_code="C10", issue_type_id=100, shop_customer_id=4),
    _o(331, "2026-09-17T08:00:00+00:00", kind="complaint", brand="Brand X", note="v1 complaint, no links"),
]
X_FINDINGS = {101, 102, 103, 104, 111, 112, 121, 131, 141, 151, 161, 162}
X_COUNTED = X_FINDINGS - {121}
X_MARKET = {171, 181}
NEVER_IN_X = {191, 201, 202, 203, 301, 311, 321, 331}

ACTIONS = [
    {"id": 1, "supplier_id": 1, "request": "Replace 24 units of B30, batch 2608", "owner": "Office Lead",
     "due_date": "2026-10-15", "status": "open", "supplier_response": None, "resolution_check": None,
     "created_at": "2026-09-18T07:00:00+00:00", "updated_at": "2026-09-18T07:00:00+00:00"},
    {"id": 2, "supplier_id": 2, "request": "Supplier Y matter", "owner": "Office Lead", "due_date": None,
     "status": "open", "supplier_response": None, "resolution_check": None,
     "created_at": "2026-09-10T07:00:00+00:00", "updated_at": "2026-09-10T07:00:00+00:00"},
    {"id": 3, "supplier_id": None, "request": "Check the print on packaging v2 for Shop Beta", "owner": "Office Lead",
     "due_date": "2026-10-20", "status": "responded",
     "supplier_response": "We will check the ink. Write to qa@supplier-x.example", "resolution_check": None,
     "created_at": "2026-09-20T07:00:00+00:00", "updated_at": "2026-09-28T07:00:00+00:00"},
    {"id": 4, "supplier_id": 1, "request": "An old matter", "owner": None, "due_date": None, "status": "resolved",
     "supplier_response": "Done", "resolution_check": "Checked",
     "created_at": "2026-07-01T07:00:00+00:00", "updated_at": "2026-08-10T07:00:00+00:00"},
    {"id": 5, "supplier_id": 1, "request": "Relabel the E06 boxes", "owner": "Office Lead", "due_date": "2026-09-30",
     "status": "resolved", "supplier_response": "Relabelled", "resolution_check": "Checked ten boxes",
     "created_at": "2026-09-02T07:00:00+00:00", "updated_at": "2026-09-25T07:00:00+00:00"},
    {"id": 6, "supplier_id": 1, "request": "Withdrawn: a mistaken request", "owner": "Office Lead", "due_date": None,
     "status": "cancelled", "supplier_response": None, "resolution_check": None,
     "created_at": "2026-09-05T07:00:00+00:00", "updated_at": "2026-09-21T07:00:00+00:00"},
]
LINKS = [{"action_id": 1, "case_id": 4}, {"action_id": 1, "case_id": 5}, {"action_id": 2, "case_id": 9},
         {"action_id": 3, "case_id": 2}, {"action_id": 5, "case_id": 12}]

# photos: colour = identity. RED is the people photo, YELLOW the receipt; neither may reach a supplier.
RED, YELLOW = (220, 20, 20), (235, 215, 30)
PHOTOS = [
    # id, observation, role, people, colour, stored as: "master" | "thumb" (only the 320 px thumbnail) | None
    (1, 101, "product", False, (40, 160, 60), "master"),
    (2, 101, "product", True, RED, "master"),
    (3, 111, "packaging", False, (30, 60, 200), "master"),
    (4, 111, "receipt", False, YELLOW, "master"),
    (5, 112, "label_batch", False, (120, 40, 160), "master"),
    (6, 131, "product", False, (90, 90, 90), None),
    (7, 141, "product", False, (0, 150, 150), "thumb"),
    (8, 151, "product", False, (240, 140, 20), "master"),
    (9, 171, "product", False, (150, 150, 150), "master"),
    (10, 191, "product", False, (100, 60, 20), "master"),
    (11, 201, "product", False, (250, 150, 200), "master"),
]
CAPTION_1 = "Cable end burnt, seen at Shop Alpha"
SHA_5 = "ab" * 32


def _picture(colour, size=(800, 600)) -> tuple[bytes, str]:
    from PIL import Image, features
    img = Image.new("RGB", size, colour)
    buf = io.BytesIO()
    if features.check("webp"):
        img.save(buf, format="WEBP", quality=80)
        return buf.getvalue(), ".webp"
    img.save(buf, format="JPEG", quality=85)
    return buf.getvalue(), ".jpg"


def _world(**over) -> _DB:
    tables = {"user_roles": USERS, "salesmen": SALESMEN, "shop_customers": SHOPS, "customers": CUSTOMERS,
              "audit_log": [], "suppliers": SUPPLIERS, "brands": BRANDS, "market_issue_types": TYPES,
              "market_cases": CASES, "market_observations": OBS, "market_photos": [], "supplier_actions": ACTIONS,
              "supplier_action_cases": LINKS, "supplier_reports": [], "market_items": [], "app_settings": []}
    tables.update(over)
    fake = _DB(tables)
    for pid, oid, role, people, colour, stored in PHOTOS:
        data, ext = _picture(colour)
        path, thumb = f"market/2026/09/p{pid}{ext}", f"market/2026/09/p{pid}-t320{ext}"
        fake.tables["market_photos"].append(
            {"id": pid, "observation_id": oid, "bucket": "finds", "path": path, "thumb_path": thumb,
             "position": pid, "role": role, "has_people_flag": people, "sha256": SHA_5 if pid == 5 else None,
             "caption": CAPTION_1 if pid == 1 else None, "blurred_path": None})
        if stored == "master":
            fake.storage.objects[("finds", path)] = (data, {})
        elif stored == "thumb":
            fake.storage.objects[("finds", thumb)] = (_picture(colour, (320, 240))[0], {})
    return fake


def _sr():
    from app import supplier_reports as sr
    return sr


def _build(edition="supplier", supplier_id=1, month=SEP, actor="clerk@example.com", role="member"):
    return _sr().build_draft(supplier_id, month, edition, actor, role)


def _row(fake, table, rid):
    return next(r for r in fake.rows(table) if r["id"] == rid)


def _raises(fn, want: str, status: int | None = None):
    sr = _sr()
    try:
        fn()
    except sr.SupplierReportError as e:
        assert want.lower() in str(e).lower(), (want, str(e))
        if status is not None:
            assert e.status == status, (status, e.status, str(e))
        return e
    raise AssertionError(f"expected SupplierReportError containing {want!r}")


def _book(data: bytes):
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(data))


def _cells(wb) -> list[str]:
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            out += [str(v) for v in row if v is not None]
    return out


def _head(ws) -> list:
    return [c.value for c in ws[1]]


def _colours(ws) -> list[tuple]:
    from PIL import Image
    out = []
    for im in ws._images:
        pic = Image.open(io.BytesIO(im._data()))
        out.append((pic.format, pic.size, pic.convert("RGB").resize((1, 1)).getpixel((0, 0))))
    return out


def _near(c, want, tol=40) -> bool:
    return all(abs(a - b) <= tol for a, b in zip(c, want))


def _ids(text: str, prefix: str) -> set[int]:
    return {int(x) for x in re.findall(rf"\b{prefix}-(\d+)\b", text)}


def _approved(fake, edition="supplier", supplier_id=1, actor="boss@example.com", role="admin"):
    rid = _build(edition, supplier_id=supplier_id)["report"]["id"]
    _sr().approve(rid, actor, role)
    return rid


def _download(rid, actor="clerk@example.com", role="member"):
    return _sr().download(rid, actor, role)


FINDING_PREFIXES = ("Recurring:", "New:", "Insufficient evidence:", "We ask:", "Competitor:", "Demand:")


# ═══════════════════════════════════════════════════════════════════════════════
# 1. the month, reads, the harness, the scrub
# ═══════════════════════════════════════════════════════════════════════════════

@test("month: YYYY-MM to the first day; junk refused; a month that has not started refused; approvable from the 8th of the next month")
def _():
    from datetime import date
    sr = _sr()
    assert sr.parse_month("2026-9", now=NOW) == date(2026, 9, 1) and sr.parse_month("2026-09-01", now=NOW) == date(2026, 9, 1)
    assert sr.parse_month("2026-10", now=NOW) == date(2026, 10, 1)          # the current month (a draft)
    for bad in ("2026-13", "Sept 2026", "", "2023-12", "2026/09"):
        _raises(lambda b=bad: sr.parse_month(b, now=NOW), "YYYY-MM")
    _raises(lambda: sr.parse_month("2026-11", now=NOW), "not started")
    assert sr.parse_month("2026-11", now=NOW, future_ok=True) == date(2026, 11, 1)
    # 31 Oct 21:30 UTC is already 1 November in Bahrain
    assert sr.parse_month("2026-11", now=datetime(2026, 10, 31, 21, 30, tzinfo=timezone.utc)) == date(2026, 11, 1)
    start, end = sr.month_bounds(date(2026, 12, 1))
    assert start.isoformat() == "2026-12-01T00:00:00+03:00" and end.isoformat() == "2027-01-01T00:00:00+03:00"
    assert sr.approvable_from(date(2026, 9, 1)).isoformat() == "2026-10-08T00:00:00+03:00"


@test("reads: paged 1,000 at a time past PostgREST's cap (2,345 rows in three reads); a read too long to finish refuses to build")
def _():
    fake = _world(suppliers=[{"id": i, "name": f"S{i}", "code": None, "merged_into": None} for i in range(1, 2346)])
    with _env(fake, CATALOG):
        sr = _sr()
        rows = sr._all(lambda: sr._q("suppliers"))
        assert len(rows) == 2345 and len({r["id"] for r in rows}) == 2345
        assert sum(1 for c in fake.calls if c[0] == "select" and c[1] == "suppliers") == 3
        _raises(lambda: sr._all(lambda: sr._q("suppliers"), max_pages=2), "too long", 503)
    shops = [{"id": i, "phone": f"9731{i:07d}", "name": f"Owner {i}", "shop": f"Store {i}"} for i in range(1, 1201)]
    fake = _world(shop_customers=shops)
    with _env(fake, CATALOG), R._Patched((_sr(), "PAGE_MAX", 1)):
        _raises(lambda: _build(), "too long", 503)
    assert not fake.rows("supplier_reports"), "nothing built from half a name list"


@test("harness: the fake refuses a column a table does not have (supplier_action_cases has no id), as PostgREST does")
def _():
    fake = _world()
    with _env(fake, CATALOG):
        try:
            fake.table("supplier_action_cases").select("*").order("id").range(0, 9).execute()
            raise AssertionError("ordered by a column that does not exist")
        except RuntimeError as e:
            assert "42703" in str(e) and "supplier_action_cases.id" in str(e)
        try:
            fake.table("supplier_reports").insert({"supplier_id": 1, "nonsense": 1}).execute()
            raise AssertionError("wrote a column that does not exist")
        except RuntimeError as e:
            assert "PGRST204" in str(e)
        assert len(fake.table("supplier_action_cases").select("*").order("action_id").order("case_id")
                   .range(0, 9).execute().data) == len(LINKS)


@test("scrub: phones with separators, bare 3-decimal amounts and price words, names in one longest-first pass (the supplier's own name never shields a longer one), other suppliers' brands and aliases")
def _():
    sr = _sr()
    s = sr.Scrubber(reps=["Rep One"], shops=["Shop Alpha", "Shop Alpha Two", "Cash Customer", "Al-Noor Mobiles"],
                    others=["Supplier Y"], other_brands=["Brand Y", "BrandY Classic"], keep=["Supplier X"])
    assert s("Shop Alpha Two and shop alpha met REP ONE") == "[shop] and [shop] met field team"
    assert s("call 3900-1234 or +973-3900-1234") == "call [number] or [number]"
    assert s("ring 3900 1234, (+973) 3900.1234 or 39 00 12 34") == "ring [number], [number] or [number]"
    assert s("we sell it at 2.500, priced at 3.250, costs 1.900, MRP 4.000") == \
        "we sell it at [amount], priced at [amount], costs [amount], MRP [amount]"
    assert s("sell at 2.5, priced at 3, costs 2, MRP 4, retail 6, margin 15") == \
        "sell at [amount], priced at [amount], costs [amount], MRP [amount], retail [amount], margin [amount]"
    big = ",".join(["1", "250"]) + ".000"            # built at run time: no such literal in a committed file
    assert s(f"paid 2.500 BD, BHD 3 and 500 fils; {big} in all") == "paid [amount], [amount] and [amount]; [amount] in all"
    assert s("sold 20 units, 4 shops, batch 2608, 65W, 1.2m, v2.010") == "sold 20 units, 4 shops, batch 2608, 65W, 1.2m, v2.010"
    assert s("Brand Y lasts longer than BrandY Classic") == "[another brand] lasts longer than [another brand]"
    assert s("Supplier X is better than Supplier Y") == "Supplier X is better than [another supplier]"
    assert s("cash customer at the counter") == "cash customer at the counter"
    assert s("Shopping alphabet soup") == "Shopping alphabet soup", "whole words only"
    assert s("Sold by Al-Noor  Mobiles, owner@shop.example") == "Sold by [shop], [email]"
    assert s(None) is None and s("") == "" and s(12) == 12
    royal = sr.Scrubber(shops=["Royal Mobile"], others=["Royal Gulf Trading"], keep=["Royal"])
    assert royal("Royal Mobile owner said Royal Gulf Trading is cheaper; Royal ships fast") == \
        "[shop] owner said [another supplier] is cheaper; Royal ships fast"


# ═══════════════════════════════════════════════════════════════════════════════
# 2. the draft: membership, counts, the rules, the narrative
# ═══════════════════════════════════════════════════════════════════════════════

@test("draft: only supplier X's September findings (links agree: case, case brand, sighting, its brand; merges followed); not substantiated, October, Y's and unlinked left out")
def _():
    fake = _world()
    with _env(fake, CATALOG):
        out = _build()
    rows = out["rows"]
    assert set(rows) == {"product_issues", "packaging"}, "a supplier edition has no competitor / demand rows"
    got = {r["id"]: r for k in rows for r in rows[k]}
    assert set(got) == X_FINDINGS, sorted(set(got) ^ X_FINDINGS)
    assert not (set(got) & NEVER_IN_X) and not (set(got) & X_MARKET)
    assert {r["id"] for r in rows["packaging"]} == {111, 112}
    assert got[111]["date"] == "2026-09-01", "22:30 UTC on 31 Aug is 1 September in Bahrain"
    assert got[121]["state"] == "insufficient" and all(got[i]["state"] == "included" for i in X_COUNTED)
    assert got[104]["case_id"] == 1, "a sighting on a merged case is read on the case it was merged into"
    assert got[101]["sku"] == "C10" and got[101]["sku_label"] == "C10"
    rep = out["report"]
    assert rep["status"] == "draft" and rep["month"] == SEP and rep["edition"] == "supplier" and rep["supplier_id"] == 1
    assert rep["included_observation_ids"] == sorted(X_COUNTED) and rep["insufficient_observation_ids"] == [121]
    assert rep["approvable_from"] == "2026-10-08" and not rep["frozen"] and out["rebuilt"] is False
    assert not any("disagree" in w for w in out["warnings"]) and not any("not closed" in w for w in out["warnings"])
    stored = fake.rows("supplier_reports")
    assert len(stored) == 1 and stored[0]["status"] == "draft" and stored[0]["month"] == "2026-09-01"
    assert stored[0]["created_by"] == "clerk@example.com" and stored[0]["included_observation_ids"] == sorted(X_COUNTED)


@test("draft: counts by issue type, new vs recurring (>= 2 cases or >= 3 shops in 90 days), verified, safety, photos, open actions (never a cancelled one)")
def _():
    with _env(_world(), CATALOG):
        out = _build()
    c = out["report"]["counts"]
    assert (c["reports"], c["product_issues"], c["packaging"], c["shops"]) == (11, 9, 2, 6), c
    assert (c["verified"], c["unverified"], c["safety"], c["photos"]) == (4, 7, 1, 6), c
    assert c["new"] == {"problems": 4, "reports": 6} and c["recurring"] == {"problems": 2, "reports": 5}, c
    assert (c["insufficient"], c["insufficient_observation_ids"], c["excluded"], c["withheld"]) == (1, [121], 0, 0)
    assert (c["open_actions"], c["action_ids"]) == (2, [1, 3, 5]), "Y's action, one resolved in August and a cancelled one are not listed"
    assert (c["competitor"], c["demand"]) == (0, 0) and c["narrative_source"] == "template"
    assert c["month"] == SEP and c["period"]["from"] == "2026-09-01" and c["period"]["to"] == "2026-09-30"
    issues = [(b["label"], b["kind"], b["reports"], b["shops"], b["verified"], b["recurring"]) for b in c["by_issue"]]
    assert issues == [("Not charging", "quality", 5, 4, 4, True), ("Loose connector", "quality", 2, 2, 0, False),
                      ("Print fading", "packaging", 2, 2, 0, False), ("Cable breaks", "quality", 1, 1, 0, True),
                      ("Swelling or heat", "quality", 1, 1, 0, False)], issues
    groups = {(g["sku"], g["issue"]["label"]): g for g in out["groups"]}
    assert groups[("C10", "Not charging")]["recurring"] and groups[("C10", "Not charging")]["reports"] == 4
    assert groups[("B30", "Cable breaks")]["recurring"]
    assert not groups[("E06", "Loose connector")]["recurring"] and groups[("E06", "Loose connector")]["reports"] == 2
    assert ("E05", "Loose connector") not in groups
    row = next(r for r in out["rows"]["product_issues"] if r["id"] == 101)
    assert (row["finding"], row["evidence_label"], row["status"], row["basis"]) == \
        ("Verified", "Inspected & verified by YQ", "Verified by YQ", "Customer said")
    assert (row["qty"], row["batch_ref"], row["recurring"]) == (2, "2608", True)
    row = next(r for r in out["rows"]["product_issues"] if r["id"] == 141)
    assert (row["finding"], row["evidence_label"]) == ("Unverified report", "Evidence submitted")
    assert [p["ref"] for p in next(r for r in out["rows"]["packaging"] if r["id"] == 111)["photos"]] == ["PH-3"], \
        "the receipt photo is not a supplier's"
    with _env(_world(), CATALOG):
        internal = _build("internal")
    assert internal["report"]["counts"]["action_ids"] == [1, 3, 5, 6], "the internal edition keeps a cancelled action"


@test("verified means a recorded verification method: a case raised to evidence 3, found not substantiated, then closed is left out; one at 3 without a method is an unverified report")
def _():
    # the path through the case routes: PATCH evidence_level=3, then not_substantiated (a reason), then closed;
    # set in the data here, since the case module will also refuse evidence >= 3 without a verification
    cases = [dict(c, evidence_level=3, status="closed", verification_method=None, status_reason="not substantiated")
             if c["id"] == 12 else dict(c, evidence_level=3, verification_method=None) if c["id"] == 8 else c
             for c in CASES]
    fake = _world(market_cases=cases)
    with _env(fake, CATALOG):
        sr = _sr()
        out = _build()
        rows = {r["id"]: r for k in out["rows"] for r in out["rows"][k]}
        assert 161 not in rows and 162 not in rows, "closed without a verification method: not substantiated"
        assert (rows[141]["finding"], rows[141]["evidence_label"], rows[141]["evidence_level"]) == \
            ("Unverified report", "Evidence submitted", 2)
        c = out["report"]["counts"]
        assert (c["verified"], c["reports"]) == (4, 9), c
        text = out["report"]["narrative"]
        assert "[MI-141]" in text and "verified by YQ [MI-141]" not in text and "MI-161" not in text
        rid = out["report"]["id"]
        sr.approve(rid, "boss@example.com", "admin")
        ws = _book(_download(rid)["data"])["Product issues"]
        got = {ws.cell(row=r, column=1).value: (ws.cell(row=r, column=6).value, ws.cell(row=r, column=11).value)
               for r in range(2, ws.max_row + 1)}
        assert got["MI-141"] == ("Unverified report", "Evidence submitted") and "MI-161" not in got
        # after approval, the same closing without a verification withdraws a frozen row: the download stops
        next(x for x in fake.tables["market_cases"] if x["id"] == 1).update(status="closed", verification_method=None)
        _raises(lambda: _download(rid), "MI-101 is no longer a finding", 409)


@test("narrative: every sentence about a finding cites real ids of supplier X; actions cite CA ids; the insufficient paragraph; the template passes its own checks")
def _():
    with _env(_world(), CATALOG):
        sr = _sr()
        out = _build()
    text = out["report"]["narrative"]
    lines = text.split("\n")
    assert lines[0] == "Supplier X: September 2026."
    assert lines[1] == ("Field reports counted this month: 11 (9 product issues, 2 packaging) from 6 shops; "
                        "4 verified by YQ; 1 flagged as a possible safety issue."), lines[1]
    finding = [ln for ln in lines if ln.startswith(FINDING_PREFIXES)]
    assert len(finding) == 6 + 1 + 1, finding
    for ln in finding:
        ids = _ids(ln, "MI")
        assert ids and ids <= X_FINDINGS, ln
    cited = _ids(text, "MI")
    assert cited == X_FINDINGS and not (cited & (NEVER_IN_X | X_MARKET))
    acts = [ln for ln in lines if ln.startswith("Open action")]
    assert len(acts) == 2 and all(_ids(ln, "CA") for ln in acts) and _ids(text, "CA") == {1, 3}
    assert lines[2].startswith('Recurring: "Not charging" on C10 (Cable C10 Type-C), batch "2608": 4 reports from '
                               "3 shops [MI-101] [MI-102] [MI-103] [MI-104]; 1 case from 3 shops in the last 90 days; "
                               "4 verified by YQ"), lines[2]
    assert any(ln.startswith('Recurring: "Cable breaks" on B30') and "2 cases from 2 shops in the last 90 days" in ln
               for ln in lines)
    assert any(ln.startswith('New: "Swelling or heat" on B31') and "flagged as a possible safety issue" in ln
               for ln in lines)
    assert ('Insufficient evidence: one report of "Loose connector" on E05 (Earphone E05) [MI-121]; no photo, a single '
            "report and not verified, so not included in the counts.") in lines
    assert ('Open action [CA-1]: "Replace 24 units of B30, batch 2608", due 15 Oct 2026; awaiting your response since '
            "18 Sep 2026; linked reports [MI-131].") in lines
    assert any(ln.startswith('Open action [CA-3]: "Check the print on packaging v2 for [shop]"') for ln in lines)
    assert lines[-1] == ('We ask: the root cause and a corrective plan for "Not charging" on C10 (Cable C10 Type-C) '
                         "[MI-101]."), "B30 already has an open action; the new problems are not verified"
    for secret in ("Shop Beta", "Rep One", "BHD", "Supplier Y"):
        assert secret not in text, secret
    with _env(_world(), CATALOG):
        ctx = sr.gather(1, sr.parse_month(SEP, now=NOW), "supplier", now=NOW)
        a = sr.assemble(ctx)
        assert sr.narrative(ctx, a) == text
        assert sr.citation_problems(text, a) == [] and sr.uncited_findings(text, a) == []


@test("draft: the internal edition adds the competitor and demand sightings (rows, counts, cited lines); Y's never")
def _():
    with _env(_world(), CATALOG):
        out = _build("internal")
    assert set(out["rows"]) == {"product_issues", "packaging", "competitor", "demand"}
    assert [r["id"] for r in out["rows"]["competitor"]] == [171] and [r["id"] for r in out["rows"]["demand"]] == [181]
    comp, dem = out["rows"]["competitor"][0], out["rows"]["demand"][0]
    assert comp["price_seen"] == "1.250" and comp["competitor"] == "Rival Co" and comp["note"] == "[shop] stocks it"
    assert dem["qty"] == 20 and dem["source"] == "Field team" and dem["asked_for"] == "Asked for C10 in white"
    c = out["report"]["counts"]
    assert (c["competitor"], c["demand"], c["reports"]) == (1, 1, 11)
    assert out["report"]["included_observation_ids"] == sorted(X_COUNTED | X_MARKET)
    lines = out["report"]["narrative"].split("\n")
    assert lines[0] == "Supplier X: September 2026. Internal edition: for YQ only, not for the supplier."
    assert "Competitor: 1 sighting of competing products this month [MI-171]." in lines
    assert "Demand: 1 request for this range this month [MI-181]." in lines
    assert [p["id"] for p in out["rows"]["packaging"][0]["photos"]] == [3, 4], "the receipt is internal-only"


@test("draft: supplier Y's report holds Y's rows only")
def _():
    with _env(_world(), CATALOG):
        out = _build(supplier_id=2)
        internal = _build("internal", supplier_id=2)
    got = {r["id"] for k in out["rows"] for r in out["rows"][k]}
    assert got == {201, 202}, got
    assert {r["id"] for r in internal["rows"]["competitor"]} == {203}
    assert not (_ids(out["report"]["narrative"], "MI") & (X_FINDINGS | X_MARKET))
    assert out["report"]["counts"]["action_ids"] == [2]


@test("links: a brand re-linked to Y leaves a sighting whose case still says X in nobody's draft, with a warning in both, and blocks both approvals until corrected")
def _():
    cases = CASES + [_case(20, "C10", 102, brand_id=10, supplier_id=1, status="under_review", evidence_level=2)]
    brands = [dict(b, supplier_id=2) if b["id"] == 10 else b for b in BRANDS]     # the admin's correction
    obs = OBS + [_o(901, "2026-09-21T08:00:00+00:00", case_id=20, yq_item_code="C10", issue_type_id=102, brand_id=10,
                    supplier_id=2, shop_customer_id=3, note="captured after the correction")]
    fake = _world(market_cases=cases, brands=brands, market_observations=obs)
    with _env(fake, CATALOG):
        sr = _sr()
        x, y = _build(supplier_id=1), _build(supplier_id=2)
        in_x = {r["id"] for k in x["rows"] for r in x["rows"][k]}
        in_y = {r["id"] for k in y["rows"] for r in y["rows"][k]}
        assert 901 not in in_x and 901 not in in_y
        for out in (x, y):
            w = [w for w in out["warnings"] if "MI-901" in w]
            assert w and "disagree" in w[0] and "Supplier X" in w[0] and "Supplier Y" in w[0], out["warnings"]
        assert {111, 112} <= in_y and not ({111, 112} & in_x), "every link of case 2 now says Y: it follows the brand"
        for rid in (x["report"]["id"], y["report"]["id"]):
            e = _raises(lambda r=rid: sr.approve(r, "boss@example.com", "admin"), "MI-901", 409)
            assert "disagree" in str(e)
        next(c for c in fake.tables["market_cases"] if c["id"] == 20)["supplier_id"] = 2      # the office corrects it
        assert sr.approve(x["report"]["id"], "boss@example.com", "admin")["status"] == "approved"
        y2 = _build(supplier_id=2)
        assert 901 in {r["id"] for k in y2["rows"] for r in y2["rows"][k]}


@test("draft: a rebuild updates the same draft; unknown or merged supplier, bad edition, future month refused; only reviewers build; a month not closed warns")
def _():
    fake = _world()
    with _env(fake, CATALOG):
        first = _build()
        again = _build(actor="boss@example.com", role="admin")
        assert again["report"]["id"] == first["report"]["id"] and again["rebuilt"] is True
        assert len(fake.rows("supplier_reports")) == 1 and fake.rows("supplier_reports")[0]["updated_by"] == "boss@example.com"
        other = _build("internal")
        assert other["report"]["id"] != first["report"]["id"] and len(fake.rows("supplier_reports")) == 2
        _raises(lambda: _build(supplier_id=99), "does not exist", 404)
        e = _raises(lambda: _build(supplier_id=3), "merged into Supplier X", 409)
        assert "build that supplier's report" in str(e)
        _raises(lambda: _build("public"), "edition", 400)
        _raises(lambda: _build(month="2026-11"), "not started", 400)
        for role in ("management", "salesman", "storekeeper"):
            _raises(lambda r=role: _build(role=r), "only the office", 403)
        october = _build(month="2026-10")
        assert any("October 2026 is not closed yet" in w and "approved from 8 Nov 2026" in w for w in october["warnings"])


@test("SKU: a code the catalog does not confirm prints as 'unverified SKU' and is masked in every text; an unreadable catalog confirms none")
def _():
    cases = CASES + [_case(30, "ZZ9", 100, supplier_id=1, evidence_level=2, title="ZZ9 Cable · Not charging")]
    obs = OBS + [_o(941, "2026-09-19T08:00:00+00:00", case_id=30, yq_item_code="ZZ9", issue_type_id=100,
                    shop_customer_id=2, note="code ZZ9 on the pack")]
    fake = _world(market_cases=cases, market_observations=obs)
    with _env(fake, CATALOG):
        out = _build()
    row = next(r for r in out["rows"]["product_issues"] if r["id"] == 941)
    assert row["sku"] is None and row["sku_label"] == "unverified SKU" and "ZZ9" not in json.dumps(row)
    line = next(ln for ln in out["report"]["narrative"].split("\n") if "[MI-941]" in ln)
    assert "a product with an unverified SKU" in line and "ZZ9" not in line, line
    with _env(_world(), {"items": {}}):
        blind = _build()
    assert {r["sku_label"] for k in blind["rows"] for r in blind["rows"][k]} == {"unverified SKU"}
    assert "C10" not in blind["report"]["narrative"]


# ═══════════════════════════════════════════════════════════════════════════════
# 3. edits, approval, supersede
# ═══════════════════════════════════════════════════════════════════════════════

@test("edit: excluding a row recounts and rewrites a template narrative; a weak row can be included; ids outside the report refused")
def _():
    fake = _world()
    with _env(fake, CATALOG):
        sr = _sr()
        rid = _build()["report"]["id"]
        out = sr.update_draft(rid, {"included_observation_ids": sorted(X_COUNTED - {102})}, "clerk@example.com", "member")
        assert out["report"]["counts"]["reports"] == 10 and "[MI-102]" not in out["report"]["narrative"]
        assert next(r for r in out["rows"]["product_issues"] if r["id"] == 102)["state"] == "excluded"
        out = sr.update_draft(rid, {"included_observation_ids": sorted(X_COUNTED | {121})}, "clerk@example.com", "member")
        c = out["report"]["counts"]
        assert c["reports"] == 12 and c["insufficient"] == 0 and "Insufficient evidence" not in out["report"]["narrative"]
        for foreign in (201, 191, 321, 171):         # Y's, not substantiated, October, a market row in a supplier copy
            _raises(lambda f=foreign: sr.update_draft(rid, {"included_observation_ids": [101, f]}, "clerk@example.com",
                                                      "member"), f"MI-{foreign}", 400)
        _raises(lambda: sr.update_draft(rid, {"included_observation_ids": ["x"]}, "clerk@example.com", "member"),
                "numbers", 400)
        assert _row(fake, "supplier_reports", rid)["included_observation_ids"] == sorted(X_COUNTED | {121})


@test("edit: a weak row or a photo can be withheld from the supplier (shown nowhere, frozen at approval); a photo of another report refused")
def _():
    fake = _world()
    with _env(fake, CATALOG):
        sr = _sr()
        rid = _build()["report"]["id"]
        out = sr.update_draft(rid, {"withheld_observation_ids": [121], "withheld_photo_ids": [5]},
                              "clerk@example.com", "member")
        rows = {r["id"]: r for k in out["rows"] for r in out["rows"][k]}
        assert rows[121]["state"] == "withheld" and rows[112]["photos"] == [] and rows[112]["photos_withheld"] == 1
        c = out["report"]["counts"]
        assert (c["withheld"], c["photos_withheld"], c["photos"], c["insufficient"]) == (1, 1, 5, 0)
        assert out["report"]["withheld_observation_ids"] == [121] and out["report"]["withheld_photo_ids"] == [5]
        assert "MI-121" not in out["report"]["narrative"] and "1 photo on the Evidence sheet" in \
            next(ln for ln in out["report"]["narrative"].split("\n") if "[MI-111]" in ln)
        _raises(lambda: sr.update_draft(rid, {"withheld_photo_ids": [11]}, "clerk@example.com", "member"), "PH-11", 400)
        _raises(lambda: sr.update_draft(rid, {"narrative": "New: see [MI-121]."}, "clerk@example.com", "member"),
                "MI-121", 400)
        sr.approve(rid, "mgmt@example.com", "management")
        text = "\n".join(_cells(_book(_download(rid)["data"])))
    assert "MI-121" not in text and "PH-5" not in text, "withheld at approval, withheld in the file"
    assert "no photo shared" in text


@test("edit: an edited narrative may cite only listed rows, and every sentence that states a finding must cite one; it is kept when rows change")
def _():
    with _env(_world(), CATALOG):
        sr = _sr()
        rid = _build()["report"]["id"]
        sr.update_draft(rid, {"included_observation_ids": sorted(X_COUNTED - {102})}, "clerk@example.com", "member")
        for text, want in (("New: see [MI-102].", "MI-102"),            # an excluded row (not printed anywhere)
                           ("New: see [MI-201].", "MI-201"),            # another supplier's
                           ("Open action [CA-2]: Y's.", "CA-2"),
                           ("Open action [CA-6]: cancelled.", "CA-6"),  # never in a supplier copy
                           ("New: C10 keeps failing.", "Add one to"),
                           ("Supplier X: September 2026.\nThe C10 cable fails after a week.", "C10 cable"),
                           ("We all saw Not charging on many units.", "Not charging"),
                           ("Open action: we wait for the supplier.", "Add one to")):
            _raises(lambda t=text: sr.update_draft(rid, {"narrative": t}, "clerk@example.com", "member"), want, 400)
        text = ("Supplier X: September 2026.\r\nNew: a hand-written line [MI-101] [MI-121]. Thanks for the quick "
                "replies.\n\nOpen action [CA-1]: due soon. Dr. Example will call.")
        out = sr.update_draft(rid, {"narrative": text}, "clerk@example.com", "member")
        assert out["report"]["narrative_source"] == "edited"
        assert out["report"]["narrative"].startswith("Supplier X: September 2026.\nNew: a hand-written line [MI-101] [MI-121].")
        out = sr.update_draft(rid, {"included_observation_ids": sorted(X_COUNTED - {102, 103})}, "clerk@example.com",
                              "member")
        assert out["report"]["narrative"].startswith("Supplier X: September 2026.\nNew: a hand-written line")
        assert any("edited by hand" in w for w in out["warnings"])
        assert any("data changed since the narrative was edited" in w for w in out["warnings"])
        _raises(lambda: sr.update_draft(rid, {"narrative": "x" * (sr.MAX_NARRATIVE + 1)}, "clerk@example.com",
                                        "member"), "limited", 400)


@test("approve: management or an admin only; not before the month's end + 7 days; a draft only; the newer approval supersedes the earlier (same supplier, month, edition only)")
def _():
    fake = _world()
    with _env(fake, CATALOG):
        sr = _sr()
        a = _build()["report"]["id"]
        for role in ("member", "salesman", "storekeeper"):
            _raises(lambda r=role: sr.approve(a, "clerk@example.com", r), "only management", 403)
        _raises(lambda: sr.approve(a, "boss@example.com", "admin", now=datetime(2026, 10, 7, 20, 0, tzinfo=timezone.utc)),
                "can be approved from 8 Oct 2026", 409)
        internal = _build("internal")["report"]["id"]
        sr.approve(internal, "boss@example.com", "admin")
        got = sr.approve(a, "boss@example.com", "admin")
        assert got["status"] == "approved" and got["superseded"] == [] and got["approved_by"] == "boss"
        assert len(got["warnings"]) == 1 and got["warnings"][0].startswith("PH-6 could not be read"), got["warnings"]
        _raises(lambda: sr.approve(a, "boss@example.com", "admin"), "only a draft", 409)
        _raises(lambda: sr.update_draft(a, {"narrative": "x"}, "clerk@example.com", "member"), "only a draft", 409)
        b = _build()["report"]["id"]
        assert b != a, "an approved report is never rebuilt: a new draft starts"
        got = sr.approve(b, "mgmt@example.com", "management")
        assert got["superseded"] == [a]
        status = {r["id"]: r["status"] for r in fake.rows("supplier_reports")}
        assert status == {a: "superseded", internal: "approved", b: "approved"}, status
        row = _row(fake, "supplier_reports", b)
        assert row["approved_by"] == "mgmt@example.com" and row["approved_at"] == NOW.isoformat()
        oct_draft = _build(month="2026-10")["report"]["id"]
        _raises(lambda: sr.approve(oct_draft, "boss@example.com", "admin"), "can be approved from 8 Nov 2026", 409)


@test("approve: a template narrative is rewritten from the data at approval (a case verified since the draft); an edited one is refused until saved again")
def _():
    fake = _world()
    with _env(fake, CATALOG):
        sr = _sr()
        rid = _build()["report"]["id"]
        next(c for c in fake.tables["market_cases"] if c["id"] == 8).update(VERIFIED)
        sr.approve(rid, "boss@example.com", "admin")
        row = _row(fake, "supplier_reports", rid)
        assert row["counts"]["verified"] == 5 and "; 5 verified by YQ" in row["narrative"].split("\n")[1]
        ws = _book(_download(rid)["data"])["Summary"]
        facts = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(3, 17)}
        assert facts["Verified by YQ"] == 5, "the narrative and the sheet tell the same story"
    fake = _world()
    with _env(fake, CATALOG):
        sr = _sr()
        rid = _build()["report"]["id"]
        sr.update_draft(rid, {"narrative": "Supplier X: September 2026.\nNew: C11 [MI-141] is not charging."},
                        "clerk@example.com", "member")
        next(c for c in fake.tables["market_cases"] if c["id"] == 8).update(VERIFIED)
        _raises(lambda: sr.approve(rid, "boss@example.com", "admin"), "data changed since the narrative was edited", 409)
        sr.update_draft(rid, {"narrative": "Supplier X: September 2026.\nVerified: C11 [MI-141] is not charging."},
                        "clerk@example.com", "member")
        assert sr.approve(rid, "boss@example.com", "admin")["status"] == "approved"


def _approving(q) -> bool:
    """The approval's own write (not a supersede, not a restore)."""
    return q.table == "supplier_reports" and q.op == "update" and "approved_by" in (q.payload or {})


@test("approve: a racing approval is a 409 and leaves the draft; a failed approval puts the superseded report back, even when that hits the index")
def _():
    from fastapi.testclient import TestClient
    import app.main as m
    fake = _world()
    with _env(fake, CATALOG):
        a = _build()["report"]["id"]
        fake.tables["supplier_reports"].append(dict(_row(fake, "supplier_reports", a), id=999))   # a second draft

        def other_lands(_q):        # after this approval read the approved reports, another approval lands
            _row(fake, "supplier_reports", 999)["status"] = "approved"
        with _Hook(fake, _approving, other_lands):
            r = TestClient(m.app, raise_server_exceptions=False).post(f"/market-intel/supplier-reports/{a}/approve",
                                                                     headers=R._h("boss"))
        assert r.status_code == 409 and "went through first" in r.json()["detail"], (r.status_code, r.text[:200])
        assert {x["id"]: x["status"] for x in fake.rows("supplier_reports")} == {a: "draft", 999: "approved"}
    fake = _world()
    with _env(fake, CATALOG):
        sr = _sr()
        old = _approved(fake)
        new = _build()["report"]["id"]

        def reset(_q):
            raise RuntimeError("connection reset")
        with _Hook(fake, _approving, reset):
            try:
                sr.approve(new, "boss@example.com", "admin")
                raise AssertionError("the failed write was swallowed")
            except RuntimeError as e:
                assert "connection reset" in str(e)
        assert {x["id"]: x["status"] for x in fake.rows("supplier_reports")} == {old: "approved", new: "draft"}
        # the restore itself hits the index: a third report was approved meanwhile; it stands, the old stays superseded
        fake.tables["supplier_reports"].append(dict(_row(fake, "supplier_reports", new), id=777))

        def third_lands(_q):
            _row(fake, "supplier_reports", 777)["status"] = "approved"
        with _Hook(fake, _approving, third_lands):
            _raises(lambda: sr.approve(new, "boss@example.com", "admin"), "went through first", 409)
        assert {x["id"]: x["status"] for x in fake.rows("supplier_reports")} == \
            {old: "superseded", new: "draft", 777: "approved"}


# ═══════════════════════════════════════════════════════════════════════════════
# 4. the frozen workbook
# ═══════════════════════════════════════════════════════════════════════════════

@test("download: an approved report only; the .xlsx has the supplier edition's five sheets and their header rows")
def _():
    fake = _world()
    with _env(fake, CATALOG):
        sr = _sr()
        rid = _build()["report"]["id"]
        _raises(lambda: _download(rid), "only an approved report", 409)
        sr.approve(rid, "boss@example.com", "admin")
        got = _download(rid)
    assert got["filename"] == "YQ-supplier-report-SUPX-2026-09-supplier.xlsx"
    wb = _book(got["data"])
    assert wb.sheetnames == ["Summary", "Product issues", "Packaging", "Evidence", "Corrective actions"]
    assert _head(wb["Product issues"]) == ["ID", "Date", "SKU", "Product", "Issue", "Finding", "Description", "Qty",
                                           "Batch / packaging", "Basis", "Evidence level", "Status", "Photo",
                                           "Evidence refs"]
    assert _head(wb["Packaging"]) == ["ID", "Date", "SKU", "Product", "Issue", "Finding", "Description", "Qty",
                                      "Batch / packaging", "Packaging version", "Basis", "Evidence level", "Status",
                                      "Photo", "Evidence refs"]
    assert _head(wb["Evidence"]) == ["Photo ref", "Report ID", "SKU", "Issue", "Role", "Caption", "SHA-256 (short)",
                                     "Photo"]
    assert _head(wb["Corrective actions"]) == ["Action", "Request", "Owner", "Due", "Supplier response", "YQ check",
                                               "Status", "Days open", "Linked reports"]
    ws = wb["Product issues"]
    ids = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    assert ids == [f"MI-{i}" for i in (101, 102, 103, 104, 131, 141, 151, 161, 162, 121)], ids
    last = [c.value for c in ws[ws.max_row]]
    assert last[5] == "Insufficient evidence (not counted)" and last[12] == "no photo"
    first = {h: v for h, v in zip(_head(ws), [c.value for c in ws[2]])}
    assert first["Description"] == ("[shop] owner says field team sold it; call [number]; paid [amount]. "
                                    "[another supplier] makes a better one."), first["Description"]
    assert (first["Finding"], first["Evidence level"], first["Evidence refs"], first["Qty"]) == \
        ("Verified", "Inspected & verified by YQ", "PH-1", 2)
    second = {h: v for h, v in zip(_head(ws), [c.value for c in ws[3]])}
    assert second["Description"] == ("[shop] returned 3 pieces; call [number] or [number]; we sell it at [amount], "
                                     "priced at [amount], costs [amount], MRP [amount]; [another brand] lasts longer "
                                     "and [another brand] too"), second["Description"]
    row131 = {h: v for h, v in zip(_head(ws), [c.value for c in ws[6]])}
    assert row131["ID"] == "MI-131" and row131["Photo"] == "photo not available"
    assert [c.value for c in wb["Packaging"][2]][9] == "v2"
    actions = [[c.value for c in r] for r in wb["Corrective actions"].iter_rows(min_row=2)]
    assert [r[0] for r in actions] == ["CA-1", "CA-3", "CA-5"], "never the cancelled CA-6 in a supplier copy"
    assert actions[0][7] == 26, "open since 18 Sep, as of the approval (14 Oct)"
    assert actions[1][4] == "We will check the ink. Write to [email]" and actions[2][7] == 23
    summary = _cells(wb)
    assert "Supplier X: product feedback, September 2026" in summary and "Supplier X: September 2026." in summary


@test("download: thumbnails and 640 px JPEGs are embedded; a photo that could not be read at approval reads 'photo not available'; the thumbnail stands in for a lost master")
def _():
    fake = _world()
    with _env(fake, CATALOG):
        rid = _approved(fake)
        data = _download(rid)["data"]
    wb = _book(data)
    ev = wb["Evidence"]
    refs = [ev.cell(row=r, column=1).value for r in range(2, ev.max_row + 1)]
    assert refs == ["PH-3", "PH-1", "PH-5", "PH-6", "PH-7", "PH-8"], refs
    by_ref = {ev.cell(row=r, column=1).value: [c.value for c in ev[r]] for r in range(2, ev.max_row + 1)}
    assert by_ref["PH-6"][7] == "photo not available"
    assert by_ref["PH-5"][6] == "abababababab" and by_ref["PH-7"][6] == "-", "a stored sha256; none from a thumbnail"
    master_1 = fake.storage.objects[("finds", next(p["path"] for p in fake.rows("market_photos") if p["id"] == 1))][0]
    assert by_ref["PH-1"][6] == hashlib.sha256(master_1).hexdigest()[:12]
    assert by_ref["PH-1"][5] == "Cable end burnt, seen at [shop]"
    pics = _colours(ev)
    assert len(pics) == 5 and all(fmt == "JPEG" for fmt, _s, _c in pics)
    assert sorted(max(size) for _f, size, _c in pics) == [320, 640, 640, 640, 640], "640 px; the thumbnail-only one 320"
    thumbs = _colours(wb["Product issues"]) + _colours(wb["Packaging"])
    assert len(thumbs) == 5 and all(max(size) <= 160 for _f, size, _c in thumbs)
    assert any(_near(col, (0, 150, 150)) for _f, _s, col in pics), "PH-7 came from its 320 px thumbnail"


@test("privacy: every cell of the supplier edition is free of shop, customer and rep names, contacts (dashes too), amounts, BHD, prices, other suppliers' ids, names and brands, and people / receipt photos")
def _():
    fake = _world()
    with _env(fake, CATALOG):
        rid = _approved(fake)
        data = _download(rid)["data"]
    wb = _book(data)
    text = "\n".join(_cells(wb))
    forbidden = ["Shop Alpha", "Shop Beta", "Shop Gamma", "Shop Delta", "Shop Epsilon", "Shop Zeta", "Owner Alpha",
                 "Focus Trader Kappa", "Rep One", "Rep Two", "R. One Old", "rep1@example.com", "rep2",
                 "qa@supplier-x.example", "3900", "1234", "2.500", "3.250", "1.900", "4.000", "1.250", "2.750",
                 "Supplier Y", "SUPY", "Brand Y", "BrandY", "Ylite", "Rival", "Supplier Old", "Withdrawn"]
    for word in forbidden:
        assert word.lower() not in text.lower(), word
    assert not re.search(r"(?i)\bbhd\b|\bbd\b|\bdinar", text), "no currency anywhere"
    headers = [str(c.value) for ws in wb.worksheets if ws.title != "Summary" for c in ws[1] if c.value]
    assert not any(re.search(r"(?i)price|cost|margin|amount|value", h) for h in headers), headers
    assert not re.search(r"(?i)\b(?:price[ds]?|costs?|margins?|mrp|rrp|sell(?:s|ing)?|paid)\b\W*(?:at|for)?\W*\d", text), \
        re.findall(r"(?i).{20}\b(?:price|cost|margin|mrp|sell|paid).{20}", text)
    assert _ids(text, "MI") <= X_FINDINGS and not (_ids(text, "MI") & (NEVER_IN_X | X_MARKET))
    assert _ids(text, "CA") == {1, 3, 5} and _ids(text, "PH") == {1, 3, 5, 6, 7, 8}
    for ws in wb.worksheets:
        for _f, _s, col in _colours(ws):
            assert not _near(col, RED), f"a people-flagged photo on {ws.title}"
            assert not _near(col, YELLOW), f"a receipt photo on {ws.title}"
    assert "[shop]" in text and "field team" in text and "[number]" in text and "[amount]" in text
    assert "[another brand]" in text and wb.properties.creator == "YQ Bahrain"


@test("privacy: a supplier whose name is part of a shop's or another supplier's keeps its name; the longer names are masked end to end")
def _():
    suppliers = [dict(s, name="Royal") if s["id"] == 1 else dict(s, name="Royal Gulf Trading") if s["id"] == 2 else s
                 for s in SUPPLIERS]
    shops = SHOPS + [{"id": 6, "phone": "97300000006", "name": "Owner Six", "shop": "Royal Mobile"}]
    obs = [dict(o, note="Royal Mobile owner said Royal Gulf Trading is cheaper; Royal ships fast")
           if o["id"] == 103 else o for o in OBS]
    fake = _world(suppliers=suppliers, shop_customers=shops, market_observations=obs)
    with _env(fake, CATALOG):
        rid = _approved(fake)
        wb = _book(_download(rid)["data"])
    text = "\n".join(_cells(wb))
    assert "[shop] owner said [another supplier] is cheaper; Royal ships fast" in text
    assert "Royal Mobile" not in text and "Gulf" not in text
    assert "Royal: product feedback, September 2026" in text


@test("download: the internal edition adds Competitor and Demand (a competitor's price allowed there), still no names, never a people photo")
def _():
    fake = _world()
    with _env(fake, CATALOG):
        rid = _approved(fake, "internal")
        got = _download(rid)
    wb = _book(got["data"])
    assert got["filename"].endswith("-2026-09-internal.xlsx")
    assert wb.sheetnames == ["Summary", "Product issues", "Packaging", "Evidence", "Corrective actions", "Competitor",
                             "Demand"]
    assert _head(wb["Competitor"]) == ["ID", "Date", "Product seen", "Brand", "Competitor", "Category",
                                       "Price seen (BHD)", "YQ SKU", "Note", "Area", "Photo"]
    assert _head(wb["Demand"]) == ["ID", "Date", "Source", "Asked for", "YQ SKU", "Qty asked", "Demand level", "Note",
                                   "Area"]
    comp = [c.value for c in wb["Competitor"][2]]
    assert comp[0] == "MI-171" and comp[6] == "1.250" and comp[8] == "[shop] stocks it"
    assert [c.value for c in wb["Demand"][2]][:6] == ["MI-181", "2026-09-26", "Field team", "Asked for C10 in white",
                                                      None, 20]
    text = "\n".join(_cells(wb))
    for word in ("Shop Alpha", "Shop Beta", "Rep One", "Supplier Y", "3900"):
        assert word not in text, word
    assert "CA-6" in text, "the internal edition keeps a cancelled action"
    colours = [col for ws in wb.worksheets for _f, _s, col in _colours(ws)]
    assert any(_near(c, YELLOW) for c in colours), "the receipt is in the internal edition"
    assert not any(_near(c, RED) for c in colours), "a people photo never is"


@test("frozen: what arrives after approval (a weak sighting, a rep's answer photo, a case verified) never reaches the file; the bytes stay the same")
def _():
    fake = _world()
    with _env(fake, CATALOG):
        rid = _approved(fake)
        first = _download(rid)
        fake.tables["market_observations"].append(_o(999, "2026-09-27T08:00:00+00:00", yq_item_code="P22",
                                                     issue_type_id=103, supplier_id=1, note="Late note after approval"))
        data, ext = _picture(YELLOW)
        fake.storage.objects[("finds", f"market/2026/10/answer{ext}")] = (data, {})
        fake.tables["market_photos"].append({"id": 77, "observation_id": 103, "bucket": "finds",
                                             "path": f"market/2026/10/answer{ext}", "thumb_path": None, "position": 9,
                                             "role": "product", "has_people_flag": False, "sha256": None,
                                             "caption": "answer to the office's question", "blurred_path": None})
        next(c for c in fake.tables["market_cases"] if c["id"] == 8).update(VERIFIED)
        again = _download(rid)
        preview = _sr().get_report(rid, "member")
    assert again["data"] == first["data"] and again["sha256"] == first["sha256"]
    text = "\n".join(_cells(_book(again["data"])))
    assert "MI-999" not in text and "Late note" not in text and "PH-77" not in text
    assert preview["frozen"] and 999 not in {r["id"] for k in preview["rows"] for r in preview["rows"][k]}
    assert next(r for r in preview["rows"]["product_issues"] if r["id"] == 141)["finding"] == "Unverified report"


@test("frozen: the download is refused (nothing served, nothing logged) when a frozen photo is flagged with people, re-classed as a receipt, deleted or changed in storage, or a row changed supplier or was withdrawn")
def _():
    def photo(fake, pid):
        return next(p for p in fake.tables["market_photos"] if p["id"] == pid)
    cases = (("PH-1 is now flagged with people", 409, lambda f: photo(f, 1).update(has_people_flag=True)),
             ("PH-3 is now a receipt photo", 409, lambda f: photo(f, 3).update(role="receipt")),
             ("PH-8 has been deleted", 409, lambda f: f.tables["market_photos"].remove(photo(f, 8))),
             ("MI-151 is no longer Supplier X's alone", 409,
              lambda f: next(c for c in f.tables["market_cases"] if c["id"] == 11).update(supplier_id=2)),
             ("MI-111 is no longer a finding", 409,
              lambda f: next(c for c in f.tables["market_cases"] if c["id"] == 2).update(status="not_substantiated")),
             ("PH-8 changed in storage", 409,
              lambda f: f.storage.objects.__setitem__(("finds", photo(f, 8)["path"]), (_picture((1, 2, 3))[0], {}))),
             ("PH-8 could not be read", 503, lambda f: f.storage.objects.pop(("finds", photo(f, 8)["path"]))))
    for want, status, change in cases:
        fake = _world()
        with _env(fake, CATALOG):
            rid = _approved(fake)
            change(fake)
            _raises(lambda r=rid: _download(r), want, status)
            assert not _row(fake, "supplier_reports", rid).get("downloads"), want
            if status == 409 and not want.startswith("PH-8 changed"):
                assert any(want in w for w in _sr().get_report(rid, "member")["warnings"]), want


@test("download: logged on the row ({at, by, sha256}) and in audit_log; the same manifest gives the same bytes; a download that cannot be logged is never served")
def _():
    fake = _world()
    with _env(fake, CATALOG):
        sr = _sr()
        rid = _approved(fake)
        c = R._client()
        r1 = c.get(f"/market-intel/supplier-reports/{rid}/download", headers=R._h("clerk"))
        assert r1.status_code == 200, r1.text[:200]
        assert r1.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml")
        assert 'filename="YQ-supplier-report-SUPX-2026-09-supplier.xlsx"' in r1.headers["content-disposition"]
        sha = hashlib.sha256(r1.content).hexdigest()
        row = _row(fake, "supplier_reports", rid)
        assert row["file_sha256"] == sha
        assert row["downloads"] == [{"at": NOW.isoformat(), "by": "clerk@example.com", "sha256": sha}]
        audit = [a for a in fake.rows("audit_log") if a["event"] == "market_intel.supplier_report.download"]
        assert len(audit) == 1 and audit[0]["detail"]["sha256"] == sha and audit[0]["user_email"] == "clerk@example.com"
        r2 = c.get(f"/market-intel/supplier-reports/{rid}/download", headers=R._h("mgmt"))
        assert r2.status_code == 200 and r2.content == r1.content, "the same frozen report, the same file"
        assert [d["by"] for d in _row(fake, "supplier_reports", rid)["downloads"]] == ["clerk@example.com", "mgmt@example.com"]
        preview = c.get(f"/market-intel/supplier-reports/{rid}", headers=R._h("clerk")).json()["report"]
        assert preview["file_sha256"] == sha and [d["by"] for d in preview["downloads"]] == ["clerk", "mgmt"]
        fake.fail[("update", "supplier_reports")] = RuntimeError("database unavailable")
        try:
            sr.download(rid, "boss@example.com", "admin")
            raise AssertionError("served without a log")
        except RuntimeError as e:
            assert "unavailable" in str(e)
        assert len(_row(fake, "supplier_reports", rid)["downloads"]) == 2


@test("download: a superseded report and a draft are refused; a control character or a formula-like note never breaks or arms the file")
def _():
    fake = _world()
    with _env(fake, CATALOG):
        old = _approved(fake)
        new = _approved(fake)
        _raises(lambda: _download(old), "superseded", 409)
        assert _download(new)["bytes"] > 0
    types = [dict(t, label_en="Not\x07 charging") if t["id"] == 100 else t for t in TYPES]
    obs = [dict(o, note='=HYPERLINK("http://example.invalid","x")') if o["id"] == 102 else o for o in OBS]
    fake = _world(market_issue_types=types, market_observations=obs)
    with _env(fake, CATALOG):
        rid = _approved(fake)
        wb = _book(_download(rid)["data"])
    ws = wb["Product issues"]
    row = {h: c for h, c in zip(_head(ws), ws[3])}
    assert row["ID"].value == "MI-102" and row["Issue"].value == "Not charging"
    assert row["Description"].value.startswith("=HYPERLINK") and row["Description"].data_type == "s", "text, not a formula"


# ═══════════════════════════════════════════════════════════════════════════════
# 5. the routes: roles, shapes, before the migration, a wrong column
# ═══════════════════════════════════════════════════════════════════════════════

@test("routes: every /market-intel/supplier-reports route is behind feature:Market Intel and nothing else")
def _():
    from fastapi.routing import APIRoute
    import app.main as m
    from app import auth
    found = {}
    for r in m.app.routes:
        if isinstance(r, APIRoute) and r.path.startswith("/market-intel/supplier-reports"):
            for meth in r.methods:
                gates = []
                for dep in r.dependant.dependencies:
                    c = dep.call
                    gates.append("user" if c is auth.get_current_user else
                                 f"feature:{c.feature}" if getattr(c, "feature", None) else f"?{c}")
                found[(meth, r.path)] = gates
    want = {("POST", "/market-intel/supplier-reports"), ("GET", "/market-intel/supplier-reports"),
            ("GET", "/market-intel/supplier-reports/{report_id}"), ("PATCH", "/market-intel/supplier-reports/{report_id}"),
            ("POST", "/market-intel/supplier-reports/{report_id}/approve"),
            ("GET", "/market-intel/supplier-reports/{report_id}/download")}
    assert set(found) == want, set(found) ^ want
    assert all(g == ["feature:Market Intel"] for g in found.values()), found
    assert _sr().APPROVE_ROUTE == "/market-intel/supplier-reports/{report_id}/approve"
    assert ("POST", _sr().APPROVE_ROUTE) in auth.READ_ONLY_WRITE_ALLOWLIST, "management approves through the gate"


@test("roles: reps and a storekeeper holding the page get 403 everywhere; a reviewer cannot approve; nothing is written")
def _():
    fake = _world()
    with _env(fake, CATALOG):
        rid = _approved(fake)
        draft = _build()["report"]["id"]
        c = R._client()
        before = len(fake.writes())
        for tok in ("rep1", "store", "rep3"):
            calls = [c.post("/market-intel/supplier-reports", headers=R._h(tok), json={"supplier_id": 1, "month": SEP}),
                     c.get("/market-intel/supplier-reports", headers=R._h(tok)),
                     c.get(f"/market-intel/supplier-reports/{rid}", headers=R._h(tok)),
                     c.patch(f"/market-intel/supplier-reports/{rid}", headers=R._h(tok), json={"narrative": "x"}),
                     c.post(f"/market-intel/supplier-reports/{rid}/approve", headers=R._h(tok)),
                     c.get(f"/market-intel/supplier-reports/{rid}/download", headers=R._h(tok))]
            assert all(r.status_code == 403 for r in calls), (tok, [(r.status_code, r.text[:80]) for r in calls])
            if tok == "rep3":
                assert "Market Intel" in calls[1].text
        r = c.post(f"/market-intel/supplier-reports/{draft}/approve", headers=R._h("clerk"))
        assert r.status_code == 403 and "Only management" in r.json()["detail"], r.text[:200]
        assert c.get("/market-intel/supplier-reports", headers=R._h("clerk")).json()["capabilities"]["can_approve"] is False
        assert len(fake.writes()) == before, fake.writes()[before:]


@test("roles: management reads (a hand-edited narrative masked), downloads and approves; its building and editing are refused centrally (read_only)")
def _():
    fake = _world()
    with _env(fake, CATALOG):
        c = R._client()
        rid = c.post("/market-intel/supplier-reports", headers=R._h("clerk"),
                     json={"supplier_id": 1, "month": SEP, "edition": "supplier"}).json()["report"]["id"]
        typed = "Supplier X: September 2026.\nNew: Shop Alpha rang 3900-1234 about C10 [MI-101]."
        assert c.patch(f"/market-intel/supplier-reports/{rid}", headers=R._h("clerk"),
                       json={"narrative": typed}).status_code == 200
        lst = c.get("/market-intel/supplier-reports", headers=R._h("mgmt"))
        assert lst.status_code == 200 and lst.json()["count"] == 1
        caps = lst.json()["capabilities"]
        assert caps == {"can_build": False, "can_edit": False, "can_approve": True, "can_download": True}, caps
        one = c.get(f"/market-intel/supplier-reports/{rid}", headers=R._h("mgmt"))
        assert one.status_code == 200 and "Shop Alpha" not in one.text and "3900" not in one.text
        assert one.json()["report"]["narrative"] == "Supplier X: September 2026.\nNew: [shop] rang [number] about C10 [MI-101]."
        assert c.get(f"/market-intel/supplier-reports/{rid}", headers=R._h("clerk")).json()["report"]["narrative"] == typed
        for r in (c.post("/market-intel/supplier-reports", headers=R._h("mgmt"), json={"supplier_id": 1, "month": SEP}),
                  c.patch(f"/market-intel/supplier-reports/{rid}", headers=R._h("mgmt"), json={"narrative": "x"})):
            assert r.status_code == 403 and r.json()["detail"]["code"] == "read_only", r.text[:200]
        r = c.post(f"/market-intel/supplier-reports/{rid}/approve", headers=R._h("mgmt"))
        assert r.status_code == 200 and r.json()["status"] == "approved", r.text[:200]
        frozen = c.get(f"/market-intel/supplier-reports/{rid}", headers=R._h("mgmt")).json()
        assert frozen["frozen"] and "Shop Alpha" not in json.dumps(frozen) and "3900" not in json.dumps(frozen)
        d = c.get(f"/market-intel/supplier-reports/{rid}/download", headers=R._h("mgmt"))
        assert d.status_code == 200 and d.content[:2] == b"PK"
        assert [a["event"] for a in fake.rows("audit_log") if a["event"].startswith("market_intel.supplier_report")] == \
            ["market_intel.supplier_report.build", "market_intel.supplier_report.edit",
             "market_intel.supplier_report.approve", "market_intel.supplier_report.download"]


@test("routes: a full round trip - a reviewer builds, lists, previews and edits; an admin approves; anyone of the office downloads")
def _():
    fake = _world()
    with _env(fake, CATALOG):
        c = R._client()
        b = c.post("/market-intel/supplier-reports", headers=R._h("clerk"), json={"supplier_id": 1, "month": SEP})
        assert b.status_code == 200, b.text[:300]
        body = b.json()
        assert body["capabilities"] == {"can_build": True, "can_edit": True, "can_approve": False, "can_download": True}
        rid = body["report"]["id"]
        assert set(body) >= {"report", "rows", "groups", "actions", "live_counts", "warnings", "rebuilt", "frozen"}
        row = next(r for r in body["rows"]["product_issues"] if r["id"] == 101)
        assert row["photos"][0]["thumb_url"].startswith("https://signed.example/finds/market/2026/09/p1-t320")
        assert row["photos"][0]["caption"] == "Cable end burnt, seen at [shop]" and row["photos_withheld"] == 1
        assert not any(k.startswith("_") for r in body["rows"]["product_issues"] for k in r)
        assert c.post("/market-intel/supplier-reports", headers=R._h("clerk"),
                      json={"supplier_id": 1, "month": SEP, "edition": "internal"}).status_code == 200
        lst = c.get("/market-intel/supplier-reports?supplier_id=1&month=2026-09", headers=R._h("clerk")).json()
        assert lst["count"] == 2 and {r["edition"] for r in lst["rows"]} == {"supplier", "internal"}
        assert all("narrative" not in r and "manifest" not in r["counts"] for r in lst["rows"])
        assert c.get("/market-intel/supplier-reports?edition=internal", headers=R._h("clerk")).json()["count"] == 1
        page = c.get("/market-intel/supplier-reports?limit=1&offset=1", headers=R._h("clerk")).json()
        assert page["count"] == 2 and len(page["rows"]) == 1
        assert c.get("/market-intel/supplier-reports?edition=public", headers=R._h("clerk")).status_code == 400
        p = c.patch(f"/market-intel/supplier-reports/{rid}", headers=R._h("clerk"),
                    json={"included_observation_ids": sorted(X_COUNTED - {161, 162}), "withheld_photo_ids": [8]})
        assert p.status_code == 200 and p.json()["report"]["counts"]["reports"] == 9
        bad = c.patch(f"/market-intel/supplier-reports/{rid}", headers=R._h("clerk"), json={"narrative": "[MI-999]"})
        assert bad.status_code == 400 and "MI-999" in bad.json()["detail"]
        assert c.get("/market-intel/supplier-reports/999", headers=R._h("clerk")).status_code == 404
        assert c.get(f"/market-intel/supplier-reports/{rid}/download", headers=R._h("clerk")).status_code == 409
        a = c.post(f"/market-intel/supplier-reports/{rid}/approve", headers=R._h("boss"))
        assert a.status_code == 200 and a.json()["approved_by"] == "boss"
        listed = c.get("/market-intel/supplier-reports?status=approved", headers=R._h("clerk")).json()["rows"][0]
        assert listed["frozen"] is True and "manifest" not in listed["counts"]
        d = c.get(f"/market-intel/supplier-reports/{rid}/download", headers=R._h("clerk"))
        assert d.status_code == 200
        cells = "\n".join(_cells(_book(d.content)))
        assert "MI-161" not in cells and "MI-101" in cells and "PH-8" not in cells
        events = [a["event"] for a in fake.rows("audit_log")]
        assert events.count("market_intel.supplier_report.build") == 2 and "market_intel.supplier_report.edit" in events


@test("before the migration every route answers 503 with the hint (a rep still 403); a table missing half-way is 503 and forgets the probe; a column we do not have is a 500 naming it")
def _():
    sr = _sr()
    fake = _world()
    fake.missing = {"supplier_reports"}
    with _env(fake, CATALOG):
        c = R._client()
        calls = [c.post("/market-intel/supplier-reports", headers=R._h("clerk"), json={"supplier_id": 1, "month": SEP}),
                 c.get("/market-intel/supplier-reports", headers=R._h("clerk")),
                 c.get("/market-intel/supplier-reports/1", headers=R._h("mgmt")),
                 c.patch("/market-intel/supplier-reports/1", headers=R._h("clerk"), json={"narrative": "x"}),
                 c.post("/market-intel/supplier-reports/1/approve", headers=R._h("boss")),
                 c.get("/market-intel/supplier-reports/1/download", headers=R._h("clerk"))]
        for r in calls:
            assert r.status_code == 503 and r.json()["detail"] == sr.MISSING_HINT, (r.status_code, r.text[:200])
        assert c.get("/market-intel/supplier-reports", headers=R._h("rep1")).status_code == 403
        assert sr.available() is False
    fake = _world()
    fake.missing = {"market_cases"}
    with _env(fake, CATALOG):
        _raises(lambda: _build(), "not set up yet", 503)
        assert sr._probe["ok"] is False and sr._probe["until"] == 0.0
    fake = _world()
    fake.missing = {"customers"}                    # an optional name list: the report still builds
    with _env(fake, CATALOG):
        assert _build()["report"]["counts"]["reports"] == 11
    fake = _world()
    with _env(fake, CATALOG):
        fake.fail[("select", "brands")] = _err("42703", "column brands.supplier_ref does not exist")
        e = _raises(lambda: _build(), "brands.supplier_ref", 500)
        assert "not a missing migration" in str(e) and sr.available() is True, "a wrong column never reads as not set up"
        fake.fail[("select", "brands")] = _err("42703", 'column "supplier_ref" of relation "brands" does not exist')
        r = R._client().post("/market-intel/supplier-reports", headers=R._h("clerk"), json={"supplier_id": 1, "month": SEP})
        assert r.status_code == 500 and "brands.supplier_ref" in r.json()["detail"], r.text[:200]
    for msg, table in (('{"code": "42P01", "message": "relation \\"public.brands\\" does not exist"}', True),
                       ('{"code": "PGRST205", "message": "Could not find the table public.brands"}', True),
                       ('{"code": "42703", "message": "column \\"x\\" of relation \\"brands\\" does not exist"}', False),
                       ('{"code": "PGRST204", "message": "Could not find the \'x\' column of \'brands\'"}', False)):
        assert sr.missing_table(RuntimeError(msg)) is table, msg


# ═══════════════════════════════════════════════════════════════════════════════
# 6. pins
# ═══════════════════════════════════════════════════════════════════════════════

@test("pins: never e-mailed; registered right after Market Intel; CI runs this suite in the api job; supplier_action_cases is never ordered by id")
def _():
    for rel in ("app/supplier_reports.py", "app/supplier_reports_api.py"):
        src = (ROOT / rel).read_text(encoding="utf-8")
        for word in ("emailer", "send_html", "notify", "smtp", "resend"):
            assert word not in src.lower(), (rel, word)
    src = (ROOT / "app" / "supplier_reports.py").read_text(encoding="utf-8")
    assert src.count('"supplier_action_cases"') == 2 and src.count("order=LINK_ORDER") == 2
    main = (ROOT / "app" / "main.py").read_text(encoding="utf-8")
    mi_at = main.index("_register_market_intel(app, limiter)")
    sr_at = main.index("_register_supplier_reports(app, limiter)")
    assert mi_at < sr_at and main[mi_at:sr_at].count("\n") <= 3, "registered right after Market Intel (and its v2 modules)"
    ci = (ROOT / ".github" / "workflows" / "ci.yml").read_text(encoding="utf-8")
    api = ci.split("\n  web:\n", 1)[0]
    assert api.index("python -m tests.test_r7b_market_intel\n") < api.index("python -m tests.test_r8m_supplier_reports\n")
    step = api.split("python -m tests.test_r8m_supplier_reports", 1)[0].rsplit("- name:", 1)[1]
    assert "SUPABASE_URL: https://ci.invalid" in step


def main() -> int:
    passed = failed = 0
    for name, fn in TESTS:
        try:
            fn()
            print(f"PASS  {name}")
            passed += 1
        except Exception:  # noqa: BLE001
            print(f"FAIL  {name}")
            traceback.print_exc()
            failed += 1
    print(f"\n{passed} passed, {failed} failed")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
