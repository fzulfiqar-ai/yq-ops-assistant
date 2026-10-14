"""R8m — Market Intel v2 end to end: the four modules working together on one database (plan Part F,
build contract §2-§3).

    python -m tests.test_r8m_mi_e2e

The REAL FastAPI app (app.main.app, through the TestClient and the token stand-in the other suites use)
drives one story on ONE in-memory PostgREST + Storage stand-in that every module shares:
app/market_intel.py (capture, My reports, the board, the people photo), app/market_quality.py (lists,
product picker, cases, the workflow, question / answer, merge / move, supplier actions, the safety alert),
app/supplier_reports.py (the monthly workbook, through its documented routes only) and
scripts/market_intel_import.py (the old finds, a dry run).

The stand-in is tests/test_r7b_market_intel.py's (1,000 rows at most per select, range() offsets), held
here to the REAL schema read from scripts/r7b_market_intel_migration.sql and
scripts/r8m_market_quality_migration.sql: a column that is not there (in a select list, a filter, an
order or a write) answers 42703 / PGRST204 as PostgREST does; the CHECK lists (statuses, events, kinds,
roles, editions), the evidence and severity ranges, NOT NULL, the identity ids, the foreign keys and the
unique keys (one approved report per supplier, month and edition, and one safety-alert claim per case,
included) hold; and market_case_events refuses every UPDATE and DELETE, as its trigger does. Nothing reaches a database, a bucket, an e-mail
provider or the network: SUPABASE_URL, SUPABASE_KEY and DATABASE_URL are pinned to ci.invalid and every
e-mail provider is silenced BEFORE any import; the safety e-mail is a stub. Every shop, rep, supplier,
brand, product, phone and amount is invented.

The story (each step is a test; they run in order on the same data):
  1. lists — an admin creates two suppliers, their brands and three issue types through the routes;
  2. capture — two reps pick products (GET /market-intel/products) and post Spotted v2's own fields
     from three shops: the second report of a product + issue joins its case; a safety report e-mails
     once (its retry and a second sighting do not) and the e-mail names no shop or rep; a 'not sure /
     not ours' report, a wrong product, packaging, another supplier's product, a competitor price;
  3. the office — the case tabs, a case opened (the people-flagged photo only as its blurred copy),
     management masked, a question the rep reads in My reports and answers with a photo;
  4. merge and move — a duplicate case merged in, a misfiled sighting moved: the counts follow;
  5. the workflow — verified (method + finding), evidence only up, a supplier action, the response,
     YQ's check, resolved, closed; refused moves are 409 (a reopen by a non-admin 403);
  6. the supplier workbook — the draft's counts are the story's, every finding sentence cites real
     [MI-...] ids, an edit, management approves through the read-only gate, the download: no shop, rep,
     phone, e-mail, price or cost in any cell or part of the file, and exactly the right supplier's
     eligible photos embedded (told apart by dHash: never the people photo, the receipt, supplier Y's);
  7. a rep — 403 on cases, list edits, actions, photo flags and the workbook; My reports and the board
     name no other rep or shop;
  8. the importer — a dry run on the same data sees the r8m photo columns, plans the thumbnails and
     writes nothing (and an in-memory --apply on a copy of the data writes only rows the schema allows);
  9. a safety sighting moved onto another case takes that case onto the Safety tab (no second e-mail);
 10. two safety reports racing on one case: the database's one-claim index refuses the second claim, the
     API reads that as "already claimed" (the report is saved, no second e-mail);
 11. the case timeline was only ever appended to.
A mismatch met here that another builder owns is listed in KNOWN_GAPS: the read is let through,
recorded, and printed as a NOTE while it is still met.
"""
from __future__ import annotations

import ast
import copy
import hashlib
import io
import json
import os
import random
import re
import sys
import traceback
import uuid
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

# nothing may reach production: the project folder's .env (found by walking up from a worktree) points at it
os.environ["SUPABASE_URL"] = "https://ci.invalid"
os.environ["SUPABASE_KEY"] = "ci-key"
os.environ["DATABASE_URL"] = "postgresql://ci:ci@ci.invalid:5432/ci"
for _k in ("RESEND_API_KEY", "BREVO_API_KEY", "SMTP_USER", "SMTP_PASS", "ALERT_EMAIL_TO", "TELEGRAM_BOT_TOKEN",
           "APP_BASE_URL"):
    os.environ[_k] = ""

# The borrowed suite re-wraps sys.stdout at import; keep both wrappers alive so the earlier one is never
# garbage-collected (which would close the shared buffer).
_KEEP: list = [sys.stdout, sys.stderr]
from tests import test_r7b_market_intel as R  # noqa: E402
_KEEP += [sys.stdout, sys.stderr]

TESTS: list[tuple[str, object]] = []


def test(name):
    def deco(fn):
        TESTS.append((name, fn))
        return fn
    return deco


# ═══════════════════════════════════════════════════════════════════════════════
# the real schema, read from the two migrations
# ═══════════════════════════════════════════════════════════════════════════════

MIGRATIONS = (ROOT / "scripts" / "r7b_market_intel_migration.sql", ROOT / "scripts" / "r8m_market_quality_migration.sql")


def _top_level(body: str) -> list[str]:
    """A create table body cut at its top-level commas (one column or constraint each)."""
    out, depth, cur = [], 0, []
    for ch in body:
        depth += (ch == "(") - (ch == ")")
        if ch == "," and depth == 0:
            out.append("".join(cur))
            cur = []
        else:
            cur.append(ch)
    out.append("".join(cur))
    return out


def _column(d: str) -> dict:
    """One column definition: NOT NULL, a default (or identity), its CHECK list / range, its reference."""
    low = d.lower()
    spec = {"notnull": " not null" in low, "default": " default " in low or " generated " in low,
            "identity": "generated always as identity" in low, "in": None, "between": None, "ref": None}
    m = re.search(r"check \(\w+ in \(([^)]*)\)\)", d, re.I)
    if m:
        spec["in"] = tuple(re.findall(r"'([^']*)'", m.group(1)))
    m = re.search(r"check \(\w+ between (\d+) and (\d+)\)", d, re.I)
    if m:
        spec["between"] = (int(m.group(1)), int(m.group(2)))
    m = re.search(r"references (\w+)\((\w+)\)", d, re.I)
    if m:
        spec["ref"] = (m.group(1), m.group(2))
    return spec


def _read_schema(*files: Path) -> dict[str, dict[str, dict]]:
    """{table: {column: spec}} for every table these files CREATE, with the columns their ALTERs add."""
    out: dict[str, dict[str, dict]] = {}
    for f in files:
        sql = re.sub(r"--[^\n]*", "", f.read_text(encoding="utf-8"))
        for m in re.finditer(r"create table if not exists (\w+)\s*\((.*?)\n\);", sql, re.S | re.I):
            cols = out.setdefault(m.group(1), {})
            for part in _top_level(m.group(2)):
                d = " ".join(part.split())
                if d and not re.match(r"(constraint|primary key|unique)\b", d, re.I):
                    cols[d.split(" ", 1)[0]] = _column(d)
        for m in re.finditer(r"alter table (\w+)\s+(add column if not exists .*?);", sql, re.S | re.I):
            if m.group(1) in out:
                for part in re.split(r",?\s*add column if not exists\s+", m.group(2), flags=re.I):
                    d = " ".join(part.split()).rstrip(",")
                    if d:
                        out[m.group(1)][d.split(" ", 1)[0]] = _column(d)
    return out


SCHEMA = _read_schema(*MIGRATIONS)

# The unique keys of the two migrations: (name, key of a row, the partial index's WHERE).
UNIQUE = {
    "market_observations": [("market_observations_client_uuid_key", lambda r: r.get("client_uuid"), None)],
    "market_photos": [("market_photos_path_key", lambda r: (r.get("bucket"), r.get("path")), None)],
    "market_issue_types": [("market_issue_types_code_key", lambda r: r.get("code"), None)],
    "suppliers": [("suppliers_active_name_key", lambda r: str(r.get("name") or "").strip().lower(),
                   lambda r: r.get("active") is not False)],
    "supplier_action_cases": [("supplier_action_cases_pkey", lambda r: (r.get("action_id"), r.get("case_id")), None)],
    "supplier_reports": [("supplier_reports_one_approved_idx",
                          lambda r: (r.get("supplier_id"), r.get("month"), r.get("edition")),
                          lambda r: r.get("status") == "approved")],
    "market_case_events": [("market_case_events_one_alert_claim_idx", lambda r: r.get("case_id"),
                            lambda r: r.get("event") == "safety_alert" and (r.get("detail") or {}).get("email") == "queued")],
}

# A mismatch this suite meets that is another builder's to fix (reported, never hidden): the read is let
# through and recorded, and main() prints it as a NOTE while it is still met.
KNOWN_GAPS: dict = {}          # the supplier_action_cases order (no id column) was fixed in app/supplier_reports.py


def _pg(code: str, message: str) -> RuntimeError:
    """An error worded as PostgREST words it (the modules read the code out of the text)."""
    return RuntimeError(json.dumps({"code": code, "message": message}))


def _check(q: "_Query") -> None:
    """The request against the real schema, before the stand-in answers it."""
    spec = SCHEMA.get(q.table)
    if spec is None:                                  # salesmen, audit_log, a view ...: not these migrations'
        return
    db = q.db
    for _kind, col, _v in q.filters:
        if col not in spec:
            raise _pg("42703", f"column {q.table}.{col} does not exist")
    for col, _desc in q.orders:
        if col not in spec:
            gap = ("order", q.table, col)
            if gap in KNOWN_GAPS:
                db.known.add(gap)
                continue
            raise _pg("42703", f"column {q.table}.{col} does not exist")
    if q.op == "select":
        for c in str(q.cols or "*").split(","):
            if c.strip() not in ("", "*") and c.strip() not in spec:
                raise _pg("42703", f"column {q.table}.{c.strip()} does not exist")
        return
    if q.op == "delete":
        return
    for p in (q.payload if isinstance(q.payload, list) else [q.payload]):
        for k, v in p.items():
            col = spec.get(k)
            if col is None:
                raise _pg("PGRST204", f"Could not find the '{k}' column of '{q.table}' in the schema cache")
            if col["identity"]:
                raise _pg("428C9", f'cannot insert a non-DEFAULT value into column "{k}"')
            if v is None:
                if col["notnull"]:
                    raise _pg("23502", f'null value in column "{k}" of relation "{q.table}" violates not-null constraint')
                continue
            if col["in"] and v not in col["in"]:
                raise _pg("23514", f'new row for relation "{q.table}" violates check constraint ({k} = {v!r})')
            if col["between"] and not col["between"][0] <= int(v) <= col["between"][1]:
                raise _pg("23514", f'new row for relation "{q.table}" violates check constraint ({k} = {v!r})')
            ref = col["ref"]
            if ref and ref[0] in db.tables and not any(R._same(r.get(ref[1]), v) for r in db.tables[ref[0]]):
                raise _pg("23503", f'insert or update on table "{q.table}" violates foreign key constraint ({k} = {v!r})')
        if q.op != "update":
            for k, col in spec.items():
                if col["notnull"] and not col["default"] and p.get(k) is None:
                    raise _pg("23502", f'null value in column "{k}" of relation "{q.table}" violates not-null constraint')


def _clash(table: str, rows: list[dict]) -> str | None:
    for name, key, where in UNIQUE[table]:
        seen: set = set()
        for r in rows:
            if where is not None and not where(r):
                continue
            k = key(r)
            if k is None or k == "" or (isinstance(k, tuple) and None in k):
                continue
            if k in seen:
                return name
            seen.add(k)
    return None


class _Query(R._Query):
    """r7b's query stand-in, held to the real schema (see the module doc) and to the append-only timeline."""

    def __init__(self, db, table):
        super().__init__(db, table)
        self.cols = "*"

    def select(self, cols="*", count=None, head=False):
        self.cols = cols
        return super().select(cols, count=count, head=head)

    def execute(self):
        db = self.db
        if self.table == "market_case_events":
            if self.op in ("update", "upsert", "delete"):
                db.event_rewrites.append((self.op, self.payload, list(self.filters)))
                raise _pg("23001", f"market_case_events is append-only (a row cannot be changed: {self.op.upper()})")
            if self.op == "select" and db.claim_race and ("eq", "event", "safety_alert") in self.filters:
                db.calls.append(("select", self.table, None, list(self.filters)))
                return SimpleNamespace(data=[], count=0)      # the other phone's claim landed after this read
        _check(self)
        rows = db.tables.setdefault(self.table, []) if self.table in UNIQUE else None
        before = [dict(r) for r in rows] if rows is not None and self.op in ("insert", "upsert", "update") else None
        res = super().execute()
        if before is not None:
            clash = _clash(self.table, rows)
            if clash:
                rows[:] = before
                raise _pg("23505", f"duplicate key value violates unique constraint {clash}")
        if self.table == "market_case_events" and self.op == "insert":
            n = len(self.payload) if isinstance(self.payload, list) else 1
            db.event_copies += copy.deepcopy(db.tables["market_case_events"][-n:])
        return res


class _DB(R._FakeDB):
    def __init__(self, tables: dict):
        super().__init__(tables)
        self.event_rewrites: list[tuple] = []          # every UPDATE / DELETE tried on market_case_events
        self.event_copies: list[dict] = []             # each timeline row as it was written
        self.known: set[tuple] = set()                 # KNOWN_GAPS met
        self.claim_race = False                        # True: a safety-claim read misses a claim another phone made
        self.clock = NOW                               # created_at for rows a module leaves to the default

    def table(self, name):
        return _Query(self, name)


# ═══════════════════════════════════════════════════════════════════════════════
# the invented world
# ═══════════════════════════════════════════════════════════════════════════════

NOW = datetime(2026, 10, 14, 9, 0, tzinfo=timezone.utc)
APPROVE_AT = datetime(2026, 11, 8, 6, 0, tzinfo=timezone.utc)   # October's approval opens 7 days after it ends
MONTH = "2026-10"
SHOP_A = {"id": 11, "name": "Shop Alpha", "phone": "97330000011", "area": "Manama"}
SHOP_B = {"id": 12, "name": "Shop Beta", "phone": "97330000012", "area": "Muharraq"}
SHOP_G = {"id": 13, "name": "Shop Gamma", "phone": "97330000013", "area": "Riffa"}
SHOPS = [{"id": s["id"], "phone": s["phone"], "name": s["name"].replace("Shop", "Owner"), "shop": s["name"],
          "area": s["area"]} for s in (SHOP_A, SHOP_B, SHOP_G)]
# what the catalog context carries beside the names: YQ's own selling price and cost (never in a workbook)
PRICE, COST = "1.900", "0.950"


def _catalog() -> dict:
    rows = (("C-10", "C-10 Type-C Cable 1m", "CABLE", "Brand X"), ("C-11", "C-11 Lightning Cable 1m", "CABLE", "Brand X"),
            ("P-22", "P-22 Power Bank 10000", "POWER BANK", "Brand X"), ("Y-30", "Y-30 Wall Charger 20W", "CHARGER", "Brand Y"))
    return {"items": {c: {"item_code": c, "display_name": n, "category": cat, "brand": b, "price": PRICE, "cost": COST,
                          "product_image_url": f"https://img.invalid/{c}.jpg", "package_image_url": None}
                      for c, n, cat, b in rows}, "order": [r[0] for r in rows]}


NOTE_1 = "Owner Alpha says it stops charging after a week; Rep One sold it. Call +973 3900 1234 or owner@shop.example, paid 2.500 BD"
NOTE_2 = "Two more dead in the box at Shop Beta"
NOTE_S = "Customer on 3900 1234 says it swelled while charging"
QUESTION = "Which batch code is printed on the box label?"


def _photo(seed: int, size=(960, 720)) -> bytes:
    """A synthetic shelf photo: coloured blocks placed by `seed` (another seed, another picture to the dHash)."""
    from PIL import Image, ImageDraw
    rnd = random.Random(seed)
    w, h = size
    img = Image.new("RGB", size, (rnd.randrange(170, 250), rnd.randrange(170, 250), rnd.randrange(170, 250)))
    d = ImageDraw.Draw(img)
    for _ in range(10):
        fx, fy = rnd.random() * 0.8, rnd.random() * 0.8
        fw, fh = 0.1 + rnd.random() * 0.35, 0.1 + rnd.random() * 0.35
        shade = rnd.randrange(0, 160)
        d.rectangle([fx * w, fy * h, (fx + fw) * w, (fy + fh) * h], fill=(shade, (shade * 3) % 200, 255 - shade))
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=88)
    return buf.getvalue()


# 1a/1b/1c rep one's report (1c: his answer), 2a/2r rep two's (2a gets people-flagged, 2r is a receipt),
# S1 the safety report, B1 the unidentified one, P1 packaging, Y1 supplier Y's, K1 a competitor's,
# I1/I2 two old finds for the importer
PHOTOS = {k: _photo(101 + i) for i, k in enumerate(("1a", "1b", "1c", "2a", "2r", "S1", "B1", "P1", "Y1", "K1", "I1", "I2"))}


class _Mail:
    """Stands in for app.emailer.send_html: records each call."""

    def __init__(self):
        self.calls: list[dict] = []

    def __call__(self, subject, html, to=None):
        self.calls.append({"subject": subject, "html": html, "to": to})
        return {"emailed": True, "via": "stub"}


W = SimpleNamespace(fake=None)


def _new_world() -> None:
    tables = {"user_roles": R.USERS, "salesmen": R.SALESMEN, "shop_customers": SHOPS, "audit_log": [],
              "app_settings": [], "market_items": [], "market_observations": [], "market_photos": [],
              "market_item_decisions": [], "product_finds": [], "field_notes": [], "shop_events": [],
              "shop_restock_requests": [],
              "shop_orders": [{"id": 1, "customer_id": 11, "status": "delivered",
                               "created_at": (NOW - timedelta(days=10)).isoformat()}],
              "shop_order_lines": [{"id": 1, "order_id": 1, "item_code": "C-10"}],
              "suppliers": [], "brands": [], "market_issue_types": [], "market_cases": [], "market_case_events": [],
              "supplier_actions": [], "supplier_action_cases": [], "supplier_reports": []}
    W.fake = _DB(tables)
    W.mail = _Mail()
    W.catalog = _catalog()
    W.obs, W.case, W.photo, W.hash = {}, {}, {}, {}


class _Env(R._env):
    """r7b's environment (the fake behind every client, token = the e-mail's local part, the limiter off,
    the catalog from memory) plus the workbook's client, ONE clock for every module, the e-mail stub and
    an empty at-shop cache."""

    def __init__(self):
        super().__init__(W.fake, W.catalog)

    def __enter__(self):
        fake = super().__enter__()
        from app import emailer, market_intel as mi, market_quality as mq, supplier_reports as sr
        self.extra = R._Patched((emailer, "send_html", W.mail), (mi, "_now", lambda: NOW),
                                (sr, "get_client", lambda: fake), (sr, "_now", lambda: NOW))
        self.extra.__enter__()
        mq._at_shop.clear()
        sr.forget_probe()
        return fake

    def __exit__(self, *exc):
        from app import supplier_reports as sr
        self.extra.__exit__(*exc)
        sr.forget_probe()
        return super().__exit__(*exc)


def _ok(r, status: int = 200):
    assert r.status_code == status, (str(r.request.url), r.status_code, r.text[:400])
    return r.json() if status == 200 else r


def _pick(c, tok: str, q: str, shop: dict) -> dict:
    """The product picker as the sheet calls it (lib/marketIntel.ts fetchProducts: the shop as hints)."""
    rows = _ok(c.get("/market-intel/products", headers=R._h(tok),
                     params={"q": q, "shop_phone": shop["phone"], "shop_name": shop["name"], "limit": 20}))
    assert rows, q
    return rows[0]


def _fields(shop: dict, at: datetime, *, kind: str = "complaint", product: dict | None = None, title: str = "",
            issue: dict | None = None, basis: str = "", qty: str = "", batch: str = "", pack: str = "",
            safety: bool = False, roles: str = "", note: str = "", price: str = "", brand: str = "") -> dict:
    """The fields Spotted v2 posts (web/src/components/MarketCapture.tsx save()); '' is left out."""
    brand_text = (product or {}).get("brand") or brand
    linked = next((b for b in W.lists["brands"] if brand_text and b["name"].lower() == brand_text.lower()
                   and b["active"] and not b.get("merged_into")), None)
    shows = (issue or {}).get("fields") or []
    q = issue is not None
    return {"note": note, "kind": kind, "title": (product or {}).get("name") or title,
            "yq_item_code": (product or {}).get("code") or "", "brand": brand_text,
            "brand_id": str(linked["id"]) if linked else "",
            "supplier_id": str(linked["supplier_id"]) if linked and linked.get("supplier_id") else "",
            "price_bhd": price, "issue_type_id": str(issue["id"]) if q else "",
            "report_basis": basis if q else "", "qty_affected": qty if q and "qty" in shows else "",
            "batch_ref": batch if q and "batch_ref" in shows else "",
            "packaging_version": pack if q and "packaging_version" in shows else "",
            "safety_flag": "true" if q and (safety or issue.get("safety_default")) else "",
            "client_observed_at": at.astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.000Z"),
            "photo_roles": roles, "shop_name": shop["name"], "shop_phone": shop["phone"], "area": shop["area"]}


def _post(c, tok: str, fields: dict, photos=(), client_uuid: str | None = None) -> dict:
    """POST /market-intel/observations as the phone's queue sends it (lib/marketIntel.ts send())."""
    data = {"client_uuid": client_uuid or str(uuid.uuid4())}
    data.update({k: v for k, v in fields.items() if v != ""})
    files = [("photos", (f"spotted-{i + 1}.jpg", PHOTOS[k], "image/jpeg")) for i, k in enumerate(photos)]
    return _ok(c.post("/market-intel/observations", headers=R._h(tok), data=data, files=files or None))


def _mine(c, tok: str) -> list[dict]:
    return _ok(c.get("/market-intel/mine", headers=R._h(tok)))["observations"]


def _row(rows: list[dict], obs_key: str) -> dict:
    return next(o for o in rows if o["id"] == W.obs[obs_key])


def _case(key: str) -> dict:
    return next(c for c in W.fake.rows("market_cases") if c["id"] == W.case[key])


def _photo_row(key: str) -> dict:
    return next(p for p in W.fake.rows("market_photos") if p["id"] == W.photo[key])


def _dhash(data: bytes) -> int:
    from PIL import Image
    from app import market_intel as mi
    return mi.dhash64(Image.open(io.BytesIO(data)))


def _ids(text: str, prefix: str) -> set[int]:
    return {int(x) for x in re.findall(rf"\b{prefix}-(\d+)\b", text or "")}


def _writes() -> list[tuple]:
    return [w for w in W.fake.writes() if w[1] != "audit_log"]


# ═══════════════════════════════════════════════════════════════════════════════
# 0. the stand-in itself
# ═══════════════════════════════════════════════════════════════════════════════

@test("stand-in: the schema read from the migrations (8 + 4 tables, the status / event / role lists); unknown columns, bad values and timeline rewrites are refused as the database refuses them; the photos are far apart to the dHash")
def _():
    assert {"market_items", "market_observations", "market_photos", "market_item_decisions", "suppliers", "brands",
            "market_issue_types", "market_cases", "market_case_events", "supplier_actions", "supplier_action_cases",
            "supplier_reports"} <= set(SCHEMA)
    assert "catalog_items" not in SCHEMA, "only tables these migrations create are checked"
    assert SCHEMA["market_cases"]["status"]["in"] == ("submitted", "under_review", "more_info_required", "verified",
                                                      "not_substantiated", "action_requested", "resolved", "closed")
    assert "safety_alert" in SCHEMA["market_case_events"]["event"]["in"]
    assert SCHEMA["market_photos"]["role"]["in"] == ("product", "packaging", "label_batch", "receipt", "other")
    assert SCHEMA["market_cases"]["evidence_level"]["between"] == (1, 5)
    assert {"case_id", "issue_type_id", "brand_text", "supplier_text", "client_observed_at"} <= set(SCHEMA["market_observations"])
    assert "id" not in SCHEMA["supplier_action_cases"], "keyed by (action_id, case_id)"
    _new_world()
    fake = W.fake
    fake.tables["suppliers"].append({"id": 1, "name": "Supplier Probe", "active": True})
    fake.tables["market_issue_types"].append({"id": 1, "kind": "quality", "code": "probe", "label_en": "Probe"})
    fake.tables["market_cases"].append({"id": 1, "issue_type_id": 1, "title": "Probe", "status": "submitted"})
    for email in ("queued", "sent", "not_sent"):          # one claim per case; outcomes are not limited
        fake.table("market_case_events").insert({"case_id": 1, "actor": "x", "event": "safety_alert",
                                                 "detail": {"email": email}}).execute()
    refused = (lambda: fake.table("market_case_events").insert({"case_id": 1, "actor": "y", "event": "safety_alert",
                                                                "detail": {"email": "queued"}}).execute(),
               lambda: fake.table("market_cases").select("id,colour").execute(),
               lambda: fake.table("market_cases").select("*").eq("colour", 1).execute(),
               lambda: fake.table("market_cases").update({"status": "lost"}).eq("id", 1).execute(),
               lambda: fake.table("market_cases").update({"evidence_level": 6}).eq("id", 1).execute(),
               lambda: fake.table("market_cases").insert({"issue_type_id": 1, "title": None}).execute(),
               lambda: fake.table("market_cases").insert({"id": 9, "issue_type_id": 1, "title": "x"}).execute(),
               lambda: fake.table("market_case_events").insert({"case_id": 99, "actor": "x", "event": "edit"}).execute(),
               lambda: fake.table("suppliers").insert({"name": " supplier probe "}).execute(),
               lambda: fake.table("market_case_events").update({"actor": "x"}).eq("id", 1).execute(),
               lambda: fake.table("market_case_events").delete().eq("id", 1).execute())
    for i, fn in enumerate(refused):
        try:
            fn()
        except RuntimeError as e:
            assert json.loads(str(e))["code"] in ("42703", "23514", "23502", "428C9", "23503", "23505", "23001"), (i, e)
        else:
            raise AssertionError(f"refused #{i} went through")
    assert len(fake.event_rewrites) == 2 and fake.rows("suppliers")[0]["name"] == "Supplier Probe"
    assert len(fake.rows("market_case_events")) == 3 and len(fake.event_copies) == 3
    keys = list(PHOTOS)
    hashes = {k: _dhash(PHOTOS[k]) for k in keys}
    from app import market_intel as mi
    closest = min(mi.hamming(hashes[a], hashes[b]) for i, a in enumerate(keys) for b in keys[i + 1:])
    assert closest > 2 * mi.PHASH_NEAR, f"two fixture photos are {closest} bits apart: a workbook picture could be misread"
    _new_world()


# ═══════════════════════════════════════════════════════════════════════════════
# 1. the lists
# ═══════════════════════════════════════════════════════════════════════════════

@test("1. lists: the admin creates suppliers, their brands and issue types (quality with qty + batch, packaging, a safety default) through the routes; the capture sheet reads them; a member cannot")
def _():
    with _Env():
        c = R._client()
        post = lambda kind, body, who="boss": c.post(f"/market-intel/lists/{kind}", headers=R._h(who), json=body)  # noqa: E731
        W.sup = {"X": _ok(post("suppliers", {"name": "Supplier X", "code": "SUPX"}))["entry"]["id"],
                 "Y": _ok(post("suppliers", {"name": "Supplier Y", "code": "SUPY"}))["entry"]["id"]}
        bx = _ok(post("brands", {"name": "Brand X", "supplier_id": W.sup["X"], "aliases": ["BX"]}))["entry"]
        by = _ok(post("brands", {"name": "Brand Y", "supplier_id": W.sup["Y"]}))["entry"]
        assert (bx["supplier_id"], by["supplier_id"], bx["aliases"]) == (W.sup["X"], W.sup["Y"], ["BX"])
        made = [_ok(post("issue-types", body))["entry"] for body in (
            {"kind": "quality", "code": "not_charging", "label_en": "Not charging", "label_ar": "لا يشحن",
             "fields": ["qty", "batch_ref", "label_photo"], "sort": 10},
            {"kind": "packaging", "code": "box_crushed", "label_en": "Box crushed", "fields": ["packaging_version"],
             "sort": 10},
            {"kind": "quality", "code": "overheating", "label_en": "Gets hot or swells", "label_ar": "يسخن أو ينتفخ",
             "safety_default": True, "fields": ["qty"], "sort": 20})]
        assert [t["code"] for t in made] == ["not_charging", "box_crushed", "overheating"]
        assert made[2]["safety_default"] is True and made[0]["fields"] == ["qty", "batch_ref", "label_photo"]
        assert post("suppliers", {"name": "Supplier Z"}, "clerk").status_code == 403, "lists are an admin's"
        W.lists = _ok(c.get("/market-intel/lists", headers=R._h("rep1")))
    assert [s["name"] for s in W.lists["suppliers"]] == ["Supplier X", "Supplier Y"]
    assert [b["name"] for b in W.lists["brands"]] == ["Brand X", "Brand Y"]
    assert [t["code"] for t in W.lists["issue_types"]] == ["not_charging", "overheating", "box_crushed"], \
        "the sheet's chips: quality by sort, then packaging"
    assert W.lists["capabilities"]["can_admin"] is False and "CABLE" in W.lists["categories"]
    W.issue = {t["code"]: t for t in W.lists["issue_types"]}
    audits = [a for a in W.fake.rows("audit_log") if a["event"] == "market_intel.list_create"]
    assert len(audits) == 7


# ═══════════════════════════════════════════════════════════════════════════════
# 2. capture
# ═══════════════════════════════════════════════════════════════════════════════

@test("2. capture: two reps pick products and post Spotted v2's fields from three shops; the second report of a product + issue joins its case; every v2 column lands as typed")
def _():
    at = lambda days, hours=0: NOW - timedelta(days=days, hours=hours)  # noqa: E731
    with _Env():
        c = R._client()
        cable = _pick(c, "rep1", "c-10", SHOP_A)
        assert (cable["code"], cable["name"], cable["brand"], cable["at_shop"]) == ("C-10", "C-10 Type-C Cable 1m", "Brand X", True)
        assert cable["thumb_url"].endswith("/catalog/thumbs/C-10-product-320.webp")
        assert PRICE not in json.dumps(cable) and COST not in json.dumps(cable), "the picker carries no price or cost"
        one = _post(c, "rep1", _fields(SHOP_A, at(6), product=cable, issue=W.issue["not_charging"], basis="customer_said",
                                       qty="3", batch="B2610", roles="product,label_batch", note=NOTE_1), ("1a", "1b"))
        two = _post(c, "rep2", _fields(SHOP_B, at(5), product=_pick(c, "rep2", "c-10", SHOP_B),
                                       issue=W.issue["not_charging"], basis="rep_saw", qty="2", batch="B2610",
                                       roles="product,receipt", note=NOTE_2), ("2a", "2r"))
    W.obs.update(o1=one["observation_id"], o2=two["observation_id"])
    W.case["A"] = one["case"]["id"]
    assert (one["result"], one["case"]["joined"], one["case"]["status_words"], one["safety_alerted"]) == \
        ("already_in_yq", False, "Received", False)
    assert one["message"] == "Quality report saved. The office will review it." and one["item"] is None
    assert two["case"]["id"] == W.case["A"] and two["case"]["joined"] and two["result"] == "seen_before"
    assert two["message"] == "Added to an open report: “C-10 Type-C Cable 1m · Not charging” (now 2 shops)."
    assert "Shop Alpha" not in json.dumps(two) and "Owner Alpha" not in json.dumps(two), "rep two never reads rep one's shop"
    case = _case("A")
    assert (case["title"], case["status"], case["evidence_level"], case["yq_item_code"]) == \
        ("C-10 Type-C Cable 1m · Not charging", "submitted", 2, "C-10")
    assert (case["brand_id"], case["supplier_id"], case["batch_ref"], case["obs_count"], case["shop_count"]) == \
        (W.lists["brands"][0]["id"], W.sup["X"], "B2610", 2, 2)
    o1 = next(o for o in W.fake.rows("market_observations") if o["id"] == W.obs["o1"])
    assert (o1["case_id"], o1["issue_type_id"], o1["qty_affected"], o1["report_basis"], o1["batch_ref"]) == \
        (W.case["A"], W.issue["not_charging"]["id"], 3, "customer_said", "B2610")
    assert (o1["brand_text"], o1["supplier_text"], o1["supplier_id"], o1["safety_flag"]) == ("Brand X", "Supplier X", W.sup["X"], False)
    assert (o1["shop_customer_id"], o1["shop_name"], o1["salesman_id"], o1["created_by"]) == (11, "Shop Alpha", 1, "rep1@example.com")
    assert o1["observed_at"] == o1["client_observed_at"] == at(6).isoformat(), "the time of the shot, not the server's"
    for obs_key, keys in (("o1", ("1a", "1b")), ("o2", ("2a", "2r"))):
        rows = sorted((p for p in W.fake.rows("market_photos") if p["observation_id"] == W.obs[obs_key]),
                      key=lambda p: p["position"])
        for k, p in zip(keys, rows):
            W.photo[k] = p["id"]
    assert [_photo_row(k)["role"] for k in ("1a", "1b", "2a", "2r")] == ["product", "label_batch", "product", "receipt"]
    for k in ("1a", "1b", "2a", "2r"):
        p = _photo_row(k)
        stored = W.fake.storage.objects[("finds", p["path"])][0]
        assert hashlib.sha256(stored).hexdigest() == p["sha256"] and p["thumb_path"].endswith("-t320.webp")
    assert [e["event"] for e in W.fake.rows("market_case_events") if e["case_id"] == W.case["A"]] == ["created", "link"]


@test("2. capture: a safety report alerts ONCE (its retry and a second sighting on the case send nothing); the e-mail names no shop or rep and masks the phone")
def _():
    at = lambda days, hours=0: NOW - timedelta(days=days, hours=hours)  # noqa: E731
    with _Env():
        c = R._client()
        bank = _pick(c, "rep2", "power bank", SHOP_B)
        assert bank["code"] == "P-22"
        fields = _fields(SHOP_B, at(4), product=bank, issue=W.issue["overheating"], basis="customer_said", qty="1",
                         roles="product", note=NOTE_S)
        assert fields["safety_flag"] == "true", "the sheet sets the flag for a safety-default chip"
        cu = str(uuid.uuid4())
        s = _post(c, "rep2", fields, ("S1",), client_uuid=cu)
        assert len(W.mail.calls) == 1, "the alert went out after the answer"
        before, objects = len(_writes()), len(W.fake.storage.objects)
        again = _post(c, "rep2", fields, ("S1",), client_uuid=cu)          # the phone's retry queue
        assert again["replayed"] and again["observation_id"] == s["observation_id"] and again["case"]["id"] == s["case"]["id"]
        assert len(_writes()) == before and len(W.fake.storage.objects) == objects, "a retry writes nothing"
        s2 = _post(c, "rep1", _fields(SHOP_A, at(3, 2), product=_pick(c, "rep1", "power bank", SHOP_A),
                                      issue=W.issue["overheating"], basis="rep_saw", qty="1",
                                      note="Same model warm at the counter"))
    W.obs.update(oS=s["observation_id"], oS2=s2["observation_id"])
    W.case["S"] = s["case"]["id"]
    W.photo["S1"] = next(p["id"] for p in W.fake.rows("market_photos") if p["observation_id"] == W.obs["oS"])
    assert s["safety_alerted"] and s["message"] == "Quality report saved. The office will review it. Manager alerted."
    assert s2["case"]["id"] == W.case["S"] and s2["case"]["joined"] and s2["safety_alerted"]
    assert len(W.mail.calls) == 1, "one e-mail per case: not per retry, not per sighting"
    mail = W.mail.calls[0]
    assert mail["subject"] == "YQ Market Intel - Safety report: P-22 Power Bank 10000 · Gets hot or swells", mail["subject"]
    text = mail["subject"] + mail["html"]
    for secret in ("Shop Beta", "Owner Beta", "Shop Alpha", "Rep Two", "Rep One", "rep2", "rep1", "3900 1234", "Muharraq"):
        assert secret not in text, secret
    assert "[number]" in mail["html"] and "Gets hot or swells" in mail["html"]
    case = _case("S")
    assert (case["safety"], case["obs_count"], case["shop_count"], case["evidence_level"]) == (True, 2, 2, 2)
    events = [(e["event"], (e.get("detail") or {}).get("email")) for e in W.fake.rows("market_case_events")
              if e["case_id"] == W.case["S"]]
    assert events == [("created", None), ("safety_alert", "queued"), ("safety_alert", "sent"), ("link", None)], events


@test("2. capture: 'not sure / not ours', a wrong product, packaging, supplier Y's product and a competitor price each land on their own track")
def _():
    at = lambda days, hours=0: NOW - timedelta(days=days, hours=hours)  # noqa: E731
    with _Env():
        c = R._client()
        b = _post(c, "rep1", _fields(SHOP_G, at(3), title="White cable, no brand", issue=W.issue["not_charging"],
                                     basis="rep_saw", qty="1", roles="product", note="Stops charging after two days"),
                  ("B1",))
        d = _post(c, "rep2", _fields(SHOP_B, at(2, 2), product=_pick(c, "rep2", "lightning", SHOP_B),
                                     issue=W.issue["not_charging"], basis="customer_said", qty="1",
                                     note="Lightning cable stopped charging"))
        p = _post(c, "rep1", _fields(SHOP_A, at(2), product=_pick(c, "rep1", "power bank", SHOP_A),
                                     issue=W.issue["box_crushed"], basis="rep_saw", pack="v2", roles="packaging",
                                     note="Box corners crushed on arrival"), ("P1",))
        y = _post(c, "rep1", _fields(SHOP_A, at(1), product=_pick(c, "rep1", "y-30", SHOP_A), issue=W.issue["not_charging"],
                                     basis="customer_said", qty="1", batch="Y77", roles="product",
                                     note="Charger dead out of the box"), ("Y1",))
        k = _post(c, "rep2", _fields(SHOP_B, at(1, -1), kind="competitor_price", title="Rival cable 1m", brand="Rival",
                                     price="1.250", roles="product", note="Shop Beta sells it cheaper than ours"), ("K1",))
    W.obs.update(oB=b["observation_id"], oD=d["observation_id"], oP=p["observation_id"], oY=y["observation_id"],
                 oK=k["observation_id"])
    W.case.update(B=b["case"]["id"], D=d["case"]["id"], P=p["case"]["id"], Y=y["case"]["id"])
    W.item_k = k["item"]["id"]
    for key, obs_key in (("B1", "oB"), ("P1", "oP"), ("Y1", "oY"), ("K1", "oK")):
        W.photo[key] = next(x["id"] for x in W.fake.rows("market_photos") if x["observation_id"] == W.obs[obs_key])
    assert len(set(W.case.values())) == 6, "A, S, B, D, P and Y are six cases"
    assert (b["result"], _case("B")["title"], _case("B")["yq_item_code"], _case("B")["supplier_id"]) == \
        ("new_find", "White cable, no brand · Not charging", None, None)
    assert (_case("D")["yq_item_code"], _case("D")["supplier_id"], _case("D")["evidence_level"]) == ("C-11", W.sup["X"], 1)
    assert (_case("P")["packaging_version"], _case("P")["issue_type_id"]) == ("v2", W.issue["box_crushed"]["id"])
    assert _photo_row("P1")["role"] == "packaging"
    assert (_case("Y")["supplier_id"], _case("Y")["brand_id"]) == (W.sup["Y"], W.lists["brands"][1]["id"])
    assert k["case"] is None and k["item"]["kind"] == "competitor_price" and k["result"] == "new_find"
    ko = next(o for o in W.fake.rows("market_observations") if o["id"] == W.obs["oK"])
    assert (ko["price_bhd"], ko.get("case_id"), ko["issue_type_id"], ko["safety_flag"]) == ("1.250", None, None, False)
    assert len(W.mail.calls) == 1
    for key in W.photo:                                      # every stored master, for telling pictures apart
        W.hash[key] = _dhash(W.fake.storage.objects[("finds", _photo_row(key)["path"])][0])


# ═══════════════════════════════════════════════════════════════════════════════
# 3. the office
# ═══════════════════════════════════════════════════════════════════════════════

@test("3. office: the case tabs; a case opened with its people-flagged photo signed ONLY as the blurred copy; management reads it masked")
def _():
    with _Env():
        c = R._client()
        everything = _ok(c.get("/market-intel/cases?tab=all", headers=R._h("clerk")))
        assert everything["count"] == 6 and everything["tabs"] == {"safety": 1, "new": 6, "need_info": 0, "identify": 1}, \
            everything["tabs"]
        ids = lambda tab: [x["id"] for x in _ok(c.get(f"/market-intel/cases?tab={tab}", headers=R._h("clerk")))["rows"]]  # noqa: E731
        assert ids("safety") == [W.case["S"]] and ids("identify") == [W.case["B"]] and ids("need_info") == []
        row = next(x for x in everything["rows"] if x["id"] == W.case["A"])
        assert (row["brand"], row["supplier"], row["obs_count"], row["shop_count"], row["evidence_label"]) == \
            ("Brand X", "Supplier X", 2, 2, "Evidence submitted")
        assert row["issue_type"]["code"] == "not_charging" and row["thumb_url"]
        flagged = _ok(c.post(f"/market-intel/photos/{W.photo['2a']}/people", headers=R._h("clerk"), json={"flag": True}))
        assert flagged["blurred"] is True and flagged["has_people"] is True
        r = c.get(f"/market-intel/cases/{W.case['A']}", headers=R._h("clerk"))
        detail = _ok(r)
        m = c.get(f"/market-intel/cases/{W.case['A']}", headers=R._h("mgmt"))
        masked = _ok(m)
    people = _photo_row("2a")
    assert people["has_people_flag"] and people["blurred_path"].endswith("-blur.webp")
    for text in (r.text, m.text):
        assert people["path"] not in text and people["thumb_path"] not in text, "the original is never signed"
    shot = next(p for o in detail["observations"] for p in o["photos"] if p["id"] == W.photo["2a"])
    assert shot["url"] == shot["thumb_url"] and shot["url"].endswith(people["blurred_path"] + "?ttl=3600")
    assert shot["blurred"] is True and shot["has_people_flag"] is True and shot["role"] == "product"
    receipt = next(p for o in detail["observations"] for p in o["photos"] if p["id"] == W.photo["2r"])
    assert receipt["role"] == "receipt" and receipt["url"].endswith(_photo_row("2r")["path"] + "?ttl=3600")
    assert [o["id"] for o in detail["observations"]] == [W.obs["o2"], W.obs["o1"]], "newest sighting first"
    assert detail["case"]["next_statuses"] == ["under_review", "more_info_required", "verified", "not_substantiated"]
    assert [e["event"] for e in detail["events"]] == ["created", "link", "edit"], "the flag is on the timeline"
    for raw in ("3900 1234", "owner@shop.example"):
        assert raw not in m.text, raw
        assert raw in r.text, "the office reads the note as typed"
    assert masked["capabilities"]["can_review"] is False and masked["case"]["next_statuses"] == []


@test("3. question / answer: the office asks, the rep reads it in My reports, answers with a label photo, the case is Under review again")
def _():
    with _Env():
        c = R._client()
        asked = _ok(c.post(f"/market-intel/cases/{W.case['A']}/question", headers=R._h("clerk"), json={"text": QUESTION}))
        assert asked["status"] == "more_info_required"
        tabs = _ok(c.get("/market-intel/cases?tab=need_info", headers=R._h("clerk")))
        assert [x["id"] for x in tabs["rows"]] == [W.case["A"]] and tabs["tabs"]["need_info"] == 1
        mine = _row(_mine(c, "rep1"), "o1")
        assert (mine["case_status_words"], mine["open_question"], mine["status_label"]) == \
            ("We need more from you", QUESTION, "We need more from you")
        assert _row(_mine(c, "rep2"), "o2")["open_question"] == QUESTION, "a question to the case reaches each rep on it"
        assert c.post(f"/market-intel/observations/{W.obs['o1']}/answer", headers=R._h("rep2"),
                      data={"text": "x"}).status_code == 403, "only the rep who sent it answers"
        answer = _ok(c.post(f"/market-intel/observations/{W.obs['o1']}/answer", headers=R._h("rep1"),
                            data={"text": "Batch B2610 is on the label", "role": "label_batch"},
                            files=[("photo", ("label.jpg", PHOTOS["1c"], "image/jpeg"))]))
        assert answer == {"ok": True, "observation_id": W.obs["o1"], "case_status_words": "Being looked at", "photo": True}
        after = _row(_mine(c, "rep1"), "o1")
        assert (after["case_status_words"], after["open_question"]) == ("Being looked at", None)
        assert _row(_mine(c, "rep2"), "o2")["open_question"] is None
    W.photo["1c"] = max(p["id"] for p in W.fake.rows("market_photos") if p["observation_id"] == W.obs["o1"])
    W.hash["1c"] = _dhash(W.fake.storage.objects[("finds", _photo_row("1c")["path"])][0])
    assert (_photo_row("1c")["role"], _photo_row("1c")["position"]) == ("label_batch", 2)
    assert _case("A")["status"] == "under_review"
    moves = [(e["from_status"], e["to_status"]) for e in W.fake.rows("market_case_events")
             if e["case_id"] == W.case["A"] and e["event"] == "status"]
    assert moves == [("submitted", "more_info_required"), ("more_info_required", "under_review")]


# ═══════════════════════════════════════════════════════════════════════════════
# 4. merge and move
# ═══════════════════════════════════════════════════════════════════════════════

@test("4. merge and move: the unidentified duplicate merged in, a misfiled sighting moved; the counts of all three cases follow; the rep's status words follow the merge")
def _():
    with _Env():
        c = R._client()
        merged = _ok(c.post(f"/market-intel/cases/{W.case['B']}/merge", headers=R._h("clerk"), json={"into_id": W.case["A"]}))
        assert merged == {"ok": True, "case_id": W.case["B"], "merged_into": W.case["A"], "moved_observations": 1}
        moved = _ok(c.post(f"/market-intel/observations/{W.obs['oD']}/move", headers=R._h("clerk"),
                           json={"case_id": W.case["A"]}))
        assert moved == {"ok": True, "changed": True, "observation_id": W.obs["oD"], "case_id": W.case["A"],
                         "from_case": W.case["D"]}
        listed = _ok(c.get("/market-intel/cases?tab=all", headers=R._h("clerk")))
        rows = {x["id"]: x for x in listed["rows"]}
        mine = _row(_mine(c, "rep1"), "oB")
    a, b, d = _case("A"), _case("B"), _case("D")
    assert (a["obs_count"], a["shop_count"], a["evidence_level"]) == (4, 3, 2)
    assert (b["status"], b["merged_into"], b["obs_count"]) == ("closed", W.case["A"], 0)
    assert (d["obs_count"], d["shop_count"]) == (0, 0)
    assert listed["count"] == 5 and W.case["B"] not in rows and listed["tabs"]["identify"] == 0
    assert (rows[W.case["A"]]["obs_count"], rows[W.case["A"]]["shop_count"], rows[W.case["D"]]["obs_count"]) == (4, 3, 0)
    on_a = sorted(o["id"] for o in W.fake.rows("market_observations") if o.get("case_id") == W.case["A"])
    assert on_a == sorted(W.obs[k] for k in ("o1", "o2", "oB", "oD"))
    assert mine["case_status_words"] == "Being looked at", "the merged sighting reads the case it joined"
    ev = [(e["case_id"], e["event"]) for e in W.fake.rows("market_case_events") if e["event"] in ("merge", "move")]
    assert ev == [(W.case["B"], "merge"), (W.case["A"], "merge"), (W.case["A"], "move"), (W.case["D"], "move")], ev


# ═══════════════════════════════════════════════════════════════════════════════
# 5. the workflow
# ═══════════════════════════════════════════════════════════════════════════════

@test("5. workflow: verified with a method and a finding (evidence 3), evidence only up, a supplier action (Action requested), the response, YQ's check, resolved, closed; refused moves 409, a reopen by a member 403")
def _():
    with _Env():
        c = R._client()
        patch = lambda body, who="clerk": c.patch(f"/market-intel/cases/{W.case['A']}", headers=R._h(who), json=body)  # noqa: E731
        assert patch({"status": "verified"}).status_code == 400, "a method and a finding are required"
        v = _ok(patch({"status": "verified", "verification_method": "sample_inspected",
                       "verification_reason": "Two returned samples fail at the port"}))
        assert (v["status"], v["evidence_level"]) == ("verified", 3)
        assert patch({"evidence_level": 2}).status_code == 409, "evidence only moves up"
        assert patch({"status": "resolved"}).status_code == 409, "verified cannot jump to resolved"
        act = _ok(c.post(f"/market-intel/cases/{W.case['A']}/action", headers=R._h("clerk"),
                         json={"request": "Replace batch B2610 and send a corrective action plan", "owner": "Quality desk",
                               "due_date": "2026-10-31"}))
        assert act["status"] == "action_requested" and act["action"]["supplier"] == "Supplier X"
        assert act["action"]["status"] == "open"
        W.action = act["action"]["id"]
        assert patch({"status": "resolved"}).status_code == 409, "resolved needs the response and YQ's check"
        r = _ok(c.patch(f"/market-intel/actions/{W.action}", headers=R._h("clerk"),
                        json={"supplier_response": "Batch B2610 will be replaced; questions to qa@supplier-x.example"}))
        assert r["action"]["status"] == "responded"
        assert _ok(patch({"evidence_level": 4}))["evidence_level"] == 4
        assert patch({"status": "resolved"}).status_code == 409, "still no YQ check"
        r = _ok(c.patch(f"/market-intel/actions/{W.action}", headers=R._h("clerk"),
                        json={"resolution_check": "Ten units of the new batch tested: all charge", "status": "resolved"}))
        assert r["action"]["status"] == "resolved" and r["cases"] == [W.case["A"]]
        done = _ok(patch({"status": "resolved", "evidence_level": 5}))
        assert (done["status"], done["evidence_level"]) == ("resolved", 5)
        assert _ok(patch({"status": "closed"}))["status"] == "closed"
        assert patch({"status": "verified"}).status_code == 409, "closed -> verified is not a move"
        assert patch({"status": "under_review", "reason": "came back"}).status_code == 403, "only an admin reopens"
        assert c.patch(f"/market-intel/cases/{W.case['A']}", headers=R._h("mgmt"),
                       json={"title": "x"}).json()["detail"]["code"] == "read_only"
    a = _case("A")
    assert (a["status"], a["evidence_level"], a["verification_method"]) == ("closed", 5, "sample_inspected")
    moves = [(e["from_status"], e["to_status"]) for e in W.fake.rows("market_case_events")
             if e["case_id"] == W.case["A"] and e["event"] == "status"]
    assert moves == [("submitted", "more_info_required"), ("more_info_required", "under_review"),
                     ("under_review", "verified"), ("verified", "action_requested"),
                     ("action_requested", "resolved"), ("resolved", "closed")], moves
    levels = [(e["detail"]["from"], e["detail"]["to"]) for e in W.fake.rows("market_case_events")
              if e["case_id"] == W.case["A"] and e["event"] == "evidence"]
    assert levels == [(2, 3), (3, 4), (4, 5)], levels
    assert [(x["action_id"], x["case_id"]) for x in W.fake.rows("supplier_action_cases")] == [(W.action, W.case["A"])]


# ═══════════════════════════════════════════════════════════════════════════════
# 6. the supplier workbook
# ═══════════════════════════════════════════════════════════════════════════════

FINDINGS = ("o1", "o2", "oS", "oS2", "oB", "oD", "oP")


@test("6. supplier report: the draft's counts are the story's, every finding sentence cites real [MI-...] ids of supplier X only; an edit; a narrative citing supplier Y's report is refused")
def _():
    with _Env():
        c = R._client()
        b = _ok(c.post("/market-intel/supplier-reports", headers=R._h("clerk"),
                       json={"supplier_id": W.sup["X"], "month": MONTH, "edition": "supplier"}))
        W.report = b["report"]["id"]
        text = b["report"]["narrative"]
        edited = text + "\nNote: the supplier replaced the batch [MI-" + str(W.obs["o1"]) + "]."
        e = _ok(c.patch(f"/market-intel/supplier-reports/{W.report}", headers=R._h("clerk"), json={"narrative": edited}))
        bad = c.patch(f"/market-intel/supplier-reports/{W.report}", headers=R._h("clerk"),
                      json={"narrative": "New: see [MI-" + str(W.obs["oY"]) + "]."})
    counts = b["report"]["counts"]
    assert (b["report"]["status"], b["report"]["month"], b["report"]["edition"]) == ("draft", MONTH, "supplier")
    assert (counts["reports"], counts["product_issues"], counts["packaging"], counts["shops"]) == (7, 6, 1, 3), counts
    assert (counts["verified"], counts["unverified"], counts["safety"], counts["photos"]) == (4, 3, 2, 6), counts
    assert counts["new"] == {"problems": 2, "reports": 3} and counts["recurring"] == {"problems": 1, "reports": 4}, counts
    assert (counts["insufficient"], counts["excluded"], counts["open_actions"], counts["action_ids"]) == (0, 0, 0, [W.action])
    issues = [(x["label"], x["kind"], x["reports"], x["shops"], x["verified"], x["recurring"]) for x in counts["by_issue"]]
    assert issues == [("Not charging", "quality", 4, 3, 4, True), ("Gets hot or swells", "quality", 2, 2, 0, False),
                      ("Box crushed", "packaging", 1, 1, 0, False)], issues
    assert b["report"]["included_observation_ids"] == sorted(W.obs[k] for k in FINDINGS)
    assert [r["id"] for r in b["rows"]["product_issues"]] == [W.obs[k] for k in ("o1", "o2", "oS", "oS2", "oB", "oD")]
    assert [r["id"] for r in b["rows"]["packaging"]] == [W.obs["oP"]]
    lines = text.split("\n")
    assert lines[0] == "Supplier X: October 2026." and lines[1] == (
        "Field reports counted this month: 7 (6 product issues, 1 packaging) from 3 shops; 4 verified by YQ; "
        "2 flagged as a possible safety issue."), lines[:2]
    findings = [ln for ln in lines if ln.startswith(("Recurring:", "New:", "Insufficient evidence:", "We ask:"))]
    assert len(findings) == 4 and all(_ids(ln, "MI") for ln in findings), findings
    assert _ids(text, "MI") == {W.obs[k] for k in FINDINGS}, "every counted report is cited, nothing else"
    assert not _ids(text, "CA") and "Open actions: none." in lines, "the action is resolved: no open action line"
    assert lines[2].startswith('Recurring: "Not charging" on C-10 (C-10 Type-C Cable 1m), batch "B2610": 4 reports '
                               "from 3 shops"), lines[2]
    for secret in ("Shop Alpha", "Shop Beta", "Rep One", "Supplier Y", "Y-30"):
        assert secret not in text, secret
    assert e["report"]["narrative_source"] == "edited" and e["report"]["narrative"] == edited
    assert bad.status_code == 400 and f"MI-{W.obs['oY']}" in bad.text, bad.text[:200]


@test("6. supplier report: management approves through the read-only gate (its other writes refused); the download is logged; the reviewer's preview")
def _():
    with _Env():
        c = R._client()
        listed = _ok(c.get(f"/market-intel/supplier-reports?supplier_id={W.sup['X']}&month={MONTH}", headers=R._h("mgmt")))
        assert listed["count"] == 1 and listed["capabilities"]["can_approve"] is True
        assert listed["capabilities"]["can_build"] is False
        refused = c.patch(f"/market-intel/supplier-reports/{W.report}", headers=R._h("mgmt"), json={"narrative": "x"})
        assert refused.status_code == 403 and refused.json()["detail"]["code"] == "read_only"
        early = c.post(f"/market-intel/supplier-reports/{W.report}/approve", headers=R._h("mgmt"))
        assert early.status_code in (400, 409), (early.status_code, early.text[:200])   # the month is not over yet
        from app import supplier_reports as sr
        real_now = sr._now
        sr._now = lambda: APPROVE_AT
        try:
            a = _ok(c.post(f"/market-intel/supplier-reports/{W.report}/approve", headers=R._h("mgmt")))
            assert (a["status"], a["approved_by"], a["superseded"]) == ("approved", "mgmt", [])
            d = c.get(f"/market-intel/supplier-reports/{W.report}/download", headers=R._h("clerk"))
            assert d.status_code == 200, d.text[:300]
            preview = _ok(c.get(f"/market-intel/supplier-reports/{W.report}", headers=R._h("mgmt")))
        finally:
            sr._now = real_now
    assert d.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml")
    assert 'filename="YQ-supplier-report-SUPX-2026-10-supplier.xlsx"' in d.headers["content-disposition"]
    W.xlsx = d.content
    sha = hashlib.sha256(W.xlsx).hexdigest()
    rep = preview["report"]
    assert (rep["status"], rep["approved_by"], rep["file_sha256"]) == ("approved", "mgmt", sha)
    assert [(x["by"], x["sha256"]) for x in rep["downloads"]] == [("clerk", sha)]
    people = _photo_row("2a")
    for secret in ("Shop Alpha", "Owner Alpha", "3900 1234", "owner@shop.example", "qa@supplier-x.example",
                   people["path"], people["thumb_path"], people["blurred_path"], _photo_row("2r")["path"]):
        assert secret not in json.dumps(preview), secret


@test("6. the workbook: five sheets; no shop, rep, phone, e-mail, price, cost or other supplier anywhere in the file; exactly supplier X's eligible photos embedded (never the people photo, the receipt or supplier Y's); the photo fingerprints are the capture's")
def _():
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(W.xlsx))
    assert wb.sheetnames == ["Summary", "Product issues", "Packaging", "Evidence", "Corrective actions"]
    cells = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells += [str(v) for v in row if v is not None]
    text = "\n".join(cells)
    with zipfile.ZipFile(io.BytesIO(W.xlsx)) as z:
        parts = "\n".join(z.read(n).decode("utf-8", "replace") for n in z.namelist() if n.endswith((".xml", ".rels")))
    forbidden = ["Shop Alpha", "Shop Beta", "Shop Gamma", "Owner Alpha", "Owner Beta", "Owner Gamma", "Rep One", "Rep Two",
                 "rep1", "rep2", "clerk", "mgmt", "@example.com", "owner@shop.example", "qa@supplier-x.example",
                 "3900 1234", "97330000011", "97330000012", "97330000013", "2.500", "1.250", PRICE, COST, "Supplier Y",
                 "SUPY", "Brand Y", "Y-30", "Rival", "Muharraq", "Riffa"]
    for word in forbidden:                            # whole words: "arrival" is not "Rival"
        rx = re.compile(r"(?<![a-z0-9])" + re.escape(word.lower()) + r"(?![a-z0-9])")
        assert not rx.search(text.lower()), ("a cell", word)
        assert not rx.search(parts.lower()), ("a part of the file", word)
    assert not re.search(r"(?i)\bbhd\b|\bbd\b|\bdinar", text), "no currency"
    assert not re.search(r"(?i)\b(price|cost|margin)s?\b", text), re.findall(r"(?i).{20}\b(?:price|cost|margin).{20}", text)
    assert "[shop]" in text and "field team" in text and "[number]" in text and "[email]" in text and "[amount]" in text
    ws = wb["Product issues"]
    assert [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)] == \
        [f"MI-{W.obs[k]}" for k in ("o1", "o2", "oS", "oS2", "oB", "oD")]
    head = [x.value for x in ws[1]]
    first = dict(zip(head, [x.value for x in ws[2]]))
    assert (first["Finding"], first["Evidence level"], first["Qty"], first["Batch / packaging"]) == \
        ("Verified", "Resolution verified", 3, "batch B2610")
    assert first["Description"].startswith("[shop] says it stops charging after a week; field team sold it."), first["Description"]
    second = dict(zip(head, [x.value for x in ws[3]]))
    assert second["Photo"] == "no photo shared", "rep two's photos: a people photo and a receipt"
    safety = dict(zip(head, [x.value for x in ws[4]]))
    assert (safety["Finding"], safety["Evidence level"]) == ("Unverified report", "Evidence submitted")
    assert [wb["Packaging"].cell(row=2, column=1).value, wb["Packaging"].cell(row=2, column=10).value] == \
        [f"MI-{W.obs['oP']}", "v2"]
    actions = [[x.value for x in r] for r in wb["Corrective actions"].iter_rows(min_row=2)]
    assert [(r[0], r[6]) for r in actions] == [(f"CA-{W.action}", "Resolved")] and "[email]" in actions[0][4]
    # the pictures, told apart by their dHash (fixture photos are > 16 bits apart; a re-encode moves < 8)
    from app import market_intel as mi
    seen: dict[str, list[str]] = {}
    for sheet in wb.worksheets:
        for im in sheet._images:
            h = _dhash(im._data())
            key, dist = min(((k, mi.hamming(h, v)) for k, v in W.hash.items()), key=lambda kv: kv[1])
            assert dist <= mi.PHASH_NEAR, (sheet.title, key, dist)
            seen.setdefault(sheet.title, []).append(key)
    assert seen["Evidence"] == ["1a", "1b", "1c", "S1", "B1", "P1"], seen
    assert seen["Product issues"] == ["1a", "S1", "B1"] and seen["Packaging"] == ["P1"], seen
    shown = {k for keys in seen.values() for k in keys}
    assert not shown & {"2a", "2r", "Y1", "K1"}, shown
    ev = wb["Evidence"]
    refs = {ev.cell(row=r, column=1).value: ev.cell(row=r, column=7).value for r in range(2, ev.max_row + 1)}
    assert set(refs) == {f"PH-{W.photo[k]}" for k in ("1a", "1b", "1c", "S1", "B1", "P1")}
    stored = W.fake.storage.objects[("finds", _photo_row("1a")["path"])][0]
    assert refs[f"PH-{W.photo['1a']}"] == hashlib.sha256(stored).hexdigest()[:12], "the capture's sha256, as stored"


# ═══════════════════════════════════════════════════════════════════════════════
# 7. a rep
# ═══════════════════════════════════════════════════════════════════════════════

@test("7. a rep: 403 on cases, list edits, actions, merges, moves, photo flags, the review queue and every supplier-report route, writing nothing; the capture sheet still reads the lists")
def _():
    A, rid = W.case["A"], W.report
    calls = (("get", "/market-intel/cases", None), ("get", f"/market-intel/cases/{A}", None),
             ("patch", f"/market-intel/cases/{A}", {"title": "x"}),
             ("post", f"/market-intel/cases/{A}/question", {"text": "x"}),
             ("post", f"/market-intel/cases/{A}/merge", {"into_id": W.case["S"]}),
             ("post", f"/market-intel/cases/{A}/action", {"request": "x"}),
             ("post", f"/market-intel/observations/{W.obs['oS2']}/move", {"case_id": A}),
             ("patch", f"/market-intel/actions/{W.action}", {"status": "cancelled"}),
             ("post", "/market-intel/lists/suppliers", {"name": "Supplier R"}),
             ("patch", f"/market-intel/lists/brands/{W.lists['brands'][0]['id']}", {"name": "x"}),
             ("post", "/market-intel/lists/issue-types", {"kind": "quality", "label_en": "x"}),
             ("post", f"/market-intel/photos/{W.photo['1a']}/people", {"flag": True}),
             ("get", "/market-intel/review", None),
             ("post", "/market-intel/supplier-reports", {"supplier_id": W.sup["X"], "month": MONTH}),
             ("get", "/market-intel/supplier-reports", None), ("get", f"/market-intel/supplier-reports/{rid}", None),
             ("patch", f"/market-intel/supplier-reports/{rid}", {"narrative": "x"}),
             ("post", f"/market-intel/supplier-reports/{rid}/approve", None),
             ("get", f"/market-intel/supplier-reports/{rid}/download", None))
    with _Env():
        c = R._client()
        before = len(W.fake.writes())
        for tok in ("rep1", "rep2"):
            for meth, url, body in calls:
                r = getattr(c, meth)(url, headers=R._h(tok), **({"json": body} if body is not None else {}))
                assert r.status_code == 403, (tok, meth, url, r.status_code, r.text[:160])
        assert len(W.fake.writes()) == before, W.fake.writes()[before:]
        assert c.get("/market-intel/mine", headers=R._h("rep3")).status_code == 403, "no page, no Market Intel"
        lists = c.get("/market-intel/lists", headers=R._h("rep1"))
        assert lists.status_code == 200, "the capture sheet's chips (contract §4)"
    for secret in ("Shop", "Owner", "Rep One", "Rep Two", "@example.com"):
        assert secret not in lists.text, secret


@test("7. a rep: My reports and the board show his own sightings with plain status words only — never another rep's or shop's name, note or photo")
def _():
    with _Env():
        c = R._client()
        one = c.get("/market-intel/mine", headers=R._h("rep1"))
        two = c.get("/market-intel/mine", headers=R._h("rep2"))
        board = c.get("/market-intel/items", headers=R._h("rep1"))
        item = c.get(f"/market-intel/items/{W.item_k}", headers=R._h("rep1"))
    for r in (one, two, board, item):
        assert r.status_code == 200, (str(r.request.url), r.text[:200])
    mine1 = {o["id"]: o for o in one.json()["observations"]}
    assert set(mine1) == {W.obs[k] for k in ("o1", "oS2", "oB", "oP", "oY")}
    assert set(o["id"] for o in two.json()["observations"]) == {W.obs[k] for k in ("o2", "oS", "oD", "oK")}
    for secret in ("Shop Beta", "Owner Beta", "Rep Two", "rep2", "Two more dead", "swelled", "Lightning cable stopped",
                   "sells it cheaper", "Muharraq", "Not charging", "Supplier X", "supplier_text"):
        assert secret not in one.text, ("rep one reads", secret)
    for secret in ("Shop Alpha", "Shop Gamma", "Owner Alpha", "Rep One", "rep1", "stops charging after a week",
                   "White cable", "Box corners", "Manama", "Supplier X", "B2610 is on the label"):
        assert secret not in two.text, ("rep two reads", secret)
    words = {o["id"]: o.get("case_status_words") for o in two.json()["observations"]}
    assert words == {W.obs["o2"]: "Closed", W.obs["oS"]: "Received", W.obs["oD"]: "Closed", W.obs["oK"]: None}, words
    assert mine1[W.obs["oB"]]["case_status_words"] == "Closed" and mine1[W.obs["o1"]]["open_question"] is None
    assert all(p["id"] != W.photo["2a"] for o in two.json()["observations"] for p in o["photos"]), \
        "a rep never receives a people-flagged photo, even his own"
    items = board.json()["items"]
    assert [i["item_id"] for i in items] == [W.item_k] and items[0]["title"] == "Rival cable 1m"
    for secret in ("Shop Beta", "Rep Two", "rep2", "sells it cheaper", "Muharraq"):
        assert secret not in board.text and secret not in item.text, secret
    assert item.json()["observations"] == [] and "shops" not in item.json(), "his own sightings only; no shop list"


# ═══════════════════════════════════════════════════════════════════════════════
# 8. the importer
# ═══════════════════════════════════════════════════════════════════════════════

@test("8. importer: a dry run on the same data sees the r8m photo columns, reads and checks every old photo, plans a 320 px thumbnail for each, flags no near-duplicate of the live captures and writes nothing")
def _():
    from scripts import market_intel_import as imp
    fake = W.fake
    fake.tables["product_finds"] = [
        {"id": 1, "name": "Brand Q desk lamp", "price_bhd": None, "currency": "BHD", "note": None, "category": None,
         "source": "WhatsApp", "image_path": "finds/old-1.jpeg", "status": "new", "promoted_item_code": None,
         "source_file": "WhatsApp Image 2026-07-05 at 4.11.20 PM.jpeg", "posted_by": "sales-team",
         "posted_at": "2026-07-12T13:42:11+00:00"},
        {"id": 2, "name": "Brand Q fan", "price_bhd": None, "currency": "BHD", "note": None, "category": None,
         "source": "WhatsApp", "image_path": "finds/old-2.jpeg", "status": "new", "promoted_item_code": None,
         "source_file": "WhatsApp Image 2026-07-06 at 9.02.00 AM.jpeg", "posted_by": "sales-team",
         "posted_at": "2026-07-12T13:42:11+00:00"}]
    fake.tables["field_notes"] = [{"id": 1, "note": "Shop Beta asks for twenty more power banks", "category": "demand",
                                   "created_by": "boss@example.com", "created_at": "2026-06-30T05:59:00+00:00",
                                   "image_path": None}]
    fake.storage.objects[("finds", "finds/old-1.jpeg")] = (PHOTOS["I1"], {})
    fake.storage.objects[("finds", "finds/old-2.jpeg")] = (PHOTOS["I2"], {})
    writes, objects, uploads = len(fake.writes()), dict(fake.storage.objects), len(fake.storage.calls)
    lines: list[str] = []
    with _Env():
        code = imp.main([], client=fake, out=lines.append)
    text = "\n".join(str(x) for x in lines)
    assert code == 0, text
    assert "DRY RUN" in text and "schema: the R7b tables with the r8m photo columns (thumb_path, role, sha256)" in text, text
    assert "read: 2 product_finds, 1 field_notes -> 3 sightings with 2 photos" in text, text
    assert "to create: 3 sightings (new_product 2, shop_asked 1)" in text, text
    assert "photos: 2, each with a 320 px thumbnail" in text and "near-duplicate photo: 0 sightings" in text, text
    assert "Dry run: nothing written." in text and "cannot be imported" not in text, text
    assert len(fake.writes()) == writes and fake.storage.objects == objects
    assert not [x for x in fake.storage.calls[uploads:] if x[0] in ("upload", "remove")]
    # the same importer with --apply, in memory, on a COPY of this world (the shared data stays as it is):
    # every row it writes meets the real schema, since the stand-in checks each write against the migrations
    twin = _DB(copy.deepcopy(fake.tables))
    twin.storage.objects = dict(fake.storage.objects)
    lines = []
    assert imp.main(["--apply"], client=twin, out=lines.append) == 0, "\n".join(lines)
    made = [o for o in twin.rows("market_observations") if o.get("source") == "import"]
    assert len(made) == 3 and all(o.get("case_id") is None and o.get("client_observed_at") is None for o in made)
    for p in (p for p in twin.rows("market_photos") if p["observation_id"] in {o["id"] for o in made}):
        assert p["role"] == "product" and p["thumb_path"].endswith("-t320.webp") and not p["has_people_flag"], p
        assert hashlib.sha256(twin.storage.objects[("finds", p["path"])][0]).hexdigest() == p["sha256"]
        assert ("finds", p["thumb_path"]) in twin.storage.objects
    lines = []
    assert imp.main(["--apply"], client=twin, out=lines.append) == 0
    assert "already in Market Intel (skipped): 3 sightings" in "\n".join(lines), "\n".join(lines)
    assert len(fake.writes()) == writes and fake.storage.objects == objects, "the shared world is untouched"


# ═══════════════════════════════════════════════════════════════════════════════
# 9. a safety sighting moved; 10. two safety reports racing; 11. the timeline; the code
# ═══════════════════════════════════════════════════════════════════════════════

@test("9. move: a safety sighting moved onto a case that was not a safety case takes that case onto the Safety tab, as a merge does; no second e-mail")
def _():
    D = W.case["D"]
    assert _case("D")["safety"] is False and _case("D")["obs_count"] == 0
    with _Env():
        c = R._client()
        before = len(W.mail.calls)
        moved = _ok(c.post(f"/market-intel/observations/{W.obs['oS2']}/move", headers=R._h("clerk"), json={"case_id": D}))
        assert (moved["case_id"], moved["from_case"]) == (D, W.case["S"])
        safety = [x["id"] for x in _ok(c.get("/market-intel/cases?tab=safety", headers=R._h("clerk")))["rows"]]
        assert _row(_mine(c, "rep1"), "oS2")["case_status_words"] == "Received"
    assert (_case("D")["safety"], _case("D")["obs_count"], _case("S")["obs_count"]) == (True, 1, 1)
    assert set(safety) == {W.case["S"], D}, safety
    assert len(W.mail.calls) == before, "the problem was alerted already"
    edit = [e["detail"] for e in W.fake.rows("market_case_events") if e["case_id"] == D and e["event"] == "edit"]
    assert edit == [{"safety": {"from": False, "to": True}, "observation_id": W.obs["oS2"]}], edit


@test("10. safety race: two phones report one problem at once (the second's claim read misses the first's claim); the database's one-claim index refuses the second claim and the API reads that as 'already claimed': the report is saved, no second e-mail")
def _():
    with _Env():
        c = R._client()
        before = len(W.mail.calls)
        bank = _pick(c, "rep2", "power bank", SHOP_B)
        W.fake.claim_race = True
        try:
            card = _post(c, "rep2", _fields(SHOP_B, NOW - timedelta(hours=2), product=bank, issue=W.issue["overheating"],
                                            basis="rep_saw", qty="1", note="Another one got hot on the shelf"))
        finally:
            W.fake.claim_race = False
    W.obs["oS3"] = card["observation_id"]
    assert card["case"]["id"] == W.case["S"] and card["case"]["joined"] and card["safety_alerted"], card
    assert len(W.mail.calls) == before, "one claim per case: the race sends nothing"
    claims = [e for e in W.fake.rows("market_case_events") if e["case_id"] == W.case["S"] and e["event"] == "safety_alert"
              and (e.get("detail") or {}).get("email") == "queued"]
    assert len(claims) == 1 and _case("S")["obs_count"] == 2


@test("11. the case timeline was only appended to: no UPDATE or DELETE was tried on market_case_events and every row is as it was written")
def _():
    assert W.fake.event_rewrites == [], W.fake.event_rewrites
    assert W.fake.rows("market_case_events") == W.fake.event_copies
    kinds = {e["event"] for e in W.fake.rows("market_case_events")}
    assert {"created", "link", "safety_alert", "edit", "question", "answer", "status", "merge", "move", "evidence",
            "action"} <= kinds, kinds


@test("code: Python 3.11 grammar, no nested f-string, no 3-decimal amount of 1,000 or more; the CI runs this suite after the other R8m suites")
def _():
    from scripts.ops import selftest
    src = Path(__file__).read_text(encoding="utf-8")
    tree = ast.parse(src, feature_version=(3, 11))
    for node in ast.walk(tree):
        if isinstance(node, ast.JoinedStr):
            assert not any(isinstance(v, ast.FormattedValue) and any(isinstance(x, ast.JoinedStr) for x in ast.walk(v.value))
                           for v in node.values), f"a nested f-string, line {node.lineno}"
    assert not selftest.scan_text(src, []), selftest.scan_text(src, [])
    ci = (ROOT / ".github" / "workflows" / "ci.yml").read_text(encoding="utf-8")
    step = ci[ci.index("python -m tests.test_r7b_market_intel\n"):].split("\n      - name:", 1)[0]
    assert step.index("python -m tests.test_r8m_market_quality\n") < step.index("python -m tests.test_r8m_mi_e2e")


def main() -> int:
    passed = failed = 0
    for name, fn in TESTS:
        try:
            fn()
            print(f"PASS  {name}")
            passed += 1
        except Exception:  # noqa: BLE001
            print(f"FAIL  {name}")
            traceback.print_exc()
            failed += 1
    for gap in sorted((W.fake.known if W.fake is not None else set())):
        print(f"NOTE  known mismatch met, owned by another builder: {KNOWN_GAPS[gap]}")
    print(f"\n{passed} passed, {failed} failed")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
