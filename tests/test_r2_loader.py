"""R2b tests (24-Sep-2026): loader hardening, costs, receivables / profitability truth, reserved stock.

    python -m tests.test_r2_loader

Same lightweight runner as tests/test_v3.py (no pytest). Every test is pure -- no database, no
.env -- except the ones headed "local:", which run against the LOCAL replay Postgres built by
`python -m scripts.local_replay_db --db r2_loader` (port 55432) inside a transaction that is always
rolled back, and SKIP (print, never fail) when that cluster or database is not there. Nothing in
this file reads or writes production.
"""
from __future__ import annotations

import io
import os
import sys
import tempfile
import traceback
import types
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import warnings  # noqa: E402
warnings.filterwarnings("ignore")

import pandas as pd  # noqa: E402

TESTS: list[tuple[str, object]] = []
LOCAL_DB = os.environ.get("YQ_REPLAY_DB", "r2_loader")
LOCAL_URL = f"postgresql://postgres@localhost:{os.environ.get('YQ_LOCAL_PG_PORT', '55432')}/{LOCAL_DB}"
DROP_240926 = Path(os.environ.get("YQ_DROP_240926", r"C:\Users\fahmed\OneDrive - YqBahrain\Desktop\YQ Bahrain Mobile Accessories\240926"))


def test(name):
    def deco(fn):
        TESTS.append((name, fn))
        return fn
    return deco


# ── synthetic Focus grids ────────────────────────────────────────────────────

def _grid(rows: list[list], width: int | None = None) -> pd.DataFrame:
    """A pandas grid like ingest.read_grid returns: object cells, None for blanks."""
    w = width or max(len(r) for r in rows)
    return pd.DataFrame([list(r) + [None] * (w - len(r)) for r in rows], dtype=object)


def _title(kind_title: str, as_of: str = "24/09/2026") -> list[list]:
    return [[None, None, None, "YQ Bahrain W.L.L"], [None, None, None, kind_title], [], [None, None, None, f"[As on date {as_of}]"], []]


# ── 1. the header contract ───────────────────────────────────────────────────

@test("header contract: every kind is found on the real layout, at any row, with extra trailing columns")
def _():
    from scripts.ingest import EXPECTED_HEADERS, HeaderMismatch, check_header, header_row_index
    for kind, labels in EXPECTED_HEADERS.items():
        g = _grid(_title(kind) + [list(labels) + ["Base Link doc. number", "Document No."]] + [["x"] * len(labels)])
        assert check_header(g, kind) == 5, kind
        g0 = _grid([list(labels), ["x"] * len(labels)])          # price books: header on row 0
        assert check_header(g0, kind) == 0, kind
        assert header_row_index(_grid([["nothing"], ["here"]]), kind) is None
    # labels are matched ignoring case and inner whitespace, never position-free
    g = _grid(_title("Summary sales register") + [["DATE", "invoice ", "Customer", "gross", "Salesman", "Payment  Mode", "Sales Account Name"]])
    assert check_header(g, "orders") == 5
    try:
        check_header(_grid(_title("x") + [["Date", "Invoice", "Customer", "Gross Amount", "Salesman", "Payment Mode", "Sales Account Name"]]), "orders")
        raise AssertionError("a renamed column must raise")
    except HeaderMismatch as e:
        msg = str(e)
        assert msg.startswith("orders:") and "col 3: expected 'Gross', found 'gross amount'" in msg and "6/7 labels" in msg, msg
        assert "nothing was loaded" in msg
    try:
        check_header(_grid([["Date", "Voucher"]]), "order_lines")
        raise AssertionError("a short header must raise")
    except HeaderMismatch as e:
        assert "order_lines:" in str(e) and "col 2: expected 'Customer Account'" in str(e)
    try:
        check_header(_grid([["a"]]), "no_such_kind")
        raise AssertionError("an unknown kind must raise")
    except HeaderMismatch:
        pass
    # a header past row 12 is not a Focus export
    try:
        check_header(_grid([[None]] * 14 + [list(EXPECTED_HEADERS["orders"])]), "orders")
        raise AssertionError("must raise")
    except HeaderMismatch:
        pass


@test("parsers: read from the located header row (title block of any height), skip totals, keep the AR Grand Total")
def _():
    from scripts.ingest import (EXPECTED_HEADERS, parse_order_lines, parse_orders, parse_profitability,
                                parse_receivables, parse_receivables_totals, parse_stock_balance)
    from datetime import datetime
    d = datetime(2026, 9, 24)
    g = _grid([["", "", "", "YQ Bahrain W.L.L"], ["", "", "", "Summary sales register"], ["", "", "", "[As on date 24/09/2026]"],
               list(EXPECTED_HEADERS["orders"]),                         # header on row 3, not 5
               [d, "SI : 1", "Shop A", 12.5, "Karrar", "Cash", "Sales"],
               ["Grand Total", None, None, 12.5],
               [d, "SI : 2", "Shop B", 7.25, "Ahmed", "Credit", "Sales"]])
    od = parse_orders(g, "f")
    assert [o["invoice_no"] for o in od] == ["SI : 1", "SI : 2"] and od[0]["order_date"] == "2026-09-24" and od[1]["gross_bhd"] == 7.25
    g = _grid(_title("Sales day book") + [list(EXPECTED_HEADERS["order_lines"]),
                                          [d, "SI : 1", "Shop A", "X01 Cable", 3, 1.1, 3.3, None, 3.0, 0.3, 3.3, None, "Karrar"],
                                          [d, "SI : 1", "Shop A", "T02 Airpod", 1, 2.95, 2.95, None, 2.68, 0.27, 2.95, None, "Karrar"],
                                          ["Total", None, None, None, 4]])
    ol = parse_order_lines(g, "f")
    assert [(r["invoice_no"], r["line_no"], r["warehouse_name"]) for r in ol] == [("SI : 1", 1, "Karrar"), ("SI : 1", 2, "Karrar")]
    g = _grid(_title("Product Profitability Report") + [list(EXPECTED_HEADERS["product_profitability"]) + ["Base Link doc. number"],
                                                        ["UK03 20W Charger (USB + Type-C Port) (VFAN)", 439.8, None, 439.8, 483.2, 43.4, 980.03, None, 43.4, 980.03, 12],
                                                        ["Grand Total", 439.8]])
    pp = parse_profitability(g, "f")
    assert len(pp) == 1 and pp[0]["cogs_bhd"] == 483.2 and pp[0]["report_date"] == "2026-09-24" and pp[0]["gross_profit_bhd"] == 43.4
    g = _grid([["", "", "YQ Bahrain W.L.L"], ["", "", "Stock balance by warehouse of Accessories Warehouse"], [], ["", "", "[As on date 24/09/2026]"], [],
               list(EXPECTED_HEADERS["stock_balance"]),
               ["Accessories Warehouse Accessories Warehouse"], ["UK12 45W Charger", 182, 4, 728], ["Total", 182, None, 728]])
    sb = parse_stock_balance(g, "f")
    assert sb == [{"item_name": "UK12 45W Charger", "warehouse_name": "Accessories Warehouse", "net_qty": 182.0,
                   "selling_rate_bhd": 4.0, "total_value_bhd": 728.0, "as_of_date": "2026-09-24", "source_file": "f"}], sb
    hdr = list(EXPECTED_HEADERS["receivables"]) + ["PDC Amount"] * 38 + ["Account Code", "LastReceiptDate", "LastReceiptAmount", "x", "y", "z", "Group Name"]
    def _ar(name, bal, b0, b91, b210, code, group):
        r = [name, bal, bal, None, None, bal, b0, None, None, b91, None, None, None, b210, bal] + [None] * 38
        return r + [code, None, None, None, None, None, group]
    g = _grid([[None] * 30 + ["YQ Bahrain W.L.L"], [None] * 30 + ["Customer summary ageing by due date"], [],
               [None] * 30 + ["[As on date 24/09/2026]"], [None] * 14 + ["Base"], hdr,
               _ar("Shop A", 101.9, 101.9, None, None, "131819-1", "Retail"),
               _ar("Shop B", 55.0, None, 28.0, 27.0, "88305-4", "Retail"),
               _ar("Grand Total", 8633.84, 2274.28, 28.0, 3200.45, None, None)])
    ar = parse_receivables(g, "f")
    assert [r["account"] for r in ar] == ["Shop A", "Shop B"] and ar[1]["bucket_91_120"] == 28.0 and ar[1]["group_name"] == "Retail"
    tot = parse_receivables_totals(g, "f")
    assert tot == {"as_of_date": "2026-09-24", "focus_total_bhd": 8633.84, "focus_over90_bhd": 3228.45,
                   "rows_total_bhd": 156.9, "source_file": "f"}, tot
    assert parse_receivables_totals(_grid([[None] * 14 + ["Base"], hdr, _ar("Shop A", 1.0, 1.0, None, None, "1", "R")]), "f") is None


@test("ingest main: a changed header stops the run (exit 2) before any CSV is written")
def _():
    from openpyxl import Workbook
    from scripts import ingest
    with tempfile.TemporaryDirectory() as d:
        src, out = Path(d) / "drop", Path(d) / "clean"
        src.mkdir()
        wb = Workbook()
        ws = wb.active
        for row in [["", "", "", "YQ Bahrain W.L.L"], ["", "", "", "Summary sales register"], [], ["", "", "", "[As on date 24/09/2026]"], [],
                    ["Date", "Invoice No", "Customer", "Gross", "Salesman", "Payment Mode", "Sales Account Name"],
                    ["2026-09-24", "SI : 1", "Shop A", 1.0, "Karrar", "Cash", "Sales"]]:
            ws.append(row)
        wb.save(src / "Summary_sales_register1_YQ_Bahrain_W_L_L.xlsx")
        old_out, old_argv = ingest.OUT_DIR, sys.argv
        ingest.OUT_DIR = out
        sys.argv = ["ingest", str(src)]
        try:
            rc = ingest.main()
        finally:
            ingest.OUT_DIR, sys.argv = old_out, old_argv
        assert rc == 2 and not list(out.glob("*.csv")), rc


# ── 2. the loader: span replace, warehouse guard, markers ─────────────────────

@test("loader: date_span / month_windows (real month ends, year boundary) / markers")
def _():
    from scripts.load_supabase import date_span, marker_present, month_windows
    assert date_span(["2026-09-24", None, float("nan"), "2025-09-21T00:00:00", "nan"]) == ("2025-09-21", "2026-09-24")
    assert date_span([]) is None and date_span([None]) is None
    assert month_windows("2026-02-10", "2026-02-20") == [("2026-02-10", "2026-02-20")]
    assert month_windows("2025-12-15", "2026-02-03") == [("2025-12-15", "2025-12-31"), ("2026-01-01", "2026-01-31"), ("2026-02-01", "2026-02-03")]
    assert month_windows("2026-02-01", "2026-02-28")[-1][1] == "2026-02-28"
    with tempfile.TemporaryDirectory() as d:
        assert not marker_present(d, "REPLACE_OK")
        (Path(d) / "replace_ok.txt").write_bytes(b"")
        assert marker_present(d, "REPLACE_OK") and not marker_present(d, "PARTIAL_OK")
        (Path(d) / "PARTIAL_OK").write_bytes(b"")
        assert marker_present(d, "PARTIAL_OK")
    assert not marker_present(None, "PARTIAL_OK")


@test("loader: disappearing invoices are listed; the guard refuses a replace over 10% (min 5) unless REPLACE_OK")
def _():
    from scripts.load_supabase import disappearing_invoices, span_guard
    db = [{"invoice_no": f"SI : {i}", "order_date": "2026-09-10"} for i in range(1, 101)]
    file_inv = {f"SI : {i}" for i in range(1, 97)}                  # 4 gone
    gone = disappearing_invoices(db + db[:2], file_inv)             # duplicates (lines) count once
    assert [r["invoice_no"] for r in gone] == ["SI : 97", "SI : 98", "SI : 99", "SI : 100"]
    assert span_guard(gone, 100) is None                            # under the minimum of 5
    gone5 = disappearing_invoices(db, {f"SI : {i}" for i in range(1, 96)})
    assert span_guard(gone5, 100) is None                           # 5 = 5%, allowed
    gone11 = disappearing_invoices(db, {f"SI : {i}" for i in range(1, 90)})
    r = span_guard(gone11, 100)
    assert r and "11 of the 100 invoices" in r and "REPLACE_OK" in r, r
    assert span_guard(gone11, 100, replace_ok=True) is None
    assert span_guard(gone11, 0) is None and span_guard([], 100) is None


@test("loader: the warehouse-set guard warns when a snapshot covers fewer warehouses than the previous one, PARTIAL_OK silences it")
def _():
    from scripts.load_supabase import warehouse_set_guard
    prev = {f"W{i}" for i in range(25)}
    r = warehouse_set_guard(prev, {"W1"})
    assert r and "covers 1 warehouse(s), the previous one 25" in r and "PARTIAL_OK" in r, r
    assert warehouse_set_guard(prev, {"W1"}, partial_ok=True) is None
    assert warehouse_set_guard({"W1"}, {"W1"}) is None
    assert warehouse_set_guard({"W1"}, prev) is None                # a wider snapshot is fine
    assert warehouse_set_guard(set(), {"W1"}) is None               # first snapshot ever
    assert warehouse_set_guard({"A", "B"}, {"A", "C"}) is None      # same size, renamed: not "smaller"


class _FakeQ:
    """In-memory PostgREST: select / eq / gte / lte / lt / is_ / not_ / in_ / order / range / limit /
    delete(count) / upsert / insert / update. Enough for the loader paths under test."""

    def __init__(self, db, table):
        self.db, self.t = db, table
        self.filters, self.op, self.payload, self.rng, self.neg = [], "select", None, None, False
        self.on_conflict, self.ignore_dup, self.order_key, self.desc, self.lim = None, False, None, False, None

    def select(self, cols="*", **k):
        self.op = "select"; self.cols = cols; return self

    def eq(self, c, v):
        self.filters.append(lambda r: str(r.get(c)) == str(v)); return self

    def gte(self, c, v):
        self.filters.append(lambda r: r.get(c) is not None and str(r.get(c))[:10] >= str(v)[:10]); return self

    def lte(self, c, v):
        self.filters.append(lambda r: r.get(c) is not None and str(r.get(c))[:10] <= str(v)[:10]); return self

    def lt(self, c, v):
        self.filters.append(lambda r: r.get(c) is not None and str(r.get(c))[:10] < str(v)[:10]); return self

    def is_(self, c, v):
        neg, self.neg = self.neg, False
        self.filters.append((lambda r: r.get(c) is not None) if neg else (lambda r: r.get(c) is None)); return self

    @property
    def not_(self):
        self.neg = True; return self

    def in_(self, c, vals):
        s = {str(x) for x in vals}; self.filters.append(lambda r: str(r.get(c)) in s); return self

    def order(self, col, desc=False, **k):
        self.order_key, self.desc = col, desc; return self

    def limit(self, n, **k):
        self.lim = n; return self

    def range(self, a, b):
        self.rng = (a, b); return self

    def delete(self, count=None, **k):
        self.op = "delete"; return self

    def upsert(self, rows, on_conflict=None, ignore_duplicates=False, **k):
        self.op, self.payload, self.on_conflict, self.ignore_dup = "upsert", rows, on_conflict, ignore_duplicates; return self

    def insert(self, payload):
        self.op, self.payload = "insert", payload; return self

    def update(self, payload):
        self.op, self.payload = "update", payload; return self

    def _match(self):
        return [r for r in self.db.setdefault(self.t, []) if all(f(r) for f in self.filters)]

    def execute(self):
        rows = self.db.setdefault(self.t, [])
        if self.op == "select":
            sel = self._match()
            if self.order_key:
                sel = sorted(sel, key=lambda r: (r.get(self.order_key) is None, str(r.get(self.order_key))), reverse=self.desc)
            if self.rng:
                sel = sel[self.rng[0]:self.rng[1] + 1]
            if self.lim is not None:
                sel = sel[:self.lim]
            return types.SimpleNamespace(data=[dict(r) for r in sel], count=len(self._match()))
        if self.op == "delete":
            gone = self._match()
            ids = {id(r) for r in gone}
            self.db[self.t] = [r for r in rows if id(r) not in ids]
            return types.SimpleNamespace(data=[dict(r) for r in gone], count=len(gone))
        if self.op == "upsert":
            keys = [k.strip() for k in (self.on_conflict or "").split(",") if k.strip()]
            n = 0
            for rec in self.payload:
                hit = next((r for r in rows if keys and all(str(r.get(k)) == str(rec.get(k)) for k in keys)), None)
                if hit is None:
                    rec = dict(rec); rec.setdefault("id", self.db.setdefault("_seq", 1000) + 1); self.db["_seq"] = rec["id"]
                    rows.append(rec); n += 1
                elif not self.ignore_dup:
                    hit.update(rec)
            return types.SimpleNamespace(data=list(self.payload), count=n)
        if self.op == "update":
            n = 0
            for r in self._match():
                r.update(self.payload); n += 1
            return types.SimpleNamespace(data=[], count=n)
        recs = self.payload if isinstance(self.payload, list) else [self.payload]
        for rec in recs:
            rows.append(dict(rec))
        return types.SimpleNamespace(data=[dict(r) for r in recs], count=None)


class _FakeDB:
    def __init__(self, tables):
        self.tables = tables

    def table(self, name):
        if name not in self.tables and name != "audit_log":
            raise RuntimeError(f"relation {name} does not exist")
        return _FakeQ(self.tables, name)


@test("loader: order_lines / orders are cleared inside the file's span only; the invoices that vanish are listed and audited")
def _():
    from scripts.load_supabase import _sales_span_replace
    db = {"orders": [
        {"id": 1, "invoice_no": "SI : Z", "order_date": "2026-08-01", "gross_bhd": 9.0},   # before the span
        {"id": 2, "invoice_no": "SI : A", "order_date": "2026-09-01", "gross_bhd": 1.0},
        {"id": 3, "invoice_no": "SI : D", "order_date": "2026-09-03", "gross_bhd": 4.0},   # in span, not in file
        {"id": 4, "invoice_no": "SI : B", "order_date": "2026-09-10", "gross_bhd": 2.0},   # after the span
    ], "audit_log": []}
    client = _FakeDB(db)
    od = pd.DataFrame([{"invoice_no": "SI : A", "order_date": "2026-09-01", "gross_bhd": "1.0"},
                       {"invoice_no": "SI : C", "order_date": "2026-09-05", "gross_bhd": "3.0"}])
    res = _sales_span_replace(client, "orders", od, "order_date", replace_ok=False)
    assert res["span"] == ("2026-09-01", "2026-09-05") and [r["invoice_no"] for r in res["disappearing"]] == ["SI : D"]
    assert res["removed"] == 2 and res["skipped"] is None, res
    assert sorted(r["invoice_no"] for r in db["orders"]) == ["SI : B", "SI : Z"]        # outside the span: untouched
    assert len(db["audit_log"]) == 1 and db["audit_log"][0]["event"] == "orders.span_replace"
    assert db["audit_log"][0]["detail"]["disappearing"] == ["SI : D"]
    # the guard: 6 of 30 invoices in the span missing from the file (20%) -> upsert only, audited, nothing removed
    db = {"orders": [{"id": i, "invoice_no": f"SI : {i}", "order_date": "2026-09-02", "gross_bhd": 1.0} for i in range(1, 31)],
          "audit_log": []}
    od = pd.DataFrame([{"invoice_no": f"SI : {i}", "order_date": "2026-09-02", "gross_bhd": "1.0"} for i in range(1, 25)])
    res = _sales_span_replace(_FakeDB(db), "orders", od, "order_date", replace_ok=False)
    assert res["skipped"] and "6 of the 30 invoices" in res["skipped"] and res["removed"] == 0, res
    assert len(db["orders"]) == 30 and db["audit_log"][0]["event"] == "orders.span_replace_skipped"
    assert db["audit_log"][0]["detail"]["disappearing"] == [f"SI : {i}" for i in range(25, 31)]
    res = _sales_span_replace(_FakeDB(db), "orders", od, "order_date", replace_ok=True)
    assert res["skipped"] is None and res["removed"] == 30
    # a handful (under 5) always replaces: 3 of 20 is the 21-Sep case (2 invoices Focus had deleted)
    db = {"orders": [{"id": i, "invoice_no": f"SI : {i}", "order_date": "2026-09-02", "gross_bhd": 1.0} for i in range(1, 21)],
          "audit_log": []}
    od = pd.DataFrame([{"invoice_no": f"SI : {i}", "order_date": "2026-09-02", "gross_bhd": "1.0"} for i in range(1, 18)])
    res = _sales_span_replace(_FakeDB(db), "orders", od, "order_date", replace_ok=False)
    assert res["skipped"] is None and res["removed"] == 20 and len(res["disappearing"]) == 3
    # a span the database cannot be read for is never deleted blind
    class _Broken(_FakeDB):
        def table(self, name):
            raise RuntimeError("connection reset")
    res = _sales_span_replace(_Broken({}), "order_lines", od.rename(columns={"order_date": "line_date"}), "line_date", False)
    assert res["skipped"] and "span read failed" in res["skipped"] and res["removed"] == 0


@test("loader: stock_balance is replaced per (as_of_date, warehouse); a narrower snapshot warns + audits, PARTIAL_OK does not")
def _():
    from scripts import load_supabase as ls
    prev = [{"id": i, "item_name": f"I{i}", "warehouse_name": w, "net_qty": 1, "total_value_bhd": 1, "as_of_date": "2026-09-21"}
            for i, w in enumerate(["Accessories Warehouse", "Karrar", "Devadas"])]
    same_day = [{"id": 90, "item_name": "OLD", "warehouse_name": "Accessories Warehouse", "net_qty": 5, "total_value_bhd": 5, "as_of_date": "2026-09-24"},
                {"id": 91, "item_name": "VAN", "warehouse_name": "Karrar", "net_qty": 7, "total_value_bhd": 7, "as_of_date": "2026-09-24"}]
    db = {"stock_balance": prev + same_day, "audit_log": []}
    sb = pd.DataFrame([{"item_name": "UK12", "warehouse_name": "Accessories Warehouse", "net_qty": "182", "selling_rate_bhd": "4",
                        "total_value_bhd": "728", "as_of_date": "2026-09-24", "source_file": "f"},
                       {"item_name": "UK12", "warehouse_name": "Accessories Warehouse", "net_qty": "18", "selling_rate_bhd": "4",
                        "total_value_bhd": "72", "as_of_date": "2026-09-24", "source_file": "f"}])   # Focus lists it twice
    res = ls._stock_balance_replace(_FakeDB(db), sb.copy(), partial_ok=False)
    assert res["pairs"] == [("2026-09-24", "Accessories Warehouse")] and len(res["warnings"]) == 1
    assert "covers 1 warehouse(s), the previous one 3" in res["warnings"][0]
    today = [r for r in db["stock_balance"] if r["as_of_date"] == "2026-09-24"]
    assert {(r["item_name"], r["warehouse_name"]) for r in today} == {("UK12", "Accessories Warehouse"), ("VAN", "Karrar")}
    assert next(r for r in today if r["item_name"] == "UK12")["net_qty"] == 200          # aggregated, OLD replaced, VAN kept
    assert len([r for r in db["stock_balance"] if r["as_of_date"] == "2026-09-21"]) == 3   # earlier snapshot untouched
    assert db["audit_log"][0]["event"] == "stock_balance.partial_snapshot" and db["audit_log"][0]["detail"]["previous_as_of"] == "2026-09-21"
    db = {"stock_balance": prev + same_day, "audit_log": []}
    res = ls._stock_balance_replace(_FakeDB(db), sb.copy(), partial_ok=True)
    assert res["warnings"] == [] and db["audit_log"] == []


# ── 3. verify: per day / per salesman to the fils, warehouse set, Focus total ──

@test("verify: group_sums / compare_groups / evaluate_checks -- 0.005 BHD absolute, row counts, notes")
def _():
    from scripts.verify_numbers import compare_groups, evaluate_checks, group_sums
    rows = [{"line_date": "2026-09-24", "warehouse_name": "Karrar", "gross_bhd": 1.234},
            {"line_date": "2026-09-24", "warehouse_name": "Karrar", "gross_bhd": 2.0},
            {"line_date": "2026-09-23", "warehouse_name": None, "gross_bhd": "3.5"}]
    fd = group_sums(rows, lambda r: str(r["line_date"])[:10], "gross_bhd")
    assert fd == {"2026-09-24": (3.234, 2), "2026-09-23": (3.5, 1)}, fd
    assert group_sums(rows, "warehouse_name", "gross_bhd") == {"Karrar": (3.234, 2), "(none)": (3.5, 1)}
    c = compare_groups(fd, {"2026-09-24": (3.238, 2), "2026-09-23": (3.5, 2)})
    assert c["max_diff"] == 0.004 and c["mismatched"] == [] and c["count_mismatched"] == [("2026-09-23", 1, 2)]
    assert c["file_rows"] == 3 and c["db_rows"] == 4
    c = compare_groups(fd, {"2026-09-24": (3.24, 2)})                     # a day missing on the DB side
    assert c["mismatched"] == [("2026-09-23", 3.5, 0.0), ("2026-09-24", 3.234, 3.24)]
    ok, out = evaluate_checks([("pct", 100.0, 100.4, 0.5), ("abs ok", 0.0, 0.005, 0.0, 0.005, "note here"),
                               ("abs bad", 10.0, 10.006, 0.0, 0.005, ""), ("count", 11956.0, 11956.0, 0.0, 0.0, "")])
    assert not ok and [r["passed"] for r in out] == [True, True, False, True] and out[1]["note"] == "note here" and "note" not in out[3]


@test("verify: the sales detail checks pass on a faithful DB and fail on a 0.02 BHD day, a line-count gap, a salesman gap")
def _():
    from scripts import verify_numbers as vn
    ol = [{"line_date": "2026-09-23", "warehouse_name": "Karrar", "gross_bhd": 10.0},
          {"line_date": "2026-09-24", "warehouse_name": "Karrar", "gross_bhd": 1.5},
          {"line_date": "2026-09-24", "warehouse_name": "Ahmed", "gross_bhd": 2.5}]
    od = [{"order_date": "2026-09-23", "gross_bhd": 10.0}, {"order_date": "2026-09-24", "gross_bhd": 4.0}]
    faithful = {"day": {"2026-09-23": (10.0, 1), "2026-09-24": (4.0, 2)}, "sm": {"Karrar": (11.5, 2), "Ahmed": (2.5, 1)},
                "inv": {"2026-09-23": (10.0, 1), "2026-09-24": (4.0, 1)}}
    calls: list[str] = []

    def fake_groups(sql, params, rest=None):
        calls.append(sql)
        assert params == ["2026-09-23", "2026-09-24"]
        if "FROM orders" in sql:
            return state["inv"]
        return state["sm"] if "salesman_raw" in sql else state["day"]
    saved = vn._db_groups
    vn._db_groups = fake_groups
    try:
        state = faithful
        checks: list = []
        vn._sales_detail_checks(ol, od, checks)
        ok, rows = vn.evaluate_checks(checks)
        assert ok and len(rows) == 5, [(r["metric"], r["passed"]) for r in rows]
        assert rows[0]["metric"].startswith("Sales per day: max |file - DB| BHD (2 days)") and rows[1]["report"] == 3.0 == rows[1]["db"]
        assert "gross_bhd" in calls[0] and "salesman_raw" in calls[1] and "FROM orders" in calls[2]
        state = {"day": {"2026-09-23": (10.02, 1), "2026-09-24": (4.0, 3)}, "sm": {"Karrar": (11.5, 2), "Ahmed": (2.51, 1)},
                 "inv": faithful["inv"]}
        checks = []
        vn._sales_detail_checks(ol, od, checks)
        ok, rows = vn.evaluate_checks(checks)
        assert not ok and [r["passed"] for r in rows] == [False, False, False, True, True], [(r["metric"], r["passed"]) for r in rows]
        assert "2026-09-23: file 10.0 vs DB 10.02" in rows[0]["note"] and "2026-09-24: file 2 vs DB 3" in rows[1]["note"]
        assert "Ahmed: file 2.5 vs DB 2.51" in rows[2]["note"]
        checks = []
        vn._sales_detail_checks(ol, None, checks)              # register not uploaded: only the day-book family
        assert len(checks) == 3
    finally:
        vn._db_groups = saved


@test("verify: a snapshot with fewer warehouses than the previous one FAILS, PARTIAL_OK in the folder makes it pass; row count check")
def _():
    from scripts import verify_numbers as vn
    from app import db_read
    sb = [{"item_name": "UK12", "warehouse_name": "Accessories Warehouse", "as_of_date": "2026-09-24", "total_value_bhd": 1}]

    def fake_params(sql, params):
        if "MAX(as_of_date)" in sql:
            assert params == ["2026-09-24"]
            return [{"d": "2026-07-08"}]
        if "DISTINCT" in sql:
            return [{"w": f"W{i}"} for i in range(24)] + [{"w": "Accessories Warehouse"}]
        if "COUNT(*)" in sql:
            assert params[1] == '["Accessories Warehouse"]'
            return [{"s": 1}]
        if "SUM(total_value_bhd)" in sql:
            return [{"s": 1.0}]
        raise AssertionError(sql)
    saved = (db_read.exec_sql_params, vn.get_client)
    db_read.exec_sql_params = fake_params
    vn.get_client = lambda: (_ for _ in ()).throw(AssertionError("the RPC answered; REST must not be used"))
    try:
        with tempfile.TemporaryDirectory() as d:
            checks: list = []
            vn._stock_snapshot_checks(sb, Path(d), checks)
            ok, rows = vn.evaluate_checks(checks)
            assert not ok and rows[0]["db"] == 24.0 and "covers 1 warehouse(s), the previous one 25" in rows[0]["note"], rows
            assert rows[1]["metric"] == "Stock snapshot 2026-09-24 rows (file vs DB, 1 warehouse(s))" and rows[1]["passed"]
            assert rows[2]["metric"] == "Stock snapshot 2026-09-24 value BHD (file warehouses)" and rows[2]["passed"]
            (Path(d) / "PARTIAL_OK").write_bytes(b"")
            checks = []
            vn._stock_snapshot_checks(sb, Path(d), checks)
            ok, rows = vn.evaluate_checks(checks)
            assert ok and rows[0]["metric"].endswith("(PARTIAL_OK)") and rows[0]["db"] == 0.0
            # production: yq_readonly cannot read stock_balance -> the RPC raises and REST pages by id
            db_read.exec_sql_params = lambda sql, params: (_ for _ in ()).throw(RuntimeError("permission denied for table stock_balance"))
            db = {"stock_balance": [{"id": 1, "as_of_date": "2026-07-08", "warehouse_name": "W1", "total_value_bhd": 5},
                                    {"id": 2, "as_of_date": "2026-09-24", "warehouse_name": "Accessories Warehouse", "total_value_bhd": 1},
                                    {"id": 3, "as_of_date": "2026-09-24", "warehouse_name": "Karrar", "total_value_bhd": 9}]}
            vn.get_client = lambda: _FakeDB(db)
            checks = []
            vn._stock_snapshot_checks(sb, Path(d), checks)
            ok, rows = vn.evaluate_checks(checks)
            assert ok and [(r["metric"][:20], r["report"], r["db"]) for r in rows[1:]] == \
                [("Stock snapshot 2026-", 1.0, 1.0), ("Stock snapshot 2026-", 1.0, 1.0)], rows
            assert "previous 2026-07-08 had 1" in rows[0]["note"]
    finally:
        db_read.exec_sql_params, vn.get_client = saved
    src = (ROOT / "scripts" / "verify_numbers.py").read_text(encoding="utf-8")
    assert '.order(_stable_order(view))' in src and 'return "line_id" if view == "v_sales" else "id"' in src
    from scripts.verify_numbers import _stable_order
    assert _stable_order("orders") == "id" and _stable_order("v_sales") == "line_id"


@test("verify: the Focus Grand Total is checked against ar_ageing_totals; before the migration it is an info row; the gap is a note")
def _():
    from scripts import verify_numbers as vn
    saved = {n: getattr(vn, n) for n in ("read_grid", "parse_receivables_totals", "_sql_sum")}
    vn.read_grid = lambda f: None
    vn.parse_receivables_totals = lambda g, s: {"as_of_date": "2026-09-24", "focus_total_bhd": 8633.84, "focus_over90_bhd": 4724.26,
                                                "rows_total_bhd": 9078.86, "source_file": s}
    ar = [{"balance_bhd": 9000.0}, {"balance_bhd": 78.86}]
    try:
        vn._sql_sum = lambda sql, params=None: 8633.84
        checks: list = []
        vn._receivables_total_checks("Customer_summary_ageing_by_due_date1.xlsx", ar, checks)
        ok, rows = vn.evaluate_checks(checks)
        assert ok and rows[0]["metric"] == "Receivables Focus total 2026-09-24 stored" and "gap +445.020" in rows[0]["note"]
        assert "credits may be shown as owed" in rows[0]["note"]
        vn._sql_sum = lambda sql, params=None: 0.0                  # table there, nothing stored for the date -> FAIL
        checks = []
        vn._receivables_total_checks("f", ar, checks)
        assert not vn.evaluate_checks(checks)[0]
        vn._sql_sum = lambda sql, params=None: None                 # table missing (migration pending) -> info row, passes
        checks = []
        vn._receivables_total_checks("f", ar, checks)
        ok, rows = vn.evaluate_checks(checks)
        assert ok and "migration pending" in rows[0]["metric"]
    finally:
        for n, f in saved.items():
            setattr(vn, n, f)


# ── 4. MRN costs: real receipt dates, weighted duplicates, workbook mode ──────

@test("mrn: Focus date serial decodes (132778260 = 2026-09-20); XML rows carry the header date; garbage is None")
def _():
    from scripts.ingest_mrn import decode_focus_date, parse_mrn_bytes
    assert decode_focus_date(132778260) == "2026-09-20" and decode_focus_date("132778260") == "2026-09-20"
    assert decode_focus_date((2025 << 16) | (11 << 8) | 10) == "2025-11-10"
    assert decode_focus_date(0) is None and decode_focus_date("x") is None and decode_focus_date(None) is None
    assert decode_focus_date((2026 << 16) | (13 << 8) | 1) is None
    xml = b"""<Root><Transaction><Header><DocNo>YQ-26-09-2</DocNo><Date>132778260</Date></Header>
      <HeaderExtra><IdNamePair><Name>PONo</Name><Tag>PO:YQ-26-09-2</Tag></IdNamePair></HeaderExtra>
      <BodyData>
        <TransBody><Sales><ProdCode>C18</ProdCode><Quantity>640</Quantity><StockValue>362.196</StockValue><Gross>281.6</Gross></Sales></TransBody>
        <TransBody><Sales><ProdCode>C18</ProdCode><Quantity>640</Quantity><StockValue>362.196</StockValue><Gross>281.6</Gross></Sales></TransBody>
        <TransBody><Sales><ProdCode>BE05</ProdCode><Quantity>3000</Quantity><StockValue>4024.546</StockValue><Gross>3129</Gross></Sales></TransBody>
      </BodyData></Transaction></Root>"""
    rows = parse_mrn_bytes(xml)
    assert [r["code"] for r in rows] == ["C18", "C18", "BE05"]
    assert all(r["eff"] == "2026-09-20" and r["eff_source"] == "xml" and r["mrn_no"] == "YQ-26-09-2" for r in rows)
    xml_nodate = xml.replace(b"<Date>132778260</Date>", b"")
    r0 = parse_mrn_bytes(xml_nodate)[0]
    assert r0["eff"] == "2026-09-01" and r0["eff_source"] == "month" and r0["xml_date"] is None


@test("mrn: duplicate lines fold to one per (doc, SKU) with a quantity-weighted unit cost (Decimal, 4 dp)")
def _():
    from scripts.ingest_mrn import fold_duplicate_lines
    rows = [{"code": "Big Product Display", "landed": 7.5014, "product": 6.909, "qty": 5, "doc_no": "YQ-25-11-1", "eff": "2025-11-10"},
            {"code": "big product display", "landed": 0.0128, "product": 0.012, "qty": 5, "doc_no": "YQ-25-11-1", "eff": "2025-11-10"},
            {"code": "T02", "landed": 1.8588, "product": 1.712, "qty": 1172, "doc_no": "YQ-25-11-1", "eff": "2025-11-10"},
            {"code": "C18", "landed": 0.5659, "product": 0.44, "qty": 640, "doc_no": "YQ-26-09-2", "eff": "2026-09-20"},
            {"code": "C18", "landed": 0.5659, "product": 0.44, "qty": 640, "doc_no": "YQ-26-09-2", "eff": "2026-09-20"}]
    out = fold_duplicate_lines(rows)
    assert [(r["code"], r["qty"], r["lines"]) for r in out] == [("Big Product Display", 10.0, 2), ("T02", 1172.0, 1), ("C18", 1280.0, 2)]
    assert out[0]["landed"] == 3.7571 and out[0]["product"] == 3.4605        # (37.507 + 0.064) / 10, ROUND_HALF_UP
    assert out[2]["landed"] == 0.5659 and out[1]["landed"] == 1.8588
    assert fold_duplicate_lines([]) == []


@test("mrn: receipt date = ledger move date > XML header date > 1st of month; payloads dated by it; older receipts never replace newer")
def _():
    from scripts.ingest_mrn import apply_receipt_dates, build_payloads
    rows = [{"code": "BE05", "landed": 1.3415, "product": 1.043, "qty": 3000, "doc_no": "YQ-26-09-2", "mrn_no": "YQ-26-09-2", "xml_date": "2026-09-20"},
            {"code": "X99", "landed": 0.5, "product": 0.4, "qty": 10, "doc_no": "YQ-26-09-2", "mrn_no": "YQ-26-09-2", "xml_date": "2026-09-20"},
            {"code": "T02", "landed": 1.8588, "product": 1.712, "qty": 1172, "doc_no": "YQ-25-11-1", "mrn_no": "YQ-25-11-1", "xml_date": None},
            {"code": "Z01", "landed": 0.1, "product": 0.1, "qty": 1, "doc_no": "YQ-26-06-1", "mrn_no": "YQ-26-06-1", "xml_date": "2026-06-07"}]
    dates = apply_receipt_dates(rows, {"YQ-26-09-2": "2026-09-20", "YQ-25-11-1": "2025-11-10"})
    assert dates == {"YQ-26-09-2": ("2026-09-20", "ledger"), "YQ-25-11-1": ("2025-11-10", "ledger"), "YQ-26-06-1": ("2026-06-07", "xml")}
    rows[2]["xml_date"] = None
    dates = apply_receipt_dates(rows, {})
    assert dates["YQ-25-11-1"] == ("2025-11-01", "month") and dates["YQ-26-09-2"] == ("2026-09-20", "xml")
    apply_receipt_dates(rows, {"YQ-26-09-2": "2026-09-20", "YQ-25-11-1": "2025-11-10"})
    p = build_payloads(rows, have={"X99": "2026-09-24", "T02": "2025-09-27"})
    assert p["skipped_older"] == ["X99"]                                  # a newer receipt is on file
    assert [c["sku_code"] for c in p["mrn_landed_costs"]] == ["BE05", "T02", "Z01"]
    assert {c["sku_code"]: c["effective_date"] for c in p["mrn_landed_costs"]} == {"BE05": "2026-09-20", "T02": "2025-11-10", "Z01": "2026-06-07"}
    assert p["newest"] == "2026-09-20" and p["purchase_costs"] == [{"sku_code": "BE05", "landed_cost_bhd": 1.3415, "currency": "BHD",
                                                                    "effective_date": "2026-09-20", "source_file": "MRN YQ-26-09-2"}]
    assert not any(c["effective_date"].endswith("-01") for c in p["mrn_landed_costs"] if c["sku_code"] != "T02")
    assert len(p["mrn_lines"]) == 4 and p["mrn_lines"][0] == {"doc_no": "YQ-26-09-2", "sku_code": "BE05", "qty": 3000,
                                                             "landed_unit_bhd": 1.3415, "product_unit_bhd": 1.043}


@test("mrn: load_mrn_costs dry run resolves dates from the ledger through the client and writes nothing")
def _():
    from scripts.ingest_mrn import load_mrn_costs
    db = {"stock_movements": [{"voucher": "MRN:YQ-26-09-2", "move_date": "2026-09-20"}, {"voucher": "MRN:YQ-26-09-2", "move_date": "2026-09-21"}],
          "mrn_landed_costs": [{"sku_code": "T02", "effective_date": "2026-06-30"}], "mrn_lines": [], "purchase_costs": []}
    rows = [{"code": "C18", "landed": 0.5659, "product": 0.44, "qty": 640, "doc_no": "YQ-26-09-2", "mrn_no": "YQ-26-09-2", "eff": "2026-09-01", "eff_source": "month", "xml_date": None},
            {"code": "C18", "landed": 0.5659, "product": 0.44, "qty": 640, "doc_no": "YQ-26-09-2", "mrn_no": "YQ-26-09-2", "eff": "2026-09-01", "eff_source": "month", "xml_date": None},
            {"code": "T02", "landed": 1.0, "product": 0.9, "qty": 10, "doc_no": "YQ-25-11-1", "mrn_no": "YQ-25-11-1", "eff": "2025-11-01", "eff_source": "month", "xml_date": None}]
    s = load_mrn_costs(rows, client=_FakeDB(db), dry_run=True)
    assert s["dry_run"] and s["receipt_dates"] == {"YQ-26-09-2": "2026-09-20", "YQ-25-11-1": "2025-11-01"}
    assert s["date_sources"] == {"YQ-26-09-2": "ledger", "YQ-25-11-1": "month"} and s["month_dated"] == ["YQ-25-11-1"]
    p = s["payloads"]
    assert p["mrn_landed_costs"] == [{"sku_code": "C18", "landed_cost_bhd": 0.5659, "product_cost_bhd": 0.44, "last_qty": 1280.0,
                                      "doc_no": "YQ-26-09-2", "effective_date": "2026-09-20"}]
    assert p["skipped_older"] == ["T02"] and s["folded_lines"] == 2 and s["skus"] == 1
    assert db["mrn_lines"] == [] and db["purchase_costs"] == []               # dry run: nothing written
    s = load_mrn_costs(rows, client=_FakeDB(db), dry_run=False)
    assert len(db["mrn_lines"]) == 2 and db["purchase_costs"][0]["effective_date"] == "2026-09-20"
    assert [r["sku_code"] for r in db["mrn_landed_costs"]] == ["T02", "C18"]


@test("mrn: the owner's workbook sheet parses by label (Item Code / Qty / Landing Cost / Unit / Purchase Rate), stops at TOTAL; floor gains")
def _():
    from openpyxl import Workbook
    from scripts.ingest_mrn import floor_gains, parse_workbook_sheet
    wb = Workbook()
    ws = wb.active
    ws.title = "YQ-25-11-1"
    ws.append(["MRN YQ-25-11-1 — Nov-2025"])
    ws.append(["Supplier VFAN | Purchase Order YQ-25-09-4 | Source file Transactions_1281.xml"])
    ws.append([])
    ws.append(["#", "Item Code", "Item Description", "ERP Item ID", "Qty", "Purchase Rate (BHD)", "Purchase Value (BHD)", "Freight (BHD)",
               "Customs (BHD)", "Misc (BHD)", "Vendor Freight (BHD)", "Total Charges (BHD)", "Landing Cost Total (BHD)",
               "Landing Cost / Unit (BHD)", "Loading % on Purchase"])
    ws.append([1, "BE05", "BE05 NeckBand", 80, 600, 1.013, 607.8, 51.17, 0.948, 0, 0, 52.118, 659.918, 1.09986333333333, 0.0857])
    ws.append([2, "Big Product Display", "Big Product Display", 152, 5, 6.909, 34.545, 2.908, 0.054, 0, 0, 2.962, 37.507, 7.5014, 0.0857])
    ws.append([3, "Big Product Display", "Big Product Display", 152, 5, 0.012, 0.06, 0.004, 0, 0, 0, 0.004, 0.064, 0.0128, 0.0667])
    ws.append([4, "UK03 C", "UK03 20W Charger + Type-C Cable", 31, 384, 0.676, 259.584, 21.854, 0.405, 0, 0, 22.259, 281.843, 0.733966145833333, 0.0857])
    ws.append(["TOTAL", "", "", "", 994, "", 901.989])
    ws.append(["Reconciliation against the ERP file"])
    ws.append(["Landing cost calculated above", "", "", "", "", "", 12696.408])
    wb.create_sheet("Summary")
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "YQ_MRN_Landing_Cost_Analysis (2).xlsx"
        wb.save(p)
        rows = parse_workbook_sheet(p, "YQ-25-11-1")
        assert [(r["code"], r["qty"]) for r in rows] == [("BE05", 600.0), ("Big Product Display", 5.0), ("Big Product Display", 5.0), ("UK03 C", 384.0)]
        assert rows[0]["landed"] == 1.09986333333333 and rows[0]["product"] == 1.013 and rows[0]["doc_no"] == "YQ-25-11-1"
        assert rows[0]["eff"] == "2025-11-01" and rows[0]["eff_source"] == "month"
        try:
            parse_workbook_sheet(p, "Summary")
            raise AssertionError("must raise: no header")
        except ValueError as e:
            assert "header" in str(e)
        try:
            parse_workbook_sheet(p, "YQ-99-99-9")
            raise AssertionError("must raise: no sheet")
        except ValueError as e:
            assert "not in workbook" in str(e)
    gains = floor_gains([r["code"] for r in rows], {"BE05", "T02"})
    assert gains == ["Big Product Display", "UK03 C"]
    assert floor_gains(["be05", " ", None], {"BE05"}) == []


# ── 5. alias autofill and the unmapped-share warning ─────────────────────────

@test("alias autofill: exact-name (1.0) proposals only, insert-only; the share is measured after; > 2% emits a warn event once")
def _():
    from scripts.alias_autofill import (DEFAULT_HOLD_BACK, EVENT_TYPE, autofill, hold_back_codes, parse_hold_back,
                                        pick_exact, summary_line, unmapped_share)
    props = [{"item_name": "T17 Open Ear Clip On - Type-C Port (VFAN) T17", "proposed_code": "T17", "product_id": 11, "confidence": 1.0, "rev_90": 766.7},
             {"item_name": "X24 CC 1Mtr something new", "proposed_code": "X24 CC 1Mtr", "product_id": 12, "confidence": 0.9, "rev_90": 5.0},
             {"item_name": "no product row", "proposed_code": "ZZ", "product_id": None, "confidence": 1.0, "rev_90": 1.0},
             {"item_name": "", "proposed_code": "T17", "product_id": 11, "confidence": 1.0, "rev_90": 0},
             {"item_name": "X24 CL 1Mtr Cable (Type-C to Lightning) (VFAN) X24 CL", "proposed_code": "X24 CL 1Mtr", "product_id": 14, "confidence": 1.0, "rev_90": 27.95}]
    assert [p["proposed_code"] for p in pick_exact(props)] == ["T17", "X24 CL 1Mtr"]
    assert [p["proposed_code"] for p in pick_exact(props, DEFAULT_HOLD_BACK)] == ["T17"]      # the owner's hold-back
    assert parse_hold_back(" x24 cc 1mtr ,X24 CL 1Mtr,, ") == {"X24 CC 1MTR", "X24 CL 1MTR"} and parse_hold_back(None) == set()
    assert hold_back_codes(_FakeDB({"app_settings": [{"key": "alias_autofill_exclude", "value": "T17"}]})) == {"T17"}
    assert hold_back_codes(_FakeDB({"app_settings": []})) == {"X24 CC 1MTR", "X24 CL 1MTR"}       # not seeded yet: the default
    assert hold_back_codes(_FakeDB({})) == {"X24 CC 1MTR", "X24 CL 1MTR"}                          # unreadable: the default
    s = unmapped_share({"lines": 682, "unmapped": 33, "rev": 1000.0, "rev_unmapped": 40.5, "data_to": "2026-09-24"})
    assert s["share"] == 0.0484 and s["over"] and s["rev_share"] == 0.0405 and s["data_to"] == "2026-09-24"
    assert not unmapped_share({"lines": 100, "unmapped": 2})["over"] and unmapped_share(None) == {
        "lines": 0, "unmapped": 0, "share": 0.0, "rev": 0.0, "rev_unmapped": 0.0, "rev_share": 0.0, "data_to": None, "over": False}

    calls: list[str] = []
    state = {"inserted": False}

    def fake_sql(sql):
        calls.append(sql)
        if "sku_code IS NULL AND sale_date" in sql:               # UNALIASED_SQL
            return [{"item_name": props[0]["item_name"], "division": "Accessories", "lines": 47, "units": 177, "rev_365": 1200.5, "rev_90": 766.7, "last_sold": "2026-09-20"},
                    {"item_name": "X24 CC 1Mtr Cable (Type-C to Type-C) (VFAN) X24 CC", "division": "Accessories", "lines": 8, "units": 8, "rev_365": 8, "rev_90": 8, "last_sold": "2026-09-22"},
                    {"item_name": "X24 CL 1Mtr Cable (Type-C to Lightning) (VFAN) X24 CL", "division": "Accessories", "lines": 8, "units": 8, "rev_365": 28, "rev_90": 28, "last_sold": "2026-09-22"}]
        if sql.startswith("SELECT id, sku_code, item_name FROM products"):
            return [{"id": 11, "sku_code": "T17", "item_name": props[0]["item_name"]},
                    {"id": 12, "sku_code": "X24 CC", "item_name": "X24 2Mtr Cable (Type-C to Type-C) (VFAN) X24 CC"},
                    {"id": 13, "sku_code": "X24 CC 1Mtr", "item_name": "X24 CC 1Mtr Cable (Type-C to Type-C) (VFAN) X24 CC"},
                    {"id": 14, "sku_code": "X24 CL 1Mtr", "item_name": "X24 CL 1Mtr Cable (Type-C to Lightning) (VFAN) X24 CL"}]
        if "COUNT(*) FILTER (WHERE sku_code IS NULL)" in sql:
            return [{"lines": 682, "unmapped": 33 if not state["inserted"] else 30, "rev": 1000.0, "rev_unmapped": 40.0, "data_to": "2026-09-24"}]
        if "GROUP BY item_name ORDER BY lines DESC" in sql:
            return [{"item_name": "ZTE-A35e MOBILE", "lines": 17, "rev": 0}]
        raise AssertionError(sql)
    # the setting holds back X24 CL 1Mtr only (an owner edit); X24 CC 1Mtr is then written
    db = {"product_aliases": [{"id": 1, "alias_text": "old", "product_id": 5}], "audit_log": [],
          "app_settings": [{"key": "alias_autofill_exclude", "value": "X24 CL 1Mtr"}]}
    emitted: list = []

    def fake_emit(emitter, event_type, **kw):
        emitted.append((emitter, event_type, kw))
        return True

    class _DB(_FakeDB):
        def table(self, name):
            q = super().table(name)
            if name == "product_aliases":
                state["inserted"] = True
            return q
    res = autofill(exec_sql=fake_sql, client=_DB(db), emit=fake_emit)
    assert res["picked"] == 2 and res["inserted"] == 2 and res["codes"] == ["T17", "X24 CC 1Mtr"] and res["held_back"] == ["X24 CL 1Mtr"], res
    assert {r["alias_text"] for r in db["product_aliases"]} == {"old", props[0]["item_name"], "X24 CC 1Mtr Cable (Type-C to Type-C) (VFAN) X24 CC"}
    assert db["audit_log"][0]["event"] == "product_aliases.autofill" and db["audit_log"][0]["detail"]["inserted"] == 2
    # no setting yet (before the migration): the default hold-back keeps BOTH X24 near-duplicates out
    db0 = {"product_aliases": [], "audit_log": [], "app_settings": []}
    res0 = autofill(exec_sql=fake_sql, client=_FakeDB(db0), emit=fake_emit)
    assert res0["codes"] == ["T17"] and res0["held_back"] == ["X24 CC 1Mtr", "X24 CL 1Mtr"] and res0["inserted"] == 1, res0
    assert "held back (owner decision): X24 CC 1Mtr, X24 CL 1Mtr" in summary_line(res0)
    assert res["warned"] and emitted[0][1] == EVENT_TYPE and emitted[0][2]["severity"] == "warn"
    assert emitted[0][2]["fingerprint"] == f"{EVENT_TYPE}:2026-09-24:4.4" and emitted[0][2]["payload"]["top"][0]["item_name"] == "ZTE-A35e MOBILE"
    assert "4.4% of the last 30 days' accessory lines" in emitted[0][2]["payload"]["summary"]
    line = summary_line(res)
    assert line.startswith("aliases: 2 exact-name inserted (T17, X24 CC 1Mtr); held back (owner decision): X24 CL 1Mtr") and "OVER 2%, warning event emitted" in line
    # dry run: nothing inserted, nothing emitted, the share still reported
    state["inserted"] = False
    db2 = {"product_aliases": [], "audit_log": [], "app_settings": []}
    emitted.clear()
    res = autofill(exec_sql=fake_sql, client=_FakeDB(db2), emit=fake_emit, dry_run=True)
    assert res["inserted"] == 0 and db2["product_aliases"] == [] and emitted == [] and res["share"]["over"] and not res["warned"]
    # a read failure never raises past the loader
    res = autofill(exec_sql=lambda s: (_ for _ in ()).throw(RuntimeError("RPC down")), client=_FakeDB({}), emit=fake_emit)
    assert res["error"].startswith("read failed") and res["inserted"] == 0
    src = (ROOT / "scripts" / "refresh.py").read_text(encoding="utf-8")
    assert "from scripts.alias_autofill import autofill, summary_line" in src and "Aliases: " in src
    ev = (ROOT / "app" / "events.py").read_text(encoding="utf-8")
    assert '"data.unmapped_share": [_react_notify_warn]' in ev


# ── 6. margin truth: agents, digest, templates, reports ──────────────────────

_UK03 = {"item_name": "UK03 20W Charger (USB + Type-C Port) (VFAN)", "sku_code": "UK03", "product_name": "UK03", "category_name": "CHARGER",
         "list_price_bhd": 1.2, "net_amount_bhd": 439.8, "cogs_bhd": 483.2, "gross_profit_bhd": 43.4, "np_margin_pct": 980.03,
         "focus_gp_margin_pct": 980.03, "net_ex_vat_bhd": 399.58, "gp_computed_bhd": -43.4, "gp_ex_vat_bhd": -83.62,
         "margin_ex_vat_pct": -20.93, "is_below_cost": True, "ex_vat_source": "day_book"}


@test("margin truth: the migrated view is read first; before it the inline formula runs; every row carries the computed gp_margin_pct")
def _():
    from app import margin_truth as mt
    seen: list[str] = []

    def q_view(sql):
        seen.append(sql)
        return [dict(_UK03)]
    rows, basis = mt.margin_rows("is_below_cost", "gp_ex_vat_bhd ASC", 20, q=q_view)
    assert basis == "view" and rows[0]["gp_margin_pct"] == -20.93 and rows[0]["is_below_cost"] is True
    assert "gp_margin_pct AS focus_gp_margin_pct" in seen[0] and "WHERE cogs_bhd IS NOT NULL AND net_amount_bhd IS NOT NULL AND (is_below_cost)" in seen[0]

    def q_old(sql):
        seen.append(sql)
        if "net_amount_bhd / 1.1" not in sql:
            raise RuntimeError('column "is_below_cost" does not exist')
        return [{**_UK03, "net_ex_vat_bhd": 399.818, "gp_ex_vat_bhd": -83.382, "margin_ex_vat_pct": -20.85, "ex_vat_source": "vat_rate"}]
    rows, basis = mt.margin_rows("is_below_cost", "gp_ex_vat_bhd ASC", 20, q=q_old)
    assert basis == "inline" and rows[0]["gp_margin_pct"] == -20.85 and "(net_amount_bhd / 1.1 < cogs_bhd) AS is_below_cost" in seen[-1]
    assert "ORDER BY gp_ex_vat_bhd ASC LIMIT 20" in seen[-1]
    assert mt.margin_rows("", limit=5, q=lambda s: [])[0] == []


@test("margin truth: anomaly_scan / fraud_scan / margin_guardian / digest flag UK03 on the computed margin, never Focus's GP %")
def _():
    from app import agents, digest, margin_truth as mt
    saved_q, saved_x = agents._q, mt.exec_sql
    calls: list[str] = []

    def fake(sql):
        calls.append(sql)
        if "v_product_margin" in sql:
            if "margin_ex_vat_pct < 5.0" in sql:
                return [dict(_UK03), {**_UK03, "item_name": "THIN", "gp_ex_vat_bhd": 1.0, "margin_ex_vat_pct": 2.5, "is_below_cost": False}]
            return [dict(_UK03)]
        return []
    agents._q = lambda sql: []
    mt.exec_sql = fake
    try:
        a = agents.anomaly_scan()
        assert a["priced_below_cost"][0]["item_name"].startswith("UK03") and a["anomaly_count"] == 1 and "1 below cost" in a["summary"]
        assert a["priced_below_cost"][0]["gp_margin_pct"] == -20.93
        f = agents.fraud_scan()
        assert f["sold_below_cost"][0]["item_name"].startswith("UK03")
        m = agents.margin_guardian()
        assert m["negative_count"] == 1 and m["thin_count"] == 1 and m["thin_margins"][0]["item_name"] == "THIN"
        d = digest.negative_margins()
        assert d[0]["item_name"].startswith("UK03") and d[0]["gp_margin_pct"] < 0
        assert not any("gp_margin_pct < 0" in c for c in calls), "the Focus GP % must never be the below-cost rule"
    finally:
        agents._q, mt.exec_sql = saved_q, saved_x
    src = (ROOT / "app" / "agents.py").read_text(encoding="utf-8") + (ROOT / "app" / "digest.py").read_text(encoding="utf-8")
    assert "WHERE gp_margin_pct < 0" not in src


@test("templates: the two margin templates compute the ex-VAT margin inline and pass the SQL validator; eval labels unchanged")
def _():
    from app.sql_validator import validate
    from app.templates import match
    label, sql = match("Which products have negative margins?")
    assert label == "Negative-margin products" and "net_amount_bhd / 1.1 < cogs_bhd" in sql and "gp_margin_pct" not in sql
    assert "margin_ex_vat_pct" in validate(sql)
    label, sql = match("Show all product margins")
    assert label == "Product margins" and "margin_ex_vat_pct" in sql and "gp_margin_pct" not in sql
    validate(sql)
    assert match("Show loss-making items")[0] == "Negative-margin products"


@test("reports: margins() counts below-cost on the computed margin; receivables() shows the Focus total and the gap; inventory() carries reserved")
def _():
    from app import margin_truth as mt, reports
    saved = (reports.exec_sql, mt.exec_sql)

    def fake(sql):
        if "FROM v_product_margin" in sql:
            return [dict(_UK03), {**_UK03, "item_name": "OK", "net_amount_bhd": 110.0, "cogs_bhd": 50.0, "net_ex_vat_bhd": 100.0,
                                  "gp_computed_bhd": 60.0, "gp_ex_vat_bhd": 50.0, "margin_ex_vat_pct": 50.0, "is_below_cost": False}]
        if "FROM v_receivables" in sql:
            return [{"account": "A", "group_name": "Retail", "outstanding_bhd": 9000.0, "overdue_bhd": 0, "over_90_bhd": 0,
                     "b_0_30": 9000.0, "b_31_60": 0, "b_61_90": 0, "b_91_120": 0, "b_121_150": 0, "b_151_180": 0, "b_181_210": 0, "b_over_210": 0},
                    {"account": "B", "group_name": "Retail", "outstanding_bhd": 78.86, "overdue_bhd": 28.0, "over_90_bhd": 28.0,
                     "b_0_30": 50.86, "b_31_60": 0, "b_61_90": 0, "b_91_120": 28.0, "b_121_150": 0, "b_151_180": 0, "b_181_210": 0, "b_over_210": 0}]
        if "FROM ar_ageing_totals" in sql:
            if state["totals"] == "missing":
                raise RuntimeError("relation ar_ageing_totals does not exist")
            return state["totals"]
        if "FROM v_catalog_reserved WHERE reserved > 0" in sql:
            if state["reserved"] is None:
                raise RuntimeError("relation v_catalog_reserved does not exist")
            return state["reserved"]
        if "FROM v_catalog_reserved LIMIT 1" in sql:
            return [{"r": 12, "t": 0, "n": 1, "o": 2, "as_of": "2026-09-24"}]
        raise AssertionError(sql)
    reports.exec_sql = mt.exec_sql = fake
    state = {"totals": [{"as_of_date": "2026-09-24", "focus_total_bhd": 8633.84, "focus_over90_bhd": 4724.26, "rows_total_bhd": 9078.86}],
             "reserved": [{"item_code": "X01", "on_hand": 100, "reserved": 12, "available": 88, "in_transit": 0, "open_orders": 2, "stock_as_of": "2026-09-24"}]}
    try:
        m = reports.margins()
        assert m["negative_count"] == 1 and m["basis"] == "view" and m["rows"][0]["is_below_cost"] and round(m["gp_pct"], 2) == round((50.0 - 83.62) / 499.58 * 100, 2)
        assert m["total_gp_bhd"] == -33.62 and m["total_net_ex_vat_bhd"] == 499.58 and m["total_gp_report_basis_bhd"] == 16.6
        r = reports.receivables()
        assert r["total"] == 9078.86 and r["focus_total"] == 8633.84 and r["focus_gap"] == 445.02 and r["focus_as_of"] == "2026-09-24"
        assert r["focus_over_90"] == 4724.26 and r["over_90"] == 28.0
        state["totals"] = "missing"
        r = reports.receivables()
        assert r["focus_total"] is None and r["focus_gap"] is None and r["total"] == 9078.86
        state["totals"] = []
        assert reports.receivables()["focus_total"] is None
        rs = reports.reserved_stock()
        assert rs["available"] and rs["units"] == 12.0 and rs["items"] == 1 and rs["rows"][0]["available"] == 88 and rs["stock_as_of"] == "2026-09-24"
        state["reserved"] = None
        assert reports.reserved_stock() == {"available": False, "rows": [], "units": 0, "in_transit": 0, "items": 0, "stock_as_of": None}
    finally:
        reports.exec_sql, mt.exec_sql = saved
    src = (ROOT / "app" / "reports.py").read_text(encoding="utf-8")
    assert '"reserved": reserved_stock(),' in src


@test("catalog: reserved / available attach to the staff catalog only; a salesman never receives available (on-hand)")
def _():
    from app import catalog
    saved = catalog.exec_sql
    catalog.exec_sql = lambda sql: [{"item_code": "X01", "reserved": 12, "available": 88, "in_transit": 0}]
    try:
        rows = [{"item_code": "X01"}, {"item_code": "T02"}]
        catalog._attach_reserved(rows, "admin")
        assert rows[0] == {"item_code": "X01", "reserved": 12.0, "in_transit": 0.0, "available": 88.0} and rows[1] == {"item_code": "T02"}
        rows = [{"item_code": "X01"}]
        catalog._attach_reserved(rows, "salesman")
        assert rows[0] == {"item_code": "X01", "reserved": 12.0, "in_transit": 0.0}
        catalog.exec_sql = lambda sql: (_ for _ in ()).throw(RuntimeError("no view"))
        rows = [{"item_code": "X01"}]
        catalog._attach_reserved(rows, "admin")
        assert rows == [{"item_code": "X01"}]
    finally:
        catalog.exec_sql = saved
    shop_src = (ROOT / "app" / "shop.py").read_text(encoding="utf-8")
    assert "v_catalog_reserved" not in shop_src, "merchants keep on-hand status: app/shop.py must not read the reserved view"


@test("migration text: economics_v2 is additive (no delete / drop table / truncate), keeps columns, appends the new ones, self-checks UK03; reverse restores")
def _():
    sql = (ROOT / "scripts" / "economics_v2_migration.sql").read_text(encoding="utf-8")
    low = sql.lower()
    assert not any(s in low for s in ("delete from", "drop table", "truncate", "drop view"))
    for v in ("v_product_economics", "v_product_margin", "v_catalog_reserved"):
        assert low.count(f"create or replace view {v} as") == 1, v
    assert "create table if not exists ar_ageing_totals" in low and "focus_total_bhd" in low and "focus_over90_bhd" in low
    assert "from mrn_landed_costs m" in low and "from purchase_costs p" in low and "order by x.pri, x.cost_effective_date desc nulls last, x.id desc" in low
    assert "voided_at is null" in low and "gp_computed_bhd" in low and "margin_ex_vat_pct" in low and "is_below_cost" in low
    assert "net_amount_bhd / 1.1" in low and "sum(ol.taxable_bhd)" in low
    assert "at time zone 'asia/bahrain'" in low and "o.issued_at is null" in low and "coalesce(o.is_test, false) = false" in low
    assert "revoke all on v_product_economics, v_product_margin, v_catalog_reserved, ar_ageing_totals from anon, authenticated" in low
    assert "grant select on v_product_economics, v_product_margin, v_catalog_reserved, ar_ageing_totals to yq_readonly" in low
    assert "uk03 20w charger (usb%" in low and "raise exception" in low and "grantee in ('anon', 'authenticated')" in low
    assert "cols[1:15] <> array['item_name', 'report_date'" in low
    rev = (ROOT / "scripts" / "economics_v2_reverse.sql").read_text(encoding="utf-8").lower()
    assert "drop view if exists v_catalog_reserved" in rev and "drop table if exists ar_ageing_totals" in rev
    assert rev.count("create view v_product_economics as") == 1 and rev.count("create view v_product_margin as") == 1
    assert "gp_computed_bhd" not in rev and "voided_at is null" in rev and "expected the 6 original columns" in rev


# ── 7. local replay Postgres (SKIP when the cluster / database is not there) ──

def _local_conn():
    try:
        import psycopg
        conn = psycopg.connect(LOCAL_URL, connect_timeout=3)
        conn.autocommit = False
        return conn
    except Exception as e:  # noqa: BLE001
        print(f"    SKIP: local replay database not reachable ({str(e)[:80]})")
        return None


def _q(conn, sql, params=None):
    with conn.cursor() as cur:
        cur.execute(sql, params)
        if cur.description:
            cols = [d.name for d in cur.description]
            return [dict(zip(cols, r)) for r in cur.fetchall()]
        return []


@test("local: economics_v2 migration applies on the production schema copy -- UK03 flagged, MRN cost preferred, reverse restores")
def _():
    conn = _local_conn()
    if conn is None:
        return
    from decimal import Decimal
    try:
        _q(conn, (ROOT / "scripts" / "economics_v2_migration.sql").read_text(encoding="utf-8"))
        uk = _q(conn, "SELECT net_amount_bhd, net_ex_vat_bhd, cogs_bhd, gp_computed_bhd, gp_ex_vat_bhd, margin_ex_vat_pct, is_below_cost, "
                      "ex_vat_source, gross_profit_bhd, gp_margin_pct FROM v_product_margin WHERE item_name LIKE 'UK03 20W Charger (USB%'")
        if uk:
            u = uk[0]
            assert u["is_below_cost"] is True and u["gp_ex_vat_bhd"] < 0 and u["gp_computed_bhd"] == u["net_amount_bhd"] - u["cogs_bhd"], u
            assert u["net_ex_vat_bhd"] < u["net_amount_bhd"] and u["margin_ex_vat_pct"] < 0
            assert u["gross_profit_bhd"] > 0, "the Focus column is sign-less on this row (the reason the view computes its own)"
            print(f"    UK03: net {u['net_amount_bhd']} ex-VAT {u['net_ex_vat_bhd']} ({u['ex_vat_source']}) COGS {u['cogs_bhd']} "
                  f"GP ex-VAT {u['gp_ex_vat_bhd']} margin {u['margin_ex_vat_pct']}% (Focus said GP {u['gross_profit_bhd']}, {u['gp_margin_pct']}%)")
        n = _q(conn, "SELECT COUNT(*) AS n, COUNT(*) FILTER (WHERE is_below_cost) AS below, COUNT(*) FILTER (WHERE ex_vat_source='day_book') AS db "
                     "FROM v_product_margin")[0]
        print(f"    v_product_margin: {n['n']} rows, {n['below']} below cost, {n['db']} costed from the day book's taxable totals")
        assert n["below"] >= 1
        econ = _q(conn, "SELECT sku_code, cost_bhd, cost_source, cost_effective_date, cost_doc_no FROM v_product_economics WHERE sku_code IN ('BE05','M04','X23','TB-D10') ORDER BY 1")
        for e in econ:
            assert e["cost_source"] == "mrn" and e["cost_doc_no"] == "YQ-26-09-2", e
        assert any(e["sku_code"] == "BE05" and e["cost_bhd"] == Decimal("1.3415") for e in econ), econ
        cov = _q(conn, "SELECT cost_source, COUNT(*) AS n FROM v_product_economics GROUP BY 1 ORDER BY 1")[0:3]
        print("    v_product_economics cost sources:", {c["cost_source"]: c["n"] for c in cov})
        stale = _q(conn, "SELECT COUNT(*) AS n FROM v_product_economics e WHERE e.cost_source = 'purchase_costs' AND EXISTS "
                         "(SELECT 1 FROM mrn_landed_costs m WHERE upper(m.sku_code) = upper(e.sku_code) AND m.landed_cost_bhd > 0)")[0]["n"]
        assert stale == 0
        r = _q(conn, "SELECT COUNT(*) AS n, COALESCE(SUM(reserved),0) AS reserved, COALESCE(SUM(in_transit),0) AS transit, MAX(stock_as_of)::text AS as_of FROM v_catalog_reserved")[0]
        print(f"    v_catalog_reserved: {r['n']} active codes, reserved {r['reserved']} u, in transit {r['transit']} u, stock as of {r['as_of']}")
        assert r["n"] > 0
        _q(conn, "INSERT INTO ar_ageing_totals (as_of_date, focus_total_bhd, focus_over90_bhd, rows_total_bhd, source_file) VALUES ('2026-09-24', 8633.84, 4724.26, 9078.86, 'test')")
        assert _q(conn, "SELECT focus_total_bhd FROM ar_ageing_totals")[0]["focus_total_bhd"] == Decimal("8633.840")
        # reverse inside the same transaction: the original shapes come back
        _q(conn, (ROOT / "scripts" / "economics_v2_reverse.sql").read_text(encoding="utf-8"))
        cols = [c["column_name"] for c in _q(conn, "SELECT column_name FROM information_schema.columns WHERE table_name='v_product_margin' ORDER BY ordinal_position")]
        assert len(cols) == 15 and "is_below_cost" not in cols
        assert _q(conn, "SELECT to_regclass('public.v_catalog_reserved') AS r")[0]["r"] is None
    finally:
        conn.rollback()
        conn.close()


@test("local: v_catalog_reserved counts open, un-issued, non-test orders created after the snapshot's end of day in Bahrain only")
def _():
    conn = _local_conn()
    if conn is None:
        return
    try:
        _q(conn, (ROOT / "scripts" / "economics_v2_migration.sql").read_text(encoding="utf-8"))
        code = _q(conn, "SELECT item_code FROM v_catalog_reserved WHERE on_hand > 50 ORDER BY item_code LIMIT 1")
        if not code:
            print("    SKIP: no stocked catalog code in the replay copy")
            return
        code = code[0]["item_code"]
        as_of = _q(conn, "SELECT MAX(as_of_date)::text AS d FROM stock_balance")[0]["d"]
        before = _q(conn, "SELECT on_hand, reserved, available FROM v_catalog_reserved WHERE item_code = %s", (code,))[0]
        def order(no, status, created, issued=None, is_test=False, qty=5):
            _q(conn, "INSERT INTO shop_orders (order_no, token, status, customer_name, customer_phone, created_at, issued_at, is_test) "
                     "VALUES (%s, %s, %s, 'TEST', '000', %s::timestamptz, %s::timestamptz, %s)", (no, no, status, created, issued, is_test))
            oid = _q(conn, "SELECT id FROM shop_orders WHERE order_no = %s", (no,))[0]["id"]
            _q(conn, "INSERT INTO shop_order_lines (order_id, item_code, qty, unit_price_bhd, line_total_bhd) VALUES (%s, %s, %s, 1, %s)", (oid, code, qty, qty))
        after_eod = f"{as_of} 21:30:00+03"          # 21:30 Bahrain on the snapshot day = same day: NOT after EOD
        next_day = _q(conn, "SELECT (%s::date + 1)::text AS d", (as_of,))[0]["d"] + " 00:00:01+03"
        order("R2T-1", "new", next_day)                       # counts: 5
        order("R2T-2", "confirmed", next_day, qty=7)          # counts: 7
        order("R2T-3", "new", after_eod)                      # same day as the snapshot: not after EOD
        order("R2T-4", "new", next_day, issued=next_day)      # issued to a rep's van: not reserved
        order("R2T-5", "new", next_day, is_test=True)         # test order
        order("R2T-6", "delivered", next_day)                 # closed
        order("R2T-7", "cancelled", next_day)
        after = _q(conn, "SELECT on_hand, reserved, available, open_orders FROM v_catalog_reserved WHERE item_code = %s", (code,))[0]
        assert after["reserved"] - before["reserved"] == 12 and after["available"] == after["on_hand"] - after["reserved"], (before, after)
        assert after["on_hand"] == before["on_hand"]
        print(f"    {code}: on hand {after['on_hand']}, reserved +12 from 2 of 7 synthetic orders (open, un-issued, real, after EOD Bahrain)")
        # in transit: a procurement order raised with the vendor, keyed by catalog code
        _q(conn, "INSERT INTO procurement_orders (title, vendor, stage, lines) VALUES ('t', 'VFAN', 'paid', %s::jsonb)", (f'[{{"item": "{code}", "qty": 300}}, {{"item_code": "NOPE-X", "qty": 5}}]',))
        _q(conn, "INSERT INTO procurement_orders (title, vendor, stage, lines) VALUES ('t2', 'VFAN', 'received', %s::jsonb)", (f'[{{"item": "{code}", "qty": 999}}]',))
        t = _q(conn, "SELECT in_transit FROM v_catalog_reserved WHERE item_code = %s", (code,))[0]["in_transit"]
        assert t == 300, t
    finally:
        conn.rollback()
        conn.close()


class _PgQ:
    """A PostgREST-shaped client over psycopg for the local replay: the methods scripts/load_supabase.py,
    scripts/verify_numbers.py and scripts/alias_autofill.py use, and nothing else."""

    def __init__(self, conn, table, types):
        self.conn, self.t, self.types = conn, table, types
        self.op, self.cols, self.where, self.params = "select", "*", [], []
        self.order_by, self.rng, self.lim, self.count, self.neg = [], None, None, None, False
        self.payload, self.on_conflict, self.ignore_dup = None, None, False

    def _cast(self, c):
        return f"::{self.types.get(c, 'text')}"

    def select(self, cols="*", count=None, **k):
        self.op, self.cols, self.count = "select", cols, count; return self

    def _f(self, c, op, v):
        self.where.append(f'"{c}" {op} %s{self._cast(c)}'); self.params.append(v); return self

    def eq(self, c, v): return self._f(c, "=", v)
    def gte(self, c, v): return self._f(c, ">=", v)
    def lte(self, c, v): return self._f(c, "<=", v)
    def lt(self, c, v): return self._f(c, "<", v)
    def gt(self, c, v): return self._f(c, ">", v)

    def ilike(self, c, v): return self._f(c, "ILIKE", v)

    def in_(self, c, vals):
        self.where.append(f'"{c}" = ANY(%s)'); self.params.append(list(vals)); return self

    def is_(self, c, v):
        neg, self.neg = self.neg, False
        self.where.append(f'"{c}" IS {"NOT " if neg else ""}NULL'); return self

    @property
    def not_(self):
        self.neg = True; return self

    def order(self, col, desc=False, **k):
        self.order_by.append(f'"{col}" {"DESC" if desc else "ASC"}'); return self

    def limit(self, n, **k):
        self.lim = n; return self

    def range(self, a, b):
        self.rng = (a, b); return self

    def delete(self, count=None, **k):
        self.op, self.count = "delete", count; return self

    def upsert(self, rows, on_conflict=None, ignore_duplicates=False, **k):
        self.op, self.payload, self.on_conflict, self.ignore_dup = "upsert", rows, on_conflict, ignore_duplicates; return self

    def insert(self, payload):
        self.op, self.payload = "insert", payload; return self

    def update(self, payload):
        self.op, self.payload = "update", payload; return self

    def _val(self, c, v):
        from psycopg.types.json import Jsonb
        if isinstance(v, (dict, list)):
            return Jsonb(v)
        if isinstance(v, float) and v != v:
            return None
        return v

    def execute(self):
        import types as _t
        cur = self.conn.cursor()
        w = (" WHERE " + " AND ".join(self.where)) if self.where else ""
        if self.op == "select":
            cols = "*" if self.cols == "*" else ", ".join(f'"{c.strip()}"' for c in self.cols.split(","))
            sql = f'SELECT {cols} FROM "{self.t}"{w}'
            if self.order_by:
                sql += " ORDER BY " + ", ".join(self.order_by)
            if self.rng:
                sql += f" OFFSET {self.rng[0]} LIMIT {self.rng[1] - self.rng[0] + 1}"
            elif self.lim is not None:
                sql += f" LIMIT {self.lim}"
            cur.execute(sql, self.params)
            names = [d.name for d in cur.description]
            data = [dict(zip(names, r)) for r in cur.fetchall()]
            for r in data:
                for k2, v in r.items():
                    if hasattr(v, "isoformat"):
                        r[k2] = v.isoformat()
            n = None
            if self.count == "exact":
                cur.execute(f'SELECT COUNT(*) FROM "{self.t}"{w}', self.params)
                n = cur.fetchone()[0]
            return _t.SimpleNamespace(data=data, count=n)
        if self.op == "delete":
            cur.execute(f'DELETE FROM "{self.t}"{w} RETURNING *', self.params)
            names = [d.name for d in cur.description]
            data = [dict(zip(names, r)) for r in cur.fetchall()]
            return _t.SimpleNamespace(data=data, count=len(data))
        if self.op == "update":
            sets = ", ".join(f'"{c}" = %s{self._cast(c)}' for c in self.payload)
            cur.execute(f'UPDATE "{self.t}" SET {sets}{w} RETURNING 1', [self._val(c, v) for c, v in self.payload.items()] + self.params)
            return _t.SimpleNamespace(data=[], count=cur.rowcount)
        rows = self.payload if isinstance(self.payload, list) else [self.payload]
        if not rows:
            return _t.SimpleNamespace(data=[], count=0)
        cols = list(rows[0].keys())
        col_sql = ", ".join(f'"{c}"' for c in cols)
        ph = ", ".join(f"%s{self._cast(c)}" for c in cols)
        sql = f'INSERT INTO "{self.t}" ({col_sql}) VALUES ({ph})'
        if self.op == "upsert" and self.on_conflict:
            keys = [k.strip() for k in self.on_conflict.split(",")]
            if self.ignore_dup:
                sql += f" ON CONFLICT ({', '.join(chr(34) + k + chr(34) for k in keys)}) DO NOTHING"
            else:
                sets = ", ".join(f'"{c}" = EXCLUDED."{c}"' for c in cols if c not in keys) or f'"{keys[0]}" = EXCLUDED."{keys[0]}"'
                sql += f" ON CONFLICT ({', '.join(chr(34) + k + chr(34) for k in keys)}) DO UPDATE SET {sets}"
        cur.executemany(sql, [[self._val(c, r.get(c)) for c in cols] for r in rows])
        return _t.SimpleNamespace(data=rows, count=len(rows))


class _PgClient:
    def __init__(self, conn):
        self.conn, self._types = conn, {}

    def table(self, name):
        if name not in self._types:
            with self.conn.cursor() as cur:
                cur.execute("SELECT column_name, data_type FROM information_schema.columns WHERE table_schema='public' AND table_name=%s", (name,))
                rows = cur.fetchall()
            if not rows:
                raise RuntimeError(f"relation {name} does not exist")
            self._types[name] = {c: ("numeric" if t == "numeric" else "date" if t == "date" else "timestamptz" if t.startswith("timestamp")
                                     else "boolean" if t == "boolean" else "bigint" if t in ("bigint", "integer") else "jsonb" if t == "jsonb" else "text")
                                 for c, t in rows}
        return _PgQ(self.conn, name, self._types[name])

    def rpc(self, *a, **k):
        raise RuntimeError("no RPC on the local replay")


def _local_exec(conn):
    """exec_sql / exec_sql_params over psycopg ($1.. -> named placeholders; params as text)."""
    import re

    def rows_of(cur):
        names = [d.name for d in cur.description]
        out = [dict(zip(names, r)) for r in cur.fetchall()]
        for r in out:
            for k, v in r.items():
                if hasattr(v, "isoformat"):
                    r[k] = v.isoformat()
        return out

    def exec_sql(sql):
        with conn.cursor() as cur:
            cur.execute(sql)
            return rows_of(cur)

    def exec_sql_params(sql, params):
        named = re.sub(r"\$(\d+)", lambda m: f"%(p{m.group(1)})s", sql)
        with conn.cursor() as cur:
            cur.execute(named, {f"p{i + 1}": str(p) for i, p in enumerate(params)})
            return rows_of(cur)
    return exec_sql, exec_sql_params


@test("local: REPLAY of the 240926 drop -- span replace keeps rows outside the span, per-warehouse snapshot, verify ALL PASS to the fils")
def _():
    conn = _local_conn()
    if conn is None:
        return
    if not DROP_240926.exists() or not list(DROP_240926.glob("Sales_day_book*.xls*")):
        print(f"    SKIP: the 240926 drop is not at {DROP_240926}")
        conn.close()
        return
    import shutil
    from decimal import Decimal
    from app import db_read
    from scripts import ingest, load_supabase as ls, verify_numbers as vn
    exec_sql, exec_sql_params = _local_exec(conn)
    saved = (ingest.OUT_DIR, ls.CLEAN, ls._client, vn.get_client, db_read.exec_sql, db_read.exec_sql_params, list(sys.argv))
    tmp = Path(tempfile.mkdtemp(prefix="r2replay_"))
    try:
        _q(conn, (ROOT / "scripts" / "economics_v2_migration.sql").read_text(encoding="utf-8"))
        # plant rows the replace must and must not touch
        _q(conn, "INSERT INTO orders (invoice_no, order_date, customer_name, gross_bhd, salesman) VALUES "
                 "('SI : R2-OUTSIDE', '2025-01-05', 'CUST-test', 9.5, 'Karrar'), ('SI : R2-INSIDE', '2026-09-10', 'CUST-test', 4.25, 'Karrar')")
        _q(conn, "INSERT INTO order_lines (invoice_no, line_no, line_date, customer_account, item_name, quantity, gross_bhd, warehouse_name) VALUES "
                 "('SI : R2-OUTSIDE', 1, '2025-01-05', 'CUST-test', 'X01 test', 1, 9.5, 'Karrar'), ('SI : R2-INSIDE', 1, '2026-09-10', 'CUST-test', 'X01 test', 1, 4.25, 'Karrar')")
        as_of = _q(conn, "SELECT MAX(as_of_date)::text AS d FROM stock_balance")[0]["d"]
        _q(conn, "INSERT INTO stock_balance (item_name, warehouse_name, net_qty, selling_rate_bhd, total_value_bhd, as_of_date) VALUES "
                 "('R2 VAN ITEM', 'R2 Test Van', 3, 1, 3, %s::date), ('R2 STALE ITEM', 'Accessories Warehouse', 3, 1, 3, %s::date)", (as_of, as_of))
        before = _q(conn, "SELECT (SELECT COUNT(*) FROM orders) AS o, (SELECT COUNT(*) FROM order_lines) AS l, (SELECT COUNT(*) FROM stock_movements) AS sm")[0]
        # ingest the real drop into a scratch clean folder, then load through the psycopg client
        ingest.OUT_DIR = tmp / "clean"
        ls.CLEAN = tmp / "clean"
        client = _PgClient(conn)
        ls._client = lambda: client
        vn.get_client = lambda: client
        db_read.exec_sql, db_read.exec_sql_params = exec_sql, exec_sql_params
        sys.argv = ["ingest", str(DROP_240926)]
        assert ingest.main() == 0
        rc = ls.main(["--staged", str(DROP_240926)])
        assert rc == 0
        # the §2 preview numbers, exactly
        o = _q(conn, "SELECT COUNT(*) AS n, ROUND(SUM(gross_bhd)::numeric, 3) AS s FROM orders WHERE invoice_no <> 'SI : R2-OUTSIDE'")[0]
        l = _q(conn, "SELECT COUNT(*) AS n, ROUND(SUM(taxable_bhd)::numeric, 3) AS s FROM order_lines WHERE invoice_no <> 'SI : R2-OUTSIDE'")[0]
        sb = _q(conn, "SELECT COUNT(*) AS n, SUM(net_qty) AS u, ROUND(SUM(total_value_bhd)::numeric, 3) AS v FROM stock_balance "
                      "WHERE as_of_date = '2026-09-24' AND warehouse_name = 'Accessories Warehouse'")[0]
        ar = _q(conn, "SELECT COUNT(*) AS n, ROUND(SUM(balance_bhd)::numeric, 3) AS s FROM ar_ageing WHERE as_of_date = '2026-09-24'")[0]
        art = _q(conn, "SELECT focus_total_bhd, focus_over90_bhd, rows_total_bhd FROM ar_ageing_totals WHERE as_of_date = '2026-09-24'")[0]
        pp = _q(conn, "SELECT COUNT(*) AS n FROM product_profitability WHERE report_date = '2026-09-24'")[0]["n"]
        sm = _q(conn, "SELECT COUNT(*) AS n, SUM(received_qty) AS r, SUM(issued_qty) AS i FROM stock_movements")[0]
        print(f"    orders {o['n']} / gross {o['s']} | lines {l['n']} / taxable {l['s']} | stock 24-Sep {sb['n']} items {sb['u']} u {sb['v']} | "
              f"AR {ar['n']} accts {ar['s']} (Focus {art['focus_total_bhd']}, >90 {art['focus_over90_bhd']}) | profitability {pp} | ledger {sm['n']} rows")
        assert (o["n"], o["s"]) == (2743, Decimal("74238.130")) and (l["n"], l["s"]) == (11956, Decimal("67490.060")), (o, l)
        assert (sb["n"], sb["u"], sb["v"]) == (132, Decimal("45516.0"), Decimal("67489.310")), sb
        assert (ar["n"], ar["s"]) == (84, Decimal("9078.860")) and art["focus_total_bhd"] == Decimal("8633.840") and art["rows_total_bhd"] == Decimal("9078.860"), (ar, art)
        assert pp == 161 and (sm["n"], sm["r"], sm["i"]) == (36267, Decimal("219477.0"), Decimal("156575.0")), (pp, sm)
        # the span rule: outside kept, inside (not in the file) gone and audited; other warehouses kept, stale rows of the covered one gone
        assert _q(conn, "SELECT COUNT(*) AS n FROM orders WHERE invoice_no = 'SI : R2-OUTSIDE'")[0]["n"] == 1
        assert _q(conn, "SELECT COUNT(*) AS n FROM order_lines WHERE invoice_no = 'SI : R2-OUTSIDE'")[0]["n"] == 1
        assert _q(conn, "SELECT COUNT(*) AS n FROM orders WHERE invoice_no = 'SI : R2-INSIDE'")[0]["n"] == 0
        assert _q(conn, "SELECT COUNT(*) AS n FROM order_lines WHERE invoice_no = 'SI : R2-INSIDE'")[0]["n"] == 0
        audits = _q(conn, "SELECT event, detail FROM audit_log WHERE event LIKE '%span_replace%' ORDER BY id")
        assert [a["event"] for a in audits] == ["order_lines.span_replace", "orders.span_replace"], audits
        assert audits[1]["detail"]["disappearing"] == ["SI : R2-INSIDE"] and audits[1]["detail"]["span"] == ["2025-09-21", "2026-09-24"]
        assert _q(conn, "SELECT COUNT(*) AS n FROM stock_balance WHERE item_name = 'R2 VAN ITEM'")[0]["n"] == 1
        assert _q(conn, "SELECT COUNT(*) AS n FROM stock_balance WHERE item_name = 'R2 STALE ITEM'")[0]["n"] == 0
        assert _q(conn, "SELECT COUNT(*) AS n FROM audit_log WHERE event = 'stock_balance.partial_snapshot'")[0]["n"] == 0
        assert before["sm"] == sm["n"]
        # verify against the same files: everything passes, to the fils
        ok, rows = vn.run_checks(DROP_240926)
        for r in rows:
            print(f"    {'PASS' if r['passed'] else 'FAIL'}  {r['metric'][:60]:60} {r['report']:>14,.3f} {r['db']:>14,.3f}")
            if r.get("note"):
                print(f"          {r['note'][:150]}")
        names = [r["metric"] for r in rows]
        assert any(m.startswith("Sales per day: max |file - DB|") for m in names) and any(m.startswith("Invoices per day") for m in names)
        assert any(m.startswith("Stock snapshot 2026-09-24: previous warehouses missing") for m in names)
        assert any(m.startswith("Receivables Focus total 2026-09-24 stored") for m in names)
        assert ok, [r["metric"] for r in rows if not r["passed"]]
        day = next(r for r in rows if r["metric"].startswith("Sales per day: max"))
        assert day["db"] <= 0.005 and next(r for r in rows if r["metric"].startswith("Sales lines"))["report"] == 11956.0
        # the alias autofill on the replay copy: nothing invents, the share is measured
        from scripts.alias_autofill import autofill
        res = autofill(exec_sql=exec_sql, client=client, emit=lambda *a, **k: True)
        print(f"    alias autofill: picked {res['picked']} inserted {res['inserted']} share {res['share']['share'] if res.get('share') else None}")
        assert res.get("error") is None and res["share"] is not None
    finally:
        conn.rollback()
        conn.close()
        ingest.OUT_DIR, ls.CLEAN, ls._client, vn.get_client, db_read.exec_sql, db_read.exec_sql_params, sys.argv = saved
        shutil.rmtree(tmp, ignore_errors=True)


def main() -> int:
    passed = failed = 0
    for name, fn in TESTS:
        try:
            fn()
            print(f"  PASS  {name}")
            passed += 1
        except Exception:  # noqa: BLE001
            print(f"  FAIL  {name}")
            traceback.print_exc(limit=3)
            failed += 1
    print(f"\n{passed} passed, {failed} failed")
    return 0 if failed == 0 else 2


if __name__ == "__main__":
    raise SystemExit(main())
