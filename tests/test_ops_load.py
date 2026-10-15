"""The load check (scripts/ops/load_check.py): what /yq-load runs before anything is loaded.

    python -m tests.test_ops_load

Same lightweight runner as the other suites (no pytest). No database, no network, no .env:
SUPABASE_URL, SUPABASE_KEY and DATABASE_URL are pinned to ci.invalid BEFORE any import, the session
and the preview are stand-ins, every folder is a temporary one. The workbooks are built here with
openpyxl in the layout the real Focus parsers accept; all names and numbers are invented (Main
Store, Van A, Shop 1, Cable one) and round.

Covered:
  * the verdict rules, pure: profitability (gap above, below and AT the limit, a cost far above the
    ledger's, blank costs of sold items only, no ledger, two holds in a row read from the plan of
    the load before), the name bridge, the price book (no change, a general change, a scoped row
    below the general price, a dropped key, a scheduled price, an end date, a status, and their
    edges: the export day itself, a price that ended yesterday, a scoped price equal to the general
    one), the stock ledger (missing item / warehouse / voucher type, changed, removed and
    earlier-dated vouchers), the preview (an invoice that changed, one that lost a line, one whose
    header is gone; each blocking code; the loader's guards), the catalog;
  * older than what is loaded: a whole drop (stops at once: nothing compared, nothing recorded),
    one older report in a current folder, a report whose rows end before production's;
  * what the check read is still there when it ends: a report saved again while it ran, a report
    changed between its verdict and its copy, a database another load moved;
  * preflight: a code worktree, the folder against the live site, the two database addresses, the
    lock (running; interrupted: the cut-off load's own GO line is handed over as the last line, with
    the place of its go.cmd on the line before; the dead lock of a rehearsal and a lock that is no
    lock at all do not stop the check; a lock that cannot be read does), and a load that failed
    after it had begun to write (the check runs as usual and its first line says so);
  * one check counts for a folder: a second check replaces the first (one live staging folder, the
    first go.cmd gone, its plan names the second), a run that started, finished or failed and a
    check of another folder are never touched, a dev run replaces nothing, the left-over of a
    stopped check is put away once it cannot still be running, and what happens when Windows will
    not move a folder or a plan cannot be rewritten; the run name is the minute of the check, and a
    second check inside one minute waits for the next;
  * the loader's marker files: in staging only on the owner's own option, the same story in
    plan.json, a report named like a marker refused, one written in a single place of the source;
  * the folder: two files of one report, a workbook open in Excel, a file still being written, a
    report with no "As on date", a file that is not a Focus export, sub-folders, a partial drop,
    reports only in sub-folders, the Ledger export never staged and never a reason to stop;
  * the stock-balance trim on a three-warehouse workbook (merged cells too), the export Focus
    filtered to one store, and every part of its proof failing by itself: the Sub Total's units,
    its value, an altered row, another date, a sheet out of line, a copy that is not identical,
    production's snapshot not readable, no store, no closing row, an item whose name reads as a
    total row;
  * the whole check with stand-ins: staging names and markers for every option, the table's words,
    the token and the GO line, go.cmd and plan.json (dated by what loads, never by the Ledger
    export or a held report), --dev (no token, no go.cmd, no GO line of any kind, nothing
    recorded), a changed header row, a loader guard predicted to fire and the option that lets it
    through, the token following a staged file, the deadline, and the drop folder left exactly as
    it was;
  * the action register (through scripts/ops/alerts.py, with no empty first run): an invoice or
    voucher that changed is listed from the first time it is found (never a baseline item), stays
    open across checks and is never cleared by one; a state (a held report) clears by itself but
    not when the drop that was loaded is merely checked again; a scoped price below the general one
    is a baseline on the rule's first run, and the line says so; a stopped check records nothing;
    sales that could not be looked up are never recorded as none, and the rows for the action list
    are looked up first when there are more rows than lookups;
  * what the load changes per salesman, not only per day;
  * status: the five states, the writer's sentences since the last call in full (text, why, next)
    and none of its machine fields, --wait, the run that holds the lock followed before the latest
    one, a check that was replaced, a rehearsal that was cut off.
A comparison with the real 7-Oct files prints SKIP when they are not on this machine.
"""
from __future__ import annotations

import argparse
import io
import json
import os
import re
import shutil
import sys
import tempfile
import threading
import traceback
from contextlib import redirect_stdout
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

os.environ["SUPABASE_URL"] = "https://ci.invalid"
os.environ["SUPABASE_KEY"] = "ci-key"
os.environ["DATABASE_URL"] = "postgresql://ci:ci@ci.invalid:5432/ci"
os.environ.pop("YQ_MAIN_CHECKOUT", None)

import warnings  # noqa: E402

warnings.filterwarnings("ignore")

from openpyxl import Workbook, load_workbook  # noqa: E402

from scripts.ops import db  # noqa: E402  (first, as every command does)
from scripts.ops import __main__ as cli  # noqa: E402
from scripts.ops import alerts, config, fmt, loadplan, paths, status  # noqa: E402
from scripts.ops import load_check as lc  # noqa: E402

TESTS: list[tuple[str, object]] = []

MAIN = "Main Store"
FIXTURE = {
    "api.health_url": "https://ci.invalid/health", "sales.kpi_steps": [],
    "load.profit_gap_max": 0.1, "load.blank_cost_share_max": 0.2, "load.backup_keep": 3,
    "brief.rep_share_min": 0.5, "brief.silent_business_days": 5, "brief.bulk_invoice_bhd": 100,
    "stock.lead_time_days": 10, "stock.display_items": [], "items.bridge_exceptions": {},
    "items.not_accessories": [], "warehouses.main": MAIN, "warehouses.aliases": {},
    "warehouses.not_counted": [],
}
API = "https://abcdefghijklmnopqrst.supabase.co"
DSN = "postgresql://postgres.abcdefghijklmnopqrst:secret@pooler.ci.invalid:5432/postgres"
NINE_DIGITS = re.compile(r"(?<![0-9])[0-9]{9}(?![0-9])|[0-9]{15,}")


def test(name):
    def deco(fn):
        TESTS.append((name, fn))
        return fn
    return deco


def raises(kind, fn) -> Exception:
    try:
        fn()
    except kind as exc:
        return exc
    raise AssertionError(f"expected {kind.__name__}")


START = datetime(2026, 1, 7, 9, 30, tzinfo=status.BAHRAIN)     # when the first check of a data root is made
WALL = {"at": START}                      # the stand-ins' wall clock: a pause moves it on, a new data root resets it


class data_root:
    """A temporary data root with the test config in force. Its clock starts at 09:30: the stand-ins
    of every check made in it share that clock, so a second check comes later than the first."""

    def __enter__(self) -> Path:
        self._tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.root = Path(self._tmp.name).resolve()
        paths.set_data_root(self.root)
        config.use(FIXTURE)
        WALL["at"] = START
        return self.root

    def __exit__(self, *exc):
        paths.set_data_root(None)
        config.use(None)
        self._tmp.cleanup()
        return False


# ── synthetic Focus workbooks (the layouts scripts/ingest.py reads) ────────────

AS_ON = "07/01/2026"                      # the title spells dd/mm/yyyy: 7 January 2026
DAY = "2026-01-07"
COMPANY = "YQ Bahrain"                    # the words scripts.ingest.sniff_focus looks for in a title block

STOCK = [("Van A", [("Cable one C1", 10, 1.5), ("Charger two C2", 4, 2.5)]),
         (MAIN, [("Cable one C1", 100, 1.5), ("Charger two C2", 40, 2.5), ("Holder three C3", 20, 0.75)]),
         ("Van B", [("Holder three C3", 6, 0.75)])]
MAIN_UNITS, MAIN_VALUE = Decimal(160), Decimal(265)

# item -> (date, voucher, received, issued, issued value, warehouse, to warehouse, voucher name)
LEDGER = {
    "Cable one C1": [
        ("2026-01-02", "MRN:T-26-01-1", 500, None, None, MAIN, None, "Material Receipt Note"),
        ("2026-01-03", "SIO:T-26-01-1", None, 50, 75, MAIN, "Van A", "Stock Issue Voucher"),
        ("2026-01-03", "STRV:T-26-01-1", 50, None, None, "Van A", None, "Stock Receive Voucher"),
        ("2026-01-04", "SI:T-26-01-1", None, 20, 100, "Van A", None, "Sales Invoice"),
        ("2026-01-05", "SI:T-26-01-2", None, 10, 50, "Van A", None, "Sales Invoice")],
    "Charger two C2": [
        ("2026-01-02", "MRN:T-26-01-1", 200, None, None, MAIN, None, "Material Receipt Note"),
        ("2026-01-04", "SI:T-26-01-1", None, 10, 100, "Van A", None, "Sales Invoice")],
    "Holder three C3": [
        ("2026-01-02", "MRN:T-26-01-1", 100, None, None, MAIN, None, "Material Receipt Note"),
        ("2026-01-06", "SI:T-26-01-3", None, 5, 50, "Van B", None, "Sales Invoice")],
}
PROFIT = [("Cable one", 300, 150), ("Charger two", 200, 100), ("Holder three", 100, 50)]     # name, net, cost
# item, code, customer, customer code, warehouse, start (m/d/yyyy), end, rate
PRICES = [("Cable one", "C1", None, None, None, "1/1/2020", "12/31/2030", 1.5),
          ("Charger two", "C2", None, None, None, "1/1/2020", "12/31/2030", 2.5),
          ("Holder three", "C3", None, None, None, "1/1/2020", "12/31/2030", 0.75),
          ("Cable one", "C1", None, None, "Shop 1", "1/1/2020", "12/31/2030", 1.75)]
LEDGER_HEADER = ["Date", "Voucher", "Received Quantity", "Rate", "Issued Quantity", "Rate", "Balance Quantity", "Value",
                 "Value", "Value", "Avg Rate", "Warehouse Name", "To Warehouse Name", "Narration", "Voucher name"]
PRICE_HEADER = ["Item Name", "Item Code", "Customer Name", "Customer Code", "Warehouse Name", "Warehouse Code",
                "Currency", "Start date", "End date", "MinQty", "MaxQty", "Unit Name", "Rate",
                *[f"Val {n}" for n in range(1, 14)], "Status", "Narration"]
NAMES = {
    "order_lines": "Sales_day_book100001_Test_Co.xlsx", "orders": "Summary_sales_register100002_Test_Co.xlsx",
    "stock_movements": "Stock_ledger100003_Test_Co.xlsx", "stock_balance": "Stock_balance_by_warehouse100004_Test_Co.xlsx",
    "receivables": "Customer_summary_ageing_by_due_date100005_Test_Co.xlsx",
    "product_profitability": "Product_Profitability_Report100006_Test_Co.xlsx",
    "selling_prices:MA_base": "MASellingPriceBook (7).xlsx", "ledger_entries": "Ledger100007_Test_Co.xlsx",
}


def sheet(title: str | None, as_on: str | None, header: list[str]):
    """A workbook with Focus's title block (company, report name, as-on date) and the header on row 6."""
    wb = Workbook()
    ws = wb.active
    if title is not None:
        ws["C1"], ws["C2"] = COMPANY, title
        if as_on:
            ws["C4"] = f"[As on date {as_on}]"
    for j, label in enumerate(header, start=1):
        ws.cell(row=6, column=j, value=label)
    return wb, ws


def put(ws, row: int, values: list) -> None:
    for j, v in enumerate(values, start=1):
        if v is not None:
            ws.cell(row=row, column=j, value=v)


def stock_balance_file(path: Path, blocks=STOCK, as_on: str | None = AS_ON, sub_total_off: int = 0,
                       closing: str | None = "Sub Total", grand: bool = True, sub_value_off: int = 0) -> Path:
    """Focus's layout: per warehouse a name row, its item rows and a closing total row ("Sub Total";
    "Total" when the export was filtered to one warehouse), then "Grand Total". sub_total_off and
    sub_value_off put the main store's closing row out by that many units / that much value."""
    wb, ws = sheet("Stock balance by warehouse", as_on, ["Particulars", "Net Quantity", "Selling Rate", "Total Value"])
    row, all_units, all_value = 7, 0, 0
    for name, items in blocks:
        put(ws, row, [f"{name} {name}"])                 # Focus doubles the warehouse name
        row += 1
        units = value = 0
        for item, qty, rate in items:
            if qty is None:
                put(ws, row, [item])                     # an item row with no figures at all
            else:
                put(ws, row, [item, qty, rate, qty * rate])
                units, value = units + qty, value + qty * rate
            row += 1
        if closing:
            put(ws, row, [closing, units + (sub_total_off if name == MAIN else 0), None,
                          value + (sub_value_off if name == MAIN else 0)])
            row += 1
        all_units, all_value = all_units + units, all_value + value
    if grand:
        put(ws, row, ["Grand Total", all_units, None, all_value])
    wb.save(path)
    return path


def ledger_file(path: Path, ledger=LEDGER, as_on: str | None = AS_ON) -> Path:
    wb, ws = sheet("Stock ledger", as_on, LEDGER_HEADER)
    row = 7
    for item, lines in ledger.items():
        put(ws, row, [item])
        row += 1
        for day, voucher, got, out, out_value, wh, to_wh, kind in lines:
            put(ws, row, [datetime.fromisoformat(day), voucher, got, 1 if got else None, out, 1 if out else None, 100,
                          got, out_value, 100, 1, wh, to_wh, None, kind])
            row += 1
        put(ws, row, ["Sub Total"])
        row += 1
    wb.save(path)
    return path


def profit_file(path: Path, rows=PROFIT, as_on: str | None = AS_ON) -> Path:
    wb, ws = sheet("Product profitability report", as_on,
                   ["Particulars", "Gross", "Discount %", "Net amount", "COGS", "Gross Profit", "GP Margin %",
                    "Misc Charges", "Net Profit", "NP Margin %"])
    for i, (name, net, cost) in enumerate(rows):
        put(ws, 7 + i, [name, net, None, net, cost, None if cost is None else net - cost])
    put(ws, 7 + len(rows), ["Grand Total", sum(r[1] for r in rows)])
    wb.save(path)
    return path


def price_file(path: Path, rows=PRICES) -> Path:
    wb = Workbook()
    ws = wb.active
    put(ws, 1, PRICE_HEADER)
    for i, (item, code, cust, cust_code, wh, start, end, rate) in enumerate(rows):
        put(ws, 2 + i, [item, code, cust, cust_code, wh, None, "Dinar", start, end, 0, 0, "Nos", rate,
                        *([0] * 13), "Authorized"])
    wb.save(path)
    return path


def plain_report(path: Path, title: str, header: list[str], as_on: str | None = AS_ON, company: bool = True) -> Path:
    wb, ws = sheet(title if company else None, as_on, header)
    if not company:
        ws["C2"] = title
    put(ws, 7, ["x", 1])
    wb.save(path)
    return path


def receivables_file(path: Path, rows, as_on: str | None = AS_ON) -> Path:
    """The customer ageing in the layout scripts.ingest.parse_receivables reads: per account its balance,
    the eight buckets (the 0-30 one carries it all) and the total."""
    from scripts import ingest as ing
    wb, ws = sheet("Customer summary ageing by due date", as_on,
                   list(ing.EXPECTED_HEADERS["receivables"]) + ["Account Code", "Group Name"])
    for i, (account, balance) in enumerate(rows):
        put(ws, 7 + i, [account, balance, balance, 0, balance, balance, balance, 0, 0, 0, 0, 0, 0, 0, balance,
                        None, "Debtors"])
    put(ws, 7 + len(rows), ["Grand Total", sum(b for _a, b in rows)])
    wb.save(path)
    return path


def make_drop(root: Path, name: str = "070126", leave_out: tuple = (), as_on: str = AS_ON, **kw) -> Path:
    """A full drop: the six daily reports, a price book and a Ledger export, each titled as on `as_on`."""
    drop = root / "Uplaods" / name
    drop.mkdir(parents=True)
    builders = {
        "order_lines": lambda p: plain_report(p, "Sales day book", ["Date", "Voucher"], as_on=as_on),
        "orders": lambda p: plain_report(p, "Summary sales register", ["Date", "Invoice"], as_on=as_on),
        "stock_movements": lambda p: ledger_file(p, kw.get("ledger", LEDGER), as_on=as_on),
        "stock_balance": lambda p: stock_balance_file(p, kw.get("stock", STOCK), as_on=as_on,
                                                      sub_total_off=kw.get("sub_total_off", 0)),
        "receivables": lambda p: plain_report(p, "Customer summary ageing", ["Account", "Balance Amount"], as_on=as_on),
        "product_profitability": lambda p: profit_file(p, kw.get("profit", PROFIT), as_on=as_on),
        "selling_prices:MA_base": lambda p: price_file(p, kw.get("prices", PRICES)),
        "ledger_entries": lambda p: plain_report(p, "Ledger", ["Date", "Voucher"], as_on=as_on),
    }
    for kind, build in builders.items():
        if kind not in leave_out:
            build(drop / NAMES[kind])
    return drop


# ── stand-ins for the database and the preview ────────────────────────────────

def live_prices(rows=PRICES, source: str = "MASellingPriceBook _6_.xlsx") -> list[dict]:
    def day(mdy: str | None) -> str | None:
        if not mdy:
            return None
        m, d, y = mdy.split("/")
        return f"{int(y):04d}-{int(m):02d}-{int(d):02d}"
    return [{"sku_code": code, "item_name": item, "price_book": "MA_base", "customer_code": cust_code,
             "customer_name": cust, "warehouse_name": wh, "start_date": day(start), "end_date": day(end),
             "rate_bhd": Decimal(str(rate)), "status": "Authorized", "source_file": source}
            for item, code, cust, cust_code, wh, start, end, rate in rows]


def db_ledger(ledger=LEDGER, upto: str = "2026-01-05") -> list[dict]:
    """Production's stock ledger as the grouped read returns it (rows up to the last loaded day)."""
    out = []
    for item, lines in ledger.items():
        for day, voucher, got, issued, _value, wh, to_wh, kind in lines:
            if day <= upto:
                out.append({"voucher": voucher, "item_name": item, "d": day, "vt": kind, "w": wh or "", "tw": to_wh or "",
                            "rq": Decimal(got or 0), "iq": Decimal(issued or 0), "n": 1})
    return out


class FakeBackend:
    """The preview's backend as the stand-in preview uses it (FakeDeps.via_backend): every call is
    recorded; a diff answers with the actions its first parameter carries, anything else with no rows."""

    mode, read_only, stage_source = "fake", True, "param"

    def __init__(self):
        self.asked: list[tuple[str, list]] = []

    def query(self, sql, params=None):
        self.asked.append((" ".join(str(sql).split())[:40], list(params or [])))
        if "jsonb_array_elements" in sql:
            return [json.loads(params[0])]
        return []


class FakeSession:
    """Answers the check's reads from canned rows; every statement passes the real read-only guard."""

    def __init__(self, backend=None, **over):
        self.backend = backend if backend is not None else object()
        self.closed = False
        self.seen: list[str] = []
        self.data = {
            "anchors": {"today": DAY, "sales_to": "2026-01-05", "stock_as_of": "2026-01-05", "ar_as_of": "2026-01-05",
                        "profit_as_of": "2026-01-05", "moves_to": "2026-01-05"},
            "profit_before": [{"item_name": n, "net_amount_bhd": Decimal(net) - 50, "cogs_bhd": Decimal(cost) - 25,
                               "report_date": "2026-01-05"} for n, net, cost in PROFIT],
            "prices": live_prices(), "warehouses_before": [{"w": MAIN}], "orders": [], "lines": [],
            "ledger": db_ledger(),
            "catalog": [{"item_code": c, "is_active": True, "hidden": False, "updated_by": lc.AUTO_SYNC}
                        for c in ("C1", "C2", "C3")],
            "book": [{"sku_code": c} for c in ("C1", "C2", "C3")],
            "candidates": [{"sku_code": c} for c in ("C1", "C2", "C3")],
            "count": {"n": 10, "max_id": 99},
            "snapshot_days": [],                      # the days production holds a stock or receivables snapshot for
            "partitions": {},                         # target -> production's rows of the file's snapshot day
        }
        self.data.update(over)

    def q(self, sql: str, params=None) -> list[dict]:
        db.check_sql(sql)
        s = " ".join(sql.split())
        self.seen.append(s)
        d = self.data
        if s.startswith("select current_date"):
            return [dict(d["anchors"])]
        if "from product_profitability where report_date" in s:
            return d["profit_before"]
        if "from selling_prices where price_book = %s and voided_at is null" in s:
            return d["prices"]
        if "from stock_balance where as_of_date = (select max(as_of_date) from stock_balance where as_of_date <" in s:
            return d["warehouses_before"]
        if s.startswith("select invoice_no, salesman from orders"):
            return d["orders"]
        if s.startswith("select distinct invoice_no, warehouse_name from order_lines"):
            return d["lines"]
        if "from stock_movements where move_date between" in s:
            return d["ledger"]
        if "from catalog_items" in s:
            return d["catalog"]
        if "from v_price_list_by_book" in s:
            return d["book"]
        if s.startswith("select distinct sku_code from selling_prices"):
            return d["candidates"]
        if s.startswith("select count(*) as n, coalesce(max(id), 0) as max_id from"):
            return [dict(d["count"])]
        if "from ingest_runs" in s:
            return [{"max_id": 5}]
        if "to_regclass" in s:
            return [{"ok": True}]
        if "from ingest_batches" in s:
            return []
        if s.startswith("select distinct 'stock_balance' as k, as_of_date::text as d from stock_balance where as_of_date > "
                        "%s::date union select distinct 'ar_ageing' as k"):
            # a list: both snapshot kinds on those days (one load brings both); a dict: the days of each kind
            days = d["snapshot_days"]
            kinds = days if isinstance(days, dict) else {"stock_balance": days, "ar_ageing": days}
            return [{"k": k, "d": day} for k, ds in kinds.items() for day in sorted(set(ds)) if day > str(params[0])]
        for target in ("ar_ageing", "stock_balance"):
            if s.startswith(f"select * from {target} where "):
                return [dict(r) for r in d["partitions"].get(target, [])]
        raise AssertionError("a statement the stand-in does not know: " + s[:90])

    def close(self) -> None:
        self.closed = True


SPAN = {"kind": "span", "from": "2026-01-01", "to": "2026-01-06", "within": {}}
TARGET_OF = {"order_lines": "order_lines", "orders": "orders", "stock_movements": "stock_movements",
             "stock_balance": "stock_balance", "receivables": "ar_ageing", "product_profitability": "product_profitability"}


class FakeDeps(lc.Deps):
    """The check's contacts, all local: the project folder on the live commit, one database."""

    deadline_s = 0                                    # run in this thread

    def __init__(self, session: FakeSession | None = None, **over):
        self.lines: list[str] = []
        self.the_session = session or FakeSession()
        self.worktree = False
        self.site = {"state": "same"}
        self.live = "abc123abc123"
        self.pair = (API, DSN)
        self.snapshot = [{"item_name": i, "net_qty": q} for i, q, _r in STOCK[1][1]]
        self.scope_sold = [{"invoices": 2, "units": 30}]
        self.actions: dict = {}                       # target -> actions the preview reports
        self.samples: dict = {}                       # target -> {updated_sample, removed_sample}
        self.exceptions: list[dict] = []
        self.hook_rows: dict = {}                     # target -> rows handed to the row hook
        self.sales: dict = {"per_day_changed": [], "per_salesman_changed": []}
        self.stock_view: dict | None = None           # the preview's stock-per-warehouse comparison
        self.loaded_rows: dict = {}                   # target -> rows production already holds in the file's scope
        self.previewed: list[list[str]] = []
        self.via_backend = False                      # the stand-in preview asks its backend, as the real one does
        self.on_pause = None
        self.ticks = 0.0
        self.pauses: list[float] = []                 # every pause the check asked for, in seconds
        for k, v in over.items():
            setattr(self, k, v)

    def say(self, line):
        self.lines.append(line)

    @property
    def steps(self) -> list[str]:
        """The progress lines of the nine steps (a line that says the check is waiting is not one)."""
        return [ln for ln in self.lines if ln.startswith("[")]

    def now(self):
        return WALL["at"]

    def clock(self):
        return self.ticks

    def pause(self, seconds):
        self.ticks += seconds
        self.pauses.append(seconds)
        WALL["at"] += timedelta(seconds=seconds)
        if self.on_pause:
            self.on_pause()

    def linked_worktree(self):
        return self.worktree

    def live_commit(self):
        return self.live

    def site_state(self, live):
        return dict(self.site)

    def untracked(self):
        return ["scripts/ops/load_check.py", "scripts/my_notes.py"]

    def addresses(self):
        return self.pair

    def read_params(self, sql, params):
        db.check_sql(sql)
        if "from stock_balance" in sql:
            assert params == [MAIN], params
            return self.snapshot
        if "from v_sales" in sql:
            return self.scope_sold
        raise AssertionError("a read the stand-in does not know")

    def session(self):
        return self.the_session

    def preview(self, folder, backend, out_dir, today, row_hook):
        from scripts import ingest as ing
        from scripts.ops import querytime
        # the session's own backend, handed over as it is or inside the check's timing proxy
        assert querytime.unwrap(backend) is self.the_session.backend and today == DAY
        self.previewed.append(sorted(p.name for p in Path(folder).iterdir()))
        targets: dict = {}
        for p in sorted(Path(folder).iterdir()):
            kind = ing.classify(p.name) if p.suffix.lower() == ".xlsx" else None
            target = TARGET_OF.get(kind) or ("selling_prices" if (kind or "").startswith("selling_prices") else None)
            if target is None:
                continue
            # the preview's own shapes: a price book voids and revives rows, every other table deletes them
            actions = ({"inserted": 2, "updated": 0, "unchanged": 5, "voided": 0, "revived": 0} if target == "selling_prices"
                       else {"inserted": 2, "updated": 0, "unchanged": 5, "deleted": 0})
            actions.update(self.actions.get(target, {}))
            if self.via_backend:                      # as app.ingest_batch._target_diff asks: columns, then the diff
                backend.query("select column_name, data_type from information_schema.columns where table_name = $1",
                              [target])
                actions = backend.query("with s as (select value as row from jsonb_array_elements($1::jsonb)) select 1",
                                        [json.dumps(actions), target])[0]
            targets[target] = {"file_rows": 7, "actions": actions, "scope": dict(SPAN),
                               "db_rows": self.loaded_rows.get(target, 0),
                               "db_sums": {"balance_bhd": "150.500"} if self.loaded_rows.get(target) else {},
                               "sums": {"balance_bhd": "150.500"}, **self.samples.get(target, {})}
        for target, rows in self.hook_rows.items():
            for r in rows:
                row_hook(target, r)
        if self.via_backend:                          # and one helper, as _sales_comparisons asks
            backend.query("select line_date::text as d from order_lines where line_date between $1::date and $2::date",
                          [SPAN["from"], SPAN["to"]])
        Path(out_dir).mkdir(parents=True, exist_ok=True)
        return {"summary": {"targets": targets, "sales": self.sales, "stock": self.stock_view,
                            "receivables": {"focus_grand_total": "140.250", "rows_sum": "150.500"}},
                "exceptions": self.exceptions, "preview_dir": str(out_dir), "batch_id": None}


def check(*argv: str, deps: FakeDeps | None = None) -> tuple[dict, FakeDeps]:
    parser = argparse.ArgumentParser()
    parser.add_argument("--json", action="store_true")
    lc.add_arguments(parser)
    deps = deps or FakeDeps()
    return lc.run_check(parser.parse_args(["check", *argv]), deps), deps


def shown(kind: str) -> str:
    """A report's file name as the table prints it (long names have their middle cut out)."""
    return lc.short_name(NAMES[kind])


def table_of(text: str) -> dict[str, str]:
    """{file name as the table prints it: its decision} from the printed block."""
    lines = text.splitlines()
    head = next(i for i, ln in enumerate(lines) if ln.startswith("File "))
    at = lines[head].index("Decision")
    out = {}
    for ln in lines[head + 1:]:
        if not ln.strip():
            break
        rest = ln[at:]
        word = next(w for w in sorted(lc.WORDS.values(), key=len, reverse=True) if rest.startswith(w))
        out[ln[:at].strip()] = word
    return out


def snapshot_of(folder: Path) -> dict:
    return {str(p.relative_to(folder)): (p.stat().st_size, p.stat().st_mtime_ns, loadplan.sha256_file(p))
            for p in sorted(folder.rglob("*")) if p.is_file()}


# ── the verdict rules ─────────────────────────────────────────────────────────

def ledger_rows(ledger=LEDGER) -> list[dict]:
    return [{"item_name": item, "move_date": d, "voucher": v, "received_qty": got, "issued_qty": out,
             "issued_value_bhd": val, "warehouse_name": wh, "to_warehouse_name": to, "voucher_type": kind,
             "row_hash": "h", "balance_qty": 1}
            for item, lines in ledger.items() for d, v, got, out, val, wh, to, kind in lines]


def profit_rows(rows=PROFIT) -> list[dict]:
    return [{"item_name": n, "net_amount_bhd": net, "cogs_bhd": cost, "report_date": DAY} for n, net, cost in rows]


@test("profitability: the gap to the stock ledger decides; a report that agrees loads")
def _():
    v = lc.profit_verdict(profit_rows(), ledger_rows(), DAY, Decimal("0.1"), Decimal("0.2"))
    assert v["ledger_cost"] == 300 and v["report_cost"] == 300 and v["gap"] == 0 and not v["hold"], v
    short = profit_rows([("Cable one", 300, 100), ("Charger two", 200, 100), ("Holder three", 100, 50)])
    v = lc.profit_verdict(short, ledger_rows(), DAY, Decimal("0.1"), Decimal("0.2"))
    assert v["gap"] == Decimal(50) / Decimal(300) and v["hold"] and "16.7%" in v["why"][0] and "10.0%" in v["why"][0], v
    assert not lc.profit_verdict(short, ledger_rows(), DAY, Decimal("0.2"), Decimal("0.2"))["hold"]     # under the limit
    # the limit itself is still accepted: the report is held back only when the gap EXCEEDS it
    edge = profit_rows([("Cable one", 300, 135), ("Charger two", 200, 90), ("Holder three", 100, 45)])   # cost 270 of 300
    v = lc.profit_verdict(edge, ledger_rows(), DAY, Decimal("0.1"), Decimal("0.2"))
    assert v["gap"] == Decimal("0.1") and not v["hold"] and v["why"] == [], v
    assert lc.profit_verdict(edge, ledger_rows(), DAY, Decimal("0.099"), Decimal("0.2"))["hold"]
    # only Sales Invoice rows up to the report's date count, and only for the items the report lists
    early = lc.ledger_cost(profit_rows(), ledger_rows(), "2026-01-04")
    assert early["ledger_cost"] == 200, early
    two = lc.ledger_cost(profit_rows(PROFIT[:2]), ledger_rows(), DAY)
    assert two["ledger_cost"] == 250 and two["items"] == 2 and two["matched"] == 2, two


@test("profitability: blank costs above their share hold it back; no ledger in the drop holds it back")
def _():
    blank = profit_rows([("Cable one", 300, 150), ("Charger two", 200, None), ("Holder three", 100, 0)])
    v = lc.profit_verdict(blank, ledger_rows(), DAY, Decimal("0.9"), Decimal("0.2"))
    assert v["blank"]["share"] == Decimal("0.5") and v["blank"]["blank_items"] == 2 and v["hold"], v
    assert "50.0%" in v["why"][-1]
    assert not lc.profit_verdict(blank, ledger_rows(), DAY, Decimal("0.9"), Decimal("0.5"))["hold"]
    v = lc.profit_verdict(profit_rows(), None, DAY, Decimal("0.1"), Decimal("0.2"))
    assert v["hold"] and v["gap"] is None and "no stock ledger" in v["why"][0]
    v = lc.profit_verdict(profit_rows(), None, DAY, Decimal("0.1"), Decimal("0.2"), no_ledger="the stock ledger stops")
    assert v["why"] == ["the stock ledger stops, so its cost of sales cannot be checked"]       # there, but not usable
    v = lc.profit_verdict(profit_rows(), [], DAY, Decimal("0.1"), Decimal("0.2"))
    assert v["hold"] and v["gap"] is None                       # a ledger with no sales cost cannot vouch
    # only SOLD items count: an item with no sale and no cost, or a return with no cost, is not a blank cost
    unsold = profit_rows([("Cable one", 300, 150), ("Charger two", 0, None), ("Holder three", -50, 0)])
    b = lc.blank_cost_share(unsold)
    assert b["blank_items"] == 0 and b["blank_net"] == 0 and b["share"] == 0 and b["net"] == 250, b
    one = lc.blank_cost_share(profit_rows([("Cable one", 300, 150), ("Charger two", 100, None)]))
    assert one["blank_items"] == 1 and one["share"] == Decimal("0.25")
    # the share itself is still accepted: held back only above it
    assert not lc.profit_verdict(profit_rows([("Cable one", 300, 150), ("Charger two", 100, None)]), ledger_rows(), DAY,
                                 Decimal("0.9"), Decimal("0.25"))["hold"]


@test("the name bridge: ledger name = report name + code, longest prefix; an exception answers first")
def _():
    names = ["Cable one", "Cable one plus", "Holder"]
    b = lc.bridge_names(["Cable one C1", "Cable one plus C9", "Holder", "Charger Z9", None], names)
    assert b == {"Cable one C1": "Cable one", "Cable one plus C9": "Cable one plus", "Holder": "Holder"}, b
    b = lc.bridge_names(["Odd name X1"], names, {"Odd name X1": "Holder", "Other": "Nowhere"})
    assert b == {"Odd name X1": "Holder"}
    assert lc.bridge_names(["Cable onex C1"], names) == {}       # a prefix must end at a space


@test("profitability: sales that moved with no cost movement are counted; two holds in a row are told apart")
def _():
    old = profit_rows([("Cable one", 250, 150), ("Charger two", 200, 100)])
    new = profit_rows([("Cable one", 300, 150), ("Charger two", 200, 100), ("Holder three", 100, 50)])
    assert lc.cost_movement(new, old) == (2, 1)                  # Cable moved without cost; Holder is new with cost
    # no plan of the load before (it was made by hand): the data dates answer
    anchors = {"stock_as_of": "2026-01-05", "ar_as_of": "2026-01-05", "profit_as_of": "2026-01-01"}
    assert lc.second_hold(True, "2026-01-07", anchors) is True
    assert lc.second_hold(False, "2026-01-07", anchors) is False
    assert lc.second_hold(True, "2026-01-07", dict(anchors, profit_as_of="2026-01-05")) is False   # the last load had it
    # the same drop checked again after its load: nothing new can be said (None = leave the register alone)
    assert lc.second_hold(True, "2026-01-05", anchors) is None
    assert lc.second_hold(True, "2026-01-04", anchors) is None
    assert lc.second_hold(False, "2026-01-05", anchors) is False           # ... unless the report now loads
    # with the plan of the load before, that plan answers and the data dates do not
    after_partial = {"stock_as_of": "2026-01-06", "ar_as_of": "2026-01-06", "profit_as_of": "2026-01-05"}
    assert lc.second_hold(True, "2026-01-07", after_partial) is True                       # the dates alone would say so
    assert lc.second_hold(True, "2026-01-07", after_partial, (False, "")) is False         # its drop brought no report
    assert lc.second_hold(True, "2026-01-07", dict(anchors, profit_as_of="2026-01-05"), (True, "2026-01-05")) is True
    assert lc.second_hold(True, "2026-01-07", anchors, (True, "2026-01-07")) is None       # that load judged this report
    assert lc.second_hold(False, "2026-01-07", anchors, (True, "2026-01-05")) is False
    # what the plan of the load before says
    assert lc.profit_before(None) == (None, "") and lc.profit_before({}) == (None, "")
    held = {"verdicts": [{"kind": "orders", "verdict": "load"}, {"kind": "product_profitability", "verdict": "hold_back"}],
            "as_on": {"orders": "2026-01-05"}, "as_on_all": {"orders": "2026-01-05", "product_profitability": "2026-01-05"}}
    assert lc.profit_before(held) == (True, "2026-01-05")
    assert lc.profit_before({"verdicts": [{"kind": "product_profitability", "verdict": "load"}],
                             "as_on": {"product_profitability": "2026-01-05"}}) == (False, "2026-01-05")
    assert lc.profit_before({"verdicts": [{"kind": "product_profitability", "verdict": "not_loaded"}], "as_on": {}}) == (False, "")
    with data_root():
        assert lc.previous_load_plan() is None                                  # no load on record
        status.write_state({"last_load": {"run": "050126"}})                    # a load made by hand: not a run name
        assert lc.previous_load_plan() is None
        status.write_state({"last_load": {"run": "050126_0900"}})               # a run with no plan on file
        assert lc.previous_load_plan() is None
        loadplan.write_plan(dict(held, run="050126_0900", drop="x", token="t"))
        assert lc.profit_before(lc.previous_load_plan()) == (True, "2026-01-05")


def price_rows(rows=PRICES, **fixed) -> list[dict]:
    out = live_prices(rows, source="MASellingPriceBook (7).xlsx")
    for r in out:
        r.update(fixed)
    return out


@test("price book: no key added, changed or dropped; other sources are never dropped")
def _():
    live = live_prices() + [dict(live_prices([("Extra", "E1", None, None, None, "1/1/2020", None, 9)])[0],
                                 source_file="Workbook import.xlsx")]
    d = lc.pricebook_diff(price_rows(), live, "MA_base", DAY)
    assert d["items"] == [] and d["file_keys"] == 4 and d["live_keys"] == 5, d
    # the same key twice in the file: the last row is the one that loads (the loader's rule)
    twice = price_rows() + price_rows([("Cable one", "C1", None, None, None, "1/1/2020", "12/31/2030", 1.5)])
    assert lc.pricebook_diff(twice, live, "MA_base", DAY)["items"] == []


@test("price book: a general change, a scoped row below the general price, a dropped key, a scheduled price, an end date")
def _():
    rows = [("Cable one", "C1", None, None, None, "1/1/2020", "12/31/2030", 1.6),          # general price changed
            ("Charger two", "C2", None, None, None, "1/1/2020", "6/30/2030", 2.5),         # end date changed
            ("Cable one", "C1", None, None, "Shop 1", "1/1/2020", "12/31/2030", 1.75),      # as before
            ("Charger two", "C2", None, None, "Van A", "1/5/2026", "12/31/2030", 2.0),      # scoped, below general
            ("Charger two", "C2", "Shop 1", "S1", None, "1/5/2026", None, 3.0),             # customer scope, above
            ("Cable one", "C1", None, None, None, "2/1/2026", None, 1.8)]                   # starts after the export
    d = lc.pricebook_diff(price_rows(rows), live_prices(), "MA_base", DAY)
    by = {(e["sku"], e["scope"], e["start"]): e for e in d["items"]}
    assert len(by) == 6, sorted(by)
    general = by[("C1", "", "2020-01-01")]
    assert general["what"] == "changed" and general["tags"] == ["general", "price"] and general["old_rate"] == Decimal("1.5")
    assert "was 1.500" in lc.price_line(general) and "general price 1.600" in lc.price_line(general)
    end = by[("C2", "", "2020-01-01")]
    assert end["tags"] == ["general", "end_date"] and "end date" in lc.price_line(end)
    van = by[("C2", "Van A", "2026-01-05")]
    assert van["what"] == "added" and van["scoped"] and van["below_general"] and van["general"] == Decimal("2.5")
    assert van["warehouse"] == "Van A" and "BELOW the general price 2.500" in lc.price_line(van)
    shop = by[("C2", "Shop 1", "2026-01-05")]
    assert shop["scoped"] and not shop["below_general"] and shop["customer"] == "Shop 1"
    later = by[("C1", "", "2026-02-01")]
    assert "scheduled" in later["tags"] and "starts after the export day" in lc.price_line(later)
    assert "live now is 1.500" in lc.price_line(later)
    gone = by[("C3", "", "2020-01-01")]
    assert gone["what"] == "dropped" and "no longer in the book" in lc.price_line(gone)
    for e in d["items"]:
        fmt.careful(lc.price_line(e))
    # the register takes the scoped row below the general price, keyed without the van's name
    unknown = lc.price_scope_items(d["items"])                 # its sales were not looked up (or the lookup failed)
    assert len(unknown) == 1 and unknown[0]["bhd"] is None and "was not looked up" in unknown[0]["detail"]
    assert "0 units" not in unknown[0]["detail"], "sales that are not known are never said to be none"
    van["units"], van["invoices"] = Decimal(30), 2
    found = lc.price_scope_items(d["items"])
    assert [f["key"] for f in found] == [f"price_scope|C2|{alerts.ref_hash('Van A')}|2026-01-05"], found
    assert found[0]["level"] == "CHECK" and found[0]["bhd"] == Decimal("15.0") and "Van A" not in found[0]["key"]
    assert "30 units already sold there at this price" in found[0]["detail"]
    van["units"], van["invoices"] = Decimal(0), 0              # looked up, and nothing was sold: that IS said
    none = lc.price_scope_items(d["items"])[0]
    assert "0 units already sold there at this price" in none["detail"] and none["bhd"] == 0


@test("price book: a row whose status changed is a changed row (the live list holds a row only while it is Authorized)")
def _():
    file_rows = price_rows()
    file_rows[1]["status"] = "Draft"                           # Charger two's general price: same price, same dates
    d = lc.pricebook_diff(file_rows, live_prices(), "MA_base", DAY)
    (e,) = d["items"]
    assert (e["sku"], e["what"], e["tags"]) == ("C2", "changed", ["general", "status"]), e
    assert lc.price_line(e) == ("C2: general price 2.500 from Wed 01 Jan (status Authorized -> Draft: it would no longer "
                                "be in force)")
    live = live_prices()
    live[1]["status"] = None                                   # the other way round: a live row that is not in force
    (e,) = lc.pricebook_diff(price_rows(), live, "MA_base", DAY)["items"]
    assert e["tags"] == ["general", "status"] and "status blank -> Authorized: it would come into force" in lc.price_line(e)
    both = price_rows()
    both[1].update(status="Pending", rate_bhd=Decimal("2.6"))  # a price and a status at once
    (e,) = lc.pricebook_diff(both, live, "MA_base", DAY)["items"]
    assert e["tags"] == ["general", "price", "status"] and "was 2.500; status blank -> Pending" in lc.price_line(e)
    assert lc.status_word({"status": " Authorized "}) == "Authorized" and lc.status_word(None) == "" and lc.status_word({}) == ""
    with data_root() as root:                 # in the whole check: not "the same as the live book", so not LEAVE OUT
        drop = make_drop(root)
        book = drop / NAMES["selling_prices:MA_base"]
        wb = load_workbook(book)
        wb.active.cell(row=3, column=27, value="Draft")        # the Status column of Charger two's general price
        wb.save(book)
        res, _ = check("070126")
        assert res["ok"] and res["token"], res["text"]
        assert table_of(res["text"])[shown("selling_prices:MA_base")] == "NEEDS YOUR DECISION"
        assert "(1 changed)" in res["text"] and "status Authorized -> Draft: it would no longer be in force" in res["text"]


@test("price book: the edges: a price from the export day is not scheduled, a scoped price equal to the general one is not below")
def _():
    rows = PRICES + [("Charger two", "C2", None, None, None, "1/7/2026", None, 2.6),        # starts ON the export day
                     ("Charger two", "C2", None, None, None, "1/8/2026", None, 2.7),        # starts the day after
                     ("Holder three", "C3", None, None, "Van A", "1/5/2026", None, 0.75)]   # scoped, same as general
    d = lc.pricebook_diff(price_rows(rows), live_prices(), "MA_base", DAY)
    by = {(e["sku"], e["scope"], e["start"]): e for e in d["items"]}
    assert len(by) == 3, sorted(by)
    today = by[("C2", "", DAY)]
    assert today["tags"] == ["general"] and "scheduled" not in lc.price_line(today)
    tomorrow = by[("C2", "", "2026-01-08")]
    assert tomorrow["tags"] == ["general", "scheduled"] and "a scheduled price" in lc.price_line(tomorrow)
    same = by[("C3", "Van A", "2026-01-05")]
    assert same["general"] == Decimal("0.75") and same["below_general"] is False
    assert "against the general price 0.750" in lc.price_line(same) and "BELOW" not in lc.price_line(same)
    assert lc.price_scope_items(d["items"]) == []
    # a general price that ended yesterday is not in force; one that ends today still is
    ended = price_rows([("Cable one", "C1", None, None, None, "1/1/2020", "1/6/2026", 1.5)])
    assert lc.general_price(ended, "C1", "2026-01-07") is None
    assert lc.general_price(ended, "C1", "2026-01-06") == Decimal("1.5")
    scoped = price_rows([("Cable one", "C1", None, None, None, "1/1/2020", "1/6/2026", 1.5),
                         ("Cable one", "C1", None, None, "Van A", "1/5/2026", None, 1.0)])
    late = lc.pricebook_diff(scoped, [], "MA_base", DAY)["items"]
    van = next(e for e in late if e["scoped"])
    assert van["general"] is None and van["below_general"] is False and "no general price in force" in lc.price_line(van)


@test("price book: the general price in force follows the live list's rule; sales at a scoped price are looked up")
def _():
    rows = price_rows([("Cable one", "C1", None, None, None, "1/1/2020", None, 1.5),
                       ("Cable one", "C1", None, None, None, "1/1/2026", None, 1.25),
                       ("Cable one", "C1", None, None, None, "1/1/2027", None, 9),
                       ("Cable one", "C1", None, None, "Shop 1", "1/1/2026", None, 0.5)])
    assert lc.general_price(rows, "C1", "2025-12-31") == Decimal("1.5")
    assert lc.general_price(rows, "C1", "2026-06-01") == Decimal("1.25")
    assert lc.general_price(rows, "C9", "2026-06-01") is None
    rows[1]["status"] = "Draft"
    assert lc.general_price(rows, "C1", "2026-06-01") == Decimal("1.5")
    seen: list = []

    class D(FakeDeps):
        def read_params(self, sql, params):
            db.check_sql(sql)
            seen.append((sql, params))
            return [{"invoices": 3, "units": "45"}]
    item = {"sku": "C2", "start": "2026-01-05", "rate": Decimal("2.0"), "warehouse": "Van A", "customer": "Shop 1"}
    lc.scope_sales(item, D())
    sql, params = seen[0]
    assert params == ["C2", "2026-01-05", "2.0", "Van A", "Shop 1"] and "salesman_raw = $4" in sql and "= $5" in sql
    assert "v_sales" in sql and item["invoices"] == 3 and item["units"] == Decimal(45)


def db_rows(rows: list[dict]) -> list[dict]:
    return [{"voucher": r["voucher"], "item_name": r["item_name"], "d": r["move_date"], "vt": r["voucher_type"],
             "w": r["warehouse_name"] or "", "tw": r["to_warehouse_name"] or "", "rq": Decimal(r["received_qty"] or 0),
             "iq": Decimal(r["issued_qty"] or 0), "n": 1} for r in rows]


@test("stock ledger: an item, a warehouse or a voucher type production holds and the file lacks is found")
def _():
    file_rows = ledger_rows()
    same = lc.ledger_compare(file_rows, db_rows(file_rows), "2026-01-06")
    assert not (same["missing_items"] or same["missing_warehouses"] or same["missing_types"])
    assert same["changed"] == same["removed"] == same["backdated"] == [] and same["new"] == 0 and same["pairs"] == 9
    held = db_rows(file_rows) + [
        {"voucher": "SI:T-26-01-9", "item_name": "Lamp four C4", "d": "2026-01-04", "vt": "Sales Invoice", "w": "Van C",
         "tw": "", "rq": Decimal(0), "iq": Decimal(2), "n": 1},
        {"voucher": "SRV:T-26-01-1", "item_name": "Cable one C1", "d": "2026-01-04", "vt": "Sales Return", "w": "Van A",
         "tw": "", "rq": Decimal(1), "iq": Decimal(0), "n": 1}]
    cmp = lc.ledger_compare(file_rows, held, "2026-01-06")
    assert cmp["missing_items"] == ["Lamp four C4"] and cmp["missing_warehouses"] == ["Van C"]
    assert cmp["missing_types"] == ["Sales Return"]


@test("stock ledger: changed, removed and earlier-dated vouchers are listed; balances, rates and hashes are ignored")
def _():
    base = ledger_rows()
    recost = [dict(r, row_hash="other", balance_qty=7, issued_value_bhd=1) for r in base]       # a re-export
    assert lc.ledger_compare(recost, db_rows(base), "2026-01-06")["changed"] == []
    file_rows = [dict(r) for r in base if r["voucher"] != "SI:T-26-01-2"]                        # one voucher gone
    for r in file_rows:
        if r["voucher"] == "SIO:T-26-01-1":
            r["issued_qty"] = 48                                                                   # one edited
    file_rows.append({"item_name": "Cable one C1", "move_date": "2026-01-03", "voucher": "SIO:T-26-01-7",
                      "received_qty": None, "issued_qty": 3, "warehouse_name": MAIN, "to_warehouse_name": "Van B",
                      "voucher_type": "Stock Issue Voucher"})                                     # new, earlier date
    file_rows.append({"item_name": "Cable one C1", "move_date": "2026-01-06", "voucher": "SI:T-26-01-8",
                      "received_qty": None, "issued_qty": 1, "warehouse_name": "Van A", "to_warehouse_name": None,
                      "voucher_type": "Sales Invoice"})                                           # new, on the last day
    cmp = lc.ledger_compare(file_rows, db_rows(base), "2026-01-06")
    assert cmp["changed"] == [("SIO:T-26-01-1", "Cable one C1")] and cmp["removed"] == [("SI:T-26-01-2", "Cable one C1")]
    assert cmp["backdated"] == [("SIO:T-26-01-7", "Cable one C1")] and cmp["new"] == 1
    found = lc.voucher_items(cmp)
    assert sorted(f["key"] for f in found) == ["voucher_changed|SI:T-26-01-2|Cable one C1",
                                               "voucher_changed|SIO:T-26-01-1|Cable one C1",
                                               "voucher_changed|SIO:T-26-01-7|Cable one C1"]
    assert all(f["level"] == "CHECK" and f["title"].endswith("to check") for f in found)
    changed = next(f for f in found if "SIO:T-26-01-1|" in f["key"])
    assert "issued 50" in changed["detail"] and "issued 48" in changed["detail"] and changed["evidence"] == ["SIO:T-26-01-1"]
    # a name that holds a refused word never stops the register: the plain line is used
    assert lc.safe_words("Anti-theft lock L1: differs", "plain") == "plain" and lc.safe_words("fine", "plain") == "fine"


def blocking(code, target=None, count=None, items=None, message="m"):
    return {"code": code, "severity": "blocking", "target": target, "count": count, "items": items or [], "message": message}


@test("the preview: changed and removed invoices are read from its samples, each blocking code has its meaning")
def _():
    summary = {"targets": {
        "orders": {"actions": {"inserted": 3, "updated": 1, "unchanged": 9, "deleted": 2},
                   "updated_sample": [{"invoice_no": "INV-2"}], "removed_sample": [{"invoice_no": "INV-5"}]},
        "order_lines": {"actions": {"inserted": 8, "updated": 60, "unchanged": 20, "deleted": 1},
                        "updated_sample": [{"invoice_no": "INV-2", "line_no": 1}, {"invoice_no": "INV-3", "line_no": 2}],
                        "removed_sample": [{"invoice_no": "INV-2", "line_no": 4}]},
        "selling_prices": {"actions": {"inserted": 1, "updated": 0, "unchanged": 4, "voided": 2, "revived": 0}}}}
    pv = lc.read_preview(summary, [
        blocking("rows_removed", "stock_movements", 400), blocking("rows_removed", "orders", 2),
        blocking("days_shrink", "order_lines", 1), blocking("rows_removed", "ar_ageing", 3),
        blocking("warehouse_set_shrinks", "stock_balance", items=["Van A"]), blocking("snapshot_void_guard", "selling_prices"),
        {"code": "unaliased_items", "severity": "warning", "message": "w"}])
    assert pv["actions"]["orders"] == {"inserted": 3, "updated": 1, "unchanged": 9, "removed": 2}
    assert pv["actions"]["selling_prices"]["removed"] == 2
    # INV-2 lost a line (its header is still in the register): changed, not removed. INV-5's header is gone.
    assert pv["invoices"] == {"INV-2": "changed", "INV-5": "removed", "INV-3": "changed"}, pv["invoices"]
    lines_only = {"targets": {"order_lines": {"actions": {"deleted": 3}, "removed_sample": [{"invoice_no": "INV-7"}]}}}
    assert lc.read_preview(lines_only, [])["invoices"] == {"INV-7": "changed"}        # only the register says 'removed'
    gone_first = {"targets": {"orders": {"actions": {}, "removed_sample": [{"invoice_no": "INV-8"}]},
                              "order_lines": {"actions": {}, "updated_sample": [{"invoice_no": "INV-8"}]}}}
    assert lc.read_preview(gone_first, [])["invoices"] == {"INV-8": "removed"}        # and 'removed' is never undone
    assert pv["more_invoices"] == 58 and pv["ledger_bulk"] == 400 and pv["stops"] == []
    assert pv["warehouse_shrinks"] == ["Van A"] and pv["void_guard"] == "m" and len(pv["infos"]) == 2
    for code in lc.HARD_CODES:                      # the preview itself is not to be trusted: always a stop
        stops = lc.read_preview(summary, [blocking(code)])["stops"]
        assert len(stops) == 1 and "cannot vouch for this drop" in stops[0] and code in stops[0], code
    stops = lc.read_preview(summary, [blocking("something_new")])["stops"]      # a code this check has never seen
    assert stops == ["the preview raised something_new: m"], stops


@test("the loader's sales guards: many invoices gone, or a salesman absent from the file, fire; --replace-ok lets through")
def _():
    held = [{"invoice_no": f"INV-{i}", "salesman": "Van A" if i < 30 else "Van B"} for i in range(40)]
    full = {"orders": [dict(r) for r in held]}
    assert lc.sales_guards(full, {"orders": held}, False) == {}
    assert lc.sales_guards({"orders": held[1:]}, {"orders": held}, False) == {}          # one invoice gone: no guard
    fired = lc.sales_guards({"orders": held[6:30] + held[30:]}, {"orders": held}, False)
    assert list(fired) == ["orders"] and "6 of the 40 invoices" in fired["orders"][0], fired
    fired = lc.sales_guards({"orders": held[:30]}, {"orders": held}, False)              # a whole salesman left out
    assert any("absent from the file (Van B)" in why for why in fired["orders"]), fired
    assert lc.sales_guards({"orders": held[:30]}, {"orders": held}, True) == {}
    lines = [{"invoice_no": r["invoice_no"], "warehouse_name": r["salesman"]} for r in held]
    fired = lc.sales_guards({"order_lines": lines[:30]}, {"order_lines": lines}, False)
    assert list(fired) == ["order_lines"]
    assert lc.sales_guards({"orders": held}, {}, False) == {}                            # nothing read: nothing claimed


@test("catalog: what the sync after a load would switch, with hidden and hand-set items marked")
def _():
    rows = [{"item_code": "C1", "is_active": True, "hidden": False, "updated_by": lc.AUTO_SYNC},
            {"item_code": "C2", "is_active": True, "hidden": True, "updated_by": "owner@example.invalid"},
            {"item_code": "C3", "is_active": False, "hidden": False, "updated_by": "owner@example.invalid"},
            {"item_code": "C4", "is_active": False, "hidden": False, "updated_by": None}]
    e = lc.catalog_effect(rows, {"C1", "C3", "C5"}, {"C1", "C3", "C5"})
    assert e == {"off": ["C2 (hidden, last set by hand)"], "on": ["C3 (last set by hand)"], "new": ["C5"]}, e
    assert "switch OFF 1 item" in lc.catalog_line(e) and "add 1 new item: C5" in lc.catalog_line(e)
    assert lc.catalog_effect(rows, set(), set()) == {"off": [], "on": [], "new": []}     # an empty list switches nothing off
    assert "switch nothing" in lc.catalog_line({"off": [], "on": [], "new": []})


# ── preflight ─────────────────────────────────────────────────────────────────

@test("the two database addresses must name one project, and neither is ever printed")
def _():
    assert lc.project_ref(API) == "abcdefghijklmnopqrst" and lc.project_ref(DSN) == "abcdefghijklmnopqrst"
    assert lc.project_ref("postgresql://postgres:p@ss:word@db.abcdefghijklmnopqrst.supabase.co:5432/postgres") == "abcdefghijklmnopqrst"
    assert lc.project_ref("https://ci.invalid") is None and lc.project_ref("postgresql://ci:ci@ci.invalid:5432/ci") is None
    assert lc.project_ref("") is None and lc.project_ref(None) is None
    assert lc.same_project(API, DSN) and not lc.same_project(API, DSN.replace("abcdefghij", "zzzzzzzzzz"))
    assert not lc.same_project("https://ci.invalid", "postgresql://ci:ci@ci.invalid:5432/ci")
    with data_root() as root:
        make_drop(root)
        res, _ = check("070126", deps=FakeDeps(pair=(API, DSN.replace("abcdefghij", "zzzzzzzzzz"))))
        assert res["exit"] == 2 and "do not name the same project" in res["text"]
        assert "secret" not in res["text"] and "abcdefghij" not in res["text"] and "zzzzzzzzzz" not in res["text"]
        assert not (root / "business_data").exists()                 # stopped before anything was staged


@test("preflight stops in a code worktree, behind or ahead of the live site, and on a lock; --dev skips the first two")
def _():
    with data_root() as root:
        make_drop(root)
        res, _ = check("070126", deps=FakeDeps(worktree=True))
        assert res["exit"] == 2 and "code worktree" in res["text"] and res["token"] is None
        res, _ = check("070126", deps=FakeDeps(site={"state": "behind"}))
        assert "behind the live site: pull it" in res["text"] and res["flags"] == ["stop"]
        assert "Left by this check" not in res["text"], "a stop before staging leaves nothing behind"
        res, _ = check("070126", deps=FakeDeps(site={"state": "changed_here", "paths": ["scripts/ingest.py", "app/x.py"]}))
        assert "scripts/ingest.py, app/x.py" in res["text"] and "tell Claude" in res["text"]
        res, _ = check("070126", deps=FakeDeps(site={"state": "release_not_live", "paths": ["app/metrics.py"]}))
        assert "merged but not live" in res["text"] and "app/metrics.py" in res["text"]

        class Down(FakeDeps):
            def live_commit(self):
                raise RuntimeError("no answer from https://ci.invalid/health")
        res, _ = check("070126", deps=Down())
        assert "live site did not answer" in res["text"] and res["exit"] == 2
        assert not (root / "business_data").exists()
        now = datetime.now(timezone.utc)
        lock = {"run": "060126_0800", "pid": os.getpid(), "step": "refresh", "started": now.isoformat(), "heartbeat": now.isoformat()}
        paths.write_text(paths.lock_file(), json.dumps(lock))
        res, _ = check("070126")
        assert "A load is running now (run 060126_0800, at refresh)" in res["text"] and res["exit"] == 2
        paths.write_text(paths.lock_file(), json.dumps(dict(lock, pid=2 ** 22 + 12345)))
        res, _ = check("070126", "--dev")
        assert res["text"].splitlines()[1].startswith("LOAD INTERRUPTED at refresh (run 060126_0800)")
        assert res["flags"] == ["stop", "interrupted"] and "GO: " not in res["text"]      # no plan: no line to hand over
        assert "business_data/ops/loads/060126_0800/go.cmd" in res["text"] and "tell Claude" in res["text"]
        loadplan.write_plan({"run": "060126_0800", "drop": "x", "token": "abcabc", "verdicts": []})
        res, _ = check("070126")                                    # only the cut-off load's own go line resumes it
        lines = res["text"].splitlines()
        assert lines[-1] == "GO: python -m scripts.yq_load_go --run 060126_0800 --token abcabc", lines
        assert "the box will say RESUME" in lines[1] and len(lines) == 3 and "No token" not in res["text"]
        assert res["exit"] == 2 and res["token"] is None and res["run"] is None
        assert loadplan.latest_run() == "060126_0800", "a stopped check opens no run of its own"
        res, _ = check("070126", "--dev")                           # a dev run prints no GO line, not even that one
        assert "GO:" not in res["text"] and res["flags"] == ["stop", "interrupted"] and res["exit"] == 2, res["text"]
        assert res["text"].splitlines()[1].startswith("LOAD INTERRUPTED at refresh (run 060126_0800)")
        assert res["text"].splitlines()[-1].startswith("DEV RUN: no GO line is printed here.") and "/yq-load" in res["text"]
        # with that run's go.cmd on disk, the line before its GO line says where it lies (for a double-click)
        go_cmd = loadplan.write_go_cmd("060126_0800", "abcabc", root)
        lines = check("070126")[0]["text"].splitlines()
        assert lines[-1] == "GO: python -m scripts.yq_load_go --run 060126_0800 --token abcabc" and len(lines) == 4, lines
        assert lines[-2] == f"If Claude may not start it: double-click {go_cmd}"
        assert "double-click" not in check("070126", "--dev")[0]["text"]
        assert not (root / "business_data" / "staging").exists(), "nothing new is checked while a load waits to be resumed"
        paths.lock_file().unlink()
        res, deps = check("070126", "--dev", deps=FakeDeps(worktree=True, site={"state": "behind"}))
        assert res["ok"] and res["exit"] == 0 and "dev run" in deps.lines[0]


def dead_lock(run: str, step: str = "refresh", beat: datetime | None = None) -> None:
    """A lock whose process is gone (no such process id on this PC)."""
    beat = beat or datetime.now(timezone.utc)
    paths.write_text(paths.lock_file(), json.dumps({"run": run, "pid": 2 ** 22 + 12345, "step": step,
                                                    "started": beat.isoformat(), "heartbeat": beat.isoformat()}))


@test("the dead lock of a rehearsal does not stop a check: one line says the load clears it, and no rehearsal go line is handed out")
def _():
    rehearsal = status.REHEARSAL_PREFIX + "0915"
    with data_root() as root:
        make_drop(root)
        loadplan.write_plan({"run": rehearsal, "drop": "x", "token": "rehearsalword", "rehearsal": True, "verdicts": []})
        loadplan.write_go_cmd(rehearsal, "rehearsalword", root)
        loadplan.journal_append(rehearsal, "started", rehearsal=True)
        loadplan.journal_append(rehearsal, "step_start", step="backup")
        dead_lock(rehearsal, "backup (rehearsal)")
        assert status.lock_state() == {"state": "none", "rehearsal_left": rehearsal, "step": "backup (rehearsal)"}
        res, deps = check("070126")
        text = res["text"]
        assert res["ok"] and res["token"] and res["run"] == "070126_0930", text
        assert text.count(lc.REHEARSAL_LEFT) == 1 and "- A cut-off rehearsal left its lock; the load clears it." in text
        assert "INTERRUPTED" not in text and rehearsal not in text and "rehearsalword" not in text
        assert text.splitlines()[-1] == "GO: " + loadplan.go_line(res["run"], res["token"])
        assert deps.steps[0] == "[1/9] Preflight: same programs as the live site (abc123ab), one database project, " \
                                "no load running."
        assert paths.lock_file().is_file(), "the check only reads the lock: the writer clears it at the go"
        # the same when the check stops for another reason: never the rehearsal's own go line
        res, _ = check("070126", deps=FakeDeps(site={"state": "behind"}))
        assert res["exit"] == 2 and "GO:" not in res["text"] and "rehearsalword" not in res["text"]
        assert "INTERRUPTED" not in res["text"] and rehearsal not in res["text"]
        # `load status` on that rehearsal says at once that it was cut off (its journal is seconds old)
        assert lc.run_state(rehearsal) == ("INTERRUPTED", "backup (rehearsal)") and lc.locked_run() is None
        seen = status_of("--run", rehearsal)
        assert seen["text"].splitlines()[-1] == "INTERRUPTED at backup (rehearsal)" and seen["exit"] == 2
        # a status that still called such a lock interrupted would get no go line either
        assert lc.preflight(False, FakeDeps(lock_state=lambda: {"state": "interrupted", "run": rehearsal, "step": "backup"})
                            )["notes"][-1] == lc.REHEARSAL_LEFT
        # a rehearsal that is running now is waited for, and called what it is
        now = datetime.now(timezone.utc)
        write_lock(rehearsal, os.getpid(), now, "backup (rehearsal)")
        res, _ = check("070126")
        assert res["exit"] == 2 and res["token"] is None, res["text"]
        assert f"A rehearsal is running now (run {rehearsal}, at backup (rehearsal)). Wait for it: python -m scripts.ops " \
               "load status --wait" in res["text"]
        assert lc.locked_run() == rehearsal and lc.run_state(rehearsal) == ("RUNNING", "backup (rehearsal)")


@test("a lock that is no lock at all, and a lock that cannot be read: the first is a left-over, the second stops the check")
def _():
    with data_root() as root:
        make_drop(root)
        paths.write_text(paths.lock_file(), "")                     # its launcher ended before it had written it
        assert status.lock_state()["state"] == "interrupted" and status.lock_state()["run"] == "?"
        res, deps = check("070126")
        assert res["ok"] and res["token"], res["text"]
        assert res["text"].count(lc.NAMELESS_LOCK) == 1 and "INTERRUPTED" not in res["text"]
        assert deps.pauses[:2] == [lc.LOCK_AGAIN_S, lc.SIZE_PAUSE_S], "it is read once more before it is taken for a left-over"
        assert paths.lock_file().is_file()

        class BeingWritten(FakeDeps):
            """The lock is being written at the very moment the check looks: the second look reads it."""
            looks = 0

            def lock_state(self):
                self.looks += 1
                return ({"state": "interrupted", "run": "?", "step": "?"} if self.looks == 2 else
                        {"state": "none"} if self.looks == 1 else {"state": "running", "run": "070126_0900", "step": "start"})
        res, _ = check("070126", deps=BeingWritten())
        assert res["exit"] == 2 and "A load is running now (run 070126_0900, at start)" in res["text"], res["text"]
        paths.lock_file().unlink()

        def busy():
            raise PermissionError(13, "The process cannot access the file")
        real = status.read_lock
        status.read_lock = busy
        try:
            res, _ = check("070126")
            assert res["exit"] == 2 and res["token"] is None and res["flags"] == ["stop"], res["text"]
            assert res["text"].splitlines()[1] == ("The load lock could not be read (PermissionError), so it cannot be told "
                                                   "whether a load is running. Nothing was checked: try again in a moment.")
            # `load status` does not end on it either: it follows the latest run and reads the lock at the next call
            assert lc.lock_now() == {"state": "unreadable", "error": "PermissionError"} and lc.locked_run() is None
            run = loadplan.latest_run()
            assert status_of()["run"] == run and status_of()["state"] == "NOT STARTED"
            # a load that started long ago and whose lock is busy at this moment is not called interrupted: the
            # lock is what tells, and the next call reads it
            long_ago = (datetime.now(timezone.utc) - timedelta(seconds=status.HEARTBEAT_STALE_S + 600)).isoformat()
            paths.write_text(loadplan.journal_file(run), json.dumps({"at": long_ago, "event": "started"}) + "\n")
            assert lc.run_state(run) == ("RUNNING", None) and status_of()["exit"] == 0
            loadplan.journal_append(run, "finished")
            assert lc.run_state(run)[0] == "FINISHED"
        finally:
            status.read_lock = real

        class Late(FakeDeps):
            """The lock cannot be read when the earlier checks are about to be put away."""
            looks = 0

            def lock_state(self):
                self.looks += 1
                if self.looks > 2:
                    raise PermissionError(13, "The process cannot access the file")
                return {"state": "none"}
        res, _ = check("070126", deps=Late())
        assert res["exit"] == 2 and res["token"] is None, res["text"]
        assert "The load lock could not be read (PermissionError), so it cannot be told whether a load has just started." \
               in res["text"] and "Nothing was staged: try again in a moment." in res["text"]


@test("a load that failed after it had begun to write: the check runs as usual and its first line says so")
def _():
    failed = "050126_0900"

    def failed_load(written: bool = True) -> None:
        loadplan.journal_append(failed, "started")
        loadplan.journal_append(failed, "step_start", step="refresh")
        loadplan.journal_append(failed, "failed", step="refresh", result="failed", written=written,
                                text="LOAD FAILED at refresh.")
    said = f"The last load (run {failed}) failed at refresh after it had begun to write: "
    with data_root() as root:
        make_drop(root)
        failed_load()
        assert status.failed_load() == {"run": failed, "step": "refresh"}
        res, deps = check("070126")
        lines = res["text"].splitlines()
        assert res["ok"] and res["token"] and "last_load_failed" in res["flags"], res["text"]
        assert lines[0] == said + "this check starts a new load that replaces it." and lines[1] == ""
        assert lines[2].startswith("LOAD CHECK  070126  (run 070126_0930") and lines[-1].startswith("GO: ")
        assert res["text"].count("The last load") == 1 and len(deps.steps) == 9, "one line, and the check runs as usual"
        assert table_of(res["text"])[shown("orders")] == "LOAD" and "last_load_failed" in loadplan.read_plan(res["run"])["flags"]
    with data_root() as root:                 # a dev run and a check that stops start no load: the line does not say they do
        make_drop(root)
        failed_load()
        res, _ = check("070126", "--dev")
        lines = res["text"].splitlines()
        assert lines[0] == said + "the database may hold part of it until a new load replaces it." and lines[1] == ""
        assert "[DEV RUN]" in lines[2] and res["token"] is None and "last_load_failed" in res["flags"]
    with data_root() as root:
        make_drop(root)
        failed_load()
        res, _ = check("070126", deps=FakeDeps(site={"state": "behind"}))          # stops at the preflight
        lines = res["text"].splitlines()
        assert lines[0] == said + "the database may hold part of it until a new load replaces it."
        assert lines[1:3] == ["", "LOAD CHECK: STOP"] and res["flags"] == ["stop", "last_load_failed"], res["text"]

        class Stuck(FakeDeps):                # ... and when the check never comes back
            deadline_s = 0.3

            def live_commit(self):
                release.wait(20)
                raise RuntimeError("late")
        release = threading.Event()
        res, _ = check("070126", deps=Stuck())
        release.set()
        assert res["flags"] == ["stop", "timeout"] and res["text"].splitlines()[:3] == [
            said + "the database may hold part of it until a new load replaces it.", "", "LOAD CHECK: STOP"], res["text"]
        # while a load is running or was cut off, the lock says it: no such line then
        dead_lock(failed)
        res, _ = check("070126")
        assert res["text"].splitlines()[1].startswith(f"LOAD INTERRUPTED at refresh (run {failed})")
        assert "The last load" not in res["text"] and lc.last_failed_load(FakeDeps()) is None
        paths.lock_file().unlink()
        assert lc.last_failed_load(FakeDeps()) == {"run": failed, "step": "refresh"}
        # a load has finished since: no such line
        status.write_state({"last_load": {"run": "060126_0900", "finished_at": "2999-01-01T00:00:00+00:00"}})
        assert lc.last_failed_load(FakeDeps()) is None
    with data_root():                         # nothing had been written: no such line
        failed_load(written=False)
        assert status.failed_load() is None and lc.last_failed_load(FakeDeps()) is None

        def unreadable():                     # a journal that cannot be read at this moment never stops a check
            raise PermissionError(13, "in use")
        assert lc.last_failed_load(FakeDeps(failed_load=unreadable)) is None
        assert lc.last_failed_load(FakeDeps(failed_load=lambda: {"step": "refresh"})) is None       # names no run
        assert lc.failed_line({"run": "x", "step": None}, True) == ("The last load (run x) failed at ? after it had begun "
                                                                    "to write: this check starts a new load that replaces it.")


# ── the folder ────────────────────────────────────────────────────────────────

@test("the folder: two files of one report, a workbook open in Excel and a file still being written each stop")
def _():
    with data_root() as root:
        drop = make_drop(root)
        ledger_file(drop / "Stock_ledger999999_Test_Co.xlsx")
        res, _ = check("070126")
        assert res["exit"] == 2 and "Stock ledger: 2 files" in res["text"]
        assert NAMES["stock_movements"] in res["text"] and "Stock_ledger999999_Test_Co.xlsx" in res["text"]
        (drop / "Stock_ledger999999_Test_Co.xlsx").unlink()
        (drop / ("~$" + NAMES["orders"][2:])).write_bytes(b"lock")
        res, _ = check("070126")
        assert "open in Excel" in res["text"] and "~$" in res["text"] and res["token"] is None
        (drop / ("~$" + NAMES["orders"][2:])).unlink()
        deps = FakeDeps()
        deps.on_pause = lambda: (drop / NAMES["orders"]).open("ab").write(b"more") and None
        res, _ = check("070126", deps=deps)
        assert "still being written" in res["text"] and NAMES["orders"] in res["text"]
        assert not (root / "business_data").exists(), "nothing is staged before the folder is sound"


@test("the folder: a report with no As on date stops; a file that is not a Focus export stops")
def _():
    with data_root() as root:
        drop = make_drop(root)
        stock_balance_file(drop / NAMES["stock_balance"], as_on=None)
        res, _ = check("070126")
        assert res["exit"] == 2 and "As on date" in res["text"] and NAMES["stock_balance"] in res["text"]
        stock_balance_file(drop / NAMES["stock_balance"])
        plain_report(drop / NAMES["orders"], "Summary sales register", ["Date", "Invoice"], company=False)
        res, _ = check("070126")
        assert "does not look like a Focus export" in res["text"]
        (drop / NAMES["orders"]).write_bytes(b"")
        res, _ = check("070126")
        assert "is empty" in res["text"]
        res, _ = check("010203")
        assert "No drop folder" in res["text"] and res["exit"] == 2


@test("the folder: other files and sub-folders are listed with where they go; a partial drop names what stays older")
def _():
    with data_root() as root:
        drop = make_drop(root, leave_out=("stock_movements", "receivables", "product_profitability",
                                          "selling_prices:MA_base", "ledger_entries"))
        (drop / "PO-17.pdf").write_bytes(b"%PDF")
        (drop / "Transactions_17.xml").write_text("<x/>", encoding="utf-8")
        (drop / "photo.jpg").write_bytes(b"jpg")
        (drop / "Stock_movement100009_Test_Co.xlsx").write_bytes(b"not read")
        (drop / "papers").mkdir()
        (drop / "papers" / "invoice.pdf").write_bytes(b"%PDF")
        (drop / "photos").mkdir()
        (drop / "photos" / "a.jpg").write_bytes(b"jpg")
        before = snapshot_of(drop)
        res, deps = check("070126")
        assert res["ok"] and res["token"], res["text"]
        text = res["text"]
        table = table_of(text)
        assert table[shown("order_lines")] == "LOAD" and table[shown("orders")] == "LOAD"
        assert table["(no Stock ledger)"] == table["(no Receivables ageing)"] == table["(no Profitability report)"] == "NOT LOADED"
        for name in ("PO-17.pdf", "Transactions_17.xml", "photo.jpg", "Stock_movement100009_Test_Co.xlsx", "papers/", "photos/"):
            assert table[name] == "NOT LOADED", name
        assert "PO-17.pdf" in text and "not loaded by this command; PO or supplier invoice: it goes on the Orders page" in text
        assert "MRN XML: it goes on the Orders page" in text and "these go on the Orders page" in text
        assert "not loaded by this command; it stays in the folder" in text
        assert "not loaded by this command (stock movement summary" in text
        assert "stock movements stay as loaded to Mon 05 Jan" in text and "receivables stay as on Mon 05 Jan" in text
        assert "partial_drop" in res["flags"]
        staged = sorted(p.name for p in loadplan.staging_path(res["run"]).iterdir())
        assert staged == sorted(["Main_Store_only_" + NAMES["stock_balance"], NAMES["order_lines"], NAMES["orders"]]), staged
        assert snapshot_of(drop) == before


@test("the folder: reports only in sub-folders: one set is taken, two sets stop with 'Which set?'")
def _():
    with data_root() as root:
        outer = root / "Uplaods" / "070126"
        inner = make_drop(root, name="070126/As on reports")
        res, deps = check("070126")
        assert res["ok"] and res["run"].startswith("070126_"), res["text"]
        assert "sub-folder 'As on reports'" in res["text"]
        assert loadplan.read_plan(res["run"])["drop"] == str(inner)
        make_drop(root, name="070126/Second set", leave_out=("ledger_entries",))
        res, _ = check("070126")
        assert res["exit"] == 2 and "Which set?" in res["text"] and "As on reports, Second set" in res["text"]
        (outer / "empty").mkdir()
        res, _ = check(str(outer / "empty"))
        assert "No Focus report was found" in res["text"]


# ── the stock-balance trim ────────────────────────────────────────────────────

@test("the trim keeps the title, the header and only the main store's block, cell for cell")
def _():
    from scripts import ingest as ing
    with data_root() as root:
        src = stock_balance_file(root / NAMES["stock_balance"])
        grid = ing.read_grid(src)
        rows = ing.parse_stock_balance(grid, src.name)
        assert sorted({r["warehouse_name"] for r in rows}) == [MAIN, "Van A", "Van B"] and len(rows) == 6
        block = lc.stock_block(grid, MAIN, {})
        assert block["warehouses"] == ["Van A", MAIN, "Van B"] and block["sub_units"] == 160 and block["sub_value"] == 265
        dst = lc.write_trimmed(src, root / "business_data" / "staging" / "x" / ("Main_Store_only_" + src.name), grid, block)
        ws = load_workbook(dst).worksheets[0]
        cells = [[c.value for c in r] for r in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=4)]
        assert cells[0][2] == COMPANY and cells[3][2] == f"[As on date {AS_ON}]" and cells[5][0] == "Particulars"
        assert cells[6] == [f"{MAIN} {MAIN}", None, None, None]
        assert cells[7:10] == [["Cable one C1", 100, 1.5, 150], ["Charger two C2", 40, 2.5, 100], ["Holder three C3", 20, 0.75, 15]]
        assert cells[10] == ["Sub Total", 160, None, 265] and len(cells) == 11
        proof = lc.prove_block(rows, dst, MAIN, {}, block)
        assert proof["problems"] == [] and (proof["items"], proof["units"], proof["value"]) == (3, MAIN_UNITS, MAIN_VALUE)
        assert proof["as_on"] == DAY and proof["raw_warehouses"] == [MAIN]
        assert ing.classify(dst.name) == "stock_balance"
        raises(PermissionError, lambda: lc.write_trimmed(src, root / "Uplaods" / "trim.xlsx", grid, block))
        # merged cells: the title's stay, any below the header go with the rows around them
        wb = load_workbook(src)
        wb.worksheets[0].merge_cells("C1:D1")
        wb.worksheets[0].merge_cells("B16:D16")                       # the name row of the third warehouse
        wb.save(root / "merged.xlsx")
        grid2 = ing.read_grid(root / "merged.xlsx")
        out = lc.write_trimmed(root / "merged.xlsx", root / "business_data" / "staging" / "x" / "merged.xlsx", grid2,
                               lc.stock_block(grid2, MAIN, {}))
        kept = load_workbook(out).worksheets[0]
        assert [str(r) for r in kept.merged_cells.ranges] == ["C1:D1"] and kept.max_row == 11
        assert lc.prove_block(ing.parse_stock_balance(grid2, "merged.xlsx"), out, MAIN, {}, lc.stock_block(grid2, MAIN, {}))["problems"] == []
        # an alias leads another spelling to the same store
        other = stock_balance_file(root / "alias.xlsx", [("Main Store 77", STOCK[1][1]), ("Van A", STOCK[0][1])])
        assert lc.stock_block(ing.read_grid(other), MAIN, {"Main Store 77": MAIN})["start"] == 6


@test("the trim fails when Focus's Sub Total disagrees, when the store is absent, and when the block has no end")
def _():
    from scripts import ingest as ing
    with data_root() as root:
        src = stock_balance_file(root / NAMES["stock_balance"], sub_total_off=1)
        grid = ing.read_grid(src)
        block = lc.stock_block(grid, MAIN, {})
        dst = lc.write_trimmed(src, root / "business_data" / "staging" / "x" / src.name, grid, block)
        proof = lc.prove_block(ing.parse_stock_balance(grid, src.name), dst, MAIN, {}, block)
        assert len(proof["problems"]) == 1 and "Sub Total says 161" in proof["problems"][0], proof["problems"]
        vans = stock_balance_file(root / "vans.xlsx", [STOCK[0], STOCK[2]])
        exc = raises(lc.Stop, lambda: lc.stock_block(ing.read_grid(vans), MAIN, {}))
        assert "does not hold the main store" in str(exc)
        twice = stock_balance_file(root / "twice.xlsx", [STOCK[1], STOCK[1]])
        assert "twice" in str(raises(lc.Stop, lambda: lc.stock_block(ing.read_grid(twice), MAIN, {})))
        open_end = stock_balance_file(root / "open.xlsx", [STOCK[1]], closing=None, grand=False)
        assert "no total row" in str(raises(lc.Stop, lambda: lc.stock_block(ing.read_grid(open_end), MAIN, {})))
        holed = stock_balance_file(root / "holed.xlsx", [(MAIN, STOCK[1][1] + [("Lamp four C4", None, None)]), STOCK[0]])
        exc = raises(lc.Stop, lambda: lc.stock_block(ing.read_grid(holed), MAIN, {}))
        assert "'Lamp four C4' has no quantity, rate or value" in str(exc)
        # a copy that holds something else is not proven either
        whole = lc.prove_block(ing.parse_stock_balance(ing.read_grid(stock_balance_file(root / "ok.xlsx")), "ok.xlsx"),
                               root / "ok.xlsx", MAIN, {}, lc.stock_block(ing.read_grid(root / "ok.xlsx"), MAIN, {}))
        assert any("not the main store alone" in p for p in whole["problems"])
    with data_root() as root:                 # Focus filtered to the one store: nothing to cut, loaded as it is
        drop = make_drop(root)
        one = stock_balance_file(drop / NAMES["stock_balance"], [STOCK[1]], closing="Total", grand=False)
        block = lc.stock_block(ing.read_grid(one), MAIN, {})
        assert block["total_label"] == "Total" and block["warehouses"] == [MAIN] and block["sub_units"] == 160
        res, deps = check("070126")
        assert res["ok"] and table_of(res["text"])[shown("stock_balance")] == "LOAD", res["text"]
        staged = loadplan.staging_path(res["run"]) / NAMES["stock_balance"]
        assert loadplan.sha256_file(staged) == loadplan.sha256_file(one)
        assert "Main Store only: 3 items / 160 units / BHD 265.000" in res["text"] and "Focus's own Total row" in res["text"]
    with data_root() as root:
        make_drop(root, sub_total_off=1)
        res, deps = check("070126")
        assert res["exit"] == 2 and res["token"] is None and "could not be proven" in res["text"]
        assert table_of(res["text"])[shown("stock_balance")] == "STOP"
        staged = [p.name for p in loadplan.staging_path(res["run"]).iterdir()]
        assert not any("Stock_balance" in n for n in staged), staged
        assert not (loadplan.run_dir(res["run"]) / "go.cmd").exists()


@test("the trim's proof, part by part: the Sub Total's value, an altered row, another date, a sheet out of line, a copy that differs")
def _():
    from scripts import ingest as ing
    with data_root() as root:
        stage = root / "business_data" / "staging" / "x"
        src = stock_balance_file(root / NAMES["stock_balance"], sub_value_off=1)        # the units agree, the value does not
        grid = ing.read_grid(src)
        block = lc.stock_block(grid, MAIN, {})
        dst = lc.write_trimmed(src, stage / src.name, grid, block)
        proof = lc.prove_block(ing.parse_stock_balance(grid, src.name), dst, MAIN, {}, block)
        assert proof["problems"] == ["the rows add up to BHD 265.000, Focus's Sub Total says 266.000"], proof["problems"]

        good = stock_balance_file(root / "good.xlsx")
        grid = ing.read_grid(good)
        rows, block = ing.parse_stock_balance(grid, good.name), lc.stock_block(grid, MAIN, {})
        copy = lc.write_trimmed(good, stage / "good.xlsx", grid, block)
        assert lc.prove_block(rows, copy, MAIN, {}, block)["problems"] == []
        wb = load_workbook(copy)
        wb.worksheets[0].cell(row=8, column=1, value="Cable one C9")                   # one item row is not the original's
        wb.save(stage / "renamed.xlsx")
        assert lc.prove_block(rows, stage / "renamed.xlsx", MAIN, {}, block)["problems"] == [
            "the copy's item rows differ from the block in the original (3 against 3)"]
        wb = load_workbook(copy)
        wb.worksheets[0]["C4"] = "[As on date 08/01/2026]"                             # the title carries another day
        wb.save(stage / "redated.xlsx")
        assert lc.prove_block(rows, stage / "redated.xlsx", MAIN, {}, block)["problems"] == [
            "the copy's as-on date is not the original's"]
        wb = load_workbook(copy)
        wb.worksheets[0]["C4"] = None                                                  # ... or no day at all
        wb.save(stage / "undated.xlsx")
        assert "the copy's as-on date is not the original's" in lc.prove_block(
            rows, stage / "undated.xlsx", MAIN, {}, block)["problems"]

        # the sheet must line up with the rows the parser read: a grid one row off is refused before anything is cut
        shifted = grid.iloc[1:].reset_index(drop=True)
        exc = raises(lc.Stop, lambda: lc.write_trimmed(good, stage / "shifted.xlsx", shifted,
                                                       lc.stock_block(shifted, MAIN, {})))
        assert "does not line up" in str(exc) and not (stage / "shifted.xlsx").exists()

        # every copy is proven by its sha256: one that differs from its original is refused
        real = loadplan.sha256_file
        loadplan.sha256_file = lambda p: real(p) + ("x" if Path(p).parent == stage else "")
        try:
            exc = raises(lc.Stop, lambda: lc.copy_checked(good, stage / "copied.xlsx"))
        finally:
            loadplan.sha256_file = real
        assert "the copy is not identical to the file" in str(exc)
        assert lc.copy_checked(good, stage / "copied.xlsx") == loadplan.sha256_file(good) == loadplan.sha256_file(stage / "copied.xlsx")
        raises(PermissionError, lambda: lc.copy_checked(good, root / "Uplaods" / "copied.xlsx"))
    with data_root() as root:                 # an item whose NAME reads as a total row closes the block too early
        named = (MAIN, [("Cable one C1", 100, 1.5), ("Total care kit K1", 5, 2), ("Holder three C3", 20, 0.75)])
        make_drop(root, stock=[STOCK[0], named, STOCK[2]])
        res, _ = check("070126")
        text = res["text"]
        assert res["exit"] == 2 and res["token"] is None and table_of(text)[shown("stock_balance")] == "STOP", text
        assert "the staged copy could not be proven" in text
        assert "The block was closed at the row 'Total care kit K1', which is read as a total row because of its name" in text
        assert "the loader would leave it out. Tell Claude" in text
        fmt.careful(text)
    with data_root() as root:                 # Focus's own closing rows get no such remark
        make_drop(root, sub_total_off=1)
        res, _ = check("070126")
        assert "could not be proven" in res["text"] and "because of its name" not in res["text"]
        assert set(lc.CLOSING_ROWS) == {"sub total", "total", "grand total"}
    with data_root() as root:                 # production's last snapshot cannot be read: the copy is not proven
        make_drop(root)

        class NoSnapshot(FakeDeps):
            def read_params(self, sql, params):
                if "from stock_balance" in sql:
                    raise RuntimeError("statement timeout on db.secret.example")
                return super().read_params(sql, params)
        res, _ = check("070126", deps=NoSnapshot())
        assert res["exit"] == 2 and res["token"] is None and table_of(res["text"])[shown("stock_balance")] == "STOP"
        assert "production's last snapshot could not be read (RuntimeError)" in res["text"] and "secret" not in res["text"]
        assert not any("Stock_balance" in p.name for p in loadplan.staging_path(res["run"]).iterdir())


# ── the whole check ───────────────────────────────────────────────────────────

@test("a clean drop: the table's words, what is staged, the token, the GO line, go.cmd and plan.json")
def _():
    with data_root() as root:
        drop = make_drop(root)
        before = snapshot_of(drop)
        res, deps = check("070126")
        text = res["text"]
        assert res["ok"] and res["exit"] == 0 and res["command"] == "load" and res["data_date"] == DAY, text
        run, token = res["run"], res["token"]
        assert run == "070126_0930" and len(token) == 32 and token.isalpha()
        assert [ln[:5] for ln in deps.lines] == [f"[{i}/9]" for i in range(1, 10)], deps.lines
        table = table_of(text)
        assert table == {shown("order_lines"): "LOAD", shown("orders"): "LOAD", shown("stock_movements"): "LOAD",
                         shown("stock_balance"): "LOAD (main-store block only)", shown("receivables"): "LOAD",
                         shown("product_profitability"): "LOAD", shown("selling_prices:MA_base"): "LEAVE OUT",
                         shown("ledger_entries"): "NOT LOADED"}, table
        assert shown("receivables") != NAMES["receivables"] and len(shown("receivables")) == 46      # a long name is cut
        staging = loadplan.staging_path(run)
        staged = sorted(p.name for p in staging.iterdir())
        assert staged == sorted([NAMES["order_lines"], NAMES["orders"], NAMES["stock_movements"], NAMES["receivables"],
                                 NAMES["product_profitability"], "Main_Store_only_" + NAMES["stock_balance"]]), staged
        assert deps.previewed == [staged], "the preview saw exactly what will load"
        for name in (NAMES["order_lines"], NAMES["stock_movements"]):
            assert loadplan.sha256_file(staging / name) == loadplan.sha256_file(drop / name)
        assert "never loaded by this command (as on Wed 07 Jan)" in text
        assert "3 items / 160 units / BHD 265.000" in text and "3 warehouses in the file" in text
        assert "entered with an earlier date 0, new 1" in text and "Loader guards: none is predicted to fire." in text
        assert "After the load you should see" in text and "Sales to Tue 06 Jan: 12 invoices, 12 lines in all." in text
        assert "BHD 150.500 on 7 accounts (Focus's own total: BHD 140.250)" in text
        assert "demand forecast" in text and "Telegram" in text and "catalog" in text and "alias" in text
        assert "scripts/my_notes.py" in text and "scripts/ops/load_check.py" not in text      # untracked, named only
        lines = text.splitlines()
        assert lines[-1] == "GO: " + loadplan.go_line(run, token) and lines[-1].startswith("GO: python -m scripts.yq_load_go --run ")
        assert [ln for ln in lines if ln.startswith("GO: ")] == [lines[-1]], "one GO line, and it is the last line"
        # the line just before it says where go.cmd lies, for a double-click when Claude may not run the GO line
        go_cmd = loadplan.run_dir(run) / "go.cmd"
        assert lines[-2] == f"If Claude may not start it: double-click {go_cmd}" and go_cmd.is_file(), lines[-2]
        assert go_cmd == root / "business_data" / "ops" / "loads" / run / "go.cmd"
        assert lines[-3].startswith(f"Token {token}: ") and lines[-4].startswith(f"Staged in business_data/staging/{run}: ")
        assert res["flags"] == ["token"] and res["warnings"] == []
        cmd = go_cmd.read_text(encoding="utf-8")
        assert loadplan.go_line(run, token) in cmd and str(paths.CODE_ROOT) in cmd
        plan = loadplan.read_plan(run)
        assert plan["token"] == token and plan["run"] == run and plan["drop"] == str(drop) and plan["staged"] == staged
        assert set(plan) >= {"run", "drop", "created_at", "verdicts", "expected", "options", "token", "token_parts", "as_on"}
        assert [v["verdict"] for v in plan["verdicts"]] == ["load", "load", "load", "trim", "load", "load", "leave_out", "not_loaded"]
        assert all(v["verdict"] in loadplan.VERDICTS for v in plan["verdicts"])
        assert plan["expected"]["stock"] == {"as_on": DAY, "items": 3, "units": "160", "value_bhd": "265.000"}
        assert plan["expected"]["invoices"] == 12 and plan["as_on"]["stock_balance"] == DAY and plan["stops"] == []
        assert plan["token_parts"]["live"] == "abc123abc123" and len(plan["token_parts"]["head"]) == 40
        assert set(plan["token_parts"]["staging"]) == set(staged) and set(plan["token_parts"]["drop"]) == set(NAMES.values())
        assert loadplan.token_of(plan["token_parts"]) == token
        assert deps.the_session.closed and snapshot_of(drop) == before
        fmt.careful(text)
        own = "\n".join(ln for ln in lines if ln != lines[-2])       # all but the go.cmd line: a path of this PC
        assert not NINE_DIGITS.search(own), "a run of digits that could read as an ID number"
        assert all(ord(ch) < 128 for ch in own), "the block is plain ASCII"
        # nothing changed, nothing on the register beyond the first-run markers: one per rule, each with nothing found
        assert alerts.open_items(include_baseline=True) == []
        marks = [json.loads(ln) for ln in alerts.alerts_file().read_text(encoding="utf-8").splitlines()]
        assert sorted((m["rule"], m["event"], m["found"]) for m in marks) == [
            ("inv_changed", "first_run", 0), ("price_scope", "first_run", 0), ("profit_held", "first_run", 0),
            ("voucher_changed", "first_run", 0)], marks
        assert all(alerts.has_run(rule) for rule in lc.EVENT_RULES + lc.STATE_RULES)


@test("the token follows the staged files, the drop and the database: any of them moving changes it")
def _():
    with data_root() as root:
        drop = make_drop(root)
        res, deps = check("070126")
        run, token = res["run"], res["token"]
        plan = loadplan.read_plan(run)

        def again(session=None) -> dict:
            return loadplan.collect_parts(drop, run, paths.CODE_ROOT, deps.live, session or FakeSession())
        assert loadplan.token_of(again()) == token
        staged = loadplan.staging_path(run) / NAMES["orders"]
        staged.write_bytes(staged.read_bytes() + b"x")
        moved = again()
        assert loadplan.token_of(moved) != token
        assert loadplan.moved_parts(plan["token_parts"], moved) == ["the staged copies: " + NAMES["orders"]]
        res2, _ = check("070126")                                    # a second check is a new run with its own plan
        assert res2["run"] == "070126_0931" and loadplan.read_plan(res2["run"])["token"] == res2["token"]
        assert loadplan.staging_path(res2["run"]).is_dir() and (loadplan.run_dir(res2["run"]) / "go.cmd").is_file()
        (drop / NAMES["receivables"]).write_bytes(b"changed")
        parts = loadplan.collect_parts(drop, res2["run"], paths.CODE_ROOT, deps.live, FakeSession())
        assert loadplan.moved_parts(loadplan.read_plan(res2["run"])["token_parts"], parts) == ["the drop folder: " + NAMES["receivables"]]
        parts = loadplan.collect_parts(drop, run, paths.CODE_ROOT, deps.live, FakeSession(count={"n": 11, "max_id": 100}))
        assert "the database table orders" in loadplan.moved_parts(moved, parts)


@test("--dev: no token, no GO line, no go.cmd, nothing recorded; the rest of the check is the same")
def _():
    with data_root() as root:
        prices = PRICES + [("Charger two", "C2", None, None, "Van A", "1/5/2026", "12/31/2030", 2.0)]
        drop = make_drop(root, prices=prices)
        before = snapshot_of(drop)
        res, deps = check("070126", "--dev", deps=FakeDeps(worktree=True, site={"state": "behind"}))
        text = res["text"]
        assert res["ok"] and res["exit"] == 0 and res["token"] is None and "dev" in res["flags"], text
        assert "GO:" not in text and "Token" not in text and "[DEV RUN]" in text.splitlines()[0]
        assert text.splitlines()[-1].startswith("DEV RUN: nothing stops this drop")
        run_dir = loadplan.run_dir(res["run"])
        assert (run_dir / "plan.json").is_file() and not (run_dir / "go.cmd").exists()
        plan = loadplan.read_plan(res["run"])
        assert plan["token"] is None and plan["token_parts"] is None and plan["options"]["dev"] is True
        assert table_of(text)[shown("selling_prices:MA_base")] == "NEEDS YOUR DECISION"
        assert "1 item for the action list (price_scope): not recorded in a dev run." in text
        assert not alerts.alerts_file().exists()
        assert snapshot_of(drop) == before


# ── one check counts for a folder: the newest ─────────────────────────────────

def live_runs() -> list[str]:
    """The runs that have a staging folder a load could work from."""
    base = paths.staging_dir()
    return sorted(p.name for p in base.iterdir() if p.is_dir() and loadplan.is_run(p.name)) if base.is_dir() else []


def put_away() -> list[str]:
    base = paths.staging_dir() / lc.SUPERSEDED
    return sorted(p.name for p in base.iterdir()) if base.is_dir() else []


def token_now(drop: Path, run: str, live: str = "abc123abc123") -> str:
    """The token the writer would work out for a run at this moment."""
    return loadplan.token_of(loadplan.collect_parts(drop, run, paths.CODE_ROOT, live, FakeSession()))


@test("a second check of a folder replaces the first: one live staging folder, the first go.cmd gone, its plan names the second")
def _():
    with data_root() as root:
        drop = make_drop(root)
        before = snapshot_of(drop)
        one, _ = check("070126")
        first, token = one["run"], one["token"]
        staged = sorted(p.name for p in loadplan.staging_path(first).iterdir())
        assert token and (loadplan.run_dir(first) / "go.cmd").is_file() and live_runs() == [first]
        assert "replaces" not in one["text"] and "replaced_earlier" not in one["flags"]
        assert loadplan.read_plan(first)["supersedes"] == [] and "superseded_by" not in loadplan.read_plan(first)
        assert token_now(drop, first) == token

        two, deps = check("070126")
        second = two["run"]
        assert two["ok"] and two["token"] and second != first, two["text"]
        assert live_runs() == [second], "one live staging folder"
        assert not (loadplan.run_dir(first) / "go.cmd").exists(), "the first run's go.cmd is gone"
        old = loadplan.read_plan(first)
        assert old["superseded_by"] == second and old["token"] == token and old["run"] == first, "its plan names the second"
        assert [v["verdict"] for v in old["verdicts"]] == [v["verdict"] for v in loadplan.read_plan(second)["verdicts"]]
        assert put_away() == [first]
        assert sorted(p.name for p in (paths.staging_dir() / lc.SUPERSEDED / first).iterdir()) == staged
        assert token_now(drop, first) != token, "the first token no longer fits what is on disk"
        assert loadplan.read_plan(second)["supersedes"] == [first] and "replaced_earlier" in two["flags"]
        assert (f"- This check replaces 1 earlier check of this folder ({first}): its go line no longer works, and the "
                "staged copies moved to business_data/staging/superseded.") in two["text"], two["text"]
        assert deps.steps[2] == f"[3/9] Staging {second}: 4 reports copied, each identical to its original; 1 earlier " \
                                "check of this folder replaced."
        assert two["text"].splitlines()[-1] == "GO: " + loadplan.go_line(second, two["token"])
        assert (loadplan.run_dir(second) / "go.cmd").is_file() and token_now(drop, second) == two["token"]

        three, _ = check("070126", "--pricebook", "skip")           # and again: only the newest counts
        third = three["run"]
        assert live_runs() == [third] and put_away() == [first, second]
        assert loadplan.read_plan(first)["superseded_by"] == second, "a check is replaced once"
        assert loadplan.read_plan(second)["superseded_by"] == third and loadplan.read_plan(third)["supersedes"] == [second]
        assert "This check replaces 1 earlier check of this folder" in three["text"]
        # `load status` says so when he asks about a check that was replaced
        seen = status_of("--run", first)
        assert seen["state"] == "NOT STARTED" and seen["text"].splitlines() == [
            f"Load {first}", f"  (this check was replaced by a newer one, run {second}: its go line no longer works)",
            "  (no new journal line)", "NOT STARTED"], seen["text"]
        assert "replaced" not in status_of("--run", third)["text"]
        assert snapshot_of(drop) == before


@test("a check that stops replaces the earlier one all the same, and says so; a check that never got to staging replaces nothing")
def _():
    with data_root() as root:
        make_drop(root)
        one, _ = check("070126")
        first = one["run"]
        res, _ = check("070126", deps=FakeDeps(site={"state": "behind"}))          # stops at the preflight
        assert res["exit"] == 2 and live_runs() == [first] and (loadplan.run_dir(first) / "go.cmd").is_file()
        assert "superseded_by" not in loadplan.read_plan(first) and "replaces" not in res["text"]

        class NoDb(FakeDeps):
            def session(self):
                raise RuntimeError("no answer")
        res, _ = check("070126", deps=NoDb())                                       # stops after staging, before its table
        lines = res["text"].splitlines()
        assert res["exit"] == 2 and res["flags"] == ["stop", "replaced_earlier"], res["text"]
        assert lines[2] == "No token was issued. Nothing was changed."
        assert lines[3].startswith(f"This check replaces 1 earlier check of this folder ({first}): its go line no longer works")
        assert lines[-1] == f"(Left by this check: business_data/staging/{res['run']}.)"
        assert loadplan.read_plan(first)["superseded_by"] == res["run"] and not (loadplan.run_dir(first) / "go.cmd").exists()
        assert live_runs() == [res["run"]] and put_away() == [first]


def copy_of_run(run: str, name: str, token: str, root: Path) -> dict:
    """The same check under another run name: its staging folder, its plan and its go.cmd. (The token binds the
    files, the programs and the database, not the name: it is the same token.)"""
    shutil.copytree(loadplan.staging_path(run), loadplan.staging_path(name))
    plan = dict(loadplan.read_plan(run), run=name, staging=str(loadplan.staging_path(name)))
    loadplan.write_plan(plan)
    loadplan.write_go_cmd(name, token, root)
    return plan


@test("never replaced: a run that started, finished or failed; a check of another folder with the same date; the run that holds the lock")
def _():
    with data_root() as root:
        drop = make_drop(root)
        one, _ = check("070126")
        token = one["token"]
        ends = {"070126_0810": ("started",), "070126_0820": ("started", "finished"), "070126_0830": ("started", "failed")}
        plans = {}
        for name, events in ends.items():     # three more checks of this folder, each one a load that went further
            plans[name] = copy_of_run(one["run"], name, token, root)
            for event in events:
                loadplan.journal_append(name, event, step="backup", written=False)
        assert [loadplan.run_result(name) for name in ends] == ["RUNNING", "FINISHED", "FAILED"]
        assert loadplan.run_result(one["run"]) == "NOT_STARTED" and all(token_now(drop, name) == token for name in ends)
        res, _ = check("070126")
        assert res["ok"] and res["token"], res["text"]
        assert loadplan.read_plan(res["run"])["supersedes"] == [one["run"]] and put_away() == [one["run"]]
        assert f"This check replaces 1 earlier check of this folder ({one['run']})" in res["text"]
        for name, plan in plans.items():
            assert loadplan.staging_path(name).is_dir() and (loadplan.run_dir(name) / "go.cmd").is_file(), name
            assert loadplan.read_plan(name) == plan and token_now(drop, name) == token, name
        assert live_runs() == sorted(list(ends) + [res["run"]])
    with data_root() as root:                 # two folders whose reports carry one date: each is its own drop
        other = make_drop(root, name="incoming")
        here = make_drop(root)
        one, _ = check(str(other))
        two, _ = check("070126")
        assert one["run"][:6] == two["run"][:6] == "070126" and one["token"] and two["token"]
        assert live_runs() == [one["run"], two["run"]] and put_away() == [] and "replaces" not in two["text"]
        assert "superseded_by" not in loadplan.read_plan(one["run"]) and (loadplan.run_dir(one["run"]) / "go.cmd").is_file()
        assert lc.same_folder(str(other), other) and lc.same_folder(here, root / "Uplaods" / "070126")
        assert not lc.same_folder(other, here) and not lc.same_folder("", here) and not lc.same_folder("x", here)
        assert lc.earlier_runs(two["run"]) == [one["run"]] and lc.earlier_runs("080126_0900") == []
    with data_root() as root:                 # a go that has just taken the lock is starting this very run
        drop = make_drop(root)
        one, _ = check("070126")

        class Starting(FakeDeps):
            """The lock appears after the preflight looked: a go took it meanwhile."""
            looks = 0

            def lock_state(self):
                self.looks += 1
                return {"state": "none"} if self.looks <= 2 else {"state": "running", "run": one["run"], "step": "start"}
        res, deps = check("070126", deps=Starting())
        assert deps.looks == 3, "asked once more just before anything is put away"
        assert loadplan.staging_path(one["run"]).is_dir() and "superseded_by" not in loadplan.read_plan(one["run"])
        assert (loadplan.run_dir(one["run"]) / "go.cmd").is_file() and "replaces" not in res["text"]
        assert token_now(drop, one["run"]) == one["token"]


@test("a dev run replaces nothing and is itself replaced by the next real check")
def _():
    with data_root() as root:
        drop = make_drop(root)
        one, _ = check("070126")
        first, token = one["run"], one["token"]
        plan = loadplan.read_plan(first)
        dev, deps = check("070126", "--dev", deps=FakeDeps(worktree=True))
        assert dev["ok"] and dev["token"] is None and "replaces" not in dev["text"], dev["text"]
        assert "replaced" not in deps.steps[2] and loadplan.read_plan(dev["run"])["supersedes"] == []
        assert live_runs() == [first, dev["run"]] and put_away() == []
        assert loadplan.read_plan(first) == plan and (loadplan.run_dir(first) / "go.cmd").is_file()
        assert token_now(drop, first) == token, "the real check can still be loaded"
        real, _ = check("070126")
        assert live_runs() == [real["run"]] and put_away() == [first, dev["run"]]
        assert loadplan.read_plan(real["run"])["supersedes"] == [first, dev["run"]]
        # the check that had a go line is named first; the dev run never had one: only its copies were put away
        assert (f"- This check replaces 1 earlier check of this folder ({first}): its go line no longer works, and the "
                "staged copies moved to business_data/staging/superseded. The staged copies of 1 earlier check that never "
                f"had a go line ({dev['run']}) were put away there as well.") in real["text"], real["text"]
        assert loadplan.read_plan(dev["run"])["superseded_by"] == real["run"]
    # the same sentences, case by case
    said = lc.retired_line
    assert said({}) is None and said({"retired": [], "no_go": [], "leftovers": [], "stuck": []}) is None
    assert said({"retired": ["a", "b"]}) == ("This check replaces 2 earlier checks of this folder (a, b): their go lines "
                                             "no longer work, and the staged copies moved to "
                                             "business_data/staging/superseded.")
    assert said({"retired": ["a", "b"], "stuck": ["b"]}) == (
        "This check replaces 2 earlier checks of this folder (a, b): their go lines no longer work; the staged copies of a "
        "moved to business_data/staging/superseded. The staging folder of b could not be moved (a file in it is open): it "
        "stays where it is and can no longer be loaded.")
    assert said({"retired": ["a", "b"], "no_go": ["a", "b"], "leftovers": ["c"]}) == (
        "The staged copies of 3 earlier checks that never had a go line (a, b, c) were put away in "
        "business_data/staging/superseded.")
    assert said({"retired": ["a"], "no_go": ["a"], "stuck": ["a"]}) == (
        "The staging folder of a could not be moved (a file in it is open): it stays where it is and can no longer be "
        "loaded.")


@test("the copies of a check that stopped before its table are put away once it cannot still be running")
def _():
    with data_root() as root:
        make_drop(root)

        class NoDb(FakeDeps):
            def session(self):
                raise RuntimeError("no answer")
        left, _ = check("070126", deps=NoDb())
        stale = left["run"]
        assert live_runs() == [stale] and not loadplan.run_dir(stale).exists()      # staged copies, no plan
        res, _ = check("070126")              # made moments later: that check may still be running, so it is left
        assert res["token"] and live_runs() == [stale, res["run"]] and "put away" not in res["text"]
        old = datetime.now(timezone.utc).timestamp() - lc.LEFTOVER_AFTER_S - 5
        os.utime(loadplan.staging_path(stale), (old, old))
        res2, _ = check("070126")
        assert live_runs() == [res2["run"]] and put_away() == [stale, res["run"]]
        assert (f"- This check replaces 1 earlier check of this folder ({res['run']}): its go line no longer works, and "
                "the staged copies moved to business_data/staging/superseded. The staged copies of 1 earlier check that "
                f"never had a go line ({stale}) were put away there as well.") in res2["text"], res2["text"]
        assert loadplan.read_plan(res2["run"])["supersedes"] == [res["run"]]
        assert lc.LEFTOVER_AFTER_S > lc.DEADLINE_S, "no check runs longer than its deadline"


@test("putting a check away: a name already used under superseded/, a folder Windows will not move, a plan that cannot be rewritten")
def _():
    def refuse_moves():
        real = os.replace

        def refuse(src, dst, *a, **k):
            if lc.SUPERSEDED in Path(dst).parts:
                raise PermissionError(13, "Access is denied")
            return real(src, dst, *a, **k)
        os.replace = refuse
        return real

    def refuse_marks():
        real = loadplan.write_plan

        def refuse(plan):
            if plan.get("superseded_by"):
                raise PermissionError(13, "Access is denied")
            return real(plan)
        loadplan.write_plan = refuse
        return real
    with data_root() as root:
        drop = make_drop(root)
        one, _ = check("070126")
        # a folder of that name already lies under superseded/: it keeps its place, the new one gets a time suffix
        (paths.staging_dir() / lc.SUPERSEDED / one["run"]).mkdir(parents=True)
        two, _ = check("070126")
        names = put_away()
        assert len(names) == 2 and names[0] == one["run"] and re.fullmatch(one["run"] + r"_\d{8}-\d{6}", names[1]), names
        assert list((paths.staging_dir() / lc.SUPERSEDED / one["run"]).iterdir()) == [], "the folder that was there is kept"
        assert len(list((paths.staging_dir() / lc.SUPERSEDED / names[1]).iterdir())) == 6 and live_runs() == [two["run"]]

        # Windows will not move the folder (a file in it is open): the plan alone ends its go line
        real = refuse_moves()
        try:
            three, _ = check("070126")
        finally:
            os.replace = real
        assert three["ok"] and three["token"], three["text"]
        assert loadplan.read_plan(two["run"])["superseded_by"] == three["run"]
        assert not (loadplan.run_dir(two["run"]) / "go.cmd").exists()
        assert live_runs() == [two["run"], three["run"]] and token_now(drop, two["run"]) == two["token"]
        assert (f"- This check replaces 1 earlier check of this folder ({two['run']}): its go line no longer works. The "
                f"staging folder of {two['run']} could not be moved (a file in it is open): it stays where it is and can "
                "no longer be loaded.") in three["text"], three["text"]
        assert "moved to business_data" not in three["text"], "what did not move is not said to have moved"
        assert loadplan.read_plan(three["run"])["supersedes"] == [two["run"]]

        # the next check moves that folder (and counts the run once only); this time a plan cannot be rewritten:
        # the folder that moved ends that token all the same
        real = refuse_marks()
        try:
            four, _ = check("070126")
        finally:
            loadplan.write_plan = real
        assert four["ok"] and four["token"] and live_runs() == [four["run"]], four["text"]
        assert loadplan.read_plan(four["run"])["supersedes"] == [three["run"]], "the run replaced before is not counted again"
        assert loadplan.read_plan(two["run"])["superseded_by"] == three["run"], "a check is replaced once"
        assert "superseded_by" not in loadplan.read_plan(three["run"]) and token_now(drop, three["run"]) != three["token"]
        assert not (loadplan.run_dir(three["run"]) / "go.cmd").exists()

        # neither can be done: two go lines for one folder must never both work, so the check stops
        real_move, real_mark = refuse_moves(), refuse_marks()
        try:
            five, _ = check("070126")
        finally:
            os.replace, loadplan.write_plan = real_move, real_mark
        assert five["exit"] == 2 and five["token"] is None and five["flags"] == ["stop"], five["text"]
        assert f"An earlier check of this folder (run {four['run']}) could not be put away" in five["text"]
        assert "two go lines for one folder must never both work" in five["text"] and "GO:" not in five["text"]
        assert f"(Left by this check: business_data/staging/{five['run']}.)" in five["text"] and "replaces" not in five["text"]
        assert token_now(drop, four["run"]) == four["token"], "the earlier check is still the one that counts"
        assert (loadplan.run_dir(four["run"]) / "go.cmd").is_file(), "and so its go.cmd is still there"
        assert "superseded_by" not in loadplan.read_plan(four["run"])
        six, _ = check("070126")              # once the files are free again, the next check puts everything away
        assert six["token"] and live_runs() == [five["run"], six["run"]], "(what the stopped check left is too new to move)"
        assert loadplan.read_plan(six["run"])["supersedes"] == [three["run"], four["run"]], "the plan left unmarked is marked now"
        assert loadplan.read_plan(three["run"])["superseded_by"] == six["run"] == loadplan.read_plan(four["run"])["superseded_by"]


@test("the run name is the minute of the check: a second check inside one minute waits for the next, and never shares a name")
def _():
    with data_root() as root:
        make_drop(root)
        one, first = check("070126")
        assert one["run"] == "070126_0930" and first.pauses == [lc.SIZE_PAUSE_S], first.pauses
        assert WALL["at"] < START + timedelta(minutes=1), "the second check is made inside the same minute"
        two, second = check("070126")
        assert two["run"] == "070126_0931" and two["token"], two["text"]
        wait = second.pauses[1]
        assert len(second.pauses) == 2 and 59 < wait <= 60.05, second.pauses
        assert second.lines[2] == f"Run name 070126_0930 is already taken by an earlier check: waiting {int(wait) + 1} s " \
                                  "for the next minute."
        assert len(second.steps) == 9 and second.steps[2].startswith("[3/9] Staging 070126_0931: ")
        assert WALL["at"].strftime("%H:%M") == "09:31", "it took the name of the minute it was made in, not one ahead"
        assert loadplan.read_plan("070126_0930")["run"] != loadplan.read_plan("070126_0931")["run"]
        assert loadplan.read_plan("070126_0931")["created_at"].startswith("2026-01-07T09:31:")
    with data_root() as root:                 # another check made the folder first (two at one instant): this one waits
        make_drop(root)
        loadplan.staging_path("070126_0930").mkdir(parents=True)                     # a staging folder takes the name
        elsewhere = {"run": "070126_0931", "drop": "elsewhere", "token": None, "verdicts": []}
        loadplan.write_plan(elsewhere)                                               # and so does a run folder alone
        res, deps = check("070126")
        assert res["run"] == "070126_0932" and res["token"] and len(deps.pauses) == 3, (res["run"], deps.pauses)
        assert sum(deps.pauses[1:]) > 119 and len([ln for ln in deps.lines if ln.startswith("Run name ")]) == 2
        assert list(loadplan.staging_path("070126_0930").iterdir()) == [], "the other check's folder is not written into"
        assert loadplan.read_plan("070126_0931") == elsewhere and "replaces" not in res["text"]
    with data_root() as root:                 # three minutes in a row are taken: it stops with a plain sentence
        drop = make_drop(root)
        for name in ("070126_0930", "070126_0931", "070126_0932"):
            loadplan.staging_path(name).mkdir(parents=True)
        before = snapshot_of(drop)
        res, deps = check("070126")
        assert res["exit"] == 2 and res["token"] is None and res["run"] is None, res["text"]
        assert res["text"].splitlines()[1] == ("Other checks of this folder took the run names up to 070126_0932, one every "
                                               "minute. Nothing was staged: wait a minute, then run the check again.")
        assert len(deps.pauses) == 1 + lc.RUN_NAME_WAITS and "Left by this check" not in res["text"]
        assert live_runs() == ["070126_0930", "070126_0931", "070126_0932"] and snapshot_of(drop) == before
        assert not paths.loads_dir().exists()


@test("the loader's marker files are in staging only on the owner's own option, and plan.json tells the same story")
def _():
    import ast
    from scripts.load_supabase import PARTIAL_OK_MARKER, REPLACE_OK_MARKER, marker_present
    assert set(lc.MARKER_OPTIONS) == set(loadplan.MARKERS) == {REPLACE_OK_MARKER, PARTIAL_OK_MARKER}
    assert lc.MARKER_OPTIONS == {REPLACE_OK_MARKER: ("replace_ok", "--replace-ok"),
                                 PARTIAL_OK_MARKER: ("partial_ok", "--partial-ok")}

    def story(run: str) -> tuple[set, set]:
        """(the markers the loader would find in the staged folder, the markers plan.json says he asked for): the
        two sets the writer compares before it shows its box."""
        staging, options = loadplan.staging_path(run), loadplan.read_plan(run)["options"]
        there = {m for m in loadplan.MARKERS if marker_present(staging, m)}
        files = {p.name.split(".")[0].upper() for p in staging.iterdir() if p.is_file()} & set(loadplan.MARKERS)
        assert there == files
        return there, {m for m, (key, _word) in lc.MARKER_OPTIONS.items() if options[key]}
    held = [{"invoice_no": f"INV-{i}", "salesman": "Van A"} for i in range(40)]
    wide = [{"w": MAIN}, {"w": "Van A"}]
    with data_root() as root:
        make_drop(root, prices=PRICES + [("Charger two", "C2", None, None, "Van A", "1/5/2026", "12/31/2030", 2.0)])

        def guarded() -> FakeDeps:            # production holds invoices the file lacks and one more warehouse:
            return FakeDeps(FakeSession(orders=held, warehouses_before=wide), hook_rows={"orders": held[8:]})   # both guards
        cases = [(guarded, (), set(), False), (guarded, ("--replace-ok",), {REPLACE_OK_MARKER}, False),
                 (guarded, ("--partial-ok",), {PARTIAL_OK_MARKER}, False),
                 (guarded, ("--replace-ok", "--partial-ok"), {REPLACE_OK_MARKER, PARTIAL_OK_MARKER}, True),
                 (FakeDeps, ("--pricebook", "load", "--with-profitability"), set(), True),
                 (FakeDeps, ("--pricebook", "skip"), set(), True), (FakeDeps, ("--dev",), set(), False),
                 (FakeDeps, ("--dev", "--replace-ok"), {REPLACE_OK_MARKER}, False),
                 (FakeDeps, ("--pricebook", "load", "--with-profitability", "--partial-ok"), {PARTIAL_OK_MARKER}, True)]
        for make, argv, want, token in cases:     # whether the check goes through, stops on a guard or is a dev run
            res, _ = check("070126", *argv, deps=make())
            there, asked = story(res["run"])
            assert there == asked == want, (argv, there, asked)
            assert bool(res["token"]) is token and "marker_mismatch" not in res["flags"], (argv, res["text"])
            if token:
                plan = loadplan.read_plan(res["run"])
                assert {m for m in loadplan.MARKERS if m in plan["token_parts"]["staging"]} == want
    with data_root() as root:                 # a report named like a marker is never staged: it stops at the folder
        drop = make_drop(root)
        (drop / NAMES["order_lines"]).rename(drop / ("REPLACE_OK." + NAMES["order_lines"]))
        res, _ = check("070126")
        assert res["exit"] == 2 and res["token"] is None and not (root / "business_data").exists(), res["text"]
        assert f"REPLACE_OK.{NAMES['order_lines']} is named like the loader's marker file REPLACE_OK" in res["text"]
        (drop / ("REPLACE_OK." + NAMES["order_lines"])).rename(drop / ("partial_ok." + NAMES["order_lines"]))
        res, _ = check("070126")
        assert res["exit"] == 2 and "is named like the loader's marker file PARTIAL_OK" in res["text"]
        assert lc.marker_of("Replace_Ok") == "REPLACE_OK" and lc.marker_of("partial_ok.txt") == "PARTIAL_OK"
        assert lc.marker_of(NAMES["orders"]) is None and lc.marker_of("REPLACE_OK_notes.xlsx") is None
    with data_root() as root:                 # a marker that got into staging some other way stops the check
        make_drop(root)

        class Planted(FakeDeps):
            def preview(self, folder, *a, **k):
                (Path(folder) / "PARTIAL_OK.txt").write_text("not asked for", encoding="utf-8")
                return super().preview(folder, *a, **k)
        res, _ = check("070126", deps=Planted())
        assert res["exit"] == 2 and res["token"] is None and "marker_mismatch" in res["flags"], res["text"]
        assert ("the staged folder holds a file the loader reads as its marker PARTIAL_OK, and --partial-ok was not "
                "given: tell Claude.") in res["text"]
        assert not (loadplan.run_dir(res["run"]) / "go.cmd").exists() and not alerts.alerts_file().exists()

        class Lost(FakeDeps):
            def preview(self, folder, *a, **k):
                (Path(folder) / "REPLACE_OK").unlink()
                return super().preview(folder, *a, **k)
        res, _ = check("070126", "--replace-ok", deps=Lost())
        assert res["exit"] == 2 and res["token"] is None, res["text"]
        assert "--replace-ok was given, but its marker file REPLACE_OK is not in the staged folder" in res["text"]
    # in the source: a marker is written in one place, each call under the option it belongs to
    tree = ast.parse((ROOT / "scripts" / "ops" / "load_check.py").read_text(encoding="utf-8"))
    under: list[tuple[str, str, str]] = []
    for fn in ast.walk(tree):
        if not isinstance(fn, ast.FunctionDef):
            continue
        for node in ast.walk(fn):
            if isinstance(node, ast.If):
                for call in (c for stmt in node.body for c in ast.walk(stmt) if isinstance(c, ast.Call)):
                    if isinstance(call.func, ast.Name) and call.func.id == "write_marker":
                        under.append((fn.name, ast.unparse(node.test), ast.unparse(call.args[1])))
    calls = [n for n in ast.walk(tree) if isinstance(n, ast.Call) and isinstance(n.func, ast.Name) and n.func.id == "write_marker"]
    assert len(calls) == 2 and sorted(under) == [("step_staging", "opt['partial_ok']", "PARTIAL_OK_MARKER"),
                                                 ("step_staging", "opt['replace_ok']", "REPLACE_OK_MARKER")], under
    names = [n.id for n in ast.walk(tree) if isinstance(n, ast.Name) and n.id in ("REPLACE_OK_MARKER", "PARTIAL_OK_MARKER")]
    assert sorted(names) == ["PARTIAL_OK_MARKER"] * 2 + ["REPLACE_OK_MARKER"] * 2, "used where it is written, and nowhere else"


@test("options: the price book and the profitability report load on the owner's word, under the right names")
def _():
    with data_root() as root:
        prices = PRICES + [("Charger two", "C2", None, None, "Van A", "1/5/2026", "12/31/2030", 2.0)]
        short = [("Cable one", 300, 100), ("Charger two", 200, 100), ("Holder three", 100, 50)]
        drop = make_drop(root, prices=prices, profit=short)
        before = snapshot_of(drop)
        earlier = [{"item_name": n, "net_amount_bhd": net, "cogs_bhd": cost, "report_date": "2026-01-05"}
                   for n, net, cost in (("Cable one", 250, 100), ("Charger two", 150, 100), ("Holder three", 50, 25))]
        res, deps = check("070126", deps=FakeDeps(FakeSession(profit_before=earlier)))   # no option: both stay out
        table = table_of(res["text"])
        assert table[shown("product_profitability")] == "HOLD BACK" and table[shown("selling_prices:MA_base")] == "NEEDS YOUR DECISION"
        assert res["ok"] and res["token"] and {"hold_back", "needs_decision", "price_scope"} <= set(res["flags"])
        staged = sorted(p.name for p in loadplan.staging_path(res["run"]).iterdir())
        assert not any("Profitability" in n or "PriceBook" in n for n in staged), staged
        assert "gap 16.7% (limit 10.0%)" in res["text"] and "production keeps the report as on Mon 05 Jan" in res["text"]
        assert "2 of the 3 items whose sales moved since Mon 05 Jan show no cost movement" in res["text"]
        assert "C2: scoped to Van A 2.000 from Mon 05 Jan (new row; BELOW the general price 2.500; already sold there at "\
               "this price: 2 invoices, 30 units)" in res["text"]
        assert "say \"load the price book\" (--pricebook load)" in res["text"]
        assert "Profitability: not loaded; it stays as on Mon 05 Jan (held back)." in res["text"]
        first = [r for r in alerts.open_items(include_baseline=True) if r["rule"] == "price_scope"]
        assert len(first) == 1 and first[0]["key"].startswith("price_scope|C2|h") and "Van A" not in first[0]["key"]
        assert first[0]["bhd"] == "15.000" and first[0]["level"] == "CHECK"

        res, deps = check("070126", "--pricebook", "load", "--with-profitability")
        run = res["run"]
        table = table_of(res["text"])
        assert table[shown("product_profitability")] == "LOAD" and table[shown("selling_prices:MA_base")] == "LOAD"
        staged = sorted(p.name for p in loadplan.staging_path(run).iterdir())
        book = f"MASellingPriceBook {run}.xlsx"
        assert book in staged and NAMES["product_profitability"] in staged and NAMES["selling_prices:MA_base"] not in staged
        assert loadplan.sha256_file(loadplan.staging_path(run) / book) == loadplan.sha256_file(drop / NAMES["selling_prices:MA_base"])
        from scripts import ingest as ing
        from scripts.load_supabase import focus_book_kind
        assert focus_book_kind(book) == "MA_base" and ing.classify(book) == "selling_prices:MA_base"
        assert "loaded on your word (--with-profitability)" in res["text"] and "loaded on your word (--pricebook load)" in res["text"]
        assert "Price book MA_base: 1 rows added, 0 changed, 0 switched off." in res["text"]
        assert "the price list this book would give" in res["text"] and res["token"]
        plan = loadplan.read_plan(run)
        assert plan["options"] == {"pricebook": "load", "with_profitability": True, "replace_ok": False,
                                   "partial_ok": False, "dev": False}
        assert [v["staged"] for v in plan["verdicts"] if v["kind"] == "selling_prices:MA_base"] == [book]

        res, _ = check("070126", "--pricebook", "skip")
        assert table_of(res["text"])[shown("selling_prices:MA_base")] == "LEAVE OUT" and "left out on your word" in res["text"]
        assert "needs_decision" not in res["flags"] and res["token"]
        assert snapshot_of(drop) == before


@test("a report whose cost of sales is far ABOVE the stock ledger's loads, and its line says so instead of 'agrees'")
def _():
    with data_root() as root:
        above = [("Cable one", 300, 600), ("Charger two", 200, 100), ("Holder three", 100, 50)]       # cost 750 against 300
        make_drop(root, profit=above)
        res, _ = check("070126")
        text, table = res["text"], table_of(res["text"])
        assert res["ok"] and res["token"] and table[shown("product_profitability")] == "LOAD", text
        assert "its cost of sales is 150.0% above the stock ledger's, not short of it: to check" in text
        assert "held back only when its cost is more than 10.0% short" in text
        assert "agrees with the stock ledger" not in text and "profit_above_ledger" in res["flags"]
        assert NAMES["product_profitability"] in [p.name for p in loadplan.staging_path(res["run"]).iterdir()]
    with data_root() as root:                 # as far above as the limit is wide: still inside it
        edge = [("Cable one", 300, 180), ("Charger two", 200, 100), ("Holder three", 100, 50)]        # cost 330 against 300
        make_drop(root, profit=edge)
        res, _ = check("070126")
        assert "its cost of sales agrees with the stock ledger (gap -10.0%, limit 10.0%)" in res["text"], res["text"]
        assert "profit_above_ledger" not in res["flags"] and res["token"]


@test("invoices that changed or went are listed and recorded; the load still goes ahead")
def _():
    with data_root() as root:
        make_drop(root)
        deps = FakeDeps()
        deps.actions = {"orders": {"updated": 1, "deleted": 1}, "order_lines": {"updated": 2, "deleted": 1}}
        deps.samples = {"orders": {"updated_sample": [{"invoice_no": "INV-2"}], "removed_sample": [{"invoice_no": "INV-5"}]},
                        "order_lines": {"updated_sample": [{"invoice_no": "INV-2", "line_no": 1}], "removed_sample": []}}
        deps.exceptions = [blocking("rows_removed", "orders", 1), blocking("days_shrink", "order_lines", 1),
                           blocking("rows_removed", "stock_movements", 9)]
        deps.sales = {"per_day_changed": [{"key": "2026-01-06", "db_rows": 0, "file_rows": 4},
                                          {"key": "2026-01-04", "db_rows": 5, "file_rows": 4, "file_gross": "40.500", "db_gross": "50.500"}],
                      "per_salesman_changed": [
                          {"key": "Van A", "file_rows": 12, "db_rows": 10, "file_gross": "125.250", "db_gross": "120.000"},
                          {"key": "Van B", "file_rows": 4, "db_rows": 5, "file_gross": "40.500", "db_gross": "50.500"},
                          {"key": "Shop 1", "file_rows": 1, "db_rows": 0, "file_gross": "0.000", "db_gross": "0.000"}]}
        res, _ = check("070126", deps=deps)
        text = res["text"]
        assert res["ok"] and res["token"] and "invoices_changed" in res["flags"], text
        assert "Invoices: 2 new, 1 changed, 1 removed. Lines: 2 new, 2 changed, 1 removed." in text
        # the rows the preview's samples leave out are counted as rows, apart from the invoices
        assert "to check): INV-2 (changed), INV-5 (removed); and 2 more changed rows not listed" in text, text
        # what he should see afterwards: the rows that go are taken off (10 loaded + 2 new - 1 removed)
        assert "Sales to Tue 06 Jan: 11 invoices, 11 lines in all." in text
        assert loadplan.read_plan(res["run"])["expected"]["invoices"] == 11
        assert "1 new day (Tue 06 Jan), 1 loaded day with different lines" in text and "file 4 lines / 40.500, loaded 5 / 50.500" in text
        assert "the loader reloads the file's whole date span (9 loaded rows are replaced)" in text and "expected" in text
        # listed from the first time they are found: this rule has no "already in the data" to count once
        items = {r["key"]: r for r in alerts.open_items()}
        assert sorted(items) == ["inv_changed|INV-2", "inv_changed|INV-5"], sorted(items)
        assert not any(r["baseline"] for r in items.values())
        assert items["inv_changed|INV-2"]["title"] == "Invoice changed in Focus after it was loaded: to check"
        assert items["inv_changed|INV-5"]["title"] == "Invoice no longer in Focus's export: to check"
        assert items["inv_changed|INV-5"]["detail"].startswith("INV-5: removed between the last load and")
        assert "- Action list (inv_changed): 2 added, 0 as before." in text, text
        # whose sales the load changes: every salesman of the preview's comparison, the largest change first
        assert ("- Sales by salesman, the file against what is loaded (lines / gross BHD): "
                "Van B: -1 / -10.000 (file 4 / 40.500, loaded 5 / 50.500); "
                "Van A: +2 / +5.250 (file 12 / 125.250, loaded 10 / 120.000); "
                "Shop 1: +1 / 0.000 (file 1 / 0.000, loaded 0 / 0.000).") in text, text
        assert "; 3 salesmen with new or changed lines." in text
        fmt.careful(text)
        # the next check finds another invoice: the two from before stay open (he has not looked at them yet)
        later = FakeDeps()
        later.actions = {"orders": {"updated": 1}}
        later.samples = {"orders": {"updated_sample": [{"invoice_no": "INV-9"}], "removed_sample": []}}
        res, _ = check("070126", deps=later)
        assert "- Action list (inv_changed): 1 added, 0 as before, 2 from earlier checks still open." in res["text"], res["text"]
        keys = sorted(r["key"] for r in alerts.open_items())
        assert keys == ["inv_changed|INV-2", "inv_changed|INV-5", "inv_changed|INV-9"], keys
        res, _ = check("070126", deps=FakeDeps(actions=later.actions, samples=later.samples))    # the same find again
        assert "Action list (inv_changed)" not in res["text"], "nothing new: no line"
        alerts.set_state("inv_changed|INV-2", "done", "checked with the salesman")
        res, _ = check("070126")                                    # a check that finds nothing clears nothing
        keys = sorted(r["key"] for r in alerts.open_items())
        assert keys == ["inv_changed|INV-5", "inv_changed|INV-9"], keys
        assert "Action list (inv_changed)" not in res["text"]
        states = {r["key"]: r["state"] for r in alerts.current().values() if r.get("rule") == "inv_changed"}
        assert states == {"inv_changed|INV-2": "done", "inv_changed|INV-5": "open", "inv_changed|INV-9": "open"}, states
        assert not any(r.get("state") == "cleared" for r in alerts.current().values()), "an event is never cleared"
    with data_root() as root:                 # an invoice that lost ONE LINE is still in the export: changed, not removed
        make_drop(root)
        deps = FakeDeps()
        deps.actions = {"orders": {"updated": 1}, "order_lines": {"deleted": 1}}
        deps.samples = {"orders": {"updated_sample": [{"invoice_no": "INV-2"}], "removed_sample": []},
                        "order_lines": {"updated_sample": [], "removed_sample": [{"invoice_no": "INV-2", "line_no": 4}]}}
        res, _ = check("070126", deps=deps)
        assert res["token"] and "to check): INV-2 (changed)" in res["text"] and "(removed)" not in res["text"], res["text"]
        assert "more changed row" not in res["text"]                # nothing beyond the samples: nothing is added
        (item,) = alerts.open_items()
        assert item["title"] == "Invoice changed in Focus after it was loaded: to check"
        assert item["detail"].startswith("INV-2: changed between the last load and") and item["baseline"] is False
    with data_root() as root:                 # more invoices than the list holds: the rest are counted as invoices
        make_drop(root)
        deps = FakeDeps()
        many = [{"invoice_no": f"INV-{i:02d}"} for i in range(1, 26)]
        deps.actions = {"orders": {"updated": 25}}
        deps.samples = {"orders": {"updated_sample": many, "removed_sample": []}}
        res, _ = check("070126", deps=deps)
        assert "INV-20 (changed) and 5 more invoices" in res["text"] and "INV-21" not in res["text"], res["text"]
        assert len(alerts.open_items()) == 25


@test("a loader guard predicted to fire stops the check and names the option; with the option the marker is written")
def _():
    from scripts.load_supabase import PARTIAL_OK_MARKER, REPLACE_OK_MARKER, marker_present
    held = [{"invoice_no": f"INV-{i}", "salesman": "Van A"} for i in range(40)]
    with data_root() as root:
        drop = make_drop(root)
        before = snapshot_of(drop)

        def deps_with_gap() -> FakeDeps:
            d = FakeDeps(FakeSession(orders=held))
            d.hook_rows = {"orders": held[8:]}
            d.actions = {"orders": {"deleted": 8}}
            d.samples = {"orders": {"updated_sample": [], "removed_sample": [{"invoice_no": "INV-0"}]}}
            return d
        res, _ = check("070126", deps=deps_with_gap())
        text = res["text"]
        assert res["exit"] == 2 and res["token"] is None and "stop" in res["flags"], text
        assert "1 item for the action list (inv_changed): not recorded while the check stops." in text
        assert not alerts.alerts_file().exists(), "a stopped check may be looking at a broken file: nothing is recorded"
        assert table_of(text)[shown("orders")] == "STOP" and "8 of the 40 invoices" in text and "--replace-ok" in text
        assert "No token was issued. Nothing was changed." in text and "GO:" not in text
        assert not (loadplan.run_dir(res["run"]) / "go.cmd").exists() and loadplan.read_plan(res["run"])["token"] is None
        assert not marker_present(loadplan.staging_path(res["run"]), REPLACE_OK_MARKER)
        res, _ = check("070126", "--replace-ok", deps=deps_with_gap())
        assert res["ok"] and res["token"] and table_of(res["text"])[shown("orders")] == "LOAD", res["text"]
        assert marker_present(loadplan.staging_path(res["run"]), REPLACE_OK_MARKER)
        assert REPLACE_OK_MARKER in loadplan.read_plan(res["run"])["token_parts"]["staging"]

        wide = FakeSession(warehouses_before=[{"w": MAIN}, {"w": "Van A"}])
        res, _ = check("070126", deps=FakeDeps(wide))
        assert res["exit"] == 2 and table_of(res["text"])[shown("stock_balance")] == "STOP" and "--partial-ok" in res["text"]
        assert "fewer warehouses than the snapshot before it" in res["text"]
        res, _ = check("070126", "--partial-ok", deps=FakeDeps(FakeSession(warehouses_before=[{"w": MAIN}, {"w": "Van A"}])))
        assert res["ok"] and res["token"] and marker_present(loadplan.staging_path(res["run"]), PARTIAL_OK_MARKER)
        assert not marker_present(loadplan.staging_path(res["run"]), REPLACE_OK_MARKER)

        seen = FakeDeps()
        seen.exceptions = [blocking("warehouse_set_shrinks", "stock_balance", items=["Van A"])]
        res, _ = check("070126", deps=seen)
        assert res["exit"] == 2 and "Van A missing" in res["text"] and "--partial-ok" in res["text"]
        hard = FakeDeps()
        hard.exceptions = [blocking("sales_join_below_80", "orders", message="the two sales reports share 50% of invoices")]
        res, _ = check("070126", deps=hard)
        assert res["exit"] == 2 and "sales_join_below_80" in res["text"] and "share 50%" in res["text"]
        # only the check that went ahead (with --replace-ok) recorded what it found; the stopped ones nothing
        assert [r["key"] for r in alerts.open_items(include_baseline=True)] == ["inv_changed|INV-0"]
        assert snapshot_of(drop) == before


@test("a stock ledger that lacks what production holds stops; a price book that would switch off too much stops")
def _():
    with data_root() as root:
        make_drop(root)
        extra = db_ledger() + [{"voucher": "SI:T-26-01-9", "item_name": "Lamp four C4", "d": "2026-01-04",
                                "vt": "Sales Invoice", "w": "Van A", "tw": "", "rq": Decimal(0), "iq": Decimal(2), "n": 1}]
        res, deps = check("070126", deps=FakeDeps(FakeSession(ledger=extra)))
        text = res["text"]
        assert res["exit"] == 2 and table_of(text)[shown("stock_movements")] == "STOP", text
        assert "1 item loaded for these dates is not in the file (Lamp four C4)" in text
        assert "export the full Stock ledger again" in text and res["token"] is None
        assert "STOP (1 item missing)" in deps.lines[7]

        changed = db_ledger()
        changed[1] = dict(changed[1], iq=Decimal(60))                    # production holds 60, the file 50
        res, _ = check("070126", deps=FakeDeps(FakeSession(ledger=changed)))
        assert res["ok"] and "vouchers_changed" in res["flags"], res["text"]
        assert "changed: SIO:T-26-01-1, Cable one C1: loaded as issued 60 on Sat 03 Jan at Main Store to Van A; the file "\
               "has issued 50 on Sat 03 Jan at Main Store to Van A" in res["text"]
        # on the list from the first time it is found (never a baseline item)
        assert [(r["key"], r["baseline"]) for r in alerts.open_items()] == [("voucher_changed|SIO:T-26-01-1|Cable one C1", False)]
    with data_root() as root:
        make_drop(root, prices=PRICES[:1])                               # three of the four live rows are gone
        res, _ = check("070126")
        assert res["ok"] and table_of(res["text"])[shown("selling_prices:MA_base")] == "NEEDS YOUR DECISION"
        assert res["text"].count("is no longer in the book") == 3
        res, _ = check("070126", "--pricebook", "load")
        assert res["exit"] == 2 and table_of(res["text"])[shown("selling_prices:MA_base")] == "STOP", res["text"]
        assert "switch off too much of the live book" in res["text"] and "--pricebook skip" in res["text"]
    with data_root() as root:                 # the loader's own count finds this book sound; the preview's guard still stops it
        make_drop(root, prices=PRICES + [("Charger two", "C2", None, None, "Van A", "1/5/2026", "12/31/2030", 2.0)])
        res, _ = check("070126", "--pricebook", "load")
        assert res["ok"] and res["token"] and table_of(res["text"])[shown("selling_prices:MA_base")] == "LOAD", res["text"]
        guarded = FakeDeps()
        guarded.exceptions = [blocking("snapshot_void_guard", "selling_prices",
                                       message="the file names 2 of the book's 9 live SKUs")]
        res, _ = check("070126", "--pricebook", "load", deps=guarded)
        assert res["exit"] == 2 and res["token"] is None, res["text"]
        assert table_of(res["text"])[shown("selling_prices:MA_base")] == "STOP"
        assert "switch off too much of the live book (the file names 2 of the book's 9 live SKUs)" in res["text"]
        assert "Loader guards predicted to fire: price book MA_base" in res["text"]


@test("a second hold in a row puts one item on the action list; a catalog switch is shown before the load")
def _():
    with data_root() as root:
        short = [("Cable one", 300, 100), ("Charger two", 200, 100), ("Holder three", 100, 50)]
        make_drop(root, profit=short)
        anchors = {"today": DAY, "sales_to": "2026-01-05", "stock_as_of": "2026-01-05", "ar_as_of": "2026-01-05",
                   "profit_as_of": "2026-01-01", "moves_to": "2026-01-05"}
        catalog = [{"item_code": "C1", "is_active": True, "hidden": False, "updated_by": lc.AUTO_SYNC},
                   {"item_code": "C2", "is_active": False, "hidden": True, "updated_by": "owner@example.invalid"},
                   {"item_code": "C9", "is_active": True, "hidden": False, "updated_by": lc.AUTO_SYNC}]
        res, deps = check("070126", deps=FakeDeps(FakeSession(anchors=anchors, catalog=catalog)))
        text = res["text"]
        assert res["ok"] and {"hold_back", "profit_held_twice", "catalog_switch"} <= set(res["flags"]), text
        assert "held back for the second load in a row" in text
        # an action from the first time it is found: alerts.open_items() is what the brief lists (no baseline items)
        item = [r for r in alerts.open_items() if r["rule"] == "profit_held"]
        assert len(item) == 1 and item[0]["key"] == f"profit_held|{DAY}" and item[0]["level"] == "ACT" and item[0]["bhd"] == "50.000"
        assert item[0]["baseline"] is False and "Action list (profit_held): 1 added" in text
        res2, _ = check("070126", "--with-profitability", deps=FakeDeps(FakeSession(anchors=anchors)))
        assert "Action list (profit_held): 0 added, 1 cleared" in res2["text"]       # a state: gone once the report loads
        assert [r for r in alerts.open_items(include_baseline=True) if r["rule"] == "profit_held"] == []
        assert "switch OFF 1 item: C9" in text and "switch ON 1 item: C2 (hidden, last set by hand)" in text
        assert "add 1 new item: C3" in text and "the live price list (no price book loads)" in text
        assert "ATTENTION" not in text


SHORT_COST = [("Cable one", 300, 100), ("Charger two", 200, 100), ("Holder three", 100, 50)]      # cost 250 of 300: a hold
BEFORE_LOAD = {"today": DAY, "sales_to": "2026-01-05", "stock_as_of": "2026-01-05", "ar_as_of": "2026-01-05",
               "profit_as_of": "2026-01-01", "moves_to": "2026-01-05"}


def held_items() -> list[tuple]:
    return [(r["key"], r["state"]) for r in alerts.current().values() if r.get("rule") == "profit_held"]


@test("a held report: checking the same drop again after its load neither clears nor adds the action")
def _():
    with data_root() as root:
        make_drop(root, profit=SHORT_COST)
        res, _ = check("070126", deps=FakeDeps(FakeSession(anchors=BEFORE_LOAD)))
        assert res["token"] and held_items() == [(f"profit_held|{DAY}", "open")], res["text"]
        # the load ran: sales, stock and receivables are now as on the drop's day; the report stayed out
        after = dict(BEFORE_LOAD, sales_to="2026-01-06", stock_as_of=DAY, ar_as_of=DAY, moves_to="2026-01-06")
        res, _ = check("070126", deps=FakeDeps(FakeSession(anchors=after)))
        assert res["ok"] and table_of(res["text"])[shown("product_profitability")] == "HOLD BACK", res["text"]
        assert "Action list (profit_held)" not in res["text"] and "second load in a row" not in res["text"]
        assert held_items() == [(f"profit_held|{DAY}", "open")], "the action is still open: the report is still held back"
        assert [r["key"] for r in alerts.open_items()] == [f"profit_held|{DAY}"]
        # ... until the report loads on his word: then the state is gone and the action clears
        res, _ = check("070126", "--with-profitability", deps=FakeDeps(FakeSession(anchors=after)))
        assert "Action list (profit_held): 0 added, 1 cleared" in res["text"] and held_items() == [(f"profit_held|{DAY}", "cleared")]
    with data_root() as root:                 # a FIRST hold, then its load, then the same drop again: still not "twice"
        make_drop(root, profit=SHORT_COST)
        first = dict(BEFORE_LOAD, profit_as_of="2026-01-05")              # the load before had its report
        res, _ = check("070126", deps=FakeDeps(FakeSession(anchors=first)))
        assert res["token"] and "profit_held_twice" not in res["flags"] and held_items() == []
        status.write_state({"last_load": {"run": res["run"]}})            # the writer published this very run
        res, _ = check("070126", deps=FakeDeps(FakeSession(anchors=dict(first, stock_as_of=DAY, ar_as_of=DAY))))
        assert "profit_held_twice" not in res["flags"] and held_items() == [], res["text"]
        # a partial drop (sales only) moves no snapshot date: the plan of that load still says it judged this report
        res, _ = check("070126", deps=FakeDeps(FakeSession(anchors=first)))
        assert "profit_held_twice" not in res["flags"] and held_items() == [], res["text"]


@test("two holds in a row are read from the plan of the load before, when there is one")
def _():
    def load_before(verdict: str | None) -> None:
        rows = [{"kind": "orders", "verdict": "load"}]
        dates = {"orders": "2026-01-05"}
        if verdict:
            rows.append({"kind": "product_profitability", "verdict": verdict})
            dates["product_profitability"] = "2026-01-05"
        loadplan.write_plan({"run": "050126_0900", "drop": "x", "token": "t", "verdicts": rows,
                             "as_on": {"orders": "2026-01-05"}, "as_on_all": dates})
        status.write_state({"last_load": {"run": "050126_0900"}})
    with data_root() as root:                 # the drop before brought no profitability report: nothing was held then
        make_drop(root, profit=SHORT_COST)
        load_before(None)
        res, _ = check("070126", deps=FakeDeps(FakeSession(anchors=BEFORE_LOAD)))      # the dates alone would say "twice"
        assert res["token"] and table_of(res["text"])[shown("product_profitability")] == "HOLD BACK"
        assert "profit_held_twice" not in res["flags"] and "second load in a row" not in res["text"] and held_items() == []
    with data_root() as root:                 # the load before loaded its report
        make_drop(root, profit=SHORT_COST)
        load_before("load")
        res, _ = check("070126", deps=FakeDeps(FakeSession(anchors=BEFORE_LOAD)))
        assert "profit_held_twice" not in res["flags"] and held_items() == []
    with data_root() as root:                 # the load before held its report back: this one is the second
        make_drop(root, profit=SHORT_COST)
        load_before("hold_back")
        level = dict(BEFORE_LOAD, profit_as_of="2026-01-05")                           # the dates alone would say "no"
        res, _ = check("070126", deps=FakeDeps(FakeSession(anchors=level)))
        assert "profit_held_twice" in res["flags"] and held_items() == [(f"profit_held|{DAY}", "open")], res["text"]
        assert [r["level"] for r in alerts.open_items()] == ["ACT"]


# ── older than what is loaded ─────────────────────────────────────────────────

@test("a drop older than what is loaded stops at once: no comparison, no token, nothing recorded")
def _():
    with data_root() as root:
        drop = make_drop(root)
        before = snapshot_of(drop)
        # production is three days ahead and holds two invoices and two vouchers entered after this export
        ahead = {"today": DAY, "sales_to": "2026-01-09", "stock_as_of": "2026-01-10", "ar_as_of": "2026-01-10",
                 "profit_as_of": "2026-01-10", "moves_to": "2026-01-10"}
        later = db_ledger(upto="2026-01-06") + [
            {"voucher": f"SI:T-26-01-{90 + i}", "item_name": "Cable one C1", "d": "2026-01-05", "vt": "Sales Invoice",
             "w": "Van A", "tw": "", "rq": Decimal(0), "iq": Decimal(2), "n": 1} for i in range(2)]

        def deps_ahead() -> FakeDeps:
            d = FakeDeps(FakeSession(anchors=ahead, ledger=later))
            d.actions = {"orders": {"deleted": 2}}
            d.samples = {"orders": {"updated_sample": [], "removed_sample": [{"invoice_no": "INV-90"}, {"invoice_no": "INV-91"}]}}
            return d
        for extra in ((), ("--dev",), ("--replace-ok",), ("--with-profitability",)):
            res, deps = check("070126", *extra, deps=deps_ahead())
            text = res["text"]
            assert res["exit"] == 2 and not res["ok"] and res["token"] is None, text
            assert res["flags"] == ["stop", "older_than_loaded"], res["flags"]
            assert "These reports are as on Wed 07 Jan; production already holds a load as on Sat 10 Jan." in text
            assert "An older drop is never loaded" in text and "GO:" not in text and "Token" not in text
            assert "No token was issued. Nothing was changed." in text
            assert deps.previewed == [] and len(deps.steps) == 3, "it ends before anything is compared"
            assert not loadplan.run_dir(res["run"]).exists(), "no plan and no go.cmd"
            assert f"(Left by this check: business_data/staging/{res['run']}.)" in text
            assert deps.the_session.closed
        assert not alerts.alerts_file().exists(), "an invoice Focus still holds is never recorded as gone"
        assert snapshot_of(drop) == before
        # the same day is not older: the drop that was loaded can be checked again
        same = dict(ahead, stock_as_of=DAY, ar_as_of=DAY, profit_as_of=DAY, sales_to="2026-01-06", moves_to="2026-01-06")
        res, _ = check("070126", deps=FakeDeps(FakeSession(anchors=same)))
        assert res["ok"] and res["token"] and "older_than_loaded" not in res["flags"], res["text"]
    with data_root() as root:                 # a folder with a price book alone carries no as-on date to be older by
        make_drop(root, leave_out=tuple(k for k in NAMES if k != "selling_prices:MA_base"))
        # (a price book is dated by the day it was exported, which is today for this file: whatever today is, it is
        # before the snapshot below, and still the rule does not read it as an older drop)
        res, _ = check("070126", deps=FakeDeps(FakeSession(anchors=dict(ahead, stock_as_of="2999-01-01"))))
        assert "An older drop" not in res["text"] and "older_than_loaded" not in res["flags"], res["text"]
        assert "nothing in this drop is ready to load" in res["text"]


@test("one older report in a current folder stops on its own line; the rest is judged as usual")
def _():
    loaded = {"today": DAY, "sales_to": "2026-01-06", "stock_as_of": DAY, "ar_as_of": DAY, "profit_as_of": DAY,
              "moves_to": "2026-01-06"}                                    # this folder's day is loaded already
    with data_root() as root:
        drop = make_drop(root)
        plain_report(drop / NAMES["order_lines"], "Sales day book", ["Date", "Voucher"], as_on="06/01/2026")
        plain_report(drop / NAMES["receivables"], "Customer summary ageing", ["Account", "Balance Amount"], as_on="06/01/2026")
        profit_file(drop / NAMES["product_profitability"], as_on="06/01/2026")
        before = snapshot_of(drop)
        res, deps = check("070126", deps=FakeDeps(FakeSession(anchors=loaded)))
        text, table = res["text"], table_of(res["text"])
        assert res["exit"] == 2 and res["token"] is None and {"stop", "older_than_loaded", "mixed_dates"} <= set(res["flags"]), text
        assert table[shown("order_lines")] == table[shown("receivables")] == table[shown("product_profitability")] == "STOP"
        assert table[shown("orders")] == table[shown("stock_movements")] == "LOAD"
        assert table[shown("stock_balance")] == "LOAD (main-store block only)"
        assert "it is as on Tue 06 Jan; production already holds a load as on Wed 07 Jan." in text        # the day book
        assert text.count("it is as on Tue 06 Jan; production already holds this report as on Wed 07 Jan.") >= 2
        assert "An older report is never loaded" in text and "export it again from Focus" in text
        staged = sorted(p.name for p in loadplan.staging_path(res["run"]).iterdir())
        assert staged == sorted([NAMES["orders"], NAMES["stock_movements"], "Main_Store_only_" + NAMES["stock_balance"]]), staged
        assert deps.previewed == [staged], "the preview never sees an older report"
        assert "Invoices: 2 new, 0 changed, 0 removed. Lines: the sales day book stops (see its line)." in text
        assert deps.lines[4] == "[5/9] Profitability: STOP (older than what is loaded)."
        assert "Receivables: not loaded; they stay as on Wed 07 Jan." in text
        assert not alerts.alerts_file().exists() and not (loadplan.run_dir(res["run"]) / "go.cmd").exists()
        assert loadplan.read_plan(res["run"])["token"] is None and snapshot_of(drop) == before
    with data_root() as root:                 # the stock balance alone is older than the snapshot production holds
        drop = make_drop(root)
        stock_balance_file(drop / NAMES["stock_balance"], as_on="05/01/2026")
        newer = dict(loaded, stock_as_of="2026-01-06", ar_as_of="2026-01-05", profit_as_of="2026-01-05")
        res, deps = check("070126", deps=FakeDeps(FakeSession(anchors=newer)))
        table = table_of(res["text"])
        assert res["exit"] == 2 and table[shown("stock_balance")] == "STOP" and table[shown("receivables")] == "LOAD"
        assert "it is as on Mon 05 Jan; production already holds this report as on Tue 06 Jan." in res["text"]
        assert deps.lines[3] == "[4/9] Stock balance: STOP (older than what is loaded)."
        assert "Stock: not loaded; it stays as on Tue 06 Jan." in res["text"]
        assert not any("Stock_balance" in p.name for p in loadplan.staging_path(res["run"]).iterdir())
        # the same file against a snapshot of its own day is not older
        res, _ = check("070126", deps=FakeDeps(FakeSession(anchors=dict(newer, stock_as_of="2026-01-05"))))
        assert res["ok"] and table_of(res["text"])[shown("stock_balance")] == "LOAD (main-store block only)", res["text"]
    with data_root() as root:                 # both sales reports and the stock ledger carry an older title
        drop = make_drop(root)
        plain_report(drop / NAMES["order_lines"], "Sales day book", ["Date", "Voucher"], as_on="06/01/2026")
        plain_report(drop / NAMES["orders"], "Summary sales register", ["Date", "Invoice"], as_on="06/01/2026")
        ledger_file(drop / NAMES["stock_movements"], as_on="06/01/2026")
        res, deps = check("070126", deps=FakeDeps(FakeSession(anchors=loaded)))
        text, table = res["text"], table_of(res["text"])
        assert res["exit"] == 2 and res["token"] is None, text
        assert table[shown("order_lines")] == table[shown("orders")] == table[shown("stock_movements")] == "STOP"
        assert text.count("it is as on Tue 06 Jan; production already holds a load as on Wed 07 Jan.") >= 3
        assert "Sales: not loaded; they stay as loaded to Tue 06 Jan." in text and "- Invoices:" not in text
        assert "Stock movements: not loaded; they stay as loaded to Tue 06 Jan." in text
        assert deps.lines[6] == "[7/9] Preview against the database: the sales register stops (see its line)."
        assert deps.lines[7] == "[8/9] Stock ledger: STOP (older than what is loaded)."
        # the profitability report cannot be weighed against a ledger that is not used, and says why
        assert table[shown("product_profitability")] == "HOLD BACK"
        assert "the stock ledger stops (see its line), so its cost of sales cannot be checked" in text
        staged = sorted(p.name for p in loadplan.staging_path(res["run"]).iterdir())
        assert staged == sorted([NAMES["receivables"], "Main_Store_only_" + NAMES["stock_balance"]]), staged
    with data_root() as root:                 # a snapshot is measured against its OWN kind, not against the newest of all
        drop = make_drop(root)
        plain_report(drop / NAMES["receivables"], "Customer summary ageing", ["Account", "Balance Amount"], as_on="06/01/2026")
        partial = dict(loaded, ar_as_of="2026-01-05")                     # the last load brought no receivables
        res, _ = check("070126", deps=FakeDeps(FakeSession(anchors=partial)))
        assert res["ok"] and res["token"] and table_of(res["text"])[shown("receivables")] == "LOAD", res["text"]
        assert "older_than_loaded" not in res["flags"]


@test("a report whose rows end before production's stops: the stock ledger, the day book, the register")
def _():
    with data_root() as root:
        make_drop(root)
        ends = FakeSession(anchors=dict(BEFORE_LOAD, profit_as_of="2026-01-05", moves_to="2026-01-09"))
        res, deps = check("070126", deps=FakeDeps(ends))
        text, table = res["text"], table_of(res["text"])
        assert res["exit"] == 2 and res["token"] is None and "older_than_loaded" in res["flags"], text
        assert table[shown("stock_movements")] == "STOP" and table[shown("orders")] == "LOAD"
        assert "its movements end on Tue 06 Jan; production already holds movements to Fri 09 Jan." in text
        assert "A report that ends earlier is never loaded" in text
        assert deps.lines[7] == "[8/9] Stock ledger: STOP (it ends before what is loaded)."
        assert NAMES["stock_movements"] not in deps.previewed[0], "the preview never sees it"
        # the profitability report cannot be weighed against a ledger that is not used, and says why
        assert table[shown("product_profitability")] == "HOLD BACK"
        assert "the stock ledger stops (see its line), so its cost of sales cannot be checked" in text
        # a ledger that ends on production's last day is fine
        res, _ = check("070126", deps=FakeDeps(FakeSession(anchors=dict(BEFORE_LOAD, profit_as_of="2026-01-05", moves_to="2026-01-06"))))
        assert res["ok"] and table_of(res["text"])[shown("stock_movements")] == "LOAD", res["text"]
    with data_root() as root:
        make_drop(root)
        sales = dict(BEFORE_LOAD, profit_as_of="2026-01-05", sales_to="2026-01-09")
        res, _ = check("070126", deps=FakeDeps(FakeSession(anchors=sales)))
        text, table = res["text"], table_of(res["text"])
        assert res["exit"] == 2 and table[shown("order_lines")] == "STOP" and table[shown("orders")] == "LOAD", text
        assert "its lines end on Tue 06 Jan; production already holds lines to Fri 09 Jan." in text
        res, _ = check("070126", deps=FakeDeps(FakeSession(anchors=dict(sales, sales_to="2026-01-06", orders_to="2026-01-08"))))
        table = table_of(res["text"])
        assert res["exit"] == 2 and table[shown("orders")] == "STOP" and table[shown("order_lines")] == "LOAD", res["text"]
        assert "its invoices end on Tue 06 Jan; production already holds invoices to Thu 08 Jan." in res["text"]
        res, _ = check("070126", deps=FakeDeps(FakeSession(anchors=dict(sales, sales_to="2026-01-06", orders_to="2026-01-06"))))
        assert res["ok"] and res["token"], res["text"]


@test("a database that cannot be reached, a preview that fails and a check that never comes back each end cleanly")
def _():
    with data_root() as root:
        drop = make_drop(root)
        before = snapshot_of(drop)

        class NoDb(FakeDeps):
            def session(self):
                raise RuntimeError("connection to host db.secret.example failed for user postgres.secretref")
        res, _ = check("070126", deps=NoDb())
        assert res["exit"] == 2 and "could not be reached (RuntimeError)" in res["text"] and "secret" not in res["text"]
        # what a stopped check leaves behind is named, and only what is really there
        assert res["text"].splitlines()[-1] == f"(Left by this check: business_data/staging/{res['run']}.)"

        class BadPreview(FakeDeps):
            def preview(self, *a, **k):
                raise ValueError("header row not found")
        res, deps = check("070126", deps=BadPreview())
        assert res["exit"] == 2 and "The preview against the database failed (ValueError: header row not found)" in res["text"]
        assert deps.the_session.closed and res["token"] is None
        import httpx
        assert lc.plain_error(httpx.ConnectError("could not reach db.secret.example")) == "ConnectError"
        assert lc.plain_error(KeyError()) == "KeyError"

        class BadRead(FakeSession):
            def q(self, sql, params=None):
                if "catalog_items" in sql:
                    raise RuntimeError("server closed the connection to db.secret.example")
                return super().q(sql, params)
        res, _ = check("070126", deps=FakeDeps(BadRead()))
        assert res["exit"] == 2 and "A read of the database failed (RuntimeError)" in res["text"] and "secret" not in res["text"]
        assert res["text"].splitlines()[-1] == (f"(Left by this check: business_data/staging/{res['run']} and "
                                                f"business_data/ops/loads/{res['run']}.)")

        release = threading.Event()

        class Stuck(FakeDeps):
            deadline_s = 0.3

            def live_commit(self):
                release.wait(20)
                raise RuntimeError("late")
        res, _ = check("070126", deps=Stuck())
        release.set()
        assert res["exit"] == 2 and res["flags"] == ["stop", "timeout"] and "did not finish" in res["text"]
        assert "step 1 of 9" in res["text"] and res["token"] is None
        assert snapshot_of(drop) == before


@test("what the check read must still be there when it ends: a report saved again, a database another load moved")
def _():
    with data_root() as root:                 # a report is saved again while the check runs (after it was staged)
        drop = make_drop(root)

        class Resaved(FakeDeps):
            def preview(self, *a, **k):
                out = super().preview(*a, **k)
                plain_report(drop / NAMES["receivables"], "Customer summary ageing", ["Account", "Balance Amount", "More"])
                return out
        res, _ = check("070126", deps=Resaved())
        text = res["text"]
        assert res["exit"] == 2 and res["token"] is None and {"stop", "moved_meanwhile"} <= set(res["flags"]), text
        assert NAMES["receivables"] + " changed in the folder while the check ran" in text and "Run the check again." in text
        assert "No token was issued." in text and "GO:" not in text
        assert not (loadplan.run_dir(res["run"]) / "go.cmd").exists() and loadplan.read_plan(res["run"])["token"] is None
        assert not alerts.alerts_file().exists()
        res, _ = check("070126")              # nothing moves this time
        assert res["ok"] and res["token"] and "moved_meanwhile" not in res["flags"], res["text"]
    with data_root() as root:                 # a report is judged, THEN copied: a file saved in between would load unjudged
        drop = make_drop(root)
        real_copy = lc.copy_checked

        def copy_after_change(src, dst):
            if "Profitability" in src.name:
                profit_file(src, [("Cable one", 300, 10), ("Charger two", 200, 10), ("Holder three", 100, 10)])
            return real_copy(src, dst)
        lc.copy_checked = copy_after_change
        try:
            res, _ = check("070126")
        finally:
            lc.copy_checked = real_copy
        assert table_of(res["text"])[shown("product_profitability")] == "LOAD"       # the verdict of the file as it was
        assert res["exit"] == 2 and res["token"] is None, res["text"]
        assert NAMES["product_profitability"] + " changed in the folder while the check ran" in res["text"]
    with data_root() as root:                 # another load commits while the check runs
        make_drop(root)

        class Moving(FakeSession):
            moved = False

            def q(self, sql, params=None):
                rows = super().q(sql, params)
                if self.moved and " ".join(sql.split()) == "select count(*) as n, coalesce(max(id), 0) as max_id from orders":
                    return [{"n": 12, "max_id": 120}]
                return rows

        class Meanwhile(FakeDeps):
            def preview(self, *a, **k):
                out = super().preview(*a, **k)
                self.the_session.moved = True
                return out
        res, _ = check("070126", deps=Meanwhile(Moving()))
        text = res["text"]
        assert res["exit"] == 2 and res["token"] is None and "moved_meanwhile" in res["flags"], text
        assert "the database changed while the check ran (invoices): another load or upload was at work. Run the check again." in text
        assert not (loadplan.run_dir(res["run"]) / "go.cmd").exists()
        assert set(lc.DB_WORDS) == set(loadplan.FINGERPRINT_TABLES) | {"ingest_runs", "ingest_batches_open"}
        res, _ = check("070126", deps=FakeDeps(Moving()))                 # a database that stays as it is: a token
        assert res["ok"] and res["token"], res["text"]

        class NoFingerprint(FakeSession):
            def q(self, sql, params=None):
                if "from ingest_runs" in sql:
                    raise RuntimeError("connection reset by db.secret.example")
                return super().q(sql, params)
        res, deps = check("070126", deps=FakeDeps(NoFingerprint()))
        assert res["exit"] == 2 and "A read of the database failed (RuntimeError)" in res["text"] and "secret" not in res["text"]
        assert deps.previewed == [] and deps.the_session.closed


@test("a scoped price whose sales could not be looked up is listed and recorded without a made-up count")
def _():
    with data_root() as root:
        make_drop(root, prices=PRICES + [("Charger two", "C2", None, None, "Van A", "1/5/2026", "12/31/2030", 2.0)])

        class NoSales(FakeDeps):
            def read_params(self, sql, params):
                if "from v_sales" in sql:
                    raise RuntimeError("statement timeout")
                return super().read_params(sql, params)
        res, _ = check("070126", deps=NoSales())
        text = res["text"]
        assert res["token"] and "C2: scoped to Van A 2.000 from Mon 05 Jan (new row; BELOW the general price 2.500)" in text
        assert "already sold there" not in text
        assert ("    What was already sold at the scoped price could not be read for 1 row (the database did not answer): "
                "those lines say nothing about sales.") in text, text
        (item,) = [r for r in alerts.open_items(include_baseline=True) if r["rule"] == "price_scope"]
        assert item["bhd"] is None and item["detail"].endswith("what was already sold there at this price was not looked up")
        assert "0 units" not in item["detail"]
        # this rule's first run is the baseline: counted once, and the line says so instead of "added"
        assert item["baseline"] is True and alerts.open_items() == []
        assert ("- Action list (price_scope): 0 added, 0 cleared, 0 as before; 1 counted once as already there (the first "
                "run of this rule: they are not listed day by day).") in text, text
        assert alerts.baseline_summary() == [{"rule": "price_scope", "count": 1, "bhd": Decimal(0)}]
        res, _ = check("070126")              # once the lookup answers, the same item carries the count and the amount
        (item,) = [r for r in alerts.open_items(include_baseline=True) if r["rule"] == "price_scope"]
        assert item["bhd"] == "15.000" and "30 units already sold there at this price" in item["detail"]
        assert "could not be read for" not in res["text"] and "was looked up for" not in res["text"]
    with data_root() as root:                 # more scoped rows than are looked up: the rows for the action list go first
        level = [("Cable one", "C1", None, None, f"Van {i:02d}", "1/5/2026", None, 1.5) for i in range(40)]   # = the general price
        lower = [("Holder three", "C3", None, None, f"Van {i:02d}", "1/5/2026", None, 0.5) for i in range(5)]   # below it
        make_drop(root, prices=PRICES + level + lower)
        asked: list[str] = []

        class Counting(FakeDeps):
            def read_params(self, sql, params):
                if "from v_sales" in sql:
                    asked.append(params[0])
                return super().read_params(sql, params)
        res, _ = check("070126", deps=Counting())
        assert res["token"] and len(asked) == lc.SCOPE_LOOKUPS_MAX == 40, len(asked)
        assert asked[:5] == ["C3"] * 5 and asked.count("C1") == 35, "the five below the general price are looked up first"
        assert ("    What was already sold at a scoped price was looked up for 40 of the 45 scoped rows (those below the "
                "general price first); for the other 5 it was not looked up.") in res["text"], res["text"]
        items = [r for r in alerts.open_items(include_baseline=True) if r["rule"] == "price_scope"]
        assert len(items) == 5 and all(r["bhd"] == "7.500" and "30 units already sold there" in r["detail"] for r in items)


@test("a changed header row stops that file, and only that file")
def _():
    with data_root() as root:
        drop = make_drop(root)
        plain_report(drop / NAMES["product_profitability"], "Product profitability report", ["Particulars", "Gross", "Cost"])
        plain_report(drop / NAMES["stock_movements"], "Stock ledger", ["Date", "Voucher"])
        plain_report(drop / NAMES["stock_balance"], "Stock balance by warehouse", ["Particulars", "Quantity"])
        wb = Workbook()
        put(wb.active, 1, ["Item Name", "Item Code", "Rate"])
        wb.save(drop / NAMES["selling_prices:MA_base"])
        before = snapshot_of(drop)
        res, deps = check("070126")
        table, text = table_of(res["text"]), res["text"]
        assert res["exit"] == 2 and res["token"] is None and len(deps.lines) == 9, deps.lines
        for kind in ("product_profitability", "stock_movements", "stock_balance", "selling_prices:MA_base"):
            assert table[shown(kind)] == "STOP", kind
        assert table[shown("order_lines")] == table[shown("orders")] == table[shown("receivables")] == "LOAD"
        assert text.count("its header row is not the one the loader reads") >= 8        # each in the table and in the reasons
        staged = sorted(p.name for p in loadplan.staging_path(res["run"]).iterdir())
        assert staged == sorted([NAMES["order_lines"], NAMES["orders"], NAMES["receivables"]]), staged
        assert "Stock ledger: STOP (its header row changed)." in deps.lines[7]
        assert snapshot_of(drop) == before

        ledger_file(drop / NAMES["stock_movements"], {})                    # the right header, no movement at all
        res, deps = check("070126")
        assert "the file holds no dated movement" in res["text"] and table_of(res["text"])[shown("stock_movements")] == "STOP"


@test("the checks also say: mixed as-on dates, items gone since the last snapshot, stock by warehouse, a day loaded twice")
def _():
    with data_root() as root:
        drop = make_drop(root, leave_out=("ledger_entries",))
        plain_report(drop / NAMES["receivables"], "Customer summary ageing", ["Account", "Balance Amount"], as_on="06/01/2026")
        for i in range(12):
            (drop / f"note{i:02d}.txt").write_text("x", encoding="utf-8")
        deps = FakeDeps()
        deps.snapshot = deps.snapshot + [{"item_name": "Lamp four C4", "net_qty": 3}]
        deps.loaded_rows = {"ar_ageing": 7}
        deps.stock_view = {"db_latest_as_of": "2026-01-05", "per_warehouse": [
            {"warehouse": MAIN, "file_items": 3, "db_items": 3, "file_qty": "160.000", "db_qty": "150.000",
             "file_value": "265.000", "db_value": "250.000", "delta_qty": "10.000", "delta_value": "15.000"}]}
        res, _ = check("070126", deps=deps)
        text = res["text"]
        assert res["ok"] and {"mixed_dates"} <= set(res["flags"]), text
        assert "ATTENTION: the reports carry different as-on dates (Tue 06 Jan, Wed 07 Jan): are they one export?" in text
        assert "1 item with stock in the last loaded snapshot is gone from this file (sold out or moved): Lamp four C4." in text
        assert "Stock against the last loaded snapshot (as on Mon 05 Jan), Main Store: file 3 items / 160 units / BHD 265.000; "\
               "loaded 3 / 150 / 250.000; difference 10 units / BHD 15.000." in text
        assert "as on Tue 06 Jan (that day is already loaded, with the same total: 2 new, 0 rewritten, 0 removed)" in text
        table = table_of(text)
        assert table["(+3 more files and folders)"] == "NOT LOADED" and table["note00.txt"] == "NOT LOADED"
        assert len([k for k in table if k.startswith("note")]) == 9


@test("a drop with nothing to load stops; a folder given by path is dated by its reports; an .xls file is refused")
def _():
    with data_root() as root:
        make_drop(root, leave_out=tuple(k for k in NAMES if k != "selling_prices:MA_base"))
        res, deps = check("070126")
        text = res["text"]
        assert res["exit"] == 2 and res["token"] is None and "nothing in this drop is ready to load" in text
        table = table_of(text)
        assert table[shown("selling_prices:MA_base")] == "LEAVE OUT" and list(table.values()).count("NOT LOADED") == 6
        assert "Sales: no sales report in this drop; they stay as loaded to Mon 05 Jan." in text
        assert list(loadplan.staging_path(res["run"]).iterdir()) == []
    with data_root() as root:
        folder = make_drop(root, name="incoming")
        res, _ = check(str(folder))
        assert res["ok"] and res["run"] == "070126_0930", res["text"]         # named by the reports' as-on date
        (folder / NAMES["orders"]).rename(folder / "Summary_sales_register100002_Test_Co.xls")
        res, _ = check(str(folder))
        assert res["exit"] == 2 and "is not an .xlsx export" in res["text"]
        (folder / "Summary_sales_register100002_Test_Co.xls").rename(folder / NAMES["orders"])
        (folder / NAMES["selling_prices:MA_base"]).write_bytes(b"not a workbook")
        res, _ = check(str(folder))
        assert res["exit"] == 2 and "MASellingPriceBook (7).xlsx cannot be opened as a workbook" in res["text"]
    with data_root():
        res, _ = check()
        assert res["exit"] == 2 and "No dated drop folder" in res["text"]
    with data_root() as root:                       # the Ledger export is never loaded, so it never stops a load
        from scripts import ingest as ing
        drop = make_drop(root)
        (drop / NAMES["ledger_entries"]).write_bytes(b"not a workbook at all")
        (drop / "Brand_new_report100008_Test_Co.xlsx").write_bytes(b"x")
        real_classify = ing.classify
        ing.classify = lambda name: "brand_new_kind" if name.startswith("Brand_new") else real_classify(name)
        try:
            res, _ = check("070126")
        finally:
            ing.classify = real_classify
        table = table_of(res["text"])
        assert res["ok"] and res["token"] and table[shown("ledger_entries")] == "NOT LOADED", res["text"]
        assert "never loaded by this command (its as-on date could not be read)" in res["text"]
        assert table["Brand_new_report100008_Test_Co.xlsx"] == "NOT LOADED" and "brand_new_kind" in res["text"]
        plan = loadplan.read_plan(res["run"])
        assert res["data_date"] == DAY and plan["as_on_all"]["ledger_entries"] == "" and "ledger_entries" not in plan["as_on"]


@test("plan.json dates only what loads: the Ledger export and a held report never date the published set")
def _():
    with data_root() as root:
        drop = make_drop(root, profit=SHORT_COST)
        plain_report(drop / NAMES["ledger_entries"], "Ledger", ["Date", "Voucher"], as_on="31/01/2026")   # a month-end export
        profit_file(drop / NAMES["product_profitability"], SHORT_COST, as_on="08/01/2026")              # held back
        res, _ = check("070126", deps=FakeDeps(FakeSession(anchors=dict(BEFORE_LOAD, profit_as_of="2026-01-05"))))
        plan = loadplan.read_plan(res["run"])
        table = table_of(res["text"])
        assert res["token"] and table[shown("ledger_entries")] == "NOT LOADED" and table[shown("product_profitability")] == "HOLD BACK"
        assert "never loaded by this command (as on Sat 31 Jan)" in res["text"]
        assert set(plan["as_on"]) == {"order_lines", "orders", "stock_movements", "stock_balance", "receivables"}, plan["as_on"]
        assert set(plan["as_on"].values()) == {DAY}, "the writer names the published folder after the newest of these"
        assert plan["as_on_all"]["ledger_entries"] == "2026-01-31" and plan["as_on_all"]["product_profitability"] == "2026-01-08"
        assert set(plan["as_on_all"]) == set(NAMES)
        # a report loaded on his word is dated; a price book that loads is listed with its export day
        res, _ = check("070126", "--with-profitability", "--pricebook", "load",
                       deps=FakeDeps(FakeSession(anchors=dict(BEFORE_LOAD, profit_as_of="2026-01-05"))))
        plan = loadplan.read_plan(res["run"])
        assert plan["as_on"]["product_profitability"] == "2026-01-08" and "selling_prices:MA_base" in plan["as_on"]
        assert "ledger_entries" not in plan["as_on"]


@test("long lists are cut after twenty with a count; a scoped price with no general price says so")
def _():
    with data_root() as root:
        vans = [("Zed nine", "Z9", None, None, f"Van {i:02d}", "1/5/2026", None, 1.0) for i in range(25)]
        make_drop(root, prices=PRICES + vans)
        gone = [{"voucher": f"SI:T-26-01-{50 + i}", "item_name": "Cable one C1", "d": "2026-01-04", "vt": "Sales Invoice",
                 "w": "Van A", "tw": "", "rq": Decimal(0), "iq": Decimal(1), "n": 1} for i in range(25)]
        res, _ = check("070126", deps=FakeDeps(FakeSession(ledger=db_ledger() + gone)))
        text = res["text"]
        assert res["ok"] and "25 rows differ from the live book (25 added)" in text and "    and 5 more rows" in text
        assert text.count("no general price in force") == 20 and "price_scope" not in res["flags"]
        assert "pairs): changed 0, removed 25, entered with an earlier date 0" in text and text.count("    removed: SI:T-26-01-") == 20
        assert "    and 5 more (each one is for the action list, to check)" in text
        assert len([r for r in alerts.open_items(include_baseline=True) if r["rule"] == "voucher_changed"]) == 25
        fmt.careful(text)


@test("--partial-ok lets a narrower stock balance through; an error that is not the check's own is raised, not hidden")
def _():
    with data_root() as root:
        make_drop(root)
        seen = FakeDeps()
        seen.exceptions = [blocking("warehouse_set_shrinks", "stock_balance", items=["Van A"])]
        res, _ = check("070126", "--partial-ok", deps=seen)
        assert res["ok"] and res["token"], res["text"]
        assert "fewer warehouses than the last snapshot (Van A missing): allowed by --partial-ok." in res["text"]

        class Broken(FakeDeps):
            deadline_s = 30

            def addresses(self):
                raise KeyError("a bug, not a stop")
        raises(KeyError, lambda: check("070126", deps=Broken()))

        class Writes(FakeSession):
            def q(self, sql, params=None):
                if "catalog_items" in sql:
                    raise db.ReadOnlyViolation("refused")
                return super().q(sql, params)
        raises(db.ReadOnlyViolation, lambda: check("070126", deps=FakeDeps(Writes())))


# ── status ────────────────────────────────────────────────────────────────────

def status_of(*argv: str, deps: FakeDeps | None = None) -> dict:
    parser = argparse.ArgumentParser()
    parser.add_argument("--json", action="store_true")
    lc.add_arguments(parser)
    return lc.run_status(parser.parse_args(["status", *argv]), deps or FakeDeps())


def write_lock(run: str, pid: int, beat: datetime, step: str) -> None:
    paths.write_text(paths.lock_file(), json.dumps({"run": run, "pid": pid, "step": step, "started": beat.isoformat(),
                                                    "heartbeat": beat.isoformat()}))


@test("status: NOT STARTED, RUNNING, FINISHED, FAILED and INTERRUPTED, with the journal lines since the last call")
def _():
    run, now = "070126_0930", datetime.now(timezone.utc)
    with data_root():
        res = status_of()
        assert res["text"].splitlines()[-1] == "NOT STARTED" and res["exit"] == 0 and res["run"] is None
        loadplan.write_plan({"run": run, "drop": "x", "token": "t", "verdicts": []})
        res = status_of()
        assert res["text"].splitlines() == [f"Load {run}", "  (no new journal line)", "NOT STARTED"] and res["state"] == "NOT STARTED"
        loadplan.journal_append(run, "started", pid=os.getpid())
        loadplan.journal_append(run, "step_start", step="backup")
        write_lock(run, os.getpid(), now, "backup")
        res = status_of("--run", run)
        lines = res["text"].splitlines()
        assert lines[-1] == "RUNNING at backup" and res["exit"] == 0 and len(lines) == 4, lines
        assert lines[1].endswith("  started") and lines[2].endswith("backup: started")      # the pid stays in the file
        assert status_of().get("text").splitlines()[1] == "  (no new journal line)"     # nothing new since the last call
        loadplan.journal_append(run, "step_ok", step="backup", text="Backup packed.", seconds=12, archive="x.7z")
        loadplan.journal_append(run, "step_failed", step="reconcile", stops=False, text="Reconcile did not pass.",
                                why="2 accounts differ.", next="Look at the Receivables page.")
        res = status_of()
        shown_lines = res["text"].splitlines()
        assert shown_lines[1].endswith("backup: done - Backup packed.") and len(shown_lines) == 4
        assert shown_lines[2].endswith("reconcile: needs a look (the load went on) - Reconcile did not pass. 2 accounts "
                                       "differ. Look at the Receivables page.")
        assert "x.7z" not in res["text"] and res["result"] is None
        write_lock(run, 2 ** 22 + 12345, now, "refresh")                                # the worker is gone
        res = status_of()
        assert res["text"].splitlines()[-1] == "INTERRUPTED at refresh" and res["exit"] == 2 and res["flags"] == ["interrupted"]
        write_lock(run, os.getpid(), now - timedelta(seconds=status.HEARTBEAT_STALE_S + 60), "refresh")
        assert status_of()["state"] == "INTERRUPTED"                                    # alive, but silent too long
        loadplan.journal_append(run, "step_failed", step="refresh", stops=True, text="The totals do not match.")
        loadplan.journal_append(run, "failed", step="refresh", result="failed", written=True, why="Verify failed.")
        paths.lock_file().unlink()
        res = status_of()
        assert res["text"].splitlines()[-1] == "FAILED at refresh" and res["exit"] == 2
        assert "refresh: FAILED - The totals do not match." in res["text"] and "refresh: FAILED - Verify failed." in res["text"]
        assert res["result"] == "failed" and res["flags"] == ["failed"]
        other = "070126_0945"
        loadplan.journal_append(other, "started")
        loadplan.journal_append(other, "step_start", step="publish")
        loadplan.journal_append(other, "finished", result="ok_needs_look", text="Loaded and published.")
        res = status_of("--run", other)
        assert res["text"].splitlines()[-1] == "FINISHED" and res["exit"] == 0 and len(res["text"].splitlines()) == 5
        assert res["text"].splitlines()[-2].endswith("FINISHED, and something needs a look - Loaded and published.")
        assert res["result"] == "ok_needs_look" and res["flags"] == ["finished", "ok_needs_look"]
        again = "070126_0950"
        loadplan.journal_append(again, "failed", step="token", result="check_again", written=False,
                                why="changed: the drop folder: x.xlsx", next="Type /yq-load for a new check.")
        res = status_of("--run", again)
        assert res["text"].splitlines()[-1] == "FAILED at token" and res["result"] == "check_again"
        assert "token: STOPPED before anything was loaded: check again - changed: the drop folder: x.xlsx. Type /yq-load "\
               "for a new check." in res["text"]            # why, then next: each a sentence of its own
        write_lock(other, 2 ** 22 + 12345, now, "publish")                              # another run's dead lock
        assert status_of("--run", run)["state"] == "FAILED"
        assert status_of("--run", "nonsense")["exit"] == 2
        assert not NINE_DIGITS.search(res["text"])


@test("status prints the writer's sentences in full, text then why then next, and none of its machine fields")
def _():
    run = "070126_0930"
    text = "LOAD FAILED at refresh. " + "The database holds this load, or a part of it, and the run is closed. " * 5
    why = "the totals of the day book do not match what the load wrote " * 5 + "(exit code 2; its last line: FAIL)"
    nxt = "This run is closed: type /yq-load for a new check, " * 4 + "then read the table again."
    assert len(text) > 300 and len(why) > 300 and len(nxt) > 200, "longer than any earlier cut"
    with data_root():
        row = loadplan.journal_append(run, "failed", step="refresh", result="failed", written=True, text=text, why=why,
                                      next=nxt, archive="D:/backups/yq-backup_pre-data.7z", path="gates/before.json",
                                      pre="pre-snapshot.json", post="post-snapshot.json", launch="launchtag", pid=4242,
                                      loaded={"orders": 12}, ingest_runs=[5, 6], seconds=12.5, metrics={"k": 1},
                                      moved=["movedpart"], encrypted=True, kept=3, reused=False)
        clock = datetime.fromisoformat(row["at"]).astimezone(status.BAHRAIN).strftime("%H:%M:%S")
        line = lc.journal_line(row)
        assert line == f"  {clock}  refresh: FAILED - {text.strip()} {why}. {nxt}", line
        assert lc.journal_says(row) == f"{text.strip()} {why}. {nxt}"            # a sentence left open is closed
        for field in ("yq-backup", "before.json", "snapshot", "launchtag", "4242", "orders", "[5, 6]", "12.5", "movedpart",
                      "True", "False"):
            assert field not in line, field
        res = status_of()
        assert res["text"].splitlines() == [f"Load {run}", line, "FAILED at refresh"], res["text"]
        # every row of every kind: only the three fields, whatever else the writer adds later
        loadplan.journal_append(run, "note", step="backup", text="One. ", why=None, next="  Two  ", new_field="machine")
        loadplan.journal_append(run, "step_ok", step="backup", archive="x.7z", encrypted=False)
        shown_now = status_of()["text"].splitlines()
        assert shown_now[1].endswith("backup: note - One. Two.") and shown_now[2].endswith("backup: done"), shown_now
        assert "machine" not in "\n".join(shown_now) and "x.7z" not in "\n".join(shown_now)
        assert lc.JOURNAL_SAYS == ("text", "why", "next") and lc.journal_says({}) == ""
        assert lc.journal_says({"text": "Is it there?", "why": "yes: ", "next": "go on!"}) == "Is it there? yes: go on!"


@test("status: a journal with no end and no lock is RUNNING while fresh, INTERRUPTED once it is stale")
def _():
    run = "070126_0930"
    with data_root():
        loadplan.journal_append(run, "started")
        loadplan.journal_append(run, "step_start", step="token")
        assert lc.run_state(run) == ("RUNNING", "token")
        later = datetime.now(timezone.utc) + timedelta(seconds=status.HEARTBEAT_STALE_S + 30)
        assert lc.run_state(run, now=later) == ("INTERRUPTED", "token")


@test("status with no --run follows the run that holds the lock, whatever was checked after it")
def _():
    a, b, now = "070126_0930", "070126_0940", datetime.now(timezone.utc)
    with data_root():
        assert lc.locked_run() is None
        loadplan.write_plan({"run": a, "drop": "x", "token": "t", "verdicts": []})
        loadplan.journal_append(a, "started")
        loadplan.journal_append(a, "step_start", step="refresh")
        write_lock(a, os.getpid(), now, "refresh")
        loadplan.write_plan({"run": b, "drop": "x", "token": None, "verdicts": []})      # a check made after the go
        later = now.timestamp() + 600
        os.utime(loadplan.run_dir(b), (later, later))
        assert loadplan.latest_run() == b and lc.locked_run() == a
        res = status_of()
        assert res["run"] == a and res["state"] == "RUNNING" and res["text"].splitlines()[-1] == "RUNNING at refresh", res["text"]
        assert res["text"].splitlines()[0] == f"Load {a}"
        assert status_of("--run", b)["state"] == "NOT STARTED"            # a run that is named is the run followed
        write_lock(a, 2 ** 22 + 12345, now, "refresh")                    # cut off: still the load he asks about
        res = status_of()
        assert res["run"] == a and res["state"] == "INTERRUPTED" and res["exit"] == 2
        paths.write_text(paths.lock_file(), "not a lock at all")          # a lock that names no run: the latest run
        os.utime(loadplan.run_dir(b), (later, later))
        assert lc.locked_run() is None and status_of()["run"] == b
        paths.lock_file().unlink()
        os.utime(loadplan.run_dir(b), (later, later))
        assert status_of()["run"] == b and status_of()["state"] == "NOT STARTED"


@test("status --wait returns as soon as the load ends, and after 100 seconds at the latest")
def _():
    run, now = "070126_0930", datetime.now(timezone.utc)
    with data_root():
        loadplan.journal_append(run, "started")
        write_lock(run, os.getpid(), now, "refresh")
        deps = FakeDeps()
        calls = {"n": 0}

        def tick():
            calls["n"] += 1
            if calls["n"] == 3:
                loadplan.journal_append(run, "step_ok", step="refresh")
                loadplan.journal_append(run, "finished")
        deps.on_pause = tick
        res = status_of("--wait", deps=deps)
        assert res["state"] == "FINISHED" and calls["n"] == 3 and deps.ticks == 3 * lc.WAIT_POLL_S
        assert "refresh: done" in res["text"] and res["text"].splitlines()[-1] == "FINISHED"
    with data_root():
        loadplan.journal_append(run, "started")
        write_lock(run, os.getpid(), now, "refresh")
        deps = FakeDeps()
        res = status_of("--wait", deps=deps)
        assert res["state"] == "RUNNING" and lc.WAIT_MAX_S <= deps.ticks < lc.WAIT_MAX_S + 2 * lc.WAIT_POLL_S
        assert res["text"].splitlines()[-1] == "RUNNING at refresh" and res["exit"] == 0
    with data_root():
        loadplan.write_plan({"run": run, "drop": "x", "token": "t", "verdicts": []})
        deps = FakeDeps()
        res = status_of("--wait", deps=deps)                              # the go was refused: nothing ever starts
        assert res["state"] == "NOT STARTED" and lc.WAIT_START_GRACE_S <= deps.ticks < lc.WAIT_MAX_S
    with data_root():                                                     # a failure ends the wait as well
        loadplan.journal_append(run, "started")
        write_lock(run, os.getpid(), now, "backup")
        deps = FakeDeps()
        deps.on_pause = lambda: loadplan.journal_append(run, "failed", step="backup", result="failed", why="No space.")
        res = status_of("--wait", deps=deps)
        assert res["state"] == "FAILED" and deps.ticks == lc.WAIT_POLL_S and res["exit"] == 2
        assert res["text"].splitlines()[-1] == "FAILED at backup" and "backup: FAILED - No space." in res["text"]
    with data_root():                                                     # and so does a worker that is gone
        loadplan.journal_append(run, "started")
        write_lock(run, 2 ** 22 + 12345, now, "refresh")
        deps = FakeDeps()
        res = status_of("--wait", deps=deps)
        assert res["state"] == "INTERRUPTED" and deps.ticks == 0 and res["exit"] == 2


@test("status after a load: RUNNING at delivery while the e-mail is made, then its one Delivery line; a cut-off one says its retry")
def _():
    run = "070126_0930"
    said = "Delivery: e-mail sent to 1 address(es) (full, auto): YQ brief"
    with data_root():
        loadplan.journal_append(run, "started")
        loadplan.journal_append(run, "finished", result="ok", text="Loaded and published.")
        loadplan.delivery_append(run, "queued", text="Preparing the brief and the e-mail.")
        res = status_of("--run", run)
        lines = res["text"].splitlines()
        assert lines[-1] == "RUNNING at delivery" and res["exit"] == 0 and res["state"] == "RUNNING", lines
        assert lines[-2].endswith("delivery: started - Preparing the brief and the e-mail."), lines
        loadplan.delivery_append(run, "pack", kind="full", manifest="x/manifest.json")
        loadplan.delivery_append(run, "sent", kind="full", how="auto", text="The full e-mail went to 1 address(es).")
        res = status_of("--run", run)                   # the sender's own row waits for the writer's last word
        assert res["text"].splitlines()[1:] == ["  (no new journal line)", "RUNNING at delivery"], res["text"]
        loadplan.delivery_append(run, "ended", text=said)
        res = status_of("--run", run)
        assert res["text"].splitlines()[1:] == ["  " + said, "FINISHED"] and res["exit"] == 0, res["text"]
        assert status_of("--run", run)["text"].splitlines()[1:] == ["  (no new journal line)", "FINISHED"]
        # a send by hand afterwards (scripts/yq_send.py) shows its own row, worded as an e-mail, never as a load
        loadplan.delivery_append(run, "failed", why="the mail service did not take it: HTTP 401",
                                 text="The e-mail could not be sent.", next=f"Its retry: {lc.SEND_RETRY.format(run=run)}")
        line = status_of("--run", run)["text"].splitlines()[1]
        assert line.endswith("delivery: e-mail NOT sent - The e-mail could not be sent. the mail service did not take it: "
                             f"HTTP 401. Its retry: python -m scripts.yq_send --run {run}"), line
        assert "FAILED" not in line
    with data_root():                                   # the PC slept while the e-mail was made: cut off, and said so
        loadplan.journal_append(run, "started")
        loadplan.journal_append(run, "finished", result="ok", text="Loaded and published.")
        old = (datetime.now(timezone.utc) - timedelta(seconds=lc.DELIVERY_WAIT_S + 60)).isoformat(timespec="seconds")
        loadplan.delivery_append(run, "queued", at=old, text="Preparing the brief and the e-mail.")
        assert lc.delivery_cut_off(run) and not lc.delivery_pending(run)
        res = status_of("--run", run)
        lines = res["text"].splitlines()
        assert lines[-1] == "FINISHED" and res["exit"] == 0, lines
        assert lines[-2] == ("  Delivery: cut off before the e-mail went (the PC slept, or the load's window was closed). "
                             f"Nothing was reloaded. Its retry: python -m scripts.yq_send --run {run}"), lines
        assert status_of("--run", run)["text"].splitlines()[-2].startswith("  Delivery: cut off")   # until it is sent
        loadplan.delivery_append(run, "sent", kind="full", how="by hand", text="The full e-mail went to 1 address(es).")
        assert not lc.delivery_cut_off(run)
        res = status_of("--run", run)
        assert "cut off" not in res["text"] and "delivery: e-mail sent - The full e-mail went" in res["text"], res["text"]
    with data_root():                                   # a rehearsal's delivery: one ended row, nothing pending
        loadplan.journal_append(run, "finished", result="ok")
        loadplan.delivery_append(run, "ended", text="Delivery: (rehearsal) a real load would now write the brief.")
        assert lc.run_state(run)[0] == "FINISHED" and not lc.delivery_cut_off(run)
        assert "  Delivery: (rehearsal) a real load would now write the brief." in status_of("--run", run)["text"]


@test("progress lines go to the screen as the check runs, and to stderr when one JSON object is asked for")
def _():
    deps = lc.Deps()
    out, err = io.StringIO(), io.StringIO()
    real_out, real_err = sys.stdout, sys.stderr
    sys.stdout, sys.stderr = out, err
    try:
        deps.say("[1/9] one")
        deps.quiet = True
        deps.say("[2/9] two")
    finally:
        sys.stdout, sys.stderr = real_out, real_err
    assert out.getvalue() == "[1/9] one\n" and err.getvalue() == "[2/9] two\n"
    assert lc.STEPS == 9 and lc.DEADLINE_S < 600 and lc.WAIT_MAX_S == 100 and lc.WAIT_POLL_S == 5
    assert set(lc.WORDS) == set(loadplan.VERDICTS)
    assert sorted(lc.WORDS.values()) == sorted(["LOAD", "LOAD (main-store block only)", "HOLD BACK", "LEAVE OUT",
                                                "NEEDS YOUR DECISION", "NOT LOADED", "STOP"])


@test("the command is registered: `load status` runs through the entry point and prints the state last")
def _():
    with data_root() as root:
        loadplan.journal_append("070126_0930", "started")
        loadplan.journal_append("070126_0930", "finished")
        buf = io.StringIO()
        with redirect_stdout(buf):
            code = cli.main(["load", "status", "--data-root", str(root)])
        assert code == 0 and buf.getvalue().strip().splitlines()[-1] == "FINISHED", buf.getvalue()
        buf = io.StringIO()
        with redirect_stdout(buf):
            code = cli.main(["load", "status", "--json", "--data-root", str(root)])
        assert json.loads(buf.getvalue())["state"] == "FINISHED" and code == 0
    assert cli.COMMANDS["load"][0] == "scripts.ops.load_check" and "load" in cli.NO_BANNER


@test("the module holds no threshold and no name from the data: limits come from config, the store from config")
def _():
    src = (ROOT / "scripts" / "ops" / "load_check.py").read_text(encoding="utf-8")
    assert 'config.number("load.profit_gap_max")' in src and 'config.number("load.blank_cost_share_max")' in src
    assert 'config.get("warehouses.main")' in src
    for word in ("Accessories", "Causeway", "Roadshow", "Acc WH"):
        assert word not in src, word
    from scripts.ops import selftest
    assert not selftest.scan_text(src, []), "a three-decimal amount of 1,000 or more in the module"
    assert not selftest.scan_text(Path(__file__).read_text(encoding="utf-8"), [])
    # it stands on the shared foundation: the one guard for what it writes, copies, moves and removes, and the one
    # register (no copy of either, no private name of another module, no empty first run to dodge the baseline)
    import ast
    tree = ast.parse(src)
    imported = {a.name for n in ast.walk(tree) if isinstance(n, ast.Import) for a in n.names}
    assert "shutil" not in imported and "json" not in imported, imported
    for gone in ("_allowed_output", "rule_has_run", "NO_BASELINE", "shutil.", ".unlink("):
        assert gone not in src, gone
    for used in ("paths.allowed_output(", "paths.copy_file(", "paths.remove_file(", "status.failed_load()"):
        assert used in src, used
    records = [n for n in ast.walk(tree) if isinstance(n, ast.Call) and ast.unparse(n.func) == "alerts.record"]
    assert len(records) == 1 and sorted(k.arg for k in records[0].keywords) == ["baseline", "clear_missing"], "one call"
    assert lc.EVENT_RULES == ("inv_changed", "voucher_changed") and lc.STATE_RULES == ("price_scope", "profit_held")
    assert lc.BASELINE_RULES == ("price_scope",), "only this rule's first run is a baseline"


# ── the real 7-Oct files, when they are on this machine ───────────────────────

@test("real files (SKIP when absent): the trimmed copy re-reads to the main-store block and Focus's Sub Total")
def _():
    from scripts import ingest as ing
    paths.set_data_root(None)
    config.use(None)
    real = paths.data_root()
    drop = real / "Uplaods" / "071026"
    cfg = real / "business_data" / "ops" / "config.json"
    src = next((p for p in drop.glob("*.xlsx") if ing.classify(p.name) == "stock_balance"), None) if drop.is_dir() else None
    if src is None or not cfg.is_file():
        print("      SKIP: the 7-Oct drop folder is not on this machine")
        return
    values = json.loads(cfg.read_text(encoding="utf-8"))
    main, aliases = values["warehouses.main"], values["warehouses.aliases"]
    before = loadplan.sha256_file(src)
    with data_root() as root:
        grid = ing.read_grid(src)
        rows = ing.parse_stock_balance(grid, src.name)
        block = lc.stock_block(grid, main, aliases)
        assert len(set(block["warehouses"])) > 1
        dst = lc.write_trimmed(src, root / "business_data" / "staging" / "real" / ("only_" + src.name), grid, block)
        proof = lc.prove_block(rows, dst, main, aliases, block)
        assert proof["problems"] == [], proof["problems"]
        assert proof["units"] == lc.dec(block["sub_units"]) and fmt.q3(proof["value"]) == fmt.q3(block["sub_value"])
        assert proof["items"] == len([r for r in rows if r["warehouse_name"] == main]) > 0
    assert loadplan.sha256_file(src) == before


# ── 14-Oct-2026: the profitability evidence, the count of holds, the check's timings, another PC's lock ──

@test("a held report: the action says what accounts must do and since when; the loads are counted; the plan keeps the evidence")
def _():
    with data_root() as root:
        make_drop(root, profit=SHORT_COST)
        loads_before = ["2026-01-03", "2026-01-05"]          # two loads since the report (by hand: their snapshots)
        res, _ = check("070126", deps=FakeDeps(FakeSession(anchors=BEFORE_LOAD, snapshot_days=loads_before)))
        text = res["text"]
        assert "held back for the third load in a row" in text, text      # two loads since the report + this one
        item = [r for r in alerts.open_items() if r["rule"] == "profit_held"][0]
        assert item["title"].startswith("Ask accounts to post cost of sales in Focus for the invoices since Fri 02 Jan")
        assert item["title"].endswith("then export the profitability report again") and item["who"] == "Accounts"
        assert "not loaded for 3 loads" in item["detail"] and "production's report is as on" in item["detail"]
        plan = loadplan.read_plan(res["run"])
        ev = plan["profitability"]
        assert ev["held"] is True and ev["verdict"] == "hold_back" and ev["loads_held"] == 3
        assert ev["gap_pct"] > ev["limit_pct"] > 0 and ev["production_as_of"] == "2026-01-01"
        assert ev["ledger_cost_bhd"] and ev["report_cost_bhd"] and ev["line"].startswith("Profitability: ")
        timings = plan["timings"]
        assert {"preflight", "folder", "staging", "profitability", "preview", "total"} <= set(timings), timings
        assert timings["parse"]["memo"] is True and timings["parse"]["files"], timings["parse"]
        assert all(r["parsed"] == 1 for r in timings["parse"]["files"].values()), "a file was parsed twice"
    assert [lc.ordinal(n) for n in (1, 2, 3, 10, 11)] == ["first", "second", "third", "tenth", "11th"]


@test("one parse per report file: the plan is the same with the memo on and off")
def _():
    def plan_of(memo: bool) -> dict:
        with data_root() as root:
            make_drop(root, profit=SHORT_COST)
            old = os.environ.pop("YQ_CHECK_NO_MEMO", None)
            try:
                if not memo:
                    os.environ["YQ_CHECK_NO_MEMO"] = "1"
                res, _ = check("070126", deps=FakeDeps(FakeSession(anchors=BEFORE_LOAD)))
            finally:
                os.environ.pop("YQ_CHECK_NO_MEMO", None)
                if old is not None:
                    os.environ["YQ_CHECK_NO_MEMO"] = old
            plan = loadplan.read_plan(res["run"])
            parse = plan.pop("timings")["parse"]
            for key in ("run", "created_at", "staging", "token", "token_parts", "drop"):
                plan.pop(key, None)
            return {"plan": plan, "parse": parse}
    on, off = plan_of(True), plan_of(False)
    assert on["plan"] == off["plan"], "the memo changed what the check found"
    assert off["parse"]["memo"] is False and on["parse"]["memo"] is True
    assert sum(r["parsed"] for r in on["parse"]["files"].values()) <= sum(r["parsed"] for r in off["parse"]["files"].values())
    from scripts import ingest as ing
    assert ing.read_grid.__module__ == "scripts.ingest", "the original parser was put back"
    # the memo itself (the preview of the tests parses nothing, so it is driven here directly)
    from scripts.ops.gridmemo import GridMemo
    with data_root() as root:
        drop = make_drop(root)
        ledger = next(drop.glob("Stock_ledger*"))
        twin = root / "twin.xlsx"
        shutil.copyfile(ledger, twin)                 # the staged copy: another path, the same content
        memo = GridMemo(memo=True)
        with memo.active():
            first, again, other = ing.read_grid(ledger), ing.read_grid(ledger), ing.read_grid(twin)
            first.iat[0, 0] = "changed by one caller"
            assert again.iat[0, 0] != "changed by one caller" != other.iat[0, 0], "each caller has its own copy"
            twin.write_bytes(next(drop.glob("Stock_balance*")).read_bytes())    # the file changed: parsed again
            changed = ing.read_grid(twin)
        assert len(changed) != len(again)
        reads = memo.summary()["files"]
        assert reads[ledger.name] == {"parsed": 1, "served": 1, "seconds": reads[ledger.name]["seconds"]}
        assert reads["twin.xlsx"]["served"] == 1 and reads["twin.xlsx"]["parsed"] == 1
        assert ing.read_grid.__module__ == "scripts.ingest"
        off_memo = GridMemo(memo=False)
        with off_memo.active():
            ing.read_grid(ledger), ing.read_grid(ledger)
        assert off_memo.summary()["files"][ledger.name]["parsed"] == 2


@test("a lock taken on another PC stops the check: no GO line, nothing resumed or cleared from this PC")
def _():
    with data_root() as root:
        make_drop(root)
        for state in ("running", "interrupted"):
            deps = FakeDeps(lock_state=lambda st=state: {"state": st, "run": "060126_0900", "step": "refresh",
                                                         "elsewhere": "OTHER-PC"})
            res, _ = check("070126", deps=deps)
            assert res["exit"] == 2 and res["token"] is None and "GO:" not in res["text"], res["text"]
            assert "taken on another PC (OTHER-PC)" in res["text"] and "finish that load on OTHER-PC" in res["text"]


# ── 15-Oct-2026: nothing new to load, the count of held loads, the preview's seconds per report ──

LOADED = {"today": DAY, "sales_to": "2026-01-06", "stock_as_of": DAY, "ar_as_of": DAY, "profit_as_of": DAY,
          "moves_to": "2026-01-06"}
SAME = {"inserted": 0, "updated": 0, "unchanged": 7, "deleted": 0}
LAST_LOAD = {"run": "070126_0900", "finished_at": "2026-01-07T09:05:00+03:00"}
NOTHING_NEW = ("Nothing new in 070126: everything this load would apply is already loaded (last load 070126_0900 at "
               "Wed 07 Jan 09:05). No load needed.")
SAME_BOOK = "- Price book: LEAVE OUT (the same as the live price book: nothing to load)."


def already_loaded(session: dict | None = None, **over) -> FakeDeps:
    """A check of the 7-Jan drop after its load: production holds every row of every report (the stock ledger
    to its last voucher), and the preview finds nothing to insert, change or remove anywhere."""
    deps = FakeDeps(FakeSession(**dict({"anchors": dict(LOADED), "ledger": db_ledger(upto="2026-01-06")}, **(session or {}))))
    deps.actions = {t: dict(SAME) for t in ("order_lines", "orders", "stock_movements", "stock_balance", "ar_ageing",
                                            "product_profitability")}
    deps.actions["selling_prices"] = {"inserted": 0, "updated": 0, "unchanged": 7, "voided": 0, "revived": 0}
    deps.loaded_rows = {"ar_ageing": 7}
    for key, value in over.items():
        if key == "actions":
            for target, change in value.items():
                deps.actions[target].update(change)
        else:
            setattr(deps, key, value)
    return deps


@test("an already-loaded drop: the whole check runs, then it ends 'Nothing new': no GO line, no token, no go.cmd")
def _():
    with data_root() as root:
        drop = make_drop(root)
        before = snapshot_of(drop)
        status.write_state({"last_load": dict(LAST_LOAD)})
        res, deps = check("070126", deps=already_loaded())
        text, lines = res["text"], res["text"].splitlines()
        assert res["ok"] and res["exit"] == 0 and res["token"] is None, text
        assert lines[-1] == NOTHING_NEW, lines[-1]
        assert not [ln for ln in lines if ln.startswith("GO: ")] and "Token " not in text and "double-click" not in text
        assert "nothing_new" in res["flags"] and not {"token", "stop", "dev"} & set(res["flags"]), res["flags"]
        assert not (loadplan.run_dir(res["run"]) / "go.cmd").exists()
        plan = loadplan.read_plan(res["run"])
        assert plan["token"] is None and plan["token_parts"] is None and plan["stops"] == []
        assert plan["nothing_new"] is True and plan["new_in_drop"] == [], plan["new_in_drop"]
        # the full check ran: its nine steps, the table and every comparison; only its end is another
        assert [ln[:5] for ln in deps.steps] == [f"[{i}/9]" for i in range(1, 10)], deps.steps
        table = table_of(text)
        assert table[shown("orders")] == "LOAD" and table[shown("stock_balance")] == "LOAD (main-store block only)"
        assert table[shown("selling_prices:MA_base")] == "LEAVE OUT" and table[shown("product_profitability")] == "LOAD"
        assert "Invoices: 0 new, 0 changed, 0 removed. Lines: 0 new, 0 changed, 0 removed." in text
        assert "Sales by day and by salesman: the file equals what is loaded." in text
        assert "changed 0, removed 0, entered with an earlier date 0, new 0" in text and "switch nothing" in text
        # no load follows, so neither what a load would show nor what it also does is said
        assert "After the load you should see" not in text and "A load also does this" not in text
        assert "Production already holds (the same as these reports)" in text and "Sales to Tue 06 Jan" in text
        # what is not part of the load is named, just above the last line
        assert lines[-4:-1] == ["Not part of this load:", SAME_BOOK, ""], lines[-6:]
        assert snapshot_of(drop) == before and deps.the_session.closed
        fmt.careful(text)
        # a developer's proof run says the same and issues nothing either
        res, _ = check("070126", "--dev", deps=already_loaded(worktree=True, site={"state": "behind"}))
        assert res["text"].splitlines()[-1] == NOTHING_NEW and res["token"] is None, res["text"]
        assert {"dev", "nothing_new"} <= set(res["flags"]) and "GO:" not in res["text"]
    with data_root() as root:                     # no load on record on this PC: the line leaves the last load out
        make_drop(root, leave_out=("selling_prices:MA_base",))
        res, _ = check("070126", deps=already_loaded())
        lines = res["text"].splitlines()
        assert lines[-1] == ("Nothing new in 070126: everything this load would apply is already loaded. "
                             "No load needed."), res["text"]
        assert "Not part of this load" not in res["text"], "nothing held back or left out: no such list"


@test("loaded already, the profitability report held back (the real 15-Oct shape): nothing new, and the held report named")
def _():
    held = dict(LOADED, profit_as_of="2026-01-01")                # production's report: older, held back since
    with data_root() as root:
        make_drop(root, profit=SHORT_COST)
        status.write_state({"last_load": dict(LAST_LOAD)})
        loads = ["2026-01-03", "2026-01-05", DAY]                 # two loads by hand before this drop's own load
        res, _ = check("070126", deps=already_loaded(session={"anchors": held, "snapshot_days": loads}))
        text, lines = res["text"], res["text"].splitlines()
        assert table_of(text)[shown("product_profitability")] == "HOLD BACK", text
        assert lines[-1] == NOTHING_NEW and res["token"] is None and "GO:" not in text, text
        assert lines[-5:-1] == ["Not part of this load:",
                                "- Profitability report: HOLD BACK. To load it anyway, say: load profitability anyway.",
                                SAME_BOOK, ""], lines[-6:]
        assert "held back for the third load in a row (this drop is already loaded: the action list stays as it is)" in text
        # on his word the report is part of the load, and the check ends as today
        deps = already_loaded(session={"anchors": held}, actions={"product_profitability": {"inserted": 3, "unchanged": 0}})
        res, _ = check("070126", "--with-profitability", deps=deps)
        assert res["token"] and res["text"].splitlines()[-1].startswith("GO: ") and "Nothing new" not in res["text"]


@test("a drop whose only difference is a line for his decision: its default loads nothing, so 'Nothing new' and his words")
def _():
    newer_prices = PRICES + [("Charger two", "C2", None, None, "Van A", "1/5/2026", "12/31/2030", 2.0)]
    with data_root() as root:
        make_drop(root, prices=newer_prices)
        status.write_state({"last_load": dict(LAST_LOAD)})
        res, _ = check("070126", deps=already_loaded())
        text, lines = res["text"], res["text"].splitlines()
        assert table_of(text)[shown("selling_prices:MA_base")] == "NEEDS YOUR DECISION", text
        assert lines[-1] == NOTHING_NEW and res["token"] is None and "GO:" not in text, text
        assert ("- Price book: NEEDS YOUR DECISION (left out unless you say load). To load it, say: load the price book."
                in lines), lines[-6:]
        assert not (loadplan.run_dir(res["run"]) / "go.cmd").exists()
        # he says "load the price book": the book is part of the load, and the check ends as today
        res, _ = check("070126", "--pricebook", "load",
                       deps=already_loaded(actions={"selling_prices": {"inserted": 1, "unchanged": 6}}))
        assert res["token"] and res["text"].splitlines()[-1].startswith("GO: ") and "Nothing new" not in res["text"]
        # left out on his word, it is named with the words that load it after all
        res, _ = check("070126", "--pricebook", "skip", deps=already_loaded())
        assert res["text"].splitlines()[-1] == NOTHING_NEW, res["text"]
        assert "- Price book: LEAVE OUT (left out on your word). To load it, say: load the price book." in res["text"]


@test("one difference anywhere and the check ends as today: the table, the token and the GO line")
def _():
    other = [{"item_code": c, "is_active": True, "hidden": False, "updated_by": lc.AUTO_SYNC} for c in ("C1", "C2", "C3", "C9")]
    day_changed = {"key": "2026-01-06", "db_rows": 3, "file_rows": 3, "file_gross": "10.000", "db_gross": "9.000"}
    rep_changed = {"key": "Van A", "db_rows": 3, "file_rows": 3, "file_gross": "10.000", "db_gross": "9.000"}
    cases = {
        "a receivables row rewritten": ((), {}, already_loaded(actions={"ar_ageing": {"updated": 1, "unchanged": 6}})),
        "a stock snapshot not loaded yet": ((), {}, already_loaded(actions={"stock_balance": {"inserted": 3, "unchanged": 0}})),
        "an invoice that went": ((), {}, already_loaded(actions={"orders": {"deleted": 1}})),
        "a price row that comes back": (("--pricebook", "load"), {},
                                        already_loaded(actions={"selling_prices": {"revived": 1}})),
        "a price row voided, nothing else": (("--pricebook", "load"), {},
                                             already_loaded(actions={"selling_prices": {"voided": 1}})),
        "a voucher production lacks": ((), {}, already_loaded(session={"ledger": db_ledger(upto="2026-01-05")})),
        "a catalog switch": ((), {}, already_loaded(session={"catalog": other})),
        "a day's sales differ, nothing else": ((), {}, already_loaded(
            sales={"per_day_changed": [day_changed], "per_salesman_changed": []})),
        "a salesman's sales differ, nothing else": ((), {}, already_loaded(
            sales={"per_day_changed": [], "per_salesman_changed": [rep_changed]})),
        "a report the preview did not compare": ((), {}, already_loaded(samples={"orders": {"actions": {}}})),
        "the last load failed after it had begun to write": ((), {"failed": True}, already_loaded()),
    }
    for name, (options, drop_kw, deps) in cases.items():
        with data_root() as root:
            failed = drop_kw.pop("failed", False)
            make_drop(root, **drop_kw)
            status.write_state({"last_load": dict(LAST_LOAD)})
            if failed:
                loadplan.journal_append("070126_0901", "started", pid=0)
                loadplan.journal_append("070126_0901", "step_start", step="refresh")
                loadplan.journal_append("070126_0901", "failed", result="failed", step="refresh", written=True)
            res, _ = check("070126", *options, deps=deps)
            lines = res["text"].splitlines()
            assert res["ok"] and res["token"] and lines[-1] == "GO: " + loadplan.go_line(res["run"], res["token"]), (name, res["text"])
            assert "Nothing new" not in res["text"] and "nothing_new" not in res["flags"], name
            assert "After the load you should see" in res["text"] and "A load also does this, every time" in res["text"]
            assert loadplan.read_plan(res["run"])["nothing_new"] is False and loadplan.read_plan(res["run"])["new_in_drop"], name
            assert (loadplan.run_dir(res["run"]) / "go.cmd").is_file(), name


AGEING = [("Shop A", 0.1), ("Shop A", 0.2), ("Shop B", 1.5)]       # Focus lists Shop A twice: the loader adds the two up
FOLDED_COLUMNS = ("balance_bhd", "bucket_0_30", "total_bhd")      # the columns of Shop A's two lines that are not zero


def ageing_loaded(root: Path, stored: dict, updated: tuple = ("Shop A",)) -> FakeDeps:
    """The drop with a receivables report the real parser reads, checked after its load: production holds the
    rows the loader wrote (Shop A as the loader's float sum of its two lines, unless `stored` says otherwise),
    and the preview counts the accounts of `updated` as rewritten."""
    from app.ingest_batch import prepare
    from scripts import ingest as ing
    drop = make_drop(root)
    path = receivables_file(drop / NAMES["receivables"], AGEING)
    prepared, notes = prepare("ar_ageing", ing.parse_receivables(ing.read_grid(path), path.name))
    assert notes["aggregated_same_name"] == 1 and len(prepared) == 2
    loader_sum = Decimal(repr(0.1 + 0.2))                         # what pandas adds up and PostgREST stores
    assert loader_sum != Decimal("0.3")
    rows = []
    for i, r in enumerate(prepared):
        row = {k: (Decimal(str(v)) if isinstance(v, float) else v) for k, v in r.items()}
        row.update(id=i + 1, as_of_date=datetime(2026, 1, 7).date(), imported_at="2026-01-07T09:01:00+00:00")
        if r["account"] == "Shop A":
            row.update({c: loader_sum for c in FOLDED_COLUMNS})
        row.update(stored.get(r["account"], {}))
        rows.append(row)
    deps = already_loaded(session={"partitions": {"ar_ageing": rows}},
                          actions={"ar_ageing": {"updated": len(updated), "unchanged": 2 - len(updated)}})
    deps.samples = {"ar_ageing": {"updated_sample": [{"account": a, "as_of_date": DAY} for a in updated],
                                  "payload_cols": list(prepared[0]),
                                  "scope": {"kind": "partition", "cols": ["as_of_date"], "partitions": [[DAY]]}}}
    deps.hook_rows = {"ar_ageing": prepared}
    status.write_state({"last_load": dict(LAST_LOAD)})
    return deps


@test("a receivables row the loader adds up from two lines of one account: its float rounding is no change, a fils is")
def _():
    with data_root() as root:                     # 151026, re-checked after its load: "0 new, 1 rewritten, 0 removed"
        res, deps = check("070126", deps=ageing_loaded(root, {}))
        text = res["text"]
        assert res["text"].splitlines()[-1] == NOTHING_NEW and res["token"] is None, text
        assert "(that day is already loaded, with the same total: 0 new, 1 rewritten, 0 removed)" in text   # the table as it was
        assert ("Receivables: 1 row the preview counts as rewritten is one the loader adds up from two or more lines of "
                "one account, and it differs from what is loaded only beyond the sixth decimal place (the loader keeps "
                "such a sum with floating-point rounding): not a change.") in text, text
        asked = [s for s in deps.the_session.seen if s.startswith("select * from ar_ageing")]
        assert asked == ["select * from ar_ageing where as_of_date::text = %s"], asked
    with data_root() as root:                     # production's sum is one fils off: that is a change
        res, _ = check("070126", deps=ageing_loaded(root, {"Shop A": {"balance_bhd": Decimal("0.301")}}))
        assert res["token"] and res["text"].splitlines()[-1].startswith("GO: "), res["text"]
        assert "Nothing new" not in res["text"] and "beyond the sixth decimal place" not in res["text"]
    with data_root() as root:                     # a row the loader does not add up is compared exactly, as by the preview
        res, _ = check("070126", deps=ageing_loaded(root, {"Shop B": {"balance_bhd": Decimal(repr(1.5 + 1e-15))}},
                                                    updated=("Shop B",)))
        assert res["token"] and "Nothing new" not in res["text"], res["text"]
    with data_root() as root:                     # more rows counted as rewritten than the preview's sample names
        deps = ageing_loaded(root, {})
        deps.actions["ar_ageing"].update(updated=2, unchanged=0)
        res, _ = check("070126", deps=deps)
        assert res["token"] and "Nothing new" not in res["text"], res["text"]
    with data_root() as root:                     # a text column that differs on the folded row is a change too
        res, _ = check("070126", deps=ageing_loaded(root, {"Shop A": {"group_name": "Other"}}))
        assert res["token"] and "Nothing new" not in res["text"], res["text"]
    with data_root() as root:                     # the stock balance: one item twice in the main store, its rate averaged
        from app.ingest_batch import prepare
        from scripts import ingest as ing
        blocks = [("Van A", [("Cable one C1", 10, 1.5)]),
                  (MAIN, [("Cable one C1", 10, 1.1), ("Cable one C1", 20, 2.2), ("Charger two C2", 40, 2.5),
                          ("Holder three C3", 20, 0.75)])]
        drop = make_drop(root, stock=blocks)
        path = drop / NAMES["stock_balance"]
        raw = [r for r in ing.parse_stock_balance(ing.read_grid(path), path.name) if r["warehouse_name"] == MAIN]
        prepared, notes = prepare("stock_balance", raw)
        assert notes["aggregated_same_item"] == 1 and prepared[0]["item_name"] == "Cable one C1"
        stored = []
        for i, r in enumerate(prepared):
            row = {k: (Decimal(str(v)) if isinstance(v, float) else v) for k, v in r.items()}
            row.update(id=i + 1, as_of_date=datetime(2026, 1, 7).date())
            stored.append(row)
        stored[0]["selling_rate_bhd"] = Decimal(repr((1.1 + 2.2) / 2))          # the loader's float mean of the two rates
        assert stored[0]["selling_rate_bhd"] != Decimal(str(prepared[0]["selling_rate_bhd"]))
        deps = already_loaded(session={"partitions": {"stock_balance": stored}},
                              actions={"stock_balance": {"updated": 1, "unchanged": 2}})
        deps.samples = {"stock_balance": {
            "updated_sample": [{"item_name": "Cable one C1", "warehouse_name": MAIN, "as_of_date": DAY}],
            "payload_cols": list(prepared[0]),
            "scope": {"kind": "partition", "cols": ["as_of_date", "warehouse_name"], "partitions": [[DAY, MAIN]]}}}
        deps.hook_rows = {"stock_balance": prepared}
        status.write_state({"last_load": dict(LAST_LOAD)})
        res, deps = check("070126", deps=deps)
        assert res["text"].splitlines()[-1] == NOTHING_NEW, res["text"]
        assert ("Stock balance: 1 row the preview counts as rewritten is one the loader adds up from two or more lines "
                "of one item in one warehouse") in res["text"], res["text"]
        asked = [s for s in deps.the_session.seen if s.startswith("select * from stock_balance")]
        assert asked == ["select * from stock_balance where as_of_date::text = %s and warehouse_name::text = %s"], asked
    rows = [{"account": "Shop A", "as_of_date": DAY}, {"account": "Shop A", "as_of_date": DAY}, {"account": "Shop B"}]
    assert lc.folded_keys("ar_ageing", rows) == {("Shop A", DAY)} and lc.folded_keys("orders", rows) == set()
    blank = [{"item_name": "Cable one", "warehouse_name": None, "as_of_date": DAY},
             {"item_name": "Cable one", "warehouse_name": "(unassigned)", "as_of_date": DAY + "T00:00:00"}]
    assert lc.folded_keys("stock_balance", blank) == {("Cable one", "(unassigned)", DAY)}     # as the loader fills it
    assert lc.same_cell(Decimal("0.3"), Decimal(repr(0.1 + 0.2)), True) and not lc.same_cell(Decimal("0.3"), Decimal("0.301"), True)
    assert not lc.same_cell(Decimal("0.3"), Decimal(repr(0.1 + 0.2)), False), "outside a folded sum: exact"
    assert lc.same_cell("2026-01-07", datetime(2026, 1, 7).date(), False) and not lc.same_cell(None, Decimal(0), True)
    for odd in (float("nan"), "NaN", "sNaN", float("inf"), "x", True):     # not a number: never the same, never an error
        assert not lc.same_cell(odd, Decimal("1.5"), True) and not lc.same_cell(odd, Decimal("1.5"), False), odd
    assert not lc.same_cell(Decimal("1.5"), Decimal("NaN"), True) and not lc.same_cell(float("inf"), Decimal("Infinity"), True)
    assert "decimal place" in lc.noise_line("stock_balance", 2) and "dinar" not in lc.noise_line("stock_balance", 2)


@test("held back for the Nth load in a row: a load by hand, then three loads (the last one twice) count 2, 3, 4, 4, 4")
def _():
    kept = "2026-01-01"                           # production's profitability report: as on this day, held back since
    short = [("Cable one", 300, 50), ("Charger two", 200, 50), ("Holder three", 100, 25)]     # short on every day
    with data_root() as root:
        days = ["2026-01-02"]                     # a load made by hand: its snapshot is in production, it left no journal
        said, counted = [], []
        for name, loads in (("030126", 1), ("050126", 1), ("070126", 1), ("070126", 1), ("070126", 0)):
            day = paths.drop_date(name).isoformat()
            if not (root / "Uplaods" / name).is_dir():
                make_drop(root, name, as_on=f"{name[:2]}/{name[2:4]}/20{name[4:]}", profit=short)
            anchors = {"today": DAY, "sales_to": "2026-01-01", "stock_as_of": max(days), "ar_as_of": max(days),
                       "profit_as_of": kept, "moves_to": "2026-01-01"}
            res, _ = check(name, deps=FakeDeps(FakeSession(anchors=anchors, snapshot_days=list(days))))
            assert res["token"] and table_of(res["text"])[shown("product_profitability")] == "HOLD BACK", res["text"]
            m = re.search(r"held back for the (\w+) load in a row", res["text"])
            said.append(m.group(1) if m else None)
            counted.append(loadplan.read_plan(res["run"])["profitability"]["loads_held"])
            if loads:                             # the writer loads it: its journal ends 'finished', state.json names it
                loadplan.journal_append(res["run"], "started", pid=0)
                loadplan.journal_append(res["run"], "finished", result="ok")
                status.write_state({"last_load": {"run": res["run"], "finished_at": "2026-01-07T10:00:00+03:00"}})
                days.append(day)
        assert said == ["second", "third", "fourth", "fourth", "fourth"], said
        assert counted == [2, 3, 4, 4, 4], counted
        # a re-check of a loaded drop counts it once and leaves the action list as it is
        res_text = res["text"]
        assert "held back for the fourth load in a row (this drop is already loaded: the action list stays as it is)" in res_text
        # a state: each new hold replaced the one before it, and the re-checks neither added nor cleared one
        assert [r["key"] for r in alerts.open_items() if r["rule"] == "profit_held"] == ["profit_held|2026-01-07"]
    # the rule itself: newest first, while each load held its report back; one entry per day; this drop once
    def both(*days) -> dict:
        return {"stock_balance": list(days), "ar_ageing": list(days)}
    assert lc.held_in_a_row("2026-01-07", kept, {"2026-01-03": True, "2026-01-05": True}, both("2026-01-02")) == 4
    assert lc.held_in_a_row("2026-01-07", kept, {"2026-01-03": True, "2026-01-05": False}, both("2026-01-02")) == 1
    assert lc.held_in_a_row("2026-01-07", kept, {"2026-01-05": True, "2026-01-07": True}, both("2026-01-07")) == 2
    assert lc.held_in_a_row("2026-01-07", "2026-01-04", {"2026-01-03": True}, both("2026-01-02", "2026-01-06")) == 2
    assert lc.held_in_a_row("2026-01-07", "", {"2026-01-05": True}, both("2026-01-02")) == 2, "no report date: no snapshot counts"
    # one load counts once, whatever its snapshots: a load whose stock balance was as on the day before its other
    # reports (its plan says so) ...
    split = {"stock_balance": ["2026-01-04"], "ar_ageing": ["2026-01-05"]}
    assert lc.held_in_a_row("2026-01-07", kept, {"2026-01-05": True}, split, {"2026-01-04", "2026-01-05"}) == 2
    # ... and a load by hand whose two snapshots carry two days: one load, not one per kind of snapshot
    assert lc.held_in_a_row("2026-01-07", kept, {}, split) == 2
    assert lc.held_in_a_row("2026-01-07", kept, {}, {"stock_balance": ["2026-01-03", "2026-01-05"],
                                                     "ar_ageing": ["2026-01-03", "2026-01-05"]}) == 3
    assert lc.held_in_a_row("2026-01-07", kept, {}, {"stock_balance": ["2026-01-03"], "ar_ageing": []}) == 2


from app.ingest_batch import Backend as _PreviewBackend  # noqa: E402  (after scripts.ops.db, as every command does)


class ParamBackend(_PreviewBackend):
    """A read-only backend as the check's (app.ingest_batch.PgBackend, read_only): the rows are handed to the diff
    as a parameter; every query is recorded and answered with canned rows."""
    mode, read_only, stage_source = "fake", True, "param"
    MONEY = ("gross_bhd", "taxable_bhd", "balance_bhd", "net_qty", "total_value_bhd", "received_qty", "issued_qty",
             "net_amount_bhd", "rate_bhd")

    def __init__(self):
        self.mem: dict = {}
        self.asked: list = []

    def tables_ready(self):
        return False

    def stage(self, batch_id, target, rows):
        self.mem[target] = rows
        return len(rows)

    def staged(self, batch_id, target):
        return self.mem.get(target, [])

    def query(self, sql, params=None):
        self.asked.append((" ".join(sql.split())[:60], [str(p)[:40] for p in params or []]))
        if "information_schema.columns" in sql:
            return [{"column_name": c, "data_type": "numeric" if c.endswith(("_bhd", "_qty")) else "text"}
                    for c in (self.mem.get(params[0]) or [{}])[0]]
        if "jsonb_array_elements" in sql:
            n = len(self.mem.get(params[1]) or [])
            return [{"removed": 0, "updated": 1, "revived": 0, "unchanged": n - 1, "inserted": 0, "db_rows": n,
                     **{f"db_sum_{m}": "1.000" for m in self.MONEY},
                     "removed_sample": None, "updated_sample": json.dumps([{"k": params[1]}])}]
        if "from order_lines" in sql:
            return [{"d": "2026-01-06", "w": "Van A", "n": 1, "g": "10.000", "t": "9.000"}]
        return []


@test("the preview's timing proxy: the real preview finds the same through it; the seconds go to each report and helper")
def _():
    from app.ingest_batch import Parsed, build_preview, prepare
    from scripts.ops import querytime

    def parsed() -> Parsed:
        p = Parsed(folder="x")
        p.rows["orders"], _n = prepare("orders", [{"invoice_no": "INV-1", "order_date": "2026-01-06", "gross_bhd": 10.0}])
        p.rows["order_lines"], _n = prepare("order_lines", [{"invoice_no": "INV-1", "line_no": 1, "line_date": "2026-01-06",
                                                             "gross_bhd": 10.0, "taxable_bhd": 9.0, "warehouse_name": "Van A",
                                                             "item_name": "Cable one"}])
        p.rows["ar_ageing"], _n = prepare("ar_ageing", [{"account": "Shop A", "balance_bhd": 4.5, "as_of_date": DAY}])
        p.rows["stock_balance"], _n = prepare("stock_balance", [{"item_name": "Cable one", "warehouse_name": MAIN,
                                                                 "net_qty": 5.0, "selling_rate_bhd": 1.5,
                                                                 "total_value_bhd": 7.5, "as_of_date": DAY,
                                                                 "source_file": "f"}])
        p.as_on = {"stock_balance": DAY, "ar_ageing": DAY}
        return p

    def run(timed: bool):
        be, told, ticks = ParamBackend(), [], iter(range(100_000))
        p = parsed()
        for target, rows in p.rows.items():
            be.stage(0, target, rows)
        use = (querytime.TimedBackend(be, targets=lc.LABEL_BY_TARGET, on_report=lambda *a: told.append(a),
                                      clock=lambda: float(next(ticks)))     # each reading of the clock: one second on
               if timed else be)
        summary, exceptions = build_preview(p, use, 0, today=DAY)
        summary.pop("generated_at", None)
        return summary, exceptions, be.asked, use, told

    plain, timed = run(False), run(True)
    assert plain[0] == timed[0] and plain[1] == timed[1], "the proxy changed what the preview found"
    assert plain[2] == timed[2] and len(plain[2]) > 8, "the proxy changed what the database was asked"
    parts = timed[3].timing_summary()
    assert parts["reports"] == {t: 2.0 for t in parsed().rows}, parts      # its column lookup and its diff: a second each
    assert {"sales_comparisons", "stock_snapshot", "item_checks", "mrn_checks"} <= set(parts["helpers"]), parts
    assert parts["queries"] == len(timed[2]) and parts["seconds"] == float(len(timed[2]))
    assert timed[4] == [(t, 2.0, True) for t in parsed().rows], timed[4]   # one line per report, as each diff answers
    # every other attribute is the backend's own, and an attribute set on the proxy is set on the backend
    proxy, be = timed[3], querytime.unwrap(timed[3])
    assert proxy.mode == "fake" and proxy.stage_source == "param" and proxy.read_only is True and proxy.mem is be.mem
    proxy.extra = 7
    assert be.extra == 7 and querytime.unwrap(be) is be and not hasattr(be, "_timed_state")

    class Down(ParamBackend):
        def query(self, sql, params=None):
            raise RuntimeError("statement timeout")
    told: list = []
    down = querytime.TimedBackend(Down(), targets={"orders"}, on_report=lambda *a: told.append(a))
    err = raises(RuntimeError, lambda: down.query("select 1", ["[]", "orders"]))
    assert str(err) == "statement timeout" and [(t, ok) for t, _s, ok in told] == [("orders", False)]
    assert down.timing_summary()["queries"] == 1 and down.timing_summary()["reports"].keys() == {"orders"}
    def cannot_print(*_args):
        raise OSError("the screen is gone")
    loud = querytime.TimedBackend(ParamBackend(), targets={"orders"}, on_report=cannot_print)
    assert loud.query("select 1", ["[]", "orders"]) == [], "a progress line that fails never reaches the preview"


@test("the timing proxy around the real run_preview (parse, stage, compare, write its files): the same with it and without")
def _():
    from app.ingest_batch import run_preview
    from scripts.ops import querytime
    with data_root() as root:
        folder = root / "reports"
        folder.mkdir()
        stock_balance_file(folder / NAMES["stock_balance"])          # the layouts the real parsers read
        ledger_file(folder / NAMES["stock_movements"])
        profit_file(folder / NAMES["product_profitability"])
        receivables_file(folder / NAMES["receivables"], AGEING)

        def run(timed: bool):
            be, told, ticks = ParamBackend(), [], iter(range(100_000))
            use = (querytime.TimedBackend(be, targets=lc.LABEL_BY_TARGET, on_report=lambda *a: told.append(a),
                                          clock=lambda: float(next(ticks))) if timed else be)
            out = root / ("timed" if timed else "plain")
            res = run_preview(folder, use, actor="ops_load_check", persist=False, out_dir=out, today=DAY)
            res["summary"].pop("generated_at", None)
            written = sorted(p.name for p in out.iterdir()) if out.is_dir() else []
            return res, be.asked, use, told, written
        plain, timed = run(False), run(True)
        assert plain[0]["summary"] == timed[0]["summary"] and plain[0]["exceptions"] == timed[0]["exceptions"]
        assert plain[0]["batch_id"] is None is timed[0]["batch_id"] and plain[4] == timed[4] and plain[4], plain[4]
        assert plain[1] == timed[1] and len(plain[1]) >= 8, "the proxy changed what the database was asked"
        targets = set(plain[0]["summary"]["targets"])
        assert targets == {"stock_balance", "stock_movements", "product_profitability", "ar_ageing"}, targets
        parts = timed[2].timing_summary()
        assert parts["reports"] == {t: 2.0 for t in targets} and parts["queries"] == len(timed[1]), parts
        assert [t for t, _s, _ok in timed[3]] == list(plain[0]["summary"]["targets"]), timed[3]   # one line per report


@test("the check's preview goes through the timing proxy: the plan is the same with it and without, its seconds are kept")
def _():
    def plan_of(timed: bool):
        with data_root() as root:
            make_drop(root)
            old = os.environ.pop("YQ_CHECK_NO_TIMING", None)
            try:
                if not timed:
                    os.environ["YQ_CHECK_NO_TIMING"] = "1"
                res, deps = check("070126", deps=FakeDeps(FakeSession(backend=FakeBackend()), via_backend=True))
            finally:
                os.environ.pop("YQ_CHECK_NO_TIMING", None)
                if old is not None:
                    os.environ["YQ_CHECK_NO_TIMING"] = old
            plan = loadplan.read_plan(res["run"])
            timings = plan.pop("timings")
            for key in ("run", "created_at", "staging", "token", "token_parts", "drop"):
                plan.pop(key, None)
            return plan, timings, deps, res
    on, off = plan_of(True), plan_of(False)
    assert on[0] == off[0], "the timing proxy changed what the check found"
    assert on[2].the_session.backend.asked == off[2].the_session.backend.asked, "... or what the database was asked"
    assert on[3]["token"] and off[3]["token"] and on[3]["text"].splitlines()[-1].startswith("GO: ")
    parts = on[1]["preview_parts"]
    order = ["ar_ageing", "stock_balance", "product_profitability", "order_lines", "stock_movements", "orders"]
    assert set(parts["reports"]) == set(order) and parts["queries"] == 2 * len(order) + 1, parts
    assert list(parts["helpers"]) == ["preview"], parts               # the stand-in preview asks its one helper itself
    assert "preview_parts" not in off[1] and {"preview", "total"} <= set(on[1]) <= set(off[1]) | {"preview_parts"}
    lines = [ln for ln in on[2].lines if ln.startswith("      preview: ")]
    assert [re.sub(r" compared in [0-9]+\.[0-9] s$", "", ln) for ln in lines] == [
        "      preview: " + lc.LABEL_BY_TARGET[t] for t in order], lines
    assert not [ln for ln in off[2].lines if ln.startswith("      preview: ")]
    assert [ln[:5] for ln in on[2].steps] == [f"[{i}/9]" for i in range(1, 10)]
    step7 = next(i for i, ln in enumerate(on[2].lines) if ln.startswith("[7/9]"))
    assert all(on[2].lines.index(ln) < step7 for ln in lines), "each report's line comes as it is done, before step 7's"


DELIVERY_ON = {"mode": "on", "from": "YQ Ops <ops@example.invalid>", "to": ["owner@example.invalid"],
               "test_to": ["tester@example.invalid"]}


@test("the retry of a cut-off e-mail: `load status` offers none exactly where the sender would refuse it, in its words")
def _():
    os.environ["YQ_SECRETS_OFF"] = "1"
    from scripts import yq_send as ys                 # the sender itself, only to hold the two rules side by side
    from scripts.ops import noretry
    run, newer = "070126_0930", "080126_0930"
    old = (datetime.now(timezone.utc) - timedelta(seconds=lc.DELIVERY_WAIT_S + 60)).isoformat(timespec="seconds")
    later = (datetime.now(timezone.utc) + timedelta(minutes=5)).isoformat(timespec="seconds")
    cut = ("  Delivery: cut off before the e-mail went (the PC slept, or the load's window was closed). "
           "Nothing was reloaded. ")
    needed = f"No retry is needed: a newer load (run {newer}) has its own e-mail."
    lost = f"No retry: the last load on record is run {newer}, not this one, so this load's brief cannot be sent."
    nobody = "No retry: no load is on record as the last load, so this load's brief cannot be sent."

    def world(last="newer", closing="finished", look=(), manifest=None, delivery="on") -> None:
        loadplan.journal_append(run, "started")
        if closing == "finished":
            loadplan.journal_append(run, "finished", result="ok", needs_look=list(look))
        else:
            loadplan.journal_append(run, "failed", result="failed", step="post_gate", written=True)
        loadplan.delivery_append(run, "queued", at=old, text="Preparing the brief and the e-mail.")
        finished_at = {"newer": later, "lost": "2026-01-01T09:00:00+03:00", "this": later}.get(last)
        if last:
            status.write_state({"last_load": {"run": run if last == "this" else newer, "finished_at": finished_at}})
        f = paths.ops_dir() / "delivery.json"
        if delivery == "broken":
            paths.write_text(f, "{not json")
        elif delivery == "bad_address":
            paths.write_text(f, json.dumps(dict(DELIVERY_ON, mode="off", to=["not an address"])))
        elif delivery is not None:
            paths.write_text(f, json.dumps(dict(DELIVERY_ON, mode=delivery)))
        if manifest is not None:
            paths.write_text(paths.exports_dir() / DAY / "delivery" / run / "manifest.json",
                             "not json" if manifest == "broken" else json.dumps({"kind": manifest}))

    cases = (                                         # (what, the world, the answer: None = the retry is offered)
        ("a newer load finished since", {}, needed),
        ("this load is the last on record", {"last": "this"}, None),
        ("no load on record at all", {"last": None}, nobody),
        ("the last load on record finished before this one", {"last": "lost"}, lost),
        ("delivery is off", {"delivery": "off"}, None),
        ("no delivery.json: off", {"delivery": None}, None),
        ("a delivery.json that cannot be read counts as on", {"delivery": "broken"}, needed),
        ("a delivery.json the sender refuses counts as on, its mode off or not", {"delivery": "bad_address"}, needed),
        ("test mode, the brief not made yet", {"delivery": "test"}, needed),
        ("test mode, the brief made: a test copy is never refused", {"delivery": "test", "manifest": "full"}, None),
        ("the manifest says full", {"manifest": "full"}, needed),
        ("the manifest says look", {"manifest": "look"}, None),
        ("a manifest that cannot be read", {"manifest": "broken"}, None),
        ("no manifest, the reconcile needs a look: a note, not a brief", {"look": ("reconcile",)}, None),
        ("no manifest, the gate needs a look: a note, not a brief", {"look": ("post_gate",)}, None),
        ("no manifest, only the publish needs a look: a full brief", {"look": ("publish",)}, needed),
        ("a load closed as failed: its note is never refused", {"closing": "failed"}, None),
    )
    for what, kw, answer in cases:
        with data_root():
            world(**kw)
            mine, theirs = noretry.no_retry_needed(run), ys.no_retry_needed(run)
            assert mine == theirs == answer, (what, mine, theirs)
            line = status_of("--run", run)["text"].splitlines()[-2]
            assert line == cut + (answer or f"Its retry: python -m scripts.yq_send --run {run}"), (what, line)
    # (scripts/ops keeps its own copy of the rule: no module of it imports the sender, tests/test_ops_readonly.py)


@test("`load check <ddmmyy> again`: the word 'again' after the folder is taken and ignored; any other word is refused")
def _():
    with data_root() as root:
        make_drop(root)
        res, _ = check("070126", "again")
        assert res["ok"] and res["token"] and res["run"].startswith("070126_"), res["text"]
        res, _ = check("Again")                                      # on its own: the latest folder, as with no word
        assert res["ok"] and res["token"] and res["text"].startswith("LOAD CHECK  070126  "), res["text"]
        res, _ = check("070126", "now")
        assert res["exit"] == 2 and res["token"] is None and res["run"] is None, res["text"]
        assert res["text"] == ("LOAD CHECK: STOP\n'now' is not a word the check takes after the folder (only 'again' "
                               "is). Nothing was checked."), res["text"]
        first = ("LOAD CHECK: STOP\n'again' comes after the folder: type /yq-load <ddmmyy> again, or /yq-load alone for "
                 "the latest folder. Nothing was checked.")
        for words in (("again", "again"), ("again", "070126"), ("AGAIN", "x")):
            res, _ = check(*words)
            assert res["exit"] == 2 and res["token"] is None and res["text"] == first, (words, res["text"])
        assert len(live_runs()) == 1, "a refused line checks nothing and stages nothing"
    # the entry point's own parser (words and options in any order, scripts/ops/__main__.py) takes it as well
    parser = argparse.ArgumentParser()
    parser.add_argument("--data-root")
    parser.add_argument("--json", action="store_true")
    lc.add_arguments(parser)
    args = parser.parse_intermixed_args(["check", "070126", "--pricebook", "load", "again", "--json"])
    assert (args.drop, args.again, args.pricebook) == ("070126", "again", "load") and lc.drop_asked(args) == ("070126", None)
    assert lc.drop_asked(parser.parse_intermixed_args(["check"])) == (None, None)
    assert lc.drop_asked(parser.parse_intermixed_args(["check", "again"])) == (None, None)


@test("a journal that is broken or holds a line that is no row: the check passes over it and runs to its end")
def _():
    def finished(run: str, day: str, held: bool = True, as_on: dict | None = None) -> None:
        loadplan.write_plan({"run": run, "drop": "x", "data_date": day, "token": "t",
                             "verdicts": [{"kind": "product_profitability", "verdict": "hold_back" if held else "load"}],
                             "as_on": as_on or {"stock_balance": day, "receivables": day}})
        loadplan.journal_append(run, "started")
        loadplan.journal_append(run, "finished", result="ok")
    with data_root() as root:
        finished("050126_0900", "2026-01-05", as_on={"stock_balance": "2026-01-04", "receivables": "2026-01-05"})
        finished("030126_0900", "2026-01-03", held=False)
        finished(status.REHEARSAL_PREFIX + "0900", "2026-01-06")                    # a rehearsal never counts
        loadplan.write_plan({"run": "060126_0900", "drop": "x", "data_date": "2026-01-06", "verdicts": []})  # never loaded
        for run, junk in (("040126_0900", b"\xff\xfe not text\n"), ("070126_0815", b"\x80\x81\n"),
                          ("020126_0900", b'[1, 2]\n"a line"\n7\n')):
            loadplan.write_plan({"run": run, "drop": "x", "data_date": paths.drop_date(run[:6]).isoformat(), "token": "t",
                                 "verdicts": [{"kind": "product_profitability", "verdict": "hold_back"}]})
            loadplan.journal_file(run).write_bytes(junk)
        loadplan.journal_append("020126_0900", "finished", result="ok")           # one good row after the junk
        assert loadplan.journal_read("020126_0900") == [r for r in loadplan.journal_read("020126_0900") if isinstance(r, dict)]
        loads, explained = lc.finished_held_loads()
        assert loads == {"2026-01-05": True, "2026-01-03": False, "2026-01-02": True}, loads
        assert {"2026-01-04", "2026-01-05", "2026-01-03", "2026-01-02"} <= explained, explained
        # a whole check over these journals: a report held back is counted, nothing breaks
        make_drop(root, profit=SHORT_COST)
        res, _ = check("070126", deps=FakeDeps(FakeSession(anchors=BEFORE_LOAD)))
        assert res["ok"] and res["token"] and table_of(res["text"])[shown("product_profitability")] == "HOLD BACK", res["text"]
        assert "held back for the second load in a row" in res["text"], res["text"]      # the 05-Jan load held its own


def main() -> int:
    passed = failed = 0
    for name, fn in TESTS:
        try:
            fn()
            print(f"PASS  {name}")
            passed += 1
        except Exception:  # noqa: BLE001
            print(f"FAIL  {name}")
            traceback.print_exc()
            failed += 1
        finally:
            paths.set_data_root(None)
            config.use(None)
    db.uninstall()
    print(f"\n{passed} passed, {failed} failed")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
