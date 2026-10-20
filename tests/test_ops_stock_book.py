"""Ops engine: the Stock & Order Plan workbook (scripts.ops.workbook.stock_book), on a synthetic world.

    python -m tests.test_ops_stock_book

Same lightweight runner as tests/test_ops_stock.py, and ITS stand-ins (imported, not copied): the invented stock
ledger (Book), the database stand-in that answers every read from it (World), the temporary data root and the
command (run). SUPABASE_URL, SUPABASE_KEY and DATABASE_URL are pinned to ci.invalid BEFORE any import; no
database, no network, no .env. All names and numbers are invented; every folder is a temporary one.

The world is the stock suite's (the data date D is Sun 15 Mar 2026): one item on the order list, ITEM-B (CB2),
20 on hand, one a day, so 85 x 1 - 20 = 65 to order at 2.5 = BHD 162.500 from VFAN; the catalog's upcoming
ZN9 is WEKOME's. The damage world adds six of ITEM-A found damaged in the main store (booked) and three of
ITEM-B sent from a van and never booked.

Covered:
  * the day's file: one per data date, made once ("made now", then "made at" its modified time), the next
    free " (2)" with /yq-stock file, never over an existing file, Excel's lock files and the engine's
    temporary files passed over, only exports/ written;
  * the seven sheets in order, the title on each, every formula with its cached value and none an error;
  * the Order sheet: Suggested equals the screen's quantity (a rate that is not whole too), Final a formula,
    the yellow cells, the validation, the supplier totals equal the screen's; typing a quantity, a budget, a
    minimum or a code (looked up on All items, case aside, wildcards escaped) moves Final, Value, the
    totals and the Summary, worked out again by the evaluator;
  * the settings per supplier on Notes feed Joins and Suggested; controls() is empty on a sound file and
    names each difference when a setting moves the file away from the screen;
  * All items: every row of the plan, the statuses coloured, With vans never below nought (and the Notes
    line that counts them); the plan missing: the order list under Other and the items of the ledger;
  * the review and the damage store missing, a stub, not set up or failing: one plain line on their sheet;
  * failures: XlsxWriter missing, a folder that cannot be written, a name taken while writing, a formula
    the evaluator refuses: one sentence, the screen still prints;
  * no van or person named on the Order, All items, Summary or Damage sheets (the Ageing sheet names the
    stores that sold most: the owner's own file);
  * release 2 (the rhythm, the season, goods on the way: facts set by with_r2 as scripts.ops.stock gives them,
    every order row worked out again with stock_season): no such facts, the file is release 1's; with them
    the tag v2, eight sheets, Rate basis and Carton after Comment (the release-1 headers unchanged), Suggested
    equal to stock_season's (controls() empty); Notes' settings per supplier and for any other supplier, the
    yellow lifts with their basis and stock-out rows; each window's days in each month across a month end, a
    year end, a leap February, the data date's month met again a year on and the 365-day cap; a lift, a lead
    time or a rhythm typed in moving Suggested as stock_season does; a Supplier typed over taking the "*"
    rule; whole cartons, a longer history, the rounding on half-way cases; the On the way sheet and its empty
    line; the saved file read back with openpyxl;
  * the figures stamp (Notes!B6, row hidden): the same figures keep the day's file, an order logged, a cost
    loaded later or a file with no stamp make the next free name with a line; team copies are never the day's
    file; the latest copy is the one named and compared;
  * the Ageing sheet's guarded paces, All items' Supplier from, the Summary's release-2 lines, and the
    command's own file with nothing to check.
"""
from __future__ import annotations

import calendar
import copy
import os
import re
import sys
import time
import traceback
import zipfile
from datetime import date, datetime, timedelta
from decimal import ROUND_CEILING, ROUND_HALF_UP, Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

os.environ["SUPABASE_URL"] = "https://ci.invalid"
os.environ["SUPABASE_KEY"] = "ci-key"
os.environ["DATABASE_URL"] = "postgresql://ci:ci@ci.invalid:5432/ci"

# The stock suite's stand-ins. Importing it also sets stdout to UTF-8 (so this file does not) and imports
# scripts.ops.db first, as every command does.
from tests.test_ops_stock import (  # noqa: E402
    CODES, D, LONG, NAMES, NINE, VAN, VAN_B, L, World, data_root, model_of, run, world)
from scripts.ops import fmt, paths, status, stock, stock_season  # noqa: E402
from scripts.ops.workbook import book as bk  # noqa: E402
from scripts.ops.workbook import formula as fx  # noqa: E402
from scripts.ops.workbook import stock_book as sb  # noqa: E402

stock.RAISE_PARTS = True              # a part that fails fails the test, never stands aside quietly

TESTS: list[tuple[str, object]] = []
FILE = f"YQ Stock and Order Plan {D.isoformat()}.xlsx"
SECOND = FILE.replace(".xlsx", " (2).xlsx")
MADE = datetime(2026, 3, 16, 7, 45, tzinfo=status.BAHRAIN)
YELLOW = "FFFFF2CC"
PEOPLE = ("Van A", "Van B", VAN, VAN_B)


def test(name):
    def deco(fn):
        TESTS.append((name, fn))
        return fn
    return deco


def raises(kind, fn) -> Exception:
    try:
        fn()
    except kind as exc:
        return exc
    raise AssertionError(f"expected {kind.__name__}")


def folder(root: Path) -> Path:
    return root / "exports" / "reports" / "stock"


def books(root: Path) -> list[str]:
    f = folder(root)
    return sorted(p.name for p in f.iterdir()) if f.is_dir() else []


def damage_world():
    w = world()
    w.to_damage(L("A"), date(2026, 3, 2), 6)                         # found in the main store, booked
    w.to_damage(L("B"), date(2026, 3, 1), 3, frm=VAN, booked=False)  # sent from a van, never booked
    return w


CUTS = [("HG3", NAMES["C"], "2026-02-01", "1.0", "0.8")]           # ITEM-C's dealer price cut by a fifth
_SHARED: dict = {}

# ── release 2's facts ─────────────────────────────────────────────────────────
# The release-1 tests take the model as release 1 gave it (v1: no season, no rule per supplier, no goods on the
# way); the release-2 tests set those facts themselves (with_r2), as scripts.ops.stock gives them, and work every
# order row out again with stock_season, as scripts.ops.stock does: so they hold whatever stock.build adds.
R2_KEYS = ("season", "supplier_rules", "on_the_way")
LIFTS = {1: "0.80", 2: "0.90", 3: "1.00", 4: "1.15", 5: "1.30", 6: "1.05", 7: "0.95", 8: "1.10", 9: "1.20",
         10: "1.25", 11: "1.40", 12: "1.60"}
Q2 = Decimal("0.01")


def v1(m):
    """The model as release 1 gave it."""
    for k in R2_KEYS:
        m.facts.pop(k, None)
    return m


def v1_copy(model):
    """A copy of the model as release 1 gave it (the screen keeps its own facts)."""
    out = copy.copy(model)
    out.facts = {k: v for k, v in model.facts.items() if k not in R2_KEYS}
    return out


def lifts_of(values: dict | None = None) -> dict:
    """The lift rows as stock_season.measure_lifts gives them on D (15 Mar 2026): April to December measured on
    2025, January to March on 2026 (March up to D: 15 days)."""
    out = {}
    for m, v in (values or LIFTS).items():
        year = 2026 if m <= D.month else 2025
        out[m] = {"lift": Decimal(v), "basis": "last year", "year": year,
                  "days": D.day if m == D.month else calendar.monthrange(year, m)[1]}
    return out


def rule(lead: int, every: int, basis: str = "measured", low: int = 30, high: int = 45) -> dict:
    return {"lead": lead, "lead_basis": basis, "every": every, "low": low, "high": high,
            "trig_days": lead + every + low, "fill_days": lead + every + high}


RULES = {"VFAN": rule(40, 30), "*": rule(40, 30, "overall")}           # VFAN joins at 100 days, fills to 115


def way(code: str, supplier: str, on: date, qty, received=0, late=False, matched=True) -> dict:
    """An on-the-way line as stock_orders.fifo gives it."""
    key = next((k for k, c in CODES.items() if c == code), None)
    qty, received = Decimal(qty), Decimal(received)
    return {"ref": f"{supplier.lower()}:{on.isoformat()}:{code}", "supplier": supplier, "code": code,
            "item_name": NAMES[key] if matched and key else None, "ledger_name": L(key) if matched and key else None,
            "date": on, "qty": qty, "received": received, "due": qty - received,
            "expected_by": on + timedelta(days=40), "late": late, "future": on > D}


WAY = [way("CB2", "VFAN", date(2026, 3, 1), 30, 10),                  # 20 still due: on the way for CB2
       way("CA1", "VFAN", date(2026, 2, 1), 40, 40),                  # received in full: not shown
       way("HG3", "VFAN", date(2026, 3, 20), 5),                      # dated after D
       way("ZN9", "WEKOME", date(2026, 1, 5), 12, late=True, matched=False)]   # no item of the main store


def with_r2(m, lifts: dict | None = None, rules: dict | None = None, lines=(), warn=(12,), packs: dict | None = None,
            longer: dict | None = None, leads: dict | None = None):
    """Release 2's facts on a model of model_of(): the season, each supplier's rule (leads: facts["leads"], the
    measured or configured lead time of a supplier the rules do not name), the goods on the way (each matched
    line's units due on its item, as stock.build sets them), and every order row worked out again by
    stock_season: joins, the quantity, whole cartons where packs names the code; longer = {code: (rate, first day)}
    for a longer history. A row that no longer joins leaves the list, as on the screen."""
    v1(m)
    lifts = lifts_of() if lifts is None else lifts
    rules = dict(RULES if rules is None else rules)
    m.facts["season"] = {"lifts": lifts, "warn": list(warn), "applied": any(r["lift"] != 1 for r in lifts.values())}
    m.facts["supplier_rules"] = rules
    m.facts["leads"] = dict(leads or {})
    m.facts["trigger_days"], m.facts["fill_days"] = rules["*"]["trig_days"], rules["*"]["fill_days"]
    due: dict = {}
    for ln in lines:
        if ln["ledger_name"] and ln["due"] > 0:
            due[ln["ledger_name"]] = due.get(ln["ledger_name"], Decimal(0)) + ln["due"]
    m.facts["on_the_way"] = {"lines": list(lines), "on_order": due, "logged": len(lines),
                             "counted": sum(1 for ln in lines if ln["ledger_name"] and ln["due"] > 0),
                             "unmatched": sum(1 for ln in lines if not ln["ledger_name"]), "log": bool(lines)}
    for it in m.items:
        it.on_order = due.get(it.ledger_name, Decimal(0))
    plan = m.facts["plan"]
    supplier_of = {x["item"].ledger_name: x["supplier"] for x in plan["order"]}

    def again(x: dict) -> dict | None:
        it = x["item"]
        r = rules.get(x.get("supplier") or supplier_of.get(it.ledger_name)) or rules["*"]
        rate6, basis = x["rate6"], "90 days"
        if longer and it.code in longer:
            rate6, basis = stock.stock_plan.round6(longer[it.code][0]), f"longer history from {longer[it.code][1]}"
        position = max(it.on_hand, Decimal(0)) + it.on_order
        if not stock_season.joins(position, rate6, m.d, r["trig_days"], lifts):
            return None
        qty = stock_season.suggested(position, rate6, m.d, r["fill_days"], lifts)
        pack = (packs or {}).get(it.code)
        if pack and pack > 1 and qty > 0:
            qty = (qty / pack).to_integral_value(rounding=ROUND_CEILING) * pack
        return dict(x, qty=qty, rate6=rate6, rate_basis=basis, pack=pack,
                    cost_bhd=fmt.q3(qty * it.cost) if it.cost is not None else None)

    m.facts["order"] = [y for y in map(again, m.facts["order"]) if y is not None]
    plan["order"] = [y for y in map(again, plan["order"]) if y is not None]
    return m


def built2(w: World | None = None, **r2):
    """(model, book): a release-2 workbook in memory (with_r2's facts), not evaluated."""
    with data_root():
        m, _inp = model_of(w)
        with_r2(m, **r2)
        return m, sb.build_book(m, MADE)


def month_count(day: date, n: int) -> list[int]:
    """The days of each calendar month (January first) among day + 1 ... day + n, counted one day at a time."""
    out = [0] * 12
    for k in range(1, n + 1):
        out[(day + timedelta(days=k)).month - 1] += 1
    return out


def shared() -> dict:
    """The damage world's day (with ITEM-C's price cut), run ONCE for the tests that only read its file (opening a
    workbook is the slow part): the screen, the files the run wrote (relative to the data root), the model, the
    file re-opened twice (its formulas, its cached values), its sheets' XML and its shared strings. The file is
    release 1's (the model's release-2 facts set aside for it): the release-2 file has tests of its own."""
    if not _SHARED:
        saved = sb.day_file
        sb.day_file = lambda model, force=False: saved(v1_copy(model), force=force)
        try:
            with data_root() as root:
                before = {p for p in root.rglob("*") if p.is_file()}
                out = run(World(book=damage_world(), cuts=CUTS))
                m, _inp = model_of(World(book=damage_world(), cuts=CUTS))
                after = {p.relative_to(root) for p in root.rglob("*") if p.is_file() and p not in before}
                path = Path(out["flags"]["file"])
                assert path == folder(root) / FILE, path
                wb, vb = opened(path)
                with zipfile.ZipFile(path) as z:
                    xml = [z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("xl/worksheets/sheet")]
                    strings = z.read("xl/sharedStrings.xml").decode("utf-8")
        finally:
            sb.day_file = saved
        _SHARED.update(out=out, model=v1(m), after=after, wb=wb, vb=vb, xml=xml, strings=strings)
    return _SHARED


def built(w: World | None = None):
    """(model, book): the release-1 workbook of a world in memory, not evaluated (inside a temporary data root)."""
    with data_root():
        m, _inp = model_of(w)
        return v1(m), sb.build_book(v1(m), MADE)


def first_line(out: dict) -> str:
    return out["text"].splitlines()[0]


def value(book, sheet: str, r: int, c: int):
    cell = book.get(sheet).cells.get((r, c))
    return None if cell is None else cell.value


def row_of(book, sheet: str, c: int, text: str) -> int:
    """The 0-based row of the first cell of column c holding `text`."""
    hits = sorted(r for (r, cc), cell in book.get(sheet).cells.items() if cc == c and cell.value == text)
    assert hits, f"{sheet}: no {text!r} in column {fx.col_letters(c)}"
    return hits[0]


def opened(path: Path):
    import openpyxl
    return openpyxl.load_workbook(path), openpyxl.load_workbook(path, data_only=True)


def cell_rows(ws, col: str, first: int):
    """Excel rows from `first` down while column `col` holds something."""
    r = first
    while ws[f"{col}{r}"].value not in (None, ""):
        yield r
        r += 1


# ── the day's file ────────────────────────────────────────────────────────────

@test("the day's file: one per data date, made now; a second run makes none (made at); /yq-stock file makes (2)")
def _():
    with data_root() as root:
        out = run(World())
        line = first_line(out)
        assert line.startswith("FILE  exports") and line.endswith(f"{FILE} (made now)"), line
        assert books(root) == [FILE]
        path = folder(root) / FILE
        assert out["flags"]["file"] == str(path) and path.is_file()
        data = path.read_bytes()
        again = run(World())
        line = first_line(again)
        assert re.search(rf"{re.escape(FILE)} \(made at \d\d:\d\d; /yq-stock file makes a new one\)$", line), line
        assert books(root) == [FILE] and path.read_bytes() == data
        forced = run(World(), "file")
        assert first_line(forced).endswith(f"{SECOND} (made now)"), first_line(forced)
        assert books(root) == sorted([FILE, SECOND]) and path.read_bytes() == data   # never overwritten
        third = run(World())
        assert f"{SECOND} (made at " in first_line(third), first_line(third)            # the newest is named
        for spec in (("reorder",), ("ageing",), ("damage",), ("CB2",)):                # only the headline makes one
            run(World(), *spec)
        assert books(root) == sorted([FILE, SECOND])


@test("day_file: its answer, made_at the file's own time, lock and temporary files passed over, the next free name")
def _():
    with data_root() as root:
        m, _inp = model_of()
        f = folder(root)
        f.mkdir(parents=True)
        for name in (f"~${FILE}", "~build-1x2y.xlsx", "YQ Stock and Order Plan 2026-03-14.xlsx"):
            (f / name).write_bytes(b"not a workbook")
        res = sb.day_file(m)
        assert set(res) == {"path", "made", "made_at", "error", "controls"}, res          # no "built" key
        assert res["made"] and res["error"] is None and res["controls"] == [] and Path(res["path"]) == f / FILE
        assert re.fullmatch(r"\d\d:\d\d", res["made_at"]), res
        made_at = res["made_at"]
        stamp = MADE.timestamp()
        os.utime(f / FILE, (stamp, stamp))               # he saves his edits: the file's modified time moves ...
        res = sb.day_file(m)
        # ... and "made at" stays the build's own time where the system keeps a creation time (Windows)
        expect = made_at if (os.name == "nt" or hasattr(os.stat(f / FILE), "st_birthtime")) else "07:45"
        assert (res["made"], res["made_at"], Path(res["path"]).name) == (False, expect, FILE), res
        res = sb.day_file(m, force=True)
        assert res["made"] and Path(res["path"]).name == SECOND
        assert sb.target_path(f, sb.STEM.format(d=D.isoformat())).name == FILE.replace(".xlsx", " (3).xlsx")
        assert sorted(p.name for p in sb.existing(f, sb.STEM.format(d=D.isoformat()))) == sorted([FILE, SECOND])


@test("a difference controls() finds is printed under the FILE line, one 'To check in the file' line each")
def _():
    saved = sb.controls
    sb.controls = lambda book, model: ["Order sheet: CB2 suggests 66 in the file, 65 on the screen"]
    try:
        with data_root():
            out = run(World())
    finally:
        sb.controls = saved
    lines = out["text"].splitlines()
    assert lines[0].endswith(f"{FILE} (made now)") and lines[2] == "", lines[:3]
    assert lines[1] == "  To check in the file: Order sheet: CB2 suggests 66 in the file, 65 on the screen", lines[1]


@test("an empty order list: the rows to add an item under the header, the totals nought, one line says so")
def _():
    with data_root({"stock.display_items": ["DS1", "CB2"]}):                  # CB2 a display item: nothing to order
        m, _inp = model_of()
        book = sb.build_book(v1(m), MADE)
    assert m.facts["order"] == [] and sb.controls(book, m) == []
    pos = book.controls["order"]
    assert pos["rows"] == [] and pos["add"] == (pos["head"] + 1, pos["head"] + sb.ADD_ROWS)
    assert [value(book, "Order", pos["total"], c) for c in (sb.B_ITEMS, sb.B_UNITS, sb.B_VALUE)] == [0, 0, 0]
    notes = [value(book, "Order", r, 0) for r in range(pos["add"][1] + 2, pos["add"][1] + 4)]
    assert notes[0] == "No item is below 70 days of cover today.", notes
    r = row_of(book, "Summary", 0, "As the plan suggests")
    assert value(book, "Summary", r, 1) == "0 items, 0 units, BHD 0.000"


@test("the file: seven sheets in order, the title on each, every formula with its cached value and none an error")
def _():
    m, book = built(World(book=damage_world(), cuts=CUTS))
    assert [ws.name for ws in book.sheets] == list(sb.SHEETS)
    assert book.evaluate() == [] and sb.controls(book, m) == []
    s = shared()
    wb, sheets, strings = s["wb"], s["xml"], s["strings"]
    assert "To check in the file" not in s["out"]["text"] and first_line(s["out"]).endswith(f"{FILE} (made now)")
    assert wb.sheetnames == ["Summary", "Order", "All items", "Fast movers", "Ageing", "Damage", "Notes"]
    for ws in wb:
        assert ws["A1"].value == "YQ Bahrain — Stock & Order Plan", ws.title
        assert ws["A2"].value.startswith("Data date Sun 15 Mar 2026 · ledger to Sun 15 Mar 2026"), ws["A2"].value
        assert ws["A3"].value, ws.title                                              # its purpose in one line
    formulas = 0
    for xml in sheets:
        for cell in re.findall(r"<c [^>]*>(.*?)</c>", xml):
            if "<f>" in cell:
                formulas += 1
                assert "<v>" in cell, cell[:120]                                     # its value goes with it
                assert not re.search(r"<v>#(N/A|VALUE!|REF!|DIV/0!|NUM!)</v>", cell), cell[:120]
    assert formulas > 100
    assert not NINE.search(strings) and not LONG.search(strings)


# ── the Order sheet ───────────────────────────────────────────────────────────

@test("Order: CB2 suggests 65 as the screen does, Final is a formula, the yellow cells, the validation, the totals")
def _():
    s = shared()
    out, m = s["out"], s["model"]
    ws, vs = s["wb"]["Order"], s["vb"]["Order"]
    head = next(r for r in range(1, 40) if ws[f"A{r}"].value == "Supplier" and ws[f"B{r}"].value == "Code")
    assert ws.freeze_panes == f"A{head + 1}" and ws.auto_filter.ref == f"A{head}:T{head + 1 + sb.ADD_ROWS}"
    cb2 = head + 1
    assert (ws[f"B{cb2}"].value, ws[f"C{cb2}"].value, ws[f"A{cb2}"].value) == (CODES["B"], NAMES["B"], "VFAN")
    assert ws[f"D{cb2}"].value == "Chargers" and ws[f"F{cb2}"].value == 20 and ws[f"J{cb2}"].value == 1
    assert ws[f"G{cb2}"].value == 32 and ws[f"I{cb2}"].value == 45 and ws[f"L{cb2}"].value == 20   # 35 - 3 damaged
    assert vs[f"N{cb2}"].value == 65 and vs[f"P{cb2}"].value == 65 and vs[f"R{cb2}"].value == 162.5
    assert vs[f"M{cb2}"].value is True and ws[f"Q{cb2}"].value == 2.5
    assert ws[f"P{cb2}"].value == f'=IF(O{cb2}="",N{cb2},O{cb2})'
    assert ws[f"R{cb2}"].value == f'=IF(Q{cb2}="","",ROUND(P{cb2}*Q{cb2},3))'
    assert ws[f"M{cb2}"].value.startswith(f"=AND(J{cb2}>0,ROUND(MAX(F{cb2},0)+H{cb2}-J{cb2}*INDEX(Notes!$F$9:$F$10,"
                                          f'MATCH(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE($A{cb2},"~","~~")'), ws[f"M{cb2}"].value
    assert ws[f"N{cb2}"].value.startswith(f"=IF(M{cb2},ROUNDUP(ROUND(J{cb2}*INDEX(Notes!$G$9:$G$10,")
    for col in ("O", "T"):                                                   # Your qty and Comment: yellow, blank
        assert ws[f"{col}{cb2}"].fill.fgColor.rgb == YELLOW and ws[f"{col}{cb2}"].value is None, col
    assert ws[f"P{cb2}"].fill.fgColor.rgb != YELLOW
    add = cb2 + 1                                                            # the first row to add an item
    assert ws[f"B{add}"].fill.fgColor.rgb == YELLOW and ws[f"N{add}"].value == 0
    assert ws[f"C{add}"].value.startswith(f'=IF($B{add}="","",IFERROR(INDEX(\'All items\'!$C$6:$C$12,')
    dv = {str(v.sqref): v for v in ws.data_validations.dataValidation}
    yours = dv[f"O{cb2}:O{cb2 + sb.ADD_ROWS}"]
    assert (yours.type, yours.operator, yours.formula1) == ("whole", "greaterThanOrEqual", "0")
    assert yours.prompt == "Leave blank to take the suggested quantity." and yours.showInputMessage
    # By supplier: VFAN's line equals the plan's totals (the screen's), WEKOME has nothing on the list
    plan = m.facts["plan"]["supplier_totals"]
    assert [(t["supplier"], t["items"], t["units"], t["value"]) for t in plan] == [("VFAN", 1, 65, 162.5)]
    assert [ws[f"{c}5"].value for c in "CDEFGHIJ"] == ["Supplier", "Items", "Units", "Value, BHD", "Budget, BHD",
                                                       "Minimum order, BHD", "Against budget", "Minimum check"]
    lines = {ws[f"C{r}"].value: (vs[f"D{r}"].value, vs[f"E{r}"].value, vs[f"F{r}"].value) for r in range(6, head)
             if ws[f"C{r}"].value}
    assert lines == {"VFAN": (1, 65, 162.5), "WEKOME": (0, 0, 0), "Total": (1, 65, 162.5)}, lines
    assert ws["E6"].value == '=SUMIFS($P$12:$P$27,$A$12:$A$27,"="&SUBSTITUTE(SUBSTITUTE(SUBSTITUTE($C6,"~","~~"),' \
                             '"*","~*"),"?","~?"))', ws["E6"].value
    assert ws["G6"].fill.fgColor.rgb == YELLOW and ws["H6"].fill.fgColor.rgb == YELLOW      # budget, minimum
    note = " ".join(str(ws[f"A{r}"].value) for r in range(head + sb.ADD_ROWS + 2, head + sb.ADD_ROWS + 9)
                    if ws[f"A{r}"].value)
    for words in ("No purchase order is drafted.", "Release 2 adds the order rhythm, the season and goods on order "
                  "when you say an order is placed.", "goods on order are not deducted unless logged"):
        assert words in note, words
    assert "To check in the file" not in out["text"]


@test("Order: typing a quantity moves Final, Value, the supplier's totals and the Summary (the evaluator)")
def _():
    m, book = built()
    pos = book.controls["order"]
    r, vf, tot = pos["rows"][0], pos["suppliers"]["vfan"], pos["total"]
    ws = book.get("Order")
    ws.put(r, sb.O_YOURS, 100, "input_qty")
    assert book.evaluate() == []
    assert value(book, "Order", r, sb.O_FINAL) == 100 and value(book, "Order", r, sb.O_VALUE) == 250
    assert value(book, "Order", r, sb.O_SUGGESTED) == 65                         # the suggestion stays beside it
    assert [value(book, "Order", vf, c) for c in (sb.B_ITEMS, sb.B_UNITS, sb.B_VALUE)] == [1, 100, 250]
    assert [value(book, "Order", tot, c) for c in (sb.B_ITEMS, sb.B_UNITS, sb.B_VALUE)] == [1, 100, 250]
    summary = book.get("Summary")
    units = next(cell for cell in summary.cells.values() if cell.formula == f"Order!E{tot + 1}")
    vfan = next(cell for cell in summary.cells.values() if cell.formula == f"Order!F{vf + 1}")
    assert units.value == 100 and vfan.value == 250
    # a nought takes the item off the order: no item, no unit, no value
    m, book = built()
    book.get("Order").put(r, sb.O_YOURS, 0, "input_qty")
    assert book.evaluate() == []
    assert [value(book, "Order", vf, c) for c in (sb.B_ITEMS, sb.B_UNITS, sb.B_VALUE)] == [0, 0, 0]
    # a budget and a minimum order: against the budget, and the check
    m, book = built()
    ws = book.get("Order")
    ws.put(vf, sb.B_BUDGET, 100, "input_money3")
    ws.put(vf, sb.B_MIN, 200, "input_money3")
    assert book.evaluate() == []
    assert value(book, "Order", vf, sb.B_AGAINST) == 62.5 and value(book, "Order", vf, sb.B_CHECK) == "below the minimum"
    assert value(book, "Order", pos["suppliers"]["wekome"], sb.B_AGAINST) == ""   # no budget typed: nothing said


@test("Order: a code typed in an add row is looked up on All items (case aside, wildcards escaped); the totals count it")
def _():
    m, book = built()
    pos = book.controls["order"]
    a = pos["add"][0]
    ws = book.get("Order")
    typed = (("zn9", 12), ("NOPE-1", 5), ("CA1", 10), ("C*", 3), ("", 7))
    for k, (code, qty) in enumerate(typed):
        if code:
            ws.text(a + k, sb.O_CODE, code, "input_note")
        ws.put(a + k, sb.O_YOURS, qty, "input_qty")
    assert book.evaluate() == []
    got = [tuple(value(book, "Order", a + k, c) for c in (sb.O_SUPPLIER, sb.O_ITEM, sb.O_CATEGORY, sb.O_MAIN,
                                                            sb.O_COST, sb.O_FINAL, sb.O_VALUE)) for k in range(5)]
    assert got[0] == ("WEKOME", "ZZ-NEW Speaker Upcoming", "Speakers", 0, None, 12, ""), got[0]  # no cost yet
    assert got[1][:2] == ("", "code not found") and got[1][5:] == (5, ""), got[1]
    assert got[2] == ("VFAN", NAMES["A"], "Cables", 394, 0.4, 10, 4.0), got[2]
    assert got[3][:2] == ("", "code not found"), got[3]                  # "C*" is no wildcard: CA1 and CB2 stay out
    assert got[4][:2] == ("", "") and got[4][5] == 7, got[4]             # no code: nothing looked up
    sup = pos["suppliers"]
    assert [value(book, "Order", sup["vfan"], c) for c in (sb.B_ITEMS, sb.B_UNITS, sb.B_VALUE)] == [2, 75, 166.5]
    assert [value(book, "Order", sup["wekome"], c) for c in (sb.B_ITEMS, sb.B_UNITS, sb.B_VALUE)] == [1, 12, 0]
    assert value(book, "Order", pos["total"], sb.B_UNITS) == 87          # a code not found counts nowhere


@test("Order: a rate that is not whole (46/45 a day) gives the screen's 67, to the unit, and controls() is empty")
def _():
    w = world()
    w.sale(L("B"), D, 1)
    m, book = built(World(book=w))
    x = m.facts["order"][0]
    assert x["qty"] == 67 and x["rate6"] == stock.stock_plan.round6(x["rate"].rate)
    assert sb.controls(book, m) == []
    r = book.controls["order"]["rows"][0]
    assert value(book, "Order", r, sb.O_RATE) == float(x["rate6"]) == 1.022222
    assert value(book, "Order", r, sb.O_SUGGESTED) == 67


@test("Notes: the layout tag in B5; the settings per supplier feed Joins and Suggested, and controls() says so")
def _():
    m, book = built()
    assert value(book, "Notes", 4, 1) == sb.LAYOUT_TAG == "stock-plan v1" and (sb.TAG_SHEET, sb.TAG_CELL) == ("Notes", "B5")
    s = book.controls["settings"]
    assert sorted(s) == ["vfan", "wekome"] and min(s.values()) == sb.NOTES_SET_HEAD + 1
    assert sb.controls(book, m) == []
    assert [value(book, "Notes", s["vfan"], c) for c in (sb.S_LEAD, sb.S_EVERY, sb.S_LOW, sb.S_HIGH, sb.S_JOIN,
                                                          sb.S_FILL)] == [40, 0, 30, 45, 70, 85]
    m, book = built()
    book.get("Notes").put(s["vfan"], sb.S_LEAD, 60, "input_qty")         # 60 + 30 = 90 to join, filled to 105
    checks = sb.controls(book, m)
    assert value(book, "Order", book.controls["order"]["rows"][0], sb.O_SUGGESTED) == 85
    assert checks == ["Order sheet: CB2 suggests 85 in the file, 65 on the screen",
                      "Order sheet: VFAN units 85 in the file, 65 on the screen",
                      "Order sheet: VFAN value BHD 212.500 in the file, BHD 162.500 on the screen"], checks
    m, book = built()
    book.get("Notes").put(s["vfan"], sb.S_LEAD, 0, "input_qty")          # 0 + 30 = 30: 20 on hand still joins
    sb.controls(book, m)
    assert value(book, "Order", book.controls["order"]["rows"][0], sb.O_SUGGESTED) == 25
    m, book = built()
    book.get("Notes").put(s["vfan"], sb.S_EVERY, 0, "input_qty")
    book.get("Notes").put(s["vfan"], sb.S_LEAD, 0, "input_qty")
    book.get("Order").put(book.controls["order"]["rows"][0], sb.O_MAIN, 40, "g_int")   # 40 on hand: above 30 days
    sb.controls(book, m)
    r = book.controls["order"]["rows"][0]
    assert value(book, "Order", r, sb.O_JOINS) is False and value(book, "Order", r, sb.O_SUGGESTED) == 0
    # a supplier typed over that the settings do not hold: its lookups come out as #N/A, and the checks say where
    m, book = built()
    r = book.controls["order"]["rows"][0]
    book.get("Order").text(r, sb.O_SUPPLIER, "Nobody", "g_text")
    checks = sb.controls(book, m)
    n = r + 1
    assert checks[:2] == ["Order sheet: CB2 suggests #N/A in the file, 65 on the screen",
                          "Order sheet: VFAN items 0 in the file, 1 on the screen"], checks
    assert f"Order!M{n} came out as #N/A" in checks and f"Order!N{n} came out as #N/A" in checks, checks


# ── All items, Fast movers, Ageing, Damage, Notes ─────────────────────────────

@test("All items: every row of the plan, statuses coloured, With vans never below nought (and the Notes line counts it)")
def _():
    s = shared()
    m, wb = s["model"], s["wb"]
    ws = wb["All items"]
    first = sb.ALL_HEAD + 2
    rows = list(cell_rows(ws, "C", first))
    plan = m.facts["plan"]["all_items"]
    assert len(rows) == len(plan) == 7 and [ws[f"C{r}"].value for r in rows] == [p["name"] for p in plan]
    assert ws.freeze_panes == f"A{first}" and ws.auto_filter.ref == f"A{first - 1}:V{rows[-1]}"
    by_code = {ws[f"B{r}"].value: r for r in rows}
    a = by_code["CA1"]
    assert next(p for p in plan if p["code"] == "CA1")["with_vans"] == -10                 # the van gave 10 back
    assert ws[f"G{a}"].value == 0 and ws[f"G{by_code['CB2']}"].value == 32                # read as nought
    assert ws[f"O{by_code['CB2']}"].value == 2.5 and ws[f"O{by_code['ZN9']}"].value is None   # landed cost
    assert ws[f"E{by_code['CB2']}"].value == "Order now" and ws[f"E{by_code['ZN9']}"].value == "Not in main store yet"
    rules = {(r.formula[0], r.operator) : r.dxf for cf in ws.conditional_formatting for r in cf.rules}
    assert {k[0] for k in rules} == {'"Sold out"', '"Order now"', '"No stock, no sale"', '"No sale 90 d"', '"OK"'}
    assert rules[('"Sold out"', "equal")].fill.bgColor.rgb == "FFF8D7DA"
    assert rules[('"Order now"', "equal")].fill.bgColor.rgb == "FFFFF3CD"
    assert rules[('"OK"', "equal")].fill.bgColor.rgb == "FFD4EDDA"
    assert rules[('"No sale 90 d"', "equal")].font.color.rgb == "FF6C757D"
    notes = " ".join(str(c.value) for row in wb["Notes"].iter_rows() for c in row if c.value)
    assert ("With vans is for information only (vans, outlets and goods on their way, from the stock ledger); where "
            "the ledger shows them below nought it reads 0 (1 item).") in notes
    assert wb["Order"]["G12"].value == 32                                                   # CB2 with the vans


@test("Fast movers, Ageing, Damage: their tables, the review's sellers (the owner's file), the decision list")
def _():
    wb = shared()["wb"]
    fast = wb["Fast movers"]
    assert [fast[f"B{r}"].value for r in (7, 8)] == ["CA1", "FO5"] and fast["I7"].value == 15 and fast["I8"].value == 0
    assert fast["I6"].value == "Days available (of 15)" and fast["J7"].value is None
    ageing = wb["Ageing"]
    rows = {ageing[f"B{r}"].value: r for r in cell_rows(ageing, "A", 7)}
    assert set(rows) == {"HG3", "CA1"}, rows
    c, a = rows["HG3"], rows["CA1"]
    assert ageing[f"C{c}"].value == "Not moving" and ageing[f"M{c}"].value == -0.2 and ageing[f"R{c}"].value == 4
    assert ageing[f"S{a}"].value == "Push sales team: 30 a week" and ageing[f"T{a}"].value == "Van A (180)"
    assert ageing[f"U{a}"].fill.fgColor.rgb == YELLOW
    words = [str(x.value) for row in ageing.iter_rows() for x in row if x.value]
    assert "The only lever is the sales team: no price cut, no bundle (your rule)." in words
    assert "Cut of Sun 01 Feb 2026" in words and "Now (Sun 15 Mar 2026)" in words
    for word in ("mark it down", "markdown", "cut the price", "lower the price"):
        assert not any(word in w.lower() for w in words), word
    damage = wb["Damage"]
    labels = {damage[f"A{r}"].value: damage[f"B{r}"].value for r in range(5, 11)}
    assert labels["Units in the damage store"] == 6 and labels["Sent and not booked in (7 days on), units"] == 3
    assert labels["Sent this month per 100 sold"] == 18.37 and labels["Trend"] == "none yet"
    dv = [v for v in damage.data_validations.dataValidation if v.type == "list"]
    assert len(dv) == 1 and dv[0].formula1 == '"Claim from supplier,Write off,Repair,Sell as B-grade,Keep"'
    head = next(r for r in range(1, 60) if damage[f"H{r}"].value == "Decision")
    assert str(dv[0].sqref) == f"H{head + 1}:H{head + 2}" and damage[f"H{head + 1}"].fill.fgColor.rgb == YELLOW
    assert [damage[f"B{r}"].value for r in (head + 1, head + 2)] == ["CA1", "CB2"]
    months = [damage[f"A{r}"].value for r in range(14, 26)]
    assert months[0] == "Apr 2025" and months[-1] == "Mar 2026 (part)"
    texts = [str(x.value) for row in damage.iter_rows() for x in row if x.value]
    assert any(t.startswith("Counted when sent") for t in texts) and any(t.startswith("Nothing here is booked") for t in texts)


@test("no van or person is named on the Order, All items, Summary or Damage sheets; only exports/ is written")
def _():
    s = shared()
    out, after, wb = s["out"], s["after"], s["wb"]
    assert out["ok"] and Path("exports", "reports", "stock", FILE) in after
    assert all(p.parts[0] == "exports" for p in after), sorted(str(p) for p in after)
    assert not [p for p in after if p.name.startswith("~")]                      # no temporary file left
    for name in ("Order", "All items", "Summary", "Damage"):
        for row in wb[name].iter_rows():
            for c in row:
                text = str(c.value or "")
                assert not any(who in text for who in PEOPLE), (name, c.coordinate, text)


@test("Summary: the file, the stock, cover, fast movers, the order list and suppliers (formulas), ageing, damage")
def _():
    m, book = built(World(book=damage_world()))
    assert book.evaluate() == []
    ws = book.get("Summary")

    def at(label: str):
        return value(book, "Summary", row_of(book, "Summary", 0, label), 1)

    assert at("Data date (the stock ledger's last day)") == fx.serial(D)
    assert at("Made at (Bahrain)") == "Mon 16 Mar 2026, 07:45" and at("Layout") == "stock-plan v1"
    h = m.facts["headline"]
    assert (at("Items (the portal's figure)"), at("Units (the portal's figure)")) == (5, h["units"])
    assert at("At landed cost, BHD") == float(h["cost_bhd"]) and at("The portal at cost, BHD") == 381.9
    assert at("Aggregate cover, days") == 74 and at("This month so far") == 0.5
    assert (at("Items"), at("Units"), at("Value at landed cost, BHD")) == (1, 65, 162.5)
    assert at("As the plan suggests") == "1 item, 65 units, BHD 162.500"
    assert at("Lead time, days") == 40
    r = row_of(book, "Summary", 0, "Lead time, days")
    assert value(book, "Summary", r, 2).startswith("Measured on the last 5 supplier invoices")
    assert (at("Joins below, days of sales"), at("Filled to, days of sales")) == (70, 85)
    vfan = row_of(book, "Summary", 0, "VFAN")
    assert ws.cells[(vfan, 2)].formula == f"Order!E{book.controls['order']['suppliers']['vfan'] + 1}"
    assert [value(book, "Summary", vfan, c) for c in (1, 2, 3)] == [1, 65, 162.5]
    assert (at("Review: Not moving"), at("Review: Slow"), at("Review: On track")) == (1, 1, 0)
    assert at("Push a week, units (the sales team)") == 34                     # HG3 4 (50 / 90 x 7) + CA1 30
    assert at("Units in the damage store") == 6 and at("Trend") == "none yet"


@test("a part missing, a stub, not set up or failing is one line on its sheet; the plan missing puts all under Other")
def _():
    with data_root():
        m, _inp = model_of()
        v1(m)
        m.facts["plan"] = {"error": "KeyError"}
        m.facts["review"] = {"clear_days": 90, "rows": [], "built": False}
        m.facts["damage"] = {"set_up": False, "store": None, "found": False}
        book = sb.build_book(m, MADE)
    assert sb.controls(book, m) == []
    pos = book.controls["order"]
    assert list(pos["suppliers"]) == ["other"] and value(book, "Order", pos["rows"][0], sb.O_SUPPLIER) == "Other"
    assert value(book, "Order", pos["rows"][0], sb.O_SUGGESTED) == 65
    first, last = book.controls["all_items"]
    assert last - first + 1 == len(m.items) == 6
    assert value(book, "All items", 3, 0).startswith("The order plan was not worked out this time (KeyError)")
    assert value(book, "Ageing", 4, 0) == "The ageing review was not worked out this time."
    assert value(book, "Damage", 4, 0).startswith("The damage store is not set up yet: config 'warehouses.damage'")
    assert not [k for k in book.get("Damage").cells if k[0] > 5]                 # nothing else on the sheet
    with data_root():
        m, _inp = model_of()
        v1(m)
        m.facts["plan"] = {"built": False}
        m.facts["damage"] = {"error": "ValueError"}
        book = sb.build_book(m, MADE)
    assert sb.controls(book, m) == []
    assert book.controls["all_items"] == (sb.ALL_HEAD + 1, sb.ALL_HEAD + 6)
    assert value(book, "Damage", 4, 0) == "The damage figures were not worked out this time (ValueError)."
    # facts that are not what the contract says: one line on that sheet, the rest of the file stands
    # (stock.RAISE_PARTS, as in every test, makes it raise instead, so a test sees it)
    with data_root():
        m, _inp = model_of()
    v1(m)
    m.facts["review"] = dict(m.facts["review"], rows=[{"bogus": 1}])
    with data_root():
        raises(KeyError, lambda: sb.build_book(m, MADE))
        stock.RAISE_PARTS = False
        try:
            book = sb.build_book(m, MADE)
        finally:
            stock.RAISE_PARTS = True
    assert value(book, "Ageing", 4, 0) == "The ageing review could not be put in this file this time (KeyError)."
    assert not [k for k in book.get("Ageing").cells if k[0] > 4] and sb.controls(book, m) == []


# ── failures ──────────────────────────────────────────────────────────────────

@test("XlsxWriter missing: one sentence; the screen still prints, with FILE  not made this time")
def _():
    had, saved = "xlsxwriter" in sys.modules, sys.modules.get("xlsxwriter")
    sys.modules["xlsxwriter"] = None                     # `import xlsxwriter` now fails, as on a PC without it
    try:
        with data_root() as root:
            out = run(World())
            files = books(root)
    finally:
        if had:
            sys.modules["xlsxwriter"] = saved
        else:
            del sys.modules["xlsxwriter"]
    assert first_line(out) == "FILE  not made this time: XlsxWriter is not installed on this PC (pip install xlsxwriter)."
    assert out["ok"] and "STOCK  main store, Main Store" in out["text"] and "ORDER LIST" in out["text"]
    assert files == [] and "file" not in out["flags"]


@test("a folder that cannot be written, a name taken meanwhile, a formula refused: one sentence each, nothing left")
def _():
    with data_root() as root:
        f = folder(root)
        f.parent.mkdir(parents=True)
        f.write_text("a file where the folder should be", encoding="utf-8")
        out = run(World())
        line = first_line(out)
        assert line.startswith("FILE  not made this time: it could not be written in exports/reports/stock "
                               "(FileExistsError"), line
        assert out["ok"] and "ORDER LIST" in out["text"]
    saved_save, saved_eval = bk.Book.save, bk.Book.evaluate

    def locked(self, target):
        raise PermissionError(13, "Permission denied")

    def taken(self, target):
        Path(target).parent.mkdir(parents=True, exist_ok=True)
        Path(target).write_bytes(b"someone else's file")
        raise FileExistsError(f"{target} appeared while the report was written: not replaced")

    def refused(self):
        raise fx.FormulaError("Order!M12 =NOPE(1): the function NOPE is not implemented here")

    try:
        for patch, words in ((("save", locked), "it could not be written in exports/reports/stock (PermissionError: "
                                                "Permission denied)"),
                             (("save", taken), "a file of that name appeared while this one was written; /yq-stock "
                                               "file makes the next one"),
                             (("evaluate", refused), "a formula could not be worked out (Order!M12 =NOPE(1): the "
                                                     "function NOPE is not implemented here)")):
            setattr(bk.Book, *patch)
            with data_root() as root:
                with_screen = run(World())
                res = sb.day_file(model_of()[0], force=True)       # the run's "taken" file is there: a new name
                left = [p.name for p in folder(root).iterdir()] if folder(root).is_dir() else []
            bk.Book.save, bk.Book.evaluate = saved_save, saved_eval
            assert first_line(with_screen) == f"FILE  not made this time: {words}.", first_line(with_screen)
            assert res == {"path": None, "made": False, "made_at": None, "error": words, "controls": []}, res
            assert not [n for n in left if n.startswith("~")] and with_screen["ok"]
    finally:
        bk.Book.save, bk.Book.evaluate = saved_save, saved_eval
    assert paths.allowed_output(ROOT / "exports" / "x.xlsx") is not None    # the guard itself is paths'


# ── release 2: the rhythm, the season, goods on the way ───────────────────────

def order_head(book) -> list:
    pos = book.controls["order"]
    return [value(book, "Order", pos["head"], c) for c in range(0, sb.O_CARTON + 1)]


R1_HEADERS = ["Supplier", "Code", "Item", "Category", "Fast", "Main stock", "With vans", "On the way", "Sold 90 d",
              "A day", "*", "Cover days", "Joins", "Suggested", "Your qty", "Final", "Landed cost", "Value, BHD",
              "Last receipt", "Comment"]


@test("no release-2 facts: the file is release 1's (tag v1, seven sheets, Joins on Join below, Suggested 65)")
def _():
    m, book = built()
    assert sb.release2(m) is None and [ws.name for ws in book.sheets] == list(sb.SHEETS)
    assert value(book, "Notes", 4, 1) == "stock-plan v1" and "season" not in book.controls
    r = book.controls["order"]["rows"][0]
    assert order_head(book)[:20] == R1_HEADERS and order_head(book)[20:] == [None, None]
    assert book.get("Order").cells[(r, sb.O_JOINS)].formula.startswith("AND(J12>0,ROUND(MAX(F12,0)+H12-J12*INDEX("
                                                                       "Notes!$F$9:$F$10,")
    assert sb.controls(book, m) == [] and value(book, "Order", r, sb.O_SUGGESTED) == 65


@test("release 2: tag v2, eight sheets (On the way after Order), Rate basis and Carton after Comment; Suggested is "
      "stock_season's (110 = 129.90 season days - 20), controls() empty, no formula an error")
def _():
    m, book = built2()
    assert sb.release2(m) is not None and [ws.name for ws in book.sheets] == list(sb.SHEETS_V2)
    assert value(book, "Notes", 4, 1) == sb.LAYOUT_TAG_V2 == "stock-plan v2"
    assert order_head(book) == R1_HEADERS + ["Rate basis", "Carton"]
    assert book.evaluate() == [] and sb.controls(book, m) == []
    x = m.facts["order"][0]
    want = stock_season.suggested(Decimal(20), x["rate6"], D, 115, m.facts["season"]["lifts"])
    r = book.controls["order"]["rows"][0]
    assert value(book, "Order", r, sb.O_SUGGESTED) == int(want) == int(x["qty"]) == 110
    assert value(book, "Order", r, sb.O_JOINS) is True and value(book, "Order", r, sb.O_VALUE) == 275
    assert value(book, "Order", r, sb.O_RATE_BASIS) == "90 days" and value(book, "Order", r, sb.O_CARTON) is None
    f = book.get("Order").cells[(r, sb.O_SUGGESTED)].formula
    assert "INDEX(Notes!$I$9:$I$11,IFERROR(MATCH(" in f and ",8),6),0),0)" in f, f      # Fill season days, 8 then 6
    assert book.get("Order").cells[(r, sb.O_JOINS)].formula.startswith(
        "AND(J12>0,ROUND(ROUND(MAX(F12,0)+H12-J12*INDEX(Notes!$H$9:$H$11,IFERROR(MATCH(")
    vf = book.controls["order"]["suppliers"]["vfan"]
    assert [value(book, "Order", vf, c) for c in (sb.B_ITEMS, sb.B_UNITS, sb.B_VALUE)] == [1, 110, 275]
    notes = " ".join(str(value(book, "Order", rr, 0)) for rr in range(book.controls["order"]["add"][1] + 2,
                                                                       book.controls["order"]["add"][1] + 12))
    for words in ("each day at its month's lift", "items that would join only come with the next file: tell Claude "
                  "to keep the new setting", "takes the rule of any other supplier", "No purchase order is drafted"):
        assert words in notes, words
    assert "Release 2 adds" not in notes


@test("release 2: Notes holds the settings per supplier and for any other supplier, the yellow lifts (two decimals) "
      "with their basis and stock-out rows, and the data date the windows count from")
def _():
    m, book = built2(leads={"WEKOME": {"days": 55, "basis": "config"}})
    assert book.evaluate() == []
    s = book.controls["season"]
    rows = s["settings"]
    assert rows == {"vfan": 8, "wekome": 9, "*": 10}, rows
    assert [value(book, "Notes", 7, c) for c in range(9)] == [
        "Supplier", "Lead days", "Order every (days)", "Safety low (days)", "Top up to (days)", "Join below (days)",
        "Fill to (days)", "Join season days", "Fill season days"]
    got = {k: [value(book, "Notes", r, c) for c in range(9)] for k, r in rows.items()}
    assert got["vfan"][:7] == ["VFAN", 40, 30, 30, 45, 100, 115], got["vfan"]
    assert [round(v, 6) for v in got["vfan"][7:]] == [114.95, 129.9], got["vfan"]     # 16 + 34.5 + 40.3 + 24.15 ...
    assert got["wekome"][:7] == ["WEKOME", 55, 30, 30, 45, 115, 130], got["wekome"]    # its own lead (facts["leads"])
    assert got["*"][:7] == [sb.ANY_SUPPLIER, 40, 30, 30, 45, 100, 115], got["*"]
    notes = book.get("Notes")
    for key, r in rows.items():
        for c in (sb.S_LEAD, sb.S_EVERY):
            assert notes.cells[(r, c)].fmt == "input_qty", (key, c)
    lifts = [value(book, "Notes", s["lifts"], c) for c in range(sb.M_FIRST, sb.M_LAST + 1)]
    assert lifts == [float(v) if Decimal(v) != int(Decimal(v)) else int(Decimal(v)) for v in LIFTS.values()], lifts
    assert all(notes.cells[(s["lifts"], c)].fmt == "input_dec2" for c in range(sb.M_FIRST, sb.M_LAST + 1))
    basis = [value(book, "Notes", s["basis"], c) for c in range(sb.M_FIRST, sb.M_LAST + 1)]
    assert basis[:4] == ["Jan 2026", "Feb 2026", "Mar 2026 (15 d)", "Apr 2025"] and basis[11] == "Dec 2025", basis
    warn = [value(book, "Notes", s["warn"], c) for c in range(sb.M_FIRST, sb.M_LAST + 1)]
    assert warn[11] == "stock-outs held this month down" and warn[:11] == [None] * 11
    dr, dc = s["date"]
    assert value(book, "Notes", dr, dc) == fx.serial(D) and value(book, "Notes", dr, 0).startswith("Data date")
    checks = [v for v in notes.validations if v[0] == s["lifts"]]
    assert checks and checks[0][4]["validate"] == "decimal" and checks[0][4]["criteria"] == ">"
    words = " ".join(str(c.value) for c in notes.cells.values() if isinstance(c.value, str))
    for text in ("tell Claude to keep the new setting", "items that would join only come with the next file",
                 "last year's months against today's pace", "a window counts at most 365 days",
                 "Any other supplier: the rule of a supplier this sheet does not name"):
        assert text in words, text


@test("release 2: each window's days in each month are formulas of the data date (a month end, a year end, a leap "
      "February, the data date's month met again a year on, the 365-day cap); season days = stock_season's")
def _():
    cases = ((D, 40, 30), (date(2026, 12, 20), 40, 30), (date(2028, 2, 10), 10, 0), (D, 300, 30), (D, 400, 30),
             (date(2026, 1, 31), 0, 0))
    for day, lead, every in cases:
        m, book = built2()
        s = book.controls["season"]
        notes = book.get("Notes")
        r = s["settings"]["vfan"]
        notes.put(*s["date"], day, "g_date")
        notes.put(r, sb.S_LEAD, lead, "input_qty")
        notes.put(r, sb.S_EVERY, every, "input_qty")
        assert book.evaluate() == []
        lifts = m.facts["season"]["lifts"]
        for w, n, c in ((s["windows"]["vfan"][0], lead + every + 30, sb.S_JSEASON),
                        (s["windows"]["vfan"][1], lead + every + 45, sb.S_FSEASON)):
            n = min(n, sb.WINDOW_CAP)
            got = [value(book, "Notes", w, k) for k in range(sb.M_FIRST, sb.M_LAST + 1)]
            assert got == month_count(day, n), (day, n, got, month_count(day, n))
            assert all(isinstance(v, int) for v in got) and value(book, "Notes", w, sb.M_TOTAL) == n
            season = Decimal(repr(value(book, "Notes", r, c))).quantize(Q2, rounding=ROUND_HALF_UP)
            assert season == stock_season.season_days(day, n, lifts), (day, n, season)
    # March met again: 16 days of March 2026 and 10 of March 2027 in a window of 360 days from 15 Mar 2026
    assert month_count(D, 360)[2] == 26 and month_count(D, 365)[2] == 31


@test("release 2: a lift, a Lead days or an Order every typed on Notes moves Suggested as stock_season works it out; "
      "controls() names the difference")
def _():
    base = lifts_of()
    april = {**base, 4: dict(base[4], lift=Decimal("2.00"))}
    for edit, lifts, fill in (("lift", april, 115), ("lead", base, 135), ("every", base, 85)):
        m, book = built2()
        s = book.controls["season"]
        notes = book.get("Notes")
        if edit == "lift":
            notes.put(s["lifts"], sb.M_FIRST + 3, 2.0, "input_dec2")              # April at 2.00
        elif edit == "lead":
            notes.put(s["settings"]["vfan"], sb.S_LEAD, 60, "input_qty")         # 60 + 30 + 45 = 135 days
        else:
            notes.put(s["settings"]["vfan"], sb.S_EVERY, 0, "input_qty")         # 40 + 0 + 45 = 85 days
        checks = sb.controls(book, m)
        x = m.facts["order"][0]
        want = stock_season.suggested(Decimal(20), x["rate6"], D, fill, lifts)
        r = book.controls["order"]["rows"][0]
        assert value(book, "Order", r, sb.O_SUGGESTED) == int(want), (edit, value(book, "Order", r, sb.O_SUGGESTED))
        assert int(want) == {"lift": 136, "lead": 129, "every": 80}[edit], (edit, want)
        assert checks[0] == f"Order sheet: CB2 suggests {int(want)} in the file, 110 on the screen", checks
        assert any(c.startswith("Notes sheet: VFAN fill season days") for c in checks), checks
    # a season that leaves the item enough: it joins no more
    m, book = built2()
    notes = book.get("Notes")
    for c in range(sb.M_FIRST, sb.M_LAST + 1):
        notes.put(book.controls["season"]["lifts"], c, 0.01, "input_dec2")
    sb.controls(book, m)
    r = book.controls["order"]["rows"][0]
    assert value(book, "Order", r, sb.O_JOINS) is False and value(book, "Order", r, sb.O_SUGGESTED) == 0


@test("release 2: a Supplier typed over to a name Notes does not hold takes any other supplier's rule (no #N/A)")
def _():
    rules = {"VFAN": rule(40, 30), "*": rule(10, 0, "overall")}
    m, book = built2(rules=rules)
    r = book.controls["order"]["rows"][0]
    book.get("Order").text(r, sb.O_SUPPLIER, "Nobody", "g_text")
    checks = sb.controls(book, m)
    lifts = m.facts["season"]["lifts"]
    rate6 = m.facts["order"][0]["rate6"]
    assert stock_season.joins(Decimal(20), rate6, D, 40, lifts)
    want = stock_season.suggested(Decimal(20), rate6, D, 55, lifts)
    assert value(book, "Order", r, sb.O_SUGGESTED) == int(want) == 43
    assert not [c for c in checks if "#N/A" in c or "came out as" in c], checks
    assert "Order sheet: VFAN items 0 in the file, 1 on the screen" in checks      # in no supplier's total


@test("release 2: whole cartons, a longer history (Rate basis, ~) and a rate that is not whole: Suggested equals the "
      "screen's to the unit")
def _():
    m, book = built2(packs={"CB2": 12})
    r = book.controls["order"]["rows"][0]
    assert m.facts["order"][0]["qty"] == 120 and sb.controls(book, m) == []     # 110 in cartons of 12
    assert value(book, "Order", r, sb.O_SUGGESTED) == 120 and value(book, "Order", r, sb.O_CARTON) == 12
    m, book = built2(longer={"CB2": (Decimal("1.5"), "2025-11-02")})
    r = book.controls["order"]["rows"][0]
    assert sb.controls(book, m) == [] and value(book, "Order", r, sb.O_SUGGESTED) == 175   # 1.5 x 129.90 - 20
    assert value(book, "Order", r, sb.O_RATE_BASIS) == "longer history from 2025-11-02"
    assert value(book, "Order", r, sb.O_STAR) == "~" and value(book, "Order", r, sb.O_RATE) == 1.5
    after = book.controls["order"]["add"][1]
    notes = [str(value(book, "Order", rr, 0)) for rr in range(after + 2, after + 14)]
    assert any(t.startswith("~ out of stock most of the last 90 days") for t in notes), notes
    w = world()
    w.sale(L("B"), D, 1)                                                         # 46/45 a day: 1.022222
    m, book = built2(World(book=w))
    x = m.facts["order"][0]
    assert x["rate6"] == Decimal("1.022222") and sb.controls(book, m) == []
    assert value(book, "Order", book.controls["order"]["rows"][0], sb.O_SUGGESTED) == int(x["qty"])


@test("release 2: the rule's rounding: rate x season days rounded to 8 decimals, then 6, then up, agrees with "
      "Python's exact rule on half-way cases a single ROUND to 6 decimals misses")
def _():
    e8 = 10 ** 8
    cases = []
    for s in range(1001, 80000, 7):              # season days 10.01 to 800, two decimals
        if s % 2 and s % 5:
            r = (50 * pow(s, -1, e8)) % e8          # rate (6 decimals) x days = position + 0.0000005, exactly
            if 0 < r <= 5_000_000 and r * s >= e8:
                cases.append((Decimal(r) / 10 ** 6, Decimal(s) / 100, Decimal(r * s // e8)))
    assert len(cases) > 100
    for rate, days, pos in cases:
        b = bk.Book()
        ws = b.sheet("T")
        for c, v in enumerate((float(rate), float(days), int(pos))):
            ws.put(0, c, v)
        joins, suggested = sb.rule_formulas("A1", "C1", "B1", "B1", "D1", "F1")
        ws.formula(0, 3, joins)
        ws.formula(0, 4, suggested)
        assert b.evaluate() == []
        want = stock.stock_plan.round6(pos - rate * days) < 0
        assert ws.cells[(0, 3)].value is want, (rate, days, pos)
        assert ws.cells[(0, 4)].value == (stock.stock_plan.units_up(rate * days - pos) if want else 0), (rate, days)


@test("release 2: On the way lists the lines still due by supplier (late, dated ahead, matching no item) with a "
      "total; nothing logged is one line; the Order sheet deducts the units due")
def _():
    m, book = built2(lines=WAY)
    assert book.evaluate() == [] and sb.controls(book, m) == []
    ws = book.get("On the way")
    head = sb.W_HEAD
    assert [value(book, "On the way", head, c) for c in range(10)] == [
        "Supplier", "Code", "Item", "Ordered on", "Ordered", "Received since", "Still due", "Expected by", "Late",
        "Ref"]
    got = [[value(book, "On the way", r, c) for c in range(10)] for r in range(head + 1, head + 4)]
    assert [g[:3] for g in got] == [["VFAN", "CB2", NAMES["B"]], ["VFAN", "HG3", NAMES["C"]],
                                    ["WEKOME", "ZN9", "no item has this code"]], got
    assert [g[4:7] for g in got] == [[30, 10, 20], [5, 0, 5], [12, 0, 12]]
    assert got[0][3] == fx.serial(date(2026, 3, 1)) and got[0][7] == fx.serial(date(2026, 4, 10))
    assert [g[8] for g in got] == [None, None, "yes"] and got[0][9] == "vfan:2026-03-01:CB2"
    assert value(book, "On the way", head + 4, 0) == "Total"
    assert [value(book, "On the way", head + 4, c) for c in (4, 5, 6)] == [47, 10, 37]
    words = " ".join(str(c.value) for c in ws.cells.values() if isinstance(c.value, str))
    for text in ("1 line dated after the data date", "1 line matches no item of the main store",
                 "1 logged line received in full is not shown"):
        assert text in words, text
    r = book.controls["order"]["rows"][0]
    assert value(book, "Order", r, sb.O_WAY) == 20 and value(book, "Order", r, sb.O_SUGGESTED) == 90   # 129.90 - 40
    m, book = built2()
    assert value(book, "On the way", head, 0) == sb.EMPTY_LOG and not [k for k in book.get("On the way").cells
                                                                         if k[0] > head]
    m, book = built2(lines=[way("CA1", "VFAN", date(2026, 2, 1), 40, 40)])
    assert value(book, "On the way", head + 1, 0) == "Every logged order has arrived: 1 line, received in full."


@test("release 2: the saved file: the tag in B5 and the figures stamp in B6 (its row hidden), lifts of two decimals "
      "and whole month days cached, the Order headers, every formula with its value (openpyxl)")
def _():
    with data_root() as root:
        m, _inp = model_of()
        with_r2(m, lines=WAY)
        res = sb.day_file(m)
        assert res["made"] and res["controls"] == [], res
        path = Path(res["path"])
        wb, vb = opened(path)
        with zipfile.ZipFile(path) as z:
            xml = [z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("xl/worksheets/sheet")]
        assert sb.read_stamp(path) == (True, sb.figures_digest(m))
    assert wb.sheetnames == list(sb.SHEETS_V2)
    notes, values = wb["Notes"], vb["Notes"]
    assert notes["B5"].value == "stock-plan v2" and notes["B6"].value == sb.figures_digest(m)
    assert notes.row_dimensions[6].hidden and re.fullmatch(r"[0-9a-f]{16}", notes["B6"].value)
    lift_row = next(r for r in range(1, 80) if notes[f"A{r}"].value == "Lift")
    assert [values.cell(lift_row, c).value for c in range(2, 14)][3] == 1.15
    join_row = next(r for r in range(1, 120) if notes[f"A{r}"].value == "VFAN: join below")
    assert [values.cell(join_row, c).value for c in range(2, 15)] == [0, 0, 16, 30, 31, 23, 0, 0, 0, 0, 0, 0, 100]
    assert values["H9"].value == 114.95 and values["I9"].value == 129.9
    order = wb["Order"]
    head = next(r for r in range(1, 40) if order[f"A{r}"].value == "Supplier" and order[f"B{r}"].value == "Code")
    assert [order.cell(head, c).value for c in range(1, 23)] == R1_HEADERS + ["Rate basis", "Carton"]
    assert vb["Order"][f"N{head + 1}"].value == 90 and order[f"U{head + 1}"].value == "90 days"
    assert order[f"B{head + 2}"].fill.fgColor.rgb == YELLOW and order.auto_filter.ref.startswith(f"A{head}:V")
    assert vb["On the way"]["G6"].value == 20 and wb["All items"]["W5"].value == "Supplier from"
    formulas = 0
    for text in xml:
        for cell in re.findall(r"<c [^>]*>(.*?)</c>", text):
            if "<f>" in cell:
                formulas += 1
                assert "<v>" in cell and not re.search(r"<v>#(N/A|VALUE!|REF!|DIV/0!|NUM!)</v>", cell), cell[:120]
    assert formulas > 200                  # 3 settings rows x 4, 6 windows x 13 on Notes, the On the way totals ...


@test("the figures stamp: the same figures keep the day's file; an order logged, a cost loaded later or a file "
      "with no stamp make the next free name, with a line that says why")
def _():
    changed = re.compile(r"the figures changed since \d\d:\d\d: a new file")
    with data_root() as root:
        f = folder(root)
        m, _inp = model_of()
        first = sb.day_file(with_r2(m))
        again = sb.day_file(with_r2(model_of()[0]))
        assert first["made"] and not again["made"] and again["path"] == first["path"], again
        logged = sb.day_file(with_r2(model_of()[0], lines=WAY))                 # an order logged: CB2 on the way
        assert logged["made"] and Path(logged["path"]).name == SECOND, logged
        assert changed.fullmatch(logged["controls"][0]) and logged["controls"][1:] == [], logged["controls"]
        same = sb.day_file(with_r2(model_of()[0], lines=WAY))
        assert not same["made"] and Path(same["path"]).name == SECOND
        m = with_r2(model_of()[0], lines=WAY)
        m.facts["order"][0]["item"].cost = Decimal("2.6")                       # the MRN cost loaded later
        m.facts["order"][0]["cost_bhd"] = m.facts["plan"]["order"][0]["cost_bhd"] = fmt.q3(90 * Decimal("2.6"))
        costed = sb.day_file(m)
        assert costed["made"] and Path(costed["path"]).name == FILE.replace(".xlsx", " (3).xlsx")
        assert changed.fullmatch(costed["controls"][0]), costed["controls"]
        assert sorted(books(root)) == sorted([FILE, SECOND, FILE.replace(".xlsx", " (3).xlsx")])
    with data_root() as root:
        f = folder(root)
        f.mkdir(parents=True)
        (f / FILE).write_bytes(b"not a workbook")                               # no stamp to compare
        res = sb.day_file(with_r2(model_of()[0]))
        assert res["made"] and Path(res["path"]).name == SECOND
        assert re.fullmatch(r"the file made at \d\d:\d\d carries no figures stamp \(an older layout\): a new file",
                            res["controls"][0]), res["controls"]
        had, saved = "openpyxl" in sys.modules, sys.modules.get("openpyxl")
        sys.modules["openpyxl"] = None                                          # nothing can be read: the file stands
        try:
            assert sb.read_stamp(f / FILE) == (False, None)
            m = with_r2(model_of()[0], lines=WAY)
            assert not sb.day_file(m)["made"]
        finally:
            if had:
                sys.modules["openpyxl"] = saved
            else:
                del sys.modules["openpyxl"]
    m1, m2 = with_r2(built()[0]), with_r2(built()[0], lifts=lifts_of({**LIFTS, 4: "1.16"}))
    assert sb.figures_digest(m1) != sb.figures_digest(m2)                      # a lift is in it
    assert sb.figures_digest(built()[0]) != sb.figures_digest(m1)              # and the layout


@test("team copies are never the day's file (no needless (2) after a team copy); the latest copy is the one named "
      "and compared, whichever copy he saved last, as readback.newest_plan reads it")
def _():
    from scripts.ops.workbook import readback
    assert sb.TEAM_MARK == readback.TEAM_MARK
    with data_root() as root:
        f = folder(root)
        assert sb.day_file(with_r2(model_of()[0]))["made"]
        stem = FILE[: -len(".xlsx")]
        team = [f"{stem} - Team.xlsx", f"{stem} - Team (2).xlsx", f"{stem} (2) - Team.xlsx"]
        for name in team:
            (f / name).write_bytes(b"a team copy")                          # newer than the plan, and no stamp
        res = sb.day_file(with_r2(model_of()[0]))
        assert not res["made"] and Path(res["path"]).name == FILE, res
        assert books(root) == sorted([FILE] + team)
        assert [p.name for p in sb.existing(f, sb.STEM.format(d=D.isoformat()))] == [FILE]
        assert Path(sb.day_file(with_r2(model_of()[0]), force=True)["path"]).name == SECOND
        later = time.time() + 120
        os.utime(f / FILE, (later, later))                                  # he saves the first copy after (2)
        res = sb.day_file(with_r2(model_of()[0]))
        assert not res["made"] and Path(res["path"]).name == SECOND, res   # (2): the latest copy, never a (3)
        assert readback.newest_plan(f).name == SECOND and books(root) == sorted([FILE, SECOND] + team)


@test("Ageing: Before / Since a week never read as a multiple: not in stock before, too few before, too soon; a cut "
      "batch's change in the screen's words")
def _():
    with data_root():
        m, _inp = model_of(World(book=damage_world(), cuts=CUTS))
        v1(m)
        review = m.facts["review"]
        hg3 = next(r for r in review["rows"] if r["item"].code == "HG3")
        assert hg3["change"] is None and hg3["in_stock_before_days"] >= 30            # nothing sold before its cut
        review["rows"] = [hg3, dict(hg3, in_stock_before_days=10),
                          dict(hg3, too_soon=True, before_week=Decimal("2.5"), since_week=Decimal("9")),
                          dict(hg3, change=Decimal("1.5"), before_week=Decimal("2"), since_week=Decimal("3"))]
        b = review["batches"][0]
        review["batches"] = [b, dict(b, too_soon=True), dict(b, change=Decimal("1.5")),
                             dict(b, in_stock_before_items=0, change=None)]
        book = sb.build_book(m, MADE)
    assert book.evaluate() == []
    cells = [(value(book, "Ageing", r, 13), value(book, "Ageing", r, 14)) for r in range(6, 10)]
    assert cells == [("too few before", 0), ("not in stock before", 0), (2.5, "too soon"), (2, 3)], cells
    head = row_of(book, "Ageing", 0, "Price cuts in the last 180 days, by cut date") + 1
    batch = [(value(book, "Ageing", r, 10), value(book, "Ageing", r, 11)) for r in range(head + 1, head + 5)]
    assert [x[1] for x in batch] == ["too few before", "too soon to judge", 1.5, "not in stock before"], batch
    assert batch[1][0] == "too soon" and value(book, "Ageing", head + 2, 8) is None


@test("All items (release 2): Supplier from at the far right (W), blank where the plan does not give it")
def _():
    with data_root():
        m, _inp = model_of()
        with_r2(m)
        for row in m.facts["plan"]["all_items"]:
            row.pop("supplier_basis", None)
            if row["code"] == "CB2":
                row["supplier_basis"] = "config"
        book = sb.build_book(m, MADE)
    first, last = book.controls["all_items"]
    assert value(book, "All items", sb.ALL_HEAD, sb.A_SUPPLIER_FROM) == "Supplier from"
    got = {value(book, "All items", r, sb.A_CODE): value(book, "All items", r, sb.A_SUPPLIER_FROM)
           for r in range(first, last + 1)}
    assert got["CB2"] == "config" and {v for k, v in got.items() if k != "CB2"} == {None}, got
    assert book.get("All items").page["autofilter"].startswith(f"A{sb.ALL_HEAD + 1}:W")
    m, book = built()                                                           # release 1: no such column
    assert value(book, "All items", sb.ALL_HEAD, sb.A_SUPPLIER_FROM) is None


@test("Summary (release 2): the layout, the order rhythm, the season and the goods on the way")
def _():
    for lines, words in (((), "nothing logged yet (/yq-stock ordered <supplier> logs an order you placed)"),
                         (WAY, "2 lines still due, deducted (the On the way sheet); 1 line matches no item")):
        m, book = built2(lines=lines)
        assert book.evaluate() == []

        def at(label: str):
            return value(book, "Summary", row_of(book, "Summary", 0, label), 1)

        assert at("Layout") == "stock-plan v2" and at("Order every, days") == 30
        assert (at("Joins below, days of sales"), at("Filled to, days of sales")) == (100, 115)   # the "*" rule
        assert at("Season") == "on" and at("Goods on the way") == words
        note = value(book, "Summary", row_of(book, "Summary", 0, "Season"), 2)
        assert note.startswith("Each day counts at its month's lift: Jan 0.80, Feb 0.90"), note
        assert "Held down by stock-outs: Dec" in note


@test("the command's own file has nothing to check, whichever release its facts are; with release 2's it is v2")
def _():
    with data_root():
        out = run(World(book=damage_world(), cuts=CUTS))
        path = Path(out["flags"]["file"])
        m, _inp = model_of(World(book=damage_world(), cuts=CUTS))
        tag = sb.LAYOUT_TAG_V2 if sb.release2(m) is not None else sb.LAYOUT_TAG
        compared, stamp = sb.read_stamp(path)
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        try:
            names, got = list(wb.sheetnames), wb["Notes"]["B5"].value
        finally:
            wb.close()
    assert "To check in the file" not in out["text"], out["text"][:600]
    assert got == tag and compared and stamp == sb.figures_digest(m)
    assert names == list(sb.SHEETS_V2 if tag == sb.LAYOUT_TAG_V2 else sb.SHEETS)


def main() -> int:
    passed = failed = 0
    for name, fn in TESTS:
        try:
            fn()
            print(f"PASS  {name}")
            passed += 1
        except Exception:  # noqa: BLE001
            print(f"FAIL  {name}")
            traceback.print_exc()
            failed += 1
    paths.set_data_root(None)
    print(f"\n{passed} passed, {failed} failed")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
