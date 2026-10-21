"""Ops engine: the workbook engine and the monthly report (python -m scripts.ops report month).

    python -m tests.test_ops_report

Same lightweight runner as the other suites (no pytest). No database, no network, no .env:
SUPABASE_URL, SUPABASE_KEY and DATABASE_URL are pinned to ci.invalid BEFORE any import and
scripts.ops.db is imported first. The report reads through a Source; here a stand-in answers from an
in-memory list of invented day-book lines (two vans, an outlet, a few shops, round amounts), so every
figure below is a sum a reader can redo by hand. Every folder is a temporary one.

Covered:
  * the formula evaluator: Excel's precedence, case-blind text, criteria with their operators and
    wildcards, numbers compared as a spreadsheet shows them, rounding half away from zero, ROUNDUP
    away from zero on the value as it is shown, dates, the array idioms the report uses (SUMPRODUCT
    ranks, LOOKUP(2, 1/(...))), errors as values, and a formula it does not know refused instead of
    given a made-up value;
  * the category rule: the longest alphabetic prefix (TB is its own row, GN a power bank), an unknown
    code through the config name table (a code key wins over a card key), a no-value item in the
    handset row, Other only when nothing else, the Notes sentence from the rule in force;
  * the data rules: the "SI : " prefix, every spelling of cash customer as one, a narration with a
    masked ID left blank, empty money as 0, the order by date, the invoice's trailing number and the
    line; the lists and their order; long labels cut at a whole word;
  * the eight sheets on the synthetic months: every formula evaluated (no error value), the figures
    against hand sums, the controls nil, the positions following the list lengths, the commentary
    DRAFT, the counts of rules, data bars, the colour scale, sparklines and charts;
  * the file: written through a temporary name, never over an existing file, the next free name, a
    valid workbook; the comparison with a reference catches a changed value and sets aside what it
    is told to; no 9-digit or 15-plus-digit run in what the owner reads;
  * the pieces added for the stock workbook, read back from a small invented grid with openpyxl: data
    validations, a hidden column, a book's own conditional-format styles, the grid and input formats,
    freeze panes and a filter, a ROUNDUP formula with its cached value; refused when a rule names a
    style nobody registered; and the monthly file carrying none of them;
  * the week as a period: the Saturday it ends on, its four windows (month-crossing, the 31st, the year's
    end), its lines read from the first day a window reaches;
  * the command: the whole-month gate (not ready: nothing written), the week and a bad month refused
    in one sentence (a day that is not a Saturday too), the week handed to its own module once that
    exists, a month with no sales refused, the build record, the text;
  * the real September file, when it is on this machine AND YQ_OPS_GOLDEN=1 is set (SKIP otherwise):
    September 2026 rebuilt from production through the read-only engine in a child process and
    compared with the reference on every sheet but Notes.
"""
from __future__ import annotations

import io
import json
import os
import re
import subprocess
import sys
import tempfile
import traceback
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

os.environ["SUPABASE_URL"] = "https://ci.invalid"
os.environ["SUPABASE_KEY"] = "ci-key"
os.environ["DATABASE_URL"] = "postgresql://ci:ci@ci.invalid:5432/ci"
os.environ.pop("YQ_MAIN_CHECKOUT", None)

from scripts.ops import db  # noqa: E402,F401  (first, as every command does)
from scripts.ops import config, paths, status  # noqa: E402
from scripts.ops.workbook import book as bk  # noqa: E402
from scripts.ops.workbook import build, categories as cats, formula as fx, golden, model as mdl  # noqa: E402
from scripts.ops.workbook import sheets_core as sc, style  # noqa: E402

TESTS: list[tuple[str, object]] = []
MONTH = "2026-03"                      # the report month; February is the comparison month
REFERENCE = Path(os.environ.get("YQ_OPS_REFERENCE") or
                 (paths.data_root() / "YQ Bahrain - Mobile Accessories Sales Report - Sep 2026 (1).xlsx"))


def test(name):
    def deco(fn):
        TESTS.append((name, fn))
        return fn
    return deco


def raises(kind, fn) -> Exception:
    try:
        fn()
    except kind as exc:
        return exc
    raise AssertionError(f"expected {kind.__name__}")


# ── a tiny book to evaluate single formulas ────────────────────────────────────────────────────

def ev(formula: str, cells: dict | None = None, names: dict | None = None):
    """Evaluate one formula on a sheet 'S' holding `cells` ({'A1': value})."""
    b = bk.Book()
    s = b.sheet("S")
    for ref, value in (cells or {}).items():
        r, c = fx.parse_cell(ref)
        if isinstance(value, str) and value.startswith("="):
            s.formula(r, c, value)
        else:
            s.put(r, c, value)
    s.formula(99, 0, formula)
    for k, v in (names or {}).items():
        b.define(k, v)
    b.evaluate()
    return s.cells[(99, 0)].value


# ── the synthetic world ────────────────────────────────────────────────────────────────────────

VAN_A, VAN_B, OUTLET, SHOW = "Van A - Acc WH", "Van B - Acc WH", "Causeway", "YQ Roadshow"
ITEMS = {
    "T01": "T01 Demo Earbuds One (VFAN) T01",
    "X01": "X01 Demo Cable One (VFAN)",
    "TB": "TB-D1 Demo Cable Box (VFAN)",
    "GN": "GN-01 Demo Power Bank",
    "WS": "WS-01 Demo Wireless Earbuds",
    "WG": "WG-01 Demo Neckband",
    "ZZ": "ZZ9 Demo Gadget",
    "PHONE": "ZTE-X1 DEMO PHONE",
}
NAMES = {"WG-01": "Headphones & Neckbands", "Wireless Audio": "TWS Earbuds"}   # the config name table
UPCOMING = {"WS-01": "Wireless Audio", "WG-01": "Wireless Audio"}               # the Upcoming cards


def line(day: date, inv: int, no: int, cust: str, item: str, qty: int, rate: float, rep: str, narr: str | None = None):
    gross = round(qty * rate, 3)
    net = round(gross / 1.1, 2) if gross else None
    return {"line_id": inv * 10 + no, "invoice_no": f"SI : SI-YQ-26-{day.month:02d}-{inv}", "line_no": no,
            "sale_date": day.isoformat(), "customer_account": cust, "customer_name": cust, "item_name": item,
            "quantity": float(qty), "rate_bhd": rate if rate else None, "gross_bhd": gross if gross else None,
            "net_bhd": net, "salesman_raw": rep, "salesman_resolved": rep, "narration": narr,
            "category_name": None, "discount_bhd": None, "vat_amount_bhd": None}


def world_rows() -> list[dict]:
    feb, mar = date(2026, 2, 1), date(2026, 3, 1)
    d = lambda base, n: base + timedelta(days=n)  # noqa: E731
    rows = [
        # February (the comparison month)
        line(d(feb, 1), 1, 1, "Shop One", ITEMS["T01"], 10, 2.2, VAN_A),
        line(d(feb, 1), 1, 2, "Shop One", ITEMS["X01"], 20, 1.1, VAN_A),
        line(d(feb, 3), 2, 1, "Cash customer", ITEMS["X01"], 10, 1.1, OUTLET),
        line(d(feb, 9), 3, 1, "Shop Three", ITEMS["GN"], 5, 5.5, VAN_B, "credit note for return"),
        line(d(feb, 20), 4, 1, "Cash Customer", ITEMS["PHONE"], 4, 0, SHOW, "ID 123456789 given"),
        # March (the report month)
        line(d(mar, 0), 9, 1, "Shop One", ITEMS["T01"], 20, 2.2, VAN_A),
        line(d(mar, 0), 10, 1, "Shop Two", ITEMS["X01"], 30, 1.1, VAN_B),
        line(d(mar, 0), 10, 2, "Shop Two", ITEMS["TB"], 2, 11.0, VAN_B),
        line(d(mar, 4), 11, 1, "CASH CUSTOMER", ITEMS["WS"], 5, 4.4, OUTLET),
        line(d(mar, 9), 12, 1, "Shop Four", ITEMS["WG"], 3, 3.3, VAN_A),
        line(d(mar, 9), 12, 2, "Shop Four", ITEMS["ZZ"], 1, 1.1, VAN_A),
        line(d(mar, 30), 13, 1, "Cash Customer", ITEMS["PHONE"], 6, 0, SHOW, "ID 987654321 given"),
        line(d(mar, 30), 14, 1, "Shop One", ITEMS["GN"], 10, 5.5, VAN_A, "credit note"),
    ]
    # v_sales masks a 9-digit run before the engine ever sees it (the R7b mask): do the same here
    for r in rows:
        if r["narration"]:
            r["narration"] = re.sub(r"(?<!\d)\d{6}(\d{3})(?!\d)", r"******\1", r["narration"])
    return rows


class FakeSource(build.Source):
    """The report's reads answered from the world (and counted, so a test can say nothing else was read)."""

    def __init__(self, rows=None, to: date = date(2026, 4, 7), uploaded: str = "2026-04-08T09:00:00+03:00",
                 upcoming_fails: bool = False):
        self.rows = world_rows() if rows is None else rows
        self.to, self.uploaded, self.upcoming_fails = to, uploaded, upcoming_fails
        self.calls: list[str] = []

    def anchors(self) -> dict:
        self.calls.append("anchors")
        return {"focus_to": self.to.isoformat(), "uploaded_at": self.uploaded}

    def lines(self, first, last):
        self.calls.append("lines")
        return [r for r in self.rows if first.isoformat() <= r["sale_date"] <= last.isoformat()]

    def totals(self, first, cur_first, last):
        self.calls.append("totals")
        out: dict = {}
        for r in self.lines(first, last):
            k = "prev" if r["sale_date"] < cur_first.isoformat() else "cur"
            t = out.setdefault(k, {"lines": 0, "qty": 0, "gross": 0, "net": 0})
            t["lines"] += 1
            t["qty"] += r["quantity"] or 0
            t["gross"] = round(t["gross"] + (r["gross_bhd"] or 0), 3)
            t["net"] = round(t["net"] + (r["net_bhd"] or 0), 3)
        return out

    def upcoming(self) -> dict:
        self.calls.append("upcoming")
        if self.upcoming_fails:
            raise RuntimeError("no session")
        return dict(UPCOMING)


def built(rows=None, names=NAMES, prepared_by: str = ""):
    return build.build_month(MONTH, FakeSource(rows), prepared_on=date(2026, 4, 8), prepared_by=prepared_by,
                             category_names=names)


def value(b, sheet: str, ref: str):
    r, c = fx.parse_cell(ref)
    cell = b.get(sheet).cells.get((r, c))
    return None if cell is None else cell.value


def formula_of(b, sheet: str, ref: str):
    r, c = fx.parse_cell(ref)
    cell = b.get(sheet).cells.get((r, c))
    return None if cell is None else cell.formula


def find_row(b, sheet: str, col: str, text: str) -> int:
    """The Excel row of the first cell in `col` holding `text`."""
    c = fx.col_index(col)
    for (r, cc), cell in sorted(b.get(sheet).cells.items()):
        if cc == c and cell.value == text:
            return r + 1
    raise AssertionError(f"{text!r} not found in {sheet}!{col}")


class data_root:
    """A temporary data root (exports/, business_data/), optionally with a last load on record."""

    def __init__(self, last_load: dict | None = None, cfg: dict | None = None):
        self.last_load, self.cfg = last_load, cfg

    def __enter__(self) -> Path:
        self._tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.root = Path(self._tmp.name).resolve()
        paths.set_data_root(self.root)
        config.use(self.cfg)
        if self.cfg is None:
            config.use(None)                    # no config.json: the report keys read their defaults
        if self.last_load:
            status.write_state({"last_load": self.last_load})
        return self.root

    def __exit__(self, *exc):
        paths.set_data_root(None)
        config.use(None)
        self._tmp.cleanup()
        return False


class Args:
    def __init__(self, spec=("month", MONTH), prepared_on="2026-04-08", prepared_by=None, compare=None):
        self.spec, self.prepared_on, self.prepared_by, self.compare = list(spec), prepared_on, prepared_by, compare


LONG_RUN = re.compile(r"(?<!\d)\d{9}(?!\d)|\d{15,}")


# ── the formula evaluator ──────────────────────────────────────────────────────────────────────

@test("formulas: Excel's precedence, text and numbers as a spreadsheet compares them")
def _():
    assert ev("1+2*3") == 7 and ev("-2^2") == 4 and ev("(1+2)*3") == 9 and ev("2^3^2") == 64
    assert ev("10/4") == 2.5 and ev("7-2-1") == 4 and ev('"a"&"b"&1&TRUE') == "ab1TRUE"
    assert ev('"Shop"="shop"') is True and ev('"a"<"B"') is True and ev("0.1+0.2=0.3") is True
    assert ev('1<"1"') is True                         # a number sorts before any text
    assert ev("A1+1", {"A1": None}) == 1 and ev('A1&"x"', {"A1": None}) == "x"
    assert isinstance(ev("1/0"), fx.Err) and ev("IFERROR(1/0,7)") == 7
    assert ev("IF(1/0,1,2)").code == "#DIV/0!"         # an error flows, it is not a crash


@test("formulas: rounding half away from zero, FIXED, dates and month names")
def _():
    assert ev("ROUND(2.675,2)") == 2.68 and ev("ROUND(-2.5,0)") == -3 and ev("ROUND(1234.5678,-2)") == 1200
    assert ev("FIXED(5961.5,0)") == "5,962" and ev("FIXED(1234.56,1)") == "1,234.6" and ev("FIXED(0.4,0)") == "0"
    jan31 = fx.serial(date(2026, 1, 31))
    assert ev(f"EDATE({jan31},1)") == fx.serial(date(2026, 2, 28)) and ev(f"EDATE({jan31},-1)") == fx.serial(date(2025, 12, 31))
    assert ev("DATE(2026,13,1)") == fx.serial(date(2027, 1, 1)) and ev(f"DAY(EDATE({jan31},1)-1)") == 27
    assert ev(f'CHOOSE(MONTH({jan31}),"Jan","Feb")&" "&YEAR({jan31})') == "Jan 2026"
    assert ev('SUBSTITUTE("Van A - Acc WH"," - Acc WH","")') == "Van A"


@test("formulas: ROUNDUP goes away from zero on the value as a spreadsheet shows it; an error argument flows")
def _():
    # ROUND as a spreadsheet holds the number: 15 significant digits first (5 x 1.2345 is 6.1724999... as a double)
    assert ev("ROUND(5*1.2345,3)") == 6.173 and ev("ROUND(1.005,2)") == 1.01 and ev("ROUND(-2.5,0)") == -3
    assert ev("ROUNDUP(2.1,0)") == 3 and ev("ROUNDUP(2.0,0)") == 2 and ev("ROUNDUP(-2.1,0)") == -3
    assert ev("ROUNDUP(0.0001,0)") == 1 and ev("ROUNDUP(1234.5,-2)") == 1300 and ev("ROUNDUP(3.14159,2)") == 3.15
    assert ev("ROUNDUP(-1234.5,-2)") == -1300 and ev("ROUNDUP(-0.0001,0)") == -1 and ev("ROUNDUP(0,0)") == 0
    assert ev("ROUNDUP((0.1+0.2)*10,0)") == 3                  # 3.0000000000000004 in binary, 3 as it is shown
    assert ev("ROUNDUP(SUM(A1:A2)*10,0)", {"A1": 0.1, "A2": 0.2}) == 3
    assert ev("ROUNDUP(0.7+0.1,1)") == 0.8 and ev("ROUNDUP(1.1*1.1,2)") == 1.21
    assert ev("ROUNDUP(6.000000000000001,0)") == 6            # the 16th digit is noise ...
    assert ev("ROUNDUP(6.00000000000001,0)") == 7             # ... the 15th is a real difference
    assert ev("ROUNDUP(3.14159,2.9)") == 3.15 and ev('ROUNDUP("3.2",0)') == 4        # digits cut, numeric text read
    assert ev("ROUNDUP(A1,0)", {"A1": None}) == 0 and ev("ROUNDUP(TRUE,0)") == 1
    assert ev("ROUNDUP(1E30,0)") == 1e30 and ev("ROUNDUP(1.5,1000)") == 1.5          # big or many digits: no crash
    assert ev("ROUNDUP(1.7976931348623157E308,-308)").code == "#NUM!"
    assert ev("ROUNDUP(1/0,0)").code == "#DIV/0!" and ev("ROUNDUP(1,1/0)").code == "#DIV/0!"
    assert ev('ROUNDUP("x",0)').code == "#VALUE!" and ev('ROUNDUP(1,"y")').code == "#VALUE!"
    raises(fx.FormulaError, lambda: ev("ROUNDUP(3.2)"))       # Excel takes both arguments: not written with one


@test("formulas: criteria carry their operator, whole-cell wildcards and case-blind text")
def _():
    cells = {"A1": "Cash Customer", "A2": "cash customer", "A3": "Shop *1", "A4": None, "A5": "Shop 2",
             "B1": 1, "B2": 2, "B3": 3, "B4": 4, "B5": 5}
    assert ev('SUMIFS(B1:B5,A1:A5,"CASH CUSTOMER")', cells) == 3
    assert ev('COUNTIFS(A1:A5,"shop*")', cells) == 2 and ev('COUNTIFS(A1:A5,"Shop ~*1")', cells) == 1
    assert ev('COUNTIFS(A1:A5,"")', cells) == 1 and ev('COUNTIFS(A1:A5,"<>")', cells) == 4
    assert ev('COUNTIFS(B1:B5,">=3")', cells) == 3 and ev('SUMIFS(B1:B5,B1:B5,"<>2")', cells) == 13
    assert ev('COUNTIFS(A1:A5,"cash*",B1:B5,">1")', cells) == 1
    assert ev("SUMIFS(B1:B5,B1:B5,3)", cells) == 3 and ev('COUNTIF(B1:B5,">0")', cells) == 5
    assert ev('COUNTIFS(A1:A5,"Shop")', cells) == 0                     # a whole cell must match


@test("formulas: MATCH, INDEX, LOOKUP(2,1/(...)) and the SUMPRODUCT rank with ties")
def _():
    cells = {"A1": "x", "A2": "Shop", "A3": "y", "A4": "shop", "B1": 10, "B2": 30, "B3": 30, "B4": 20}
    assert ev('MATCH("SHOP",A1:A4,0)', cells) == 2 and ev("INDEX(B1:B4,3)", cells) == 30
    assert ev('MATCH("none",A1:A4,0)', cells).code == "#N/A"
    assert ev('LOOKUP(2,1/(A1:A4="shop"),B1:B4)', cells) == 20          # the last row the condition holds for
    ranks = [ev(f"SUMPRODUCT(--($B$1:$B$4>B{i}))+SUMPRODUCT(--($B$1:B{i}=B{i}))", cells) for i in range(1, 5)]
    assert ranks == [4, 1, 2, 3], ranks                                 # ties ranked by their order
    assert ev("COUNTA(A1:A4)", cells) == 4 and ev("SUM(B1:B4)", cells) == 90


@test("formulas: an unknown function or a circular reference is refused, never given a value")
def _():
    raises(fx.FormulaError, lambda: ev("VLOOKUP(1,A1:A2,1,0)"))
    raises(fx.FormulaError, lambda: ev("NOTANAME+1"))
    exc = raises(fx.FormulaError, lambda: ev("1", {"A1": "=A2", "A2": "=A1"}))
    assert "circular" in str(exc)
    assert ev("Total", names={"Total": "S!$A$1"}) is None
    raises(fx.FormulaError, lambda: ev("1+"))


# ── the category rule ──────────────────────────────────────────────────────────────────────────

@test("categories: the longest alphabetic prefix; TB is its own row and never TWS")
def _():
    assert cats.code_prefix("GN-08C") == "GN" and cats.code_prefix("UK12") == "UK" and cats.code_prefix("TB-D3") == "TB"
    assert cats.by_code("T17") == "TWS Earbuds" and cats.by_code("TB-D3") == cats.BUNDLES
    assert cats.by_code("GN-08C") == "Power Banks" and cats.by_code("G176-1") == "Power Banks"
    assert cats.by_code("BE02") == "Headphones & Neckbands" and cats.by_code("BS06") == "Bluetooth Speakers"
    assert cats.by_code("ZTE-A35e") is None and cats.by_code("WDC-C37") is None and cats.by_code("9X") is None


@test("categories: an unknown code through the name table (code beats card), no value to the handset row, else Other")
def _():
    assert cats.resolve("WG-01", ["Wireless Audio"], NAMES) == ("Headphones & Neckbands", "name table")
    assert cats.resolve("WS-01", ["Wireless Audio"], NAMES) == ("TWS Earbuds", "name table")
    assert cats.resolve("ws-01", ["wireless audio"], {}) == (cats.OTHER, "none")
    assert cats.resolve("ZTE-A35e", [], {}, no_value=True) == (cats.HANDSET, "no value")
    assert cats.resolve("T17", [], {"T17": "Cables"}) == ("TWS Earbuds", "code")       # the rule comes first
    text = cats.rule_paragraph({})
    for part in ("T = TWS earbuds", "X, P and L = cables", "F, G, GN and K = power banks", "TB = bundle sets"):
        assert part in text, part
    assert "WG-01 = Headphones & Neckbands" in cats.rule_paragraph(NAMES)


@test("config: the two report keys read their defaults when config.json has none, and are type-checked")
def _():
    with data_root():
        assert config.get("report.category_names") == {} and config.get("report.prepared_by") == ""
        assert "report.category_names" not in config.missing_keys()
        paths.write_text(paths.config_file(), json.dumps({"report.category_names": ["not", "a", "table"]}))
        config.use(None)
        raises(config.ConfigMissing, lambda: config.get("report.category_names"))
        paths.write_text(paths.config_file(), json.dumps({"report.category_names": {"WG-01": "Headphones & Neckbands"}}))
        config.use(None)
        assert config.get("report.category_names") == {"WG-01": "Headphones & Neckbands"}


# ── the data rules ─────────────────────────────────────────────────────────────────────────────

@test("data: SI prefix off, one cash customer, a masked ID left out, empty money 0, ordered by date and invoice number")
def _():
    lines = mdl.tidy(world_rows())
    assert [ln.invoice for ln in lines][:3] == ["SI-YQ-26-02-1", "SI-YQ-26-02-1", "SI-YQ-26-02-2"]
    march = [ln.invoice for ln in lines if ln.sale_date == date(2026, 3, 1)]
    assert march == ["SI-YQ-26-03-9", "SI-YQ-26-03-10", "SI-YQ-26-03-10"], march   # 9 before 10, by the number
    cash = {ln.customer for ln in lines if "ash" in ln.customer.lower()}
    assert cash == {mdl.CASH_ACCOUNT}, cash
    phone = [ln for ln in lines if ln.item == ITEMS["PHONE"]]
    assert all(ln.narration == "" and ln.narration_masked for ln in phone)
    assert all(ln.rate == 0 and ln.gross == 0 and ln.net == 0 for ln in phone)
    credit = [ln for ln in lines if "credit note" in ln.narration]
    assert len(credit) == 2 and not any(ln.narration_masked for ln in credit)
    assert mdl.product_name(ITEMS["T01"], "T01") == "Demo Earbuds One" and lines[0].rep == "Van A"


@test("data: the lists and their order; categories by the rule; long labels cut at a whole word")
def _():
    m = mdl.build(mdl.Period.month(MONTH), world_rows(), category_names=NAMES, upcoming=UPCOMING)
    assert m.period.prev_start == date(2026, 2, 1) and m.period.end == date(2026, 3, 31)
    assert m.reps == ["Van A", "Van B", "Causeway"], m.reps            # by March, then February, then name
    assert m.customers[0] == mdl.CASH_ACCOUNT and m.customers[1:] == ["Shop One", "Shop Two", "Shop Four", "Shop Three"]
    cat_of = {it.code: it.category for it in m.items.values()}
    assert cat_of["TB-D1"] == cats.BUNDLES and cat_of["WS-01"] == "TWS Earbuds"
    assert cat_of["WG-01"] == "Headphones & Neckbands" and cat_of["ZZ9"] == cats.OTHER
    assert cat_of["ZTE-X1"] == cats.HANDSET and [it.code for it in m.not_counted] == ["ZTE-X1"]
    assert m.categories[-1] == cats.OTHER and cats.HANDSET not in m.categories
    assert m.categories == ["TWS Earbuds", "Power Banks", "Cables", cats.BUNDLES, "Headphones & Neckbands",
                            cats.OTHER], m.categories
    long = "X99 " + "Demo Very Long Cable Name With Many Words + Extras " * 3
    label = mdl.dashboard_label("X99", long[4:])
    assert label.endswith(" ...") and mdl.label_width(label) <= mdl.LABEL_ROOM_PRODUCT
    assert not label[:-4].rstrip().endswith("+") and long.startswith(label[:-4])
    assert mdl.dashboard_label("X01", "Short") == "X01 Short"
    raises(ValueError, lambda: mdl.month_bounds("2026-13"))
    raises(ValueError, lambda: mdl.month_bounds("Sep 2026"))


@test("period: a week ends on a Saturday, against the four weeks before, the same days a month back and the month to date")
def _():
    # the Saturday -> the week, the four weeks before, the same seven days one month back, the month to date,
    # and the first day read: the four windows on every tile of the weekly workbook
    cases = {
        "2026-10-17": ("2026-10-11", "2026-10-17", "2026-09-13", "2026-10-10", "2026-09-11", "2026-09-17",
                       "2026-10-01", "2026-09-11"),                              # a week inside a month
        "2026-10-03": ("2026-09-27", "2026-10-03", "2026-08-30", "2026-09-26", "2026-08-27", "2026-09-02",
                       "2026-10-01", "2026-08-27"),                              # a week across two months
        "2026-10-31": ("2026-10-25", "2026-10-31", "2026-09-27", "2026-10-24", "2026-09-25", "2026-10-01",
                       "2026-10-01", "2026-09-25"),                              # a Saturday on the 31st
        "2026-06-06": ("2026-05-31", "2026-06-06", "2026-05-03", "2026-05-30", "2026-04-30", "2026-05-06",
                       "2026-06-01", "2026-04-30"),                              # Sunday the 31st: April has 30
        "2026-04-04": ("2026-03-29", "2026-04-04", "2026-03-01", "2026-03-28", "2026-02-28", "2026-03-06",
                       "2026-04-01", "2026-02-28"),                              # one month back is February's end
        "2027-01-02": ("2026-12-27", "2027-01-02", "2026-11-29", "2026-12-26", "2026-11-27", "2026-12-03",
                       "2027-01-01", "2026-11-27"),                              # the year's end
    }
    for sat, want in cases.items():
        p = mdl.Period.week(date.fromisoformat(sat))
        got = (p.start, p.end, p.prev_start, p.prev_end, p.lm_start, p.lm_end, p.mtd_start, p.data_first)
        assert got == tuple(date.fromisoformat(d) for d in want), (sat, got)
        assert p.kind == "week" and p.key == sat and (p.end - p.start).days == 6 == (p.lm_end - p.lm_start).days
    assert mdl.edate(date(2026, 3, 31), -1) == date(2026, 2, 28) and mdl.edate(date(2028, 3, 30), -1) == date(2028, 2, 29)
    assert mdl.edate(date(2026, 1, 15), -1) == date(2025, 12, 15) and mdl.edate(date(2026, 12, 31), 2) == date(2027, 2, 28)
    w = mdl.Period.week(date(2026, 10, 17))
    assert (w.label, w.long) == ("w/e 17 Oct 2026", "the week ending 17 October 2026")
    assert mdl.Period.week(date(2027, 1, 2)).label == "w/e 2 Jan 2027"
    assert mdl.Period.week(datetime(2026, 10, 17, 9, 30)) == w
    d = date.fromisoformat
    assert w.in_cur(d("2026-10-11")) and not w.in_cur(d("2026-10-10")) and w.in_prev(d("2026-10-10"))
    assert not w.in_prev(d("2026-09-12")) and w.in_lm(d("2026-09-11")) and not w.in_lm(d("2026-09-18"))
    assert w.in_mtd(d("2026-10-01")) and not w.in_mtd(d("2026-09-30")) and not w.in_mtd(d("2026-10-18"))
    for day, ends in (("2026-10-16", "2026-10-17"), ("2026-10-18", "2026-10-24")):
        said = str(raises(ValueError, lambda x=day: mdl.Period.week(d(x))))
        assert said.endswith(f"{day} is a {mdl.DAYS[d(day).weekday()]}, and the week that holds it ends on {ends}"), said
    raises(ValueError, lambda: mdl.Period.week("2026-10-17"))
    # a month is what it always was: no week window, its lines read from the month before
    m = mdl.Period.month(MONTH)
    assert (m.lm_start, m.lm_end, m.mtd_start, m.data_first) == (None, None, None, None)
    assert (m.key, m.label, m.long) == ("2026-03", "Mar 2026", "March 2026") and not m.in_lm(date(2026, 3, 5))
    # the week's lines: every window's, and none beyond; an item's comparison is the four weeks before only
    item = ITEMS["X01"]
    rows = [line(d("2026-09-10"), 1, 1, "Shop One", item, 1, 1.1, VAN_A),      # before every window: not read
            line(d("2026-09-11"), 2, 1, "Shop One", item, 2, 1.1, VAN_A),      # the same days last month only
            line(d("2026-09-20"), 3, 1, "Shop Two", item, 3, 1.1, VAN_B),      # the four weeks before
            line(d("2026-10-12"), 4, 1, "Shop One", item, 4, 1.1, VAN_A),      # the week
            line(d("2026-10-18"), 5, 1, "Shop One", item, 5, 1.1, VAN_A)]      # after the week: not read
    model = mdl.build(w, rows)
    assert [ln.sale_date for ln in model.lines] == [d("2026-09-11"), d("2026-09-20"), d("2026-10-12")]
    assert (model.items[item].cur_qty, model.items[item].prev_qty) == (4, 3)


# ── the eight sheets ───────────────────────────────────────────────────────────────────────────

@test("sheets: the eight sheets in the Sep 2026 order, every formula with a value and none an error")
def _():
    b, m, _f = built()
    assert [s.name for s in b.sheets] == ["Dashboard", "Salesmen", "Categories", "Products", "Customers", "Daily",
                                          "Data", "Notes"]
    assert not m.warnings, m.warnings
    n = 0
    for ws in b.sheets:
        for (_r, _c), cell in ws.cells.items():
            if cell.kind == "f":
                n += 1
                assert not isinstance(cell.value, fx.Err), (ws.name, cell.formula, cell.value)
    assert n > 300, n


@test("sheets: the Dashboard tiles are the hand sums (the handset left out of units, invoices and accounts)")
def _():
    b, m, _f = built()
    rows = world_rows()
    mar = [r for r in rows if r["sale_date"] >= "2026-03-01"]
    feb = [r for r in rows if r["sale_date"] < "2026-03-01"]
    net = lambda rs: sum(r["net_bhd"] or 0 for r in rs)  # noqa: E731
    acc = lambda rs: [r for r in rs if r["item_name"] != ITEMS["PHONE"]]  # noqa: E731
    assert abs(value(b, "Dashboard", "B9") - net(mar)) < 1e-9 and abs(value(b, "Dashboard", "C10") - net(feb)) < 1e-9
    assert value(b, "Dashboard", "F9") == sum(r["quantity"] for r in acc(mar)) == 71
    assert value(b, "Dashboard", "H9") == len({r["invoice_no"] for r in acc(mar)}) == 5
    assert value(b, "Dashboard", "O9") == 3                     # Shop One, Shop Two, Shop Four (cash never counts)
    assert value(b, "Dashboard", "Q9") == 3 and value(b, "Dashboard", "R10") == 3
    assert value(b, "Dashboard", "K3") == "Prepared 8 April 2026"
    assert value(b, "Dashboard", "B5") == "March 2026 compared with February 2026"
    foot = value(b, "Dashboard", "B78")
    assert foot == ("Units, invoices and accounts cover accessories only. ZTE-X1 handsets issued through the YQ "
                    "Roadshow carry no value in the day book (4 units in February, 6 in March) and are excluded."), foot
    assert "1 February to 31 March 2026" in value(b, "Dashboard", "B77")


@test("sheets: every sheet's total less the Data sheet is nil, and the Data sheet equals v_sales")
def _():
    b, _m, _f = built()
    top = find_row(b, "Notes", "B", "Sheet totals less Data (net sales)")
    for k in range(1, 6):
        assert value(b, "Notes", f"C{top + k}") == 0 and value(b, "Notes", f"D{top + k}") == 0
    exp = find_row(b, "Notes", "B", "Data sheet against the export file")
    for k in range(1, 5):
        assert value(b, "Notes", f"G{exp + k}") == 0, (k, value(b, "Notes", f"G{exp + k}"))
    assert value(b, "Notes", f"C{exp + 1}") == 5 and value(b, "Notes", f"D{exp + 1}") == 8   # every line, the zero-value ones too
    db_row = find_row(b, "Notes", "B", "Day book (Data sheet) less v_sales")
    assert value(b, "Notes", f"C{db_row}") == 0 and value(b, "Notes", f"D{db_row}") == 0
    left = find_row(b, "Notes", "B", "Narrations left out (they hold an ID number)")
    assert value(b, "Notes", f"C{left}") == 1 and value(b, "Notes", f"D{left}") == 1
    assert value(b, "Notes", f"E{left + 1}") == 0
    unmapped = find_row(b, "Notes", "B", "Day book lines with no category")
    assert value(b, "Notes", f"C{unmapped}") == 0 and value(b, "Notes", f"D{unmapped}") == 0


@test("sheets: the Data sheet's calculated columns: month, salesman, category, line type, units and the flags")
def _():
    b, m, _f = built()
    assert value(b, "Data", "A5") == fx.serial(date(2026, 2, 2)) and value(b, "Data", "K5") == fx.serial(date(2026, 2, 1))
    assert value(b, "Data", "B5") == "SI-YQ-26-02-1" and value(b, "Data", "L5") == "Van A"
    for i, ln in enumerate(m.lines):
        r = 5 + i
        handset = ln.item == ITEMS["PHONE"]
        assert value(b, "Data", f"O{r}") == ("Handset" if handset else "Accessory")
        assert value(b, "Data", f"P{r}") == (0 if handset else ln.qty)
        assert value(b, "Data", f"M{r}") == m.items[ln.item].category
        assert value(b, "Data", f"J{r}") == ("" if handset else (ln.narration or ""))
    assert formula_of(b, "Data", "Q6") == 'IF(O6="Accessory",IF(COUNTIFS($B$5:B6,B6,$O$5:O6,"Accessory")=1,1,0),0)'
    flags = [value(b, "Data", f"Q{5 + i}") for i in range(len(m.lines))]
    assert sum(flags) == len({ln.invoice for ln in m.lines if ln.item != ITEMS["PHONE"]})
    assert b.names["d_Net"] == f"Data!$H$5:$H${4 + len(m.lines)}"


@test("sheets: Salesmen, Categories, Products and Customers follow the lists, totals at the end")
def _():
    b, m, _f = built()
    S, P, C = len(m.reps), len(m.products), len(m.customers)
    tot = 6 + S
    assert value(b, "Salesmen", f"B{tot}") == "Total" and formula_of(b, "Salesmen", f"C{tot}") == f"SUM(C6:C{5 + S})"
    assert value(b, "Salesmen", "B6") == "Van A" and formula_of(b, "Salesmen", f"G6") == f"IF(D${tot}=0,0,D6/D${tot})"
    head = tot + 4
    assert value(b, "Salesmen", f"B{head}").startswith("Net sales by seven-day period")
    assert value(b, "Salesmen", f"C{head + 1}") == fx.serial(date(2026, 2, 1))
    assert len(b.get("Salesmen").sparklines) == S
    week_sum = sum(value(b, "Salesmen", f"{fx.col_letters(c)}{head + 2 + S}") for c in range(2, 2 + 9))
    assert abs(week_sum - value(b, "Salesmen", f"C{tot}") - value(b, "Salesmen", f"D{tot}")) < 1e-9
    cat_tot = 6 + len(m.categories)
    assert value(b, "Categories", f"B{cat_tot}") == "Total"
    assert value(b, "Categories", f"B{cat_tot - 1}") == cats.OTHER
    assert abs(value(b, "Categories", f"G{cat_tot}") - value(b, "Dashboard", "B9")) < 1e-9
    ptot = 6 + P
    assert value(b, "Products", f"B{ptot}") == "Total"
    assert value(b, "Products", f"B{ptot + 1}").startswith("Not included above: handsets")
    assert value(b, "Products", f"B{ptot + 2}") == "ZTE-X1" and value(b, "Products", f"F{ptot + 2}") == 6
    assert b.names["p_Cat"] == f"Products!$D$6:$D${ptot + 2}"
    assert sorted(value(b, "Products", f"N{6 + i}") for i in range(P)) == list(range(1, P + 1))
    ctot = 6 + C
    assert value(b, "Customers", "B6") == mdl.CASH_ACCOUNT and value(b, "Customers", "J6") == "Cash"
    assert value(b, "Customers", "K6") == "" and value(b, "Customers", "D6") == "Various"
    status_of = {value(b, "Customers", f"B{6 + i}"): value(b, "Customers", f"J{6 + i}") for i in range(C)}
    assert status_of == {mdl.CASH_ACCOUNT: "Cash", "Shop One": "Feb & Mar", "Shop Two": "Mar only",
                         "Shop Four": "Mar only", "Shop Three": "Feb only"}, status_of
    assert value(b, "Customers", "D7") == "Van A"                      # the salesman of the latest invoice
    assert value(b, "Customers", f"B{ctot}") == "Total"


@test("sheets: the Daily sheet has 31 days, the week blocks add up to each month, the days a month lacks are blank")
def _():
    b, _m, _f = built()
    assert value(b, "Daily", "B36") == 31 and formula_of(b, "Daily", "D36") == "SUMIFS(d_Net,d_Date,CurMonth+$B36-1)"
    assert formula_of(b, "Daily", "C34") is None and formula_of(b, "Daily", "C36") is None   # February has 28 days
    assert abs(value(b, "Daily", "I11") - value(b, "Daily", "C37")) < 1e-9
    assert abs(value(b, "Daily", "K11") - value(b, "Daily", "D37")) < 1e-9
    assert value(b, "Daily", "J11") == 1 and value(b, "Daily", "L11") == 1


@test("sheets: the commentary is DRAFT; the Notes sheet's sentences come from the data and the rule in force")
def _():
    b, _m, _f = built()
    for ref in b.commentary["Dashboard"]:
        assert str(value(b, "Dashboard", ref)).startswith("DRAFT"), ref
    assert str(value(b, "Customers", b.commentary["Customers"][0])).startswith("DRAFT")
    text = " ".join(str(c.value) for c in b.get("Notes").cells.values() if c.kind == "s")
    assert "Two invoices carry \"credit note\" in the narration" in text, text[:2000]
    assert "ZTE-X1 handsets issued through the YQ Roadshow carry no value in the day book (2 lines, 10 units)" in text
    assert "The three spellings of Cash Customer in the ERP are treated as one" in text
    assert "TB = bundle sets (a display group)" in text and "WG-01 = Headphones & Neckbands" in text
    assert "The narration on the handset invoices holds customer ID numbers" in text
    words = sc.Commentary(highlights=["h"] * 5, points=["p"] * 5)
    assert not words.is_draft and sc.Commentary().is_draft


@test("sheets: positions follow the list lengths (one more salesman moves every reference to the totals)")
def _():
    rows = world_rows() + [line(date(2026, 3, 20), 30, 1, "Shop Five", ITEMS["X01"], 1, 1.1, "Van C - Acc WH")]
    b, m, _f = built(rows)
    S = len(m.reps)
    assert S == 4 and value(b, "Salesmen", f"B{6 + S}") == "Total"
    top = find_row(b, "Notes", "B", "Sheet totals less Data (net sales)")
    assert formula_of(b, "Notes", f"C{top + 1}") == f"ROUND(Salesmen!C{6 + S}-SUMIFS(d_Net,d_Month,PrevMonth),2)"
    assert formula_of(b, "Dashboard", "Q9") == f'COUNTIF(Salesmen!D6:D{5 + S},">0")'
    assert value(b, "Categories", f"B{6 + len(m.categories) + 4 + 2 + S}") == "Total"


@test("sheets: a January report reaches back into December of the year before")
def _():
    rows = []
    for r in world_rows():
        r = dict(r)
        d = date.fromisoformat(r["sale_date"])
        r["sale_date"] = (date(2025, 12, d.day) if d.month == 2 else date(2026, 1, d.day)).isoformat()
        r["invoice_no"] = r["invoice_no"].replace("-02-", "-12-").replace("-03-", "-01-")
        rows.append(r)
    b, m, _f = build.build_month("2026-01", FakeSource(rows, to=date(2026, 2, 5)), prepared_on=date(2026, 2, 6),
                                 category_names=NAMES)
    assert m.period.span_text() == "1 December 2025 to 31 January 2026"
    assert value(b, "Notes", "C62") == "Dec 2025" and value(b, "Notes", "C61") == "Jan 2026"
    assert value(b, "Dashboard", "B5") == "January 2026 compared with December 2025"
    assert value(b, "Notes", "C6") == fx.serial(date(2025, 12, 1))
    text = " ".join(str(c.value) for c in b.get("Notes").cells.values() if c.kind == "s")
    assert "This file covers December 2025 and January 2026 only" in text
    assert not [c for ws in b.sheets for c in ws.cells.values() if c.kind == "f" and isinstance(c.value, fx.Err)]


# ── the file ───────────────────────────────────────────────────────────────────────────────────

@test("file: written through a temporary name, never over an existing file; the next free name; a valid workbook")
def _():
    b, m, _f = built()
    with data_root() as root:
        first = build.target_path(m.period, draft=True)
        assert first.name == "YQ Bahrain - Mobile Accessories Sales Report - Mar 2026 - DRAFT.xlsx"
        assert first.parent == root / "exports" / "reports" / "monthly"
        b.save(first)
        assert first.is_file() and not [p for p in first.parent.iterdir() if p.name.startswith("~build-")]
        raises(FileExistsError, lambda: b.save(first))
        second = build.target_path(m.period, draft=True)
        assert second.name.endswith("DRAFT (2).xlsx")
        with zipfile.ZipFile(first) as z:
            assert "xl/worksheets/sheet8.xml" in z.namelist() and len([n for n in z.namelist() if n.startswith("xl/charts/")]) == 4
            xml = z.read("xl/workbook.xml").decode("utf-8")
            assert 'calcPr' in xml and "fullCalcOnLoad" in xml       # the file is recalculated when it opens
        part = golden.read_book(first)
        assert part.sheet("Dashboard").pane is None and part.sheet("Data").pane["topLeftCell"] == "C5"
        assert part.print_areas["Dashboard"] == "Dashboard!$A$1:$S$80" and part.sheet("Data").autofilter.startswith("A4:S")


@test("file: the counts of rules, data bars, the colour scale, sparklines and charts")
def _():
    b, m, _f = built()
    with data_root():
        path = b.save(build.target_path(m.period, draft=True))
        part = golden.read_book(path)
    total = lambda what: sum(getattr(s, what) for s in part.sheets)  # noqa: E731
    assert (total("cell_rules"), total("data_bars"), total("color_scales"), total("charts")) == (20, 6, 1, 4)
    assert total("sparklines") == len(m.reps)


@test("file: the comparison with a reference sees a changed value and sets aside the cells it is told to")
def _():
    b, m, _f = built()
    with data_root() as root:
        path = b.save(build.target_path(m.period, draft=True))
        same = golden.compare(path, path)
        assert same.ok and all(s.ok for s in same.sheets)
        changed = root / "exports" / "changed.xlsx"
        with zipfile.ZipFile(path) as zin, zipfile.ZipFile(changed, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":          # the Dashboard: one cached value moved
                    text = data.decode("utf-8")
                    text = re.sub(r'(<c r="F9"[^>]*><f>[^<]*</f><v>)([^<]+)(</v>)', r"\g<1>12345\3", text, count=1)
                    data = text.encode("utf-8")
                zout.writestr(item, data)
        diff = golden.compare(changed, path)
        dash = next(s for s in diff.sheets if s.name == "Dashboard")
        assert not diff.ok and [v[0] for v in dash.value] == ["F9"], dash.value
        aside = golden.compare(changed, path, set_aside={"Dashboard": ["F9"]})
        assert aside.ok and next(s for s in aside.sheets if s.name == "Dashboard").set_aside == 1
        assert any("F9" in ln for ln in golden.summary(diff))


@test("file: no 9-digit or 15-plus-digit run in what the owner reads (the masked IDs are left out)")
def _():
    b, m, _f = built()
    with data_root():
        path = b.save(build.target_path(m.period, draft=True))
        with zipfile.ZipFile(path) as z:
            strings = z.read("xl/sharedStrings.xml").decode("utf-8")
    assert not LONG_RUN.search(strings), LONG_RUN.search(strings).group(0)
    assert "******" not in strings                                   # not even the masked form


# ── the pieces added for the stock workbook ────────────────────────────────────────────────────

GRID_FORMATS = ["g_head", "g_head_r", "g_text", "g_text_wrap", "g_int", "g_dec2", "g_money3", "g_pct", "g_date",
                "g_total_int", "g_total_money3", "input_qty", "input_note", "input_dec2", "input_money3"]


def grid_book() -> bk.Book:
    """A small grid on invented codes: typed-in quantities, a hidden helper column, a status list, totals,
    ROUNDUP formulas, rules in the new styles, a filter and frozen headings."""
    b = bk.Book()
    for name in style.DXF_EXTRA:
        b.dxf(name)                                           # the styles of style.py, by name
    b.dxf("cf_mine", {"bg_color": "#00FF00", "bold": True})  # and one of this book's own
    ws = b.sheet("Grid")
    ws.col(0, 0, 14)
    ws.col(1, 1, 10)
    ws.col(2, 2, 12, hidden=True)                             # a helper nobody needs to read
    ws.col(3, 5, 12)
    heads = [("Code", "g_head"), ("Qty", "g_head_r"), ("Helper", "g_head"), ("Cost", "g_head_r"),
             ("Status", "g_head"), ("Note", "g_head")]
    for c, (text, fmt) in enumerate(heads):
        ws.text(3, c, text, fmt)
    rows = [("T01", 5, "x", 1.2345, "ok"), ("T02", 0, "y", 2.5, "low"), ("T03", 12, "z", 0.5, "ok")]
    for i, (code, qty, helper, cost, status) in enumerate(rows):
        r = 4 + i
        ws.text(r, 0, code, "g_text")
        ws.put(r, 1, qty, "input_qty")
        ws.text(r, 2, helper, "g_text")
        ws.put(r, 3, cost, "g_money3")
        ws.text(r, 4, status, "g_text")
        ws.put(r, 5, None, "input_note")
    ws.text(7, 0, "Total", "g_text")
    ws.formula(7, 1, "SUM(B5:B7)", "g_total_int")
    ws.formula(7, 3, "SUM(D5:D7)", "g_total_money3")
    ws.formula(9, 1, "ROUNDUP(D5*10,0)", "g_int")             # 12.345 goes up to 13
    ws.formula(9, 3, "ROUNDUP((0.1+0.2)*10,0)", "g_int")      # 3 as it is shown, not 4
    ws.formula(9, 4, "ROUNDUP(1234.5,-2)", "g_dec2")          # 1300
    ws.put(10, 0, 0.256, "g_pct")
    ws.put(10, 1, date(2026, 10, 5), "g_date")
    ws.put(10, 3, 2.5, "input_dec2")
    ws.put(10, 4, 0.125, "input_money3")
    ws.put(10, 5, "a note that wraps onto a second line", "g_text_wrap")
    ws.cond("B5:B7", {"type": "cell", "criteria": "less than", "value": 1, "format": "cf_red"})
    ws.cond("B5:B7", {"type": "cell", "criteria": "greater than", "value": 10, "format": "cf_green"})
    ws.cond("E5:E7", {"type": "cell", "criteria": "equal to", "value": '"low"', "format": "cf_amber"})
    ws.cond("F5:F7", {"type": "blanks", "format": "cf_grey"})
    ws.cond("A5:A7", {"type": "cell", "criteria": "equal to", "value": '"T03"', "format": "cf_mine"})
    ws.validate(4, 1, 6, 1, {"validate": "integer", "criteria": ">=", "value": 0, "error_title": "Quantity",
                              "error_message": "Whole units, nought or more."})
    ws.validate(4, 4, 6, 4, {"validate": "list", "source": ["ok", "low", "out"]})
    ws.page.update({"freeze": (4, 1), "autofilter": "A4:F7"})      # frozen headings and a filter: the page keys
    return b


@test("grid: validation, a hidden column, a book's own conditional styles, the grid and input formats, freeze and filter")
def _():
    import openpyxl
    b = grid_book()
    assert b.evaluate() == []
    with data_root() as root:
        path = b.save(paths.reports_dir("stock") / "Grid test.xlsx")
        assert path.parent == root / "exports" / "reports" / "stock"
        ws = openpyxl.load_workbook(path)["Grid"]
        shown = openpyxl.load_workbook(path, data_only=True)["Grid"]
    assert ws.freeze_panes == "B5" and ws.auto_filter.ref == "A4:F7"
    cols = ws.column_dimensions                                   # the hidden column keeps its width
    assert cols["C"].hidden is True and cols["C"].width == 12
    assert not cols["A"].hidden and cols["A"].width == 14 and not cols["D"].hidden and not cols["B"].hidden
    dv = {str(v.sqref): v for v in ws.data_validations.dataValidation}
    assert sorted(dv) == ["B5:B7", "E5:E7"], sorted(dv)
    qty, status = dv["B5:B7"], dv["E5:E7"]
    assert (qty.type, qty.operator, qty.formula1) == ("whole", "greaterThanOrEqual", "0")
    assert qty.errorTitle == "Quantity" and qty.error == "Whole units, nought or more." and qty.showErrorMessage
    assert status.type == "list" and status.formula1 == '"ok,low,out"'
    rules = {(str(cf.sqref), r.operator or r.type): r.dxf for cf in ws.conditional_formatting for r in cf.rules}
    assert len(rules) == 5, sorted(rules)
    fill = lambda d: d.fill.bgColor.rgb  # noqa: E731
    red = rules[("B5:B7", "lessThan")]
    assert fill(red) == "FFF8D7DA" and red.font.color.rgb == "FF721C24"
    assert fill(rules[("B5:B7", "greaterThan")]) == "FFD4EDDA" and fill(rules[("E5:E7", "equal")]) == "FFFFF3CD"
    assert rules[("F5:F7", "containsBlanks")].font.color.rgb == "FF6C757D"
    mine = rules[("A5:A7", "equal")]                              # the style only this book registered
    assert fill(mine) == "FF00FF00" and mine.font.b
    written = {(r, c): cell for (r, c), cell in b.get("Grid").cells.items() if cell.fmt in GRID_FORMATS}
    assert {cell.fmt for cell in written.values()} == set(GRID_FORMATS)         # every new format is in use
    for r, c in written:                                          # a thin grey border all round, 9.5 points
        cell = ws[bk.a1(r, c)]
        edge = cell.border
        assert (edge.left.style, edge.right.style, edge.bottom.style) == ("thin", "thin", "thin"), bk.a1(r, c)
        assert edge.top.style == ("medium" if written[(r, c)].fmt.startswith("g_total") else "thin"), bk.a1(r, c)
        assert edge.left.color.rgb == "FFC3C2B7" and cell.font.sz == 9.5, bk.a1(r, c)
    typed = ws["B5"]                                              # a typed-in quantity: yellow, blue, whole number
    assert typed.fill.fill_type == "solid" and typed.fill.fgColor.rgb == "FFFFF2CC"
    assert typed.font.color.rgb == "FF0000FF" and typed.number_format == '#,##0;-#,##0;"–"'
    assert ws["F5"].fill.fgColor.rgb == "FFFFF2CC" and ws["F5"].alignment.wrap_text         # the note: yellow, wraps
    assert ws["D11"].number_format == '#,##0.00;-#,##0.00;"–"' and ws["E11"].number_format == '#,##0.000;-#,##0.000;"–"'
    assert ws["D11"].fill.fgColor.rgb == "FFFFF2CC" and ws["E11"].font.color.rgb == "FF0000FF"
    assert ws["D5"].number_format == '#,##0.000;-#,##0.000;"–"' and ws["A11"].number_format == "0.0%"
    assert ws["B11"].number_format == "d mmm yyyy" and ws["F11"].alignment.wrap_text
    head = ws["A4"]
    assert head.font.b and head.alignment.wrap_text and head.fill.fgColor.rgb == "FFE1E0D9"
    assert ws["B4"].alignment.horizontal == "right" and head.alignment.horizontal is None
    total = ws["B8"]
    assert total.font.b and total.border.top.color.rgb == "FF0B0B0B" and ws["D8"].font.b
    assert ws["D8"].number_format == '#,##0.000;-#,##0.000;"–"'
    assert ws["B10"].value == "=ROUNDUP(D5*10,0)"                 # the formula stays a formula ...
    assert shown["B10"].value == 13 and shown["D10"].value == 3 and shown["E10"].value == 1300      # ... with its value
    assert shown["B8"].value == 17 and abs(shown["D8"].value - 4.2345) < 1e-9


@test("grid: a rule naming a style nobody registered, a validation XlsxWriter would skip and a name clash are refused")
def _():
    import warnings
    b = bk.Book()
    ws = b.sheet("S")
    ws.put(0, 0, 1)
    ws.cond("A1", {"type": "cell", "criteria": "less than", "value": 1, "format": "cf_red"})     # never registered
    with data_root() as root:
        target = root / "exports" / "never.xlsx"
        exc = raises(KeyError, lambda: b.save(target))
        assert "S!A1" in str(exc) and "cf_red" in str(exc) and "Book.dxf" in str(exc), exc
        assert not target.exists() and not list((root / "exports").iterdir())     # nothing written, no temp left
        b.dxf("cf_red")
        assert b.save(target) == target and target.is_file()                       # registered: the same book writes
        ws.validate(1, 0, 1, 0, {"validate": "list", "source": ["a" * 300]})        # past Excel's 255 characters
        later = root / "exports" / "later.xlsx"
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")                       # XlsxWriter only warns and skips the rule
            exc = raises(ValueError, lambda: b.save(later))
        assert "S!A2" in str(exc) and "data validation" in str(exc), exc
        assert not later.exists() and [p.name for p in (root / "exports").iterdir()] == ["never.xlsx"]
    raises(ValueError, lambda: bk.Book().dxf("up", {"bg_color": "#FF0000"}))      # a style name of the monthly report
    raises(KeyError, lambda: bk.Book().dxf("cf_unknown"))
    two = bk.Book()
    two.dxf("cf_amber")
    two.dxf("cf_red")
    assert list(two.dxf_styles) == ["cf_amber", "cf_red"]                         # kept in the order registered


@test("grid: every new format has a thin grey border all round and 9.5 points; the conditional styles stay out of DXF")
def _():
    for name in GRID_FORMATS:
        p = style.props(name)
        assert list(p)[0] == "border" and p["border"] == 1 and p["border_color"] == "#C3C2B7", name   # first: a 'top' after it wins
        assert p["font_size"] == 9.5, name
        yellow_and_blue = p.get("bg_color") == "#FFF2CC" and p.get("font_color") == "#0000FF"
        assert yellow_and_blue == name.startswith("input_"), name              # the typed-in cells, and only they
    assert style.props("g_total_int")["top"] == 2 and style.props("g_total_money3")["bold"]
    assert style.props("g_head")["text_wrap"] and style.props("g_head_r")["align"] == "right"
    assert set(style.DXF) == {"na", "up", "down"}                                # the three that go into every file
    assert set(style.DXF_EXTRA) == {"cf_red", "cf_amber", "cf_green", "cf_grey"}
    assert style.DXF_EXTRA["cf_red"]["bg_color"] == "#F8D7DA" and "font_color" in style.DXF_EXTRA["cf_red"]
    assert style.DXF_EXTRA["cf_amber"]["bg_color"] == "#FFF3CD" and style.DXF_EXTRA["cf_green"]["bg_color"] == "#D4EDDA"
    assert style.DXF_EXTRA["cf_grey"]["font_color"] == "#6C757D"


@test("grid: the monthly report's file carries none of it: three conditional styles, no hidden column, no validation")
def _():
    b, m, _f = built()
    assert b.dxf_styles == {} and all(not ws.hidden_cols and not ws.validations for ws in b.sheets)
    with data_root():
        path = b.save(build.target_path(m.period, draft=True))
        with zipfile.ZipFile(path) as z:
            styles = z.read("xl/styles.xml").decode("utf-8")
            sheets = [z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("xl/worksheets/sheet")]
    assert '<dxfs count="3">' in styles and len(sheets) == 8
    assert not any("hidden=" in xml or "dataValidation" in xml for xml in sheets)


# ── the command ────────────────────────────────────────────────────────────────────────────────

LAST_LOAD ={"run": "080426_0900", "finished_at": "2026-04-08T09:00:00+03:00", "result": "ok", "data_date": "2026-04-07"}


@test("command: a whole month builds the DRAFT workbook and a record, and says what it checked")
def _():
    with data_root(last_load=LAST_LOAD) as root:
        src = FakeSource()
        out = build.run(Args(), source=src)
        assert out["ok"] and out.get("exit") == 0, out["text"]
        assert out["data_date"] == "2026-04-07" and out["flags"]["month"] == MONTH and out["flags"]["draft"]
        path = Path(out["files"][0])
        assert path.is_file() and path.parent == root / "exports" / "reports" / "monthly"
        record = json.loads(Path(out["files"][1]).read_text(encoding="utf-8"))
        assert record["month"] == MONTH and record["checks_ok"] and record["draft"]
        text = out["text"]
        assert text.startswith("Monthly sales report, March 2026, against February 2026")
        assert "Checks: the Data sheet equals v_sales in both months" in text and "DRAFT: the ten comment lines" in text
        assert not LONG_RUN.search(text)
        assert sorted(set(src.calls)) == ["anchors", "lines", "totals", "upcoming"]
        again = build.run(Args(), source=FakeSource())
        assert Path(again["files"][0]).name.endswith("DRAFT (2).xlsx")      # the first file is kept


@test("command: not a whole month yet (no load after the month, or data short of its end): nothing is written")
def _():
    early = {"run": "250326_0900", "finished_at": "2026-03-25T09:00:00+03:00", "result": "ok"}
    with data_root(last_load=early) as root:
        out = build.run(Args(), source=FakeSource(to=date(2026, 3, 24)))
        assert not out["ok"] and out["exit"] == 2 and "not whole yet" in out["text"], out["text"]
        assert not (root / "exports" / "reports").exists()
    with data_root() as root:                 # no load on record: the database's own upload time stands in
        out = build.run(Args(), source=FakeSource(uploaded="2026-03-30T10:00:00+03:00"))
        assert not out["ok"] and "newest successful upload" in out["text"]
        out = build.run(Args(), source=FakeSource())
        assert out["ok"], out["text"]


@test("command: the week, a bad month and a month with no sales are refused in one sentence")
def _():
    import importlib.util
    with data_root(last_load=LAST_LOAD):
        if importlib.util.find_spec(build.WEEK_MODULE) is None:          # until the weekly workbook's module exists
            for words in (("week",), ("week", "2026-03-28")):
                out = build.run(Args(spec=words), source=FakeSource())
                assert out["exit"] == 1 and "later phase" in out["text"]
        for words, said in ((("week", "2026-03-27"), "2026-03-27 is a Friday, and the week that holds it ends on "
                                                     "2026-03-28"),
                            (("week", "28-03-2026"), "written YYYY-MM-DD"), (("week", "2026-02-30"), "written YYYY-MM-DD"),
                            (("week", "2026-03-28", "2026-04-04"), "one week only")):
            out = build.run(Args(spec=words), source=FakeSource())
            assert out["exit"] == 1 and said in out["text"] and "\n" not in out["text"], out["text"]
        out = build.run(Args(spec=("month", "2026-13")), source=FakeSource())
        assert out["exit"] == 1 and "YYYY-MM" in out["text"]
        out = build.run(Args(spec=("monthly",)), source=FakeSource())
        assert out["exit"] == 1 and "report month" in out["text"]
        out = build.run(Args(spec=("month", "2026-01")), source=FakeSource())
        assert out["exit"] == 2 and "nothing to report" in out["text"], out["text"]
    assert build.parse_spec([]) == ("month", None) and build.parse_spec(["month 2026-03"]) == ("month", "2026-03")
    assert build.parse_spec(["week"]) == ("week", None) and build.parse_spec(["WEEK", "2026-03-28"]) == \
        ("week", date(2026, 3, 28))
    assert build.default_month(date(2026, 4, 7)) == "2026-03" and build.default_month(date(2026, 1, 3)) == "2025-12"


@test("command: report week is handed to the weekly workbook's module once it exists, the Saturday as a date")
def _():
    import importlib.machinery
    import types
    seen = []
    stand_in = types.ModuleType(build.WEEK_MODULE)
    stand_in.__spec__ = importlib.machinery.ModuleSpec(build.WEEK_MODULE, None)
    stand_in.run = lambda args, saturday: seen.append(saturday) or {"ok": True, "exit": 0, "text": "built"}
    saved = sys.modules.get(build.WEEK_MODULE)
    sys.modules[build.WEEK_MODULE] = stand_in
    try:
        with data_root(last_load=LAST_LOAD):
            src = FakeSource()
            assert build.run(Args(spec=("week", "2026-03-28")), source=src) == {"ok": True, "exit": 0, "text": "built"}
            assert build.run(Args(spec=("week",)), source=src)["text"] == "built"
            out = build.run(Args(spec=("week", "2026-03-27")), source=src)          # not a Saturday: refused first
            assert out["exit"] == 1 and "Saturday" in out["text"]
            assert src.calls == [], src.calls                     # the week reads through its own module, not this one
    finally:
        if saved is None:
            sys.modules.pop(build.WEEK_MODULE, None)
        else:
            sys.modules[build.WEEK_MODULE] = saved
    assert seen == [date(2026, 3, 28), None], seen


@test("command: names a spreadsheet reads as one (capitals) or as a pattern are said, and the controls catch the sum")
def _():
    rows = world_rows() + [line(date(2026, 3, 5), 40, 1, "shop one", ITEMS["X01"], 1, 1.1, VAN_A),
                           line(date(2026, 3, 6), 41, 1, "Shop Six", "X77 Box 10*20*30MM (VFAN)", 1, 1.1, VAN_A)]
    with data_root(last_load=LAST_LOAD):
        out = build.run(Args(), source=FakeSource(rows))
    assert not out["ok"] and out["exit"] == 2, out["text"]
    assert "CHECKS TO LOOK AT: Customers sheet" in out["text"], out["text"]       # SUMIFS adds the two spellings
    assert "Note: two customer names differ only in capitals" in out["text"]
    assert "Note: an item name holds * ? or ~" in out["text"]


@test("command: without the Upcoming cards the build goes on and says so")
def _():
    b, m, _f = build.build_month(MONTH, FakeSource(upcoming_fails=True), prepared_on=date(2026, 4, 8),
                                 category_names=NAMES)
    assert any("Upcoming cards could not be read" in w for w in m.warnings)
    assert m.items[ITEMS["WS"]].category == cats.OTHER                  # no card: the name table has no key for it
    assert m.items[ITEMS["WG"]].category == "Headphones & Neckbands"     # the code key still holds


# ── the real September file (local only) ───────────────────────────────────────────────────────

@test("golden: September 2026 from production against the reference (SKIP without the file or YQ_OPS_GOLDEN=1)")
def _():
    if not REFERENCE.is_file():
        print(f"      SKIP: the reference workbook is not on this machine ({REFERENCE.name})")
        return
    if os.environ.get("YQ_OPS_GOLDEN") != "1":
        print("      SKIP: set YQ_OPS_GOLDEN=1 to rebuild September 2026 from production (read-only) and compare")
        return
    env = {k: v for k, v in os.environ.items() if k not in ("SUPABASE_URL", "SUPABASE_KEY", "DATABASE_URL")}
    env["PYTHONUTF8"] = "1"
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
        cmd = [sys.executable, "-m", "scripts.ops", "report", "month", "2026-09", "--data-root", tmp,
               "--prepared-on", "2026-10-04", "--compare", str(REFERENCE), "--json"]
        if os.environ.get("YQ_OPS_GOLDEN_BY"):
            cmd += ["--prepared-by", os.environ["YQ_OPS_GOLDEN_BY"]]
        proc = subprocess.run(cmd, cwd=str(ROOT), env=env, capture_output=True, text=True, encoding="utf-8",
                              timeout=900)
    last = (proc.stdout.strip().splitlines() or [""])[-1]
    try:
        out = json.loads(last)
    except ValueError:
        print("      SKIP: the database could not be read: " + (proc.stdout + proc.stderr).strip()[-200:])
        return
    print("\n".join("      " + ln for ln in out.get("text", "").splitlines()[-12:]))
    assert out.get("flags", {}).get("golden_ok") is True, out.get("text")


def main() -> int:
    passed = failed = 0
    for name, fn in TESTS:
        try:
            fn()
            print(f"PASS  {name}")
            passed += 1
        except Exception:  # noqa: BLE001
            print(f"FAIL  {name}")
            traceback.print_exc()
            failed += 1
    paths.set_data_root(None)
    config.use(None)
    print(f"\n{passed} passed, {failed} failed")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
