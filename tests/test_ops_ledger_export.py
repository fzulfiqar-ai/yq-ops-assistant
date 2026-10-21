"""Ops engine: the month-end Ledger export reader (scripts/ops/ledger_export.py).

    python -m tests.test_ops_ledger_export

Same lightweight runner as the other suites (no pytest). No database, no network, no .env (SUPABASE_URL /
SUPABASE_KEY / DATABASE_URL are pinned to ci.invalid before any import). The Ledger is a workbook built here,
in a temporary folder, in the layout Focus exports (title block with "[As on date ...]", the header row, one
section per account with its Sub Total, an Opening Balance row, a Grand Total) and read by the REAL reader and
parser of scripts/ingest.py. It is cumulative, like the export: September and November entries stand beside
October's. All names and numbers are invented.

Covered:
  * kind_of and prefix_of: SI sale, RCTV receipt, SRV return, CN credit_note, PV payment, JV journal, anything
    else (a stock voucher, no prefix, blank) other;
  * read: the rows of a cumulative file in the shape the module note gives; Sub Total, Grand Total and
    Opening Balance rows are no rows; money is Decimal to the fils; the narration, the ID-shaped numbers and
    the card-shaped numbers typed into it are in no row, no sentence and no file (nothing is written); the
    noise payment modes ('0' as text and as a number, 'Payment Mode') are None; unknown prefixes are listed as
    "type not known: <prefix> (n rows)" with their count, a voucher with no prefix is noted; Focus's Grand
    Total is kept and compared with the rows; a changed header, a file that cannot be opened, a title without
    an "As on date" and a missing file each stop the read in words;
  * movements: a cumulative file filtered by date (September and November stay out of October), per account
    and kind, every kind kept; windows that are wrong or run past the as-on date are refused;
  * balance: the running balance at the end of a day (two rows on one day, before the first row, an Opening
    Balance row carried, a credit balance, a name spelt another way, an account that is not there, a day past
    the as-on date); opening + movements = closing for every customer;
  * balance, blank cells: the real export leaves the Balance off some credit legs (the payment inside a cash
    sale, a run of receipts, an account's last row). The balance is worked forward from the last stated one,
    back from the next one for an account's first rows, and None only for an account that states none;
    blanking cells changes no balance on any account on any day; a stated balance is never replaced;
  * balance, no sign: the real export prints a magnitude and leaves a nought blank. balance() is signed (a
    credit balance negative), the sign taken from the rows, an opening balance above the first row found from
    the figures; written three ways (signed, as the export prints it, and with credit legs blank as well) every
    account gives the same balance on every day; when nothing ties the figures to the rows they stay as printed;
  * customer_legs: the accounts the caller passes, matched by name however spelt, with or without a window;
  * returns_ex_vat_by_salesman and returns_by_voucher: a return's VAT from its own legs, else by rate; the
    salesman from any leg; no salesman last; windows;
  * find: the Ledger as on the month's last day, else the smallest as-on after it; an earlier as-on refused with
    its reason; the folders searched (the last day to ten days after); names that are not the Ledger, lock
    files and sub-folders left alone; a file that cannot be opened; the newest of two with one as-on; a named
    file overrides the search and not the as-on rule; the year's end;
  * the real export, when YQ_OPS_LEDGER_SAMPLE names one (SKIP otherwise): the header labels, the as-on date
    and the voucher prefixes are what the reader expects, and every run of blank Balance cells worked forward
    lands on the next stated balance; aggregate counts are printed (the blank cells, the balances printed
    without a sign, and for each month the customer accounts whose opening + movements close on balance()),
    never a name, an account, a voucher number or an amount.
"""
from __future__ import annotations

import io
import itertools
import os
import re
import sys
import tempfile
import time
import traceback
from collections import Counter
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

os.environ["SUPABASE_URL"] = "https://ci.invalid"
os.environ["SUPABASE_KEY"] = "ci-key"
os.environ["DATABASE_URL"] = "postgresql://ci:ci@ci.invalid:5432/ci"

from scripts.ops import ledger_export as le  # noqa: E402
from scripts.ops import paths  # noqa: E402

TESTS: list[tuple[str, object]] = []

NINE = re.compile(r"(?<![0-9])[0-9]{9}(?![0-9])")            # shaped like a personal ID number
LONG = re.compile(r"(?<![0-9])[0-9]{15,}(?![0-9])")          # shaped like a card or account number
OLD_XLS = bytes.fromhex("d0cf11e0a1b11ae1") + bytes(504)      # how an old-style .xls begins: not a zip, so not an .xlsx

# Invented digit runs, put together here so that this file holds none: two that look like personal ID numbers
# and one that looks like a card number. Every narration also carries a word found nowhere else.
ID_A = "12345" + "6789"
ID_B = "98765" + "4789"
CARD = "4" + "1" * 15
WORDS = ("zebra", "walrus", "quokka", "yak", "ibis", "lemur")

HEADER = ["Date", "Voucher", "Account", "Debit", "Credit", "Balance", "Currency", "Payment Mode", "Salesman",
          "Narration"]

CUST_A, CUST_B, CUST_C = "Customer A", "Customer B", "Customer C"
CUST_D, CUST_E, CUST_F = "Customer D", "Customer E", "Customer F"
SALES, RETURNS, VAT, TAXCTL = "Sales Accessories", "Sales Returns", "Output VAT", "Tax Control"
CASH, BANK, SUPPLIER = "Cash in Hand", "Main Bank", "Supplier One"
REP1, REP2 = "Rep One", "Rep Two"
ORDER = [CUST_A, CUST_B, CUST_C, SALES, RETURNS, VAT, TAXCTL, CASH, BANK, SUPPLIER]
ORDER_ALL = ORDER + [CUST_D, CUST_E, CUST_F]
AS_ON = date(2026, 11, 3)

_narrations = itertools.count()


def d(month: int, dom: int) -> datetime:
    return datetime(2026, month, dom)


def post(when, voucher, legs, mode=None):
    """One voucher: (day, voucher, legs, payment mode, narration). A leg is (account, debit, credit) or
    (account, debit, credit, salesman). The narration holds a word, two ID-shaped numbers and, now and
    then, a card-shaped one."""
    k = next(_narrations)
    text = f"{WORDS[k % len(WORDS)]} {ID_A} {ID_B}" + (f" card {CARD}" if k % 3 == 0 else "")
    return (when, voucher, legs, mode, text)


POSTINGS = [
    # September: before the month
    post(d(9, 15), "SRV : 6", [(RETURNS, 40, 0), (VAT, 4, 0), (CUST_B, 0, 44, REP2)]),
    post(d(9, 20), "SI : 100", [(CUST_A, 220, 0, REP1), (SALES, 0, 200), (VAT, 0, 20)], "Credit"),
    post(d(9, 25), "SI : 101", [(CUST_B, 55, 0, REP2), (SALES, 0, 50), (VAT, 0, 5)], "Credit"),
    # October
    post(d(10, 5), "SI : 102", [(CUST_A, 330, 0, REP1), (SALES, 0, 300), (VAT, 0, 30)], "Credit"),
    post(d(10, 5), "SI : 103", [(CUST_A, 110, 0, REP1), (SALES, 0, 100), (VAT, 0, 10)], "Credit"),
    post(d(10, 7), "SI : 104", [(CUST_B, 550, 0, REP2), (SALES, 0, 500), (VAT, 0, 50)], "Credit"),
    post(d(10, 12), "RCTV : 50", [(CASH, 200, 0), (CUST_A, 0, 200, REP1)], "Cash"),
    post(d(10, 20), "SRV : 7", [(RETURNS, 100, 0), (VAT, 10, 0), (CUST_A, 0, 110, REP1)]),
    post(d(10, 22), "SRV : 8", [(RETURNS, 50, 0), (VAT, 5, 0), (CUST_B, 0, 55, REP2)]),
    post(d(10, 25), "CN : 3", [(SALES, 20, 0), (CUST_A, 0, 20, REP1)]),
    post(d(10, 28), "SRV : 9", [(RETURNS, 200, 0, REP1), (TAXCTL, 20, 0), (CUST_A, 0, 220)]),   # VAT not named VAT;
    #                                                                       the salesman stands on the returns leg only
    post(d(10, 29), "SRV : 10", [(RETURNS, 30, 0), (VAT, 3, 0), (CUST_C, 0, 33)]),              # no salesman
    post(d(10, 30), "RCTV : 51", [(BANK, 550, 0), (CUST_B, 0, 550, REP2)], "0"),
    post(d(10, 30), "PV : 9", [(SUPPLIER, 1000, 0), (BANK, 0, 1000)], "Payment Mode"),
    post(d(10, 31), "JV : 1", [(CASH, 10, 0), (BANK, 0, 10)], 0),
    post(d(10, 31), "XYZ : 5", [(CASH, 5, 0), (BANK, 0, 5)]),
    post(d(10, 31), "xyz : 6", [(CASH, 5, 0), (BANK, 0, 5)]),
    post(d(10, 31), "ABC : 1", [(CASH, 5, 0), (BANK, 0, 5)]),
    post(d(10, 31), "77", [(CASH, 1, 0), (BANK, 0, 1)]),                                         # no prefix
    # November: the report is cumulative up to its as-on date
    post(d(11, 2), "SI : 110", [(CUST_A, 55, 0, REP1), (SALES, 0, 50), (VAT, 0, 5)], "Credit"),
    post(d(11, 2), "SRV : 11", [(RETURNS, 10, 0), (VAT, 1, 0), (CUST_B, 0, 11, REP2)]),
]
OPENING = {CUST_B: 40}                      # a row above the account's first entry, as the report can write it
N_ROWS = sum(len(p[2]) for p in POSTINGS)
TOTAL_DEBIT = sum(leg[1] for p in POSTINGS for leg in p[2])
TOTAL_CREDIT = sum(leg[2] for p in POSTINGS for leg in p[2])
ROW_KEYS = {"account", "date", "voucher", "prefix", "kind", "counter", "debit", "credit", "balance", "currency",
            "payment_mode", "salesman"}

SMALL = [post(d(10, 5), "SI : 1", [(CUST_A, 11, 0, REP1), (SALES, 0, 11)])]

# What the real export does: a credit leg can show no Balance, the payment inside a cash sale being the usual one
# (the sale debits the customer and the payment credits it, in one voucher; the report prints the balance after the
# debit and leaves it off the credit). The writer's `blank` rule below leaves the credit legs without one.
RUNS = [
    # D: two cash sales, a run of two receipts, a sale on credit, and a last receipt: the account's last row
    # shows no balance either
    post(d(10, 3), "SI : 201", [(CUST_D, 110, 0, REP1), (CUST_D, 0, 110, REP1), (SALES, 0, 100), (VAT, 0, 10),
                                (CASH, 110, 0)], "Cash"),
    post(d(10, 4), "SI : 202", [(CUST_D, 55, 0, REP1), (CUST_D, 0, 55, REP1), (SALES, 0, 50), (VAT, 0, 5),
                                (CASH, 55, 0)], "Cash"),
    post(d(10, 5), "RCTV : 61", [(CASH, 20, 0), (CUST_D, 0, 20, REP1)], "Cash"),
    post(d(10, 5), "RCTV : 62", [(CASH, 30, 0), (CUST_D, 0, 30, REP1)], "Cash"),
    post(d(10, 6), "SI : 203", [(CUST_D, 110, 0, REP1), (SALES, 0, 100), (VAT, 0, 10)], "Credit"),
    post(d(10, 7), "RCTV : 63", [(CASH, 60, 0), (CUST_D, 0, 60, REP1)], "Cash"),
    # E: its first row is a receipt, then a sale
    post(d(10, 3), "RCTV : 64", [(CASH, 20, 0), (CUST_E, 0, 20, REP2)], "Cash"),
    post(d(10, 4), "SI : 204", [(CUST_E, 120, 0, REP2), (SALES, 0, 109), (VAT, 0, 11)], "Credit"),
    # F: one receipt, and nothing else
    post(d(10, 5), "RCTV : 65", [(CASH, 5, 0), (CUST_F, 0, 5, REP2)], "Cash"),
]


def blank_credits(account, voucher, debit, credit) -> bool:
    """The writer's `blank` rule: every credit leg shows no Balance."""
    return credit > 0


def test(name):
    def deco(fn):
        TESTS.append((name, fn))
        return fn
    return deco


def raises(kind, fn) -> Exception:
    try:
        fn()
    except kind as exc:
        return exc
    raise AssertionError(f"expected {kind.__name__}")


def write_ledger(path: Path, postings, *, as_on: str = "03/11/2026", opening: dict | None = None,
                 header: list[str] | None = None, title_line: str | None = None, lead: list | None = None,
                 grand: tuple | None = None, order: list[str] | None = None, balances: bool = True,
                 blank=None, magnitude: bool = False) -> None:
    """A Ledger workbook as Focus exports it: the title block (company, 'Ledger', the as-on line), the header
    row, then per account its name, an optional Opening Balance row, its rows in date order with a running
    balance and a Sub Total; a Grand Total closes the sheet. The Account column of a row names the other side
    of its voucher. `lead` rows stand above the first account; `balances=False` leaves the Balance cells
    blank; `blank(account, voucher, debit, credit)` leaves the Balance cell of the rows it says yes to empty
    (the running balance goes on without them, as in the real export); `magnitude=True` prints the balance as
    the real export does, without its sign, and leaves it empty when it is nought."""
    from openpyxl import Workbook

    legs_of: dict[str, list[tuple]] = {}
    for n, (when, voucher, legs, mode, text) in enumerate(postings):
        for k, leg in enumerate(legs):
            account, debit, credit = leg[0], leg[1], leg[2]
            salesman = leg[3] if len(leg) > 3 else None
            counter = next((o[0] for j, o in enumerate(legs) if j != k and o[0] != account), None)
            legs_of.setdefault(account, []).append((when, n, k, voucher, counter, debit, credit, salesman, mode, text))
    wb = Workbook()
    ws = wb.active
    ws.title = "Ledger"
    pad = [None] * 7
    ws.append(pad + ["Example Trading W.L.L"])
    ws.append(pad + ["Ledger"])
    ws.append([])
    ws.append(pad + [title_line if title_line is not None else f"[As on date {as_on}]"])
    ws.append(list(HEADER) if header is None else header)
    for row in lead or []:
        ws.append(row)

    def printed(value):
        """What the report prints for a balance: the figure, its magnitude, or nothing for nought."""
        if not magnitude:
            return value
        return abs(value) if value else None

    total_debit = total_credit = 0
    for account in order or ORDER:
        if account not in legs_of:
            continue
        ws.append([account])
        balance = (opening or {}).get(account, 0)
        if balance:
            ws.append(["Opening Balance", None, None, balance if balance > 0 else None,
                       -balance if balance < 0 else None, printed(balance) if balances else None])
        sub_debit = sub_credit = 0
        for when, _n, _k, voucher, counter, debit, credit, salesman, mode, text in sorted(
                legs_of[account], key=lambda t: (t[0], t[1], t[2])):
            balance += debit - credit
            sub_debit += debit
            sub_credit += credit
            shown = printed(balance) if balances and not (blank and blank(account, voucher, debit, credit)) else None
            ws.append([when, voucher, counter, debit or None, credit or None, shown, "BHD", mode, salesman, text])
        ws.append(["Sub Total", None, None, sub_debit, sub_credit, balance if balances else None])
        total_debit += sub_debit
        total_credit += sub_credit
    ws.append(["Grand Total", None, None, total_debit if grand is None else grand[0],
               total_credit if grand is None else grand[1]])
    wb.save(path)
    wb.close()


class data_root:
    """A temporary data root (the drop folders live under it)."""

    def __enter__(self) -> Path:
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name).resolve()
        paths.set_data_root(self.root)
        return self.root

    def __exit__(self, *exc):
        paths.set_data_root(None)
        for _ in range(20):
            try:
                self._tmp.cleanup()
                break
            except OSError:               # Windows: a scanner still holds a file written a moment ago
                time.sleep(0.1)
        return False


class book:
    """The full invented Ledger, written once into a temporary folder and read."""

    def __enter__(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.folder = Path(self._tmp.name).resolve()
        self.path = self.folder / "Ledger_Example.xlsx"
        write_ledger(self.path, POSTINGS, opening=OPENING)
        self.before = sorted(p.name for p in self.folder.iterdir())
        self.export = le.read(self.path)
        return self

    def __exit__(self, *exc):
        for _ in range(20):
            try:
                self._tmp.cleanup()
                break
            except OSError:
                time.sleep(0.1)
        return False


def leaks(blob: str) -> list[str]:
    """What must never be in a row, a sentence or a printed export: narration words and the ID-shaped and
    card-shaped numbers typed into it. A file hash is told apart by its 'sha256:' prefix and left out of the
    scan (its digits are chance)."""
    blob = re.sub(r"sha256:[0-9a-f]{64}", "sha256:(hash)", blob)
    low = blob.lower()
    found = [w for w in WORDS + (ID_A, ID_B, CARD) if w in low]
    if NINE.search(blob):
        found.append("a 9-digit run")
    if LONG.search(blob):
        found.append("a run of 15 digits or more")
    return found


def drop(root: Path, folder: str, name: str, as_on: str, postings=None) -> Path:
    """A small Ledger as on `as_on` ('dd/mm/yyyy') in the drop folder Uplaods/<folder>."""
    path = root / "Uplaods" / folder / name
    path.parent.mkdir(parents=True, exist_ok=True)
    write_ledger(path, SMALL if postings is None else postings, as_on=as_on)
    return path


def money(*figures) -> tuple[Decimal, ...]:
    return tuple(Decimal(x) for x in figures)


# -- the vouchers -------------------------------------------------------------

@test("kind_of: SI sale, RCTV receipt, SRV return, CN credit_note, PV payment, JV journal; the rest other")
def _():
    table = [("SI : 1234", "sale"), ("si:7", "sale"), ("  RCTV : 5", "receipt"), ("SRV : 3", "return"),
             ("CN : 1", "credit_note"), ("PV : 9", "payment"), ("JV : 2", "journal"), ("XYZ : 1", "other"),
             ("SIO:X-26-01-1", "other"), ("12345", "other"), ("SI", "other"), ("", "other"), (None, "other")]
    for voucher, kind in table:
        assert le.kind_of(voucher) == kind, (voucher, le.kind_of(voucher), kind)
    assert [le.prefix_of(v) for v in ("SI : 1234", "rctv:7", " Srv : 3", "77", "", None)] == \
        ["SI", "RCTV", "SRV", "", "", ""]
    assert le.KINDS == ("sale", "receipt", "return", "credit_note", "payment", "journal", "other")
    assert le.fold("  Customer   A ") == le.fold("customer a") == "customer a" and le.fold(None) == ""


# -- read ---------------------------------------------------------------------

@test("read: a cumulative file gives every account row, in file order, in the module's shape; totals are no rows")
def _():
    with book() as b:
        e = b.export
        assert len(e.rows) == N_ROWS == 54
        assert all(set(r) == ROW_KEYS for r in e.rows), "a row holds exactly the documented keys"
        assert e.source == "Ledger_Example.xlsx" and e.sha and e.sha.startswith("sha256:")
        assert e.as_on == AS_ON and e.first_date == date(2026, 9, 15) and e.last_date == date(2026, 11, 2)
        assert not [r for r in e.rows if r["account"] in ("Sub Total", "Grand Total") or
                    r["voucher"] in ("Opening Balance", "Sub Total", "Grand Total")]
        assert e.accounts == [CUST_A, CUST_B, CUST_C, SALES, RETURNS, VAT, TAXCTL, CASH, BANK, SUPPLIER]
        assert all(isinstance(r["date"], date) and not isinstance(r["date"], datetime) for r in e.rows)
        # file order: an account's rows stand together and run in date order
        seen = [r["account"] for r in e.rows]
        assert seen == sorted(seen, key=lambda a: e.accounts.index(a))
        for rows in e.by_account().values():
            assert [r["date"] for r in rows] == sorted(r["date"] for r in rows)
        # one row, whole: Customer A's first
        first = e.by_account()[le.fold(CUST_A)][0]
        assert first == {"account": CUST_A, "date": date(2026, 9, 20), "voucher": "SI : 100", "prefix": "SI",
                         "kind": "sale", "counter": SALES, "debit": Decimal(220), "credit": Decimal(0),
                         "balance": Decimal(220), "currency": "BHD", "payment_mode": "Credit", "salesman": REP1}


@test("read: money is Decimal to the fils (a blank cell is 0, a blank Balance is None); balances run with the rows")
def _():
    with book() as b:
        for r in b.export.rows:
            for key in ("debit", "credit"):
                assert isinstance(r[key], Decimal) and r[key].as_tuple().exponent >= -3 and r[key] >= 0
            assert isinstance(r["balance"], Decimal)
        # a running balance follows from the row before it: the check the real-file test prints counts of
        for rows in b.export.by_account().values():
            for before, row in zip(rows, rows[1:]):
                assert row["balance"] == before["balance"] + row["debit"] - row["credit"]
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "Ledger_blank.xlsx"
        write_ledger(path, SMALL, balances=False)
        rows = le.read(path).rows
        assert rows and all(r["balance"] is None for r in rows)
        assert all(r["debit"] == Decimal(0) or r["credit"] == Decimal(0) for r in rows)


@test("read: money is held to the fils: three decimals stay, a float's stray digits and a longer figure are rounded")
def _():
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "Ledger_fils.xlsx"
        write_ledger(path, [post(d(10, 5), "SI : 1", [(CUST_A, 12.345, 0), (SALES, 0, 12.345)]),
                            post(d(10, 6), "SI : 2", [(CUST_A, 0.1 + 0.2, 0), (SALES, 0, 0.1 + 0.2)]),
                            post(d(10, 7), "SI : 3", [(CUST_A, 1.0005, 0), (SALES, 0, 1.0005)])])
        e = le.read(path)
        a = e.by_account()[le.fold(CUST_A)]
        assert [r["debit"] for r in a] == [Decimal("12.345"), Decimal("0.3"), Decimal("1.001")], "half up, at the fils"
        assert [r["balance"] for r in a] == [Decimal("12.345"), Decimal("12.645"), Decimal("13.646")]
        for r in e.rows:
            for key in ("debit", "credit", "balance"):
                assert r[key].as_tuple().exponent >= -3, (key, r[key])
        assert e.rows_total() == (Decimal("13.646"), Decimal("13.646")) and e.grand_total_gap() == (0, 0)


@test("read: the narration is in no row, no sentence and no file; nothing is written")
def _():
    with book() as b:
        e = b.export
        assert all("narration" not in r and "row_hash" not in r for r in e.rows)
        blob = repr(e) + repr(e.rows) + " ".join(e.unknown + e.notes) + repr(e.by_account()) + e.source
        assert leaks(blob) == [], leaks(blob)
        # the helpers hand out the same rows or numbers, never the words
        blob = repr(le.movements(e, "2026-09-01", "2026-11-03")) + repr(le.customer_legs(e, [CUST_A, CUST_B]))
        blob += repr(le.returns_by_voucher(e, "2026-09-01", "2026-11-03"))
        assert leaks(blob) == [], leaks(blob)
        assert sorted(p.name for p in b.folder.iterdir()) == b.before, "reading left a file behind"
        # and the words really were in the workbook (the test would pass vacuously otherwise)
        from scripts import ingest as ing
        raw = ing.parse_ledger(ing.read_grid(b.path), b.path.name)
        assert any(r.get("narration") and ID_A in r["narration"] for r in raw)


@test("read: the noise payment modes ('0' as text and as a number, 'Payment Mode') are None; real ones stay")
def _():
    with book() as b:
        by_voucher = {}
        for r in b.export.rows:
            by_voucher.setdefault(r["voucher"], set()).add(r["payment_mode"])
        assert by_voucher["SI : 100"] == {"Credit"}
        assert by_voucher["RCTV : 50"] == {"Cash"}
        assert by_voucher["RCTV : 51"] == {None}, "the text '0'"
        assert by_voucher["JV : 1"] == {None}, "the number 0"
        assert by_voucher["PV : 9"] == {None}, "the header's own words"
        assert by_voucher["SRV : 7"] == {None}, "a blank cell"
        assert {r["payment_mode"] for r in b.export.rows} == {"Credit", "Cash", None}
    assert [le._payment_mode(x) for x in ("0", " 0.00 ", "payment  MODE", "", None, 0, "Cash", "Bank Transfer")] == \
        [None, None, None, None, None, None, "Cash", "Bank Transfer"]


@test("read: prefixes counted; unknown ones listed as 'type not known: <prefix> (n rows)'; no prefix is noted")
def _():
    with book() as b:
        e = b.export
        assert e.prefixes == {"SRV": 18, "SI": 18, "RCTV": 4, "CN": 2, "PV": 2, "JV": 2, "XYZ": 4, "ABC": 2, "": 2}
        assert sum(e.prefixes.values()) == len(e.rows)
        assert e.unknown == ["type not known: ABC (2 rows)", "type not known: XYZ (4 rows)"]
        assert any("2 rows carry a voucher that does not begin 'PREFIX :'" in n for n in e.notes), e.notes
        kinds = {}
        for r in e.rows:
            kinds[r["kind"]] = kinds.get(r["kind"], 0) + 1
        assert kinds == {"return": 18, "sale": 18, "receipt": 4, "credit_note": 2, "payment": 2, "journal": 2,
                         "other": 8}
        assert all(r["kind"] == le.kind_of(r["voucher"]) and r["prefix"] == le.prefix_of(r["voucher"])
                   for r in e.rows)
    with tempfile.TemporaryDirectory() as tmp:           # a file of known vouchers lists nothing
        path = Path(tmp) / "Ledger_known.xlsx"
        write_ledger(path, SMALL)
        e = le.read(path)
        assert e.unknown == [] and e.notes == []


@test("read: Focus's Grand Total is kept and compared with the rows; a difference shows as a gap")
def _():
    with book() as b:
        e = b.export
        assert e.grand_total == (Decimal(TOTAL_DEBIT), Decimal(TOTAL_CREDIT))
        assert e.rows_total() == (Decimal(TOTAL_DEBIT), Decimal(TOTAL_CREDIT)) and TOTAL_DEBIT == TOTAL_CREDIT
        assert e.grand_total_gap() == (Decimal(0), Decimal(0))
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "Ledger_gap.xlsx"
        write_ledger(path, SMALL, grand=(15, 11))
        e = le.read(path)
        assert e.grand_total == (Decimal(15), Decimal(11))
        assert e.grand_total_gap() == (Decimal(-4), Decimal(0))
        path = Path(tmp) / "Ledger_none.xlsx"
        write_ledger(path, SMALL)
        # a file that ends without a Grand Total row has nothing to compare with
        from openpyxl import load_workbook
        wb = load_workbook(path)
        ws = wb.active
        ws.delete_rows(ws.max_row)
        wb.save(path)
        wb.close()
        e = le.read(path)
        assert e.grand_total is None and e.grand_total_gap() is None and len(e.rows) == 2


@test("read: an Opening Balance row is no row, but its balance is carried by the running balance")
def _():
    with book() as b:
        rows = b.export.by_account()[le.fold(CUST_B)]
        assert [r["voucher"] for r in rows] == ["SRV : 6", "SI : 101", "SI : 104", "SRV : 8", "RCTV : 51", "SRV : 11"]
        assert rows[0]["balance"] == Decimal(40 - 44), "40 opened, 44 returned"


@test("read: rows above the first account are left out and noted")
def _():
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "Ledger_lead.xlsx"
        write_ledger(path, SMALL, lead=[[d(10, 1), "SI : 9", "Nowhere", 5, None, 5]])
        e = le.read(path)
        assert len(e.rows) == 2 and all(r["account"] for r in e.rows)
        assert e.notes == ["1 rows stand above the first account's header and were left out"], e.notes


@test("read: a changed header, a file that cannot be opened, no 'As on date' and a missing file each stop in words")
def _():
    from scripts import ingest as ing
    with tempfile.TemporaryDirectory() as tmp:
        folder = Path(tmp)
        bad = folder / "Ledger_header.xlsx"
        write_ledger(bad, SMALL, header=["Date", "Voucher", "Account", "Dr", "Cr", "Balance", "Currency",
                                         "Payment Mode", "Salesman", "Narration"])
        exc = raises(ing.HeaderMismatch, lambda: le.read(bad))
        assert "ledger_entries" in str(exc) and "Debit" in str(exc), str(exc)[:200]

        old = folder / "Ledger_old.xlsx"
        old.write_bytes(OLD_XLS)
        exc = raises(le.LedgerExportError, lambda: le.read(old))
        text = str(exc)
        assert text.startswith("the Ledger export Ledger_old.xlsx could not be opened (") and "BadZipFile" in text, text
        assert str(folder) not in text and "\n" not in text, "a name, never a path; one sentence"

        for line in ("[From 01/09/2025 To 31/10/2026]", "Ledger of everything"):
            titled = folder / "Ledger_title.xlsx"
            write_ledger(titled, SMALL, title_line=line)
            exc = raises(le.LedgerExportError, lambda: le.read(titled))
            assert "has no \"As on date\" in its title block" in str(exc) and "Ledger_title.xlsx" in str(exc), str(exc)

        exc = raises(le.LedgerExportError, lambda: le.read(folder / "nothing.xlsx"))
        assert str(exc) == "no such Ledger export: nothing.xlsx"
        exc = raises(le.LedgerExportError, lambda: le.read(folder))                      # a folder is no file
        assert "no such Ledger export" in str(exc)
        assert issubclass(le.LedgerExportError, ValueError)


# -- movements ----------------------------------------------------------------

@test("movements: a cumulative file is filtered by date; per account and kind; every kind is kept")
def _():
    with book() as b:
        e = b.export
        m = le.movements(e, date(2026, 10, 1), date(2026, 10, 31))
        a = m[CUST_A]
        assert set(a) == {"sale", "receipt", "return", "credit_note"}
        assert a["sale"] == le.Movement(Decimal(440), Decimal(0), 2), "the September and November sales stay out"
        assert a["receipt"] == le.Movement(Decimal(0), Decimal(200), 1)
        assert a["return"] == le.Movement(Decimal(0), Decimal(330), 2)
        assert a["credit_note"] == le.Movement(Decimal(0), Decimal(20), 1)
        assert a["return"].net == Decimal(-330) and a["sale"].net == Decimal(440)
        assert m[CUST_B] == {"sale": le.Movement(Decimal(550), Decimal(0), 1),
                             "return": le.Movement(Decimal(0), Decimal(55), 1),
                             "receipt": le.Movement(Decimal(0), Decimal(550), 1)}
        assert m[CUST_C] == {"return": le.Movement(Decimal(0), Decimal(33), 1)}
        assert m[SALES]["sale"].credit == Decimal(900) and m[SALES]["credit_note"].debit == Decimal(20)
        assert m[VAT]["sale"].credit == Decimal(90) and m[VAT]["return"].debit == Decimal(18)
        assert m[RETURNS]["return"] == le.Movement(Decimal(380), Decimal(0), 4)
        assert m[SUPPLIER] == {"payment": le.Movement(Decimal(1000), Decimal(0), 1)}
        bank = m[BANK]
        assert bank["payment"].credit == Decimal(1000) and bank["receipt"].debit == Decimal(550)
        assert bank["journal"].credit == Decimal(10) and bank["other"] == le.Movement(Decimal(0), Decimal(16), 4)
        assert m[CASH]["other"] == le.Movement(Decimal(16), Decimal(0), 4)
        assert list(m) == [CUST_A, CUST_B, CUST_C, SALES, RETURNS, VAT, TAXCTL, CASH, BANK, SUPPLIER]
        # the other months of the same file
        assert le.movements(e, "2026-09-01", "2026-09-30")[CUST_A] == {"sale": le.Movement(Decimal(220), Decimal(0), 1)}
        assert le.movements(e, "2026-11-01", "2026-11-03")[CUST_A] == {"sale": le.Movement(Decimal(55), Decimal(0), 1)}
        assert le.movements(e, "2026-08-01", "2026-08-31") == {}, "no entry in August"
        # the whole file: the months add up to it
        whole = le.movements(e, "2026-09-01", "2026-11-03")
        every = sum((mv.count for kinds in whole.values() for mv in kinds.values()), 0)
        assert every == len(e.rows)
        assert sum((mv.debit for kinds in whole.values() for mv in kinds.values()), Decimal(0)) == Decimal(TOTAL_DEBIT)
        one_day = le.movements(e, date(2026, 10, 5), date(2026, 10, 5))
        assert one_day[CUST_A]["sale"].count == 2 and set(one_day) == {CUST_A, SALES, VAT}


@test("movements: a window that is wrong, or runs past the as-on date, is refused")
def _():
    with book() as b:
        e = b.export
        for first, last in (("2026-10-31", "2026-10-01"), ("x", "2026-10-31"), ("2026-10-01", None),
                            ("2026-10-01", "2026-11-04"), ("2026-10-01", "2026-12-31")):
            raises(ValueError, lambda: le.movements(e, first, last))
        msg = str(raises(ValueError, lambda: le.movements(e, "2026-10-01", "2026-11-30")))
        assert "as on 2026-11-03" in msg and "2026-11-30" in msg and "Ledger_Example.xlsx" in msg, msg
        le.movements(e, "2026-10-01", "2026-11-03")                  # the as-on day itself is fine


# -- balance ------------------------------------------------------------------

@test("balance: the running balance at the end of a day; opening + movements = closing for every customer")
def _():
    with book() as b:
        e = b.export
        A = lambda *day: le.balance(e, CUST_A, date(2026, *day))      # noqa: E731
        assert A(9, 30) == Decimal(220) and A(10, 4) == Decimal(220)
        assert A(10, 5) == Decimal(660), "two sales on one day: the later row of the file"
        assert A(10, 12) == Decimal(460) and A(10, 31) == Decimal(110) and A(11, 3) == Decimal(165)
        assert A(9, 1) == Decimal(0), "before its first row: what that row opened with"
        assert le.balance(e, CUST_B, date(2026, 9, 1)) == Decimal(40), "the Opening Balance row, carried"
        assert le.balance(e, CUST_B, date(2026, 10, 31)) == Decimal(-4), "a credit balance is negative"
        assert le.balance(e, CUST_C, date(2026, 10, 1)) == Decimal(0)
        assert le.balance(e, CUST_C, "2026-10-31") == Decimal(-33), "an ISO text day"
        assert le.balance(e, "  customer   a ", date(2026, 10, 31)) == Decimal(110), "a name spelt another way"
        assert le.balance(e, "Nobody", date(2026, 10, 31)) is None and le.balance(e, None, date(2026, 10, 31)) is None
        assert isinstance(A(10, 31), Decimal)
        msg = str(raises(ValueError, lambda: le.balance(e, CUST_A, date(2026, 11, 4))))
        assert "as on 2026-11-03" in msg and "2026-11-04" in msg, msg
        raises(ValueError, lambda: le.balance(e, CUST_A, "soon"))

        first, last = date(2026, 10, 1), date(2026, 10, 31)
        moves = le.movements(e, first, last)
        closing = {}
        for account in (CUST_A, CUST_B, CUST_C):
            opening = le.balance(e, account, date(2026, 9, 30))
            net = sum((mv.net for mv in moves[account].values()), Decimal(0))
            assert opening + net == le.balance(e, account, last), account
            closing[account] = le.balance(e, account, last)
        assert closing == {CUST_A: Decimal(110), CUST_B: Decimal(-4), CUST_C: Decimal(-33)}
        # "opening + SI - RCTV - SRV - CN = closing" for Customer A, kind by kind
        a = moves[CUST_A]
        assert Decimal(220) + a["sale"].debit - a["receipt"].credit - a["return"].credit - a["credit_note"].credit \
            == Decimal(110)


@test("balance: an account none of whose rows states a Balance has none (None), before the first row as well")
def _():
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "Ledger_blank.xlsx"
        write_ledger(path, SMALL, balances=False)
        e = le.read(path)
        assert le.balance(e, CUST_A, date(2026, 10, 31)) is None
        assert le.balance(e, CUST_A, date(2026, 9, 1)) is None
        assert le.balance(e, "Nobody", date(2026, 10, 31)) is None


def run_rows(blank=blank_credits, postings=None, order=None):
    """The invented Ledger of RUNS (with the credit legs showing no Balance, unless `blank` says otherwise),
    written and read."""
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "Ledger_runs.xlsx"
        write_ledger(path, RUNS if postings is None else postings, order=order or ORDER_ALL, blank=blank)
        return le.read(path)


def blank_runs(legs: list[dict]) -> tuple[int, int]:
    """(the runs of rows that show no Balance and are closed by a row that does, how many of those runs the
    balance worked forward from the last stated one reaches: the stated balance of the row that closes it).
    A run before the first stated balance, and one that ends the account, are not counted: nothing closes them."""
    runs = reached = 0
    last = None                       # the last stated balance
    since = Decimal(0)                # debit less credit of the rows after it
    blanks = 0
    for r in legs:
        net = r["debit"] - r["credit"]
        if r["balance"] is None:
            if last is not None:
                since += net
                blanks += 1
            continue
        if blanks:
            runs += 1
            reached += last + since + net == r["balance"]
        last, since, blanks = r["balance"], Decimal(0), 0
    return runs, reached


@test("balance: a blank Balance (a credit leg: the payment inside a cash sale, a run of receipts, a last row) is "
      "worked forward from the last stated one")
def _():
    e = run_rows()
    d_rows = e.by_account()[le.fold(CUST_D)]
    # the shape of the real export: the sale's debit states its balance, its credit leg (the last row of the
    # voucher) does not; two receipts in a row show none; the account's last row shows none
    assert [(r["voucher"], r["debit"] or r["credit"], r["balance"]) for r in d_rows] == [
        ("SI : 201", 110, Decimal(110)), ("SI : 201", 110, None),
        ("SI : 202", 55, Decimal(55)), ("SI : 202", 55, None),
        ("RCTV : 61", 20, None), ("RCTV : 62", 30, None),
        ("SI : 203", 110, Decimal(60)), ("RCTV : 63", 60, None)]
    assert d_rows[-1]["balance"] is None, "the account's latest row shows no balance"
    day = lambda *a: date(2026, *a)                                    # noqa: E731
    want = {(10, 2): 0, (10, 3): 0, (10, 4): 0, (10, 5): -50, (10, 6): 60, (10, 7): 0, (11, 3): 0}
    for (m, dom), figure in want.items():
        assert le.balance(e, CUST_D, day(m, dom)) == Decimal(figure), (m, dom)
    assert isinstance(le.balance(e, CUST_D, day(10, 3)), Decimal)
    # the blank rows are not guesses: working forward through a cash sale's payment and through the run of
    # receipts lands on the stated balance of the row that closes each run (see blank_runs)
    assert blank_runs(d_rows) == (2, 2), "every run of blank rows closed by a stated balance lands on it"
    assert le.balance(e, "  customer  D ", day(10, 5)) == Decimal(-50)

    # an account whose first row shows no balance: worked back from the first row that does
    e_rows = e.by_account()[le.fold(CUST_E)]
    assert [r["balance"] for r in e_rows] == [None, Decimal(100)]
    assert le.balance(e, CUST_E, day(10, 3)) == Decimal(-20), "worked back from the next stated balance"
    assert le.balance(e, CUST_E, day(10, 4)) == Decimal(100)
    assert le.balance(e, CUST_E, day(10, 1)) == Decimal(0), "before its first row: what that row opened with"
    assert blank_runs(e_rows) == (0, 0), "a run no stated balance comes before cannot be worked forward"

    # an account that states no balance at all has none, which is the only time balance() says None
    f_rows = e.by_account()[le.fold(CUST_F)]
    assert [r["balance"] for r in f_rows] == [None]
    assert le.balance(e, CUST_F, day(10, 5)) is None and le.balance(e, CUST_F, day(10, 1)) is None


@test("balance: leaving cells blank changes no balance, on any account on any day, unless no row states one")
def _():
    postings = POSTINGS + RUNS
    with tempfile.TemporaryDirectory() as tmp:
        plain, gappy = Path(tmp) / "Ledger_plain.xlsx", Path(tmp) / "Ledger_gappy.xlsx"
        write_ledger(plain, postings, opening=OPENING, order=ORDER_ALL)
        write_ledger(gappy, postings, opening=OPENING, order=ORDER_ALL, blank=blank_credits)
        full, part = le.read(plain), le.read(gappy)
    assert all(r["balance"] is not None for r in full.rows)
    assert sum(1 for r in part.rows if r["balance"] is None) == sum(1 for r in full.rows if r["credit"] > 0) > 25
    none_at_all = {a for a in part.accounts if all(r["balance"] is None for r in part.by_account()[le.fold(a)])}
    assert none_at_all == {CUST_C, CUST_F}, none_at_all
    assert part.accounts == full.accounts and len(part.accounts) == 13
    checked, day = 0, date(2026, 9, 1)
    while day <= AS_ON:
        for account in full.accounts:
            want = None if account in none_at_all else le.balance(full, account, day)
            assert le.balance(part, account, day) == want, (account, day.isoformat())
            checked += 1
        day += timedelta(days=1)
    assert checked == 13 * 64
    # the same, row by row: every run of blank rows that a stated balance closes is reached by working forward
    # (Customer A's receipt, returns and credit note before its last sale; the Output VAT account's sales credits
    # before a return, twice; Customer D's two)
    runs = reached = 0
    for rows in part.by_account().values():
        n, ok = blank_runs(rows)
        runs, reached = runs + n, reached + ok
    assert (runs, reached) == (5, 5), (runs, reached)


def read_variant(postings, **kw):
    """`postings` written as a Ledger with the given writer options, and read."""
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "Ledger_variant.xlsx"
        write_ledger(path, postings, order=ORDER_ALL, **kw)
        return le.read(path)


def every_day(first=date(2026, 9, 1), last=AS_ON):
    day = first
    while day <= last:
        yield day
        day += timedelta(days=1)


@test("balance: the real export prints a magnitude (never a sign, a zero left blank): a credit balance comes "
      "out negative, every blank is the worked-out nought, an opening balance above the first row is found")
def _():
    postings = POSTINGS + RUNS
    for opening, label in ((None, "no opening balance, as in the real export"), (OPENING, "Customer B opens at 40")):
        signed = read_variant(postings, opening=opening)                                  # signed figures
        printed = read_variant(postings, opening=opening, magnitude=True)                # as the real export prints
        gappy = read_variant(postings, opening=opening, magnitude=True, blank=blank_credits)   # credit legs blank too
        assert all(r["balance"] >= 0 for r in printed.rows if r["balance"] is not None), "no sign, as in the export"
        assert any(r["balance"] is not None and r["balance"] < 0 for r in signed.rows)
        assert any(r["balance"] is None for r in printed.rows), "a nought is left blank"
        # every account on every day: the same figure whichever way the report prints it
        none_at_all = {a for a in gappy.accounts
                       if all(r["balance"] is None for r in gappy.by_account()[le.fold(a)])}
        assert none_at_all == {CUST_C, CUST_F}, none_at_all
        checked = 0
        for account in signed.accounts:
            for day in every_day():
                want = le.balance(signed, account, day)
                assert le.balance(printed, account, day) == want, (label, account, day.isoformat())
                assert le.balance(gappy, account, day) == (None if account in none_at_all else want), \
                    (label, account, day.isoformat())
                checked += 1
        assert checked == 13 * 64
        # a few by name: a customer in credit is negative, a sale account (a credit nature) too, a debtor positive
        assert le.balance(printed, CUST_C, date(2026, 10, 31)) == Decimal(-33), "its one row is a return"
        assert le.balance(printed, CUST_F, date(2026, 10, 31)) == Decimal(-5)
        assert le.balance(printed, CUST_A, date(2026, 10, 31)) == Decimal(110)
        assert le.balance(printed, SALES, date(2026, 10, 31)) < 0 and le.balance(printed, CASH, date(2026, 10, 31)) > 0
        assert le.balance(printed, CUST_E, date(2026, 10, 3)) == Decimal(-20), "in credit before its first sale"
        assert le.balance(printed, CUST_E, date(2026, 10, 4)) == Decimal(100)
        # the size of a printed figure is the report's own: only its sign is restored, only a blank is filled
        for rows in printed.by_account().values():
            for stated, figure in zip(rows, le.running_balances(rows)):
                assert stated["balance"] is None or abs(figure) == stated["balance"]
    # Customer B's opening balance is a row the reader leaves out: -4 after its first return, printed as 4. The sign
    # and the opening come from the later figures, which only the right opening makes equal the rows' totals.
    printed = read_variant(postings, opening=OPENING, magnitude=True)
    assert le.balance(printed, CUST_B, date(2026, 10, 31)) == Decimal(-4)
    assert le.balance(printed, CUST_B, date(2026, 9, 1)) == Decimal(40), "before its first row: the opening it carried"


@test("running_balances: a stated balance is kept, a blank one is worked from the last stated, leading blanks "
      "from the next, none at all gives None")
def _():
    def row(debit, credit, balance):
        return {"debit": Decimal(debit), "credit": Decimal(credit),
                "balance": None if balance is None else Decimal(balance)}

    rows = [row(0, 20, None), row(120, 0, 100), row(0, 30, None), row(0, 10, None), row(50, 0, 110), row(0, 5, 999)]
    # -20 worked back from 100; 100 stated; 70 and 60 worked forward; 110 stated; 999 stated, though 105 follows
    assert le.running_balances(rows) == [Decimal(-20), Decimal(100), Decimal(70), Decimal(60), Decimal(110),
                                         Decimal(999)]
    assert le.running_balances([row(0, 5, None), row(7, 0, None)]) == [None, None]
    assert le.running_balances([]) == []
    assert le.running_balances([row(3, 0, None), row(0, 1, 9), row(0, 2, None)]) == \
        [Decimal(10), Decimal(9), Decimal(7)], "the first row's balance is the next stated one less the movements"
    figures = le.running_balances(rows)
    assert all(isinstance(x, Decimal) for x in figures)

    # a balance printed without its sign gets it from the rows (nothing is owed before the first row)
    # (the last stated figure, 40, is the report's own: the rows add up to 0, and its size is kept)
    printed = [row(0, 30, 30), row(100, 0, 70), row(0, 150, 80), row(0, 20, None), row(100, 0, 40)]
    assert le.running_balances(printed) == [Decimal(-30), Decimal(70), Decimal(-80), Decimal(-100), Decimal(40)]
    # an opening balance above the first row: 40 owed, a return of 44 puts it 4 in credit, printed as 4
    assert le.running_balances([row(0, 44, 4), row(55, 0, 51), row(0, 55, None)]) == \
        [Decimal(-4), Decimal(51), Decimal(-4)]
    # a negative stated cell means the report prints its signs: taken as printed, whatever the rows say
    assert le.running_balances([row(0, 30, -30), row(100, 0, 70)]) == [Decimal(-30), Decimal(70)]
    # nothing ties the printed figures to the rows (no opening explains them): they stay as printed
    assert le.running_balances([row(10, 0, 500), row(10, 0, 3), row(0, 5, 77)]) == \
        [Decimal(500), Decimal(3), Decimal(77)]
    # ... and a minority that a (credit) opening does tie to them is not enough to give the whole account a sign:
    # two of four figures fit an opening of 90 in credit, the other two fit nothing
    assert le.running_balances([row(0, 10, 100), row(0, 5, 105), row(0, 5, 1), row(0, 5, 2)]) == \
        [Decimal(100), Decimal(105), Decimal(1), Decimal(2)]
    # a printed negative means the report prints its signs: taken as printed, though the other figures match the rows
    assert le.running_balances([row(0, 30, -30), row(100, 0, 70), row(10, 0, 80), row(10, 0, 90)]) == \
        [Decimal(-30), Decimal(70), Decimal(80), Decimal(90)]
    # a figure that fits both a debit from nought and a credit under a larger opening is read as the debit
    assert le.running_balances([row(20, 0, 20)]) == [Decimal(20)]
    # a single stated figure: the rows explain it from nought, so it is a credit balance; a figure they cannot
    # explain from nought is read as a debit balance (an opening of 9009 less the 9 returned)
    assert le.running_balances([row(0, 9, 9)]) == [Decimal(-9)]
    assert le.running_balances([row(0, 9, 9000)]) == [Decimal(9000)]


# -- customer_legs ------------------------------------------------------------

@test("customer_legs: the accounts the caller passes, however spelt; none held gives []; with or without a window")
def _():
    with book() as b:
        e = b.export
        legs = le.customer_legs(e, [CUST_A, "customer  b", "Nobody"])
        assert list(legs) == [CUST_A, "customer  b", "Nobody"], "keyed as the caller wrote them"
        assert [r["voucher"] for r in legs[CUST_A]] == ["SI : 100", "SI : 102", "SI : 103", "RCTV : 50", "SRV : 7",
                                                         "CN : 3", "SRV : 9", "SI : 110"]
        assert len(legs["customer  b"]) == 6 and legs["Nobody"] == []
        assert all(r["account"] in (CUST_A, CUST_B) and "narration" not in r for rows in legs.values() for r in rows)
        october = le.customer_legs(e, [CUST_A, CUST_B], "2026-10-01", date(2026, 10, 31))
        assert [r["voucher"] for r in october[CUST_A]] == ["SI : 102", "SI : 103", "RCTV : 50", "SRV : 7", "CN : 3",
                                                           "SRV : 9"]
        assert [r["voucher"] for r in october[CUST_B]] == ["SI : 104", "SRV : 8", "RCTV : 51"]
        assert len(le.customer_legs(e, CUST_A)[CUST_A]) == 8, "a lone name works as a list of one"
        assert [r["voucher"] for r in le.customer_legs(e, [CUST_A], last="2026-09-30")[CUST_A]] == ["SI : 100"]
        assert [r["voucher"] for r in le.customer_legs(e, [CUST_A], first="2026-11-01")[CUST_A]] == ["SI : 110"]
        assert le.customer_legs(e, []) == {}
        raises(ValueError, lambda: le.customer_legs(e, [CUST_A], "2026-10-31", "2026-10-01"))
        raises(ValueError, lambda: le.customer_legs(e, [CUST_A], None, "2026-11-04"))
        # the Sales account's rows are not customer legs unless the caller names it
        assert all(r["account"] != SALES for rows in legs.values() for r in rows)


# -- returns ------------------------------------------------------------------

@test("returns_ex_vat_by_salesman: VAT from the return's own legs, else by rate; salesman from any leg")
def _():
    with book() as b:
        e = b.export
        got = le.returns_ex_vat_by_salesman(e, "2026-10-01", "2026-10-31")
        assert got == {REP1: Decimal(300), REP2: Decimal(50), le.NO_SALESMAN: Decimal(30)}, got
        assert list(got) == [REP1, REP2, le.NO_SALESMAN], "by name, no salesman last"
        v = {x.voucher: x for x in le.returns_by_voucher(e, "2026-10-01", "2026-10-31")}
        assert sorted(v) == ["SRV : 10", "SRV : 7", "SRV : 8", "SRV : 9"], "September's and November's stay out"
        assert (v["SRV : 7"].gross, v["SRV : 7"].vat, v["SRV : 7"].ex_vat) == money(110, 10, 100)
        assert (v["SRV : 8"].gross, v["SRV : 8"].vat, v["SRV : 8"].ex_vat) == money(55, 5, 50)
        assert (v["SRV : 9"].gross, v["SRV : 9"].vat, v["SRV : 9"].ex_vat) == money(220, 20, 200)
        assert (v["SRV : 10"].gross, v["SRV : 10"].vat, v["SRV : 10"].ex_vat) == money(33, 3, 30)
        assert {k: x.basis for k, x in v.items()} == {"SRV : 7": "legs", "SRV : 8": "legs", "SRV : 9": "rate",
                                                      "SRV : 10": "legs"}
        assert all(x.legs == 3 for x in v.values())
        assert v["SRV : 7"].salesman == REP1 and v["SRV : 10"].salesman is None
        assert v["SRV : 9"].salesman == REP1, "named on the returns leg only: any leg will do"
        assert v["SRV : 9"].date == date(2026, 10, 28) and v["SRV : 8"].salesman == REP2
        # the VAT account named by the caller
        named = {x.voucher: x for x in le.returns_by_voucher(e, "2026-10-01", "2026-10-31",
                                                             vat_accounts=["tax  control"])}
        assert {k: x.basis for k, x in named.items()} == {"SRV : 7": "rate", "SRV : 8": "rate", "SRV : 9": "legs",
                                                          "SRV : 10": "rate"}
        assert named["SRV : 9"].vat == Decimal(20) and named["SRV : 9"].ex_vat == Decimal(200)
        assert le.returns_ex_vat_by_salesman(e, "2026-10-01", "2026-10-31", vat_accounts=[VAT, TAXCTL]) == got
        # the rate is a parameter: at 100% a rate-based return halves
        assert le.returns_ex_vat_by_salesman(e, "2026-10-01", "2026-10-31", vat_rate=1)[REP1] == Decimal(210)
        assert le.returns_ex_vat_by_salesman(e, "2026-10-01", "2026-10-31", vat_rate="0.1") == got
        # the other months of the same file
        assert le.returns_ex_vat_by_salesman(e, "2026-09-01", "2026-09-30") == {REP2: Decimal(40)}
        assert le.returns_ex_vat_by_salesman(e, "2026-11-01", "2026-11-03") == {REP2: Decimal(10)}
        assert le.returns_ex_vat_by_salesman(e, "2026-08-01", "2026-08-31") == {}
        raises(ValueError, lambda: le.returns_ex_vat_by_salesman(e, "2026-10-01", "2026-11-30"))
        assert sum(le.returns_ex_vat_by_salesman(e, "2026-09-01", "2026-11-03").values()) == Decimal(430)


@test("returns_by_voucher: a customer whose name holds the word VAT is no VAT leg (a return debits its VAT)")
def _():
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "Ledger_vat.xlsx"
        legs = [(RETURNS, 100, 0), (VAT, 10, 0), ("Vat Experts", 0, 110, REP1)]
        write_ledger(path, [post(d(10, 5), "SRV : 1", legs)], order=["Vat Experts", RETURNS, VAT])
        e = le.read(path)
        (v,) = le.returns_by_voucher(e, "2026-10-01", "2026-10-31")
        assert (v.gross, v.vat, v.ex_vat) == money(110, 10, 100) and v.basis == "legs" and v.legs == 3
        assert le.returns_ex_vat_by_salesman(e, "2026-10-01", "2026-10-31") == {REP1: Decimal(100)}
        # the same voucher with no VAT leg at all: by rate, and the customer still is not mistaken for one
        write_ledger(path, [post(d(10, 5), "SRV : 1", [(RETURNS, 110, 0), ("Vat Experts", 0, 110, REP1)])],
                     order=["Vat Experts", RETURNS])
        (v,) = le.returns_by_voucher(le.read(path), "2026-10-01", "2026-10-31")
        assert (v.gross, v.vat, v.ex_vat, v.basis, v.legs) == money(110, 10, 100) + ("rate", 2)


# -- find ---------------------------------------------------------------------

@test("find: the Ledger as on the month's last day, in the folder of the first morning; other files are left alone")
def _():
    with data_root() as root:
        want = drop(root, "011126", "Ledger_Oct.xlsx", "31/10/2026")
        (want.parent / "Stock_ledger_Oct.xlsx").write_bytes(b"not read")       # classified as the stock ledger
        (want.parent / "Ledger_detail_Oct.xlsx").write_bytes(b"not read")      # a subset of the Ledger: skipped
        (want.parent / "Sales_day_book_Oct.xlsx").write_bytes(b"not read")
        (want.parent / "~$Ledger_Oct.xlsx").write_bytes(b"Excel's lock")
        (want.parent / "Ledger_Oct.pdf").write_bytes(b"not a workbook")
        (want.parent / "Archive").mkdir()
        write_ledger(want.parent / "Archive" / "Ledger_old.xlsx", SMALL, as_on="31/10/2026")     # not entered
        found, why = le.find("2026-10")
        assert found == want, (found, why)
        assert why == "the Ledger as on 2026-10-31, in 011126/Ledger_Oct.xlsx", why


@test("find: an as-on on the last day beats a later one, whichever folder holds it")
def _():
    with data_root() as root:
        later = drop(root, "011126", "Ledger_a.xlsx", "01/11/2026")
        exact = drop(root, "021126", "Ledger_b.xlsx", "31/10/2026")
        found, why = le.find("2026-10")
        assert found == exact and "as on 2026-10-31" in why, (found, why)
        assert "not used: 011126/Ledger_a.xlsx (as on 2026-11-01)" in why, why
        assert later.is_file()


@test("find: with none on the last day, the smallest as-on after it; the report is read up to the month end")
def _():
    with data_root() as root:
        drop(root, "031126", "Ledger_c.xlsx", "03/11/2026")
        smallest = drop(root, "021126", "Ledger_b.xlsx", "01/11/2026")
        drop(root, "011126", "Ledger_a.xlsx", "02/11/2026")
        found, why = le.find("2026-10")
        assert found == smallest, (found, why)
        assert why.startswith("no Ledger is as on 2026-10-31; the earliest after it is as on 2026-11-01, in "
                              "021126/Ledger_b.xlsx, and it is read up to 2026-10-31"), why
        assert "011126/Ledger_a.xlsx (as on 2026-11-02)" in why and "031126/Ledger_c.xlsx (as on 2026-11-03)" in why


@test("find: an earlier as-on is refused with the reason; a later one beside it is used and the earlier one named")
def _():
    with data_root() as root:
        drop(root, "311026", "Ledger_early.xlsx", "30/10/2026")
        found, why = le.find("2026-10")
        assert found is None
        assert why == ("the Ledger found is as on 2026-10-30 (311026/Ledger_early.xlsx), before the month ends "
                       "(2026-10-31): it cannot cover 2026-10. Export the Ledger again as on 2026-10-31"), why
        later = drop(root, "011126", "Ledger_ok.xlsx", "01/11/2026")
        found, why = le.find("2026-10")
        assert found == later, why
        assert "311026/Ledger_early.xlsx (as on 2026-10-30, before the month ends)" in why, why


@test("find: only folders dated from the last day to ten days after are searched")
def _():
    nowhere = ("no drop folder is dated 2026-10-31 to 2026-11-10, where the Ledger of 2026-10 is looked for")
    with data_root() as root:
        drop(root, "301026", "Ledger_before.xlsx", "31/10/2026")        # the day before the month ends
        drop(root, "111126", "Ledger_after.xlsx", "31/10/2026")         # eleven days after
        found, why = le.find("2026-10")
        assert found is None and why == nowhere, why
        edge = drop(root, "101126", "Ledger_edge.xlsx", "31/10/2026")   # ten days after: still searched
        found, why = le.find("2026-10")
        assert found == edge, why
    with data_root() as root:                                            # the last day itself is searched
        day = drop(root, "311026", "Ledger_day.xlsx", "31/10/2026")
        assert le.find("2026-10")[0] == day
    with data_root():                                                    # no drop folder anywhere
        found, why = le.find("2026-10")
        assert found is None and why == nowhere, why


@test("find: folders in range with no Ledger file say so, and name the folders")
def _():
    with data_root() as root:
        folder = root / "Uplaods" / "011126"
        folder.mkdir(parents=True)
        (folder / "Sales_day_book_Oct.xlsx").write_bytes(b"x")
        (root / "Uplaods" / "021126").mkdir()
        found, why = le.find("2026-10")
        assert found is None
        assert why == ("no Ledger file lies directly in the drop folders dated 2026-10-31 to 2026-11-10 "
                       "(011126, 021126)"), why


@test("find: a Ledger that cannot be opened is passed over and named; alone it leaves no usable file")
def _():
    with data_root() as root:
        folder = root / "Uplaods" / "011126"
        folder.mkdir(parents=True)
        (folder / "Ledger_broken.xlsx").write_bytes(OLD_XLS)
        found, why = le.find("2026-10")
        assert found is None
        assert why == ("no usable Ledger file in the drop folders dated 2026-10-31 to 2026-11-10; not usable: "
                       "011126/Ledger_broken.xlsx could not be opened (BadZipFile)"), why
        good = drop(root, "011126", "Ledger_good.xlsx", "31/10/2026")
        found, why = le.find("2026-10")
        assert found == good
        assert "not used: 011126/Ledger_broken.xlsx could not be opened (BadZipFile)" in why, why
        titled = folder / "Ledger_untitled.xlsx"
        write_ledger(titled, SMALL, title_line="[From 01/10/2026 To 31/10/2026]")
        found, why = le.find("2026-10")
        assert found == good and "011126/Ledger_untitled.xlsx has no \"As on date\" in its title block" in why, why


@test("find: of two with one as-on the newest by modified time wins, in one folder or across two")
def _():
    with data_root() as root:
        old = drop(root, "011126", "Ledger_b.xlsx", "31/10/2026")
        new = drop(root, "011126", "Ledger_a.xlsx", "31/10/2026")          # the name that sorts first is the newer
        now = time.time()
        os.utime(old, (now - 3600, now - 3600))
        os.utime(new, (now, now))
        found, why = le.find("2026-10")
        assert found == new and "not used: 011126/Ledger_b.xlsx (as on 2026-10-31)" in why, (found, why)
        os.utime(old, (now + 60, now + 60))
        assert le.find("2026-10")[0] == old
    with data_root() as root:
        first = drop(root, "011126", "Ledger_x.xlsx", "31/10/2026")
        second = drop(root, "021126", "Ledger_x.xlsx", "31/10/2026")
        now = time.time()
        os.utime(first, (now, now))
        os.utime(second, (now - 3600, now - 3600))
        assert le.find("2026-10")[0] == first, "the newer file, not the later folder"
        os.utime(second, (now + 60, now + 60))
        assert le.find("2026-10")[0] == second


@test("find: a named file replaces the search, not the as-on rule; any name will do")
def _():
    with data_root() as root:
        named = root / "somewhere" / "october_all_accounts.xlsx"
        named.parent.mkdir(parents=True)
        write_ledger(named, SMALL, as_on="31/10/2026")
        found, why = le.find("2026-10", path=named)
        assert found == named and why == "the file named, the Ledger as on 2026-10-31", (found, why)
        write_ledger(named, SMALL, as_on="04/11/2026")
        found, why = le.find("2026-10", path=str(named))
        assert found == named and "as on 2026-11-04" in why and "read up to the month's last day, 2026-10-31" in why
        write_ledger(named, SMALL, as_on="30/10/2026")
        found, why = le.find("2026-10", path=named)
        assert found is None and "october_all_accounts.xlsx is as on 2026-10-30, before the month ends" in why, why
        found, why = le.find("2026-10", path=root / "gone.xlsx")
        assert found is None and why == "no such file: gone.xlsx"
        broken = root / "broken.xlsx"
        broken.write_bytes(OLD_XLS)
        found, why = le.find("2026-10", path=broken)
        assert found is None and why == "broken.xlsx could not be opened (BadZipFile)", why
        # a Ledger in a drop folder is not looked for when a file is named
        drop(root, "011126", "Ledger_Oct.xlsx", "31/10/2026")
        write_ledger(named, SMALL, as_on="31/10/2026")
        assert le.find("2026-10", path=named)[0] == named


@test("find: `drops` names the folders to look in; the data root is not consulted")
def _():
    with data_root() as root, tempfile.TemporaryDirectory() as other:
        elsewhere = Path(other) / "first morning"
        elsewhere.mkdir()
        write_ledger(elsewhere / "Ledger_Oct.xlsx", SMALL, as_on="31/10/2026")
        decoy = drop(root, "011126", "Ledger_decoy.xlsx", "31/10/2026")
        found, why = le.find("2026-10", drops=[(date(2026, 11, 1), elsewhere)])
        assert found == elsewhere / "Ledger_Oct.xlsx"
        assert why == "the Ledger as on 2026-10-31, in first morning/Ledger_Oct.xlsx"
        assert found != decoy
        found, why = le.find("2026-10", drops=[(date(2026, 11, 11), elsewhere)])
        assert found is None and "no drop folder is dated" in why
        found, why = le.find("2026-10", drops=[])
        assert found is None and "no drop folder is dated" in why


@test("find: the year's end (December -> January folders); a month that is not YYYY-MM is refused")
def _():
    with data_root() as root:
        want = drop(root, "020127", "Ledger_Dec.xlsx", "31/12/2026")
        found, why = le.find("2026-12")
        assert found == want and why == "the Ledger as on 2026-12-31, in 020127/Ledger_Dec.xlsx", (found, why)
        assert le.find("2026-11")[0] is None, "a folder of January is nowhere near November's month end"
        assert le.find("2026-02")[0] is None
    for bad in ("2026-13", "26-10", "October", "2026-10-31", "", None):
        exc = raises(ValueError, lambda: le.find(bad))
        assert "YYYY-MM" in str(exc), str(exc)


# -- the real export ----------------------------------------------------------

def counts(export) -> dict:
    """Aggregate counts of a real export. Numbers and codes only: never an account, a name, a voucher number
    or an amount."""
    rows = export.rows
    vouchers: dict[tuple, list[dict]] = {}
    for r in rows:
        vouchers.setdefault((" ".join(r["voucher"].split()).upper(), r["date"]), []).append(r)
    sections = sum(1 for legs in vouchers.values() if len({le.fold(r["account"]) for r in legs}) >= 2)
    balanced = sum(1 for legs in vouchers.values()
                   if sum((r["debit"] for r in legs), Decimal(0)) == sum((r["credit"] for r in legs), Decimal(0)))
    pairs = follows = ordered = 0
    for legs in export.by_account().values():
        for before, row in zip(legs, legs[1:]):
            ordered += row["date"] >= before["date"]
            if before["balance"] is not None and row["balance"] is not None:
                pairs += 1
                follows += row["balance"] == before["balance"] + row["debit"] - row["credit"]
    gap = export.grand_total_gap()
    returns = le.returns_by_voucher(export, export.first_date or export.as_on, export.as_on) if rows else []
    return {
        "rows": len(rows), "accounts": len(export.accounts), "vouchers": len(vouchers),
        "vouchers in two or more sections": sections, "vouchers with debit = credit": balanced,
        "rows with a balance": sum(1 for r in rows if r["balance"] is not None),
        "pairs of rows following debit/credit": f"{follows} of {pairs}",
        "pairs of rows in date order": f"{ordered} of {sum(max(len(v) - 1, 0) for v in export.by_account().values())}",
        "rows with a payment mode": sum(1 for r in rows if r["payment_mode"]),
        "rows with a salesman": sum(1 for r in rows if r["salesman"]),
        "currencies": dict(sorted(Counter(r["currency"] for r in rows).items(), key=lambda kv: str(kv[0]))),
        "Grand Total row": "none" if gap is None else ("rows add up to it" if gap == (0, 0) else
                                                       f"debit {'matches' if gap[0] == 0 else 'differs'}, "
                                                       f"credit {'matches' if gap[1] == 0 else 'differs'}"),
        "returns (SRV vouchers)": len(returns),
        "returns: VAT legs found / by rate": f"{sum(1 for x in returns if x.basis == 'legs')} / "
                                             f"{sum(1 for x in returns if x.basis == 'rate')}",
        "returns: fewer than two legs": sum(1 for x in returns if x.legs < 2),
        "returns: no salesman": sum(1 for x in returns if not x.salesman),
    }


CUSTOMER_ROWS = 1000                  # an account of more rows than this is not a customer in the real-file counts


def month_end_before(day: date) -> date:
    return date(day.year, day.month, 1) - timedelta(days=1)


def balance_counts(export) -> tuple[dict, tuple[int, int]]:
    """How the report's Balance column behaves, as counts: ({label: figure}, (runs of blank rows a stated balance
    closes, how many of them the balance worked forward reaches)). The blank cells, the missing sign, and the
    roll-forward a reconciliation does with balance(), month by month. Numbers only."""
    accounts = export.by_account()
    rows = export.rows
    blank_rows = [r for r in rows if r["balance"] is None]
    runs = reached = 0
    for legs in accounts.values():
        n, ok = blank_runs(legs)
        runs, reached = runs + n, reached + ok
    stated_somewhere = [legs for legs in accounts.values() if any(r["balance"] is not None for r in legs)]
    out = {
        "rows with no Balance": f"{len(blank_rows)} of {len(rows)} (credit legs: "
                                f"{sum(1 for r in blank_rows if r['credit'] > 0 and r['debit'] == 0)}, "
                                f"debit legs: {sum(1 for r in blank_rows if r['debit'] > 0 and r['credit'] == 0)})",
        "accounts whose last row shows no Balance":
            f"{sum(1 for legs in accounts.values() if legs[-1]['balance'] is None)} of {len(accounts)}",
        "accounts that state no Balance at all": len(accounts) - len(stated_somewhere),
        "accounts whose first rows show none before a stated one": sum(1 for legs in stated_somewhere
                                                                      if legs[0]["balance"] is None),
        "runs of blank rows closed by a stated Balance": runs,
        "runs the forward balance reaches (equals the stated one)": f"{reached} of {runs}",
    }
    for label, day in (("as on", export.as_on), ("at the month end before", month_end_before(export.as_on))):
        month_first = date(day.year, day.month, 1)
        with_row = raw_blank = unstated = 0
        debited = debited_raw = debited_unstated = 0
        for legs in accounts.values():
            latest = None
            for i, r in enumerate(legs):
                if r["date"] <= day and (latest is None or r["date"] >= legs[latest]["date"]):
                    latest = i
            if latest is None:
                continue
            with_row += 1
            is_raw_blank = legs[latest]["balance"] is None
            is_none = le.balance(export, legs[0]["account"], day) is None
            raw_blank += is_raw_blank
            unstated += is_none
            if any(r["kind"] == "sale" and r["debit"] > 0 and month_first <= r["date"] <= day for r in legs):
                debited += 1
                debited_raw += is_raw_blank
                debited_unstated += is_none
        out[f"{label} {day.isoformat()}: accounts with a row / latest row blank / balance() None"] = \
            f"{with_row} / {raw_blank} / {unstated}"
        out[f"{label} {day.isoformat()}: accounts a sales invoice debited that month, the same three"] = \
            f"{debited} / {debited_raw} / {debited_unstated}"
    # the sign: the report prints a magnitude, balance() gives the signed figure
    negatives = sum(1 for r in rows if r["balance"] is not None and r["balance"] < 0)
    restored_accounts = restored_rows = 0
    for legs in accounts.values():
        given = [(r["balance"], b) for r, b in zip(legs, le.running_balances(legs)) if r["balance"] is not None]
        flipped = sum(1 for printed, b in given if b != printed)
        restored_rows += flipped
        restored_accounts += bool(flipped)
    out["printed balances below nought (a report that prints its signs)"] = negatives
    out["balances given a minus sign from the rows: accounts / rows"] = f"{restored_accounts} / {restored_rows}"
    # what the reconciliation does with it: opening + movements = closing for the accounts a sales invoice debited
    # in each month of the file. The accounts of up to CUSTOMER_ROWS rows are the customers; the few with thousands
    # of rows (cash sales, up to hundreds of rows a day) are counted apart, their printed balances do not follow
    # the file's row order.
    month_first = date(export.first_date.year, export.first_date.month, 1)
    while month_first <= export.as_on:
        month_next = date(month_first.year + (month_first.month == 12), month_first.month % 12 + 1, 1)
        last = min(month_next - timedelta(days=1), export.as_on)
        eve = month_first - timedelta(days=1)
        moves = le.movements(export, month_first, last)
        debited = [legs for legs in accounts.values()
                   if any(r["kind"] == "sale" and r["debit"] > 0 and month_first <= r["date"] <= last for r in legs)]
        small = [legs for legs in debited if len(legs) <= CUSTOMER_ROWS]
        in_credit = off_fils = off_five = unknown = 0
        for legs in small:
            account = legs[0]["account"]
            opening, closing = le.balance(export, account, eve), le.balance(export, account, last)
            if opening is None or closing is None:
                unknown += 1
                continue
            gap = abs(opening + sum((m.net for m in moves.get(account, {}).values()), Decimal(0)) - closing)
            off_fils += gap > Decimal("0.001")
            off_five += gap > Decimal("0.05")
            in_credit += closing < 0
        out[f"{month_first:%Y-%m}: customer accounts / in credit at the close / no balance / opening + movements != "
            f"closing by over 0.001 BHD / by over 0.05 BHD / bigger accounts apart"] = \
            f"{len(small)} / {in_credit} / {unknown} / {off_fils} / {off_five} / {len(debited) - len(small)}"
        month_first = month_next
    return out, (runs, reached)


@test("the real export (SKIP unless YQ_OPS_LEDGER_SAMPLE names one): header, as-on date and prefixes are as expected")
def _():
    sample = os.environ.get("YQ_OPS_LEDGER_SAMPLE", "").strip()
    if not sample:
        print("  SKIP: YQ_OPS_LEDGER_SAMPLE names no Ledger export on this machine")
        return
    path = Path(sample)
    assert path.is_file(), "YQ_OPS_LEDGER_SAMPLE names no file"
    import pandas as pd
    from scripts import ingest as ing

    top = pd.read_excel(path, header=None, dtype=object, engine="openpyxl", nrows=12)
    assert ing.header_row_index(top, "ledger_entries") is not None, "the header labels are not those the parser expects"
    t0 = time.perf_counter()
    export = le.read(path)
    took = time.perf_counter() - t0

    assert isinstance(export.as_on, date), "no as-on date"
    assert export.rows and export.accounts, "no account rows"
    assert all("narration" not in r for r in export.rows)
    with_prefix = sum(1 for r in export.rows if r["prefix"])
    without = len(export.rows) - with_prefix
    assert with_prefix >= 0.99 * len(export.rows), f"rows whose voucher does not begin 'PREFIX :': {without}"

    c = counts(export)
    blanks, (runs, reached) = balance_counts(export)
    print(f"    read in {took:.1f}s; as on {export.as_on.isoformat()}, entries {export.first_date.isoformat()} to "
          f"{export.last_date.isoformat()}")
    for key, value in {**c, **blanks}.items():
        print(f"    {key}: {value}")
    # balance() rests on this: a blank Balance worked forward from the last stated one lands on the next stated one
    assert reached == runs, f"the forward balance reaches the next stated one in {reached} of {runs} runs of blank rows"
    print("    voucher prefixes: " + ", ".join(f"{p or '(none)'} {n}" for p, n in sorted(export.prefixes.items())))
    print("    unknown prefixes: " + ("; ".join(export.unknown) if export.unknown else "none"))
    for note in export.notes:
        print(f"    note: {note}")


def main() -> int:
    passed = failed = 0
    for name, fn in TESTS:
        try:
            fn()
            print(f"PASS  {name}")
            passed += 1
        except Exception:  # noqa: BLE001
            print(f"FAIL  {name}")
            traceback.print_exc()
            failed += 1
    paths.set_data_root(None)
    print(f"\n{passed} passed, {failed} failed")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
