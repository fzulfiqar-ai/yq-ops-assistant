"""R1 data-honesty tests (24-Sep-2026): the day-first twin rule, the Focus price-book snapshot
rule (un-void, staged gate, mass-void guard, audit rows), the badge evidence floors, honest social
proof, the covered-lines discount cap, the honest coupon message, the purchase_costs floor fallback
and cost flags, the clearance age from the first price date, the verify price checks and the alias
matcher.

    python -m tests.test_r1_data

Same lightweight runner as tests/test_v3.py (no pytest). Every test here is pure -- no database
-- except the last one, which reads production read-only through the yq_readonly RPC and SKIPs
(prints, never fails) without credentials. Nothing in this file writes anywhere.
"""
from __future__ import annotations

import io
import os
import sys
import tempfile
import traceback
import types
from datetime import datetime, timedelta, timezone
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from dotenv import load_dotenv  # noqa: E402
load_dotenv(ROOT / ".env")

import warnings  # noqa: E402
warnings.filterwarnings("ignore")

TESTS: list[tuple[str, object]] = []


def test(name):
    def deco(fn):
        TESTS.append((name, fn))
        return fn
    return deco


# ── synthetic context (same shape as tests/test_shop.py, plus the evidence columns) ──

def _item(code, price, stock=100, cat="CABLE", moq=1, sold_90d=0, sold_30d=0, prev_30d=0, customers_30d=0,
          created_at="2025-07-03T00:00:00+00:00", **extra):
    it = {"item_code": code, "display_name": code, "spec": f"{code} spec", "category": cat, "brand": "VFAN",
          "standard_rate": price, "b2c_rate": None, "product_image_url": None, "package_image_url": None,
          "sort_order": None, "created_at": created_at, "moq": moq, "pack_size": None,
          "stock_qty": stock, "stock_as_of": "2026-09-21", "sold_30d": sold_30d, "prev_30d": prev_30d,
          "sold_90d": sold_90d, "customers_30d": customers_30d}
    it.update(extra)
    return it


def _rule(id, kind, name=None, item_codes=(), categories=(), referral_codes=(), **kw):
    r = {"id": id, "kind": kind, "name": name or f"rule{id}", "stackable": False, "min_qty": None,
         "min_value_bhd": None, "pct_off": None, "amount_off_bhd": None, "fixed_price_bhd": None,
         "coupon_code": None, "starts_at": None, "ends_at": None, "max_uses": None, "uses": 0, "priority": 100,
         "scope": {"item_codes": [c.upper() for c in item_codes], "categories": [c.upper() for c in categories],
                   "referral_codes": [c.lower() for c in referral_codes]}}
    r.update(kw)
    return r


def _ctx(items, rules=(), costs=None, sources=None, **settings):
    from app.shop import SETTING_DEFAULTS
    vals = dict(SETTING_DEFAULTS)
    vals.update({k: str(v) for k, v in settings.items()})
    return {"settings": vals, "items": {i["item_code"]: i for i in items},
            "order": [i["item_code"] for i in items], "costs": dict(costs or {}), "cost_sources": dict(sources or {}),
            "rules": list(rules), "salesmen": [], "loaded_at": ""}   # loaded_at "" = synthetic: never touches the database


def _sp(sku, start, rate=1.0, book="MA_base", src="MASellingPriceBook (40).xlsx", cust=None, wh=None, id=None,
        voided=None):
    return {"id": id, "sku_code": sku, "price_book": book, "customer_code": cust, "warehouse_name": wh,
            "start_date": start, "rate_bhd": rate, "source_file": src, "voided_at": voided}


# ── 1. the day-first twin rule ────────────────────────────────────────────────

@test("twin rule: a (40) row with day/month swapped, same key and rate, is the twin of a _36_ row")
def _():
    from scripts.load_supabase import is_dayfirst_twin
    ph = _sp("UK10 C", "2026-09-06", 2.0, src="MASellingPriceBook _36_.xlsx")      # day-first reading of 9-Jun
    assert is_dayfirst_twin(ph, _sp("UK10 C", "2026-06-09", 2.0))
    assert is_dayfirst_twin(_sp("M20", "2026-10-05", 1.7, src="x"), _sp("M20", "2026-05-10", 1.7))   # the 5-Oct change
    assert is_dayfirst_twin(_sp("F04", "2026-11-01", 3.0, src="x"), _sp("F04", "2026-01-11", 3.0))   # the 1-Nov reverts
    assert not is_dayfirst_twin(ph, _sp("UK10 C", "2026-06-09", 1.9))                   # rate differs
    assert not is_dayfirst_twin(ph, _sp("UK10 L", "2026-06-09", 2.0))                   # SKU differs
    assert not is_dayfirst_twin(ph, _sp("UK10 C", "2026-06-09", 2.0, wh="Causeway"))    # layer differs
    assert not is_dayfirst_twin(ph, _sp("UK10 C", "2026-06-09", 2.0, book="modern_trade"))
    assert not is_dayfirst_twin(ph, _sp("UK10 C", "2026-09-06", 2.0))                   # same date = no swap
    assert not is_dayfirst_twin(_sp("A", "2026-09-17", 1.0), _sp("A", "2026-09-17", 1.0))   # day > 12: unambiguous
    assert not is_dayfirst_twin(_sp("A", "2026-03-03", 1.0), _sp("A", "2026-03-03", 1.0))   # day == month
    assert not is_dayfirst_twin(_sp("A", None, 1.0), _sp("A", "2026-03-04", 1.0))


@test("migration text: 5 views filter voided rows, twin against any MA book, loader rule once, asserts 165, no deletes")
def _():
    sql = (ROOT / "scripts" / "selling_prices_void_migration.sql").read_text(encoding="utf-8")
    low = sql.lower()
    assert "'dayfirst-parse 14-sep (twin of masellingpricebook (40))'" in low
    assert "n_voided <> 165" in low and "raise exception" in low
    assert not any(s in low for s in ("delete from", "drop table", "drop view", "drop column", "truncate"))
    assert "add column if not exists voided_at" in low and "add column if not exists void_reason" in low
    assert "grantee in ('anon', 'authenticated')" in low
    assert "insert into audit_log" in low
    for v in ("v_price_list_by_book", "v_price_change", "v_price_list", "v_price_history", "v_product_margin"):
        assert low.count(f"create or replace view {v} as") == 1, v
    # every view body carries the filter (5 view bodies + the twin subquery + the newest-export CTE/update)
    assert low.count("voided_at is null") >= 8
    # the twin is looked for in ANY live MA Focus book other than _36_, not only '(40)'
    assert "t.source_file like 'masellingpricebook%'" in low and "t.source_file <> 'masellingpricebook _36_.xlsx'" in low
    assert "t.source_file = 'masellingpricebook (40).xlsx'" not in low
    # the loader's snapshot rule, run once, keyed on import time and self-checked to leave 0 stale rows
    assert "selling_prices_void_migration.sql)'" in low and "max(imported_at) desc" in low
    assert "n_stale <> 0" in low and "n_newest <> 0" in low
    assert "foreach v in array array['v_price_list_by_book', 'v_price_change', 'v_price_list', 'v_price_history', 'v_product_margin']" in low
    # the reverse restores all five views and un-voids exactly what the migration voided
    rev = (ROOT / "scripts" / "selling_prices_void_reverse.sql").read_text(encoding="utf-8").lower()
    for v in ("v_price_list_by_book", "v_price_change", "v_price_list", "v_price_history", "v_product_margin"):
        assert f"create or replace view {v} as" in rev, v
    assert "set voided_at = null" in rev and "%selling_prices_void_migration.sql%" in rev
    views = rev[rev.index("create or replace view"):rev.index("revoke all on")]   # the five restored view bodies
    assert "voided_at" not in views, "a restored view still filters on voided_at"
    vel = (ROOT / "scripts" / "catalog_velocity_v2_migration.sql").read_text(encoding="utf-8").lower()
    assert "shops_30d" in vel and "shops_90d" in vel and "coalesce(v.is_cash_customer, false) = false" in vel
    assert (ROOT / "scripts" / "catalog_velocity_v2_reverse.sql").exists()


# ── 2. the Focus book snapshot rule (loader) ──────────────────────────────────

@test("snapshot: file names -> book kind (portal underscore variant included); the download counter is never read")
def _():
    from scripts import load_supabase as ls
    assert ls.focus_book_kind("MASellingPriceBook (40).xlsx") == "MA_base"
    assert ls.focus_book_kind("MASellingPriceBook _36_.xlsx") == "MA_base"
    assert ls.focus_book_kind("ModernTradeSellerBook _7_.xlsx") == "modern_trade"
    assert ls.focus_book_kind("ModernTradeSellerBook.xlsx") == "modern_trade"
    assert ls.focus_book_kind("YQ_MRN_Landing_Cost_Analysis Sep 2026.xlsx") is None
    assert ls.focus_book_kind(None) is None and ls.focus_book_kind("") is None
    assert not hasattr(ls, "focus_book_seq"), "the browser download counter must not order exports"
    k = ls.pricebook_key({"sku_code": " T02 ", "price_book": "MA_base", "customer_code": float("nan"),
                          "warehouse_name": None, "start_date": "2026-03-26T00:00:00"})
    assert k == ("T02", "MA_base", "", "", "2026-03-26"), k
    assert k == ls.pricebook_key({"sku_code": "T02", "price_book": "MA_base", "customer_code": "",
                                  "warehouse_name": "", "start_date": "2026-03-26"})
    with tempfile.TemporaryDirectory() as d:
        (Path(d) / "MASellingPriceBook (41).xlsx").write_bytes(b"x")
        (Path(d) / "Sales_day_book.xlsx").write_bytes(b"x")
        (Path(d) / "notes.txt").write_bytes(b"x")
        assert ls.staged_files(d) == {"MASellingPriceBook (41).xlsx", "Sales_day_book.xlsx"}
    assert ls.staged_files(Path(tempfile.gettempdir()) / "does-not-exist-yq") == set()


@test("snapshot: the loaded file is newest by definition -- older rows not in it are stale whatever their counter")
def _():
    from scripts.load_supabase import stale_pricebook_rows
    loaded = [_sp("T02", "2026-03-26", 2.95, src="MASellingPriceBook _41_.xlsx"),
              _sp("T02", "2026-03-26", 3.55, src="MASellingPriceBook _41_.xlsx", wh="Causeway"),
              _sp("UK10 C", "2026-06-09", 2.0, src="MASellingPriceBook _41_.xlsx")]
    existing = [
        _sp("T02", "2026-03-26", 2.95, src="MASellingPriceBook (40).xlsx", id=1),          # still in file -> keep
        _sp("T02", "2026-03-26", 3.55, src="MASellingPriceBook (40).xlsx", wh="Causeway", id=2),   # keep
        _sp("UK10 C", "2026-09-06", 2.0, src="MASellingPriceBook _36_.xlsx", id=3),        # phantom, not in file -> void
        _sp("UK10 C", "2026-06-09", 2.0, src="MASellingPriceBook (40).xlsx", id=4),        # in file -> keep
        _sp("M20", "2026-05-10", 1.7, src="MASellingPriceBook (40).xlsx", id=5),           # dropped by Focus -> void
        _sp("M20", "2026-09-14", 1.6, src="YQ_MRN_Landing_Cost_Analysis Sep 2026.xlsx", id=6),  # workbook -> keep
        _sp("X01", "2026-09-17", 0.6, src="MASellingPriceBook (42).xlsx", id=7),           # HIGHER counter, still older -> void
        _sp("T02", "2026-05-13", 3.9, src="ModernTradeSellerBook _7_.xlsx", book="modern_trade", id=8),  # other book -> keep
        _sp("T02", "2026-03-26", 2.95, src="MASellingPriceBook _41_.xlsx", id=9),          # this file -> keep
    ]
    stale = stale_pricebook_rows(loaded, existing)
    assert sorted(r["id"] for r in stale) == [3, 5, 7], [r["id"] for r in stale]
    # a load with no Focus book (workbook import only) voids nothing
    assert stale_pricebook_rows([_sp("M20", "2026-09-14", 1.6, src="YQ_MRN_x.xlsx")], existing) == []
    # a plain 'MASellingPriceBook.xlsx' (another PC, cleared Downloads) is as newest as any other
    plain = [_sp("T02", "2026-03-26", 2.95, src="MASellingPriceBook.xlsx")]
    assert [r["id"] for r in stale_pricebook_rows(plain, existing[:1] + existing[4:5])] == [5]
    # the staged gate: only a book whose file ingest read THIS run may void
    assert stale_pricebook_rows(loaded, existing, staged={"MASellingPriceBook _41_.xlsx", "Sales_day_book.xlsx"}) and \
        sorted(r["id"] for r in stale_pricebook_rows(loaded, existing, staged={"MASellingPriceBook _41_.xlsx"})) == [3, 5, 7]
    assert stale_pricebook_rows(loaded, existing, staged={"Sales_day_book.xlsx"}) == []      # leftover CSV: nothing
    assert stale_pricebook_rows(loaded, existing, staged=set()) == []


@test("snapshot: a voided key the file carries again is revived (the blocker: re-loading a voided key)")
def _():
    from scripts.load_supabase import revived_pricebook_rows
    loaded = [_sp("C01", "2026-11-01", 1.2, src="MASellingPriceBook _41_.xlsx"),        # a real future price
              _sp("T02", "2026-03-26", 2.95, src="MASellingPriceBook _41_.xlsx")]
    voided = [_sp("C01", "2026-11-01", 1.2, src="MASellingPriceBook _41_.xlsx", id=1, voided="2026-09-25T00:00:00Z"),  # phantom key, just upserted
              _sp("M20", "2026-10-05", 1.7, src="MASellingPriceBook _36_.xlsx", id=2, voided="2026-09-25T00:00:00Z"),  # not in file -> stays voided
              _sp("T02", "2026-03-26", 2.95, src="x.xlsx", book="modern_trade", id=3, voided="2026-09-25T00:00:00Z")]  # other book
    assert [r["id"] for r in revived_pricebook_rows(loaded, voided)] == [1]
    assert revived_pricebook_rows([_sp("C01", "2026-11-01", 1.2, src="YQ_MRN_x.xlsx")], voided) == []   # workbook: not a Focus book


@test("snapshot: the mass-void guard -- over 10% stale rows, or a file naming under 90% of the live SKUs, voids nothing")
def _():
    from scripts.load_supabase import snapshot_guard
    live = [_sp(f"S{i:02}", "2026-01-01", 1.0, id=i) for i in range(1, 101)]      # 100 live Focus rows, 100 SKUs
    file_skus = {f"S{i:02}" for i in range(1, 101)}
    assert snapshot_guard([], live, file_skus) is None
    assert snapshot_guard(live[:10], live, file_skus) is None                    # exactly 10% is allowed
    r = snapshot_guard(live[:11], live, file_skus)
    assert r and "11 stale rows" in r and "10%" in r, r
    r = snapshot_guard(live[:5], live, {f"S{i:02}" for i in range(1, 90)})       # 89 of 100 SKUs named
    assert r and "89 of the book's 100 live SKUs" in r, r
    assert snapshot_guard(live[:5], live, {f"S{i:02}" for i in range(1, 91)}) is None   # 90 of 100
    assert snapshot_guard(live[:1], [], set()) is None                           # a first-ever load: nothing to compare


class _FakeQ:
    """A tiny in-memory PostgREST: select / eq / is_ / not_.is_ / in_ / order / range / limit / update / insert."""

    def __init__(self, db, table):
        self.db, self.t = db, table
        self.filters, self.used, self.op, self.payload, self.rng, self.neg = [], [], "select", None, None, False

    def select(self, cols="*", **k):
        self.op = "select"
        self.used += [c.strip() for c in str(cols).split(",")]
        return self

    def eq(self, c, v):
        self.used.append(c); self.filters.append(lambda r: r.get(c) == v); return self

    def is_(self, c, v):
        self.used.append(c)
        neg, self.neg = self.neg, False
        self.filters.append((lambda r: r.get(c) is not None) if neg else (lambda r: r.get(c) is None))
        return self

    @property
    def not_(self):
        self.neg = True; return self

    def in_(self, c, vals):
        s = set(vals); self.filters.append(lambda r: r.get(c) in s); return self

    def order(self, *a, **k): return self
    def limit(self, *a, **k): return self

    def range(self, a, b):
        self.rng = (a, b); return self

    def update(self, payload):
        self.op, self.payload = "update", payload; return self

    def insert(self, payload):
        self.op, self.payload = "insert", payload; return self

    def execute(self):
        rows = self.db.setdefault(self.t, [])
        if self.t == "selling_prices" and not self.db.get("_has_void", True) and \
                ("voided_at" in self.used or (self.payload and "voided_at" in self.payload)):
            raise RuntimeError("column selling_prices.voided_at does not exist")
        if self.op == "select":
            sel = sorted((r for r in rows if all(f(r) for f in self.filters)), key=lambda r: r.get("id") or 0)
            if self.rng:
                sel = sel[self.rng[0]:self.rng[1] + 1]
            return types.SimpleNamespace(data=[dict(r) for r in sel], count=None)
        if self.op == "update":
            n = 0
            for r in rows:
                if all(f(r) for f in self.filters):
                    r.update(self.payload); n += 1
            return types.SimpleNamespace(data=[], count=n)
        rows.append(dict(self.payload))
        return types.SimpleNamespace(data=[self.payload], count=None)


class _FakeDB:
    def __init__(self, tables):
        self.tables = tables

    def table(self, name):
        if name not in self.tables and not name.startswith("_") and name != "audit_log":
            raise RuntimeError(f"table {name} unavailable")
        return _FakeQ(self.tables, name)


def _snapshot_db(has_void=True):
    rows = [_sp(f"S{i:02}", "2026-01-01", 1.0, id=i) for i in range(1, 13)]                  # 12 live (40) rows
    rows.append(_sp("M20", "2026-05-10", 1.7, id=13))                                          # live, about to vanish
    rows.append(_sp("C01", "2026-11-01", 1.2, src="MASellingPriceBook _36_.xlsx", id=14,
                    voided="2026-09-24T10:00:00+00:00"))                                       # a voided phantom key
    rows.append(_sp("M20", "2026-09-14", 1.6, src="YQ_MRN_Landing_Cost_Analysis Sep 2026.xlsx", id=15))  # workbook
    rows.append(_sp("T02", "2026-05-13", 3.9, src="ModernTradeSellerBook _7_.xlsx", book="modern_trade", id=16))
    for r in rows:
        r["void_reason"] = "phantom" if r["voided_at"] else None
    return {"selling_prices": rows, "audit_log": [], "_has_void": has_void}


@test("snapshot: apply -- void in load N, the key returns in load N+1 and is live again; one audit row per batch")
def _():
    from scripts.load_supabase import _apply_pricebook_snapshot
    db = _snapshot_db()
    client = _FakeDB(db)
    file_n = "MASellingPriceBook _41_.xlsx"
    loaded = [_sp(f"S{i:02}", "2026-01-01", 1.0, src=file_n) for i in range(1, 13)]
    loaded.append(_sp("C01", "2026-11-01", 1.2, src=file_n))       # Focus now really holds C01 on 1-Nov
    res = _apply_pricebook_snapshot(client, loaded, staged={file_n, "Sales_day_book.xlsx"})
    by_id = {r["id"]: r for r in db["selling_prices"]}
    assert res["revived"] == 1 and res["voided"] == 1 and not res["skipped"], res
    assert by_id[14]["voided_at"] is None and by_id[14]["void_reason"] is None          # C01 lives again
    assert by_id[13]["voided_at"] and by_id[13]["void_reason"].startswith("superseded by MASellingPriceBook _41_.xlsx")
    assert by_id[15]["voided_at"] is None and by_id[16]["voided_at"] is None            # workbook + other book untouched
    assert all(by_id[i]["voided_at"] is None for i in range(1, 13))
    assert len(db["audit_log"]) == 1
    a = db["audit_log"][0]
    assert a["event"] == "selling_prices.void" and a["detail"]["count"] == 1 and a["detail"]["skus"] == ["M20"]
    assert a["detail"]["book"] == "MA_base" and a["detail"]["file"] == file_n and a["user_email"] == "load_supabase"
    # load N+1: M20 on 2026-05-10 is back in the book -> revived; nothing else stale
    file_n1 = "MASellingPriceBook.xlsx"
    loaded = [_sp(f"S{i:02}", "2026-01-01", 1.0, src=file_n1) for i in range(1, 13)]
    loaded += [_sp("C01", "2026-11-01", 1.2, src=file_n1), _sp("M20", "2026-05-10", 1.7, src=file_n1)]
    res = _apply_pricebook_snapshot(client, loaded, staged={file_n1})
    assert res["revived"] == 1 and res["voided"] == 0, res
    assert by_id[13]["voided_at"] is None and by_id[13]["void_reason"] is None
    assert len(db["audit_log"]) == 1                                                    # no void, no audit row
    # a re-run of the same file changes nothing
    res = _apply_pricebook_snapshot(client, loaded, staged={file_n1})
    assert res == {"revived": 0, "voided": 0, "skipped": {}}, res


@test("snapshot: apply -- the guard skips a filtered export (audit 'void_skipped'), an unstaged or unflagged book never voids")
def _():
    from scripts.load_supabase import _apply_pricebook_snapshot
    db = _snapshot_db()
    client = _FakeDB(db)
    partial = [_sp("S01", "2026-01-01", 1.0, src="MASellingPriceBook _41_.xlsx"),      # one item group exported
               _sp("S02", "2026-01-01", 1.0, src="MASellingPriceBook _41_.xlsx")]
    res = _apply_pricebook_snapshot(client, partial, staged={"MASellingPriceBook _41_.xlsx"})
    assert res["voided"] == 0 and "MA_base" in res["skipped"] and "stale rows" in res["skipped"]["MA_base"], res
    assert all(r["voided_at"] is None for r in db["selling_prices"] if r["id"] != 14)
    assert len(db["audit_log"]) == 1 and db["audit_log"][0]["event"] == "selling_prices.void_skipped"
    assert db["audit_log"][0]["detail"]["stale_rows"] == 11 and db["audit_log"][0]["detail"]["live_rows"] == 13
    # the file is not part of this run's upload (a leftover CSV): un-void still happens, the void does not
    db = _snapshot_db()
    client = _FakeDB(db)
    full = [_sp(f"S{i:02}", "2026-01-01", 1.0, src="MASellingPriceBook _41_.xlsx") for i in range(1, 13)]
    full.append(_sp("C01", "2026-11-01", 1.2, src="MASellingPriceBook _41_.xlsx"))
    res = _apply_pricebook_snapshot(client, full, staged={"Sales_day_book.xlsx"})
    assert res["revived"] == 1 and res["voided"] == 0 and res["skipped"] == {"MA_base": "file not staged this run"}, res
    assert db["audit_log"] == []
    res = _apply_pricebook_snapshot(client, full, staged=None)
    assert res["voided"] == 0 and res["skipped"] == {"MA_base": "no staged folder"}, res
    # before the migration (no void columns): nothing changes, the load is unaffected
    db = _snapshot_db(has_void=False)
    res = _apply_pricebook_snapshot(_FakeDB(db), full, staged={"MASellingPriceBook _41_.xlsx"})
    assert res["revived"] == 0 and res["voided"] == 0 and "voided_at" in res["skipped"]["MA_base"], res
    assert db["audit_log"] == []
    # a workbook-only load never touches the snapshot machinery
    assert _apply_pricebook_snapshot(_FakeDB(_snapshot_db()), [_sp("M20", "2026-09-14", 1.6, src="YQ_MRN_x.xlsx")],
                                     staged={"YQ_MRN_x.xlsx"}) == {"revived": 0, "voided": 0, "skipped": {}}


@test("snapshot: refresh clears data/clean and passes --staged; the workbook import un-voids the row it updates")
def _():
    from scripts.load_supabase import void_columns_present
    assert void_columns_present(_FakeDB(_snapshot_db())) is True
    assert void_columns_present(_FakeDB(_snapshot_db(has_void=False))) is False
    rf = (ROOT / "scripts" / "refresh.py").read_text(encoding="utf-8")
    assert 'shutil.rmtree(ROOT / "data" / "clean", ignore_errors=True)' in rf
    assert '_run("scripts.load_supabase", "--staged", src_path)' in rf
    iw = (ROOT / "scripts" / "import_workbook_prices.py").read_text(encoding="utf-8")
    assert "void_columns_present(client)" in iw and '{"voided_at": None, "void_reason": None}' in iw
    assert ".update({**row, **unvoid})" in iw


# ── 3. badges, social proof, settings ─────────────────────────────────────────

@test("badges: best seller needs invoices_90d >= 10 and shops_90d >= 5 when known; unknown keeps the volume rule")
def _():
    from app.shop import _badges
    items = [_item("BIG", 1.0, sold_90d=900, invoices_90d=3, shops_90d=1),     # one shop's restock
             _item("REAL", 1.0, sold_90d=300, invoices_90d=40, shops_90d=12),
             _item("THIN", 1.0, sold_90d=250, invoices_90d=12, shops_90d=4)]   # too few shops
    b = _badges(_ctx(items, shop_best_seller_top_n="2"))
    assert "best_seller" in b["REAL"] and "best_seller" not in b["BIG"] and "best_seller" not in b["THIN"], b
    # the floors are settings
    b = _badges(_ctx(items, shop_best_seller_top_n="2", shop_best_seller_min_invoices="1", shop_best_seller_min_shops="1"))
    assert "best_seller" in b["BIG"] and "best_seller" in b["REAL"] and "best_seller" not in b["THIN"]
    # no evidence columns yet (pre-migration rows) -> the pre-R1 rule: top N by units
    b = _badges(_ctx([_item("A", 1.0, sold_90d=50), _item("B", 1.0, sold_90d=5)], shop_best_seller_top_n="1"))
    assert "best_seller" in b["A"] and "best_seller" not in b["B"]


@test("badges: selling fast needs invoices_90d >= 10 when known")
def _():
    from app.shop import _badges
    fast = _item("F", 1.0, stock=20, sold_90d=300, invoices_90d=30)           # 6 days of cover
    lone = _item("L", 1.0, stock=20, sold_90d=300, invoices_90d=2)            # same cover, 2 invoices
    old = _item("O", 1.0, stock=20, sold_90d=300)                             # no evidence column
    b = _badges(_ctx([fast, lone, old], shop_best_seller_top_n="0"))
    assert "selling_fast" in b["F"] and "selling_fast" not in b["L"] and "selling_fast" in b["O"], b
    b = _badges(_ctx([lone], shop_best_seller_top_n="0", shop_selling_fast_min_invoices="2"))
    assert "selling_fast" in b["L"]


@test("badges: clearance age runs from the first price date (first_seen), then created_at; default 180; 0 disables")
def _():
    from app.shop import _badges
    young = (datetime.now(timezone.utc) - timedelta(days=40)).date().isoformat()
    launch = (datetime.now(timezone.utc) - timedelta(days=86)).date().isoformat()      # the 30-Jun VFAN lines
    items = [_item("NEWLINE", 1.0, stock=600, sold_90d=10, first_seen=young),        # priced 40 days ago
             _item("JULY", 1.0, stock=600, sold_90d=10, first_seen=launch, created_at="2026-07-03T00:00:00+00:00"),
             _item("OLDLINE", 1.0, stock=600, sold_90d=10, first_seen="2025-09-17", created_at="2026-07-03T00:00:00+00:00"),
             _item("SEEDED", 1.0, stock=600, sold_90d=10, first_seen=None)]         # no price date: created_at (a year old)
    b = _badges(_ctx(items, shop_best_seller_top_n="0"))
    assert "clearance" in b["OLDLINE"] and "clearance" in b["SEEDED"], b
    assert "clearance" not in b["NEWLINE"] and "clearance" not in b["JULY"], b
    # a bulk-seeded created_at (2026-07-03 for 190 of 194 items) no longer hides an old line
    b = _badges(_ctx(items, shop_best_seller_top_n="0", shop_clearance_min_age_days="30"))
    assert "clearance" in b["NEWLINE"] and "clearance" in b["JULY"]
    b = _badges(_ctx(items, shop_best_seller_top_n="0", shop_clearance_min_age_days="0"))
    assert "clearance" in b["NEWLINE"]
    # an item with no date at all cannot be judged -> not blocked by the guard
    b = _badges(_ctx([_item("NODATE", 1.0, stock=600, sold_90d=10, created_at=None, first_seen=None)],
                     shop_best_seller_top_n="0"))
    assert "clearance" in b["NODATE"]


@test("social proof: named B2B shops only, 'in the last 30 days', never 0, customers_30d until the migration")
def _():
    from app.shop import social_proof_text
    ctx = _ctx([])
    assert social_proof_text(ctx, _item("A", 1.0, customers_30d=9, shops_30d=7), 5) == "Ordered by 7 shops in the last 30 days"
    assert social_proof_text(ctx, _item("A", 1.0, customers_30d=9, shops_30d=4), 5) is None     # 9 names, 4 real shops
    assert social_proof_text(ctx, _item("A", 1.0, customers_30d=6), 5) == "Ordered by 6 shops in the last 30 days"
    assert social_proof_text(ctx, _item("A", 1.0, customers_30d=0, shops_30d=0), 0) is None
    assert "this month" not in (social_proof_text(ctx, _item("A", 1.0, shops_30d=12), 5) or "")


@test("settings: the four badge keys are defaults and the retail anchor default is off (owner 24-Sep)")
def _():
    from app.shop import SETTING_DEFAULTS as D
    assert D["shop_clearance_min_age_days"] == "180" and D["shop_best_seller_min_invoices"] == "10"
    assert D["shop_best_seller_min_shops"] == "5" and D["shop_selling_fast_min_invoices"] == "10"
    assert D["shop_clearance_show_retail"] == "0"
    sql = (ROOT / "scripts" / "catalog_velocity_v2_migration.sql").read_text(encoding="utf-8")
    for k in ("shop_clearance_min_age_days", "shop_best_seller_min_invoices",
              "shop_best_seller_min_shops", "shop_selling_fast_min_invoices"):
        assert f"('{k}'," in sql, k


@test("evidence: loaded in _build_context (background), never on a request; a synthetic context has none")
def _():
    from app import shop
    ctx = _ctx([_item("A", 1.0)])
    real = shop.exec_sql
    shop.exec_sql = lambda sql: (_ for _ in ()).throw(RuntimeError("must not query on the request path"))
    try:
        assert shop._velocity_extra(ctx) == {} and ctx["_velocity_extra"] == {}
        ctx["_velocity_extra"] = {"A": {"shops_30d": 8}}
        assert shop._evidence(ctx, ctx["items"]["A"], "shops_30d") == 8
        assert shop._evidence(ctx, ctx["items"]["A"], "invoices_90d") is None
        assert shop._badges(ctx) is not None                                  # badges never hit the database
        # the loaders fall back cleanly when the view predates the migration / the RPC is down
        assert shop._load_velocity_extra() == {}
        assert shop._load_first_seen() == {}
    finally:
        shop.exec_sql = real
    src = (ROOT / "app" / "shop.py").read_text(encoding="utf-8")
    ctx_src = src[src.index("def _build_context"):src.index("def _refresh_in_background")]
    assert "_load_velocity_extra()" in ctx_src and "_load_first_seen()" in ctx_src and "prices_updated_date()" in ctx_src
    # first price date wins, then first sale, and the voided filter is tried first
    calls: list[str] = []

    def fake(sql):
        calls.append(sql)
        if "voided_at" in sql:
            raise RuntimeError("column voided_at does not exist")
        if "selling_prices" in sql:
            return [{"sku_code": "T17", "d": "2026-06-30"}]
        return [{"sku_code": "T17", "d": "2026-07-09"}, {"sku_code": "X99", "d": "2025-01-05"}]
    shop.exec_sql = fake
    try:
        assert shop._load_first_seen() == {"T17": "2026-06-30", "X99": "2025-01-05"}
    finally:
        shop.exec_sql = real
    assert "voided_at IS NULL" in calls[0] and "voided_at" not in calls[1]


@test("prices_updated: skips voided rows once the column exists, plain query before")
def _():
    from app import catalog
    real = catalog.exec_sql
    calls: list[str] = []

    def before(sql):
        calls.append(sql)
        if "voided_at" in sql:
            raise RuntimeError("column voided_at does not exist")
        return [{"d": "2026-09-20"}]
    catalog.exec_sql = before
    try:
        assert catalog.prices_updated_date() == "2026-09-20" and len(calls) == 2
        catalog.exec_sql = lambda sql: [{"d": "2026-09-19"}] if "voided_at IS NULL" in sql else [{"d": "wrong"}]
        assert catalog.prices_updated_date() == "2026-09-19"
    finally:
        catalog.exec_sql = real
    src = (ROOT / "app" / "shop.py").read_text(encoding="utf-8")
    assert "SELECT MAX(start_date)" not in src, "shop.py must take the stamp from catalog.prices_updated_date"


# ── 4. margin floor: covered-lines cap, honest coupon, the purchase_costs fallback ──

@test("pricing: the cart cap comes from covered lines only -- an uncovered line adds 0 headroom, never unlimited")
def _():
    from app.shop import price_cart
    coupon = _rule(8, "coupon", "HALF", coupon_code="HALF", pct_off=50)
    # A: cost 7 -> floor 8.4 -> 1.6/unit headroom; B: no cost on file -> nothing
    ctx = _ctx([_item("A", 10.0), _item("B", 10.0)], rules=[coupon], costs={"A": 7.0},
               shop_min_margin_pct="0.20", shop_vat_rate="0")
    q = price_cart([{"item_code": "A", "qty": 2}, {"item_code": "B", "qty": 2}], coupon_code="HALF", ctx=ctx)
    assert q["discount_bhd"] == 3.2 and q["total_bhd"] == 36.8 and "cart:8" in q["_clamped"], q
    # nothing covered -> no cart-level discount at all (before R1 this cart got BHD 20 off)
    ctx = _ctx([_item("A", 10.0), _item("B", 10.0)], rules=[coupon], costs={})
    q = price_cart([{"item_code": "A", "qty": 2}, {"item_code": "B", "qty": 2}], coupon_code="HALF", ctx=ctx)
    assert q["discount_bhd"] == 0 and q["total_bhd"] == 40.0 and "cart:8" in q["_clamped"], q
    # fully covered -> the full headroom, as before
    ctx = _ctx([_item("A", 10.0), _item("B", 10.0)], rules=[coupon], costs={"A": 7.0, "B": 7.0},
               shop_min_margin_pct="0.20", shop_vat_rate="0")
    q = price_cart([{"item_code": "A", "qty": 2}, {"item_code": "B", "qty": 2}], coupon_code="HALF", ctx=ctx)
    assert q["discount_bhd"] == 6.4, q
    # a stackable cart rule shares the same cap
    auto = _rule(7, "cart_value", "5% over 10", min_value_bhd=10, pct_off=5, stackable=True)
    ctx = _ctx([_item("A", 10.0), _item("B", 10.0)], rules=[coupon, auto], costs={"A": 7.0},
               shop_min_margin_pct="0.20", shop_vat_rate="0")
    q = price_cart([{"item_code": "A", "qty": 2}, {"item_code": "B", "qty": 2}], coupon_code="HALF", ctx=ctx)
    assert q["discount_bhd"] == 3.2, q


@test("pricing: a coupon the floor caps to 0 says so, is not 'applied' and is not spent; a partial cap says how much")
def _():
    from app.shop import price_cart
    coupon = _rule(8, "coupon", "HALF", coupon_code="HALF", pct_off=50, max_uses=1)
    # cap 0: a cart of floorless lines only
    ctx = _ctx([_item("A", 10.0), _item("B", 10.0)], rules=[coupon], costs={})
    q = price_cart([{"item_code": "A", "qty": 2}, {"item_code": "B", "qty": 2}], coupon_code="HALF", ctx=ctx)
    assert q["discount_bhd"] == 0 and q["coupon"]["valid"] is False and q["_coupon_rule_id"] is None, q["coupon"]
    assert "cannot lower these items further" in q["coupon"]["message"] and "applied" not in q["coupon"]["message"]
    assert q["discounts"] == []
    # partial cap: the message carries the capped amount, the code is valid and spent
    ctx = _ctx([_item("A", 10.0), _item("B", 10.0)], rules=[coupon], costs={"A": 7.0},
               shop_min_margin_pct="0.20", shop_vat_rate="0")
    q = price_cart([{"item_code": "A", "qty": 2}, {"item_code": "B", "qty": 2}], coupon_code="HALF", ctx=ctx)
    assert q["discount_bhd"] == 3.2 and q["coupon"]["valid"] is True and q["_coupon_rule_id"] == 8, q["coupon"]
    assert "capped at BHD 3.200" in q["coupon"]["message"] and "applied" in q["coupon"]["message"]
    # no cap: the plain 'applied'
    ctx = _ctx([_item("A", 10.0)], rules=[coupon], costs={"A": 1.0}, shop_min_margin_pct="0.20", shop_vat_rate="0")
    q = price_cart([{"item_code": "A", "qty": 2}], coupon_code="HALF", ctx=ctx)
    assert q["discount_bhd"] == 10.0 and q["coupon"] == {"code": "HALF", "valid": True, "message": "50% off applied"}, q["coupon"]


class _FakeQuery:
    def __init__(self, rows):
        self._rows = rows

    def select(self, *a, **k):
        return self

    def order(self, *a, **k):
        return self

    def limit(self, *a, **k):
        return self

    def execute(self):
        return types.SimpleNamespace(data=self._rows)


class _FakeClient:
    def __init__(self, tables):
        self.tables = tables

    def table(self, name):
        rows = self.tables.get(name)
        if rows is None:
            raise RuntimeError(f"table {name} unavailable")
        return _FakeQuery(rows)


@test("costs: MRN wins; purchase_costs fills the gaps by MAX(id); every cost under exact AND UPPER keys with its source")
def _():
    from app import shop
    real = shop.get_client
    shop.get_client = lambda: _FakeClient({
        "mrn_landed_costs": [{"sku_code": "T02", "landed_cost_bhd": 2.0, "effective_date": "2026-09-14", "id": 9},
                             {"sku_code": "X05 UL-1Mtr", "landed_cost_bhd": 0.3, "effective_date": "2026-09-14", "id": 8}],
        "purchase_costs": [                                              # ordered id desc, as the query asks
            {"id": 30, "sku_code": "t02", "landed_cost_bhd": 9.9},       # covered by MRN -> ignored
            {"id": 29, "sku_code": "X05 UL-1MTR", "landed_cost_bhd": 0.31},   # covered by MRN (case) -> ignored
            {"id": 28, "sku_code": "Tb-d1", "landed_cost_bhd": 0.25},
            {"id": 27, "sku_code": "TB-D1", "landed_cost_bhd": 0.2},         # older row -> ignored
            {"id": 26, "sku_code": "M20", "landed_cost_bhd": 0},             # no real cost -> ignored
        ]})
    sources: dict = {}
    try:
        costs = shop._load_costs(sources)
    finally:
        shop.get_client = real
    assert costs["T02"] == 2.0 and costs["X05 UL-1Mtr"] == 0.3 and costs["X05 UL-1MTR"] == 0.3, costs
    assert costs["Tb-d1"] == 0.25 and costs["TB-D1"] == 0.25 and "M20" not in costs, costs
    assert sources["T02"] == "mrn" and sources["X05 UL-1MTR"] == "mrn" and sources["TB-D1"] == "purchase_costs"
    ctx = _ctx([_item("X05 UL-1Mtr", 0.6), _item("T02", 2.95), _item("M20", 1.0), _item("tb-d1", 0.5)],
               costs=costs, sources=sources, shop_min_margin_pct="0.20", shop_vat_rate="0.10")
    assert shop.cost_for(ctx, "tb-d1") == 0.25 and shop.cost_source_for(ctx, "tb-d1") == "purchase_costs"
    assert shop.cost_for(ctx, "T02") == 2.0 and shop.cost_for(ctx, "M20") is None and shop.cost_source_for(ctx, "M20") is None
    assert shop._floor_for(ctx, "X05 UL-1Mtr") == 0.396 and shop._floor_for(ctx, "T02") == 2.64
    assert shop._floor_for(ctx, "M20") is None and shop._floor_for(ctx, "tb-d1") == 0.33
    # the fallback survives purchase_costs being unavailable; a plain call still returns the dict
    shop.get_client = lambda: _FakeClient({"mrn_landed_costs": [{"sku_code": "T02", "landed_cost_bhd": 2.0}]})
    try:
        assert shop._load_costs() == {"T02": 2.0}
    finally:
        shop.get_client = real


@test("costs: a cost under 10% of the list price is flagged; margin_health reads cost_for and carries the flag")
def _():
    from app import shop
    costs = {"X10 LT": 0.012571, "OK": 7.0, "UK20 (New)": 0.0114}
    sources = {"X10 LT": "purchase_costs", "OK": "mrn", "UK20 (New)": "mrn", "UK20 (NEW)": "mrn"}
    ctx = _ctx([_item("X10 LT", 0.4), _item("OK", 10.0), _item("UK20 (New)", 2.8), _item("NOCOST", 1.0)],
               costs=costs, sources=sources, shop_vat_rate="0.10", shop_min_margin_pct="0.20")
    flags = shop.cost_flags(ctx)
    assert set(flags) == {"X10 LT", "UK20 (New)"}, flags
    assert flags["X10 LT"].startswith("purchase_costs fallback cost BHD 0.0126 is under 10% of the BHD 0.400")
    assert flags["UK20 (New)"].startswith("landed cost BHD 0.0114 is under 10% of the BHD 2.800")
    ctx["cost_flags"] = flags
    real = shop.context
    shop.context = lambda: ctx
    try:
        mh = shop.margin_health()
    finally:
        shop.context = real
    rows = {r["item_code"]: r for r in mh["rows"]}
    assert rows["X10 LT"]["cost_flag"] == flags["X10 LT"] and rows["X10 LT"]["cost_source"] == "purchase_costs"
    assert rows["X10 LT"]["landed_cost_bhd"] == 0.013 and rows["X10 LT"]["status"] == "ok"   # the number is shown, the flag says why not to trust it
    assert rows["OK"]["cost_flag"] is None and rows["OK"]["cost_source"] == "mrn"
    assert rows["NOCOST"]["status"] == "no_cost" and rows["NOCOST"]["cost_source"] is None
    assert mh["summary"]["flagged_costs"] == 2 and mh["summary"]["with_cost"] == 3
    # margin_health finds a cost stored under the UPPER key for a mixed-case catalog code
    ctx2 = _ctx([_item("Tb-d1", 0.5)], costs={"TB-D1": 0.25}, sources={"TB-D1": "purchase_costs"},
                shop_vat_rate="0", shop_min_margin_pct="0.20")
    ctx2["cost_flags"] = shop.cost_flags(ctx2)
    shop.context = lambda: ctx2
    try:
        r = shop.margin_health()["rows"][0]
    finally:
        shop.context = real
    assert r["status"] == "ok" and r["landed_cost_bhd"] == 0.25 and r["floor_bhd"] == 0.3 and r["cost_source"] == "purchase_costs"


# ── 5. verify_numbers: price checks ───────────────────────────────────────────

def _pb(sku, start, rate, end=None, wh=None, status="Authorized", src="MASellingPriceBook (40).xlsx"):
    return {"sku_code": sku, "item_name": sku, "price_book": "MA_base", "customer_code": None, "warehouse_name": wh,
            "start_date": start, "end_date": end, "rate_bhd": rate, "status": status, "source_file": src}


@test("verify: current_prices mirrors v_price_list_by_book (base layer first, latest start, later row, live window)")
def _():
    from scripts.verify_numbers import current_prices
    rows = [
        _pb("T02", "2026-03-26", 3.55, end="2027-11-10"),
        _pb("T02", "2026-03-26", 2.95),                         # same key later in the file -> the id-desc winner
        _pb("T02", "2026-03-26", 4.50, wh="Causeway"),          # outlet layer never beats the base layer
        _pb("UK10 C", "2026-06-09", 2.0),
        _pb("UK10 C", "2026-09-06", 1.9),                       # later start wins
        _pb("UK10 C", "2026-11-01", 1.5),                       # not started yet
        _pb("OLD", "2025-01-01", 9.0, end="2026-01-01"),        # ended
        _pb("NOPE", "2026-01-01", 1.0, status="Cancelled"),
        _pb("ZERO", "2026-01-01", 0),
    ]
    cur = current_prices(rows, today="2026-09-24")
    assert cur == {"T02": 2.95, "UK10 C": 1.9}, cur
    assert current_prices(rows, today="2026-11-02")["UK10 C"] == 1.5


@test("verify: export_date is the earlier of docProps.created and the file mtime; the newest export in a folder is checked")
def _():
    from openpyxl import Workbook
    from scripts.verify_numbers import _find_newest_book, export_date
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "MASellingPriceBook (40).xlsx"
        wb = Workbook()
        wb.properties.created = datetime(2026, 9, 21, 11, 48, 38)
        wb.save(p)
        assert export_date(p) == "2026-09-21", export_date(p)       # today's mtime is later
        p2 = Path(d) / "MASellingPriceBook (1).xlsx"                 # a lower counter, a newer export
        wb.properties.created = datetime(2026, 9, 23, 9, 0, 0)
        wb.save(p2)
        assert os.path.basename(_find_newest_book("masellingpricebook", Path(d))) == "MASellingPriceBook (1).xlsx"
        assert _find_newest_book("moderntradesellerbook", Path(d)) is None
        wb.properties.created = datetime(2099, 1, 1)
        wb.save(p)
        assert export_date(p) == datetime.now().date().isoformat()  # a clock ahead of the file never wins


@test("verify: the price-book checks pass on a matching load; a ghost, a stale older-export row and a skipped void all FAIL")
def _():
    from scripts import verify_numbers as vn
    rows = [_pb("T02", "2026-03-26", 2.95), _pb("UK10 C", "2026-06-09", 2.0), _pb("F04", "2026-11-01", 3.0)]
    names = ("_find_newest_book", "read_grid", "parse_pricebook", "export_date", "_db_book_prices",
             "_db_live_focus_rows", "_db_today")
    saved = {n: getattr(vn, n) for n in names}
    vn._find_newest_book = lambda key, src: "MASellingPriceBook (40).xlsx" if key == "masellingpricebook" else None
    vn.read_grid = lambda f: None
    vn.parse_pricebook = lambda grid, src, book: rows
    vn.export_date = lambda f: "2026-09-21"
    vn._db_today = lambda: "2026-09-24"
    try:
        vn._db_book_prices = lambda book: {"T02": 2.95, "UK10 C": 2.0}
        vn._db_live_focus_rows = lambda book, key: [_pb("F04", "2026-11-01", 3.0)]     # in the file: a scheduled price
        checks: list = []
        vn._price_book_checks(Path("."), checks)
        assert [(c[0].split(" ")[0], c[1], c[2]) for c in checks] == \
            [("MA_base", 2.0, 2.0), ("MA_base", 0.0, 0.0), ("MA_base", 0.0, 0.0)], checks
        assert checks[1][0] == "MA_base rows after 2026-09-21 not in file"
        assert checks[2][0] == "MA_base live rows not in file (older exports)"
        # a phantom future row the file lacks, and an older-export row the void skipped (both live)
        vn._db_book_prices = lambda book: {"T02": 2.95, "UK10 C": 1.9}
        vn._db_live_focus_rows = lambda book, key: [_pb("F04", "2026-11-01", 3.0),
                                                    _pb("M20", "2026-10-05", 1.7, src="MASellingPriceBook _36_.xlsx"),
                                                    _pb("F25", "2026-03-26", 6.0, wh="Moideen KP", src="MASellingPriceBook (32).xlsx")]
        checks = []
        vn._price_book_checks(Path("."), checks)
        assert checks[0][1:3] == (2.0, 1.0) and checks[1][1:3] == (0.0, 1.0) and checks[2][1:3] == (0.0, 2.0), checks
        # no price book in the upload -> no price check at all
        vn._find_newest_book = lambda key, src: None
        checks = []
        vn._price_book_checks(Path("."), checks)
        assert checks == []
    finally:
        for n, f in saved.items():
            setattr(vn, n, f)
    assert "selling_prices" not in vn._UNCHECKED
    # the ghost/stale query is scoped to Focus exports of the book, never workbook rows
    src = (ROOT / "scripts" / "verify_numbers.py").read_text(encoding="utf-8")
    assert "source_file ILIKE $2" in src and "AND voided_at IS NULL" in src and "CURRENT_DATE::text" in src


# ── 6. alias matcher and the backfill preview ─────────────────────────────────

@test("matcher: exact name, longest code on a token boundary, whole-name prefix; the 40-char shortcut is gone")
def _():
    from scripts.reconcile_products import build_index, match_alias
    idx = build_index([
        ("T17", "T17 Open Ear Clip On - Type-C Port (VFAN) T17"),
        ("X24 CC", "X24 2Mtr Cable (Type-C to Type-C) (VFAN) X24 CC"),
        ("X24 CC 1Mtr", "X24 CC 1Mtr Cable (Type-C to Type-C) (VFAN) X24 CC"),
        ("X02-M", "X02 3A 1.8Mtr NYLON BRAIDED MICRO USB Cable (VFAN)"),
        ("X02-L", "X02 3A 1.8Mtr NYLON BRAIDED LIGHTNING USB Cable (VFAN)"),
        ("LONG", "A very long product name that runs past forty characters and then keeps going"),
    ])
    assert match_alias("T17 Open Ear Clip On - Type-C Port (VFAN)  T17", idx)[:2] == ("T17", "exact_name")
    assert match_alias("X24 CC 1Mtr Cable (Type-C to Type-C) (VFAN) X24 CC", idx)[:2] == ("X24 CC 1Mtr", "exact_name")
    assert match_alias("X24 CC 1Mtr something new", idx)[:2] == ("X24 CC 1Mtr", "code")      # longest code, not X24 CC
    assert match_alias("X24 CC 2Mtr something new", idx)[:2] == ("X24 CC", "code")
    assert match_alias("T170 other", idx)[0] is None                                          # T17 is not a prefix token
    assert match_alias("X02-M 3A cable", idx)[:2] == ("X02-M", "code")
    assert match_alias("X02 3A 1.8Mtr NYLON BRAIDED LIGHTNING USB Cable (VFAN) X02-L", idx)[:2] == ("X02-L", "name_prefix")
    assert match_alias("X02 something else entirely", idx)[:2] == ("X02-M", "code_family")
    # a string sharing only the first 40 characters of a longer product name is NOT that product any more
    forty = "A very long product name that runs past forty characters and then diverges"
    assert match_alias(forty, idx)[0] is None, match_alias(forty, idx)
    assert match_alias("A very", idx)[0] is None                       # a short string no product name starts
    assert match_alias("", idx) == (None, "unmatched", 0.0)
    sku, method, conf = match_alias("T17 Open Ear Clip On - Type-C Port (VFAN) T17", idx)
    assert conf == 1.0 and match_alias("X02 zzz", idx)[2] == 0.6


@test("preview: propose() keeps the exact item string; --commit is exact-name only, insert-only, honest about the cache")
def _():
    from scripts.alias_backfill_preview import propose
    products = [{"id": 11, "sku_code": "T17", "item_name": "T17 Open Ear Clip On - Type-C Port (VFAN) T17"},
                {"id": 12, "sku_code": "X31 CC 1 Mtr", "item_name": "X31 CC 1 Mtr PD Fast (60W) (Type-C to Type-C) (VFAN) X31 CC"}]
    una = [{"item_name": "T17 Open Ear Clip On - Type-C Port (VFAN) T17", "division": "Accessories", "lines": 47,
            "units": 177, "rev_365": 1200.5, "rev_90": 766.7, "last_sold": "2026-09-20"},
           {"item_name": "X34 CC 1 Mtr Invisible Stand Cable (240W) (VFAN) X34 CC", "division": "Accessories",
            "lines": 55, "units": 300, "rev_365": 480.5, "rev_90": 480.5, "last_sold": "2026-09-22"},
           {"item_name": "Batelco Vanilla Sim", "division": "SIM", "lines": 67, "units": 8288, "rev_365": 2900.8,
            "rev_90": 700.0, "last_sold": "2026-09-24"}]
    out = propose(una, products)
    assert [o["proposed_code"] for o in out] == ["T17", None, None]
    assert out[0]["product_id"] == 11 and out[0]["confidence"] == 1.0 and out[0]["rev_90"] == 766.7
    assert out[0]["item_name"] == una[0]["item_name"]              # alias_text must be the exact Focus string
    assert out[2]["division"] == "SIM" and out[1]["method"] == "unmatched"
    src = (ROOT / "scripts" / "alias_backfill_preview.py").read_text(encoding="utf-8")
    assert 'ap.add_argument("--min-confidence", type=float, default=1.0' in src
    assert 'on_conflict="alias_text",\n                                               ignore_duplicates=True' in src
    assert "shop cache invalidated" not in src and "invalidate()" not in src
    assert "within 60 s" in src


# ── 7. production, read-only (SKIPs without credentials) ──────────────────────

@test("live (read-only): the twin rule (any MA book other than _36_) selects exactly 165 rows today")
def _():
    if not (os.getenv("SUPABASE_URL") and (os.getenv("SUPABASE_SERVICE_KEY") or os.getenv("SUPABASE_KEY"))):
        print("    SKIP: no Supabase credentials in this checkout")
        return
    from app.db_read import exec_sql
    n = exec_sql("""
        SELECT COUNT(*) AS n FROM selling_prices p
         WHERE p.source_file = 'MASellingPriceBook _36_.xlsx' AND p.price_book = 'MA_base'
           AND EXISTS (SELECT 1 FROM selling_prices t
                        WHERE t.source_file LIKE 'MASellingPriceBook%' AND t.source_file <> 'MASellingPriceBook _36_.xlsx'
                          AND t.price_book = p.price_book
                          AND t.sku_code = p.sku_code
                          AND t.customer_code IS NOT DISTINCT FROM p.customer_code
                          AND t.warehouse_name IS NOT DISTINCT FROM p.warehouse_name
                          AND t.rate_bhd = p.rate_bhd AND t.start_date <> p.start_date
                          AND t.start_date = CASE WHEN EXTRACT(day FROM p.start_date) <= 12
                                THEN make_date(EXTRACT(year FROM p.start_date)::int, EXTRACT(day FROM p.start_date)::int,
                                               EXTRACT(month FROM p.start_date)::int) END)""")[0]["n"]
    assert int(n) == 165, f"twin rule selects {n} rows, the migration asserts 165"


def main() -> int:
    passed = failed = 0
    for name, fn in TESTS:
        try:
            fn()
            print(f"  PASS  {name}")
            passed += 1
        except Exception:  # noqa: BLE001
            print(f"  FAIL  {name}")
            traceback.print_exc(limit=2)
            failed += 1
    print(f"\n{passed} passed, {failed} failed")
    return 0 if failed == 0 else 2


if __name__ == "__main__":
    raise SystemExit(main())
