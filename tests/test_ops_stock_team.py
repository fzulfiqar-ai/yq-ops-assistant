"""Ops engine: his edited Stock & Order Plan read back (scripts.ops.workbook.readback) and the team copy made from
it (scripts.ops.workbook.stock_team), on a synthetic world.

    python -m tests.test_ops_stock_team

Same lightweight runner as tests/test_ops_stock.py, and ITS stand-ins (imported, not copied): the invented stock
ledger, the database stand-in that answers every read from it (World), the temporary data root and the command
(run) that makes the day's plan. SUPABASE_URL, SUPABASE_KEY and DATABASE_URL are pinned to ci.invalid BEFORE any
import; no database, no network, no .env. All names and numbers are invented; every folder is a temporary one.

The plan is the one /yq-stock makes for the stock suite's world (the data date D is Sun 15 Mar 2026), made once
and copied into each test's own data root: one item on the order list, ITEM-B (CB2, VFAN) suggesting 65, the 15
rows to add an item under it; the ageing review holds HG3 (Not moving) and CA1 (Slow, sold most by Van A). A test
types in it with openpyxl as he would in Excel; openpyxl keeps no saved value of a formula, so reading such a file
back also proves the engine works the formulas out. Rows and columns are found as readback finds them: by the
header texts.

Covered:
  * the file as made: its tag, data date, sha and saved time; CB2's row read from the saved values alone;
  * an edited file: his quantity and comment, items added by code (case aside), a code alone is no line, a
    nought a line of nothing; the finals; the sha moves with an edit and not without;
  * each problem in one sentence naming its row: a formula, not whole, below nought, not a number (text, TRUE),
    a quantity with no code, a code not on All items (a wildcard is no wildcard), a code twice, no code, no
    supplier, a Suggested with an error value or no value at all;
  * final_lines per supplier: case and spaces ignored, sheet order, finals above nought with a code;
  * newest_plan: the data date, then the copy number; team copies, lock and temporary files, other names
    passed over;
  * open_in_excel; a file held by another program (read through a copy, else one sentence); PlanFileError for a
    tag unknown or missing, a team copy, no workbook, a folder, a missing file, no item table; v1, v2 and v3 read,
    each as stock_book makes it (its quantities typed in too); the header found by its texts (another row, other
    columns, a note under the table);
  * team_file: "<plan> - Team.xlsx" beside it with his quantities and comments, the push list and the notes;
    no cost, value, BHD, seller or long digit run anywhere (leaks() empty); a second copy " (2)", never over a
    file, his plan untouched, only exports/ written; refused in one sentence with nothing written (a problem,
    money in a comment, a plan outside exports/, no workbook, XlsxWriter missing);
  * leaks() lists every header naming a cost, a value, money or a person, in a book and in a file;
  * the team copy of each layout: Vans and outlets read by its header (blank for a v1 or v2 plan, the plan's units
    for a v3 one, worked out again after an edit), the Notes sentence on what Suggested rests on, no leak and no
    store name;
  * readback's names agree with stock_book's (the file name, the tag's cell, the layouts, "code not found", the
    Vans and outlets header).
"""
from __future__ import annotations

import hashlib
import io
import os
import re
import sys
import traceback
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

os.environ["SUPABASE_URL"] = "https://ci.invalid"
os.environ["SUPABASE_KEY"] = "ci-key"
os.environ["DATABASE_URL"] = "postgresql://ci:ci@ci.invalid:5432/ci"

# The stock suite's stand-ins. Importing it also sets stdout to UTF-8 (so this file does not) and imports
# scripts.ops.db first, as every command does.
from tests.test_ops_stock import (  # noqa: E402
    D, LONG, NAMES, NINE, VAN, VAN_B, World, data_root, model_of, run)
from tests import test_ops_stock_book as tb  # noqa: E402  (its facts of each layout: v1, with_r2, with_v3)
from scripts.ops import config, paths, stock  # noqa: E402
from scripts.ops.workbook import book as bk  # noqa: E402
from scripts.ops.workbook import readback  # noqa: E402
from scripts.ops.workbook import stock_book as sb  # noqa: E402
from scripts.ops.workbook import stock_team  # noqa: E402

stock.RAISE_PARTS = True              # a part that fails fails the test, never stands aside quietly

TESTS: list[tuple[str, object]] = []
FILE = f"YQ Stock and Order Plan {D.isoformat()}.xlsx"
TEAM = FILE.replace(".xlsx", " - Team.xlsx")
YELLOW = "FFFFF2CC"
PEOPLE = ("Van A", "Van B", VAN, VAN_B)
KEYS = {"path", "sha256", "layout", "data_date", "rows", "problems", "open_in_excel", "saved_at"}
ORDER_HEADER = ["Supplier", "Code", "Item", "Category", "Main stock", "Vans and outlets", "On the way", "Sold 90 d",
                "Cover days", "Suggested", "His quantity", "His comment", "Team comment"]
PUSH_HEADER = ["Item", "Code", "Status", "Sellable", "A week", "Push a week", "Action"]
BOOK_TAGS = tuple(getattr(sb, "LAYOUT_TAGS", (sb.LAYOUT_TAG,)))     # the tags stock_book writes (v1, v2, v3)
V1, V2, V3 = "stock-plan v1", "stock-plan v2", "stock-plan v3"
# CB2 in the plan of each layout (layout_bytes): Suggested and Cover days. v1: release 1's rule (85 - 20); v2: the
# rhythm and the season (129.90 - 20); v3: the same with the second van's 35 in the position (129.90 - 55)
CB2_BY_LAYOUT = {V1: (65, 20), V2: (110, 20), V3: (75, 55)}
_PLAN: dict = {}
_LAYOUT_FILES: dict = {}


def test(name):
    def deco(fn):
        TESTS.append((name, fn))
        return fn
    return deco


def raises(kind, fn) -> Exception:
    try:
        fn()
    except kind as exc:
        return exc
    raise AssertionError(f"expected {kind.__name__}")


# ── the plan and his edits ────────────────────────────────────────────────────

def plan_bytes() -> bytes:
    """The day's plan of the stock suite's world, made once by the command itself (/yq-stock)."""
    if "bytes" not in _PLAN:
        with data_root():
            out = run(World())
            _PLAN["bytes"] = Path(out["flags"]["file"]).read_bytes()
    return _PLAN["bytes"]


def layout_bytes(tag: str) -> bytes:
    """A plan of that layout as stock_book makes it for the stock suite's world (each made once): the model with
    release 1's facts (v1), release 2's (v2) or release 2's and the order position (v3), as the workbook suite sets
    them (tests.test_ops_stock_book: v1, with_r2, with_v3)."""
    if tag not in _LAYOUT_FILES:
        with data_root():
            m, _inp = model_of()
            {V1: tb.v1, V2: tb.with_r2, V3: tb.with_v3}[tag](m)
            res = sb.day_file(m)
            assert res["made"] and res["error"] is None and res["controls"] == [], (tag, res)
            _LAYOUT_FILES[tag] = Path(res["path"]).read_bytes()
    return _LAYOUT_FILES[tag]


def folder(root: Path) -> Path:
    return root / "exports" / "reports" / "stock"


def placed(root: Path, name: str = FILE, data: bytes | None = None, where: Path | None = None) -> Path:
    """The plan written into this test's data root (exports/reports/stock, or `where`)."""
    f = where or folder(root)
    f.mkdir(parents=True, exist_ok=True)
    p = f / name
    p.write_bytes(plan_bytes() if data is None else data)
    return p


def opened(path_or_bytes, data_only: bool = False):
    import openpyxl
    src = io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, bytes) else path_or_bytes
    return openpyxl.load_workbook(src, data_only=data_only)


def layout(ws) -> tuple[int, dict]:
    """(header row, {header: column letter}) of the Order sheet's item table, found by its texts."""
    for row in ws.iter_rows(min_row=1, max_row=80):
        texts = {c.value: c.column_letter for c in row if isinstance(c.value, str)}
        if all(h in texts for h in readback.ORDER_HEADS):
            return row[0].row, texts
    raise AssertionError("the Order sheet has no item table")


def head_of(path) -> int:
    return layout(opened(path)["Order"])[0]


def edit(path: Path, *cells, notes: dict | None = None) -> None:
    """Type in his file as he would in Excel, then save it with openpyxl (which keeps no saved value of a formula).
    cells: (k, header, value) on the Order item table, k counted from its first row (1 = CB2, 2 = the first row to
    add an item ...); notes: {cell: value} on the Notes sheet."""
    wb = opened(path)
    ws = wb["Order"]
    head, col = layout(ws)
    for k, header, value in cells:
        ws[f"{col[header]}{head + k}"] = value
    for ref, value in (notes or {}).items():
        wb["Notes"][ref] = value
    wb.save(path)


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def files(root: Path) -> set:
    return {p.relative_to(root) for p in root.rglob("*") if p.is_file()}


def values(path) -> dict:
    """{sheet: {cell: value}} of every cell that holds something."""
    wb = opened(path)
    return {ws.title: {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
            for ws in wb.worksheets}


def row_values(ws, r: int, n: int) -> list:
    return [ws.cell(r, c).value for c in range(1, n + 1)]


# ── read_plan ─────────────────────────────────────────────────────────────────

@test("read_plan, the file as made: its tag, data date, sha and saved time; CB2 suggests 65, read from the saved "
      "values alone")
def _():
    saved = readback.PlanBook._worked_out

    def never(self, sheet, r, c):
        raise AssertionError(f"{sheet} row {r} was worked out: the file carries its saved values")

    with data_root() as root:
        p = placed(root)
        head = head_of(p)
        readback.PlanBook._worked_out = never
        try:
            plan = readback.read_plan(p)
        finally:
            readback.PlanBook._worked_out = saved
    assert set(plan) == KEYS, sorted(plan)
    assert plan["layout"] in BOOK_TAGS and plan["layout"] in readback.LAYOUTS, plan["layout"]
    assert (plan["path"], plan["data_date"], plan["sha256"]) == (str(p), "2026-03-15", sha(plan_bytes()))
    assert plan["rows"] == [{"row": head + 1, "supplier": "VFAN", "code": "CB2", "item": NAMES["B"], "suggested": 65,
                             "your_qty": None, "final": 65, "comment": None, "added": False}], plan["rows"]
    assert plan["problems"] == [] and plan["open_in_excel"] is False
    assert re.fullmatch(r"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\d\+03:00", plan["saved_at"]), plan["saved_at"]
    assert readback.final_lines(plan, "VFAN") == [("CB2", 65)] and readback.final_lines(plan, "WEKOME") == []


@test("read_plan, an edited file: his quantity and comment, items added by code (case aside), a code alone is no "
      "line, a nought a line of nothing; the sha moves with an edit")
def _():
    with data_root() as root:
        p = placed(root)
        head = head_of(p)
        edit(p, (1, "Your qty", 100), (1, "Comment", "Call VFAN first\nthen confirm"),
             (2, "Code", "zn9"), (2, "Your qty", 12), (2, "Comment", "  New model  "),
             (3, "Code", "CA1"), (3, "Your qty", 10.0),
             (4, "Code", "HG3"),                                       # a code alone: a look-up, no line
             (5, "Code", "SN4"), (5, "Your qty", 0))
        ws, col = opened(p, data_only=True)["Order"], layout(opened(p)["Order"])[1]
        assert ws[f"{col['Suggested']}{head + 1}"].value is None       # no saved value: the engine works it out
        plan = readback.read_plan(p)
        again = readback.read_plan(p)
        edited = p.read_bytes()
    got = [(x["row"], x["supplier"], x["code"], x["item"], x["suggested"], x["your_qty"], x["final"], x["comment"],
            x["added"]) for x in plan["rows"]]
    assert got == [(head + 1, "VFAN", "CB2", NAMES["B"], 65, 100, 100, "Call VFAN first\nthen confirm", False),
                   (head + 2, "WEKOME", "zn9", "ZZ-NEW Speaker Upcoming", 0, 12, 12, "New model", True),
                   (head + 3, "VFAN", "CA1", NAMES["A"], 0, 10.0, 10, None, True),
                   (head + 5, "VFAN", "SN4", NAMES["N"], 0, 0, 0, None, True)], got
    assert plan["problems"] == [], plan["problems"]
    assert plan["sha256"] == sha(edited) != sha(plan_bytes()) and again["sha256"] == plan["sha256"]
    assert readback.final_lines(plan, "VFAN") == [("CB2", 100), ("CA1", 10)]
    assert readback.final_lines(plan, " wekome ") == [("zn9", 12)]
    assert readback.final_lines(plan, "Other") == [] and readback.final_lines(plan, "") == []


@test("read_plan names each problem in one sentence, by its row: a formula, not whole, below nought, not a number, "
      "a quantity with no code, a code not on All items, a code twice")
def _():
    with data_root() as root:
        p = placed(root)
        h = head_of(p)
        edit(p, (1, "Your qty", "=50+15"),
             (2, "Code", "CB2"), (2, "Your qty", 5),                   # CB2 again, a quantity on both rows
             (3, "Code", "SN4"), (3, "Your qty", 2.5),
             (4, "Code", "FO5"), (4, "Your qty", -2),
             (5, "Code", "DS1"), (5, "Your qty", "ten"),
             (6, "Your qty", 7),
             (7, "Code", "NOPE-1"), (7, "Your qty", 3),
             (8, "Code", "C*"), (8, "Your qty", 4),                    # no wildcard: CA1 and CB2 stay out
             (9, "Code", "HG3"), (9, "Your qty", True))
        plan = readback.read_plan(p)
    assert plan["problems"] == [
        f"Your qty in row {h + 1} holds a formula: type a number",
        f"Code CB2 is on rows {h + 1} and {h + 2}, each with a quantity: keep it on one row",
        f"Your qty in row {h + 3} is not a whole number (2.5): type a whole number of units",
        f"Your qty in row {h + 4} is below nought (-2): type nought or more",
        f"Your qty in row {h + 5} is not a number ('ten'): type a whole number of units, or clear the cell",
        f"Row {h + 6} has a quantity but no code: type the item's code in its Code cell, or clear the quantity",
        f"Row {h + 7}: the code NOPE-1 is not on the All items sheet: check the code",
        f"Row {h + 8}: the code C* is not on the All items sheet: check the code",
        f"Your qty in row {h + 9} is not a number (TRUE): type a whole number of units, or clear the cell",
    ], plan["problems"]
    got = [(x["row"], x["supplier"], x["your_qty"], x["final"]) for x in plan["rows"]]
    assert got == [(h + 1, "VFAN", "=50+15", 65), (h + 2, "VFAN", 5, 5), (h + 3, "VFAN", 2.5, 0),
                   (h + 4, "VFAN", -2, 0), (h + 5, "VFAN", "ten", 0), (h + 7, "", 3, 3), (h + 8, "", 4, 4),
                   (h + 9, "VFAN", True, 0)], got                      # a Your qty refused: Final is Suggested
    assert [x["item"] for x in plan["rows"]][5:7] == [readback.NOT_FOUND, readback.NOT_FOUND]
    assert readback.final_lines(plan, "VFAN") == [("CB2", 65), ("CB2", 5)]   # the caller stops on the problems


@test("read_plan names a row to order with no code or no supplier, and a Suggested with an error value or none")
def _():
    h = None
    out = {}
    with data_root() as root:
        p = placed(root)
        h = head_of(p)
        for key, cells in (("no code", [(1, "Code", None)]),
                           ("no supplier", [(1, "Supplier", None), (1, "Your qty", 30)]),
                           ("an error value", [(1, "Suggested", "=1/0")]),
                           ("no value", [(1, "Suggested", "=VLOOKUP(1,A1:A2,1,0)")])):   # a function it does not know
            p.write_bytes(plan_bytes())
            edit(p, *cells)
            out[key] = readback.read_plan(p)
    assert out["no code"]["problems"] == [f"Row {h + 1} ({NAMES['B']}) has no code, so it cannot be logged: type its "
                                          "code in the Code cell, or 0 in Your qty"], out["no code"]["problems"]
    assert out["no code"]["rows"][0]["final"] == 65 and readback.final_lines(out["no code"], "VFAN") == []
    assert out["no supplier"]["problems"] == [f"Row {h + 1} (CB2) names no supplier: type it in the Supplier cell"]
    assert out["an error value"]["problems"] == [f"Suggested in row {h + 1} is '#DIV/0!', not a whole number: type "
                                                 "your quantity in Your qty"], out["an error value"]["problems"]
    assert out["an error value"]["rows"][0]["suggested"] == "#DIV/0!"
    assert out["an error value"]["rows"][0]["final"] is None
    assert out["no value"]["problems"] == [f"Suggested in row {h + 1} has no value: open the file in Excel, save it "
                                           "and try again"], out["no value"]["problems"]
    assert out["no value"]["rows"][0]["suggested"] is None and out["no value"]["rows"][0]["final"] is None


@test("final_lines: one supplier's rows with a code and a final above nought, in sheet order, case and spaces aside")
def _():
    plan = {"rows": [{"row": 12, "supplier": "VFAN", "code": "CB2", "final": 65},
                     {"row": 13, "supplier": " vfan ", "code": " ca1 ", "final": 10},
                     {"row": 14, "supplier": "VFAN", "code": "HG3", "final": 0},
                     {"row": 15, "supplier": "VFAN", "code": "", "final": 5},
                     {"row": 16, "supplier": "VFAN", "code": "SN4", "final": None},
                     {"row": 17, "supplier": "WEKOME", "code": "ZN9", "final": 12},
                     {"row": 18, "supplier": "VFAN  Co", "code": "X05  UC", "final": 3},
                     {"row": 19, "supplier": "VFAN", "code": "FO5", "final": True}]}
    assert readback.final_lines(plan, "VFAN") == [("CB2", 65), ("ca1", 10)]
    assert readback.final_lines(plan, "vfan co") == [("X05 UC", 3)] and readback.final_lines(plan, "wekome") == [
        ("ZN9", 12)]
    assert readback.final_lines(plan, "   ") == [] and readback.final_lines({"rows": []}, "VFAN") == []


# ── which file ────────────────────────────────────────────────────────────────

@test("newest_plan: the latest data date, then the copy number; team copies, lock and temporary files and other "
      "names passed over")
def _():
    with data_root() as root:
        f = folder(root)
        f.mkdir(parents=True)
        for name in ("YQ Stock and Order Plan 2026-03-14 (12).xlsx", "YQ Stock and Order Plan 2026-03-15.xlsx",
                     "YQ Stock and Order Plan 2026-03-15 (2).xlsx", "YQ Stock and Order Plan 2026-03-15 (9).xlsx",
                     "YQ Stock and Order Plan 2026-03-15 (10).xlsx",
                     "YQ Stock and Order Plan 2026-03-16 - Team.xlsx",
                     "yq stock and order plan 2026-03-16 - team (2).xlsx",
                     "YQ Stock and Order Plan 2026-03-15 (10) - Team.xlsx", "~$YQ Stock and Order Plan 2026-03-17.xlsx",
                     "~build-1x2y.xlsx", "YQ Stock and Order Plan 2026-13-40.xlsx",
                     "YQ Stock and Order Plan 2026-03-18.csv", "notes.txt"):
            (f / name).write_bytes(b"")
        (f / "YQ Stock and Order Plan 2026-03-19.xlsx").mkdir()                 # a folder of that name
        assert readback.newest_plan(f) == f / "YQ Stock and Order Plan 2026-03-15 (10).xlsx"
        for name in ("YQ Stock and Order Plan 2026-03-15 (10).xlsx", "YQ Stock and Order Plan 2026-03-15 (9).xlsx"):
            (f / name).unlink()
        assert readback.newest_plan(f) == f / "YQ Stock and Order Plan 2026-03-15 (2).xlsx"
        (f / "YQ Stock and Order Plan 2026-03-15 (2).xlsx").unlink()
        (f / "YQ Stock and Order Plan 2026-03-15.xlsx").unlink()
        assert readback.newest_plan(f) == f / "YQ Stock and Order Plan 2026-03-14 (12).xlsx"
        (f / "YQ Stock and Order Plan 2026-03-14 (12).xlsx").unlink()
        assert readback.newest_plan(f) is None and readback.newest_plan(root / "nowhere") is None


@test("open_in_excel: Excel's ~$ file beside it; the plan still read (its last save) and still the newest")
def _():
    with data_root() as root:
        p = placed(root)
        assert readback.read_plan(p)["open_in_excel"] is False
        (p.parent / ("~$" + FILE)).write_bytes(b"owner")
        plan = readback.read_plan(p)
        assert plan["open_in_excel"] is True and plan["rows"][0]["final"] == 65
        assert readback.newest_plan(p.parent) == p


@test("a file another program holds: read through a copy (removed after); when the copy fails too, one sentence")
def _():
    real_read, real_copy = readback._read_bytes, readback._copy_aside

    def held(path):
        if Path(path).name == FILE:
            raise PermissionError(13, "Permission denied")
        return real_read(path)

    def no_copy(path, tmp):
        raise PermissionError(13, "Permission denied")

    with data_root() as root:
        p = placed(root)
        readback._read_bytes = held
        try:
            plan = readback.read_plan(p)
            left = [x for x in files(root) if x.name.startswith("~")]
            readback._copy_aside = no_copy
            exc = raises(readback.PlanFileError, lambda: readback.read_plan(p))
        finally:
            readback._read_bytes, readback._copy_aside = real_read, real_copy
    assert plan["sha256"] == sha(plan_bytes()) and plan["rows"][0]["code"] == "CB2" and plan["path"] == str(p)
    assert left == [], left
    assert str(exc) == f"{FILE} is held by another program: save and close it in Excel, then try again.", str(exc)


@test("PlanFileError, one sentence each: a tag unknown or missing, a team copy, no workbook, a folder, a missing "
      "file, no item table; v1, v2 and v3 are read")
def _():
    words = {}
    with data_root() as root:
        p = placed(root)
        for tag in ("stock-plan v1", "stock-plan v2", "stock-plan v3"):
            p.write_bytes(plan_bytes())
            edit(p, notes={"B5": tag})
            plan = readback.read_plan(p)
            assert plan["layout"] == tag and plan["rows"][0]["final"] == 65, tag
        for key, notes in (("unknown", {"B5": "stock-plan v9"}), ("missing", {"B5": None}),
                           ("team", {"B5": stock_team.TEAM_TAG})):
            p.write_bytes(plan_bytes())
            edit(p, notes=notes)
            words[key] = str(raises(readback.PlanFileError, lambda: readback.read_plan(p)))
        p.write_bytes(plan_bytes())
        wb = opened(p)
        head, col = layout(wb["Order"])
        wb["Order"][f"{col['Your qty']}{head}"] = "Qty"
        wb.save(p)
        words["no table"] = str(raises(readback.PlanFileError, lambda: readback.read_plan(p)))
        p.write_bytes(b"not a workbook")
        words["no workbook"] = str(raises(readback.PlanFileError, lambda: readback.read_plan(p)))
        words["folder"] = str(raises(readback.PlanFileError, lambda: readback.read_plan(p.parent)))
        words["missing file"] = str(raises(readback.PlanFileError, lambda: readback.read_plan(p.parent / "gone.xlsx")))
    assert issubclass(readback.PlanFileError, ValueError)
    assert words == {
        "unknown": f"{FILE} is not a Stock & Order Plan this engine reads: its layout tag (Notes!B5) is "
                   "'stock-plan v9', not stock-plan v1, stock-plan v2 or stock-plan v3.",
        "missing": f"{FILE} is not a Stock & Order Plan this engine reads: its layout tag (Notes!B5) is missing, not "
                   "stock-plan v1, stock-plan v2 or stock-plan v3.",
        "team": f"{FILE} is a team copy (no costs, his quantities as values): read your own plan file instead.",
        "no table": f"{FILE} has no order list this engine reads: the Order sheet's item table (its header Supplier, "
                    "Code, Item, Suggested, Your qty, Final, Comment) is not found.",
        "no workbook": f"{FILE} cannot be read as an Excel workbook (BadZipFile).",
        "folder": "stock is a folder, not a Stock & Order Plan workbook.",
        "missing file": "gone.xlsx is not there: it was moved or deleted.",
    }, words


@test("the header is found by its texts: another row, other columns, more columns; a note under the table passed over")
def _():
    import openpyxl
    with data_root() as root:
        p = folder(root) / "YQ Stock and Order Plan 2026-03-20 (3).xlsx"
        p.parent.mkdir(parents=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Order"
        ws["A1"] = "A plan laid out another way"
        heads = ["Comment", "Item", "Note", "Code", "Supplier", "Final", "Your qty", "Suggested"]
        for k, text in enumerate(heads):
            ws.cell(row=7, column=2 + k, value=text)
        for i, row in enumerate((("ring first", "Widget One", None, "W1", "Acme", 9, 9, 4),
                                 (None, "Widget Two", "x", "W2", "acme ", 4, None, 4))):
            for k, v in enumerate(row):
                ws.cell(row=8 + i, column=2 + k, value=v)
        ws.cell(row=12, column=2, value="A note under the table, in its first column")
        notes = wb.create_sheet("Notes")
        notes["A5"], notes["B5"] = "Layout of this file", "stock-plan v2"
        wb.save(p)
        plan = readback.read_plan(p)
    got = [(x["row"], x["supplier"], x["code"], x["item"], x["suggested"], x["your_qty"], x["final"], x["comment"],
            x["added"]) for x in plan["rows"]]
    assert got == [(8, "Acme", "W1", "Widget One", 4, 9, 9, "ring first", False),
                   (9, "acme", "W2", "Widget Two", 4, None, 4, None, False)], got
    assert plan["problems"] == [] and plan["data_date"] == "2026-03-20" and plan["layout"] == "stock-plan v2"
    assert readback.final_lines(plan, "ACME") == [("W1", 9), ("W2", 4)]


# ── the team copy ─────────────────────────────────────────────────────────────

def team_edits(p: Path) -> None:
    edit(p, (1, "Your qty", 100), (1, "Comment", "Call VFAN first"),
         (2, "Code", "zn9"), (2, "Your qty", 12), (2, "Comment", "New model, sample first"),
         (3, "Code", "CA1"), (3, "Your qty", 10))


@test("team_file: '<plan> - Team.xlsx' beside it, his quantities and comments, the push list, the notes; his plan "
      "untouched, only exports/ written")
def _():
    his = opened(plan_bytes(), data_only=True)
    head, col = layout(opened(plan_bytes())["Order"])
    cb2 = [his["Order"][f"{col[h]}{head + 1}"].value if h in col else None      # a v2 plan has no Vans and outlets
           for h in ("Category", "Main stock", "Vans and outlets", "On the way", "Sold 90 d", "Cover days")]
    ag = his["Ageing"]
    ag_head = next(r for r in range(1, 40) if ag[f"A{r}"].value == "Item" and ag[f"C{r}"].value == "Status")
    ag_col = {ag.cell(ag_head, c).value: c for c in range(1, ag.max_column + 1) if ag.cell(ag_head, c).value}
    review = []
    r = ag_head + 1
    while ag.cell(r, ag_col["Item"]).value:
        review.append([ag.cell(r, ag_col[h]).value for h in PUSH_HEADER])
        r += 1
    with data_root() as root:
        p = placed(root)
        team_edits(p)
        mine = p.read_bytes()
        before = files(root)
        res = stock_team.team_file(p)
        after = files(root)
        path = folder(root) / TEAM
        assert res == {"path": str(path), "made": True, "error": None}, res
        assert after - before == {Path("exports", "reports", "stock", TEAM)} and p.read_bytes() == mine
        wb = opened(path)
        found = stock_team.leaks(path)
        cells = values(path)
        newest = readback.newest_plan(p.parent)
        team_read = str(raises(readback.PlanFileError, lambda: readback.read_plan(path)))
    assert wb.sheetnames == ["Order", "Push list", "Notes"], wb.sheetnames
    ws = wb["Order"]
    assert ws["A1"].value == stock_team.TITLE and ws["A2"].value.startswith(f"Data date Sun 15 Mar 2026 · made from "
                                                                            f"{FILE}, saved "), ws["A2"].value
    assert ws["A4"].value == "His quantities by supplier: VFAN 110 units (2 items); WEKOME 12 units (1 item).", \
        ws["A4"].value
    assert row_values(ws, 5, 14) == ORDER_HEADER + [None]
    assert row_values(ws, 6, 13) == ["VFAN", "CB2", NAMES["B"], *cb2, 65, 100, "Call VFAN first", None], \
        row_values(ws, 6, 13)
    assert cb2[:2] == ["Chargers", 20]
    assert row_values(ws, 7, 13) == ["WEKOME", "zn9", "ZZ-NEW Speaker Upcoming", "Speakers", 0, None, None, None, None,
                                     0, 12, "New model, sample first", None], row_values(ws, 7, 13)
    assert row_values(ws, 8, 13) == ["VFAN", "CA1", NAMES["A"], "Cables", 394, None, None, None, None, 0, 10, None,
                                     None]
    assert ws["A9"].value is None
    assert all(ws[f"M{r}"].fill.fgColor.rgb == YELLOW for r in (6, 7, 8)) and ws["K6"].fill.fgColor.rgb != YELLOW
    assert ws.freeze_panes == "A6" and ws.auto_filter.ref == "A5:M8", (ws.freeze_panes, ws.auto_filter.ref)
    push = wb["Push list"]
    assert row_values(push, 5, 8) == PUSH_HEADER + [None]
    assert [row_values(push, 6 + k, 7) for k in range(len(review))] == review and push.cell(6 + len(review), 1).value \
        is None
    assert [x[1:3] for x in review] == [["HG3", "Not moving"], ["CA1", "Slow"]], review
    rules = {r.formula[0] for cf in push.conditional_formatting for r in cf.rules}
    assert rules == {'"Not moving"', '"Slow"', '"On track"', '"Cleared"'}, rules
    assert push.freeze_panes == "A6" and push.auto_filter.ref == f"A5:G{5 + len(review)}"
    notes = wb["Notes"]
    assert notes["B5"].value == stock_team.TEAM_TAG == "stock-plan team v1"
    assert notes["A5"].value == "Layout of this file"
    words = " ".join(str(v) for v in cells["Notes"].values())
    assert "Costs and values are left out; reply with your comments in the Team comment column." in words
    assert "Sun 15 Mar 2026 (the stock ledger's last day)" in words and FILE in words
    assert newest == p and team_read.endswith("is a team copy (no costs, his quantities as values): read your own "
                                              "plan file instead.")
    assert found == [], found


@test("the team copy shows no cost, value, BHD, seller or long digit run in any cell; every number of its Order "
      "sheet is whole")
def _():
    with data_root() as root:
        p = placed(root)
        team_edits(p)
        res = stock_team.team_file(p)
        cells = values(res["path"])
        his = values(p)
    for sheet, at in cells.items():
        for ref, v in at.items():
            text = str(v)
            assert "BHD" not in text, (sheet, ref, text)
            assert not any(who in text for who in PEOPLE), (sheet, ref, text)
            assert not NINE.search(text) and not LONG.search(text), (sheet, ref, text)
    numbers = [v for v in cells["Order"].values() if isinstance(v, (int, float)) and not isinstance(v, bool)]
    assert numbers and all(float(v).is_integer() for v in numbers), numbers      # no landed cost (2.5, 0.4), no value
    heads = [v for v in list(cells["Order"].values()) + list(cells["Push list"].values()) if isinstance(v, str)]
    for word in ("Landed cost", "Value, BHD", "Budget, BHD", "Minimum order, BHD", "Sold most by", "With vans"):
        assert word not in heads, word
    assert any("Sold most by" == v for v in his["Ageing"].values())            # his own file names the sellers


@test("team_file never writes over a file: a second copy is ' (2)', a name taken is passed over")
def _():
    with data_root() as root:
        p = placed(root)
        team_edits(p)
        one = stock_team.team_file(p)
        first = Path(one["path"]).read_bytes()
        two = stock_team.team_file(p)
        theirs = folder(root) / FILE.replace(".xlsx", " - Team (3).xlsx")
        theirs.write_bytes(b"someone else's file")
        four = stock_team.team_file(p)
        names = sorted(x.name for x in folder(root).iterdir())
        assert Path(one["path"]).read_bytes() == first and theirs.read_bytes() == b"someone else's file"
    assert Path(one["path"]).name == TEAM and Path(two["path"]).name == FILE.replace(".xlsx", " - Team (2).xlsx")
    assert Path(four["path"]).name == FILE.replace(".xlsx", " - Team (4).xlsx") and two["made"] and four["made"]
    assert names == sorted([FILE, TEAM] + [FILE.replace(".xlsx", f" - Team ({k}).xlsx") for k in (2, 3, 4)]), names


@test("team_file refuses in one sentence, nothing written: a problem, money in a comment, outside exports/, no "
      "workbook, XlsxWriter missing")
def _():
    out = {}
    with data_root() as root:
        p = placed(root)
        h = head_of(p)
        for key, cells in (("a problem", [(1, "Your qty", "=50+15")]),
                           ("two problems", [(1, "Your qty", 2.5), (2, "Your qty", 7)]),
                           ("BHD", [(1, "Comment", "VFAN quoted BHD 2.400")]),
                           ("fils", [(1, "Comment", "quoted 2.400 a unit")])):
            p.write_bytes(plan_bytes())
            edit(p, *cells)
            out[key] = stock_team.team_file(p)
        p.write_bytes(b"not a workbook")
        out["no workbook"] = stock_team.team_file(p)
        p.write_bytes(plan_bytes())
        desk = placed(root, where=root / "Desktop")
        out["outside"] = stock_team.team_file(desk)
        had, saved = "xlsxwriter" in sys.modules, sys.modules.get("xlsxwriter")
        sys.modules["xlsxwriter"] = None                 # `import xlsxwriter` now fails, as on a PC without it
        try:
            out["no XlsxWriter"] = stock_team.team_file(p)
        finally:
            if had:
                sys.modules["xlsxwriter"] = saved
            else:
                del sys.modules["xlsxwriter"]
        left = sorted(x.name for x in folder(root).iterdir()) + sorted(x.name for x in (root / "Desktop").iterdir())
    for key, res in out.items():
        assert res["path"] is None and res["made"] is False and not res["error"].endswith("."), (key, res)
    assert out["a problem"]["error"] == f"{FILE} has 1 thing to fix first: Your qty in row {h + 1} holds a formula: " \
                                        "type a number", out["a problem"]
    assert out["two problems"]["error"] == (f"{FILE} has 2 things to fix first: Your qty in row {h + 1} is not a whole "
                                            "number (2.5): type a whole number of units (and 1 more)")
    assert out["BHD"]["error"] == (f"your comment in row {h + 1} names money (BHD): the team copy carries none; "
                                   "reword it in your file, then try again"), out["BHD"]
    assert out["fils"]["error"].startswith(f"your comment in row {h + 1} names money (2.400)"), out["fils"]
    assert out["no workbook"]["error"] == f"{FILE} cannot be read as an Excel workbook (BadZipFile)"
    assert out["outside"]["error"] == (f"the team copy goes beside the plan, and {root / 'Desktop'} is outside "
                                       "exports/ and business_data/, the only places scripts/ops writes")
    assert out["no XlsxWriter"]["error"] == "XlsxWriter is not installed on this PC (pip install xlsxwriter)"
    assert left == [FILE, FILE], left                     # no team copy, no temporary file, here or there


@test("leaks() lists every header naming a cost, a value, money or a person: in a book and in a file")
def _():
    book = bk.Book()
    ws = book.sheet("Sheet")
    for c, text in enumerate(("Item", "Landed cost", "Value, BHD", "Margin %", "Budget", "Unit price", "Salesman",
                              "Sold most by", "Minimum order", "Code")):
        ws.text(4, c, text, "g_head_r" if c % 2 else "g_head")
    ws.text(5, 0, "cost and value, in a cell that is no header", "g_text")
    assert stock_team.leaks(book) == ["Sheet!B5 'Landed cost'", "Sheet!C5 'Value, BHD'", "Sheet!D5 'Margin %'",
                                      "Sheet!E5 'Budget'", "Sheet!F5 'Unit price'", "Sheet!G5 'Salesman'",
                                      "Sheet!H5 'Sold most by'", "Sheet!I5 'Minimum order'"], stock_team.leaks(book)
    with data_root() as root:
        p = placed(root)
        target = folder(root) / "leaky.xlsx"
        book.evaluate()
        book.save(target)
        in_file = stock_team.leaks(target)
        in_plan = stock_team.leaks(p)
        head, col = layout(opened(p)["Order"])
        team = stock_team.build_team(readback.load(p), readback.read_plan(p))
    assert in_file == stock_team.leaks(book), in_file
    assert f"Order!{col['Landed cost']}{head} 'Landed cost'" in in_plan, in_plan
    assert any(x.startswith("Ageing!") and x.endswith("'Sold most by'") for x in in_plan), in_plan
    assert stock_team.leaks(team) == [] and [ws.name for ws in team.sheets] == list(stock_team.SHEETS)


@test("read_plan reads every layout stock_book makes, v1, v2 and v3: the tag, CB2's row and its final (v3 counts the "
      "second van's 35: 75), then his quantities typed in")
def _():
    with data_root() as root:
        for tag, (qty, _cover) in CB2_BY_LAYOUT.items():
            p = placed(root, data=layout_bytes(tag))
            plan = readback.read_plan(p)
            assert plan["layout"] == tag and plan["problems"] == [], (tag, plan["problems"])
            assert [(x["code"], x["supplier"], x["suggested"], x["final"], x["added"]) for x in plan["rows"]] == [
                ("CB2", "VFAN", qty, qty, False)], (tag, plan["rows"])
            assert readback.final_lines(plan, "VFAN") == [("CB2", qty)], tag
            edit(p, (1, "Your qty", 30), (2, "Code", "CA1"), (2, "Your qty", 10))
            plan = readback.read_plan(p)
            assert plan["problems"] == [] and readback.final_lines(plan, "VFAN") == [("CB2", 30), ("CA1", 10)], tag
            assert plan["rows"][0]["suggested"] == qty, (tag, plan["rows"][0])     # worked out again, unsaved


@test("the team copy of each layout: Vans and outlets by its header (blank for a v1 or v2 plan, the plan's 35 for a "
      "v3 one, also worked out again after an edit); the Notes sentence names the vans and the outlets for v3 only; "
      "no leak, BHD or store name")
def _():
    said = {V1: "the main store.", V2: "the main store.",
            V3: "the main store and the stores counted beside it (its Vans and outlets column)."}
    with data_root() as root:
        for tag, (qty, cover) in CB2_BY_LAYOUT.items():
            for edited in (False, True):
                p = placed(root, data=layout_bytes(tag))
                if edited:                                     # saved by openpyxl: no formula keeps a value
                    edit(p, (1, "Your qty", 100))
                res = stock_team.team_file(p)
                assert res["made"] and res["error"] is None, (tag, res)
                path = Path(res["path"])
                ws = opened(path)["Order"]
                cells = values(path)
                found = stock_team.leaks(path)
                assert row_values(ws, 5, 13) == ORDER_HEADER, (tag, row_values(ws, 5, 13))
                row = dict(zip(ORDER_HEADER, row_values(ws, 6, 13)))
                got = (row["Code"], row["Main stock"], row["Suggested"], row["Cover days"])
                assert got == ("CB2", 20, qty, cover), (tag, edited, row)
                assert row["Vans and outlets"] == (35 if tag == V3 else None), (tag, edited, row)
                assert row["His quantity"] == (100 if edited else qty), (tag, edited, row)
                sentence = next(v for v in cells["Notes"].values() if str(v).startswith("What he will order"))
                assert sentence.endswith("Suggested is worked out from each item's sales a day and its stock in "
                                         + said[tag]), (tag, sentence)
                assert found == [], (tag, found)
                for sheet, at in cells.items():
                    for ref, v in at.items():
                        text = str(v)
                        assert "BHD" not in text and not any(who in text for who in PEOPLE), (tag, sheet, ref, text)
                path.unlink()


@test("readback's names agree with stock_book's: the file name, the tag's cell, the layouts, 'code not found', the "
      "Vans and outlets header and the layout that counts it")
def _():
    assert sb.STEM.format(d=D.isoformat()) + ".xlsx" == FILE and FILE.startswith(readback.STEM + " ")
    assert (sb.TAG_SHEET, sb.TAG_CELL) == (readback.TAG_SHEET, readback.TAG_CELL)
    assert set(BOOK_TAGS) <= set(readback.LAYOUTS) and stock_team.TEAM_TAG not in readback.LAYOUTS
    assert readback.LAYOUTS == sb.LAYOUT_TAGS == (V1, V2, V3)
    assert (stock_team.VANS_HEAD, stock_team.POSITION_LAYOUT) == (sb.VANS_HEAD, sb.LAYOUT_TAG_V3)
    assert stock_team.TEAM_TAG.startswith(readback.TEAM_TAG_START) and stock_team.team_name(FILE) == TEAM[:-5]
    ws = opened(plan_bytes())["Order"]
    head, col = layout(ws)
    item = ws[f"{col['Item']}{head + 2}"].value                       # the first row to add an item
    assert isinstance(item, str) and f'"{readback.NOT_FOUND}"' in item, item


def main() -> int:
    passed = failed = 0
    for name, fn in TESTS:
        try:
            fn()
            print(f"PASS  {name}")
            passed += 1
        except Exception:  # noqa: BLE001
            print(f"FAIL  {name}")
            traceback.print_exc()
            failed += 1
    paths.set_data_root(None)
    config.use(None)
    print(f"\n{passed} passed, {failed} failed")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
